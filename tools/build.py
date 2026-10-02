#!/usr/bin/env python3
"""Build and verify course lessons.

    python tools/build.py 1.4              # build + verify one lesson
    python tools/build.py 1.4 2.1          # several lessons
    python tools/build.py all              # every lesson that has a builder module
    python tools/build.py 1.4 --no-verify  # skip the LibreOffice verification
    python tools/build.py readme           # refresh course + module READMEs (syllabus, lesson lists)

Each lesson builder lives in tools/lessons/lesson_<major>_<minor>.py and defines build() -> Lesson.
Verification needs LibreOffice Calc with the Python UNO bridge (see tools/README.md).
"""
from __future__ import annotations

import argparse
import importlib
import os
import re
import sys
import tempfile
import time
from datetime import date, datetime, timedelta
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import curriculum as CUR  # noqa: E402
from xlcourse.lesson import CORRECT, ROOT, Lesson, Task, md_escape_dollars  # noqa: E402
from xlcourse.data import excel_serial  # noqa: E402

MARK = "<!-- BEGIN GENERATED: {0} -->"
END = "<!-- END GENERATED: {0} -->"


# ---------------------------------------------------------------------------
# README handling
# ---------------------------------------------------------------------------
def nav_block(code: str) -> str:
    info = CUR.by_code(code)
    prev, nxt = CUR.neighbors(code)
    here = ROOT / info.path

    def link(other):
        rel = os.path.relpath(ROOT / other.path / "README.md", here)
        return f"[{other.code} {other.title}]({rel.replace(os.sep, '/')})"
    parts = []
    if prev:
        parts.append(f"⬅️ **Previous:** {link(prev)}")
    parts.append("🏠 [Course home](../../README.md)")
    if nxt:
        parts.append(f"**Next:** {link(nxt)} ➡️")
    return "---\n\n" + " · ".join(parts)


def skeleton(lesson: Lesson) -> str:
    wb = lesson.workbook_name
    return f"""# Lesson {lesson.code} · {lesson.title}

> **Level:** {lesson.level} · **Time:** about {lesson.minutes} minutes · **Workbook:** [`{wb}`]({wb.replace(' ', '%20')})
> **Data:** {lesson.data_note or 'Bluestone Health System (synthetic)'}

TODO: one-paragraph hook that explains why this skill matters in a hospital setting.

## What you'll learn

""" + "\n".join(f"- {o}" for o in lesson.objectives) + f"""

## 📖 Guide

TODO: the teaching content.

## 🧪 Hands-on practice

Download [`{wb}`]({wb.replace(' ', '%20')}) and open the **Practice** sheet. Type each answer in the yellow cell — the **Check**
column turns green when you're right.

{MARK.format('practice')}
{END.format('practice')}

## ✅ Answer key

The workbook has a hidden **Answer Key** sheet (right-click any sheet tab → **Unhide…**). The same answers are below,
collapsed so you don't see them by accident.

{MARK.format('answers')}
{END.format('answers')}

## 🏆 Bonus challenge

{MARK.format('bonus')}
{END.format('bonus')}

{MARK.format('bonus-answers')}
{END.format('bonus-answers')}

## Key takeaways

- TODO

{MARK.format('nav')}
{END.format('nav')}
"""


def inject(text: str, blocks: dict[str, str], strict: bool = True) -> str:
    for name, content in blocks.items():
        start, end = MARK.format(name), END.format(name)
        pat = re.compile(re.escape(start) + r".*?" + re.escape(end), re.S)
        if not pat.search(text):
            if strict:
                raise SystemExit(f"README is missing the block markers for '{name}':\n  {start}\n  {end}")
            continue
        text = pat.sub(lambda _m: f"{start}\n{content}\n{end}" if content else f"{start}\n{end}", text)
    return text


# ---------------------------------------------------------------------------
# Verification
# ---------------------------------------------------------------------------
def expected_number(t: Task):
    v = t.key_value()
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (datetime, date)):
        return excel_serial(v)
    if isinstance(v, (int, float)):
        return float(v)
    return None


def matches(t: Task, info: dict) -> bool:
    kind = t.kind()
    if kind == "custom" and isinstance(t.answer, str):
        kind = "text"
    if kind == "text":
        if info["kind"] != "text":
            return False
        got = info["text"].strip().lower()
        opts = [str(t.answer).strip().lower()] + [str(a).strip().lower() for a in (t.accept or [])]
        return got in opts
    exp = expected_number(t)
    if exp is None or info["kind"] != "value":
        return False
    tol = max(t.tolerance(), 1e-9 * abs(exp), 1e-9)
    if kind == "percent":
        return abs(info["value"] - exp) <= tol or abs(info["value"] - exp * 100) <= tol * 100
    if abs(info["value"] - exp) <= tol:
        return True
    return any(isinstance(a, (int, float)) and abs(info["value"] - a) <= tol for a in (t.accept or []))


def verify(lesson: Lesson, path: Path) -> list[str]:
    from openpyxl import load_workbook
    from xlcourse import lo

    problems, warnings = [], []
    wb = load_workbook(path)
    for name in (lesson.key_sheet, lesson.bonus_key_sheet):
        if name in wb.sheetnames and wb[name].sheet_state != "hidden":
            problems.append(f"sheet '{name}' is not hidden")
    tmpdir = Path(tempfile.mkdtemp(prefix="xlc_selftest_"))
    st_path = lesson.save(tmpdir / ("selftest-" + lesson.workbook_name), selftest=True)
    # rebuild normal workbook state (save(selftest) mutated task cell info identically, so addresses still valid)
    lesson.build_workbook(selftest=False)

    all_tasks = [(lesson.practice_sheet, lesson.key_sheet, t) for t in lesson.tasks] + \
                [(lesson.bonus_sheet, lesson.bonus_key_sheet, t) for t in lesson.bonus]
    cells = []
    for psheet, ksheet, t in all_tasks:
        cells.append((psheet, t.check_cell))
        cells.append((psheet, t.answer_cell))
        if t.live_cell:
            cells.append((ksheet, t.live_cell))
    scan = [(lesson.practice_sheet, "A1:E200"), (lesson.key_sheet, "A1:F200")] + list(lesson.verify_scan)
    if lesson.bonus:
        scan += [(lesson.bonus_sheet, "A1:E200"), (lesson.bonus_key_sheet, "A1:F200")]
    with lo.office() as desk:
        pristine = lo.read_cells(desk, path, cells, scan)
        selft = lo.read_cells(desk, st_path, cells, [])

    def fmt(info):
        return info["text"] if info["kind"] != "value" else f"{info['value']:.6g}"

    rows = []
    for psheet, ksheet, t in all_tasks:
        kind = t.kind()
        live_res = "—"
        if t.live_cell:
            info = pristine["cells"][(ksheet, t.live_cell)]
            if kind == "manual":
                live_res = "n/a"
            elif matches(t, info):
                live_res = "OK"
            else:
                live_res = "FAIL"
                problems.append(f"[{t.number}] live formula in key gives {fmt(info)!r}, expected {t.answer!r}")
        chk = pristine["cells"][(psheet, t.check_cell)]
        if chk["kind"] not in ("empty",) and kind != "manual":
            if t.summary:
                if chk["text"] == CORRECT:
                    problems.append(f"[{t.number}] summary task already shows ✔ before the learner does anything")
                elif chk["text"]:
                    warnings.append(f"[{t.number}] summary check shows {chk['text']!r} before any work (prefer a summary that returns \"\" until work starts)")
            else:
                problems.append(f"[{t.number}] check cell is not blank in the pristine workbook: {fmt(chk)!r}")
        st = "—"
        summary_opt_in = t.self_test == "summary"
        if t.summary and not t.fill and not summary_opt_in and kind != "manual":
            warnings.append(f"[{t.number}] summary task has no fill, so the self-test can't check it (add fill=, or self_test='summary' "
                            "if a customize hook simulates the work)")
        if t.self_test and kind != "manual" and (t.is_formula or t.fill or t.answer is not None or summary_opt_in) \
                and not (t.summary and not t.fill and not summary_opt_in):
            sinfo = selft["cells"][(psheet, t.check_cell)]
            if sinfo["text"] == CORRECT:
                st = "OK"
            else:
                st = "FAIL"
                ans = selft["cells"][(psheet, t.answer_cell)]
                problems.append(f"[{t.number}] self-test: check shows {fmt(sinfo)!r}; answer cell evaluated to {fmt(ans)!r}, expected {t.answer!r}")
        rows.append((t.number, kind, live_res, st, (t.title or t.prompt)[:60]))
    for sheet, addr, txt in pristine["errors"]:
        if sheet in (lesson.key_sheet, lesson.bonus_key_sheet):
            problems.append(f"error value {txt} in {sheet}!{addr}")
        else:
            warnings.append(f"error value {txt} in {sheet}!{addr} (pristine workbook)")
    print(f"\n  {'#':>4}  {'check':8} {'key-live':8} {'self-test':9} task")
    for n, k, lr, st, title in rows:
        print(f"  {n:>4}  {k:8} {lr:8} {st:9} {title}")
    for w in warnings:
        print("  warning:", w)
    for p in problems:
        print("  PROBLEM:", p)
    return problems


def lint_readme(path: Path, text: str) -> list[str]:
    """Cheap checks for things that render badly or break on GitHub."""
    out = []
    if re.search(r"^\s*(?:[-*]\s*)?TODO:", text, re.M):
        out.append("README still contains a TODO: placeholder")
    in_fence = False
    for n, line in enumerate(text.splitlines(), 1):
        if line.lstrip().startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        bare = re.sub(r"`[^`]*`", "", line).replace("\\$", "")
        # GitHub renders $...$ as math when the opening $ is followed by non-space and the closing $ preceded by non-space
        if re.search(r"(?<!\\)\$(?=\S)[^$\n]*?(?<=[^\s\\])\$", bare):
            out.append(f"line {n}: two '$' outside code may render as math on GitHub — wrap references in backticks or escape as \\$")
    no_code = re.sub(r"```.*?```", "", text, flags=re.S)
    no_code = re.sub(r"`[^`\n]*`", "", no_code)
    for m in re.finditer(r"\]\(([^)\s]+)\)", no_code):
        target = m.group(1).split("#")[0]
        if not target or re.match(r"^[a-z]+:", target):
            continue
        if not (path.parent / target.replace("%20", " ")).exists():
            out.append(f"broken relative link: {m.group(1)}")
    return out


# ---------------------------------------------------------------------------
# Lesson discovery & build
# ---------------------------------------------------------------------------
def builder_module(code: str):
    name = "lessons.lesson_" + code.replace(".", "_")
    return importlib.import_module(name)


def available_codes() -> list[str]:
    codes = []
    for l in CUR.LESSONS:
        if (HERE / "lessons" / f"lesson_{l.code.replace('.', '_')}.py").exists():
            codes.append(l.code)
    return codes


def build_lesson(code: str, do_verify: bool = True) -> list[str]:
    t0 = time.time()
    info = CUR.by_code(code)
    mod = builder_module(code)
    lesson: Lesson = mod.build()
    assert lesson.code == info.code, f"builder code {lesson.code} != curriculum {info.code}"
    assert lesson.module_dir == info.module and lesson.slug == info.slug, "builder folder doesn't match curriculum"
    lesson.dir.mkdir(parents=True, exist_ok=True)
    path = lesson.save(lesson.dir / lesson.workbook_name)
    lesson.write_extra_files()
    readme = lesson.dir / "README.md"
    if not readme.exists():
        readme.write_text(skeleton(lesson), encoding="utf-8")
        print(f"  created README skeleton at {readme.relative_to(ROOT)}")
    text = readme.read_text(encoding="utf-8")
    text = inject(text, lesson.readme_blocks(nav=nav_block(code)))
    readme.write_text(text, encoding="utf-8")
    size = path.stat().st_size / 1024
    print(f"Lesson {code}: wrote {path.relative_to(ROOT)} ({size:,.0f} KB), {len(lesson.tasks)} tasks + {len(lesson.bonus)} bonus")
    problems = []
    for w in lint_readme(readme, text):
        print("  warning:", w)
    if do_verify:
        problems = verify(lesson, path)
        print(f"  verification: {'PASS' if not problems else f'{len(problems)} problem(s)'}  ({time.time() - t0:.1f}s)")
    return problems


# ---------------------------------------------------------------------------
# Course README syllabus
# ---------------------------------------------------------------------------
def refresh_course_readmes():
    lines = []
    for mod, (mtitle, level, blurb) in CUR.MODULES.items():
        lines += [f"### [{mtitle}]({mod}/README.md)", "", f"*{level}* — {blurb}", "", "| # | Lesson | Time | You'll learn to… |", "|:-:|---|:-:|---|"]
        for l in [x for x in CUR.LESSONS if x.module == mod]:
            learn = "; ".join(o[0].lower() + o[1:] for o in l.objectives[:3])
            lines.append(f"| {l.code} | [{l.title}]({l.path}/README.md) | {l.minutes} min | {learn} |")
        lines.append("")
        # module README
        mlines = [f"# {mtitle}", "", f"**Level:** {level}", "", blurb, "", "| # | Lesson | Time | Objectives |", "|:-:|---|:-:|---|"]
        for l in [x for x in CUR.LESSONS if x.module == mod]:
            objs = "<br>".join("• " + o for o in l.objectives)
            mlines.append(f"| {l.code} | [{l.title}]({l.slug}/README.md) | {l.minutes} min | {objs} |")
        mlines += ["", "🏠 [Course home](../README.md)", ""]
        (ROOT / mod).mkdir(exist_ok=True)
        (ROOT / mod / "README.md").write_text(md_escape_dollars("\n".join(mlines)), encoding="utf-8")
    total = sum(l.minutes for l in CUR.LESSONS)
    lines.append(f"**{len(CUR.LESSONS)} lessons · about {total // 60} hours of guided practice.**")
    readme = ROOT / "README.md"
    text = readme.read_text(encoding="utf-8")
    text = inject(text, {"syllabus": md_escape_dollars("\n".join(lines))}, strict=True)
    readme.write_text(text, encoding="utf-8")
    print("refreshed README.md and module READMEs")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("targets", nargs="+", help="lesson codes like 1.4, 'all', or 'readme'")
    ap.add_argument("--no-verify", action="store_true")
    args = ap.parse_args()
    failures = {}
    for target in args.targets:
        if target == "readme":
            refresh_course_readmes()
            continue
        codes = available_codes() if target == "all" else [target]
        for code in codes:
            probs = build_lesson(code, do_verify=not args.no_verify)
            if probs:
                failures[code] = probs
    if failures:
        print("\nFAILED:", ", ".join(failures))
        sys.exit(1)


if __name__ == "__main__":
    main()
