"""Lesson 1.3 · Formatting Cells & Number Formats.

Formatting changes what a cell DISPLAYS, never what it STORES, so almost every task here asks the
learner to apply a format and then type exactly what the cell now shows (``($572,762.00)``, ``-4.6%``,
``412:47``…). Two consequences for the builder:

* The yellow answer cells for those tasks are formatted as Text (``fmt="@"``). Otherwise Excel would
  turn a typed ``-4.6%`` or ``(15.7%)`` back into a number and the text check could never match.
* Every expected display string is rendered in Python below (Excel-style rounding: 15 significant
  digits, then half away from zero) and cross-checked in the hidden key by a live
  ``=TEXT(value, "format code")`` formula, which LibreOffice evaluates during verification.

Three tasks ask for stored values (read from the formula bar) and are checked as numbers.
The practice sheets (Budget, Registry, Stays) and the bonus sheet (Dec Report) are plain ranges, not
Excel Tables, because Tables can't contain merged cells and their built-in styling would get in the way
of a formatting lesson. They are built in the ``customize`` hook, together with two hidden reference
sheets (Budget Key, Report Key) that show the finished, formatted versions.
"""
from __future__ import annotations

from datetime import date
from decimal import ROUND_HALF_UP, Decimal

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from xlcourse import Lesson, Task, data
from xlcourse.lesson import NAVY

CODE = "1.3"

# ---------------------------------------------------------------------------- format codes used in the lesson
CUR2_PAREN = '$#,##0.00_);($#,##0.00)'          # what Ctrl+Shift+$ applies (Windows)
PCT1 = "0.0%"
K0 = '$#,##0,"K"'                                # thousands with a K (practice, Budget sheet)
ACCT2 = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'   # Home > Accounting Number Format ($ button)
MRN8 = "00000000"
PHONE = "(000) 000-0000"
DT_CODE = "ddd mm/dd/yyyy h:mm AM/PM"
ELAPSED = "[h]:mm"
DAYS_SIGNED = '[Red]+0.0 "days";-0.0 "days"'
# bonus (Dec Report)
K1 = '$#,##0.0,"K"'
K1_RED = '$#,##0.0,"K";[Red]($#,##0.0,"K")'
PCT1_RED = "0.0%;[Red](0.0%)"
MK = '[>=1000000]$#,##0.00,,"M";[>=1000]$#,##0,"K";$#,##0'

TITLE_FONT = Font(bold=True, size=13, color=NAVY)
NOTE_FONT = Font(italic=True, size=10, color="595959")
RAW_HDR_FONT = Font(bold=True)
RAW_HDR_FILL = PatternFill("solid", fgColor="E7E6E6")
KEY_HDR_FONT = Font(bold=True, color="FFFFFF")
KEY_HDR_FILL = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="A6A6A6")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------------------- Excel-style display rendering
def _dec(x) -> Decimal:
    """The value Excel formats: the stored double, rounded to 15 significant digits."""
    return Decimal(str(x)) if isinstance(x, int) else Decimal(f"{x:.15g}")


def _round(d: Decimal, nd: int) -> Decimal:
    """Round half away from zero (Excel's display rounding) and refuse exact ties, so a reviewer never
    has to wonder which way a value rounds."""
    q = Decimal(1).scaleb(-nd)
    shifted = abs(d) * (Decimal(10) ** nd)
    assert shifted - int(shifted) != Decimal("0.5"), f"{d} is a rounding tie at {nd} decimals; pick another row"
    return d.quantize(q, rounding=ROUND_HALF_UP)


def _num(d: Decimal, nd: int, comma: bool = True) -> str:
    r = abs(_round(d, nd))
    return f"{r:,.{nd}f}" if comma else f"{r:.{nd}f}"


def show_currency_paren(v) -> str:          # $#,##0.00_);($#,##0.00)
    d = _dec(v)
    return f"(${_num(d, 2)})" if d < 0 else f"${_num(d, 2)}"


def show_pct(v, nd: int = 1, paren: bool = False) -> str:   # 0.0%  /  0.0%;(0.0%)
    d = _dec(v) * 100
    body = f"{_num(d, nd, comma=False)}%"
    if d < 0:
        return f"({body})" if paren else f"-{body}"
    return body


def show_k(v, nd: int = 0, paren: bool = False) -> str:     # $#,##0,"K"  /  $#,##0.0,"K";($#,##0.0,"K")
    d = _dec(v) / 1000
    body = f"${_num(d, nd)}K"
    if d < 0:
        return f"({body})" if paren else f"-{body}"
    return body


def show_mk(v) -> str:                                     # [>=1000000]$#,##0.00,,"M";[>=1000]$#,##0,"K";$#,##0
    d = _dec(v)
    assert d >= 1000
    return f"${_num(d / 1_000_000, 2)}M" if d >= 1_000_000 else f"${_num(d / 1000, 0)}K"


def show_days_signed(v) -> str:                            # [Red]+0.0 "days";-0.0 "days"
    d = _dec(v)
    return f"{'+' if d >= 0 else '-'}{_num(d, 1, comma=False)} days"


def show_elapsed(minutes: int) -> str:                     # [h]:mm
    return f"{minutes // 60}:{minutes % 60:02d}"


def show_datetime(dt) -> str:                              # ddd mm/dd/yyyy h:mm AM/PM
    h12 = dt.hour % 12 or 12
    return f"{dt:%a %m/%d/%Y} {h12}:{dt.minute:02d} {'AM' if dt.hour < 12 else 'PM'}"


def q(code: str) -> str:
    """A format code as a string literal inside an Excel formula (double the quotes)."""
    return '"' + code.replace('"', '""') + '"'


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="01-foundations", slug="03-formatting-cells",
        title="Formatting Cells & Number Formats", level="Beginner", minutes=45,
        objectives=[
            "Apply number formats: Number, Currency vs Accounting, Percentage, Date, Time, Text",
            "Write custom number formats (leading zeros, units, colors, thousands)",
            "Format fonts, fills, borders, alignment, and wrap text — and avoid merged cells",
            "Use Format Painter, cell styles, and themes for consistent reports",
            "Understand that formatting changes how a value looks, not the value itself",
        ],
        data_note="Bluestone Memorial Hospital's 2025 expense budget vs actual by department, two days of ED registrations "
                  "(12/30–12/31/2025), December 2025 discharges from Cardiac Step-Down, and the Emergency Department's "
                  "December 2025 budget report (bonus). Every number arrives unformatted, the way it comes out of a source system.",
    )

    # ======================================================================== source data
    depts = data.index(data.load("departments"), "DeptID")
    patients = data.index(data.load("patients"), "PatientID")
    payers = data.index(data.load("payers"), "PayerID")
    diagnoses = data.index(data.load("diagnoses"), "DxCode")
    claims = {c["EncounterID"]: c for c in data.load("claims")}
    budget = data.load("budget")

    # ---- Budget sheet: 2025 operating EXPENSE budget vs actual, one row per Bluestone Memorial (F01) department
    agg: dict[str, list[int]] = {}
    for r in budget:
        if r["FiscalYear"] == 2025 and r["FacilityID"] == "F01" and r["LineType"] == "Expense":
            a = agg.setdefault(r["DeptID"], [0, 0])
            a[0] += r["BudgetAmount"]
            a[1] += r["ActualAmount"]
    brows = sorted(({"ServiceLine": depts[d]["ServiceLine"], "Department": depts[d]["DeptName"], "Budget": b, "Actual": a,
                     "Variance": b - a, "VarPct": (b - a) / b} for d, (b, a) in agg.items()),
                   key=lambda r: (r["ServiceLine"], r["Department"]))
    B_HDR, B_FIRST = 4, 5
    B_LAST = B_FIRST + len(brows) - 1
    B_TOTAL = B_LAST + 1
    tot_b = sum(r["Budget"] for r in brows)
    tot_a = sum(r["Actual"] for r in brows)
    totals = {"Budget": tot_b, "Actual": tot_a, "Variance": tot_b - tot_a, "VarPct": (tot_b - tot_a) / tot_b}
    # service-line groups -> vertical merged blocks in column A (the "pretty but harmful" layout)
    groups: list[tuple[str, int, int]] = []
    for i, r in enumerate(brows):
        row = B_FIRST + i
        if groups and groups[-1][0] == r["ServiceLine"]:
            groups[-1] = (groups[-1][0], groups[-1][1], row)
        else:
            groups.append((r["ServiceLine"], row, row))
    merged_blocks = [g for g in groups if g[2] > g[1]]
    n_labels = len(groups)

    def brow(dept_name: str) -> tuple[int, dict]:
        i = next(k for k, r in enumerate(brows) if r["Department"] == dept_name)
        return B_FIRST + i, brows[i]

    worst_var_row, worst_var = min(((B_FIRST + i, r) for i, r in enumerate(brows)), key=lambda t: t[1]["Variance"])
    worst_pct_row, worst_pct = min(((B_FIRST + i, r) for i, r in enumerate(brows)), key=lambda t: t[1]["VarPct"])
    pharm_row, pharm = brow("Pharmacy")
    first_b = brows[0]

    # ---- Registry sheet: Bluestone Memorial ED registrations on 12/30 and 12/31/2025
    ed = sorted((v for v in data.load("ed_visits") if v["FacilityID"] == "F01"
                 and v["ArrivalDateTime"].date() in (date(2025, 12, 30), date(2025, 12, 31))),
                key=lambda v: v["ArrivalDateTime"])
    reg = []
    for v in ed:
        p = patients[v["PatientID"]]
        c = claims[v["EncounterID"]]
        reg.append({"Arrived": v["ArrivalDateTime"], "MRN": int(p["MRN"]), "MRNText": p["MRN"],
                    "Patient": f"{p['LastName']}, {p['FirstName']}",
                    "Phone": int("".join(ch for ch in p["Phone"] if ch.isdigit())),
                    "Payer": payers[c["PayerID"]]["PayerName"], "PatientDue": c["PatientResponsibility"]})
    R_HDR, R_FIRST = 4, 5
    DUE = "F"                                               # PatientDue column on the Registry sheet
    R_LAST = R_FIRST + len(reg) - 1
    zero_i = next(i for i, r in enumerate(reg) if r["PatientDue"] == 0)
    mrn_i = next((i for i, r in enumerate(reg) if r["MRNText"].startswith("00")), 0)
    phone_i = 0

    # ---- Stays sheet: Cardiac Step-Down inpatient discharges, December 2025
    enc = data.load("encounters")
    stays_src = sorted((e for e in enc if e["DeptID"] == "D120" and e["EncounterType"] == "Inpatient"
                        and e["DischargeDateTime"].year == 2025 and e["DischargeDateTime"].month == 12),
                       key=lambda e: (e["DischargeDateTime"], e["EncounterID"]))
    stays = []
    for e in stays_src:
        mins = round((e["DischargeDateTime"] - e["AdmitDateTime"]).total_seconds() / 60)
        los = mins / 1440                                    # days, exactly what =Discharged-Admitted stores
        exp = diagnoses[e["PrimaryDxCode"]]["ExpectedLOS"]
        stays.append({"EncounterID": e["EncounterID"], "Diagnosis": diagnoses[e["PrimaryDxCode"]]["DxDescription"],
                      "AdmitDT": e["AdmitDateTime"], "Admitted": data.excel_serial(e["AdmitDateTime"]),
                      "Discharged": e["DischargeDateTime"], "LOS": los, "Minutes": mins, "Expected": exp,
                      "VsExpected": round(los - exp, 4)})
    S_HDR, S_FIRST = 4, 5
    S_LAST = S_FIRST + len(stays) - 1
    dt_i = next(i for i, s in enumerate(stays) if 7 <= s["AdmitDT"].hour <= 11)
    stored_i = next(i for i, s in enumerate(stays) if abs(s["LOS"] - round(s["LOS"], 1)) > 0.03)
    long_i = max(range(len(stays)), key=lambda i: stays[i]["LOS"])
    neg_i = next(i for i, s in enumerate(stays) if s["VsExpected"] < -0.5)
    pos_n = sum(1 for s in stays if s["VsExpected"] > 0)

    # ---- Bonus: Emergency Department (D100) December 2025 operating report, by expense category
    cat_m: dict[str, list[int]] = {}
    cat_y: dict[str, list[int]] = {}
    for r in budget:
        if r["DeptID"] == "D100" and r["FiscalYear"] == 2025 and r["LineType"] == "Expense":
            y = cat_y.setdefault(r["Category"], [0, 0])
            y[0] += r["BudgetAmount"]
            y[1] += r["ActualAmount"]
            if r["FiscalMonth"] == 12:
                m = cat_m.setdefault(r["Category"], [0, 0])
                m[0] += r["BudgetAmount"]
                m[1] += r["ActualAmount"]
    rep = []
    for cat in sorted(cat_m):
        mb, ma = cat_m[cat]
        yb, ya = cat_y[cat]
        rep.append({"Category": cat, "DecBudget": mb, "DecActual": ma, "DecVar": mb - ma, "DecVarPct": (mb - ma) / mb,
                    "YtdBudget": yb, "YtdActual": ya, "YtdVar": yb - ya, "YtdVarPct": (yb - ya) / yb})
    tm_b, tm_a = sum(r["DecBudget"] for r in rep), sum(r["DecActual"] for r in rep)
    ty_b, ty_a = sum(r["YtdBudget"] for r in rep), sum(r["YtdActual"] for r in rep)
    rep_total = {"Category": "Total", "DecBudget": tm_b, "DecActual": tm_a, "DecVar": tm_b - tm_a,
                 "DecVarPct": (tm_b - tm_a) / tm_b, "YtdBudget": ty_b, "YtdActual": ty_a, "YtdVar": ty_b - ty_a,
                 "YtdVarPct": (ty_b - ty_a) / ty_b}
    assert ty_a == agg["D100"][1], "Dec Report YTD actual should tie to the Budget sheet's ED actual"
    RP_COLS = ["Category", "DecBudget", "DecActual", "DecVar", "DecVarPct", "YtdBudget", "YtdActual", "YtdVar", "YtdVarPct"]
    RP_HEAD = ["Expense category", "Dec Budget", "Dec Actual", "Dec Variance", "Dec Var %", "YTD Budget", "YTD Actual",
               "YTD Variance", "YTD Var %"]
    rcol = {c: get_column_letter(j) for j, c in enumerate(RP_COLS, 1)}
    P_HDR, P_FIRST = 4, 5
    P_LAST = P_FIRST + len(rep) - 1
    P_TOTAL = P_LAST + 1

    def prow(cat: str) -> tuple[int, dict]:
        i = next(k for k, r in enumerate(rep) if r["Category"] == cat)
        return P_FIRST + i, rep[i]

    b1_row, b1 = prow("Employee Benefits")
    b2_row, b2 = prow("Salaries & Wages")
    b3_row, b3 = prow("Equipment & Maintenance")
    b5_row, b5 = prow("Other Operating")
    assert b2["DecVar"] < 0 and b3["YtdVarPct"] < 0 and b5["YtdBudget"] < 1_000_000 <= rep_total["YtdActual"]
    assert b1["DecVar"] < 0 and b1["DecVarPct"] < 0, "Rules 3-4 quote Employee Benefits' Dec variance as the negative example"
    n_red_dec = sum(1 for r in rep + [rep_total] if r["DecVar"] < 0)

    # ======================================================================== tasks
    L.practice_intro = (
        "Tasks 1–4 and 12–13 use the Budget sheet, tasks 5–7 the Registry sheet, and tasks 8–11 the Stays sheet. Work in order, "
        "because some tasks build on earlier ones. Most tasks ask you to apply a format and then type exactly what the cell "
        "displays, including any dollar sign, commas, parentheses, minus sign, or % sign. Those yellow cells are formatted as Text, "
        "so Excel keeps your entry exactly as you type it. Type the answer rather than pasting a copied cell, because pasting brings "
        "the cell's number and format along. Tasks 3, 9, and 12 ask for a plain number instead. Answers assume US regional settings.")

    def bcell(col: str, row: int) -> str:
        return f"Budget!{col}{row}"

    s9, sL, sN = stays[stored_i], stays[long_i], stays[neg_i]
    phone_digits = str(reg[phone_i]["Phone"])
    phone_shown = f"({phone_digits[:3]}) {phone_digits[3:6]}-{phone_digits[6:]}"

    L.tasks = [
        # ------------------------------------------------------------------ built-in formats & shortcuts (Budget)
        Task(f"On the Budget sheet, select the Variance cells E{B_FIRST}:E{B_TOTAL} and press Ctrl + Shift + $ "
             f"(Mac: ⌃ + Shift + $). What does E{worst_var_row} ({worst_var['Department']}, which spent more than its "
             f"budget) display now?",
             answer=show_currency_paren(worst_var["Variance"]), fmt="@", title="Currency shortcut on a negative variance",
             solution=f"1. Select **Budget!E{B_FIRST}:E{B_TOTAL}**.\n"
                      "2. Press **Ctrl + Shift + $** (Mac: **⌃ + Shift + $**).\n"
                      f"3. Read **E{worst_var_row}**.",
             live=f"=TEXT({bcell('E', worst_var_row)},{q(CUR2_PAREN)})",
             hint="Negative amounts in this format don't use a minus sign",
             explanation=f"The shortcut applies the Currency format with two decimals, whose code is `{CUR2_PAREN}`. The part after the "
                         "semicolon is used for negative numbers, and it wraps them in parentheses instead of showing a minus sign, "
                         f"the way accountants write losses. The stored value is still {worst_var['Variance']}. Click the cell and look "
                         "at the formula bar to confirm. (The key's live formula uses TEXT, a Lesson 2.2 function that returns what a "
                         "format would display.)"),
        Task(f"Select the Variance % cells F{B_FIRST}:F{B_TOTAL}, press Ctrl + Shift + % (Mac: ⌃ + Shift + %), then click "
             f"Home → Increase Decimal once. What does F{worst_pct_row} ({worst_pct['Department']}) display?",
             answer=show_pct(worst_pct["VarPct"]), fmt="@", title="Percentage with one decimal place",
             solution=f"1. Select **Budget!F{B_FIRST}:F{B_TOTAL}**.\n"
                      "2. Press **Ctrl + Shift + %** (Mac: **⌃ + Shift + %**). The cells show whole percentages.\n"
                      "3. Click **Home → Increase Decimal** (the .00 button with the left arrow) once.\n"
                      f"4. Read **F{worst_pct_row}**.",
             live=f"=TEXT({bcell('F', worst_pct_row)},{q(PCT1)})",
             hint="The shortcut shows 0 decimals; each Increase Decimal click adds one",
             explanation=f"The cell stores {worst_pct['VarPct']:.6f}… The Percentage format multiplies by 100 for display only and adds the % "
                         f"sign, so it shows {show_pct(worst_pct['VarPct'], 0)} at first and {show_pct(worst_pct['VarPct'])} after one "
                         "Increase Decimal click. Excel rounds the display, but the stored value keeps every digit."),
        Task(f"C{B_FIRST} already has a finance format: it shows the {first_b['Department']} budget in thousands with a K. "
             f"What number is actually stored in C{B_FIRST}? Click the cell and read the formula bar.",
             answer=first_b["Budget"], title="Stored value behind a thousands (K) display",
             solution=f"Click **Budget!C{B_FIRST}** and read the formula bar. It shows the full dollar amount.",
             live=f"={bcell('C', B_FIRST)}",
             hint="The cell shows one thing; the formula bar shows another",
             explanation=f"The cell displays `{show_k(first_b['Budget'])}`, but it stores {first_b['Budget']:,}. The custom code `{K0}` "
                         "ends with a comma after the last 0, which tells Excel to *display* the number divided by 1,000. Any formula that "
                         f"uses C{B_FIRST} still gets the full amount, so totals stay exact even when every cell is shown in thousands."),
        Task(f"Copy C{B_FIRST}'s format to the rest of the money columns with Format Painter: select C{B_FIRST}, double-click "
             f"Home → Format Painter, drag over C{B_FIRST}:E{B_LAST}, then drag over the Total row C{B_TOTAL}:E{B_TOTAL}, and press Esc. "
             f"What does E{pharm_row} (Pharmacy's variance) display now?",
             answer=show_k(pharm["Variance"]), fmt="@", title="Format Painter: thousands format on a negative number",
             solution=f"1. Select **Budget!C{B_FIRST}**.\n"
                      "2. **Double-click** **Home → Format Painter** (the paintbrush). Double-clicking keeps it switched on.\n"
                      f"3. Drag over **C{B_FIRST}:E{B_LAST}**, then drag over the Total row, **C{B_TOTAL}:E{B_TOTAL}**.\n"
                      "4. Press **Esc** (or click the paintbrush again) to switch it off.\n"
                      f"5. Read **E{pharm_row}**.",
             live=f"=TEXT({bcell('E', pharm_row)},{q(K0)})",
             hint="A format with only one section puts a minus sign in front of negatives",
             explanation=f"Format Painter copies *all* of C{B_FIRST}'s formatting (number format, font, fill, borders, alignment) onto the cells "
                         f"you paint. `{K0}` has a single section, so Excel uses it for every number and simply puts a minus sign in front "
                         f"of negatives: {pharm['Variance']:,} shows as `{show_k(pharm['Variance'])}`. The parentheses from task 1 are "
                         "gone because each cell has only one number format, and the newest one wins. A single click on Format Painter "
                         "paints once. A double-click lets you paint several ranges until you press Esc."),
        # ------------------------------------------------------------------ Accounting & custom codes (Registry)
        Task(f"On the Registry sheet, select the PatientDue cells {DUE}{R_FIRST}:{DUE}{R_LAST} and click Home → Accounting "
             f"Number Format (the $ button). {reg[zero_i]['Patient']} (row {R_FIRST + zero_i}) owes nothing for this visit. "
             f"What character does {DUE}{R_FIRST + zero_i} show where the 0 used to be?",
             answer="-", fmt="@", accept=["$ -", "$-", "dash", "a dash", "hyphen", "a hyphen", "–", "$ –"],
             title="Accounting format: how a zero looks",
             solution=f"1. Select **Registry!{DUE}{R_FIRST}:{DUE}{R_LAST}**.\n"
                      "2. Click the **$** button in the **Number** group of the **Home** tab (its ScreenTip says Accounting Number Format).\n"
                      f"3. Look at **{DUE}{R_FIRST + zero_i}**: the $ sits at the left edge and a dash sits near the right.",
             live=False,
             hint="Compare Accounting with Currency in the guide's table",
             explanation=f"The Accounting code is `{ACCT2}`. Its third section (zero) prints a dash, so a column of balances shows "
                         "which patients owe nothing at a glance. Accounting also pins the $ to the left edge of the cell and lines up "
                         "the decimal points, and it always shows negatives in parentheses. Currency puts the $ right next to the "
                         "number and shows `$0.00` for zero. The cell still stores 0."),
        Task(f"The MRN column lost its leading zeros on the way out of the registration system. Bluestone MRNs are always 8 digits. "
             f"Select B{R_FIRST}:B{R_LAST}, open Format Cells (Ctrl + 1; Mac: ⌘ + 1), choose Custom, and type the code "
             f"{MRN8} in the Type box. What does B{R_FIRST + mrn_i} display?",
             answer=reg[mrn_i]["MRNText"], fmt="@", title="Custom format 00000000 for MRNs",
             solution=f"1. Select **Registry!B{R_FIRST}:B{R_LAST}** and press **Ctrl + 1** (Mac: **⌘ + 1**).\n"
                      f"2. On the **Number** tab choose **Custom**, replace the Type box with `{MRN8}`, and click **OK**.\n"
                      f"3. Read **B{R_FIRST + mrn_i}**.",
             live=f"=TEXT(Registry!B{R_FIRST + mrn_i},{q(MRN8)})",
             hint="Each 0 in the code is a digit that always shows, even when it's a zero",
             explanation=f"The cell stores the number {reg[mrn_i]['MRN']}. Each `0` in `{MRN8}` is a placeholder that forces a digit, so "
                         f"Excel pads the display to 8 digits: `{reg[mrn_i]['MRNText']}`. This fixes how the MRN *looks* on a printed list. "
                         "The value is still a number, though, so it won't match the text MRN `" + reg[mrn_i]["MRNText"] + "` in "
                         "another system or in a lookup. When you control data entry, store IDs as text instead (Lesson 1.2)."),
        Task(f"The Phone column (D) stores 10-digit numbers. Write a custom number format that shows each one the usual US way: "
             f"the first three digits in parentheses, a space, three digits, a hyphen, and the last four digits. For example, "
             f"D{R_FIRST + phone_i} ({phone_digits}) should display as {phone_shown}. Apply your format to D{R_FIRST}:D{R_LAST}, "
             f"then type the format code you used.",
             answer=PHONE, fmt="@", check="custom",
             custom_check=('OR(LOWER(TRIM(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({cell}&"","\\",""),CHAR(34),""),"#","0")))'
                           '="(000) 000-0000",LOWER(TRIM(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({cell}&"","\\",""),CHAR(34),""),"#","0")))'
                           '="[<=9999999]000-0000;(000) 000-0000")'),
             title="Write a phone-number format code",
             solution=f"1. Select **Registry!D{R_FIRST}:D{R_LAST}** and press **Ctrl + 1** (Mac: **⌘ + 1**).\n"
                      f"2. Choose **Custom** and type `{PHONE}` in the Type box. Click **OK**.\n"
                      f"3. Type `{PHONE}` in the answer cell.",
             live=False,
             hint="0 is a digit placeholder; parentheses, spaces, and hyphens can be typed as they are",
             explanation=f"`{PHONE}` has ten 0 placeholders, filled from the right with the number's ten digits. Parentheses, the space, "
                         "and the hyphen are characters Excel prints as-is, without quotes. The check also accepts `(###) ###-####` "
                         "(the same for 10-digit numbers) and Excel's built-in **Special → Phone Number** format, "
                         "`[<=9999999]###-####;(###) ###-####`, which you'll see in the Custom box if you used Special."),
        # ------------------------------------------------------------------ dates, times, durations (Stays)
        Task(f"On the Stays sheet, the Admitted column (C) shows serial numbers because the export lost its date format. "
             f"Select C{S_FIRST}:C{S_LAST} and apply the custom format {DT_CODE}. What does C{S_FIRST + dt_i} display?",
             answer=show_datetime(stays[dt_i]["AdmitDT"]), fmt="@", title="Custom date-and-time format",
             solution=f"1. Select **Stays!C{S_FIRST}:C{S_LAST}** and press **Ctrl + 1** (Mac: **⌘ + 1**).\n"
                      f"2. Choose **Custom**, type `{DT_CODE}`, and click **OK**.\n"
                      f"3. Read **C{S_FIRST + dt_i}**.",
             live=f"=TEXT(Stays!C{S_FIRST + dt_i},{q(DT_CODE)})",
             hint="ddd is the short day name; mm right after h means minutes",
             explanation=f"The cell stores {stays[dt_i]['Admitted']:.5f}: the whole number is the date and the decimal is the time of "
                         "day. `ddd` gives the short day name, `mm/dd/yyyy` the date with leading zeros, and `h:mm AM/PM` a 12-hour "
                         "time without a leading zero. Because `mm` follows `h`, Excel reads it as minutes, not months."),
        Task(f"The LOS (days) column is formatted to show one decimal place. E{S_FIRST + stored_i} shows "
             f"{_num(_dec(s9['LOS']), 1, comma=False)}. What value is actually stored in that cell? Read the formula bar and enter "
             f"at least 4 decimal places.",
             answer=s9["LOS"], tol=0.0001, title="Stored value behind a one-decimal display",
             answer_display=f"{s9['LOS']:.8f}… (any entry within 0.0001 is accepted)",
             solution=f"Click **Stays!E{S_FIRST + stored_i}** and read the formula bar.",
             live=f"=Stays!E{S_FIRST + stored_i}",
             hint="Formatting rounds the display, not the value",
             explanation=f"The stay lasted {s9['Minutes']:,} minutes, and {s9['Minutes']:,} ÷ 1,440 minutes per day = "
                         f"{s9['LOS']:.8f}… days. The `0.0` format only rounds what you see. If you add up the LOS column, Excel adds "
                         "the full values, so a total can differ slightly from the sum of the rounded numbers on screen. Use ROUND "
                         "(Lesson 1.4) when the rounded number is the one you mean."),
        Task(f"Hospitals often track length of stay in hours. Select E{S_FIRST}:E{S_LAST} and apply the custom format {ELAPSED}. "
             f"What does E{S_FIRST + long_i} (the longest stay of the month) display?",
             answer=show_elapsed(sL["Minutes"]), fmt="@", title="Elapsed hours with [h]:mm",
             solution=f"1. Select **Stays!E{S_FIRST}:E{S_LAST}** and press **Ctrl + 1** (Mac: **⌘ + 1**).\n"
                      f"2. Choose **Custom**, type `{ELAPSED}`, and click **OK**.\n"
                      f"3. Read **E{S_FIRST + long_i}**.",
             live=f"=TEXT(Stays!E{S_FIRST + long_i},{q(ELAPSED)})",
             hint="Square brackets let the hours keep counting past 24",
             explanation=f"The cell stores {sL['LOS']:.4f} days. One day is 24 hours, so that's {sL['Minutes'] // 60} hours and "
                         f"{sL['Minutes'] % 60} minutes. `[h]` shows *elapsed* hours. Plain `h:mm` would show "
                         f"`{(sL['Minutes'] // 60) % 24}:{sL['Minutes'] % 60:02d}`, the clock time, because ordinary hours roll back to "
                         "0 every 24 hours. Use `[h]:mm` for any duration that can pass a day: LOS, shift hours, ED boarding time."),
        Task(f"Column G (LOS vs Expected) stores LOS minus the diagnosis's expected LOS, in days, so a positive value means the "
             f"patient stayed longer than expected. Write a two-section custom format for G{S_FIRST}:G{S_LAST}: positive values in red "
             f"with a plus sign and the word days (like +2.1 days), and negative values with a minus sign (like -0.4 days), each with "
             f"one decimal. What does G{S_FIRST + neg_i} display?",
             answer=show_days_signed(sN["VsExpected"]), fmt="@", title="Two sections, a color, and a unit",
             solution=f"1. Select **Stays!G{S_FIRST}:G{S_LAST}** and press **Ctrl + 1** (Mac: **⌘ + 1**).\n"
                      f"2. Choose **Custom**, type `{DAYS_SIGNED}`, and click **OK**.\n"
                      f"3. Read **G{S_FIRST + neg_i}**.",
             live=f"=TEXT(Stays!G{S_FIRST + neg_i},{q(DAYS_SIGNED)})",
             hint="Sections are separated by semicolons: positive;negative. The second section needs its own minus sign",
             explanation=f"`{DAYS_SIGNED}` has two sections. The first (positive and zero) starts with the color `[Red]`, then a literal +, "
                         "the number, and the text \"days\" in quotes. The second section (negative) has no color. When a format has a "
                         "negative section, Excel stops adding the minus sign for you, so you type it yourself. Without it, "
                         f"{sN['VsExpected']} would display as `{show_days_signed(sN['VsExpected'])[1:]}` and look like a positive "
                         f"number. {pos_n} of the {len(stays)} stays show in red because they lasted longer than expected."),
        # ------------------------------------------------------------------ alignment, merged cells, styles (Budget)
        Task(f"Back on the Budget sheet, column A labels each department's service line with merged cells. Select A{B_FIRST}:A{B_LAST} "
             f"and read Count on the status bar (it counts cells that aren't empty). How many of those {len(brows)} cells actually "
             f"contain a service-line label?",
             answer=n_labels, title="What merged cells really store",
             solution=f"Select **Budget!A{B_FIRST}:A{B_LAST}** and read **Count** on the status bar.",
             live=f"=COUNTA(Budget!A{B_FIRST}:A{B_LAST})",
             hint="A merged block stores its text in one cell only",
             explanation=f"There are {len(brows)} departments but only {n_labels} labels, because a merged block keeps its value in the "
                         f"top-left cell and the other cells are empty. For example, only the first {merged_blocks[1][0]} row "
                         f"really says {merged_blocks[1][0]}. "
                         "Sort, filter, copy, or count by service line and the rows under each block behave as if they had no service line. "
                         "Excel even refuses to sort a range whose merged cells are different sizes. Task 13 replaces the merges with a "
                         f"label on every row. (After task 13, the key's live COUNTA shows {len(brows)}, because every row then has its "
                         "label.)"),
        Task(f"Finish the Budget report so it looks like the hidden Budget Key sheet: "
             f"(1) unmerge A1:F1 and center the title with Center Across Selection; "
             f"(2) unmerge column A and fill each service line down so every row has its label; "
             f"(3) make header row {B_HDR} bold with white text on a dark blue fill, wrapped and centered; "
             f"(4) add All Borders to A{B_HDR}:F{B_TOTAL}; "
             f"(5) apply the Total cell style to A{B_TOTAL}:F{B_TOTAL}; (6) select A{B_HDR}:F{B_TOTAL} and AutoFit the columns "
             f"with Home → Format → AutoFit Column Width. "
             f"Then zoom in on row {B_TOTAL}: what kind of line does the Total style draw along the bottom of the row? "
             f"Answer in a word or two.",
             answer="double", fmt="@",
             accept=["double line", "a double line", "double border", "a double border", "double bottom border",
                     "a double bottom border", "bottom double border", "double line border", "double underline",
                     "double-line", "two lines", "2 lines", "double lines"],
             title="Style the report (Center Across Selection, header, borders, Total style)",
             solution=("1. **Title:** select **A1** and click **Home → Merge & Center** to unmerge it. Select **A1:F1**, press "
                       "**Ctrl + 1** (Mac: **⌘ + 1**), and on the **Alignment** tab set **Horizontal** to **Center Across Selection**.\n"
                       f"2. **Service lines:** select **A{B_FIRST}:A{B_LAST}** and choose **Home → Merge & Center ▾ → Unmerge Cells**. "
                       "Then, for each block that now has blank cells under its label ("
                       + ", ".join(f"A{a}:A{b}" for _, a, b in merged_blocks)
                       + "), select the block and press **Ctrl + D** (Mac: **⌘ + D**) to fill the label down.\n"
                       f"3. **Header:** select **A{B_HDR}:F{B_HDR}**. Press **Ctrl + B** (Mac: **⌘ + B**), pick **Home → Font Color ▾ → "
                       "White**, **Home → Fill Color ▾ → Dark Blue** (or any dark theme color), then click **Wrap Text** and **Center**.\n"
                       f"4. **Borders:** select **A{B_HDR}:F{B_TOTAL}** and choose **Home → Borders ▾ → All Borders**.\n"
                       f"5. **Total row:** select **A{B_TOTAL}:F{B_TOTAL}** and choose **Home → Cell Styles → Total** "
                       "(in the *Titles and Headings* group).\n"
                       f"6. **AutoFit:** select **A{B_HDR}:F{B_TOTAL}** (just the table, not whole columns) and choose **Home → Format → "
                       "AutoFit Column Width** (Windows KeyTips **Alt, H, O, I**). Selecting whole columns would also measure the long "
                       "note in A2, and column A would grow as wide as that sentence.\n"
                       "7. Unhide **Budget Key** (right-click a sheet tab → **Unhide…**) and compare."),
             live=False,
             hint="Cell Styles is on the Home tab; look closely at the bottom edge of the total row",
             explanation=("The built-in **Total** style makes the text bold and draws a thin line above the row and a **double** line "
                          "below it, the accounting convention for a grand total. Because it's a *style*, every total row you apply it to "
                          "looks the same. Its lines use a theme color, so if you switched the workbook to another theme in "
                          "**Page Layout → Themes**, the Total lines would change color while a header fill picked from Standard Colors "
                          "(Dark Blue) would stay put. Center Across Selection looks exactly like Merge & Center but leaves every cell "
                          "independent, so sorting, selecting columns, and copying keep working."))
    ]

    # ======================================================================== bonus
    L.bonus_title = "Bonus: Format the ED's monthly budget report"
    L.bonus_scenario = (
        "The Emergency Department director presents the December 2025 operating report to the CFO next week. Finance's style "
        "guide has five rules for department reports. Apply them on the Dec Report sheet, writing every custom format code "
        "yourself. Then type what each listed Dec Report cell displays into the yellow cells on the Bonus sheet. When you finish, "
        "compare your sheet with the hidden Report Key sheet.\n\n"
        "- Rule 1: Center the title across A1:I1 without merging. Make the header row bold white text on a dark blue fill, "
        "wrapped and centered. Give the Total row the Total cell style.\n"
        "- Rule 2: Show Dec Budget and Dec Actual in thousands with one decimal place, a dollar sign, and a K, so "
        f"{b2['DecBudget']:,} displays as {show_k(b2['DecBudget'], 1)}.\n"
        "- Rule 3: Show both Variance columns the same way, except that negative variances (over budget) are red and in "
        f"parentheses instead of having a minus sign, so {b1['DecVar']:,} displays as {show_k(b1['DecVar'], 1, paren=True)}.\n"
        "- Rule 4: Show both Var % columns with one decimal place, with negatives red and in parentheses, so "
        f"{b1['DecVarPct']:.4f}… displays as {show_pct(b1['DecVarPct'], 1, paren=True)}.\n"
        "- Rule 5: Show YTD Budget and YTD Actual with a dollar sign: in millions with two decimals and an M when the amount is "
        f"1,000,000 or more ({b2['YtdActual']:,} displays as {show_mk(b2['YtdActual'])}), and otherwise in thousands with no "
        "decimals and a K.")
    L.bonus = [
        Task(f"Rule 2: on the Dec Report sheet, what does {rcol['DecActual']}{b1_row} (Employee Benefits, Dec Actual) display?",
             answer=show_k(b1["DecActual"], 1), fmt="@", title="Thousands with one decimal",
             solution=f"1. **Rule 1** uses the same steps as practice task 13: select **'Dec Report'!A1:I1** and set **Center Across "
                      f"Selection** (**Ctrl + 1**, Mac: **⌘ + 1** → **Alignment** tab), style the header **A{P_HDR}:I{P_HDR}** (bold, white font, dark blue fill, "
                      f"Wrap Text, Center), and apply **Home → Cell Styles → Total** to **A{P_TOTAL}:I{P_TOTAL}**.\n"
                      f"2. **Rule 2:** select **'Dec Report'!B{P_FIRST}:C{P_TOTAL}**, press **Ctrl + 1** (Mac: **⌘ + 1**), choose "
                      f"**Custom**, and type `{K1}`.\n"
                      f"3. Read **{rcol['DecActual']}{b1_row}**.",
             live=f"=TEXT('Dec Report'!{rcol['DecActual']}{b1_row},{q(K1)})",
             hint="One comma after the last digit placeholder divides the display by 1,000",
             explanation=f"{b1['DecActual']:,} ÷ 1,000 = {b1['DecActual'] / 1000:,.3f}, shown with one decimal as "
                         f"`{show_k(b1['DecActual'], 1)}`. The comma between `#` and `##0` is the thousands separator. The comma after "
                         "`0.0` is the scaling comma."),
        Task(f"Rule 3: on the Dec Report sheet, what does {rcol['DecVar']}{b2_row} (Salaries & Wages, Dec Variance) display?",
             answer=show_k(b2["DecVar"], 1, paren=True), fmt="@", title="Negative variance: red, parentheses, thousands",
             solution=f"Select **'Dec Report'!D{P_FIRST}:D{P_TOTAL}**, then hold **Ctrl** (Mac: **⌘**) and drag over "
                      f"**H{P_FIRST}:H{P_TOTAL}** to add the YTD Variance column. Apply the custom format `{K1_RED}`.",
             live=f"=TEXT('Dec Report'!{rcol['DecVar']}{b2_row},{q(K1_RED)})",
             hint="Two sections: positive;negative. Put the color first in the negative section",
             explanation=f"The first section formats positive variances (under budget). The second section starts with `[Red]` and "
                         f"wraps the same thousands pattern in parentheses, so {b2['DecVar']:,} shows as "
                         f"`{show_k(b2['DecVar'], 1, paren=True)}` in red. Because the format has a negative section, Excel adds no "
                         f"minus sign. In December, {n_red_dec} of the {len(rep) + 1} Dec Variance cells (including the total) turn red."),
        Task(f"Rule 4: on the Dec Report sheet, what does {rcol['YtdVarPct']}{b3_row} (Equipment & Maintenance, YTD Var %) display?",
             answer=show_pct(b3["YtdVarPct"], 1, paren=True), fmt="@", title="Negative percentage in red parentheses",
             solution=f"Select **'Dec Report'!E{P_FIRST}:E{P_TOTAL}**, then hold **Ctrl** (Mac: **⌘**) and drag over "
                      f"**I{P_FIRST}:I{P_TOTAL}**. Apply the custom format `{PCT1_RED}`.",
             live=f"=TEXT('Dec Report'!{rcol['YtdVarPct']}{b3_row},{q(PCT1_RED)})",
             hint="A % in a custom code multiplies by 100, in every section where it appears",
             explanation=f"The cell stores {b3['YtdVarPct']:.6f}… Each section has its own `%`, so both multiply by 100. The negative "
                         f"section adds `[Red]` and parentheses: `{show_pct(b3['YtdVarPct'], 1, paren=True)}`."),
        Task(f"Rule 5: on the Dec Report sheet, what does {rcol['YtdActual']}{P_TOTAL} (the Total row's YTD Actual) display?",
             answer=show_mk(rep_total["YtdActual"]), fmt="@", title="Conditional format: millions",
             solution=f"Select **'Dec Report'!F{P_FIRST}:G{P_TOTAL}** and apply the custom format `{MK}`.",
             live=f"=TEXT('Dec Report'!{rcol['YtdActual']}{P_TOTAL},{q(MK)})",
             hint="Conditions go in square brackets at the start of a section, like [>=1000000]",
             explanation=f"A condition in square brackets replaces the usual positive/negative meaning of a section. "
                         f"`[>=1000000]` sends {rep_total['YtdActual']:,} to the first section, where two scaling commas divide by "
                         f"1,000,000: `{show_mk(rep_total['YtdActual'])}`. This total matches the Emergency Department's 2025 actual on "
                         "the Budget sheet, because December closes the fiscal year."),
        Task(f"Rule 5: on the Dec Report sheet, what does {rcol['YtdBudget']}{b5_row} (Other Operating, YTD Budget) display?",
             answer=show_mk(b5["YtdBudget"]), fmt="@", title="Conditional format: thousands",
             solution="Same format as the previous part. Check that values under 1,000,000 fall through to the K section.",
             live=f"=TEXT('Dec Report'!{rcol['YtdBudget']}{b5_row},{q(MK)})",
             hint="A value that fails the first condition moves on to the next section",
             explanation=f"{b5['YtdBudget']:,} is less than 1,000,000, so Excel skips the first section and uses `[>=1000]$#,##0,\"K\"`: "
                         f"`{show_mk(b5['YtdBudget'])}`. The full code is `{MK}`. Its third section catches anything under 1,000. "
                         f"Without conditions, every value would get the same scale, and a `{show_mk(b5['YtdBudget'])}` line would show "
                         f"as `${_num(_dec(b5['YtdBudget']) / 1_000_000, 2)}M`."),
    ]

    L.start_notes = [
        "Work sheets: Budget (tasks 1–4 and 12–13), Registry (tasks 5–7), Stays (tasks 8–11), and Dec Report (bonus).",
        "Two more hidden sheets, Budget Key and Report Key, show the finished, formatted versions of the Budget and Dec Report "
        "sheets. Unhide them to compare once you've tried.",
        "Answers assume US regional settings ($, comma thousands separators, month/day/year dates).",
    ]
    L.sheet_order = ["Start Here", "Practice", "Budget", "Registry", "Stays", "Bonus", "Dec Report",
                     "Answer Key", "Bonus Key", "Budget Key", "Report Key"]

    # ======================================================================== sheets (built in customize)
    def titled(ws, title: str, note: str):
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        ws["A2"] = note
        ws["A2"].font = NOTE_FONT

    def header(ws, row: int, labels: list[str], raw: bool = False):
        for j, h in enumerate(labels, 1):
            c = ws.cell(row=row, column=j, value=h)
            if not raw:
                c.font = RAW_HDR_FONT
                c.fill = RAW_HDR_FILL
                c.border = GRID

    def widths(ws, ws_widths: dict[str, float]):
        for col, w in ws_widths.items():
            ws.column_dimensions[col].width = w

    BUDGET_HEAD = ["Service Line", "Department", "Budget", "Actual", "Variance", "Variance %"]
    # Kept short on purpose: a centered title wider than A:F spills off the left edge of column A and gets cut off,
    # both in the Budget Key (Title style, 18 pt) and on the learner's sheet after task 13's AutoFit narrows the columns.
    BUDGET_TITLE = "Bluestone Memorial Hospital · 2025 Expense Budget vs Actual"
    BUDGET_NOTE = ("Finance system export, 12/31/2025. Variance = Budget − Actual (negative = over budget). "
                   "Variance % = Variance ÷ Budget.")
    REPORT_TITLE = "Emergency Department · Monthly Operating Expense Report · December 2025"
    REPORT_NOTE = ("Bluestone Memorial Hospital, cost center 6100. Variance = Budget − Actual (negative = over budget). "
                   "YTD = January–December 2025.")

    def budget_values(ws):
        for i, r in enumerate(brows):
            row = B_FIRST + i
            for j, k in enumerate(["ServiceLine", "Department", "Budget", "Actual", "Variance", "VarPct"], 1):
                ws.cell(row=row, column=j, value=r[k])
        ws.cell(row=B_TOTAL, column=1, value="Total")
        for j, k in zip((3, 4, 5, 6), ("Budget", "Actual", "Variance", "VarPct")):
            ws.cell(row=B_TOTAL, column=j, value=totals[k])

    def report_values(ws):
        for i, r in enumerate(rep + [rep_total]):
            for j, k in enumerate(RP_COLS, 1):
                ws.cell(row=P_FIRST + i, column=j, value=r[k])

    def style_header(ws, row: int, ncols: int):
        for j in range(1, ncols + 1):
            c = ws.cell(row=row, column=j)
            c.font = KEY_HDR_FONT
            c.fill = KEY_HDR_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 32

    def grid(ws, rng: str):
        for row in ws[rng]:
            for c in row:
                c.border = GRID

    @L.customize
    def _sheets(wb, lesson, selftest):
        # Practice answer column: a little wider so long display strings fit
        wb[lesson.practice_sheet].column_dimensions["D"].width = 26
        wb[lesson.bonus_sheet].column_dimensions["D"].width = 26
        # Answer Key "Live result" column: wide enough for task 8's "Tue 11/25/2025 11:58 AM" (the library default is 16)
        wb[lesson.key_sheet].column_dimensions["E"].width = 25

        # ------------------------------------------------------------------ Budget (raw, half-finished by someone else)
        ws = wb.create_sheet("Budget")
        ws.sheet_properties.tabColor = "2E75B6"
        ws["A1"] = BUDGET_TITLE
        ws["A1"].font = Font(bold=True, size=12)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.merge_cells(f"A1:F1")
        ws["A2"] = BUDGET_NOTE
        ws["A2"].font = Font(italic=True, size=10)
        header(ws, B_HDR, BUDGET_HEAD, raw=True)
        budget_values(ws)
        ws[f"C{B_FIRST}"].number_format = K0
        for label, a, b in merged_blocks:
            ws.merge_cells(f"A{a}:A{b}")
            ws[f"A{a}"].alignment = Alignment(horizontal="center", vertical="center")
        for _, a, b in groups:
            if a == b:
                ws[f"A{a}"].alignment = Alignment(horizontal="center", vertical="center")
        widths(ws, {"A": 18, "B": 26, "C": 17, "D": 17, "E": 17, "F": 13})
        ws.freeze_panes = f"A{B_FIRST}"

        # ------------------------------------------------------------------ Registry
        ws = wb.create_sheet("Registry")
        ws.sheet_properties.tabColor = "2E75B6"
        titled(ws, "Bluestone Memorial ED · registrations, 12/30–12/31/2025 (export)",
               "MRN and Phone arrived as plain numbers, so the leading zeros and punctuation are gone. "
               "PatientDue = the patient's share of this visit's claim, in dollars.")
        header(ws, R_HDR, ["Arrived", "MRN", "Patient", "Phone", "Payer", "PatientDue"])
        for i, r in enumerate(reg):
            row = R_FIRST + i
            ws.cell(row=row, column=1, value=r["Arrived"]).number_format = "mm/dd/yyyy hh:mm"
            ws.cell(row=row, column=2, value=r["MRN"])
            ws.cell(row=row, column=3, value=r["Patient"])
            ws.cell(row=row, column=4, value=r["Phone"])
            ws.cell(row=row, column=5, value=r["Payer"])
            ws.cell(row=row, column=6, value=r["PatientDue"])
        widths(ws, {"A": 17, "B": 12, "C": 22, "D": 17, "E": 30, "F": 14})
        ws.freeze_panes = f"A{R_FIRST}"

        # ------------------------------------------------------------------ Stays
        ws = wb.create_sheet("Stays")
        ws.sheet_properties.tabColor = "2E75B6"
        titled(ws, "Cardiac Step-Down (Bluestone Memorial) · inpatient discharges · December 2025",
               "Admitted lost its date format in the export. LOS = Discharged − Admitted, in days. Expected LOS is the diagnosis "
               "benchmark. LOS vs Expected = LOS − Expected LOS (positive = stayed longer than expected).")
        header(ws, S_HDR, ["EncounterID", "Diagnosis", "Admitted", "Discharged", "LOS (days)", "Expected LOS", "LOS vs Expected"])
        for i, s in enumerate(stays):
            row = S_FIRST + i
            ws.cell(row=row, column=1, value=s["EncounterID"])
            ws.cell(row=row, column=2, value=s["Diagnosis"])
            ws.cell(row=row, column=3, value=s["Admitted"])
            ws.cell(row=row, column=4, value=s["Discharged"]).number_format = "mm/dd/yyyy hh:mm"
            ws.cell(row=row, column=5, value=s["LOS"]).number_format = "0.0"
            ws.cell(row=row, column=6, value=s["Expected"]).number_format = "0.0"
            ws.cell(row=row, column=7, value=s["VsExpected"])
        widths(ws, {"A": 13, "B": 44, "C": 25, "D": 17, "E": 11, "F": 13, "G": 16})
        ws.freeze_panes = f"A{S_FIRST}"

        # ------------------------------------------------------------------ Dec Report (bonus, raw)
        ws = wb.create_sheet("Dec Report")
        ws.sheet_properties.tabColor = "BF9000"
        ws["A1"] = REPORT_TITLE
        ws["A2"] = REPORT_NOTE
        header(ws, P_HDR, RP_HEAD, raw=True)
        report_values(ws)
        widths(ws, {"A": 26, **{get_column_letter(j): 13 for j in range(2, 10)}})
        ws.freeze_panes = f"B{P_FIRST}"

        # ------------------------------------------------------------------ Budget Key (hidden): the finished report
        ws = wb.create_sheet("Budget Key")
        ws.sheet_properties.tabColor = "C00000"
        ws["A1"] = BUDGET_TITLE
        ws["A1"].style = "Title"
        for col in "ABCDEF":
            ws[f"{col}1"].alignment = Alignment(horizontal="centerContinuous")
        ws["A2"] = BUDGET_NOTE + "  ·  Finished version for task 13: compare with your Budget sheet."
        ws["A2"].font = NOTE_FONT
        for j, h in enumerate(BUDGET_HEAD, 1):
            ws.cell(row=B_HDR, column=j, value=h)
        budget_values(ws)
        style_header(ws, B_HDR, 6)
        grid(ws, f"A{B_HDR}:F{B_LAST}")
        for col in "ABCDEF":                             # the Total style replaces the grid borders on the total row
            ws[f"{col}{B_TOTAL}"].style = "Total"
        for row in range(B_FIRST, B_TOTAL + 1):          # number formats after the style (openpyxl styles reset them)
            for col in "CDE":
                ws[f"{col}{row}"].number_format = K0
            ws[f"F{row}"].number_format = PCT1
        widths(ws, {"A": 18, "B": 27, "C": 13, "D": 13, "E": 13, "F": 11})
        ws.sheet_state = "hidden"

        # ------------------------------------------------------------------ Report Key (hidden): the finished bonus
        ws = wb.create_sheet("Report Key")
        ws.sheet_properties.tabColor = "C00000"
        ws["A1"] = REPORT_TITLE
        ws["A1"].style = "Title"
        for j in range(1, 10):
            ws.cell(row=1, column=j).alignment = Alignment(horizontal="centerContinuous")
        ws["A2"] = REPORT_NOTE
        ws["A2"].font = NOTE_FONT
        for j, h in enumerate(RP_HEAD, 1):
            ws.cell(row=P_HDR, column=j, value=h)
        report_values(ws)
        style_header(ws, P_HDR, 9)
        for j in range(1, 10):
            ws.cell(row=P_TOTAL, column=j).style = "Total"
        codes = {"DecBudget": K1, "DecActual": K1, "DecVar": K1_RED, "DecVarPct": PCT1_RED,
                 "YtdBudget": MK, "YtdActual": MK, "YtdVar": K1_RED, "YtdVarPct": PCT1_RED}
        for row in range(P_FIRST, P_TOTAL + 1):
            for k, code in codes.items():
                ws[f"{rcol[k]}{row}"].number_format = code
        ws.cell(row=P_TOTAL + 2, column=1, value="Format codes used").font = Font(bold=True)
        for i, (label, code) in enumerate([("Dec Budget, Dec Actual", K1), ("Dec Variance, YTD Variance", K1_RED),
                                           ("Dec Var %, YTD Var %", PCT1_RED), ("YTD Budget, YTD Actual", MK)]):
            ws.cell(row=P_TOTAL + 3 + i, column=1, value=label)
            c = ws.cell(row=P_TOTAL + 3 + i, column=2, value=code)
            c.data_type = "s"
            c.font = Font(name="Consolas", size=10)
            c._style.quotePrefix = 1
        widths(ws, {"A": 27, **{get_column_letter(j): 12 for j in range(2, 10)}})
        ws.sheet_state = "hidden"

    return L
