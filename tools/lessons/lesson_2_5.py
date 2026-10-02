"""Lesson 2.5 · Conditional Counting & Summing.

Sheets
  Encounters  Every encounter that ended (was discharged) in 2025 for a fixed one-in-four sample of Bluestone patients
              (PatientID divisible by 4): 2,724 rows. Facility, department, service line, diagnosis description and
              category, and payer name are pre-joined so the lesson needs no lookups (those are Lesson 2.6).
              LOSDays = midnights between admission and discharge, filled only for Inpatient and Observation stays.
  Claims      The claims for those encounters that had been submitted by the course as-of date (12/31/2025). Payer name
              and encounter type are pre-joined. A yellow DaysToPay column is the bonus helper column.
  Payer Mix   (customize) payer x encounter-type grid the learner fills with ONE COUNTIFS formula (mixed references).
  Scorecard   (customize, bonus) inpatient claims scorecard by payer: adjudicated claims, denied-or-appealed claims,
              denial rate (with a divide-by-zero guard: Workers' Compensation has no inpatient claims), and average
              days to pay.

Every answer is computed in Python from data/*.csv.
"""
from __future__ import annotations

from collections import Counter
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from statistics import mean

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from xlcourse import Lesson, Task, data
from xlcourse.lesson import BOX, HEADER_FILL, INPUT_BORDER, INPUT_FILL, NAVY, PREFILL_FILL

CODE = "2.5"

SAMPLE_MOD = 4                      # PatientID number divisible by 4 -> about one patient in four
YEAR = 2025
AS_OF = data.AS_OF                  # 12/31/2025
TYPES = ["Emergency", "Inpatient", "Observation", "Outpatient"]

TITLE_FONT = Font(bold=True, size=13, color=NAVY)
NOTE_FONT = Font(italic=True, size=10, color="595959")
HDR_FONT = Font(bold=True, color="FFFFFF")
LABEL_FILL = PatternFill("solid", fgColor="D9E1F2")


def xround(x: float, n: int) -> float:
    """Round half away from zero, like Excel's ROUND (Python's round() is banker's rounding)."""
    q = Decimal(1).scaleb(-n)
    d = Decimal(repr(x)).quantize(q, rounding=ROUND_HALF_UP)
    return float(d)


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="02-formulas-functions", slug="05-conditional-aggregation",
        title="Conditional Counting & Summing", level="Beginner → Intermediate", minutes=55,
        objectives=[
            "Count and sum with conditions using COUNTIF(S), SUMIF(S), and AVERAGEIF(S)",
            "Find conditional extremes with MAXIFS and MINIFS",
            "Write criteria with operators, cell references, wildcards, and date ranges",
            "Build summary grids that fill with a single formula",
        ],
        data_note="2,724 encounters that ended in 2025 across all four Bluestone Health System facilities (every 2025 "
                  "encounter for a fixed sample of about 900 patients), with facility, department, service line, "
                  "diagnosis, payer, charges, and length of stay. The 2,666 claims for those encounters that had been "
                  "submitted by 12/31/2025, with status, denial reason, and payment dates.",
    )

    # ================================================================== source data
    fac = data.index(data.load("facilities"), "FacilityID")
    dept = data.index(data.load("departments"), "DeptID")
    payers = data.index(data.load("payers"), "PayerID")
    dx = data.index(data.load("diagnoses"), "DxCode")
    payer_names = [payers[p]["PayerName"] for p in sorted(payers)]          # PY01..PY08 order

    enc_rows = []
    for e in data.load("encounters"):
        if e["DischargeDateTime"].year != YEAR or int(e["PatientID"][2:]) % SAMPLE_MOD:
            continue
        d = dept[e["DeptID"]]
        stay = e["EncounterType"] in ("Inpatient", "Observation")
        enc_rows.append({
            "EncounterID": e["EncounterID"], "PatientID": e["PatientID"], "EncounterType": e["EncounterType"],
            "FacilityName": fac[e["FacilityID"]]["FacilityName"], "DeptName": d["DeptName"], "ServiceLine": d["ServiceLine"],
            "AdmitDate": e["AdmitDateTime"].date(), "DischargeDate": e["DischargeDateTime"].date(),
            "LOSDays": (e["DischargeDateTime"].date() - e["AdmitDateTime"].date()).days if stay else None,
            "PrimaryDxCode": e["PrimaryDxCode"], "DxDescription": dx[e["PrimaryDxCode"]]["DxDescription"],
            "DxCategory": dx[e["PrimaryDxCode"]]["DxCategory"], "DischargeDisposition": e["DischargeDisposition"],
            "PayerName": payers[e["PayerID"]]["PayerName"], "TotalCharges": e["TotalCharges"], "Readmit30": e["Readmit30"],
        })
    enc_rows.sort(key=lambda r: r["EncounterID"])
    etype = {r["EncounterID"]: r["EncounterType"] for r in enc_rows}

    claim_rows = []
    for c in data.load("claims"):
        if c["EncounterID"] not in etype or c["SubmitDate"] > AS_OF:
            continue        # not yet billed on the as-of date
        claim_rows.append({
            "ClaimID": c["ClaimID"], "EncounterID": c["EncounterID"], "PayerName": payers[c["PayerID"]]["PayerName"],
            "EncounterType": etype[c["EncounterID"]], "ServiceDate": c["ServiceDate"], "SubmitDate": c["SubmitDate"],
            "BilledAmount": c["BilledAmount"], "PaidAmount": c["PaidAmount"], "ClaimStatus": c["ClaimStatus"],
            "DenialReason": c["DenialReason"], "PaidDate": c["PaidDate"],
        })
    claim_rows.sort(key=lambda r: r["ClaimID"])
    unbilled = len(enc_rows) - len(claim_rows)

    enc = L.add_table_sheet(
        "Encounters", enc_rows, table="tblEncounters",
        columns=["EncounterID", "PatientID", "EncounterType", "FacilityName", "DeptName", "ServiceLine", "AdmitDate",
                 "DischargeDate", "LOSDays", "PrimaryDxCode", "DxDescription", "DxCategory", "DischargeDisposition",
                 "PayerName", "TotalCharges", "Readmit30"],
        formats={"TotalCharges": "#,##0.00", "LOSDays": "0"},
        widths={"FacilityName": 30, "DeptName": 26, "DxDescription": 44, "DischargeDisposition": 26, "PayerName": 29,
                "TotalCharges": 13, "ServiceLine": 18},
    )
    clm = L.add_table_sheet(
        "Claims", claim_rows, table="tblClaims",
        columns=["ClaimID", "EncounterID", "PayerName", "EncounterType", "ServiceDate", "SubmitDate", "BilledAmount",
                 "PaidAmount", "ClaimStatus", "DenialReason", "PaidDate"],
        extra_cols=["DaysToPay"],
        formats={"BilledAmount": "#,##0.00", "PaidAmount": "#,##0.00", "DaysToPay": "0"},
        widths={"PayerName": 29, "DenialReason": 24, "DaysToPay": 11, "BilledAmount": 13},
    )
    assert [enc.col(h) for h in ("EncounterType", "FacilityName", "LOSDays", "PrimaryDxCode", "PayerName", "TotalCharges",
                                 "Readmit30")] == ["C", "D", "I", "J", "N", "O", "P"]
    assert [clm.col(h) for h in ("PayerName", "EncounterType", "SubmitDate", "BilledAmount", "ClaimStatus", "DenialReason",
                                 "PaidDate", "DaysToPay")] == ["C", "D", "F", "G", "I", "J", "K", "L"]

    def _rng(sd, col, absolute):
        # Sheet!C2:C2725 the way a learner types it (no quotes: these sheet names have no spaces).
        c = sd.col(col)
        return f"{sd.name}!${c}${sd.first_row}:${c}${sd.last_row}" if absolute else f"{sd.name}!{c}{sd.first_row}:{c}{sd.last_row}"

    def E(col, absolute=False):
        return _rng(enc, col, absolute)

    def C(col, absolute=False):
        return _rng(clm, col, absolute)

    # ================================================================== practice answers (Python)
    n_ed = sum(1 for r in enc_rows if r["EncounterType"] == "Emergency")
    cedar_ip = sum(1 for r in enc_rows if r["EncounterType"] == "Inpatient" and r["FacilityName"] == "Cedar Ridge Medical Center")
    medicare_chg = sum(r["TotalCharges"] for r in enc_rows if r["PayerName"] == "Medicare")
    medicare_like = sum(r["TotalCharges"] for r in enc_rows if "Medicare" in r["PayerName"])
    auth = [c for c in claim_rows if c["ClaimStatus"] == "Denied" and c["DenialReason"] == "Authorization Required"]
    auth_billed = sum(c["BilledAmount"] for c in auth)
    auth_any = sum(1 for c in claim_rows if c["DenialReason"] == "Authorization Required")
    # What happened to the Authorization Required claims that are no longer plain Denied (for the explanation).
    auth_other = Counter(c["ClaimStatus"] for c in claim_rows
                         if c["DenialReason"] == "Authorization Required" and c["ClaimStatus"] != "Denied")
    _status_words = {"Appealed": "under appeal", "Partially Paid": "partially paid", "Paid": "paid", "Pending": "pending"}
    assert sum(auth_other.values()) == auth_any - len(auth) > 0
    if len(auth_other) == 1:
        (_s, _n), = auth_other.items()
        auth_other_txt = f"The other {_n} are all {_status_words[_s]} (ClaimStatus = {_s})"
    else:
        auth_other_txt = f"Of the other {auth_any - len(auth)}, " + " and ".join(
            f"{n} {'is' if n == 1 else 'are'} {_status_words[s]}" for s, n in sorted(auth_other.items(), key=lambda kv: -kv[1]))
    hf = [r["LOSDays"] for r in enc_rows if r["EncounterType"] == "Inpatient" and r["PrimaryDxCode"] == "I50.9"]
    hf_all = [r["LOSDays"] for r in enc_rows if r["PrimaryDxCode"] == "I50.9" and r["LOSDays"] is not None]
    hf_alos = xround(mean(hf), 1)
    assert xround(mean(hf_all), 1) != hf_alos, "the EncounterType criterion should change the answer"
    q3_start, q3_end = date(YEAR, 7, 1), date(YEAR, 9, 30)
    q3 = sum(1 for c in claim_rows if q3_start <= c["SubmitDate"] <= q3_end)
    circ = sum(1 for r in enc_rows if r["PrimaryDxCode"].startswith("I"))
    assert circ == sum(1 for r in enc_rows if r["DxCategory"] == "Circulatory")
    ip = [r for r in enc_rows if r["EncounterType"] == "Inpatient"]
    ip_avg = mean(r["TotalCharges"] for r in ip)
    ip_above = sum(1 for r in ip if r["TotalCharges"] > ip_avg)
    bmh_stays = sum(1 for r in enc_rows if r["FacilityName"] == "Bluestone Memorial Hospital"
                    and r["EncounterType"] in ("Inpatient", "Observation"))
    sepsis_max = max(r["TotalCharges"] for r in enc_rows if "sepsis" in r["DxDescription"].lower())
    pend_ins = [c["SubmitDate"] for c in claim_rows if c["ClaimStatus"] == "Pending" and c["PayerName"] != "Self-Pay"]
    pend_all = [c["SubmitDate"] for c in claim_rows if c["ClaimStatus"] == "Pending"]
    oldest_pending = min(pend_ins)
    assert min(pend_all) != oldest_pending, "the <>Self-Pay criterion should change the answer"
    grid = Counter((r["PayerName"], r["EncounterType"]) for r in enc_rows)
    grid_total = sum(grid[(p, t)] for p in payer_names for t in TYPES)
    assert grid_total == len(enc_rows)
    cv = [r for r in enc_rows if r["ServiceLine"] == "Cardiovascular" and r["EncounterType"] == "Inpatient"]
    cv_y = sum(1 for r in cv if r["Readmit30"] == "Y")
    cv_rate = cv_y / len(cv)
    cv_all = sum(1 for r in enc_rows if r["ServiceLine"] == "Cardiovascular")
    assert cv_all > len(cv), "the Cardiology Clinic visits make the Inpatient criterion matter"

    e_first, e_last = enc.first_row, enc.last_row
    c_first, c_last = clm.first_row, clm.last_row

    # Payer Mix sheet layout (built in the customize hook)
    PM = "'Payer Mix'"
    PM_HDR, PM_FIRST = 4, 5
    PM_LAST = PM_FIRST + len(payer_names) - 1                 # 12
    PM_C0, PM_C1 = "B", get_column_letter(1 + len(TYPES))     # B..E
    pm_rng = f"{PM}!{PM_C0}{PM_FIRST}:{PM_C1}{PM_LAST}"
    grid_formula = (f"=COUNTIFS({E('PayerName', True)},$A{PM_FIRST},"
                    f"{E('EncounterType', True)},{PM_C0}${PM_HDR})")

    L.practice_intro = (
        f"The Encounters sheet holds rows {e_first}–{e_last} and the Claims sheet holds rows {c_first}–{c_last}. Type text "
        "criteria exactly as the task spells them. Upper and lower case don't matter, but every other character does. "
        "Task 12 is filled on the Payer Mix sheet."
    )

    L.tasks = [
        Task("How many encounters were Emergency visits (EncounterType = Emergency)?",
             answer=n_ed, solution=f'=COUNTIF({E("EncounterType")},"Emergency")',
             hint="COUNTIF(range, criteria)", title="Emergency visits",
             explanation="COUNTIF checks every cell in the range against one criterion and counts the matches. A text criterion "
                         "goes in double quotes and must match the whole cell, but upper and lower case don't matter, so "
                         "\"emergency\" works too."),
        Task("How many Inpatient encounters were at Cedar Ridge Medical Center?",
             answer=cedar_ip,
             solution=f'=COUNTIFS({E("EncounterType")},"Inpatient",{E("FacilityName")},"Cedar Ridge Medical Center")',
             hint="COUNTIFS takes range/criteria pairs, and a row must pass every pair",
             title="Inpatient stays at Cedar Ridge",
             explanation="COUNTIFS joins its conditions with AND: a row counts only when EncounterType is Inpatient *and* "
                         "FacilityName is Cedar Ridge Medical Center. The order of the pairs doesn't matter."),
        Task("What were the total charges (TotalCharges) of encounters billed to Medicare? Use PayerName = Medicare exactly, "
             "because Silverline Medicare Advantage is a separate payer.",
             answer=round(medicare_chg, 2), fmt="#,##0.00",
             solution=f'=SUMIF({E("PayerName")},"Medicare",{E("TotalCharges")})',
             hint="SUMIF(range, criteria, sum_range)", title="Total charges billed to Medicare",
             explanation="SUMIF tests one column (PayerName) and adds the matching rows of another (TotalCharges). Without "
                         "wildcards, \"Medicare\" matches only cells that say exactly Medicare. The criterion `\"*Medicare*\"` "
                         f"would also pick up Silverline Medicare Advantage and return {medicare_like:,.2f}."),
        Task("How much was billed (BilledAmount) on claims with ClaimStatus = Denied and DenialReason = Authorization Required? "
             "Use the Claims sheet.",
             answer=round(auth_billed, 2), fmt="#,##0.00",
             solution=(f'=SUMIFS({C("BilledAmount")},{C("ClaimStatus")},"Denied",'
                       f'{C("DenialReason")},"Authorization Required")'),
             hint="SUMIFS puts the sum_range FIRST", title="Billed dollars denied for missing authorization",
             explanation=f"SUMIFS starts with the column to add, then lists range/criteria pairs. The status criterion matters: "
                         f"{auth_any} claims carry the reason Authorization Required, but only {len(auth)} of them are still "
                         f"plain Denied. {auth_other_txt}, so they don't belong in a Denied total."),
        Task("What was the average length of stay (LOSDays) of Inpatient encounters with PrimaryDxCode I50.9 (heart failure)? "
             "Round to 1 decimal place with ROUND.",
             answer=hf_alos, fmt="0.0", tol=0.0001,
             solution=(f'=ROUND(AVERAGEIFS({E("LOSDays")},{E("EncounterType")},"Inpatient",'
                       f'{E("PrimaryDxCode")},"I50.9"),1)'),
             hint="AVERAGEIFS(average_range, criteria_range1, criteria1, …), wrapped in ROUND",
             title="Heart failure average length of stay",
             explanation=f"AVERAGEIFS averages only the rows that pass every condition. Leave out the Inpatient criterion and the "
                         f"heart failure Observation stays join the average, which drops to {xround(mean(hf_all), 1)}. "
                         f"Bluestone's benchmark (ExpectedLOS) for I50.9 is {dx['I50.9']['ExpectedLOS']} days, so these "
                         "stays run long."),
        Task("How many claims were submitted in Q3 2025 (SubmitDate from 07/01/2025 through 09/30/2025)? Build the dates with DATE.",
             answer=q3,
             solution=(f'=COUNTIFS({C("SubmitDate")},">="&DATE(2025,7,1),'
                       f'{C("SubmitDate")},"<="&DATE(2025,9,30))'),
             hint="Use the same column twice, once for the start and once for the end",
             title="Claims submitted in Q3 2025",
             explanation="A date range needs two conditions on the same column. `\">=\"&DATE(2025,7,1)` joins the operator to the "
                         "date's serial number, so the criterion works in any regional date setting. Because SubmitDate holds "
                         "whole dates, `\"<=\"&DATE(2025,9,30)` is safe here. `\"<\"&DATE(2025,10,1)` gives the same answer and "
                         "also works on columns that include times."),
        Task("How many encounters had a primary diagnosis code in the circulatory chapter of ICD-10, meaning a PrimaryDxCode "
             "that starts with the letter I?",
             answer=circ, solution=f'=COUNTIF({E("PrimaryDxCode")},"I*")',
             hint="An asterisk matches any number of characters", title="Circulatory diagnoses (I codes)",
             explanation="The asterisk stands for \"anything, or nothing,\" so `\"I*\"` matches I10, I21.4, I50.9, and every other "
                         "code that starts with I. Wildcards only work on text, which is fine here because ICD-10 codes are "
                         f"text. Cross-check: COUNTIF on DxCategory = Circulatory also returns {circ}."),
        Task("How many Inpatient encounters had TotalCharges above the average TotalCharges of all Inpatient encounters? "
             "Calculate the average inside the formula rather than typing it.",
             answer=ip_above,
             solution=(f'=COUNTIFS({E("EncounterType")},"Inpatient",{E("TotalCharges")},'
                       f'">"&AVERAGEIF({E("EncounterType")},"Inpatient",{E("TotalCharges")}))'),
             hint="Join an operator to a function: \">\"&AVERAGEIF(…)",
             title="Inpatient stays above the average inpatient charge",
             explanation=f"AVERAGEIF works out the average Inpatient charge ({ip_avg:,.2f}), and `\">\"&` turns it into the "
                         f"criterion text \">{ip_avg:.2f}…\". Only {ip_above} of {len(ip)} stays ({ip_above / len(ip):.0%}) "
                         "beat the average, because a few very expensive stays pull the mean up (Lesson 2.4)."),
        Task("How many encounters at Bluestone Memorial Hospital were hospital stays, meaning EncounterType is Inpatient OR "
             "Observation?",
             answer=bmh_stays,
             solution=(f'=SUM(COUNTIFS({E("FacilityName")},"Bluestone Memorial Hospital",'
                       f'{E("EncounterType")},{{"Inpatient","Observation"}}))'),
             hint="COUNTIFS joins conditions with AND. For OR, add two counts together",
             title="Hospital stays at Bluestone Memorial (OR logic)",
             explanation="One COUNTIFS can't say OR, because every row would need to be Inpatient and Observation at once. Add "
                         "two counts instead: `=COUNTIFS(…,\"Inpatient\")+COUNTIFS(…,\"Observation\")`. The SUM version does the "
                         "same thing in one formula. The array constant {\"Inpatient\",\"Observation\"} makes COUNTIFS return "
                         "two counts, and SUM adds them. This is safe because no row can match both values."),
        Task("What was the highest TotalCharges for an encounter whose DxDescription contains the word sepsis?",
             answer=sepsis_max, fmt="#,##0.00",
             solution=f'=MAXIFS({E("TotalCharges")},{E("DxDescription")},"*sepsis*")',
             hint="MAXIFS(max_range, criteria_range1, criteria1, …). \"Contains\" needs an asterisk on both sides",
             title="Most expensive sepsis encounter",
             explanation="MAXIFS returns the largest value among the rows that pass every condition. `\"*sepsis*\"` means "
                         "\"sepsis with anything before or after it,\" which is how you write *contains* in a criterion. "
                         "`\"sepsis\"` alone would only match a cell holding exactly that one word."),
        Task("Collections wants the oldest claim still waiting on an insurance company. What is the earliest SubmitDate among "
             "claims with ClaimStatus = Pending and a PayerName other than Self-Pay? Enter it as a date.",
             answer=oldest_pending, fmt="mm/dd/yyyy",
             solution=(f'=MINIFS({C("SubmitDate")},{C("ClaimStatus")},"Pending",'
                       f'{C("PayerName")},"<>Self-Pay")'),
             hint="MINIFS works on dates too, and \"<>\" means not equal to",
             title="Oldest pending insurance claim",
             explanation=f"Dates are numbers, so the smallest SubmitDate is the oldest. `\"<>Self-Pay\"` keeps every payer except "
                         f"Self-Pay. Without it, the answer would be {min(pend_all):%m/%d/%Y}: a self-pay balance the patient "
                         "still owes, which belongs to a different work queue. If the cell shows a number like "
                         f"{int(data.excel_serial(oldest_pending))}, format it as a date."),
        Task(f"On the Payer Mix sheet, fill the yellow grid {PM_C0}{PM_FIRST}:{PM_C1}{PM_LAST} with ONE COUNTIFS formula: the "
             f"number of encounters for the payer in column A and the encounter type in row {PM_HDR}. Type it in "
             f"{PM_C0}{PM_FIRST}, then copy it across and down. Type the data ranges as cell addresses with $ signs instead "
             "of selecting Table columns (guide section 11 explains why). The gray cell adds up your grid.",
             answer=grid_total, title="Payer mix grid filled with one formula",
             solution=grid_formula,
             summary=f'=IF(COUNT({pm_rng})=0,"",SUM({pm_rng}))',
             fill={"range": pm_rng, "formula": grid_formula},
             live=f'=COUNTA({E("EncounterID")})',
             hint="Lock the data ranges fully. For each label, lock only the part (row or column) that must stay put",
             explanation=f"The data ranges never move, so they get full $ signs. `$A{PM_FIRST}` lets the row change but always "
                         f"reads column A, and `{PM_C0}${PM_HDR}` lets the column change but always reads row {PM_HDR}. Every "
                         "encounter has exactly one payer and one type, so a correct grid adds up to the number of rows on the "
                         "Encounters sheet. Each row also matches the gray COUNTIF check in column G."),
        Task("What was the 30-day readmission rate for the Cardiovascular service line? Divide the Cardiovascular Inpatient "
             "stays with Readmit30 = Y by all Cardiovascular Inpatient stays. Enter it as a percentage.",
             answer=cv_rate, fmt="0.0%",
             solution=(f'=COUNTIFS({E("ServiceLine")},"Cardiovascular",{E("EncounterType")},"Inpatient",'
                       f'{E("Readmit30")},"Y")/COUNTIFS({E("ServiceLine")},"Cardiovascular",{E("EncounterType")},"Inpatient")'),
             hint="Rate = COUNTIFS(numerator) / COUNTIFS(denominator). Same filters, plus one more on top",
             title="Cardiovascular 30-day readmission rate",
             explanation=f"The numerator is the denominator plus one extra condition (Readmit30 = Y). That's the pattern for "
                         f"every rate. The Inpatient criterion matters in both: the Cardiovascular service line also includes "
                         f"Cardiology Clinic visits, so `COUNTIF(ServiceLine,\"Cardiovascular\")` alone returns {cv_all} instead "
                         f"of {len(cv)}. Formal CMS measures also leave out some stays, such as deaths, transfers, and planned "
                         "readmissions. This is the simple internal version."),
    ]

    # ================================================================== bonus: inpatient claims scorecard by payer
    paid = [c for c in claim_rows if c["PaidDate"] is not None]
    avg_dtp = mean((c["PaidDate"] - c["SubmitDate"]).days for c in paid)

    SC = "Scorecard"
    SC_HDR, SC_FIRST = 4, 5
    SC_LAST = SC_FIRST + len(payer_names) - 1                 # 12
    row_of = {p: SC_FIRST + i for i, p in enumerate(payer_names)}
    stats = {}
    for p in payer_names:
        cl = [c for c in claim_rows if c["PayerName"] == p and c["EncounterType"] == "Inpatient"]
        adj = [c for c in cl if c["ClaimStatus"] != "Pending"]
        den = [c for c in adj if c["ClaimStatus"] in ("Denied", "Appealed")]
        dtp = [(c["PaidDate"] - c["SubmitDate"]).days for c in cl if c["PaidDate"] is not None]
        stats[p] = dict(adj=len(adj), den=len(den), rate=len(den) / len(adj) if adj else None,
                        dtp=mean(dtp) if dtp else None, at_risk=sum(c["BilledAmount"] for c in den))
    no_claims = [p for p in payer_names if stats[p]["adj"] == 0]
    assert no_claims == ["Workers' Compensation"], no_claims          # the divide-by-zero row
    ranked = sorted((p for p in payer_names if stats[p]["rate"] is not None), key=lambda p: stats[p]["rate"], reverse=True)
    worst, runner = ranked[0], ranked[1]
    assert stats[worst]["rate"] - stats[runner]["rate"] > 0.05, "B3 needs a clear winner"
    assert worst == "Silverline Medicare Advantage"
    worst_rate = stats[worst]["rate"]
    worst_at_risk = round(stats[worst]["at_risk"], 2)
    dtp_gap = xround(stats[worst]["dtp"] - stats["Medicare"]["dtp"], 1)
    assert 3.5 < worst_rate / stats["Medicare"]["rate"] < 4.5 and 12.5 < dtp_gap < 15.5, "B5's 'four times / two weeks' wording"
    # The two denial-rate shortcuts a learner might take must not change the winner silently: report them in the key.
    denied_only = {p: (sum(1 for c in claim_rows if c["PayerName"] == p and c["EncounterType"] == "Inpatient"
                           and c["ClaimStatus"] == "Denied") / stats[p]["adj"]) for p in ranked}

    def cr(col):          # absolute Claims range for the scorecard formulas
        return C(col, True)

    sc_adj = f'=COUNTIFS({cr("PayerName")},$A{SC_FIRST},{cr("EncounterType")},"Inpatient",{cr("ClaimStatus")},"<>Pending")'
    sc_den = (f'=SUM(COUNTIFS({cr("PayerName")},$A{SC_FIRST},{cr("EncounterType")},"Inpatient",'
              f'{cr("ClaimStatus")},{{"Denied","Appealed"}}))')
    sc_rate = f'=IF(B{SC_FIRST}=0,"n/a",C{SC_FIRST}/B{SC_FIRST})'
    sc_dtp = f'=IFERROR(AVERAGEIFS({cr("DaysToPay")},{cr("PayerName")},$A{SC_FIRST},{cr("EncounterType")},"Inpatient"),"n/a")'
    dtp_fill = f'=IF({clm.col("PaidDate")}{c_first}="","",{clm.col("PaidDate")}{c_first}-{clm.col("SubmitDate")}{c_first})'
    dtp_rng = f"Claims!{clm.col('DaysToPay')}{c_first}:{clm.col('DaysToPay')}{c_last}"
    rate_rng = f"{SC}!D{SC_FIRST}:D{SC_LAST}"

    def rate_formula(p):
        crit = f'{C("PayerName")},"{p}",{C("EncounterType")},"Inpatient"'
        return (f'SUM(COUNTIFS({crit},{C("ClaimStatus")},{{"Denied","Appealed"}}))'
                f'/COUNTIFS({crit},{C("ClaimStatus")},"<>Pending")')

    def avg_dtp_nohelper(p):
        crit = f'{C("PayerName")},"{p}",{C("EncounterType")},"Inpatient",{C("PaidDate")},"<>"'
        return (f'(SUMIFS({C("PaidDate")},{crit})-SUMIFS({C("SubmitDate")},{crit}))'
                f'/COUNTIFS({crit})')

    def avg_dtp_helper(p):
        return f'AVERAGEIFS({C("DaysToPay")},{C("PayerName")},"{p}",{C("EncounterType")},"Inpatient")'

    L.bonus_title = "Bonus: Inpatient payer scorecard"
    L.bonus_scenario = (
        "Bluestone's revenue cycle director is preparing for contract talks and wants a 2025 scorecard of inpatient claims by "
        "payer. Use these definitions. An **adjudicated** claim is any claim whose ClaimStatus is not Pending (the payer has "
        "made a decision). A **denied** claim is one whose ClaimStatus is Denied or Appealed, because an appealed claim was "
        "denied first. **Denial rate** = denied ÷ adjudicated. **Days to pay** = PaidDate − SubmitDate for every claim that "
        "has a PaidDate. Build the scorecard on the Scorecard sheet, one row per payer, using only claims whose "
        "EncounterType is Inpatient."
    )
    L.bonus = [
        Task(f"On the Claims sheet, fill the yellow DaysToPay column ({clm.col('DaysToPay')}{c_first}:"
             f"{clm.col('DaysToPay')}{c_last}) with PaidDate − SubmitDate, or an empty text string (\"\") when PaidDate is "
             f"blank. Type one formula in {clm.col('DaysToPay')}{c_first}. Claims is a Table, so Excel fills it down the whole "
             "column for you. If you click cells instead of typing their addresses, Excel writes [@PaidDate], which means "
             "the PaidDate on the same row. The gray cell averages your column across all claim types.",
             answer=avg_dtp, fmt="0.00", title="DaysToPay helper column",
             solution=dtp_fill,
             summary=f'=IF(COUNT({dtp_rng})=0,"",AVERAGE({dtp_rng}))',
             fill={"range": dtp_rng, "formula": dtp_fill},
             live=(f'=(SUM({C("PaidDate")})-SUMIFS({C("SubmitDate")},{C("PaidDate")},"<>"))'
                   f'/COUNT({C("PaidDate")})'),
             hint="IF(PaidDate=\"\", \"\", PaidDate − SubmitDate)",
             explanation=f"The -IFS functions can only test and average columns that already exist. None of them can subtract "
                         f"two columns row by row, so a **helper column** does the subtraction first, and then AVERAGEIFS can "
                         f"average it by payer. Without the IF, an unpaid claim computes 0 − SubmitDate, "
                         f"a large negative number of days that wrecks every average. The empty text \"\" is "
                         f"skipped by AVERAGE and AVERAGEIFS. The key's live formula shows a no-helper version: the average of "
                         f"the differences is (sum of PaidDates − sum of their SubmitDates) ÷ the number of paid claims."),
        Task(f"On the Scorecard sheet, fill columns B (adjudicated inpatient claims), C (denied inpatient claims), and D "
             f"(denial rate) for all {len(payer_names)} payers, one formula per column copied down. A payer with no "
             "adjudicated inpatient claims must show n/a in column D instead of #DIV/0!. The gray cell shows the highest "
             "rate in your column D. (If it shows #DIV/0!, a row still needs the guard.)",
             answer=worst_rate, fmt="0.0%", title="Denial rates (with a divide-by-zero guard)",
             solution=(f"B{SC_FIRST}: `{sc_adj}`\n\nC{SC_FIRST}: `{sc_den}`\n\nD{SC_FIRST}: `{sc_rate}`\n\n"
                       f"Copy each one down to row {SC_LAST}."),
             solution_lang="markdown",
             summary=f'=IF(COUNTA({rate_rng})=0,"",MAX({rate_rng}))',
             fill={"range": rate_rng, "formula": sc_rate},
             live="=" + rate_formula(worst),
             hint="Column B: \"<>Pending\". Column C: the OR pattern from guide section 4. Column D: test column B for 0 "
                  "before you divide",
             explanation=f"Column B uses `\"<>Pending\"` to count every decision. Column C needs OR logic (Denied or Appealed), "
                         f"so it wraps COUNTIFS with an array constant in SUM. `$A{SC_FIRST}` keeps each formula reading its own "
                         f"payer as you copy down. Workers' Compensation had no inpatient claims, so its row divides 0 by 0. "
                         f"`IF(B{SC_FIRST}=0,\"n/a\",…)` tests for exactly that case. `IFERROR(C{SC_FIRST}/B{SC_FIRST},\"n/a\")` also "
                         f"works, but it would hide any other mistake too. Counting only Denied (not Appealed) would understate "
                         f"the worst payer's rate at {denied_only[worst]:.1%}."),
        Task("Which payer has the highest inpatient denial rate? Type its PayerName.",
             answer=worst, accept=["Silverline", "Silverline MA"],
             solution="Read it from column D of your scorecard: the largest rate belongs to **" + worst + "**.",
             live=False, hint="Compare the rates in column D", title="Payer with the highest denial rate",
             explanation=f"{worst} denied {worst_rate:.1%} of its adjudicated inpatient claims. Next is {runner} at "
                         f"{stats[runner]['rate']:.1%}, and traditional Medicare is at {stats['Medicare']['rate']:.1%}. A formula "
                         f"can find the name for you too: `=INDEX(A{SC_FIRST}:A{SC_LAST},MATCH(MAX(D{SC_FIRST}:D{SC_LAST}),"
                         f"D{SC_FIRST}:D{SC_LAST},0))` on the Scorecard sheet. Lesson 2.6 teaches INDEX and MATCH."),
        Task("Dollars at risk: for the payer you named in task B3, what is the total BilledAmount of its inpatient claims "
             "that are Denied or Appealed?",
             answer=worst_at_risk, fmt="#,##0.00",
             solution=(f'=SUM(SUMIFS({C("BilledAmount")},{C("PayerName")},"{worst}",{C("EncounterType")},"Inpatient",'
                       f'{C("ClaimStatus")},{{"Denied","Appealed"}}))'),
             hint="The same OR trick works with SUMIFS", title="Billed dollars at risk with the worst payer",
             explanation=f"SUMIFS with the array constant returns two totals (Denied and Appealed), and SUM adds them. "
                         f"About {worst_at_risk / 1e6:.1f} million dollars of inpatient charges are tied up in denials with this "
                         "one payer, which is the number that gets a contract meeting's attention."),
        Task("Fill column E of the Scorecard with each payer's average DaysToPay for inpatient claims (show n/a when a payer "
             "has none). How many more days, on average, does the payer you named in task B3 take to pay an inpatient claim than "
             "Medicare does? Subtract the full-precision averages (point at the cells instead of retyping the 1-decimal "
             "values that column E displays), then round the difference to 1 decimal place.",
             answer=dtp_gap, fmt="0.0", tol=0.0001, title="How much slower the worst payer pays",
             solution=f"=ROUND({avg_dtp_helper(worst)}-{avg_dtp_helper('Medicare')},1)",
             live=f"=ROUND({avg_dtp_nohelper(worst)}-{avg_dtp_nohelper('Medicare')},1)",
             hint="AVERAGEIFS on your DaysToPay column, wrapped in IFERROR for the payer with no claims",
             explanation=f"E{SC_FIRST} is `{sc_dtp}`, copied down. AVERAGEIFS skips the \"\" cells of unpaid claims, and "
                         f"IFERROR turns Workers' Compensation's #DIV/0! into n/a. Here IFERROR is the right guard, because "
                         f"AVERAGEIFS has no separate denominator you could test. {worst} averages "
                         f"{stats[worst]['dtp']:.2f} days and Medicare {stats['Medicare']['dtp']:.2f}, so "
                         f"`=ROUND({SC}!E{row_of[worst]}-{SC}!E{row_of['Medicare']},1)` in the answer cell gives the same "
                         f"{dtp_gap}. "
                         f"Compared with traditional Medicare, {worst} denies about four times as often *and* takes about "
                         "two weeks longer to pay."),
    ]

    L.start_notes = [
        "Encounters and Claims are Excel Tables. If you select a whole Table column while writing a formula, Excel may write "
        "it as tblEncounters[EncounterType] instead of Encounters!C2:C2725. Both give the same result.",
        f"{unbilled} encounters from 2025 had not been billed by 12/31/2025, so the Claims sheet has {len(claim_rows):,} rows.",
        "Practice task 12 is filled on the Payer Mix sheet, and the bonus scorecard on the Scorecard sheet. Their gray "
        "cells are pre-filled totals and checks.",
    ]
    L.sheet_order = ["Start Here", "Practice", "Encounters", "Claims", "Payer Mix", "Bonus", SC, "Answer Key", "Bonus Key"]

    # ================================================================== extra sheets
    def titled(ws, title, *notes, cols, width):
        """Title in A1 and wrapped notes in rows 2.., each merged across the table's `cols` columns
        (`width` = about how many characters of note text fit on one line of the merge)."""
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        for i, n in enumerate(notes, 2):
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=cols)
            c = ws.cell(row=i, column=1, value=n)
            c.font = NOTE_FONT
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[i].height = 14 * max(1, -(-len(n) // width)) + 4
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    def header(ws, row, labels):
        for j, h in enumerate(labels, 1):
            c = ws.cell(row=row, column=j, value=h)
            c.font = HDR_FONT
            c.fill = HEADER_FILL
            c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 32

    def yellow(ws, rng, fmt):
        for row in ws[rng]:
            for c in row:
                c.fill = INPUT_FILL
                c.border = INPUT_BORDER
                c.number_format = fmt
                c.alignment = Alignment(horizontal="right")

    def gray(ws, coord, formula, fmt="#,##0"):
        L.set_formula(ws, coord, formula, dynamic=False)
        c = ws[coord]
        c.fill = PREFILL_FILL
        c.font = Font(italic=True, color="404040")
        c.border = BOX
        c.number_format = fmt

    def labels(ws, first, names):
        for i, p in enumerate(names):
            c = ws.cell(row=first + i, column=1, value=p)
            c.font = Font(bold=True)
            c.fill = LABEL_FILL
            c.border = BOX

    @L.customize
    def _sheets(wb, lesson, selftest):
        # ---------------------------------------------------------- Payer Mix
        ws = wb.create_sheet("Payer Mix")
        ws.sheet_properties.tabColor = "548235"
        titled(ws, "Payer mix · 2025 encounters by payer and encounter type",
               f"Fill the yellow grid {PM_C0}{PM_FIRST}:{PM_C1}{PM_LAST} with ONE COUNTIFS formula typed in {PM_C0}{PM_FIRST} "
               "(practice task 12), then copy it across and down. Each cell counts the Encounters rows whose PayerName matches "
               f"column A and whose EncounterType matches row {PM_HDR}. Type the data ranges as cell addresses with $ signs, "
               "not as Table references.",
               "Gray cells are already filled in. Column F totals your row, and column G counts the payer's encounters with "
               "COUNTIF, so F and G match when a row is right.", cols=3 + len(TYPES), width=105)
        header(ws, PM_HDR, ["PayerName ↓   EncounterType →"] + TYPES + ["Row total", "COUNTIF check"])
        labels(ws, PM_FIRST, payer_names)
        yellow(ws, pm_rng.split("!")[1], "#,##0")
        tc = get_column_letter(2 + len(TYPES))           # F
        gc = get_column_letter(3 + len(TYPES))           # G
        for r in range(PM_FIRST, PM_LAST + 1):
            gray(ws, f"{tc}{r}", f"=SUM({PM_C0}{r}:{PM_C1}{r})")
            gray(ws, f"{gc}{r}", f"=COUNTIF({E('PayerName', True)},$A{r})")
        tot = PM_LAST + 1
        c = ws.cell(row=tot, column=1, value="Total")
        c.font = Font(bold=True)
        c.border = BOX
        for j in range(2, 4 + len(TYPES)):
            col = get_column_letter(j)
            gray(ws, f"{col}{tot}", f"=SUM({col}{PM_FIRST}:{col}{PM_LAST})")
        ws.column_dimensions["A"].width = 31
        for j in range(2, 4 + len(TYPES)):
            ws.column_dimensions[get_column_letter(j)].width = 13
        ws.freeze_panes = f"B{PM_FIRST}"

        # ---------------------------------------------------------- Scorecard (bonus)
        ws = wb.create_sheet(SC)
        ws.sheet_properties.tabColor = "BF9000"
        titled(ws, "Inpatient claims scorecard by payer · 2025 (bonus)",
               "Count only Claims rows whose EncounterType is Inpatient. Adjudicated = ClaimStatus is not Pending. Denied = "
               "ClaimStatus is Denied or Appealed. Denial rate = Denied ÷ Adjudicated (n/a when there are no adjudicated "
               "claims). Avg days to pay = average of your DaysToPay helper column (n/a when there are no paid claims).",
               f"Fill each yellow column with ONE formula typed in row {SC_FIRST} and copied down to row {SC_LAST}.",
               cols=5, width=85)
        header(ws, SC_HDR, ["Payer", "Adjudicated claims", "Denied or appealed", "Denial rate", "Avg days to pay"])
        labels(ws, SC_FIRST, payer_names)
        yellow(ws, f"B{SC_FIRST}:C{SC_LAST}", "#,##0")
        yellow(ws, f"D{SC_FIRST}:D{SC_LAST}", "0.0%")
        yellow(ws, f"E{SC_FIRST}:E{SC_LAST}", "0.0")
        ws.column_dimensions["A"].width = 31
        for col in "BCDE":
            ws.column_dimensions[col].width = 14
        ws.freeze_panes = f"B{SC_FIRST}"

        # ---------------------------------------------------------- Markdown clean-up inside Excel cells
        # The library writes the bonus scenario (Bonus sheet) and each solution (key sheets, column D) verbatim, so the
        # README's **bold** and `code` markers would show up literally in Excel. Strip them in the workbook only.
        for sheet, col in ((lesson.bonus_sheet, "A"), (lesson.bonus_key_sheet, "D"), (lesson.key_sheet, "D")):
            if sheet not in wb.sheetnames:
                continue
            for cell in wb[sheet][col]:
                v = cell.value
                if isinstance(v, str) and not v.startswith("=") and ("**" in v or "`" in v):
                    cell.value = v.replace("**", "").replace("`", "")

        if selftest:
            # Simulate the learner's scorecard columns B, C, and E (task B2's own fill writes column D).
            for i in range(len(payer_names)):
                r = SC_FIRST + i
                for col, f in (("B", sc_adj), ("C", sc_den), ("E", sc_dtp)):
                    lesson.set_formula(ws, f"{col}{r}", f.replace(f"$A{SC_FIRST}", f"$A{r}"))

    return L
