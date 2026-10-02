"""Lesson 4.6 · Building Interactive Dashboards.

The builder turns the raw CSVs into a monthly KPI fact table (KPI_Monthly: one row per hospital per month, with
additive components such as LWBS and EDVisits next to the non-additive rates and medians), a visit-level ED table
for medians (EDWaits), and a KPI dictionary (Targets). Customize hooks add the dashboard layers:

  Dashboard      the practice canvas: two yellow selectors named SelFacility / SelMonth (data-validation dropdowns),
                 a finished example card (Inpatient discharges), and space for the learner's chart and combo box.
  Calc           the model layer: selection helpers, the combo-box link cell (task 10), and the 12-month trend block
                 the learner fills and charts (task 11).
  Lists          the dropdown sources (facilities incl. "All facilities", 24 months).
  Board          a blank canvas for the bonus build.
  Dashboard Key  (hidden, protected without a password) a finished, formula-driven reference dashboard with eight
                 KPI cards, direction-aware status colors, two openpyxl charts, and its model area. It is set to the
                 bonus selection, and the Bonus Key's live formulas for B3/B4 read it.

Practice tasks 1-9 are formulas typed on the Practice sheet that refer to SelFacility and SelMonth (pre-set to
Cedar Ridge Medical Center / Nov 2025). Tasks 10-11 are summary tasks (the self-test simulates the learner's combo
box and trend block). Tasks 12-13 are PivotTable/slicer builds: the learner types what the pivots show, the self-test
types the expected value, and the key's live SUMIFS formulas cross-check the Python answers.
"""
from __future__ import annotations

import statistics
from collections import defaultdict
from datetime import date

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from xlcourse import Lesson, Task, data
from xlcourse.lesson import BOX, HEADER_FILL, INPUT_BORDER, INPUT_FILL, NAVY, PREFILL_FILL, _plain

CODE = "4.6"

FACILITIES = [("F01", "Bluestone Memorial Hospital"), ("F02", "Ashby Falls Community Hospital"),
              ("F03", "Cedar Ridge Medical Center")]
FAC_NAME = dict(FACILITIES)
ALL = "All facilities"
FACILITY_LIST = [ALL] + [n for _, n in FACILITIES]
MONTHS = [date(y, m, 1) for y in (2024, 2025) for m in range(1, 13)]
IMMATURE = date(2025, 12, 1)            # readmissions, surveys, and claims for this month aren't final yet

# Practice selection (pre-set in the Dashboard selectors) and the slices used by the build tasks.
SEL_FAC = "Cedar Ridge Medical Center"
SEL_MONTH = date(2025, 11, 1)
COMBO_PICK = "Ashby Falls Community Hospital"
PIVOT_FAC = "Ashby Falls Community Hospital"
PIVOT_FROM, PIVOT_TO = date(2025, 7, 1), date(2025, 9, 1)      # 2025 Q3
# Bonus selection (and the Dashboard Key's default).
B_MONTH = date(2025, 10, 1)

MONTH_FMT = "mmm yyyy"
GOOD_FILL = PatternFill("solid", fgColor="C6EFCE")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
CARD_FILL = PatternFill("solid", fgColor="F2F2F2")
NOTE_FONT = Font(italic=True, color="595959")
WHITE_BOLD = Font(bold=True, color="FFFFFF")

KPI_COLUMNS = ["Month", "Facility", "EDVisits", "LWBS", "LWBSRate", "MedianDTP", "IPDischarges", "LOSDays", "ALOS",
               "IndexStays", "Readmits", "ReadmitRate", "PatientDays", "BedDays", "Occupancy", "Surveys", "TopBox",
               "TopBoxRate", "ClaimsAdjudicated", "ClaimsDenied", "DenialRate"]

# The KPI dictionary: (KPI, unit, question, definition, direction, target, owner, numerator, denominator)
TARGETS = [
    ("LWBS %", "%", "Are ED patients leaving before a provider sees them?",
     "LWBS ÷ EDVisits. ED arrivals by arrival month.", "Lower is better", 0.02, "ED Medical Director", "LWBS", "EDVisits"),
    ("Median door-to-provider", "minutes", "How long do ED patients wait to see a provider?",
     "Median minutes from arrival to first provider contact, over patients seen by a provider. Medians can't be added "
     "or averaged across hospitals: compute roll-ups from EDWaits.", "Lower is better", 30, "ED Nurse Director", None, None),
    ("ALOS", "days", "Are inpatients staying longer than they need to?",
     "LOSDays ÷ IPDischarges. Inpatient stays by discharge month.", "Lower is better", 4.5,
     "Director of Case Management", "LOSDays", "IPDischarges"),
    ("30-day readmission rate", "%", "Are discharged patients coming back within 30 days?",
     "Readmits ÷ IndexStays. Inpatient discharges by month, excluding deaths. Blank for Dec 2025 (window still open).",
     "Lower is better", 0.15, "Chief Quality Officer", "Readmits", "IndexStays"),
    ("Occupancy %", "%", "Will we have a bed for the next admission?",
     "PatientDays ÷ BedDays over inpatient and critical care units. Above the ceiling leaves too little surge room.",
     "Lower is better", 0.85, "Chief Nursing Officer", "PatientDays", "BedDays"),
    ("HCAHPS top-box %", "%", "Would patients rate their stay 9 or 10 out of 10?",
     "TopBox ÷ Surveys. Surveys by discharge month. Blank for Dec 2025 (most surveys not returned yet).",
     "Higher is better", 0.50, "Director of Patient Experience", "TopBox", "Surveys"),
    ("Denial rate", "%", "Are payers refusing our claims?",
     "ClaimsDenied ÷ ClaimsAdjudicated (Denied or Appealed ÷ every claim except Pending), by service month. "
     "Blank for Dec 2025 (most claims still pending).", "Lower is better", 0.10, "Revenue Cycle Director",
     "ClaimsDenied", "ClaimsAdjudicated"),
]
TARGET_BY_KPI = {t[0]: t for t in TARGETS}


# ---------------------------------------------------------------------------------------------------------- data
def _month(d) -> date:
    return date(d.year, d.month, 1)


def wait_rows() -> list[dict]:
    """Every 2024-2025 ED visit at the three hospitals that was seen by a provider, with door-to-provider minutes."""
    rows = []
    for v in data.load("ed_visits"):
        if v["FacilityID"] not in FAC_NAME or v["ProviderSeenDateTime"] is None:
            continue
        mins = (v["ProviderSeenDateTime"] - v["ArrivalDateTime"]).total_seconds() / 60
        assert mins == int(mins)
        rows.append({"EDVisitID": v["EDVisitID"], "ArrivalDateTime": v["ArrivalDateTime"],
                     "Facility": FAC_NAME[v["FacilityID"]], "DoorToProviderMin": int(mins)})
    rows.sort(key=lambda r: (r["ArrivalDateTime"], r["EDVisitID"]))
    return rows


def kpi_rows(waits: list[dict]) -> list[dict]:
    """One row per hospital per month. Components are stored next to the rates, so roll-ups can be rebuilt."""
    acc = {(m, f): defaultdict(float) for m in MONTHS for f, _ in FACILITIES}
    for v in data.load("ed_visits"):
        k = (_month(v["ArrivalDateTime"]), v["FacilityID"])
        if k in acc:
            acc[k]["ED"] += 1
            acc[k]["LWBS"] += v["EDDisposition"] == "LWBS"
    dtp = defaultdict(list)
    fid = {n: f for f, n in FACILITIES}
    for w in waits:
        dtp[(_month(w["ArrivalDateTime"]), fid[w["Facility"]])].append(w["DoorToProviderMin"])
    for e in data.load("encounters"):
        if e["EncounterType"] != "Inpatient":
            continue
        k = (_month(e["DischargeDateTime"]), e["FacilityID"])
        if k not in acc:
            continue
        acc[k]["Dis"] += 1
        acc[k]["LOS"] += (e["DischargeDateTime"] - e["AdmitDateTime"]).total_seconds() / 86400
        if e["DischargeDisposition"] != "Expired":
            acc[k]["Idx"] += 1
            acc[k]["Re"] += e["Readmit30"] == "Y"
    unit_type = {d["DeptID"]: d["UnitType"] for d in data.load("departments")}
    for c in data.load("daily_census"):
        if unit_type[c["DeptID"]] not in ("Inpatient", "Critical Care"):
            continue
        k = (_month(c["CensusDate"]), c["FacilityID"])
        acc[k]["PD"] += c["MidnightCensus"]
        acc[k]["BD"] += c["StaffedBeds"]
    for s in data.load("patient_satisfaction"):
        k = (_month(s["DischargeDate"]), s["FacilityID"])
        if k in acc:
            acc[k]["Sv"] += 1
            acc[k]["Top"] += s["OverallRating"] >= 9
    enc_fac = {e["EncounterID"]: e["FacilityID"] for e in data.load("encounters")}
    for c in data.load("claims"):
        k = (_month(c["ServiceDate"]), enc_fac[c["EncounterID"]])
        if k not in acc or c["ClaimStatus"] == "Pending":
            continue
        acc[k]["Adj"] += 1
        acc[k]["Den"] += c["ClaimStatus"] in ("Denied", "Appealed")

    rows = []
    for m in MONTHS:
        for f, name in FACILITIES:
            a = acc[(m, f)]
            r = {"Month": m, "Facility": name}
            r["EDVisits"], r["LWBS"] = int(a["ED"]), int(a["LWBS"])
            r["LWBSRate"] = r["LWBS"] / r["EDVisits"]
            r["MedianDTP"] = float(statistics.median(dtp[(m, f)]))
            r["IPDischarges"] = int(a["Dis"])
            r["LOSDays"] = round(a["LOS"], 2)
            r["ALOS"] = r["LOSDays"] / r["IPDischarges"]
            r["PatientDays"], r["BedDays"] = int(a["PD"]), int(a["BD"])
            r["Occupancy"] = r["PatientDays"] / r["BedDays"]
            if m == IMMATURE:
                for c in ("IndexStays", "Readmits", "ReadmitRate", "Surveys", "TopBox", "TopBoxRate",
                          "ClaimsAdjudicated", "ClaimsDenied", "DenialRate"):
                    r[c] = None
            else:
                r["IndexStays"], r["Readmits"] = int(a["Idx"]), int(a["Re"])
                r["ReadmitRate"] = r["Readmits"] / r["IndexStays"]
                r["Surveys"], r["TopBox"] = int(a["Sv"]), int(a["Top"])
                r["TopBoxRate"] = r["TopBox"] / r["Surveys"]
                r["ClaimsAdjudicated"], r["ClaimsDenied"] = int(a["Adj"]), int(a["Den"])
                r["DenialRate"] = r["ClaimsDenied"] / r["ClaimsAdjudicated"]
            rows.append(r)
    return rows


# ---------------------------------------------------------------------------------------------------------- build
def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="04-advanced-analysis", slug="06-dashboards",
        title="Building Interactive Dashboards", level="Advanced", minutes=70,
        objectives=[
            "Plan a dashboard around audience, questions, and KPIs",
            "Build KPI cards and selector-driven formulas",
            "Connect slicers to multiple PivotTables and drive charts from selections",
            "Polish layout, interactivity, and performance",
        ],
        data_note="Monthly KPIs for Bluestone Health's three hospitals, Jan 2024 – Dec 2025 (72 hospital-months), "
                  "every ED visit's door-to-provider time, and a KPI dictionary with targets and owners.",
    )

    waits = wait_rows()
    kpis = kpi_rows(waits)

    L.add_table_sheet(
        "KPI_Monthly", kpis, table="tblKPI", columns=KPI_COLUMNS, start_row=5,
        notes=["KPI_Monthly · one row per hospital per month, January 2024 – December 2025",
               "Count and day columns are additive: SUMIFS them across hospitals and months. Rate and median columns are "
               "not: rebuild a rate from its two components, and compute a roll-up median from EDWaits.",
               "Readmission, survey, and claims columns are blank for December 2025 because those numbers aren't final "
               "yet (the 30-day window is still open, surveys are still arriving, and most claims are pending)."],
        formats={"Month": MONTH_FMT, "LWBSRate": "0.0%", "MedianDTP": "0.0", "LOSDays": "#,##0.00", "ALOS": "0.00",
                 "ReadmitRate": "0.0%", "Occupancy": "0.0%", "TopBoxRate": "0.0%", "DenialRate": "0.0%",
                 "PatientDays": "#,##0", "BedDays": "#,##0"},
        widths={"Month": 11, "Facility": 31})
    L.add_table_sheet(
        "EDWaits", waits, table="tblEDWaits", start_row=4,
        columns=["EDVisitID", "ArrivalDateTime", "Facility", "DoorToProviderMin"],
        notes=["EDWaits · every 2024–2025 ED visit that was seen by a provider",
               "DoorToProviderMin = minutes from arrival to first provider contact. Visits that left without being seen "
               "(LWBS) have no provider time, so they aren't listed. Use this table for medians across hospitals or months."],
        widths={"Facility": 31, "ArrivalDateTime": 18, "DoorToProviderMin": 19})
    target_rows = [{"KPI": t[0], "Unit": t[1], "Question": t[2], "Definition": t[3], "Direction": t[4],
                    "Target": t[5], "Owner": t[6]} for t in TARGETS]
    L.add_table_sheet(
        "Targets", target_rows, table="tblTargets", start_row=3,
        columns=["KPI", "Unit", "Question", "Definition", "Direction", "Target", "Owner"],
        notes=["Targets · the KPI dictionary: what each KPI answers, how it's calculated, which way is good, "
               "the target, and who owns it",
               "On target means at or below a 'Lower is better' target, or at or above a 'Higher is better' target."],
        widths={"KPI": 24, "Unit": 9, "Question": 46, "Definition": 70, "Direction": 16, "Target": 9, "Owner": 28})

    # ------------------------------------------------------------------ Python answers
    by = {(r["Facility"], r["Month"]): r for r in kpis}

    def tot(col, month, fac=None):
        return sum(r[col] or 0 for r in kpis if r["Month"] == month and (fac is None or r["Facility"] == fac))

    def ratio(num, den, month, fac=None):
        d = tot(den, month, fac)
        return tot(num, month, fac) / d if d else None

    def months_between(a, b):
        return [m for m in MONTHS if a <= m <= b]

    def edate(m: date, k: int) -> date:
        n = m.year * 12 + (m.month - 1) + k
        return date(n // 12, n % 12 + 1, 1)

    def median_dtp(month, fac=None):
        vals = [w["DoorToProviderMin"] for w in waits if _month(w["ArrivalDateTime"]) == month
                and (fac is None or w["Facility"] == fac)]
        return float(statistics.median(vals))

    def kpi_value(kpi, month, fac=None):
        """A KPI for one hospital (fac) or all three (fac=None), rebuilt the way the dashboard does it."""
        if kpi == "Median door-to-provider":
            return median_dtp(month, fac)
        t = TARGET_BY_KPI[kpi]
        return ratio(t[7], t[8], month, fac)

    def on_target(kpi, v):
        t = TARGET_BY_KPI[kpi]
        assert abs(v - t[5]) > 1e-9 or kpi == "HCAHPS top-box %", f"{kpi} sits on its target; float risk"
        return v <= t[5] if t[4] == "Lower is better" else v >= t[5]

    sel = by[(SEL_FAC, SEL_MONTH)]
    prev_m, ly_m = edate(SEL_MONTH, -1), edate(SEL_MONTH, -12)
    t1 = sel["EDVisits"]
    t2 = sel["LWBS"] / sel["EDVisits"]
    t3 = t2 - TARGET_BY_KPI["LWBS %"][5]
    t4 = sel["EDVisits"] / by[(SEL_FAC, ly_m)]["EDVisits"] - 1
    occ_now = sel["PatientDays"] / sel["BedDays"]
    occ_prev = by[(SEL_FAC, prev_m)]["PatientDays"] / by[(SEL_FAC, prev_m)]["BedDays"]
    assert occ_now != occ_prev
    t5 = "▲" if occ_now > occ_prev else "▼"
    dtp_sel = sel["MedianDTP"]
    assert dtp_sel == median_dtp(SEL_MONTH, SEL_FAC)
    t6 = "On target" if on_target("Median door-to-provider", dtp_sel) else "Off target"
    t7 = ratio("LWBS", "EDVisits", SEL_MONTH)
    t7_avg = statistics.mean(by[(n, SEL_MONTH)]["LWBSRate"] for _, n in FACILITIES)
    assert abs(t7 - t7_avg) > 0.001
    t8 = median_dtp(SEL_MONTH)
    t8_avg = statistics.mean(by[(n, SEL_MONTH)]["MedianDTP"] for _, n in FACILITIES)
    assert t8 != t8_avg
    t9 = sum(by[(SEL_FAC, m)]["EDVisits"] for m in months_between(edate(SEL_MONTH, -2), SEL_MONTH))
    t10 = FACILITY_LIST.index(COMBO_PICK) + 1
    trend_months = months_between(edate(SEL_MONTH, -11), SEL_MONTH)
    assert len(trend_months) == 12
    t11 = sum(by[(SEL_FAC, m)]["EDVisits"] for m in trend_months)
    t11_peak = max(trend_months, key=lambda m: by[(SEL_FAC, m)]["EDVisits"])
    q3 = months_between(PIVOT_FROM, PIVOT_TO)
    t12 = sum(by[(PIVOT_FAC, m)]["EDVisits"] for m in q3)
    t13 = sum(by[(PIVOT_FAC, m)]["LWBS"] for m in q3) / t12
    t13_avg = statistics.mean(by[(PIVOT_FAC, m)]["LWBSRate"] for m in q3)
    ip_now, ip_ly = sel["IPDischarges"], by[(SEL_FAC, ly_m)]["IPDischarges"]

    # Bonus (All facilities, Oct 2025)
    kpi_names = [t[0] for t in TARGETS]
    b1 = ratio("Readmits", "IndexStays", B_MONTH)
    sys_vals = {k: kpi_value(k, B_MONTH) for k in kpi_names}
    b2 = sum(on_target(k, v) for k, v in sys_vals.items())
    assert sys_vals["HCAHPS top-box %"] == TARGET_BY_KPI["HCAHPS top-box %"][5]   # the B2 explanation relies on the tie
    hosp_counts = {n: sum(on_target(k, kpi_value(k, B_MONTH, n)) for k in kpi_names) for _, n in FACILITIES}
    fewest = min(hosp_counts.values())
    assert list(hosp_counts.values()).count(fewest) == 1, hosp_counts
    b3 = min(hosp_counts, key=hosp_counts.get)
    b_trend = months_between(edate(B_MONTH, -11), B_MONTH)
    lwbs_trend = {m: ratio("LWBS", "EDVisits", m) for m in b_trend}
    peak = max(lwbs_trend.values())
    assert list(lwbs_trend.values()).count(peak) == 1
    b4 = max(lwbs_trend, key=lwbs_trend.get)

    # ------------------------------------------------------------------ formulas
    def f_sel(col, month="SelMonth", fac="SelFacility"):
        return f"SUMIFS(tblKPI[{col}],tblKPI[Facility],{fac},tblKPI[Month],{month})"

    def f_sys(col, month="SelMonth"):
        return f"SUMIFS(tblKPI[{col}],tblKPI[Month],{month})"

    month_label = SEL_MONTH.strftime("%b %Y")
    L.practice_intro = (
        f"The yellow selectors on the Dashboard sheet are named SelFacility (Dashboard!C4) and SelMonth (Dashboard!C5). "
        f"They start at {SEL_FAC} and {month_label}. Write every formula with SelFacility and SelMonth, never typed-in "
        f"names or dates, and keep that selection while you check your answers. (Change it afterwards and watch your "
        f"formulas follow.) Tasks 1–9 go in the yellow cells below. Tasks 10–13 have you build on other sheets.")

    occ_now_f = f"{f_sel('PatientDays')}/{f_sel('BedDays')}"
    occ_prev_f = f"{f_sel('PatientDays', 'EDATE(SelMonth,-1)')}/{f_sel('BedDays', 'EDATE(SelMonth,-1)')}"
    status_f = ('=LET(val,XLOOKUP(1,(tblKPI[Facility]=SelFacility)*(tblKPI[Month]=SelMonth),tblKPI[MedianDTP]),'
                'tgt,XLOOKUP("Median door-to-provider",tblTargets[KPI],tblTargets[Target]),'
                'dir,XLOOKUP("Median door-to-provider",tblTargets[KPI],tblTargets[Direction]),'
                'IF(IF(dir="Lower is better",val<=tgt,val>=tgt),"On target","Off target"))')
    median_f = ("=MEDIAN(FILTER(tblEDWaits[DoorToProviderMin],"
                "(tblEDWaits[ArrivalDateTime]>=SelMonth)*(tblEDWaits[ArrivalDateTime]<EDATE(SelMonth,1))))")
    rolling_f = ('=SUMIFS(tblKPI[EDVisits],tblKPI[Facility],SelFacility,'
                 'tblKPI[Month],">="&EDATE(SelMonth,-2),tblKPI[Month],"<="&SelMonth)')
    q3_crit = (f'tblKPI[Facility],"{PIVOT_FAC}",tblKPI[Month],">="&DATE({PIVOT_FROM.year},{PIVOT_FROM.month},1),'
               f'tblKPI[Month],"<="&DATE({PIVOT_TO.year},{PIVOT_TO.month},1)')

    # Calc sheet layout (built in customize), referenced by tasks 10 and 11.
    LINK_CELL = "C12"
    TREND_HDR, TREND_FIRST, TREND_LAST = 16, 17, 28
    trend_m = f"Calc!$B${TREND_FIRST}:$B${TREND_LAST}"
    trend_v = f"Calc!$C${TREND_FIRST}:$C${TREND_LAST}"
    trend_month_formula = f"=EDATE(SelMonth,ROWS(B${TREND_FIRST}:B{TREND_FIRST})-12)"
    trend_value_formula = f"=SUMIFS(tblKPI[EDVisits],tblKPI[Facility],SelFacility,tblKPI[Month],B{TREND_FIRST})"

    L.tasks = [
        Task("Write a formula that returns the number of ED visits for the selected facility and month "
             "(the EDVisits column of tblKPI).",
             answer=t1, solution=f"={f_sel('EDVisits')}",
             hint="SUMIFS with two criteria: Facility = SelFacility and Month = SelMonth",
             explanation="tblKPI has exactly one row per hospital per month, so SUMIFS with both criteria returns that "
                         "row's value. SUMIFS is the workhorse of selector-driven dashboards. It never returns #N/A, it "
                         "adds up correctly when a criterion matches several rows (a quarter, or every hospital), and it "
                         "recalculates the moment a selector changes. Because the criteria are the named cells, choosing "
                         "a different facility on the Dashboard updates this number without anyone touching the formula."),
        Task("Return the LWBS % (left without being seen) for the selection. Build it from its components, LWBS ÷ "
             "EDVisits, instead of reading the LWBSRate column. Enter it as a percentage.",
             answer=t2, fmt="0.00%", solution=f"={f_sel('LWBS')}/{f_sel('EDVisits')}",
             hint="One SUMIFS for the numerator divided by one SUMIFS for the denominator",
             explanation=f"{sel['LWBS']} of {sel['EDVisits']} patients left before a provider saw them. For one "
                         "hospital-month this matches the LWBSRate column. The components version is still the better "
                         "habit, because the same formula keeps working when the selection covers several rows (a "
                         "quarter, or all hospitals). There, SUMIFS on LWBSRate would add percentages together, which "
                         "is meaningless."),
        Task("LWBS variance to target: the selected LWBS % minus the LWBS % target from tblTargets. Look the target "
             "up with a formula instead of typing 2%. Enter the result as a percentage. A positive result means the hospital "
             "is above (worse than) this lower-is-better target.",
             answer=t3, fmt="0.00%",
             solution=f'={f_sel("LWBS")}/{f_sel("EDVisits")}-XLOOKUP("LWBS %",tblTargets[KPI],tblTargets[Target])',
             hint='XLOOKUP("LWBS %", tblTargets[KPI], tblTargets[Target])',
             explanation=f"{t2:.2%} − {TARGET_BY_KPI['LWBS %'][5]:.2%} = {t3:+.2%}. Strictly, that's "
                         f"{t3 * 100:+.2f} percentage points: the gap between two rates. (The relative variance, actual "
                         f"÷ target − 1, would be {t2 / TARGET_BY_KPI['LWBS %'][5] - 1:+.1%}.) Keeping targets in a table "
                         "and looking them up means one edit to tblTargets updates every card that uses the target. On a "
                         "real card this cell would subtract the target cell from the value cell."),
        Task("ED visits change versus the same month last year: this year ÷ last year − 1. Find last year's month with "
             "EDATE so the formula works for any selected month. Enter it as a percentage.",
             answer=t4, fmt="0.0%",
             solution=f"={f_sel('EDVisits')}/{f_sel('EDVisits', 'EDATE(SelMonth,-12)')}-1",
             hint="EDATE(SelMonth,-12) is the same month one year earlier",
             explanation=f"{t1} visits this year against {by[(SEL_FAC, ly_m)]['EDVisits']} in "
                         f"{ly_m.strftime('%b %Y')}. Comparing with the same month last year removes seasonality, which a "
                         "month-over-month change can't do: November always differs from October in an ED. EDATE moves a "
                         "date by whole months and keeps it on the first of the month, so it matches the Month column "
                         "exactly. Subtracting 365 days would not."),
        Task("Occupancy trend arrow. Compare the selected month's occupancy (PatientDays ÷ BedDays) with the previous "
             "month's. Return ▲ with UNICHAR(9650) if it rose, ▼ with UNICHAR(9660) if it fell, or ▬ with "
             "UNICHAR(9644) if it didn't change.",
             answer=t5,
             solution=f"=LET(cur,{occ_now_f},prev,{occ_prev_f},"
                      "IF(cur>prev,UNICHAR(9650),IF(cur<prev,UNICHAR(9660),UNICHAR(9644))))",
             hint="LET(cur, …, prev, …, IF(cur>prev, UNICHAR(9650), …)). EDATE(SelMonth,-1) is the previous month",
             explanation=f"Occupancy went from {occ_prev:.1%} in {prev_m.strftime('%b %Y')} to {occ_now:.1%}. LET names "
                         "the two rates once, so the IF reads like a sentence instead of repeating four SUMIFS. A "
                         "formula-made arrow is ordinary text, so you can color it with conditional formatting. On a "
                         "dashboard, rising occupancy deserves amber or red as it approaches the ceiling, even though "
                         "the arrow itself only says which way it moved."),
        Task("Status cell for median door-to-provider: return the text On target or Off target for the selection. Read "
             "the target and the Direction of \"Median door-to-provider\" from tblTargets, so the same pattern works for "
             "any KPI. A median isn't additive, so look up the hospital's MedianDTP value instead of adding it up.",
             answer=t6, solution=status_f,
             hint="LET + XLOOKUP(1, (tblKPI[Facility]=SelFacility)*(tblKPI[Month]=SelMonth), tblKPI[MedianDTP]); "
                  "then IF on the direction",
             explanation=f"The median is {dtp_sel:.1f} minutes against a target of "
                         f"{TARGET_BY_KPI['Median door-to-provider'][5]} or less, so it's off target, even if only "
                         "barely. XLOOKUP(1, (condition)*(condition), …) finds the one row where both conditions are "
                         "TRUE (TRUE×TRUE = 1). The inner IF flips the comparison by direction: lower-is-better KPIs pass "
                         "at or below the target, and higher-is-better KPIs pass at or above it. Binary statuses hide "
                         "near misses like this one, which is why many dashboards add an amber band (for example, within "
                         "10% of target)."),
        Task("System-wide LWBS % for the selected month: all three hospitals combined, ignoring the facility selector. "
             "Divide total LWBS by total ED visits instead of averaging the three hospitals' rates. Enter it as a percentage.",
             answer=t7, fmt="0.00%", solution=f"={f_sys('LWBS')}/{f_sys('EDVisits')}",
             hint="Drop the Facility criterion from both SUMIFS",
             explanation=f"{tot('LWBS', SEL_MONTH)} ÷ {tot('EDVisits', SEL_MONTH)} = {t7:.2%}. Averaging the three "
                         f"hospital rates gives {t7_avg:.2%}, because a simple average gives the small hospitals the "
                         "same weight as Bluestone Memorial, which sees several times as many patients. Rates roll up as "
                         "total numerator ÷ total denominator. That's why tblKPI stores the components."),
        Task("System-wide median door-to-provider, in minutes, for the selected month. Medians can't be added or "
             "averaged across hospitals, so compute it from the visit-level table: the MEDIAN of DoorToProviderMin for "
             "every tblEDWaits visit whose ArrivalDateTime falls in the selected month. Keep one decimal place.",
             answer=t8, fmt="0.0", solution=median_f,
             hint="MEDIAN(FILTER(…)) with ArrivalDateTime >= SelMonth and < EDATE(SelMonth,1)",
             explanation=f"The true system median is {t8:.1f} minutes. Averaging the three hospital medians gives "
                         f"{t8_avg:.1f}, which is not the median of anything. A median depends on every individual value, "
                         "so it can only be computed from the detail rows. The date test uses a half-open window "
                         "(>= first day, < first day of next month), which catches every arrival time on the last day "
                         "of the month. In Excel 2019 or earlier, use =AGGREGATE(17,6,values/(condition),2) instead, "
                         "where function 17 is QUARTILE.INC and quartile 2 is the median."),
        Task("Rolling 3-month ED visits for the selected facility: the selected month plus the two months before it "
             "(Sep–Nov 2025 for the default selection). Use one SUMIFS with a date window, not OFFSET.",
             answer=t9, solution=rolling_f,
             hint='Two criteria on the Month column: ">="&EDATE(SelMonth,-2) and "<="&SelMonth',
             explanation="Joining an operator to a date (\">=\"&EDATE(SelMonth,-2)) turns the date into criteria text "
                         "that SUMIFS understands. A common older pattern is SUM(OFFSET(…)). OFFSET is volatile, so Excel "
                         "recalculates it after every edit anywhere in the workbook, and dozens of them make a dashboard "
                         "sluggish. SUMIFS recalculates only when its inputs change, and it doesn't depend on the table "
                         "being sorted."),
        Task("Insert a Form Controls combo box on the Dashboard. In Format Control, set its Input range to "
             "Lists!$A$2:$A$5 and its Cell link to Calc!$C$12. Then choose Ashby Falls Community Hospital in the combo "
             "box. The gray cell shows the number your combo box writes to the cell link.",
             answer=t10, title="Combo box cell link",
             summary=f'=IF(Calc!{LINK_CELL}="","",Calc!{LINK_CELL})',
             fill={"range": f"Calc!{LINK_CELL}:{LINK_CELL}", "values": [t10]},
             live=f'=MATCH("{COMBO_PICK}",Lists!$A$2:$A$5,0)',
             hint="Developer → Insert → Combo Box (Form Control), then right-click it → Format Control → Control tab",
             solution=("1. Show the **Developer** tab if you haven't (see the guide).\n"
                       "2. **Developer → Insert → Combo Box (Form Control)** (Mac: **Developer → Combo Box**), then drag "
                       "a box onto the Dashboard.\n"
                       "3. Right-click the combo box → **Format Control** → **Control** tab. Input range: "
                       "`Lists!$A$2:$A$5`. Cell link: `Calc!$C$12`. Drop down lines: 4. Click **OK**.\n"
                       "4. Click a cell to deselect the control, then choose **Ashby Falls Community Hospital**.\n\n"
                       "Calc!C13 turns the number back into a name: `=INDEX(Lists!$A$2:$A$5,Calc!C12)`."),
             explanation=f"A combo box writes the position of the chosen item, not its text: {COMBO_PICK} is item {t10} "
                         "of the input range. INDEX(list, position) converts it back, which Calc!C13 does. To make the "
                         "combo box drive the whole dashboard, point your formulas (or the SelFacility name) at that "
                         "INDEX cell. A combo box floats above the grid, so it can't be typed over by accident, but it "
                         "doesn't work in Excel for the web, and the position-number link breaks if someone re-sorts "
                         "the list."),
        Task(f"On the Calc sheet, fill the yellow 12-month trend block. In B{TREND_FIRST}:B{TREND_LAST}, return the 12 "
             f"months ending at SelMonth, oldest first (use EDATE). In C{TREND_FIRST}:C{TREND_LAST}, return ED visits "
             f"for SelFacility in each of those months. Then select B{TREND_HDR}:C{TREND_LAST}, insert a line chart, "
             "and move it to the Dashboard. The gray cell totals your block: the trailing-12-month ED visits.",
             answer=t11, title="12-month trend block and chart (trailing-12-month ED visits)",
             summary=f'=IF(COUNT({trend_v})=0,"",SUM({trend_v}))',
             fill={"range": f"Calc!C{TREND_FIRST}:C{TREND_LAST}", "formula": trend_value_formula},
             live=('=SUMIFS(tblKPI[EDVisits],tblKPI[Facility],SelFacility,'
                   'tblKPI[Month],">="&EDATE(SelMonth,-11),tblKPI[Month],"<="&SelMonth)'),
             hint="EDATE(SelMonth,-11) is the oldest month. A counter such as ROWS(B$17:B17) lets one formula step forward "
                  "as you copy it down",
             solution=(f"1. In **Calc!B{TREND_FIRST}** type `{trend_month_formula}` and copy it down to "
                       f"B{TREND_LAST}. The first row gives EDATE(SelMonth,-11) and the last gives SelMonth itself.\n"
                       f"2. In **C{TREND_FIRST}** type `{trend_value_formula}` and copy it down.\n"
                       f"3. Select **B{TREND_HDR}:C{TREND_LAST}** → **Insert → Charts → Line**. The corner cell "
                       f"B{TREND_HDR} is blank on purpose, so Excel uses the months as the axis.\n"
                       "4. Cut the chart (Ctrl + X, Mac: ⌘ + X), click a cell on the Dashboard, and paste. The chart "
                       "keeps pointing at Calc.\n"
                       "5. Change the selectors on the Dashboard and watch the line redraw."),
             explanation=f"The chart's series points at formula cells, and those cells point at the selectors, so one "
                         f"dropdown redraws the chart. The block for {SEL_FAC} runs "
                         f"{trend_months[0].strftime('%b %Y')}–{trend_months[-1].strftime('%b %Y')} and peaks at "
                         f"{by[(SEL_FAC, t11_peak)]['EDVisits']} visits in {t11_peak.strftime('%b %Y')}. In Microsoft "
                         "365 you can instead spill the months with =EDATE(SelMonth,SEQUENCE(12,1,-11)). A chart built from "
                         "a spill keeps a fixed range, so if a spill can change size, chart it through a defined name "
                         "that refers to it (for example =Calc!$B$17#). If "
                         "you pick an early month, the window reaches back before January 2024 and SUMIFS returns 0. "
                         "Wrapping it as IF(COUNTIFS(…)=0,NA(),SUMIFS(…)) makes the line chart leave a gap instead of "
                         "plunging to zero."),
        Task("Insert a new sheet named Pivots. Build PivotTable 1 from tblKPI with Month in Rows and Sum of EDVisits in "
             "Values. Insert a Facility slicer and a Month timeline for it. Select Ashby Falls Community Hospital in the "
             "slicer, and select 2025 Q3 in the timeline (switch it to QUARTERS). What is PivotTable 1's grand total?",
             answer=t12, title="PivotTable 1 with a slicer and a timeline",
             live=f"=SUMIFS(tblKPI[EDVisits],{q3_crit})",
             hint="PivotTable Analyze → Insert Slicer, and PivotTable Analyze → Insert Timeline",
             solution=("1. Click in tblKPI → **Insert → PivotTable** → **New Worksheet** → **OK**. Rename the sheet "
                       "**Pivots**.\n"
                       "2. Drag **Month** to Rows and **EDVisits** to Values (Sum of EDVisits). Excel may group the "
                       "dates into Years and Quarters. That's fine.\n"
                       "3. **PivotTable Analyze → Insert Slicer** → tick **Facility** → **OK**. Click **Ashby Falls "
                       "Community Hospital**.\n"
                       "4. **PivotTable Analyze → Insert Timeline** → tick **Month** → **OK**. Set the time level to "
                       "**QUARTERS** and click **2025 Q3**.\n"
                       "5. Read the Grand Total.\n\n"
                       f"Formula twin: `=SUMIFS(tblKPI[EDVisits],{q3_crit})`"),
             explanation="A slicer is a selector made of buttons, and a timeline is a selector for dates. Both filter "
                         "the pivot they were inserted from. Visible selections are what make pivots dashboard-friendly: "
                         "a Filters-area dropdown hides what's selected, but a slicer shows it. The formula twin proves "
                         "the number: ED visits at Ashby Falls in July, August, and September 2025."),
        Task("On the Pivots sheet, build PivotTable 2 from tblKPI with Facility in Rows, plus Sum of LWBS and Sum of "
             "EDVisits in Values. Connect the slicer and the timeline to it (Report Connections). Then add a calculated "
             "field LWBSPct = LWBS / EDVisits. With Ashby Falls and 2025 Q3 still selected, what LWBS % does PivotTable "
             "2 show? Enter it as a percentage.",
             answer=t13, fmt="0.00%", title="PivotTable 2 connected to the same slicer and timeline",
             live=f"=SUMIFS(tblKPI[LWBS],{q3_crit})/SUMIFS(tblKPI[EDVisits],{q3_crit})",
             hint="Select the slicer → Slicer → Report Connections. Then PivotTable Analyze → Fields, Items & Sets → "
                  "Calculated Field",
             solution=("1. Click in tblKPI → **Insert → PivotTable** → **Existing Worksheet**, pick a cell on Pivots a "
                       "few columns right of PivotTable 1 → **OK**.\n"
                       "2. Drag **Facility** to Rows, then **LWBS** and **EDVisits** to Values.\n"
                       "3. Select the Facility slicer → **Slicer → Report Connections** → tick both PivotTables → "
                       "**OK**. Do the same for the timeline (**Timeline → Report Connections**). PivotTable 2 now "
                       "shows only Ashby Falls, Q3 2025.\n"
                       "4. Click in PivotTable 2 → **PivotTable Analyze → Fields, Items & Sets → Calculated Field**. "
                       "Name: `LWBSPct`. Formula: `=LWBS/EDVisits` → **Add** → **OK**. Format it as a percentage.\n\n"
                       f"Formula twin: `=SUMIFS(tblKPI[LWBS],{q3_crit})/SUMIFS(tblKPI[EDVisits],{q3_crit})`"),
             explanation=f"Report Connections is what lets one slicer filter several pivots. If you skipped it, "
                         "PivotTable 2 would still show every hospital and month. A pivot calculated field sums each "
                         "field first and then divides, so LWBSPct is total LWBS ÷ total visits for the quarter. That's "
                         f"exactly the right rate here: {t13:.2%}. Averaging the three monthly rates would give "
                         f"{t13_avg:.2%}. Both pivots must come from the same source (tblKPI) to share a slicer."),
    ]

    # ------------------------------------------------------------------ bonus
    B_LABEL = B_MONTH.strftime("%b %Y")
    L.bonus_title = "Bonus: The board's one-page operations dashboard"
    L.bonus_scenario = (
        "Each month the COO presents one page to the board's Quality & Operations Committee. It must answer three "
        "questions at a glance: Are we on target? Where are we missing? Which way are things heading? Build it on the "
        "Board sheet to this spec:\n\n"
        "1. Facility and Month dropdowns fed by the Lists sheet, named BoardFacility and BoardMonth. Facility must "
        "allow All facilities.\n"
        "2. One card for each of the seven KPIs in tblTargets, showing the value, the target, a status colored by "
        "conditional formatting, and an arrow versus the same month last year. Every card must work for All "
        "facilities, so rebuild rates from their components and compute the median door-to-provider from tblEDWaits.\n"
        "3. A scorecard line such as '3 of 7 KPIs on target'.\n"
        "4. A 12-month LWBS % trend block and a line chart that follow both dropdowns.\n"
        "5. Polish: gridlines and headings off, only the two dropdowns unlocked, the sheet protected, and one "
        "landscape page when printed.\n\n"
        "The hidden Dashboard Key sheet is a finished reference build. Compare your numbers with it when you're done.")

    readmit_card = ('=SUMIFS(tblKPI[Readmits],tblKPI[Facility],FacCrit,tblKPI[Month],BoardMonth)'
                    '/SUMIFS(tblKPI[IndexStays],tblKPI[Facility],FacCrit,tblKPI[Month],BoardMonth)')
    b_date = f"DATE({B_MONTH.year},{B_MONTH.month},1)"
    b2_live = (f"=LET(m,{b_date},"
               "lwbs,SUMIFS(tblKPI[LWBS],tblKPI[Month],m)/SUMIFS(tblKPI[EDVisits],tblKPI[Month],m),"
               "dtp,MEDIAN(FILTER(tblEDWaits[DoorToProviderMin],"
               "(tblEDWaits[ArrivalDateTime]>=m)*(tblEDWaits[ArrivalDateTime]<EDATE(m,1)))),"
               "alos,SUMIFS(tblKPI[LOSDays],tblKPI[Month],m)/SUMIFS(tblKPI[IPDischarges],tblKPI[Month],m),"
               "readm,SUMIFS(tblKPI[Readmits],tblKPI[Month],m)/SUMIFS(tblKPI[IndexStays],tblKPI[Month],m),"
               "occ,SUMIFS(tblKPI[PatientDays],tblKPI[Month],m)/SUMIFS(tblKPI[BedDays],tblKPI[Month],m),"
               "tbox,SUMIFS(tblKPI[TopBox],tblKPI[Month],m)/SUMIFS(tblKPI[Surveys],tblKPI[Month],m),"
               "den,SUMIFS(tblKPI[ClaimsDenied],tblKPI[Month],m)/SUMIFS(tblKPI[ClaimsAdjudicated],tblKPI[Month],m),"
               "tg,tblTargets[Target],"
               "(lwbs<=INDEX(tg,1))+(dtp<=INDEX(tg,2))+(alos<=INDEX(tg,3))+(readm<=INDEX(tg,4))"
               "+(occ<=INDEX(tg,5))+(tbox>=INDEX(tg,6))+(den<=INDEX(tg,7)))")
    sys_lines = "\n".join(
        f"| {k} | {_fmt_kpi(k, v)} | {'≤' if TARGET_BY_KPI[k][4] == 'Lower is better' else '≥'} "
        f"{_fmt_kpi(k, TARGET_BY_KPI[k][5])} | {'✔ On target' if on_target(k, v) else '✘ Off target'} |"
        for k, v in sys_vals.items())
    hosp_lines = "\n".join(f"| {n} | {c} of 7 |" for n, c in hosp_counts.items())

    # Dashboard Key model cells read by the bonus key (see _dashboard_key below).
    DK = "'Dashboard Key'"
    L.bonus = [
        Task(f"Set your Board dashboard to All facilities and {B_LABEL}. What is the system-wide 30-day readmission "
             "rate? Enter it as a percentage.",
             answer=b1, fmt="0.00%", title=f"System-wide readmission rate, {B_LABEL}",
             live=(f'=SUMIFS(tblKPI[Readmits],tblKPI[Facility],"*",tblKPI[Month],{b_date})'
                   f'/SUMIFS(tblKPI[IndexStays],tblKPI[Facility],"*",tblKPI[Month],{b_date})'),
             hint='Swap "All facilities" for the asterisk wildcard in the Facility criterion',
             solution=("1. Make a helper cell named **FacCrit**: `=IF(BoardFacility=\"All facilities\",\"*\",BoardFacility)`.\n"
                       "2. The readmission card's value:\n\n"
                       f"```\n{readmit_card}\n```"),
             explanation=f"{tot('Readmits', B_MONTH)} readmissions ÷ {tot('IndexStays', B_MONTH)} index stays = "
                         f"{b1:.2%}. No row in tblKPI says \"All facilities\", so SUMIFS with that text returns 0. The "
                         "wildcard `*` matches any text, so the same SUMIFS adds all three hospitals. Every additive "
                         "component works this way, and every rate is then rebuilt as total ÷ total."),
        Task(f"With All facilities and {B_LABEL} still selected, how many of the seven KPIs are on target?",
             answer=b2, title=f"Scorecard, All facilities, {B_LABEL}", live=b2_live,
             hint="Give each card a status cell that returns 1 or 0, then SUM them. Remember which KPIs are "
                  "higher-is-better",
             solution=("Give each card a numeric status cell, for example for LWBS:\n\n"
                       "```\n=IF(ISNUMBER(val),--IF(dir=\"Lower is better\",val<=tgt,val>=tgt),\"\")\n```\n\n"
                       "Then the scorecard is `=SUM(statuses)&\" of \"&COUNT(statuses)&\" KPIs on target\"`.\n\n"
                       "The system values for this month:\n\n"
                       "| KPI | Value | Target | Status |\n|---|---|---|---|\n" + sys_lines),
             explanation=f"{b2} of 7. The median door-to-provider must come from tblEDWaits, because the three hospital "
                         "medians can't be combined. HCAHPS top-box is the one higher-is-better KPI, and this month it "
                         f"lands at exactly {sys_vals['HCAHPS top-box %']:.1%}: on target with >=, off target with >. "
                         "Decide ties in the KPI dictionary, not in each formula. Status cells that return 1 or 0 (not "
                         "text) make the scorecard a plain SUM and make conditional formatting rules simple."),
        Task(f"Keep {B_LABEL} and switch the Facility dropdown to each hospital in turn. Which hospital has the fewest "
             "KPIs on target?",
             answer=b3, accept=[b3.replace(" Hospital", "")], title=f"Hospital with the fewest KPIs on target, {B_LABEL}",
             live=f"={DK}!$O$46",
             hint="Your scorecard line answers this. Change only the Facility dropdown",
             solution=("Read the scorecard line after each switch:\n\n| Hospital | On target |\n|---|---|\n" + hosp_lines
                       + "\n\nThe Dashboard Key's model area has this comparison table (columns N–V), with "
                         "`=INDEX(N42:N44,MATCH(MIN(V42:V44),V42:V44,0))` picking the lowest."),
             explanation=f"{b3} meets {fewest} of 7 targets in {B_LABEL}. A selector-driven dashboard answers \"which "
                         "hospital?\" in a few clicks without a separate report per hospital. The large hospital being "
                         "furthest off target is also why the system-wide scorecard lands where it does: its volumes "
                         "dominate every total ÷ total rate."),
        Task(f"Back on All facilities and {B_LABEL}, your 12-month LWBS % trend runs "
             f"{b_trend[0].strftime('%b %Y')}–{b_trend[-1].strftime('%b %Y')}. In which month was the system-wide LWBS % "
             "highest? Enter the first day of that month as a date.",
             answer=b4, fmt="mm/dd/yyyy", title="Peak system LWBS % month in the 12-month trend",
             live=f"={DK}!$O$36",
             hint="INDEX(months, MATCH(MAX(rates), rates, 0)) on your trend block, or just read the chart's peak",
             solution=("Trend block on the Board sheet (12 rows): month `=EDATE(BoardMonth,ROWS(first:current)-12)`, "
                       "and the rate:\n\n"
                       "```\n=IFERROR(SUMIFS(tblKPI[LWBS],tblKPI[Facility],FacCrit,tblKPI[Month],month)\n"
                       "  /SUMIFS(tblKPI[EDVisits],tblKPI[Facility],FacCrit,tblKPI[Month],month),NA())\n```\n\n"
                       "Then `=INDEX(months,MATCH(AGGREGATE(4,6,rates),rates,0))`. AGGREGATE(4,6,…) is MAX that "
                       "skips error values, so the #N/A gaps of an early month can't break it."),
             explanation=f"{b4.strftime('%B %Y')} peaks at {peak:.2%}, against a target of 2%. "
                         "Every other month in the window is lower. A target line on the chart (a second series that "
                         "repeats the target in every row) makes the misses visible at a glance, and the dynamic title "
                         "tells the reader which facility and months they're looking at."),
    ]

    # ------------------------------------------------------------------ workbook extras
    L.sheet_order = ["Start Here", "Practice", "Dashboard", "Calc", "KPI_Monthly", "EDWaits", "Targets", "Lists",
                     "Board", "Bonus", "Answer Key", "Bonus Key", "Dashboard Key"]
    L.start_notes = [
        f"The Dashboard sheet's yellow selectors are named SelFacility (Dashboard!C4) and SelMonth (Dashboard!C5). "
        f"They start at {SEL_FAC} and {month_label}. Tasks 1–9 and 11 are checked against that selection, so set it back "
        "if you've been exploring.",
        "Needs Microsoft 365 or Excel 2021+ for XLOOKUP, LET, and FILTER. The guide shows alternatives for older "
        "versions. Form Controls (task 10) need Excel for Windows or Mac, not Excel for the web.",
        "Dashboard Key (hidden) is a finished reference dashboard set to the bonus selection. It's protected without a "
        "password so its selectors are the only cells you can change. Review → Unprotect Sheet lets you edit it.",
        "KPI_Monthly leaves readmission, survey, and claims numbers blank for December 2025 on purpose: those numbers "
        "aren't final yet.",
    ]

    trend_month_cells = (TREND_FIRST, TREND_LAST, trend_month_formula)
    card_example = {"now": ip_now, "ly": ip_ly}

    @L.customize
    def _custom(wb, lesson, selftest):
        _lists(wb)
        _targets_formats(wb)
        _dashboard(wb, lesson)
        _calc(wb, lesson, LINK_CELL, TREND_HDR, TREND_FIRST, TREND_LAST, selftest, trend_month_cells)
        _board(wb)
        _dashboard_key(wb, lesson)
        wb.defined_names["SelFacility"] = DefinedName("SelFacility", attr_text="Dashboard!$C$4")
        wb.defined_names["SelMonth"] = DefinedName("SelMonth", attr_text="Dashboard!$C$5")
        # Show Markdown solutions as plain text inside Excel.
        for key_name in (lesson.key_sheet, lesson.bonus_key_sheet):
            ks = wb[key_name]
            for row in range(5, ks.max_row + 1):
                v = ks.cell(row=row, column=4).value
                if isinstance(v, str) and not v.startswith("="):
                    ks.cell(row=row, column=4).value = _plain(v)
        # Text answers in the Practice sheet: keep arrows and statuses readable.
        practice = wb[lesson.practice_sheet]
        for t in lesson.tasks:
            if t.kind() == "text":
                practice[t.answer_cell].alignment = Alignment(horizontal="center", vertical="top")
                practice[t.answer_cell].font = Font(size=12)

    L._answers = dict(t1=t1, t2=t2, t3=t3, t4=t4, t5=t5, t6=t6, t7=t7, t7_avg=t7_avg, t8=t8, t8_avg=t8_avg, t9=t9,
                      t10=t10, t11=t11, t12=t12, t13=t13, t13_avg=t13_avg, b1=b1, b2=b2, b3=b3, b4=b4,
                      hosp_counts=hosp_counts, sys_vals=sys_vals, card_example=card_example, lwbs_trend=lwbs_trend)
    return L


# ---------------------------------------------------------------------------------------------------------- helpers
def _fmt_kpi(kpi: str, v) -> str:
    unit = TARGET_BY_KPI[kpi][1]
    if v is None:
        return "n/a"
    if unit == "%":
        return f"{v:.1%}"
    if unit == "minutes":
        return f"{v:.1f} min"
    return f"{v:.2f} days"


def _hdr(ws, cell, text, width_cols: int = 1):
    c = ws[cell]
    c.value = text
    c.font = WHITE_BOLD
    c.fill = HEADER_FILL
    c.border = BOX
    if width_cols > 1:
        from openpyxl.utils import column_index_from_string, get_column_letter
        col = "".join(ch for ch in cell if ch.isalpha())
        row = int("".join(ch for ch in cell if ch.isdigit()))
        start = column_index_from_string(col)
        for k in range(1, width_cols):
            x = ws[f"{get_column_letter(start + k)}{row}"]
            x.fill = HEADER_FILL
            x.border = BOX


def _landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _lists(wb):
    ws = wb.create_sheet("Lists")
    ws.sheet_properties.tabColor = "7F7F7F"
    _hdr(ws, "A1", "Facility list")
    for i, name in enumerate(FACILITY_LIST, 2):
        ws.cell(row=i, column=1, value=name).border = BOX
    _hdr(ws, "C1", "Month list")
    for i, m in enumerate(MONTHS, 2):
        c = ws.cell(row=i, column=3, value=m)
        c.number_format = MONTH_FMT
        c.border = BOX
        c.alignment = Alignment(horizontal="left")
    ws["E1"] = "Dropdown sources. The Dashboard's selectors use Data Validation lists that point here:"
    ws["E2"] = "Facility: =Lists!$A$2:$A$5      Month: =Lists!$C$2:$C$25"
    ws["E3"] = ("The month cells hold real dates (the first of each month) formatted as mmm yyyy, so the dropdown "
                "shows 'Nov 2025' but the selector stores a date that SUMIFS can match.")
    for c in ("E1", "E2", "E3"):
        ws[c].font = NOTE_FONT
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 3
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 3
    ws.column_dimensions["E"].width = 100
    ws.freeze_panes = "A2"
    _landscape(ws)


def _targets_formats(wb):
    ws = wb["Targets"]
    hdr_row = 3
    headers = [c.value for c in ws[hdr_row]]
    tcol = headers.index("Target") + 1
    for i, t in enumerate(TARGETS):
        c = ws.cell(row=hdr_row + 1 + i, column=tcol)
        c.number_format = {"%": "0.0%", "minutes": "0", "days": "0.0"}[t[1]]
        for col in range(1, len(headers) + 1):
            ws.cell(row=hdr_row + 1 + i, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[hdr_row + 1 + i].height = 45


def _selector(ws, cell, value, fmt=None, list_ref=None, width_note=None):
    c = ws[cell]
    c.value = value
    c.fill = INPUT_FILL
    c.border = INPUT_BORDER
    c.font = Font(bold=True, size=12)
    c.alignment = Alignment(horizontal="left", vertical="center")
    if fmt:
        c.number_format = fmt
    if list_ref:
        dv = DataValidation(type="list", formula1=list_ref, allow_blank=False, showErrorMessage=True,
                            errorTitle="Pick from the list", error="Choose a value from the dropdown list.")
        ws.add_data_validation(dv)
        dv.add(cell)


def _dashboard(wb, lesson):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_properties.tabColor = "1F4E79"
    for col, w in {"A": 2, "B": 26, "C": 34, "D": 3, "E": 70}.items():
        ws.column_dimensions[col].width = w
    ws["B1"] = "Hospital Operations Dashboard · practice canvas"
    ws["B1"].font = Font(bold=True, size=18, color=NAVY)
    ws["B2"] = "Change the yellow selectors. Every number on this sheet, and your Practice formulas, should follow them."
    ws["B2"].font = NOTE_FONT
    for r, (label, cell) in enumerate([("Facility  (named SelFacility)", "C4"), ("Month  (named SelMonth)", "C5")], 4):
        ws[f"B{r}"] = label
        ws[f"B{r}"].font = Font(bold=True)
        ws[f"B{r}"].alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 20
    _selector(ws, "C4", SEL_FAC, list_ref="Lists!$A$2:$A$5")
    _selector(ws, "C5", SEL_MONTH, fmt=MONTH_FMT, list_ref="Lists!$C$2:$C$25")

    # Example card: Inpatient discharges.
    top = 7
    ws.merge_cells(f"B{top}:C{top}")
    _hdr(ws, f"B{top}", "INPATIENT DISCHARGES  ·  example card", width_cols=2)
    ws.merge_cells(f"B{top + 1}:C{top + 1}")
    lesson.set_formula(ws, f"B{top + 1}",
                       "=SUMIFS(tblKPI[IPDischarges],tblKPI[Facility],SelFacility,tblKPI[Month],SelMonth)", dynamic=False)
    big = ws[f"B{top + 1}"]
    big.font = Font(bold=True, size=24, color=NAVY)
    big.alignment = Alignment(horizontal="center", vertical="center")
    big.number_format = "#,##0"
    ws.row_dimensions[top + 1].height = 36
    ws[f"B{top + 2}"] = "Same month last year"
    lesson.set_formula(ws, f"C{top + 2}",
                       "=SUMIFS(tblKPI[IPDischarges],tblKPI[Facility],SelFacility,tblKPI[Month],EDATE(SelMonth,-12))",
                       dynamic=False)
    ws[f"C{top + 2}"].number_format = "#,##0"
    ws[f"B{top + 3}"] = "Change vs last year"
    lesson.set_formula(ws, f"C{top + 3}", f'=IFERROR(B{top + 1}/C{top + 2}-1,"n/a")', dynamic=False)
    ws[f"C{top + 3}"].number_format = '"▲ "0.0%;"▼ "0.0%;"▬ "0.0%'
    for r in range(top, top + 4):
        for col in "BC":
            ws[f"{col}{r}"].border = BOX
            if r > top:
                ws[f"{col}{r}"].fill = CARD_FILL
    for r in (top + 2, top + 3):
        ws[f"C{r}"].alignment = Alignment(horizontal="right")
    notes = [
        "← A finished card. Click each cell and read its formula in the formula bar.",
        "Big number: SUMIFS on the selectors.  Context: the same month last year (EDATE(SelMonth,-12)).",
        "The ▲/▼ comes from a custom number format, \"▲ \"0.0%;\"▼ \"0.0%;\"▬ \"0.0%, so the cell is still a number.",
        "It isn't colored: more discharges are neither good nor bad, so color would mislead.",
    ]
    for i, n in enumerate(notes):
        ws[f"E{top + i}"] = n
        ws[f"E{top + i}"].font = NOTE_FONT
    ws[f"B{top + 6}"] = "▸ Task 10: put your combo box here."
    ws[f"B{top + 8}"] = "▸ Task 11: paste your 12-month trend chart here (it reads Calc!B16:C28)."
    for r in (top + 6, top + 8):
        ws[f"B{r}"].font = Font(italic=True, color="7F7F7F")
    ws.freeze_panes = "A6"
    _landscape(ws)


def _calc(wb, lesson, link_cell, hdr_row, first, last, selftest, month_cells):
    ws = wb.create_sheet("Calc")
    ws.sheet_properties.tabColor = "548235"
    for col, w in {"A": 2, "B": 36, "C": 30, "D": 3, "E": 80}.items():
        ws.column_dimensions[col].width = w
    ws["B1"] = "Calc · the model layer"
    ws["B1"].font = Font(bold=True, size=16, color=NAVY)
    ws["B2"] = ("Formulas here turn the selectors into the numbers the Dashboard shows. Viewers never need to see "
                "this sheet, so a finished dashboard usually hides it.")
    ws["B2"].font = NOTE_FONT
    _hdr(ws, "B4", "Selection (read from the Dashboard)", width_cols=2)
    rows = [("Facility", "=SelFacility", None), ("Month", "=SelMonth", MONTH_FMT),
            ("Prior month", "=EDATE(SelMonth,-1)", MONTH_FMT), ("Same month last year", "=EDATE(SelMonth,-12)", MONTH_FMT),
            ("Facility criterion for SUMIFS", '=IF(SelFacility="All facilities","*",SelFacility)', None)]
    for i, (label, f, fmt) in enumerate(rows, 5):
        ws[f"B{i}"] = label
        lesson.set_formula(ws, f"C{i}", f, dynamic=False)
        ws[f"C{i}"].fill = PREFILL_FILL
        ws[f"C{i}"].alignment = Alignment(horizontal="left")
        if fmt:
            ws[f"C{i}"].number_format = fmt
        for col in "BC":
            ws[f"{col}{i}"].border = BOX
    ws["E9"] = 'Turns "All facilities" into the wildcard *, which SUMIFS matches to every hospital.'
    ws["E9"].font = NOTE_FONT

    _hdr(ws, "B11", "Form control (task 10)", width_cols=2)
    ws["B12"] = "Combo box cell link"
    c = ws[link_cell]
    c.fill = INPUT_FILL
    c.border = INPUT_BORDER
    c.alignment = Alignment(horizontal="left")
    ws["B13"] = "Facility picked in the combo box"
    lesson.set_formula(ws, "C13", f'=IF({link_cell}="","",INDEX(Lists!$A$2:$A$5,{link_cell}))', dynamic=False)
    ws["C13"].fill = PREFILL_FILL
    for r in (12, 13):
        for col in "BC":
            ws[f"{col}{r}"].border = BOX if not (col == "C" and r == 12) else INPUT_BORDER
    ws["E12"] = "Your combo box writes the position of the chosen item here (1 = first item in the list)."
    ws["E13"] = "INDEX turns the position back into the facility name."
    for c_ in ("E12", "E13"):
        ws[c_].font = NOTE_FONT

    _hdr(ws, f"B{hdr_row - 1}", "12-month trend for the chart (task 11)", width_cols=2)
    ws[f"B{hdr_row}"] = None
    ws[f"C{hdr_row}"] = "ED visits"
    ws[f"C{hdr_row}"].font = Font(bold=True)
    for col in "BC":
        ws[f"{col}{hdr_row}"].border = BOX
    for r in range(first, last + 1):
        for col, fmt in (("B", MONTH_FMT), ("C", "#,##0")):
            x = ws[f"{col}{r}"]
            x.fill = INPUT_FILL
            x.border = INPUT_BORDER
            x.number_format = fmt
            x.alignment = Alignment(horizontal="right")
    ws[f"E{hdr_row}"] = "B16 stays blank on purpose: with an empty corner cell, Excel uses column B as the chart's axis."
    ws[f"E{first}"] = "B: the 12 months ending at SelMonth, oldest first.  C: ED visits for SelFacility in each month."
    for c_ in (f"E{hdr_row}", f"E{first}"):
        ws[c_].font = NOTE_FONT
    ws.freeze_panes = "A4"
    _landscape(ws)

    if selftest:   # simulate the learner's month column (task 11 fills column C through its `fill`)
        first_, last_, formula = month_cells
        from openpyxl.formula.translate import Translator
        from xlcourse.xlfn import to_file_formula
        f0 = to_file_formula(formula)
        for r in range(first_, last_ + 1):
            ws[f"B{r}"] = f0 if r == first_ else Translator(f0, origin=f"B{first_}").translate_formula(f"B{r}")


def _board(wb):
    ws = wb.create_sheet("Board")
    ws.sheet_properties.tabColor = "BF9000"
    ws.column_dimensions["A"].width = 2
    ws["B1"] = "Board dashboard (bonus)"
    ws["B1"].font = Font(bold=True, size=18, color=NAVY)
    lines = [
        "Build the one-page board dashboard here. The Bonus sheet has the full spec. A layout that works:",
        "Rows 1–2: title and a dynamic subtitle (facility · month).   Rows 4–5: the two dropdowns (name them "
        "BoardFacility and BoardMonth).",
        "Rows 7–17: two rows of four cards.   Rows 19–33: the trend chart.   Columns N onward: your model cells "
        "(helpers, the KPI table, the 12-month block).",
        "Delete these notes when you're done, then hide gridlines and headings, unlock the dropdowns, protect the "
        "sheet, and set it to print on one landscape page.",
    ]
    for i, t in enumerate(lines, 2):
        ws[f"B{i}"] = t
        ws[f"B{i}"].font = NOTE_FONT
    _landscape(ws)


# ---------------------------------------------------------------------------------------------------------- reference
CARDS = [  # (title, model row, value format, kind)
    ("ED visits", 10, "#,##0", "count"),
    ("LWBS %", 11, "0.0%", "pct"),
    ("Median door-to-provider", 12, '0.0" min"', "min"),
    ("ALOS", 13, '0.00" days"', "days"),
    ("30-day readmission rate", 14, "0.0%", "pct"),
    ("Occupancy %", 15, "0.0%", "pct"),
    ("HCAHPS top-box %", 16, "0.0%", "pct"),
    ("Denial rate", 17, "0.0%", "pct"),
]
CARD_COLS = [("B", "C"), ("E", "F"), ("H", "I"), ("K", "L")]


def _dashboard_key(wb, lesson):
    ws = wb.create_sheet("Dashboard Key")
    ws.sheet_properties.tabColor = "C00000"
    sf = lambda cell, f: lesson.set_formula(ws, cell, f)  # noqa: E731
    widths = {"A": 2, "B": 14, "C": 17, "D": 2, "E": 14, "F": 17, "G": 2, "H": 14, "I": 17, "J": 2, "K": 14, "L": 17,
              "M": 4, "N": 32}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for col in "OPQRSTUV":
        ws.column_dimensions[col].width = 14

    # ---------------- header and selectors
    ws["B1"] = "Bluestone Health · Hospital Operations Dashboard"
    ws["B1"].font = Font(bold=True, size=20, color=NAVY)
    sf("B2", '="Selection: "&$C$4&"  ·  "&TEXT($C$5,"mmmm yyyy")&"  ·  reference build for the Lesson 4.6 bonus"')
    ws["B2"].font = NOTE_FONT
    for r, label in ((4, "Facility"), (5, "Month")):
        ws[f"B{r}"] = label
        ws[f"B{r}"].font = Font(bold=True)
        ws.merge_cells(f"C{r}:F{r}")
        ws.row_dimensions[r].height = 20
    _selector(ws, "C4", ALL, list_ref="Lists!$A$2:$A$5")
    _selector(ws, "C5", B_MONTH, fmt=MONTH_FMT, list_ref="Lists!$C$2:$C$25")
    for cell in ("C4", "C5"):
        ws[cell].protection = Protection(locked=False)
    ws.merge_cells("H4:L5")
    sf("H4", '=$O$18&" of "&$P$18&" KPIs on target"')
    ws["H4"].font = Font(bold=True, size=18, color=NAVY)
    ws["H4"].alignment = Alignment(horizontal="center", vertical="center")
    ws["H4"].fill = CARD_FILL

    # ---------------- model area (normally its own Calc sheet)
    ws["N1"] = "MODEL  ·  on a real dashboard this lives on its own (hidden) Calc sheet"
    ws["N1"].font = Font(bold=True, color="7B2C2C")
    model_rows = [
        (3, "Facility criterion", '=IF($C$4="All facilities","*",$C$4)', None),
        (4, "Month", "=$C$5", MONTH_FMT),
        (5, "Same month last year", "=EDATE($O$4,-12)", MONTH_FMT),
        (6, "Month end (exclusive)", "=EDATE($O$4,1)", "mm/dd/yyyy"),
        (7, "Last year's month end (exclusive)", "=EDATE($O$5,1)", "mm/dd/yyyy"),
    ]
    for r, label, f, fmt in model_rows:
        ws[f"N{r}"] = label
        sf(f"O{r}", f)
        if fmt:
            ws[f"O{r}"].number_format = fmt
        for col in "NO":
            ws[f"{col}{r}"].border = BOX
    for j, h in enumerate(["KPI", "This month", "Last year", "Target", "Direction", "On target (1/0)"]):
        _hdr(ws, f"{'NOPQRS'[j]}9", h)
    fac_ok = '((tblEDWaits[Facility]=$C$4)+($C$4="All facilities"))'

    def comp(num, den, m):
        return (f'=IFERROR(SUMIFS(tblKPI[{num}],tblKPI[Facility],$O$3,tblKPI[Month],{m})'
                f'/SUMIFS(tblKPI[{den}],tblKPI[Facility],$O$3,tblKPI[Month],{m}),"n/a")')

    def med(m, end):
        return (f'=IFERROR(MEDIAN(FILTER(tblEDWaits[DoorToProviderMin],(tblEDWaits[ArrivalDateTime]>={m})'
                f'*(tblEDWaits[ArrivalDateTime]<{end})*{fac_ok})),"n/a")')

    ws["N10"] = "ED visits"
    sf("O10", "=SUMIFS(tblKPI[EDVisits],tblKPI[Facility],$O$3,tblKPI[Month],$O$4)")
    sf("P10", "=SUMIFS(tblKPI[EDVisits],tblKPI[Facility],$O$3,tblKPI[Month],$O$5)")
    for col in "OP":
        ws[f"{col}10"].number_format = "#,##0"
    for i, t in enumerate(TARGETS):
        r = 11 + i
        ws[f"N{r}"] = t[0]
        if t[0] == "Median door-to-provider":
            sf(f"O{r}", med("$O$4", "$O$6"))
            sf(f"P{r}", med("$O$5", "$O$7"))
        else:
            sf(f"O{r}", comp(t[7], t[8], "$O$4"))
            sf(f"P{r}", comp(t[7], t[8], "$O$5"))
        sf(f"Q{r}", f"=XLOOKUP($N{r},tblTargets[KPI],tblTargets[Target])")
        sf(f"R{r}", f"=XLOOKUP($N{r},tblTargets[KPI],tblTargets[Direction])")
        sf(f"S{r}", f'=IF(ISNUMBER(O{r}),--IF(R{r}="Lower is better",O{r}<=Q{r},O{r}>=Q{r}),"")')
        fmt = {"%": "0.0%", "minutes": "0.0", "days": "0.00"}[t[1]]
        for col in "OPQ":
            ws[f"{col}{r}"].number_format = fmt
    for r in range(10, 18):
        for col in "NOPQRS":
            ws[f"{col}{r}"].border = BOX
    ws["N18"] = "KPIs on target (of those with data)"
    sf("O18", "=SUM(S11:S17)")
    sf("P18", "=COUNT(S11:S17)")
    for col in "NOP":
        ws[f"{col}18"].font = Font(bold=True)
        ws[f"{col}18"].border = BOX

    # 12-month trend block
    ws["N21"] = "12-month trend (selected facility)"
    ws["N21"].font = Font(bold=True, color=NAVY)
    for j, h in enumerate(["Month", "ED visits", "LWBS %", "LWBS target"]):
        _hdr(ws, f"{'NOPQ'[j]}22", h)
    for r in range(23, 35):
        sf(f"N{r}", f"=EDATE($O$4,ROWS(N$23:N{r})-12)")
        sf(f"O{r}", f"=IF(COUNTIFS(tblKPI[Facility],$O$3,tblKPI[Month],N{r})=0,NA(),"
                    f"SUMIFS(tblKPI[EDVisits],tblKPI[Facility],$O$3,tblKPI[Month],N{r}))")
        sf(f"P{r}", f"=IFERROR(SUMIFS(tblKPI[LWBS],tblKPI[Facility],$O$3,tblKPI[Month],N{r})"
                    f"/SUMIFS(tblKPI[EDVisits],tblKPI[Facility],$O$3,tblKPI[Month],N{r}),NA())")
        sf(f"Q{r}", "=$Q$11")
        ws[f"N{r}"].number_format = "mmm yy"
        ws[f"O{r}"].number_format = "#,##0"
        ws[f"P{r}"].number_format = "0.0%"
        ws[f"Q{r}"].number_format = "0.0%"
        for col in "NOPQ":
            ws[f"{col}{r}"].border = BOX
    ws["N36"] = "Peak LWBS % month in the window"
    sf("O36", "=INDEX(N23:N34,MATCH(AGGREGATE(4,6,P23:P34),P23:P34,0))")
    ws["O36"].number_format = "mm/dd/yyyy"
    for col in "NO":
        ws[f"{col}36"].border = BOX
        ws[f"{col}36"].font = Font(bold=True)

    # Hospital comparison for the selected month
    ws["N38"] = "Hospital comparison (selected month)"
    ws["N38"].font = Font(bold=True, color=NAVY)
    kcols = "OPQRSTU"
    _hdr(ws, "N39", "Facility")
    for j, t in enumerate(TARGETS):
        _hdr(ws, f"{kcols[j]}39", t[0])
        ws[f"{kcols[j]}39"].alignment = Alignment(wrap_text=True, vertical="top")
    _hdr(ws, "V39", "On target")
    ws.row_dimensions[39].height = 32
    ws["N40"], ws["N41"] = "Target", "Direction"
    for j, t in enumerate(TARGETS):
        col = kcols[j]
        sf(f"{col}40", f"=XLOOKUP({col}$39,tblTargets[KPI],tblTargets[Target])")
        sf(f"{col}41", f"=XLOOKUP({col}$39,tblTargets[KPI],tblTargets[Direction])")
        ws[f"{col}40"].number_format = {"%": "0.0%", "minutes": "0", "days": "0.0"}[t[1]]
        ws[f"{col}41"].font = Font(size=8, italic=True)
    for k, (_, name) in enumerate(FACILITIES):
        r = 42 + k
        ws[f"N{r}"] = name
        for j, t in enumerate(TARGETS):
            col = kcols[j]
            if t[0] == "Median door-to-provider":
                sf(f"{col}{r}", f'=IFERROR(XLOOKUP(1,(tblKPI[Facility]=$N{r})*(tblKPI[Month]=$O$4),'
                                f'tblKPI[MedianDTP]),"n/a")')
            else:
                sf(f"{col}{r}", f'=IFERROR(SUMIFS(tblKPI[{t[7]}],tblKPI[Facility],$N{r},tblKPI[Month],$O$4)'
                                f'/SUMIFS(tblKPI[{t[8]}],tblKPI[Facility],$N{r},tblKPI[Month],$O$4),"n/a")')
            ws[f"{col}{r}"].number_format = {"%": "0.0%", "minutes": "0.0", "days": "0.00"}[t[1]]
        sf(f"V{r}", f'=SUMPRODUCT(ISNUMBER(O{r}:U{r})*(((O$41:U$41="Lower is better")*(O{r}:U{r}<=O$40:U$40))'
                    f'+((O$41:U$41="Higher is better")*(O{r}:U{r}>=O$40:U$40))))')
        ws[f"V{r}"].font = Font(bold=True)
    for r in range(39, 45):
        for col in "NOPQRSTUV":
            ws[f"{col}{r}"].border = BOX
    rng = "O42:U44"
    ok = 'IF(O$41="Lower is better",O42<=O$40,O42>=O$40)'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER(O42),{ok})"], fill=GOOD_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER(O42),NOT({ok}))"], fill=BAD_FILL))
    ws["N46"] = "Fewest KPIs on target"
    sf("O46", "=INDEX(N42:N44,MATCH(MIN(V42:V44),V42:V44,0))")
    for col in "NO":
        ws[f"{col}46"].border = BOX
        ws[f"{col}46"].font = Font(bold=True)

    # ---------------- KPI cards
    for idx, (title, mrow, fmt, kind) in enumerate(CARDS):
        c1, c2 = CARD_COLS[idx % 4]
        top = 7 if idx < 4 else 13
        ws.merge_cells(f"{c1}{top}:{c2}{top}")
        _hdr(ws, f"{c1}{top}", title.upper(), width_cols=2)
        ws.merge_cells(f"{c1}{top + 1}:{c2}{top + 1}")
        sf(f"{c1}{top + 1}", f"=$O${mrow}")
        v = ws[f"{c1}{top + 1}"]
        v.number_format = fmt
        v.font = Font(bold=True, size=22, color=NAVY)
        v.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[top + 1].height = 34
        ws[f"{c1}{top + 2}"] = "Target"
        ws[f"{c1}{top + 3}"] = "Status"
        ws[f"{c1}{top + 4}"] = "vs last year"
        o, p = f"$O${mrow}", f"$P${mrow}"
        arrow = f"IF({o}>{p},UNICHAR(9650),IF({o}<{p},UNICHAR(9660),UNICHAR(9644)))"
        both = f"AND(ISNUMBER({o}),ISNUMBER({p}))"
        if kind == "count":
            ws[f"{c2}{top + 2}"] = "—"
            ws[f"{c2}{top + 3}"] = "—"
            sf(f"{c2}{top + 4}", f'=IF(AND({both},{p}<>0),{arrow}&" "&TEXT(ABS({o}/{p}-1),"0.0%"),"n/a")')
        else:
            tfmt = {"pct": "0.0%", "min": "0", "days": "0.0"}[kind]
            unit = {"pct": "", "min": " min", "days": " days"}[kind]
            sf(f"{c2}{top + 2}", f'=IF($R${mrow}="Lower is better","≤ ","≥ ")&TEXT($Q${mrow},"{tfmt}")&"{unit}"')
            sf(f"{c2}{top + 3}", f'=IF(ISNUMBER($S${mrow}),IF($S${mrow}=1,"✔ On target","✘ Off target"),"n/a")')
            if kind == "pct":
                change = f'TEXT(ABS({o}-{p})*100,"0.0")&" pts"'
            elif kind == "min":
                change = f'TEXT(ABS({o}-{p}),"0.0")&" min"'
            else:
                change = f'TEXT(ABS({o}-{p}),"0.00")&" days"'
            sf(f"{c2}{top + 4}", f'=IF({both},{arrow}&" "&{change},"n/a")')
            st = f"{c2}{top + 3}"
            ws.conditional_formatting.add(st, FormulaRule(formula=[f'LEFT({st},1)="✔"'], fill=GOOD_FILL,
                                                          font=Font(bold=True, color="006100")))
            ws.conditional_formatting.add(st, FormulaRule(formula=[f'LEFT({st},1)="✘"'], fill=BAD_FILL,
                                                          font=Font(bold=True, color="9C0006")))
            vs = f"{c2}{top + 4}"
            better = f'IF($R${mrow}="Lower is better",{o}<{p},{o}>{p})'
            worse = f'IF($R${mrow}="Lower is better",{o}>{p},{o}<{p})'
            ws.conditional_formatting.add(vs, FormulaRule(formula=[f"AND({both},{better})"],
                                                          font=Font(color="006100")))
            ws.conditional_formatting.add(vs, FormulaRule(formula=[f"AND({both},{worse})"],
                                                          font=Font(color="9C0006")))
        for r in range(top + 1, top + 5):
            for col in (c1, c2):
                ws[f"{col}{r}"].fill = CARD_FILL
                ws[f"{col}{r}"].border = BOX
            ws[f"{c2}{r}"].alignment = Alignment(horizontal="right")
        ws[f"{c2}{top + 1}"].alignment = Alignment(horizontal="center")
        for r in range(top + 2, top + 5):
            ws[f"{c1}{r}"].font = Font(color="595959", size=9)

    # ---------------- charts
    cats = Reference(ws, min_col=14, min_row=23, max_row=34)
    bar = BarChart()
    bar.type = "col"
    bar.title = "ED visits, last 12 months"
    bar.add_data(Reference(ws, min_col=15, min_row=22, max_row=34), titles_from_data=True)
    bar.set_categories(cats)
    bar.legend = None
    bar.y_axis.scaling.min = 0
    bar.y_axis.number_format = "#,##0"
    bar.y_axis.majorGridlines = None
    bar.x_axis.number_format = "mmm yy"
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bar.height, bar.width = 7.2, 15.8
    bar.series[0].graphicalProperties.solidFill = "2E75B6"
    bar.series[0].graphicalProperties.line.solidFill = "2E75B6"
    ws.add_chart(bar, "B19")

    line = LineChart()
    line.title = "LWBS % vs target, last 12 months"
    line.add_data(Reference(ws, min_col=16, max_col=17, min_row=22, max_row=34), titles_from_data=True)
    line.set_categories(cats)
    line.y_axis.scaling.min = 0
    line.y_axis.number_format = "0.0%"
    line.x_axis.number_format = "mmm yy"
    line.x_axis.delete = False
    line.y_axis.delete = False
    line.legend.position = "b"
    line.series[0].graphicalProperties.line.solidFill = "1F4E79"
    line.series[0].graphicalProperties.line.width = 28000
    line.series[0].smooth = False
    line.series[1].graphicalProperties.line.solidFill = "C00000"
    line.series[1].graphicalProperties.line.dashStyle = "dash"
    line.series[1].smooth = False
    line.height, line.width = 7.2, 15.8
    ws.add_chart(line, "H19")

    ws["B34"] = ("Every number and both charts follow the two selectors. Model cells start at column N →. "
                 "Status uses text and color, so it still reads in grayscale.")
    ws["B34"].font = NOTE_FONT

    # ---------------- polish
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False
    ws.sheet_view.zoomScale = 90
    ws.print_area = "A1:L34"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.protection.sheet = True
    ws.sheet_state = "hidden"
