"""Lesson 1.2 · Data Entry, AutoFill & Editing.

Most tasks are "do it in Excel" skills (AutoFill, Fill Series, Flash Fill, Ctrl+Enter, Paste Special,
Find & Replace). Each one is checked by a gray summary formula on the Practice sheet that reads the
learner's work on a worksheet built in the `customize` hook below; the self-test simulates that work
with `fill={"range": ..., "values": [...]}`. Every expected value is computed in Python from the data.
"""
from __future__ import annotations

from datetime import date, datetime, time, timedelta

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from xlcourse import Lesson, Task, data
from xlcourse.lesson import BOX, HEADER_FILL, INPUT_BORDER, INPUT_FILL, NAVY

CODE = "1.2"

MONTH_START = date(2025, 12, 1)
DAYS_IN_MONTH = 31
SHIFT_CODES = {"Day 12h": "D", "Night 12h": "N"}
NIGHT_NEW = "N12"


def _frac(minutes: int) -> float:
    """Excel time value (fraction of a day) for a number of minutes after midnight."""
    return minutes / 1440


def _whiteboard(name: str) -> str:
    """'Last, First' -> 'First L.' (the privacy-friendly label used on unit whiteboards)."""
    last, first = [p.strip() for p in name.split(",", 1)]
    return f"{first} {last[0]}."


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="01-foundations", slug="02-data-entry-autofill",
        title="Data Entry, AutoFill & Editing", level="Beginner", minutes=40,
        objectives=[
            "Recognize how Excel stores text, numbers, dates, times, and TRUE/FALSE",
            "Enter and edit data efficiently (F2, Ctrl+Enter, Ctrl+D, Alt+Enter)",
            "Create series with AutoFill, the Fill Series dialog, and Flash Fill",
            "Use Copy, Paste Special (values, transpose, formats), and Find & Replace",
            "Avoid classic traps: lost leading zeros, numbers stored as text, accidental dates",
        ],
        data_note="A December 2025 rotating-RN schedule for Medical-Surgical 4 West, a registration record, "
                  "a Q15 observation form, 4 West's bed list, supply room, and December admissions, plus a "
                  "January 2026 clinic template for the bonus.",
    )

    # ================================================================== source data
    patients = data.index(data.load("patients"), "PatientID")
    employees = data.index(data.load("employees"), "EmployeeID")
    encounters = data.load("encounters")
    ed_by_enc = {v["EncounterID"]: v for v in data.load("ed_visits")}

    # --- Entries: one real 4 West admission whose MRN starts with "00" (so the leading zeros matter)
    dec_4w = sorted((e for e in encounters if e["DeptID"] == "D110" and e["EncounterType"] == "Inpatient"
                     and e["AdmitDateTime"].year == 2025 and e["AdmitDateTime"].month == 12),
                    key=lambda e: e["AdmitDateTime"])
    reg_enc = next(e for e in dec_4w if patients[e["PatientID"]]["MRN"].startswith("00")
                   and e["EncounterID"] in ed_by_enc and patients[e["PatientID"]]["DOB"].year < 1990
                   and e["AdmitDateTime"].day >= 15)
    reg_pt = patients[reg_enc["PatientID"]]
    reg_ed = ed_by_enc[reg_enc["EncounterID"]]
    mrn_text = reg_pt["MRN"]                      # e.g. '00874160'
    mrn_number = int(mrn_text)                    # what Excel keeps if you type it plainly
    arrival = reg_ed["ArrivalDateTime"]
    member_typed = "4410287366519087"             # fictional 16-digit insurance member ID
    member_stored = int(member_typed[:15] + "0")  # Excel keeps 15 significant digits
    pod_bed_typed = "3-12"                        # ED pod 3, bed 12 -> Excel turns it into March 12
    pod_bed_date = date(2025, 3, 12)
    patient_name = f"{reg_pt['LastName']}, {reg_pt['FirstName']}"

    # (field, stored value, number format or None, what was typed)
    entries = [
        ("Patient name", patient_name, None, "Typed it"),
        ("MRN", mrn_number, None, f"Typed {mrn_text}"),
        ("Date of birth", reg_pt["DOB"], "mm/dd/yyyy", f"Typed {reg_pt['DOB'].month}/{reg_pt['DOB'].day}/{reg_pt['DOB'].year}"),
        ("ED arrival", arrival, "mm/dd/yyyy hh:mm", f"Typed {arrival.month}/{arrival.day}/{arrival.year} {arrival.hour}:{arrival.minute:02d}"),
        ("ED pod-bed", pod_bed_date, "d-mmm", f"Typed {pod_bed_typed} (pod 3, bed 12)"),
        ("Temperature (°F)", float(reg_ed["TempF"]), None, f"Typed {reg_ed['TempF']:.1f}"),
        ("Weight (lb)", f"{reg_pt['WeightLb']}", None, "Pasted from the bedside scale's export file"),
        ("Isolation needed", True, None, "Typed true"),
        ("Allergies", "Penicillin", None, "Typed it"),
        ("Insurance member ID", member_stored, None, f"Typed the 16-digit ID {member_typed}"),
        ("Copay", 25, '"$"#,##0', "Typed $25"),      # Excel applies a no-decimals currency format, so it shows $25
        ("Pain score (0–10)", reg_ed["PainScore"], None, "Typed it"),
    ]
    ENTRY_FIRST = 4
    entry_row = {f: ENTRY_FIRST + i for i, (f, *_rest) in enumerate(entries)}
    ENTRY_LAST = ENTRY_FIRST + len(entries) - 1

    def _is_number(v):  # what COUNT / the status bar's Numerical Count count (TRUE/FALSE and text are not numbers)
        return isinstance(v, (int, float, date, datetime)) and not isinstance(v, bool)
    n_numbers = sum(1 for _, v, _, _ in entries if _is_number(v))
    arrival_serial = data.excel_serial(arrival)

    # --- Schedule: 4 West's rotating RNs, December 2025, one code per day (D = day 12h, N = night 12h)
    shifts = [s for s in data.load("shifts") if s["DeptID"] == "D110" and s["ShiftDate"].year == 2025
              and s["ShiftDate"].month == 12]
    rn_ids = sorted({s["EmployeeID"] for s in shifts
                     if employees[s["EmployeeID"]]["JobTitle"] == "Registered Nurse"
                     and employees[s["EmployeeID"]]["ShiftPreference"] == "Rotating"},
                    key=lambda i: (employees[i]["LastName"], employees[i]["FirstName"]))
    nurses = [f"{employees[i]['FirstName']} {employees[i]['LastName']}" for i in rn_ids]
    days = [MONTH_START + timedelta(days=k) for k in range(DAYS_IN_MONTH)]
    grid = {(i, d): "" for i in rn_ids for d in days}
    for s in shifts:
        if s["EmployeeID"] in rn_ids:
            key = (s["EmployeeID"], s["ShiftDate"])
            assert grid[key] == "", "one shift per nurse per day"
            grid[key] = SHIFT_CODES[s["ShiftType"]]
    SCH_HDR, SCH_FIRST = 4, 5
    SCH_LAST = SCH_FIRST + DAYS_IN_MONTH - 1
    NURSE_COL0 = 3                                   # column C
    nurse_col = {i: get_column_letter(NURSE_COL0 + k) for k, i in enumerate(rn_ids)}
    LAST_NURSE_COL = get_column_letter(NURSE_COL0 + len(rn_ids) - 1)
    FLOAT_COL = get_column_letter(NURSE_COL0 + len(rn_ids))
    sched_rows = [[grid[(i, d)] for i in rn_ids] for d in days]
    n_nights = sum(row.count("N") for row in sched_rows)
    weekend = [d.weekday() >= 5 for d in days]
    n_weekend = sum(weekend)
    code_rng = f"Schedule!C{SCH_FIRST}:{LAST_NURSE_COL}{SCH_LAST}"
    date_rng = f"Schedule!A{SCH_FIRST}:A{SCH_LAST}"
    float_rng = f"Schedule!{FLOAT_COL}{SCH_FIRST}:{FLOAT_COL}{SCH_LAST}"

    # transpose target: Roster!A3 -> 'Date' row, 'Day' row, then one row per nurse
    ROSTER_ANCHOR = 3
    check_k = 6 if len(rn_ids) > 6 else 0
    check_id = rn_ids[check_k]
    check_name = nurses[check_k]
    check_row = ROSTER_ANCHOR + 2 + check_k
    check_shifts = sum(1 for d in days if grid[(check_id, d)])
    roster_last_col = get_column_letter(1 + DAYS_IN_MONTH)          # A + 31 date columns
    roster_last_row = ROSTER_ANCHOR + 2 + len(rn_ids) - 1

    # --- Q15 observation log: 07:00 to 18:45 every 15 minutes
    Q_FIRST = 4
    q_times = list(range(7 * 60, 18 * 60 + 45 + 1, 15))
    Q_LAST = Q_FIRST + len(q_times) - 1
    q_last_time = timedelta(minutes=q_times[-1])

    # --- Beds: 4 West has 36 staffed beds, two per room
    beds_n = next(d for d in data.load("departments") if d["DeptID"] == "D110")["StaffedBeds"]
    bed_ids = [f"4W-{k:02d}" for k in range(1, beds_n + 1)]
    B_FIRST = 4
    B_LAST = B_FIRST + beds_n - 1

    # --- Admissions (Flash Fill): first 24 December admissions to 4 West
    adm = []
    for e in dec_4w[:24]:
        p = patients[e["PatientID"]]
        adm.append({"Admitted": e["AdmitDateTime"], "EncounterID": e["EncounterID"],
                    "PatientName": f"{p['LastName']}, {p['FirstName']}"})
    wb_labels = [_whiteboard(r["PatientName"]) for r in adm]

    def _initials(name: str) -> tuple[str, str]:
        last, first = [p.strip() for p in name.split(",", 1)]
        return first[0].upper(), last[0].upper()
    # Row 1's first and last names share an initial (e.g. Richardson, Rebecca -> "Rebecca R."), so one example
    # can't tell Flash Fill which initial is wanted. The learner types a second, unambiguous example from row 2.
    ff_ambiguous = len(set(_initials(adm[0]["PatientName"]))) == 1
    if ff_ambiguous:
        assert len(set(_initials(adm[1]["PatientName"]))) == 2, "row 2 must disambiguate the Flash Fill pattern"
    ff_ex = 1 if ff_ambiguous else 0          # the example shown in the prompt is always unambiguous

    # --- Supplies (Paste Special → Values): 4 West supply room, restock-to-par suggestion
    supplies = sorted((r for r in data.load("supply_inventory") if r["LocationDeptID"] == "D110"), key=lambda r: r["SKU"])
    SUP_HDR = 3
    sup_first = SUP_HDR + 1
    for k, r in enumerate(supplies):
        rr = sup_first + k
        r["SuggestedOrder"] = f"=MAX(0,G{rr}-F{rr})"
    suggested = [max(0, r["ParLevel"] - r["QtyOnHand"]) for r in supplies]
    total_order = sum(suggested)

    # ================================================================== data sheets (Excel Tables)
    sup = L.add_table_sheet(
        "Supplies", supplies, table="tblSupplies", start_row=SUP_HDR,
        columns=["SKU", "ItemDescription", "Vendor", "UnitOfMeasure", "UnitCost", "QtyOnHand", "ParLevel", "SuggestedOrder"],
        formats={"UnitCost": "#,##0.00"}, widths={"ItemDescription": 40, "Vendor": 26, "SuggestedOrder": 16},
        notes=["4 West supply room · restock-to-par list · 12/31/2025",
               "SuggestedOrder is a formula: =MAX(0, ParLevel − QtyOnHand). Click any cell in that column and read the formula bar."],
    )
    assert sup.first_row == sup_first and sup.col("SuggestedOrder") == "H"
    sug_rng = f"Supplies!H{sup.first_row}:H{sup.last_row}"

    ad = L.add_table_sheet(
        "Admissions", adm, table="tblAdmissions", columns=["Admitted", "EncounterID", "PatientName"],
        extra_cols=["WhiteboardName"], formats={"Admitted": "mm/dd/yyyy hh:mm"},
        widths={"Admitted": 18, "PatientName": 24, "WhiteboardName": 20},
    )
    a_name = f"Admissions!{ad.col('PatientName')}{ad.first_row}:{ad.col('PatientName')}{ad.last_row}"
    a_wb = f"Admissions!{ad.col('WhiteboardName')}{ad.first_row}:{ad.col('WhiteboardName')}{ad.last_row}"

    # Order sheet layout (built in customize): rows line up with the Supplies table rows on purpose
    O_FIRST = sup.first_row
    O_LAST = sup.last_row
    order_rng = f"Order!D{O_FIRST}:D{O_LAST}"

    # ================================================================== bonus grid (computed in Python)
    clinic_days = []
    d = date(2026, 1, 5)
    while len(clinic_days) < 10:
        if d.weekday() < 5:
            clinic_days.append(d)
        d += timedelta(days=1)
    slot_mins = list(range(8 * 60, 16 * 60 + 45 + 1, 15))
    G_HDR = 3
    G_FIRST = G_HDR + 1
    G_LAST = G_FIRST + len(slot_mins) - 1
    G_COL0 = 2  # column B
    gcol = {dd: get_column_letter(G_COL0 + k) for k, dd in enumerate(clinic_days)}
    G_LASTCOL = gcol[clinic_days[-1]]
    body = f"'Clinic Grid'!B{G_FIRST}:{G_LASTCOL}{G_LAST}"

    def slot_state(dd: date, m: int) -> str:
        if 12 * 60 <= m < 13 * 60:
            return "Lunch"
        if m >= 13 * 60 and dd.weekday() == 2:
            return "Admin"
        if m >= 13 * 60 and dd.weekday() == 4:
            return "Telehealth"
        return "Open"
    final = [[slot_state(dd, m) for dd in clinic_days] for m in slot_mins]
    flat_final = [v for row in final for v in row]
    n_lunch = flat_final.count("Lunch")
    n_admin = flat_final.count("Admin")
    n_tele = flat_final.count("Telehealth")
    n_open = flat_final.count("Open")
    lunch_r1 = G_FIRST + slot_mins.index(12 * 60)
    lunch_r2 = G_FIRST + slot_mins.index(12 * 60 + 45)
    pm_r1 = G_FIRST + slot_mins.index(13 * 60)
    weds = [dd for dd in clinic_days if dd.weekday() == 2]
    fris = [dd for dd in clinic_days if dd.weekday() == 4]
    lunch_ok = f"COUNTIF('Clinic Grid'!B{lunch_r1}:{G_LASTCOL}{lunch_r2},\"Lunch\")"
    admin_ok = "+".join(f"COUNTIF('Clinic Grid'!{gcol[w]}{pm_r1}:{gcol[w]}{G_LAST},\"Admin\")" for w in weds)
    tele_ok = "+".join(f"COUNTIF('Clinic Grid'!{gcol[f]}{pm_r1}:{gcol[f]}{G_LAST},\"Telehealth\")" for f in fris)
    provider = next(p for p in data.load("providers") if p["ProviderID"] == "PRV1125")
    doc = f"{provider['FirstName']} {provider['LastName']}, {provider['Credential']}"   # not "Dr. ..., MD" (redundant)

    # ================================================================== tasks
    L.practice_intro = ("Most tasks are done on the other sheets (Entries, Schedule, Q15 Log, Beds, Admissions, Supplies, Order, "
                        "Roster). Type in the yellow cells here. A gray cell is a pre-filled formula that reads your work "
                        "on another sheet, and its Check turns green when that work is right. Do the tasks in order, because the Schedule tasks build on "
                        "each other.")

    def sched_cell(col, k):
        return f"Schedule!{col}{SCH_FIRST + k}"

    L.tasks = [
        # ---------------------------------------------------------------- 1 · what Excel stores
        Task(f"On the Entries sheet, select the Entry column (B{ENTRY_FIRST}:B{ENTRY_LAST}). How many of those 12 entries did Excel store "
             f"as numbers? Dates and times count as numbers. Read the status bar (right-click it to turn on Numerical Count) "
             f"or use the alignment clue.",
             answer=n_numbers, title="How many entries are stored as numbers?",
             solution=f"1. Select **Entries!B{ENTRY_FIRST}:B{ENTRY_LAST}**.\n"
                      "2. Right-click the status bar and tick **Numerical Count** if it isn't showing.\n"
                      "3. Read Numerical Count (Count shows 12, because it counts every non-empty cell).",
             # No live formula: Excel's =COUNT(Entries!B4:B15) ignores TRUE/FALSE and returns the Python answer, but
             # LibreOffice stores booleans as numbers and counts the TRUE cell too (verified per cell: 8 numbers + 1 logical).
             live=False,
             hint="By default, numbers line up on the right of a cell and text on the left",
             explanation=f"Two entries surprise most people. The pod-bed `{pod_bed_typed}` became a **date** (March 12), so it's a number. "
                         f"The weight was pasted in as **text**, so it isn't a number even though it looks like one. That's why it sits on "
                         "the left with a green triangle. TRUE is a logical value (centered), and the name and allergy are text. "
                         "Everything else, including the date of birth, the arrival time, the copay, and the long member ID, is a number."),
        Task(f"Entries!B{entry_row['ED arrival']} shows the ED arrival as a date and time. What number does Excel actually store in that "
             f"cell? Switch the cell to General format to see it, and enter it rounded to 2 decimal places.",
             answer=round(arrival_serial, 2), fmt="0.00", tol=0.0051, title="Stored value of a date and time",
             solution=f"Select **Entries!B{entry_row['ED arrival']}** and press **Ctrl + Shift + ~** (Mac: **⌃ + Shift + ~**), or choose "
                      "**Home → Number Format → General**. Read the number, then press **Ctrl + Z** (Mac: **⌘ + Z**) to put the date "
                      "format back.",
             live=f"=Entries!B{entry_row['ED arrival']}",
             hint="The whole part counts days and the decimal part is the time of day",
             explanation=f"Excel stores {arrival:%m/%d/%Y} as the serial number {int(arrival_serial):,} (day 1 is 1/1/1900). "
                         f"The time {arrival:%H:%M} is {arrival.hour * 60 + arrival.minute:,} minutes out of 1,440 in a day, which is "
                         f"{(arrival.hour * 60 + arrival.minute) / 1440:.4f}. Add them and you get {arrival_serial:,.4f}. Formatting only "
                         "changes how that one number is displayed."),
        # ---------------------------------------------------------------- 2 · traps and entry keys
        Task(f"The MRN on the Entries sheet lost its leading zeros. In the yellow cell, enter the patient's real 8-digit MRN, "
             f"{mrn_text}, so that Excel keeps it as text with both zeros.",
             answer=mrn_text, title="Enter an MRN with leading zeros",
             solution=f"Type **'{mrn_text}** (start with an apostrophe) and press **Enter**. The apostrophe tells Excel \"this is text\" and "
                      "is not shown in the cell.",
             live=f'=TEXT(Entries!B{entry_row["MRN"]},"00000000")',
             hint="Start the entry with an apostrophe",
             explanation=f"Typed plainly, `{mrn_text}` becomes the number {mrn_number}, because numbers don't have leading zeros. "
                         "An apostrophe (or formatting the cell as **Text** *before* you type) stores the characters exactly as typed. "
                         "A custom number format like `00000000` only *displays* zeros on top of the number, so the stored value would "
                         "still be wrong for matching and lookups. (The live formula in the key rebuilds the text with TEXT, which you'll "
                         "meet in Lesson 2.2.)"),
        Task("Type this two-line handoff note into the yellow cell, with \"Allergy: Penicillin\" on the first line and "
             "\"Isolation: Contact\" on the second line of the same cell.",
             answer="Allergy: Penicillin\nIsolation: Contact", check="custom",
             custom_check='SUBSTITUTE(LOWER({cell}&"")," ","")=SUBSTITUTE(LOWER({key}&"")," ","")',
             answer_display="Allergy: Penicillin ⏎ Isolation: Contact (two lines in one cell)",
             title="A two-line note with a line break",
             solution="Type `Allergy: Penicillin`, press **Alt + Enter** (Mac: **⌃ + ⌥ + Return**), type `Isolation: Contact`, "
                      "then press **Enter**.",
             live=False, hint="Enter on its own leaves the cell, so you need a different key combination",
             explanation="Alt + Enter inserts a line-break character inside the cell and turns on Wrap Text for you. Pressing "
                         "Enter on its own would finish the entry and jump to the next cell. The check ignores capitals and spaces, "
                         "but it needs the line break."),
        # ---------------------------------------------------------------- 3 · AutoFill
        Task(f"On the Schedule sheet, A{SCH_FIRST} holds 12/01/2025. Use the fill handle to fill the yellow cells "
             f"A{SCH_FIRST + 1}:A{SCH_LAST} with the rest of December, one day per row. The gray cell shows your last date.",
             answer=days[-1], fmt="mm/dd/yyyy", title="AutoFill the December dates",
             solution=f"Select **A{SCH_FIRST}**, point at the fill handle (the small square at its bottom-right corner) until the "
                      f"pointer becomes a thin **+**, and drag down to **A{SCH_LAST}**.",
             summary=f'=IF(COUNT(Schedule!A{SCH_FIRST + 1}:A{SCH_LAST})=0,"",Schedule!A{SCH_LAST})',
             fill={"range": f"Schedule!A{SCH_FIRST + 1}:A{SCH_LAST}", "values": days[1:]},
             live=f"=Schedule!A{SCH_FIRST}+{DAYS_IN_MONTH - 1}",
             hint="Drag the fill handle. A single date counts up one day at a time",
             explanation="AutoFill recognizes a date and adds one day per cell. If every cell shows 12/01/2025 instead, you held "
                         "Ctrl while dragging (which copies) or chose Copy Cells from the Auto Fill Options button."),
        Task(f"B{SCH_FIRST} holds \"Mon\". Fill the weekday names down to B{SCH_LAST} by double-clicking the fill handle "
             f"instead of dragging. The gray cell shows the day name in B{SCH_LAST}.",
             answer=days[-1].strftime("%a"), accept=[days[-1].strftime("%A")], title="AutoFill weekday names (double-click)",
             solution=f"Select **B{SCH_FIRST}** and **double-click** its fill handle. Excel fills down as far as the dates in the "
                      "neighboring column A go.",
             summary=f'=IF(COUNTA(Schedule!B{SCH_FIRST + 1}:B{SCH_LAST})=0,"",Schedule!B{SCH_LAST})',
             fill={"range": f"Schedule!B{SCH_FIRST + 1}:B{SCH_LAST}", "values": [d.strftime("%a") for d in days[1:]]},
             live=f'=TEXT(Schedule!A{SCH_FIRST}+{DAYS_IN_MONTH - 1},"ddd")',
             hint="Finish task 5 first: double-click fills as far as the neighboring column goes",
             explanation="Weekday names (Mon, Tue… and Monday, Tuesday…) are a built-in **custom list**, so AutoFill cycles through "
                         "them. Double-clicking the fill handle copies down to the last row of the adjacent column's data, which is "
                         "why column A had to be filled first."),
        Task(f"On the Beds sheet, A{B_FIRST} holds the first bed label, 4W-01. Fill the yellow cells below it so the list runs "
             f"through all {beds_n} of 4 West's staffed beds. The gray cell shows the label in the last row (A{B_LAST}).",
             answer=bed_ids[-1], title="AutoFill an ID series (4W-01 …)",
             solution=f"Select **Beds!A{B_FIRST}** and double-click the fill handle (column B is full, so Excel fills to row {B_LAST}), "
                      f"or drag it down to **A{B_LAST}**.",
             summary=f'=IF(COUNTA(Beds!A{B_FIRST + 1}:A{B_LAST})=0,"",Beds!A{B_LAST})',
             fill={"range": f"Beds!A{B_FIRST + 1}:A{B_LAST}", "values": bed_ids[1:]},
             live=False, hint="AutoFill increases the number at the end of a text entry",
             explanation="When an entry is text that ends in a number, AutoFill increases that number and keeps the rest of the text, "
                         "including the leading zero (4W-01, 4W-02 … 4W-36). Text with no number in it is just copied."),
        Task(f"On the Q15 Log sheet, A{Q_FIRST} holds 07:00. Fill the yellow cells below it with a check time every 15 minutes, "
             f"ending at 18:45. Use Home → Fill → Series, or the two-cell AutoFill pattern. The gray cell shows the latest time "
             f"in column A.",
             answer=q_last_time, fmt="hh:mm", title="Fill Series: Q15 check times from 07:00 to 18:45",
             solution=f"**Option A (Fill Series):** select **A{Q_FIRST}**, choose **Home → Fill → Series…**, pick **Columns** and "
                      "**Linear**, type **0:15** as the Step value and **18:50** as the Stop value, then click **OK**.\n\n"
                      f"**Option B (AutoFill):** type **7:15** in A{Q_FIRST + 1}, select A{Q_FIRST}:A{Q_FIRST + 1}, and drag the fill "
                      f"handle down to **A{Q_LAST}** (18:45).",
             summary=f"=IF(COUNT('Q15 Log'!A{Q_FIRST + 1}:A{Q_LAST + 20})=0,\"\",MAX('Q15 Log'!A{Q_FIRST}:A{Q_LAST + 20}))",
             fill={"range": f"'Q15 Log'!A{Q_FIRST + 1}:A{Q_LAST}", "values": [_frac(m) for m in q_times[1:]]},
             live=False, hint="Step value 0:15, with the Stop value a few minutes past the last time",
             explanation=f"Times are fractions of a day, so 15 minutes is 0:15 (0.0104…). The Series dialog adds that step until it "
                         "reaches the Stop value. A stop of 18:50 rather than 18:45 protects you from tiny rounding errors that can drop "
                         f"the last time. With two starting cells, AutoFill copies the gap between them. Either way you get "
                         f"{len(q_times)} check times. A single time dragged on its own steps by a whole **hour**, not 15 minutes."),
        # ---------------------------------------------------------------- 4 · selection tricks
        Task(f"4 West gets one float-pool RN on every Saturday and Sunday. In the Float RN column ({FLOAT_COL}) of the Schedule "
             f"sheet, put FLOAT in every weekend row and nowhere else, using a single entry: Ctrl+click (Mac: ⌘+click) the weekend "
             f"cells, type FLOAT, and press Ctrl+Enter (Mac: ⌘+Return). The gray cell counts correctly placed FLOATs minus any "
             f"entries on weekdays.",
             answer=n_weekend, title="Ctrl+Enter into a non-adjacent selection",
             solution=f"1. Click the first weekend cell in column {FLOAT_COL} (the first Sat row), then **Ctrl+click** "
                      "(Mac: **⌘+click**) every other Sat and Sun row.\n"
                      "2. Type `FLOAT` (it appears in the last cell you clicked).\n"
                      "3. Press **Ctrl + Enter** (Mac: **⌘ + Return**) to enter it in every selected cell at once.",
             summary=(f'=IF(COUNTA({float_rng})=0,"",SUMPRODUCT(({float_rng}="FLOAT")*({date_rng}<>"")*(WEEKDAY({date_rng},2)>5))'
                      f'-SUMPRODUCT(({float_rng}<>"")*(1-({date_rng}<>"")*(WEEKDAY({date_rng},2)>5))))'),
             fill={"range": float_rng, "values": ["FLOAT" if w else None for w in weekend]},
             live=False, hint="Finish tasks 5–6 first so you can see which rows are weekends",
             explanation="Ctrl + Enter puts the same entry into every selected cell, even when the cells aren't next to each other. "
                         "Plain Enter would fill only the active cell. The gray formula subtracts any FLOAT typed on a weekday, so it "
                         "only reaches the full count when every weekend row is marked and no weekday is."),
        Task(f"On the Admissions sheet, use Flash Fill to fill the yellow WhiteboardName column with each patient's first name and "
             f"last initial, such as \"{wb_labels[ff_ex]}\" for \"{adm[ff_ex]['PatientName']}\". The gray cell counts how many of the "
             f"{len(adm)} labels are exactly right.",
             answer=len(adm), title="Flash Fill whiteboard names (First L.)",
             solution=((f"1. In **Admissions!{ad.col('WhiteboardName')}{ad.first_row}**, type `{wb_labels[0]}` and press **Enter**.\n"
                        f"2. In **{ad.col('WhiteboardName')}{ad.first_row + 1}**, type `{wb_labels[1]}` and press **Enter**. "
                        "If Excel shows a gray preview of the remaining labels while you type, you can press Enter to accept it.\n"
                        "3. Otherwise press **Ctrl + E** (or choose **Data → Flash Fill**, on both Windows and Mac).\n"
                        "4. Scan the results, and correct any row Flash Fill got wrong.")
                       if ff_ambiguous else
                       (f"1. In **Admissions!{ad.col('WhiteboardName')}{ad.first_row}**, type `{wb_labels[0]}` and press **Enter**.\n"
                        "2. Press **Ctrl + E** (or choose **Data → Flash Fill**, on both Windows and Mac).\n"
                        "3. Scan the results, and correct any row Flash Fill got wrong.")),
             summary=(f'=IF(COUNTA({a_wb})=0,"",SUMPRODUCT(--({a_wb}=MID({a_name},FIND(",",{a_name})+2,50)'
                      f'&" "&LEFT({a_name},1)&".")))'),
             fill={"range": a_wb, "values": wb_labels},
             live=False, hint="Type an example that can only mean one thing, then press Ctrl + E",
             explanation=("Flash Fill studies your examples, finds the pattern (\"the text after the comma, a space, the first letter, "
                          "a period\"), and applies it to every row. "
                          + (f"The first row needs help. In `{adm[0]['PatientName']}` both names start with "
                             f"{_initials(adm[0]['PatientName'])[0]}, so `{wb_labels[0]}` doesn't show *which* initial you want, and "
                             f"Flash Fill could turn `{adm[1]['PatientName']}` into "
                             f"`{adm[1]['PatientName'].split(', ')[1]} {_initials(adm[1]['PatientName'])[0]}.`. The second "
                             f"example, `{wb_labels[1]}`, settles it. "
                             if ff_ambiguous else "")
                          + "The results are typed-in values, not formulas, so they won't update if a name changes. A formula "
                            "(Lesson 2.2) would. Many units use first name plus last initial on hallway whiteboards to protect "
                            "patient privacy.")),
        # ---------------------------------------------------------------- 5 · Paste Special and Find & Replace
        Task(f"The buyer needs 4 West's suggested restock quantities on the Order sheet as plain numbers. Copy "
             f"Supplies!H{sup.first_row}:H{sup.last_row} (SuggestedOrder) and paste only the values into the yellow cells "
             f"Order!D{O_FIRST}:D{O_LAST}. The gray cell totals your Order Qty column.",
             answer=total_order, title="Paste Special → Values",
             solution=f"1. Select **Supplies!H{sup.first_row}:H{sup.last_row}** and press **Ctrl + C** (Mac: **⌘ + C**).\n"
                      f"2. Click **Order!D{O_FIRST}**.\n"
                      "3. Press **Ctrl + Alt + V** (Mac: **⌃ + ⌘ + V**) to open Paste Special, choose **Values**, and click **OK**. "
                      "Or use **Home → Paste ▾ → Values (123)**.",
             summary=f'=IF(COUNTA({order_rng})=0,"",SUM({order_rng}))',
             fill={"range": order_rng, "values": suggested},
             live=f"=SUM({sug_rng})",
             hint="A normal paste brings the formulas, and their references move",
             explanation=f"SuggestedOrder holds formulas such as `=MAX(0,G{sup.first_row}-F{sup.first_row})`. A normal paste copies the *formula*, and because its "
                         "references are relative, it ends up pointing at the wrong cells on the Order sheet, so you see #VALUE! or wrong "
                         "numbers. Paste Special → Values pastes only the *results*, which is what you want when the numbers must stay "
                         "fixed or leave the workbook."),
        Task(f"Most people read schedules with staff down the side and dates across the top. Copy Schedule!A{SCH_HDR}:"
             f"{LAST_NURSE_COL}{SCH_LAST} and use Paste Special → Transpose with the top-left corner in Roster!A{ROSTER_ANCHOR}. "
             f"The gray cell checks that {check_name} landed in column A, then counts that nurse's December shifts.",
             answer=check_shifts, title="Paste Special → Transpose",
             solution=f"1. Select **Schedule!A{SCH_HDR}:{LAST_NURSE_COL}{SCH_LAST}** and press **Ctrl + C** (Mac: **⌘ + C**).\n"
                      f"2. Click **Roster!A{ROSTER_ANCHOR}**.\n"
                      "3. Press **Ctrl + Alt + V** (Mac: **⌃ + ⌘ + V**), tick **Transpose**, and click **OK**. "
                      "Or use **Home → Paste ▾ → Transpose**.",
             summary=(f'=IF(COUNTA(Roster!A{ROSTER_ANCHOR}:AH{roster_last_row + 5})=0,"",IF(Roster!A{check_row}="{check_name}",'
                      f'COUNTA(Roster!B{check_row}:{roster_last_col}{check_row}),"Not transposed into A{ROSTER_ANCHOR} yet"))'),
             fill={"range": f"Roster!A{ROSTER_ANCHOR}:{roster_last_col}{roster_last_row}",
                   "values": (["Date"] + days + ["Day"] + [d.strftime("%a") for d in days]
                              + [v for k, i in enumerate(rn_ids) for v in [nurses[k]] + [grid[(i, d)] or None for d in days]])},
             live=f"=COUNTA(Schedule!{nurse_col[check_id]}{SCH_FIRST}:{nurse_col[check_id]}{SCH_LAST})",
             hint="It's a checkbox in the Paste Special dialog",
             explanation=f"Transpose turns the copied block on its side: row {SCH_HDR} (the names) becomes column A, and each date "
                         f"row becomes a column. {check_name}'s column of codes becomes row {check_row}. Transpose isn't available "
                         "after **Cut**, and the paste area must not overlap the copied cells."),
        Task(f"The new scheduling system uses N12 instead of N for a 12-hour night shift. On the Schedule sheet, use Find & Replace "
             f"to change every N code to N12 without touching any names or day labels. The gray cell counts N12 codes in the grid "
             f"(C{SCH_FIRST}:{LAST_NURSE_COL}{SCH_LAST}) and warns you if a name or day label changed too.",
             answer=n_nights, title="Find & Replace with Match entire cell contents",
             solution="1. Click any single cell on the Schedule sheet (so Excel searches the whole sheet).\n"
                      "2. Press **Ctrl + H** (Mac: **⌃ + H**). Find what: `N`, Replace with: `N12`.\n"
                      "3. Tick **Match entire cell contents** (click **Options >>** first if you can't see it). On a Mac the box is "
                      "called **Find entire cells only**.\n"
                      "4. Click **Replace All**. Excel reports how many replacements it made, which matches the gray cell.",
             summary=(f'=IF(COUNTIF({code_rng},"{NIGHT_NEW}")=0,"",IF(COUNTIF(Schedule!A{SCH_HDR}:{FLOAT_COL}{SCH_HDR},"*{NIGHT_NEW}*")'
                      f'+COUNTIF(Schedule!B{SCH_FIRST}:B{SCH_LAST},"*{NIGHT_NEW}*")>0,"Undo: names or days changed too",'
                      f'COUNTIF({code_rng},"{NIGHT_NEW}")))'),
             fill={"range": code_rng, "values": [(NIGHT_NEW if v == "N" else (v or None)) for row in sched_rows for v in row]},
             live=f'=COUNTIF({code_rng},"N")',
             hint="Look under Options >> before you click Replace All",
             explanation="Without **Match entire cell contents**, Excel replaces the letter n *anywhere* in a cell, and it ignores "
                         "capitals, so \"Mon\" would become \"MoN12\" and every name with an n would be damaged. With the option ticked, "
                         "only cells containing exactly N change. Selecting just the code grid before Replace All is another safe "
                         "approach, because Excel then searches only the selection."),
    ]

    # ================================================================== bonus
    L.bonus_title = "Bonus: Build a two-week clinic template"
    L.bonus_scenario = (
        f"{doc}, at the Primary Care Clinic in the Bluestone Outpatient Pavilion, needs an appointment template for the two weeks "
        f"of Monday 01/05/2026 through Friday 01/16/2026. The clinic books 15-minute slots from 08:00 to 16:45 on weekdays only. "
        "Lunch (12:00–12:45) is blocked every day, Wednesday afternoons (13:00 onward) are admin time, and Friday afternoons "
        "become video visits. Build the whole grid on the Clinic Grid sheet with AutoFill, Fill Series, Ctrl+Enter, and Find & "
        f"Replace, without typing cell by cell. B{G_HDR} (01/05/2026) and A{G_FIRST} (08:00) are filled in for you.")
    L.bonus = [
        Task(f"Fill the date header C{G_HDR}:{G_LASTCOL}{G_HDR} with the next nine weekdays, skipping Saturdays and Sundays. "
             f"The gray cell shows the date in {G_LASTCOL}{G_HDR}.",
             answer=clinic_days[-1], fmt="mm/dd/yyyy", title="Weekday-only date header",
             solution=f"Select **B{G_HDR}:{G_LASTCOL}{G_HDR}**, then choose **Home → Fill → Series…**: Rows, Type **Date**, Date unit "
                      f"**Weekday**, Step 1, and click **OK**. Or drag B{G_HDR}'s fill handle to {G_LASTCOL}{G_HDR} and choose "
                      "**Auto Fill Options → Fill Weekdays**.",
             summary=f"=IF(COUNT('Clinic Grid'!C{G_HDR}:{G_LASTCOL}{G_HDR})=0,\"\",'Clinic Grid'!{G_LASTCOL}{G_HDR})",
             fill={"range": f"'Clinic Grid'!C{G_HDR}:{G_LASTCOL}{G_HDR}", "values": clinic_days[1:]},
             live=f"=WORKDAY('Clinic Grid'!B{G_HDR},{len(clinic_days) - 1})",
             hint="Date unit: Weekday (or Fill Weekdays)",
             explanation="Ten weekdays from Monday 01/05 end on Friday 01/16. If you see 01/14 in the last cell, the series included "
                         "the weekend of 01/10–01/11. (The live formula in the key uses WORKDAY, a date function from Lesson 2.3.)"),
        Task(f"Fill the time column A{G_FIRST + 1}:A{G_LAST} with 15-minute slots after 08:00, ending at 16:45. "
             "The gray cell shows the latest time in column A.",
             answer=timedelta(minutes=slot_mins[-1]), fmt="hh:mm", title="15-minute slot times",
             solution=f"Select **A{G_FIRST}**, choose **Home → Fill → Series…**: Columns, Linear, Step **0:15**, Stop **16:50**, and click "
                      f"**OK**. Or type 8:15 in A{G_FIRST + 1}, select both cells, and drag the fill handle to **A{G_LAST}**.",
             summary=f"=IF(COUNT('Clinic Grid'!A{G_FIRST + 1}:A{G_LAST + 20})=0,\"\",MAX('Clinic Grid'!A{G_FIRST}:A{G_LAST + 20}))",
             fill={"range": f"'Clinic Grid'!A{G_FIRST + 1}:A{G_LAST}", "values": [_frac(m) for m in slot_mins[1:]]},
             live=False, hint="Same technique as the Q15 Log, with a different stop",
             explanation=f"08:00 to 16:45 every 15 minutes is {len(slot_mins)} slots per day, so the grid has "
                         f"{len(slot_mins)} × {len(clinic_days)} = {len(slot_mins) * len(clinic_days)} cells."),
        Task(f"Fill the whole grid B{G_FIRST}:{G_LASTCOL}{G_LAST} with Open in one entry. Then type Lunch over the 12:00–12:45 rows "
             "for every day, and Admin over both Wednesday afternoons (13:00–16:45). Use one Ctrl+Enter for each step. "
             "How many slots are blocked (Lunch plus Admin)? The gray cell counts blocks in the right places, minus any in the wrong place.",
             answer=n_lunch + n_admin, title="Block lunch and admin time with Ctrl+Enter",
             solution=f"1. Select **B{G_FIRST}:{G_LASTCOL}{G_LAST}**, type `Open`, and press **Ctrl + Enter** (Mac: **⌘ + Return**).\n"
                      f"2. Select **B{lunch_r1}:{G_LASTCOL}{lunch_r2}** (12:00–12:45), type `Lunch`, and press **Ctrl + Enter**.\n"
                      f"3. Select **{gcol[weds[0]]}{pm_r1}:{gcol[weds[0]]}{G_LAST}**, Ctrl+drag (Mac: ⌘+drag) "
                      f"**{gcol[weds[1]]}{pm_r1}:{gcol[weds[1]]}{G_LAST}** (the two Wednesday afternoons), type `Admin`, and press "
                      "**Ctrl + Enter**.",
             summary=(f'=IF(COUNTA({body})=0,"",2*({lunch_ok}+{admin_ok})-COUNTIF({body},"Lunch")-COUNTIF({body},"Admin"))'),
             fill={"range": body, "values": flat_final},
             live=False, hint="Holding Ctrl (Mac: ⌘) while you drag adds a second block to the selection",
             explanation=f"Lunch is 4 slots × {len(clinic_days)} days = {n_lunch}, and Admin is 16 afternoon slots × {len(weds)} "
                         f"Wednesdays = {n_admin}, so {n_lunch + n_admin} slots are blocked. Typing over a selection with Ctrl + Enter "
                         "replaces whatever was there, which is why you can paint Open everywhere first and then overwrite the exceptions."),
        Task("Select both Friday-afternoon blocks (13:00–16:45), then use Find & Replace to change Open to Telehealth inside that "
             "selection only. How many replacements does Excel report? The gray cell counts Telehealth slots in the right places.",
             answer=n_tele, title="Find & Replace inside a selection",
             solution=f"1. Select **{gcol[fris[0]]}{pm_r1}:{gcol[fris[0]]}{G_LAST}**, then Ctrl+drag (Mac: ⌘+drag) "
                      f"**{gcol[fris[1]]}{pm_r1}:{gcol[fris[1]]}{G_LAST}**.\n"
                      "2. Press **Ctrl + H** (Mac: **⌃ + H**). Find what: `Open`, Replace with: `Telehealth`. Tick **Match entire "
                      "cell contents** for safety.\n"
                      "3. Click **Replace All**. Because more than one cell is selected, Excel searches only the selection. "
                      "(You can also do one Friday at a time and add the two counts.)",
             summary=f'=IF(COUNTA({body})=0,"",2*({tele_ok})-COUNTIF({body},"Telehealth"))',
             fill={"range": body, "values": flat_final},
             live=False, hint="With several cells selected, Replace All stays inside the selection",
             explanation=f"Two Fridays × 16 afternoon slots = {n_tele}. If you had only one cell selected, Excel would have replaced "
                         "every Open on the sheet, so the whole template would have turned into Telehealth."),
        Task("After all the steps above, how many slots are still Open for in-person booking? The gray cell counts them.",
             answer=n_open, title="Open slots remaining",
             solution="No new steps: the gray cell counts the cells that still say Open.",
             summary=f'=IF(COUNTA({body})=0,"",COUNTIF({body},"Open"))',
             fill={"range": body, "values": flat_final},
             live=False, hint="Total slots minus everything you blocked or converted",
             explanation=f"{len(slot_mins) * len(clinic_days)} slots − {n_lunch} Lunch − {n_admin} Admin − {n_tele} Telehealth = "
                         f"{n_open}. If your count is higher, part of a block is missing. If it's lower, something extra was overwritten."),
    ]

    L.start_notes = [
        "Work sheets: Entries (tasks 1–2), Schedule (tasks 5–6, 9, 13), Beds (7), Q15 Log (8), Admissions (10), Supplies and "
        "Order (11), Roster (12), and Clinic Grid (bonus).",
    ]
    L.sheet_order = ["Start Here", "Practice", "Entries", "Schedule", "Beds", "Q15 Log", "Admissions", "Supplies", "Order",
                     "Roster", "Bonus", "Clinic Grid", "Answer Key", "Bonus Key"]

    # ================================================================== custom sheets
    TITLE_FONT = Font(bold=True, size=13, color=NAVY)
    NOTE_FONT = Font(italic=True, size=10, color="595959")
    HDR_FONT = Font(bold=True, color="FFFFFF")
    CENTER = Alignment(horizontal="center", vertical="center")

    def titled(ws, title, *notes):
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        for k, n in enumerate(notes, 2):
            ws.cell(row=k, column=1, value=n).font = NOTE_FONT

    def header(ws, row, labels, start_col=1, wrap=False):
        for j, h in enumerate(labels, start_col):
            c = ws.cell(row=row, column=j, value=h)
            c.font = HDR_FONT
            c.fill = HEADER_FILL
            c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)

    def yellow(ws, rng, fmt=None):
        for row in ws[rng]:
            for c in row:
                c.fill = INPUT_FILL
                c.border = INPUT_BORDER
                if fmt:
                    c.number_format = fmt

    @L.customize
    def _sheets(wb, lesson, selftest):
        # ---------------------------------------------------------- Entries
        ws = wb.create_sheet("Entries")
        ws.sheet_properties.tabColor = "2E75B6"
        titled(ws, f"Registration entries · {patient_name} · ED visit {arrival:%m/%d/%Y}",
               "A registration clerk entered these values. Column C says exactly what was typed. Compare it with what Excel stored in column B.")
        header(ws, 3, ["Field", "Entry", "What the clerk typed or did"])
        for k, (field_, val, fmt, how) in enumerate(entries):
            r = ENTRY_FIRST + k
            ws.cell(row=r, column=1, value=field_).border = BOX
            c = ws.cell(row=r, column=2, value=val)
            c.border = BOX
            if fmt:
                c.number_format = fmt
            h = ws.cell(row=r, column=3, value=how)
            h.border = BOX
            h.font = Font(italic=True, color="404040")
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 48
        ws.freeze_panes = "A4"

        # ---------------------------------------------------------- Schedule
        ws = wb.create_sheet("Schedule")
        ws.sheet_properties.tabColor = "2E75B6"
        titled(ws, "4 West (Medical-Surgical) · Rotating RN schedule · December 2025",
               "Codes: D = Day 12h (07:00–19:30) · N = Night 12h (19:00–07:30) · blank = off. Yellow cells are yours to fill.")
        header(ws, SCH_HDR, ["Date", "Day"] + nurses + ["Float RN"], wrap=True)
        ws.row_dimensions[SCH_HDR].height = 32
        ws.cell(row=SCH_FIRST, column=1, value=days[0]).number_format = "mm/dd/yyyy"
        ws.cell(row=SCH_FIRST, column=2, value=days[0].strftime("%a"))
        yellow(ws, f"A{SCH_FIRST + 1}:A{SCH_LAST}", "mm/dd/yyyy")
        yellow(ws, f"B{SCH_FIRST + 1}:B{SCH_LAST}")
        yellow(ws, f"{FLOAT_COL}{SCH_FIRST}:{FLOAT_COL}{SCH_LAST}")
        for k, row in enumerate(sched_rows):
            r = SCH_FIRST + k
            for j, v in enumerate(row):
                c = ws.cell(row=r, column=NURSE_COL0 + j, value=v or None)
                c.alignment = CENTER
                c.border = BOX
            for col in (1, 2):
                ws.cell(row=r, column=col).border = INPUT_BORDER if k else BOX
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 7
        for j in range(len(rn_ids)):
            ws.column_dimensions[get_column_letter(NURSE_COL0 + j)].width = 10
        ws.column_dimensions[FLOAT_COL].width = 10
        ws.freeze_panes = f"C{SCH_FIRST}"

        # ---------------------------------------------------------- Beds
        ws = wb.create_sheet("Beds")
        ws.sheet_properties.tabColor = "2E75B6"
        titled(ws, f"4 West · bed list ({beds_n} staffed beds, two per room)",
               "Fill the yellow Bed column from 4W-01 down to the last bed.")
        header(ws, B_FIRST - 1, ["Bed", "Room", "Position"])
        ws.cell(row=B_FIRST, column=1, value=bed_ids[0])
        yellow(ws, f"A{B_FIRST + 1}:A{B_LAST}")
        for k in range(beds_n):
            r = B_FIRST + k
            ws.cell(row=r, column=2, value=401 + k // 2).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=3, value="Window" if k % 2 == 0 else "Door")
            for col in (2, 3):
                ws.cell(row=r, column=col).border = BOX
        ws.cell(row=B_FIRST, column=1).border = BOX
        for col, w in zip("ABC", (10, 8, 10)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = f"A{B_FIRST}"

        # ---------------------------------------------------------- Q15 Log
        ws = wb.create_sheet("Q15 Log")
        ws.sheet_properties.tabColor = "2E75B6"
        titled(ws, "Behavioral Health · Q15 safety observation record · 12/16/2025 day shift",
               "Staff check on the patient every 15 minutes from 07:00 until 18:45 and record location, behavior, and initials.")
        header(ws, Q_FIRST - 1, ["Time", "Location", "Behavior", "Staff initials"])
        c = ws.cell(row=Q_FIRST, column=1, value=time(7, 0))
        c.number_format = "hh:mm"
        c.border = BOX
        yellow(ws, f"A{Q_FIRST + 1}:A{Q_LAST}", "hh:mm")
        for r in range(Q_FIRST, Q_LAST + 1):
            for col in (2, 3, 4):
                ws.cell(row=r, column=col).border = BOX
        for col, w in zip("ABCD", (9, 18, 28, 14)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = f"A{Q_FIRST}"

        # ---------------------------------------------------------- Order
        ws = wb.create_sheet("Order")
        ws.sheet_properties.tabColor = "2E75B6"
        titled(ws, "4 West · supply requisition · 12/31/2025",
               "Order Qty must hold plain numbers (values), not formulas.")
        header(ws, O_FIRST - 1, ["SKU", "Item", "Unit", "Order Qty"])
        for k, r_ in enumerate(supplies):
            r = O_FIRST + k
            for col, v in ((1, r_["SKU"]), (2, r_["ItemDescription"]), (3, r_["UnitOfMeasure"])):
                ws.cell(row=r, column=col, value=v).border = BOX
        yellow(ws, f"D{O_FIRST}:D{O_LAST}", "#,##0")
        for col, w in zip("ABCD", (12, 40, 8, 11)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = f"A{O_FIRST}"

        # ---------------------------------------------------------- Roster (transpose target)
        ws = wb.create_sheet("Roster")
        ws.sheet_properties.tabColor = "2E75B6"
        titled(ws, "4 West · December 2025 roster (staff down the side, dates across the top)",
               f"Paste the transposed Schedule here, with its top-left corner in the yellow cell A{ROSTER_ANCHOR}.")
        yellow(ws, f"A{ROSTER_ANCHOR}:A{ROSTER_ANCHOR}")
        ws.column_dimensions["A"].width = 18

        # ---------------------------------------------------------- Clinic Grid (bonus)
        ws = wb.create_sheet("Clinic Grid")
        ws.sheet_properties.tabColor = "BF9000"
        titled(ws, f"Primary Care Clinic · {doc} · appointment template, 01/05/2026 – 01/16/2026",
               "Build the grid with AutoFill, Fill Series, Ctrl+Enter, and Find & Replace. No typing cell by cell.")
        header(ws, G_HDR, ["Time"])
        c = ws.cell(row=G_HDR, column=G_COL0, value=clinic_days[0])
        c.number_format = "ddd mm/dd"
        c.font = Font(bold=True)
        c.alignment = CENTER
        c.border = BOX
        yellow(ws, f"C{G_HDR}:{G_LASTCOL}{G_HDR}", "ddd mm/dd")
        for row in ws[f"C{G_HDR}:{G_LASTCOL}{G_HDR}"]:
            for cc in row:
                cc.font = Font(bold=True)
                cc.alignment = CENTER
        c = ws.cell(row=G_FIRST, column=1, value=time(8, 0))
        c.number_format = "hh:mm"
        c.border = BOX
        yellow(ws, f"A{G_FIRST + 1}:A{G_LAST}", "hh:mm")
        yellow(ws, f"B{G_FIRST}:{G_LASTCOL}{G_LAST}")
        for row in ws[f"B{G_FIRST}:{G_LASTCOL}{G_LAST}"]:
            for cc in row:
                cc.alignment = Alignment(horizontal="center")
        ws.column_dimensions["A"].width = 8
        for dd in clinic_days:
            ws.column_dimensions[gcol[dd]].width = 11
        ws.freeze_panes = f"B{G_FIRST}"

    return L
