"""Lesson 3.5 · Charts & Data Visualization.

Learners build charts from small, pre-summarized Bluestone tables (one table per sheet, so every chart sits next to
its data) and answer a question that the finished chart shows. A chart can't be checked by a formula, so:

* Each chart task's ``solution`` is Markdown build steps; the answer is computed in Python from the same rows that
  are written to the workbook.
* Each task's ``live`` formula is an independent cross-check on the Table (XLOOKUP, SLOPE, RSQ, COUNTIF…), so the
  verifier proves every Python answer in LibreOffice and learners see the "formula twin" of every chart read.
* The self-test types the expected value into each answer cell (the default for non-formula solutions). Task 13
  (dynamic title) is a real formula, so the self-test types the formula itself.
* A customize hook adds a visible "Makeover" sheet with a deliberately misleading chart (Task 4), yellow sparkline
  cells under the ED_Monthly table (Task 12), and a hidden "Chart Key" sheet with reference charts drawn with
  openpyxl. openpyxl can't draw Excel 2016+ chart types (histogram, Pareto, waterfall) or sparklines, so the key
  shows the classic equivalents: a pre-binned column chart, a column + cumulative-line combo, and a stacked-column
  waterfall with an invisible base series.
"""
from __future__ import annotations

import math
import re
import statistics as st
from collections import Counter, defaultdict
from datetime import date

from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.trendline import Trendline, TrendlineLabel
from openpyxl.styles import Alignment, Font

from xlcourse import Lesson, Task, data
from xlcourse.data import excel_serial
from xlcourse.lesson import BOX, HEADER_FILL, INPUT_BORDER, INPUT_FILL, NAVY, _plain

CODE = "3.5"

FAC_COLS = {"F01": "Bluestone Memorial", "F02": "Ashby Falls", "F03": "Cedar Ridge"}
T2_FACILITY = "Cedar Ridge"
AXIS_MIN = 0.05            # the misleading axis minimum on the Makeover chart
LWBS_TARGET = 0.02         # bonus: LWBS rate target (2% or less)
WATERFALL_DEPT = "D110"    # Medical-Surgical 4 West, Bluestone Memorial
WATERFALL_YEAR = 2025
SAMPLE_N = 300
HIST_UNDER, HIST_OVER = 1, 10
PARETO_CUT = 0.80

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTHS_LONG = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
               "November", "December"]

# Budget categories in the order the waterfall shows them, with the label used on the chart.
BUDGET_STEPS = [
    ("Revenue", "Net Patient Service Revenue", "Net patient revenue"),
    ("Expense", "Salaries & Wages", "Salaries & wages"),
    ("Expense", "Employee Benefits", "Employee benefits"),
    ("Expense", "Medical Supplies", "Medical supplies"),
    ("Expense", "Pharmaceuticals", "Pharmaceuticals"),
    ("Expense", "Purchased Services", "Purchased services"),
    ("Expense", "Equipment & Maintenance", "Equipment & maintenance"),
    ("Expense", "Other Operating", "Other operating"),
]

# openpyxl writes the Office 2007 theme, whose accents run blue, red, green, purple. Every chart a learner builds in
# this workbook takes its default colors from the theme, so swap in the Office 2013-2022 colors (blue, orange, gray,
# gold, ...) that match current Excel and the lesson's "avoid red with green" advice. Only the color scheme changes.
OFFICE_2013_COLORS = {"1F497D": "44546A", "EEECE1": "E7E6E6", "4F81BD": "4472C4", "C0504D": "ED7D31",
                      "9BBB59": "A5A5A5", "8064A2": "FFC000", "4BACC6": "5B9BD5", "F79646": "70AD47",
                      "0000FF": "0563C1", "800080": "954F72"}


def _office_2013_theme() -> bytes:
    from openpyxl.writer.theme import theme_xml
    start, end = theme_xml.index("<a:clrScheme"), theme_xml.index("</a:clrScheme>")
    scheme = theme_xml[start:end]
    for old, new in OFFICE_2013_COLORS.items():
        assert scheme.count(f'val="{old}"') == 1, old
        scheme = scheme.replace(f'val="{old}"', f'val="{new}"')
    return (theme_xml[:start] + scheme + theme_xml[end:]).encode("utf-8")


# Okabe-Ito colour-blind-safe colours used in the reference charts.
BLUE, ORANGE, SKY, GREY, VERMILLION = "0072B2", "E69F00", "56B4E9", "A6A6A6", "D55E00"


class _Labels(DataLabelList):
    """Data labels that show only what's asked for, with a number format Excel won't link back to the source.

    openpyxl writes <c:numFmt formatCode=".."/> without sourceLinked="0", and leaves the other show* flags unset,
    which some renderers read as "show everything" (category name, series name, legend key)."""

    tagname = "dLbls"
    # openpyxl's metaclass only collects descriptors defined on the class itself, so reuse the parent's lists.
    __attrs__ = DataLabelList.__attrs__
    __nested__ = DataLabelList.__nested__
    __elements__ = DataLabelList.__elements__

    def __init__(self, fmt=None, val=False, pct=False, cat=False, **kw):
        super().__init__(numFmt=fmt, showVal=val, showPercent=pct, showCatName=cat, showSerName=False,
                         showLegendKey=False, **kw)

    def to_tree(self, tagname=None, idx=None, namespace=None):
        el = super().to_tree(tagname, idx, namespace)
        for child in el:
            if child.tag.endswith("numFmt"):
                child.set("sourceLinked", "0")
        return el


def _age(dob: date, on: date) -> int:
    return on.year - dob.year - ((on.month, on.day) < (dob.month, dob.day))


def _month_label(d: date) -> str:
    return f"{MONTHS[d.month - 1]} {d.year}"


# ---------------------------------------------------------------------------------------------- data
def ed_monthly_rows() -> list[dict]:
    counts = defaultdict(Counter)
    lwbs = Counter()
    for r in data.load("ed_visits"):
        m = date(r["ArrivalDateTime"].year, r["ArrivalDateTime"].month, 1)
        counts[m][r["FacilityID"]] += 1
        if r["EDDisposition"] == "LWBS":
            lwbs[m] += 1
    rows = []
    for m in sorted(counts):
        row = {"Month": m}
        for fid, name in FAC_COLS.items():
            row[name] = counts[m][fid]
        row["Total"] = sum(counts[m].values())
        row["LWBS"] = lwbs[m]
        rows.append(row)
    assert len(rows) == 24 and rows[0]["Month"] == date(2024, 1, 1) and rows[-1]["Month"] == date(2025, 12, 1)
    assert all(set(counts[m]) <= set(FAC_COLS) for m in counts), "only the three hospitals have EDs"
    return rows


def ed_hourly_rows() -> list[dict]:
    n = Counter()
    waits = defaultdict(list)
    for r in data.load("ed_visits"):
        if r["ArrivalDateTime"].year != 2025:
            continue
        h = r["ArrivalDateTime"].hour
        n[h] += 1
        if r["ProviderSeenDateTime"] is not None:
            waits[h].append((r["ProviderSeenDateTime"] - r["ArrivalDateTime"]).total_seconds() / 60)
    return [{"ArrivalHour": f"{h:02d}:00", "Hour": h, "Arrivals": n[h],
             "AvgDoorToProviderMin": round(sum(waits[h]) / len(waits[h]), 1)} for h in range(24)]


def readmit_rows() -> list[dict]:
    dep = data.index(data.load("departments"), "DeptID")
    by_sl = defaultdict(lambda: [0, 0])
    for e in data.load("encounters"):
        if e["EncounterType"] == "Inpatient" and e["AdmitDateTime"].year == 2025:
            s = by_sl[dep[e["DeptID"]]["ServiceLine"]]
            s[0] += 1
            s[1] += e["Readmit30"] == "Y"
    return [{"ServiceLine": k, "IndexStays": v[0], "Readmissions": v[1], "ReadmitRate": v[1] / v[0]}
            for k, v in sorted(by_sl.items())]


def payer_mix_rows() -> list[dict]:
    pay = data.index(data.load("payers"), "PayerID")
    c = Counter(pay[e["PayerID"]]["PayerType"] for e in data.load("encounters") if e["AdmitDateTime"].year == 2025)
    return [{"PayerType": k, "Encounters": v} for k, v in sorted(c.items(), key=lambda kv: (-kv[1], kv[0]))]


def stay_rows() -> list[dict]:
    """A systematic sample of 2025 inpatient stays: every (n/300)-th stay in admission order."""
    pat = data.index(data.load("patients"), "PatientID")
    dep = data.index(data.load("departments"), "DeptID")
    ip = [e for e in data.load("encounters") if e["EncounterType"] == "Inpatient" and e["AdmitDateTime"].year == 2025]
    ip.sort(key=lambda e: (e["AdmitDateTime"], e["EncounterID"]))
    n = len(ip)
    rows = []
    for i in range(SAMPLE_N):
        e = ip[i * n // SAMPLE_N]
        rows.append({
            "EncounterID": e["EncounterID"], "ServiceLine": dep[e["DeptID"]]["ServiceLine"],
            "AgeAtAdmit": _age(pat[e["PatientID"]]["DOB"], e["AdmitDateTime"].date()),
            "LOSDays": round((e["DischargeDateTime"] - e["AdmitDateTime"]).total_seconds() / 86400, 1),
            "TotalCharges": e["TotalCharges"],
        })
    return rows, n


def denial_rows() -> list[dict]:
    """2025 initial denials: claims with a 2025 service date that were denied (ClaimStatus Denied or Appealed)."""
    cnt, dollars = Counter(), defaultdict(float)
    for c in data.load("claims"):
        if c["ClaimStatus"] in ("Denied", "Appealed") and c["ServiceDate"].year == 2025:
            assert c["DenialReason"], c["ClaimID"]
            cnt[c["DenialReason"]] += 1
            dollars[c["DenialReason"]] += c["BilledAmount"]
    return [{"DenialReason": k, "Claims": cnt[k], "DeniedCharges": round(dollars[k], 2)} for k in sorted(cnt)]


def budget_rows() -> tuple[list[dict], dict]:
    agg = defaultdict(lambda: [0, 0])
    for r in data.load("budget"):
        if r["DeptID"] == WATERFALL_DEPT and r["FiscalYear"] == WATERFALL_YEAR:
            a = agg[(r["LineType"], r["Category"])]
            a[0] += r["BudgetAmount"]
            a[1] += r["ActualAmount"]
    assert set(agg) == {(lt, cat) for lt, cat, _ in BUDGET_STEPS}, sorted(agg)

    def margin(i):
        return sum(v[i] if k[0] == "Revenue" else -v[i] for k, v in agg.items())
    b_margin, a_margin = margin(0), margin(1)
    rows = [{"Step": "Budgeted operating margin", "Amount": b_margin}]
    for lt, cat, label in BUDGET_STEPS:
        bud, act = agg[(lt, cat)]
        rows.append({"Step": label, "Amount": (act - bud) if lt == "Revenue" else (bud - act)})
    rows.append({"Step": "Actual operating margin", "Amount": a_margin})
    assert b_margin + sum(r["Amount"] for r in rows[1:-1]) == a_margin
    return rows, agg


# ---------------------------------------------------------------------------------------------- build
def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="03-data-analysis", slug="05-charts-visualization",
        title="Charts & Data Visualization", level="Intermediate", minutes=55,
        objectives=[
            "Choose the right chart for comparisons, trends, parts of a whole, distributions, and relationships",
            "Build and format column, line, bar, scatter, histogram, combo, and Pareto charts",
            "Add trendlines, dynamic titles, and sparklines",
            "Design clear, accessible charts that tell one story",
        ],
    )

    monthly = ed_monthly_rows()
    hourly = ed_hourly_rows()
    readmits = readmit_rows()
    payers = payer_mix_rows()
    stays, n_ip = stay_rows()
    denials = denial_rows()
    budget, _ = budget_rows()
    total_ed = sum(r["Total"] for r in monthly)

    L.data_note = (
        "Bluestone Health System, pre-summarized for charting: monthly ED visits by hospital with counts of patients who "
        "left without being seen (LWBS), Jan 2024 – Dec 2025, 2025 ED arrivals and average door-to-provider minutes by hour of day, 2025 readmission rates by "
        f"service line, 2025 payer mix, a sample of {SAMPLE_N} of the {n_ip:,} inpatient stays that began in 2025, 2025 "
        "claim denials by reason, "
        "and 4 West's FY2025 budget-to-actual margin bridge.")

    # ------------------------------------------------------------------ data sheets
    hr = L.add_table_sheet(
        "ED_Hourly", hourly, table="tblEDHourly", columns=["ArrivalHour", "Arrivals", "AvgDoorToProviderMin"],
        formats={"Arrivals": "#,##0", "AvgDoorToProviderMin": "0.0"},
        widths={"ArrivalHour": 13, "Arrivals": 11, "AvgDoorToProviderMin": 23})
    mo = L.add_table_sheet(
        "ED_Monthly", monthly, table="tblEDMonthly",
        columns=["Month"] + list(FAC_COLS.values()) + ["Total", "LWBS"], extra_cols=["LWBSRate"],
        formats={"Month": "mmm yyyy", "Total": "#,##0", "LWBS": "0", "LWBSRate": "0.00%",
                 **{c: "#,##0" for c in FAC_COLS.values()}},
        widths={"Month": 11, "Bluestone Memorial": 19, "Ashby Falls": 13, "Cedar Ridge": 13, "Total": 9, "LWBS": 8,
                "LWBSRate": 11})
    ra = L.add_table_sheet(
        "Readmits", readmits, table="tblReadmits",
        columns=["ServiceLine", "IndexStays", "Readmissions", "ReadmitRate"],
        formats={"IndexStays": "#,##0", "Readmissions": "#,##0", "ReadmitRate": "0.0%"},
        widths={"ServiceLine": 20, "IndexStays": 12, "Readmissions": 14, "ReadmitRate": 13})
    pm = L.add_table_sheet(
        "PayerMix", payers, table="tblPayerMix", columns=["PayerType", "Encounters"],
        formats={"Encounters": "#,##0"}, widths={"PayerType": 21, "Encounters": 12})
    sy = L.add_table_sheet(
        "Stays", stays, table="tblStays",
        columns=["EncounterID", "ServiceLine", "AgeAtAdmit", "LOSDays", "TotalCharges"],
        formats={"AgeAtAdmit": "0", "LOSDays": "0.0", "TotalCharges": "#,##0.00"},
        widths={"EncounterID": 13, "ServiceLine": 19, "AgeAtAdmit": 12, "LOSDays": 10, "TotalCharges": 14})
    dn = L.add_table_sheet(
        "Denials", denials, table="tblDenials", columns=["DenialReason", "Claims", "DeniedCharges"],
        extra_cols=["CumulativePct"],
        formats={"Claims": "#,##0", "DeniedCharges": "#,##0.00", "CumulativePct": "0.0%"},
        widths={"DenialReason": 25, "Claims": 9, "DeniedCharges": 15, "CumulativePct": 15})
    bg = L.add_table_sheet(
        "Budget", budget, table="tblBudget", columns=["Step", "Amount"], start_row=4,
        notes=["Medical-Surgical 4 West (Bluestone Memorial) · FY2025 operating margin, budget to actual",
               "Positive amounts are favorable: revenue above budget, or an expense category that came in under budget. "
               "Negative amounts are unfavorable."],
        formats={"Amount": "#,##0;-#,##0"}, widths={"Step": 30, "Amount": 14})

    def col(sd, name: str) -> str:
        """Structured reference to a Table column."""
        assert name in sd.headers, name
        return f"{sd.table}[{name}]"

    def cross(formula: str) -> str:
        return f"Cross-check with a formula: `{formula}`"

    mfirst, mlast = mo.first_row, mo.last_row

    # ------------------------------------------------------------------ answers (computed in Python)
    # T1 busiest arrival hour (must be unique)
    by_arr = sorted(hourly, key=lambda r: -r["Arrivals"])
    assert by_arr[0]["Arrivals"] > by_arr[1]["Arrivals"]
    t1 = by_arr[0]["Hour"]
    quiet = min(hourly, key=lambda r: r["Arrivals"])

    # T2 Cedar Ridge's busiest month (unique)
    cr = sorted(monthly, key=lambda r: -r[T2_FACILITY])
    assert cr[0][T2_FACILITY] > cr[1][T2_FACILITY]
    t2 = cr[0]["Month"]
    mem_peak = max(monthly, key=lambda r: r["Bluestone Memorial"])

    # T3 sorted bar chart: the LAST row (lowest rate) is drawn at the top before the axis is reversed
    ra_sorted = sorted(readmits, key=lambda r: -r["ReadmitRate"])
    assert ra_sorted[0]["ReadmitRate"] > ra_sorted[1]["ReadmitRate"]
    assert ra_sorted[-1]["ReadmitRate"] < ra_sorted[-2]["ReadmitRate"]
    t3 = ra_sorted[-1]["ServiceLine"]
    hi_sl, lo_sl = ra_sorted[0], ra_sorted[-1]

    # T4 truncated axis: apparent ratio of bar heights measured from the 5% axis minimum
    t4_true = hi_sl["ReadmitRate"] / lo_sl["ReadmitRate"]
    t4_raw = (hi_sl["ReadmitRate"] - AXIS_MIN) / (lo_sl["ReadmitRate"] - AXIS_MIN)
    t4 = round(t4_raw, 1)
    assert lo_sl["ReadmitRate"] > AXIS_MIN

    # T5 pie: Commercial share of 2025 encounters
    pm_total = sum(r["Encounters"] for r in payers)
    commercial = next(r["Encounters"] for r in payers if r["PayerType"] == "Commercial")
    t5 = commercial / pm_total
    smallest_payer = payers[-1]

    # T6 histogram: overflow bin (> 10 days)
    los = [r["LOSDays"] for r in stays]
    t6 = sum(1 for x in los if x > HIST_OVER)
    hist_bins = [("≤" + str(HIST_UNDER), sum(1 for x in los if x <= HIST_UNDER))]
    for k in range(HIST_UNDER, HIST_OVER):
        hist_bins.append((f"({k}, {k + 1}]", sum(1 for x in los if k < x <= k + 1)))
    hist_bins.append((">" + str(HIST_OVER), t6))
    assert sum(c for _, c in hist_bins) == SAMPLE_N
    tallest_bin = max(hist_bins, key=lambda b: b[1])
    los_median, los_mean = st.median(los), st.mean(los)

    # T7 slope of charges on LOS ($ per extra day); T8 R² of LOS on age
    charges = [r["TotalCharges"] for r in stays]
    ages = [r["AgeAtAdmit"] for r in stays]
    reg_ch = st.linear_regression(los, charges)
    t7 = reg_ch.slope
    r2_ch = st.correlation(los, charges) ** 2
    reg_age = st.linear_regression(ages, los)
    t8 = st.correlation(ages, los) ** 2

    # T9 combo: arrival hour with the longest average door-to-provider time (unique)
    by_wait = sorted(hourly, key=lambda r: -r["AvgDoorToProviderMin"])
    assert by_wait[0]["AvgDoorToProviderMin"] > by_wait[1]["AvgDoorToProviderMin"]
    t9 = by_wait[0]["Hour"]
    short_wait = min(hourly, key=lambda r: r["AvgDoorToProviderMin"])
    corr_hour = st.correlation([r["Arrivals"] for r in hourly], [r["AvgDoorToProviderMin"] for r in hourly])

    # T10 Pareto by claim count: cumulative % of the top two reasons
    dn_by_claims = sorted(denials, key=lambda r: -r["Claims"])
    assert len({r["Claims"] for r in denials}) == len(denials), "Pareto order by count must have no ties"
    dn_total = sum(r["Claims"] for r in denials)
    cum_claims, run = [], 0
    for r in dn_by_claims:
        run += r["Claims"]
        cum_claims.append(run / dn_total)
    t10 = cum_claims[1]

    # T11 waterfall: biggest drop
    steps = budget[1:-1]
    t11_row = min(steps, key=lambda r: r["Amount"])
    assert sorted(r["Amount"] for r in steps)[1] > t11_row["Amount"]
    t11 = t11_row["Step"]
    biggest_up = max(steps, key=lambda r: r["Amount"])
    b_margin, a_margin = budget[0]["Amount"], budget[-1]["Amount"]

    # T12 sparklines: facilities whose Dec 2025 value is above their Jan 2024 value
    first_m, last_m = monthly[0], monthly[-1]
    up = [f for f in FAC_COLS.values() if last_m[f] > first_m[f]]
    assert all(last_m[f] != first_m[f] for f in FAC_COLS.values())
    t12 = len(up)
    down = [f for f in FAC_COLS.values() if f not in up]

    # T13 dynamic title
    t13 = (f"ED visits by facility, {_month_label(monthly[0]['Month'])} to {_month_label(monthly[-1]['Month'])} "
           f"({total_ed:,} visits)")
    t13_accept = [
        f"ED visits by facility, {MONTHS_LONG[0]} 2024 to {MONTHS_LONG[11]} 2025 ({total_ed:,} visits)",
        f"ED visits by facility, Jan 2024 to Dec 2025 ({total_ed} visits)",
    ]
    assert monthly[0]["Month"] == date(2024, 1, 1) and monthly[-1]["Month"] == date(2025, 12, 1)

    # ------------------------------------------------------------------ bonus answers
    rates = [(r["Month"], r["LWBS"] / r["Total"]) for r in monthly]
    b1 = sum(1 for _, x in rates if x > LWBS_TARGET)
    assert all(abs(x - LWBS_TARGET) > 1e-4 for _, x in rates), "no month should sit on the target"
    rates_sorted = sorted(rates, key=lambda mx: -mx[1])
    assert rates_sorted[0][1] > rates_sorted[1][1]
    b2, b2_rate = rates_sorted[0]
    b2_row = next(r for r in monthly if r["Month"] == b2)
    busiest_month = max(monthly, key=lambda r: r["Total"])
    quietest_month = min(monthly, key=lambda r: r["Total"])
    b2_note = "the quietest month of the two years" if b2 == quietest_month["Month"] else "a quieter-than-average month"
    assert b2_row["Total"] < st.mean(r["Total"] for r in monthly)
    busiest_rate = busiest_month["LWBS"] / busiest_month["Total"]
    b3 = st.correlation([r["Total"] for r in monthly], [x for _, x in rates]) ** 2
    reg_lwbs = st.linear_regression([r["Total"] for r in monthly], [x for _, x in rates])

    dn_by_dollars = sorted(denials, key=lambda r: -r["DeniedCharges"])
    dollars_total = sum(r["DeniedCharges"] for r in denials)
    cum_dollars, run = [], 0.0
    for r in dn_by_dollars:
        run += r["DeniedCharges"]
        cum_dollars.append(run / dollars_total)
    b4 = next(i for i, c in enumerate(cum_dollars, 1) if c >= PARETO_CUT)
    assert all(abs(c - PARETO_CUT) > 0.005 for c in cum_dollars), "no cumulative share should sit on the 80% line"
    rank_claims = {r["DenialReason"]: i for i, r in enumerate(dn_by_claims, 1)}
    rank_dollars = {r["DenialReason"]: i for i, r in enumerate(dn_by_dollars, 1)}
    movers = [k for k in rank_dollars if rank_dollars[k] < rank_claims[k]]
    assert len(movers) == 1, movers
    b5 = movers[0]
    b5_row = next(r for r in denials if r["DenialReason"] == b5)
    swapped = next(k for k in rank_dollars if rank_dollars[k] > rank_claims[k])
    swapped_row = next(r for r in denials if r["DenialReason"] == swapped)
    avg_claim = {r["DenialReason"]: r["DeniedCharges"] / r["Claims"] for r in denials}

    # Qualitative claims made in explanations and Chart Key titles must hold for the data.
    assert 12 <= t1 < 18, "Chart Key title says arrivals peak in the afternoon"
    assert len(payers) == 5 and {r["PayerType"] for r in payers[:2]} == {"Government", "Commercial"}
    assert (payers[0]["Encounters"] + payers[1]["Encounters"]) / pm_total > 0.6, "two payer types should dominate"
    assert t8 < 0.05 and reg_age.slope > 0, "Task 8 text: weak but upward trend"
    assert corr_hour > 0.8, "Task 9 text: waits rise and fall with arrivals"
    assert reg_lwbs.slope > 0 and b3 < 0.25, "B3 text: slight upward trend, weak fit"
    assert t11_row["Step"] == "Net patient revenue" and biggest_up["Step"] == "Salaries & wages"
    assert (a_margin - b_margin - t11_row["Amount"]) > 0.5 * -t11_row["Amount"], "expense savings offset most of it"

    # ------------------------------------------------------------------ cross-check formulas
    H_HOUR, H_ARR, H_WAIT = col(hr, "ArrivalHour"), col(hr, "Arrivals"), col(hr, "AvgDoorToProviderMin")
    M_MONTH, M_TOTAL, M_LWBS = col(mo, "Month"), col(mo, "Total"), col(mo, "LWBS")
    R_SL, R_RATE = col(ra, "ServiceLine"), col(ra, "ReadmitRate")
    P_TYPE, P_ENC = col(pm, "PayerType"), col(pm, "Encounters")
    S_AGE, S_LOS, S_CH = col(sy, "AgeAtAdmit"), col(sy, "LOSDays"), col(sy, "TotalCharges")
    D_REASON, D_CLAIMS, D_DOLLARS = col(dn, "DenialReason"), col(dn, "Claims"), col(dn, "DeniedCharges")
    B_STEP, B_AMT = col(bg, "Step"), col(bg, "Amount")
    CR = col(mo, T2_FACILITY)

    f1 = f"=VALUE(LEFT(XLOOKUP(MAX({H_ARR}),{H_ARR},{H_HOUR}),2))"
    f2 = f"=XLOOKUP(MAX({CR}),{CR},{M_MONTH})"
    f3 = f"=XLOOKUP(MIN({R_RATE}),{R_RATE},{R_SL})"
    f4 = f"=(MAX({R_RATE})-{AXIS_MIN})/(MIN({R_RATE})-{AXIS_MIN})"
    f5 = f'=XLOOKUP("Commercial",{P_TYPE},{P_ENC})/SUM({P_ENC})'
    f6 = f'=COUNTIF({S_LOS},">{HIST_OVER}")'
    f7 = f"=SLOPE({S_CH},{S_LOS})"
    f8 = f"=RSQ({S_LOS},{S_AGE})"
    f9 = f"=VALUE(LEFT(XLOOKUP(MAX({H_WAIT}),{H_WAIT},{H_HOUR}),2))"
    f10 = f"=SUM(LARGE({D_CLAIMS},{{1,2}}))/SUM({D_CLAIMS})"
    f11 = f"=XLOOKUP(MIN({B_AMT}),{B_AMT},{B_STEP})"
    fac_first = f"ED_Monthly!$B${mfirst}:$D${mfirst}"
    fac_last = f"ED_Monthly!$B${mlast}:$D${mlast}"
    f12 = f"=SUMPRODUCT(--({fac_last}>{fac_first}))"
    f13 = (f'="ED visits by facility, "&TEXT(MIN({M_MONTH}),"mmm yyyy")&" to "&TEXT(MAX({M_MONTH}),"mmm yyyy")'
           f'&" ("&TEXT(SUM({M_TOTAL}),"#,##0")&" visits)"')
    rate_arr = f"{M_LWBS}/{M_TOTAL}"
    lwbs_rng = f"ED_Monthly!${mo.col('LWBSRate')}${mfirst}:${mo.col('LWBSRate')}${mlast}"
    fb1_live = f"=SUMPRODUCT(--({rate_arr}>{LWBS_TARGET}))"
    fb2 = f"=XLOOKUP(MAX({rate_arr}),{rate_arr},{M_MONTH})"
    fb3 = f"=RSQ({rate_arr},{M_TOTAL})"
    fb4 = (f'=SUMPRODUCT(--(SUMIF({D_DOLLARS},">="&{D_DOLLARS})<{PARETO_CUT}*SUM({D_DOLLARS})))+1')
    fb5 = (f'=INDEX({D_REASON},MATCH(1,--(COUNTIF({D_DOLLARS},">"&{D_DOLLARS})'
           f'<COUNTIF({D_CLAIMS},">"&{D_CLAIMS})),0))')

    # ------------------------------------------------------------------ practice
    L.practice_intro = (
        "Each task names the sheet to work on. Build each chart on the same sheet as its data, to the right of the table, "
        "then type the number or name your chart shows into the yellow cell. The Check column can't see your chart, so "
        "every task asks a question the finished chart answers. The hidden Chart Key sheet shows a reference version of "
        "each chart. The yellow LWBSRate and CumulativePct columns are for the Bonus.")
    L.start_notes = [
        "Build every chart next to its table. Each data sheet has empty columns on the right for that.",
        "The Makeover sheet holds a deliberately misleading chart for Task 4. Fix it in place.",
        "The hidden 'Chart Key' sheet shows a reference version of every chart. Unhide it the same way as the Answer Key. "
        "It draws the histogram, Pareto, and waterfall the classic way (column charts with helper data), so they look "
        "slightly different from Excel 2016's built-in versions.",
        "Histogram, Pareto, and waterfall charts need Excel 2016 or later (Windows or Mac) or Microsoft 365.",
        "On the Stays sheet, LOSDays is the exact length of stay (discharge time minus admit time, in days, rounded to 1 decimal "
        "place), so it can differ slightly from the count of midnights used in Lesson 3.4.",
    ]

    hour_txt = lambda h: f"{h:02d}:00"  # noqa: E731
    L.tasks = [
        Task("On ED_Hourly, select A1:B25 (ArrivalHour and Arrivals) and insert a clustered column chart. Which hour of the "
             "day had the most ED arrivals in 2025? Type the hour as a number from 0 to 23 (for example, 9 for the "
             "09:00 hour).",
             answer=t1, accept=[t1 / 24], title="Column chart: busiest arrival hour",
             hint="Insert → Insert Column or Bar Chart → Clustered Column. Hover over the tallest column",
             solution="1. On **ED_Hourly**, select **A1:B25**: the ArrivalHour labels and the Arrivals numbers, headers "
                      "included.\n2. Choose **Insert → Insert Column or Bar Chart → 2-D Column → Clustered Column**. "
                      "On Windows you can also press **Alt + F1**, which inserts the default chart type (clustered "
                      "column unless someone changed it).\n3. Drag the chart to the right of the table. Hover over the "
                      f"tallest column: the tooltip reads *Series \"Arrivals\" Point \"{hour_txt(t1)}\" Value: "
                      f"{by_arr[0]['Arrivals']}*.",
             live=f1,
             explanation=f"The {hour_txt(t1)} hour had {by_arr[0]['Arrivals']} arrivals, just ahead of "
                         f"{hour_txt(by_arr[1]['Hour'])} with {by_arr[1]['Arrivals']}. The quietest hour was "
                         f"{hour_txt(quiet['Hour'])} with {quiet['Arrivals']}. A column chart suits this question because "
                         "it compares one number across categories, and the eye compares column heights quickly. "
                         "ArrivalHour holds text labels such as 16:00, so Excel uses it as the category axis "
                         "automatically. If the hours were numbers, Excel would plot them as a second series of short "
                         "columns. " + cross(f1)),
        Task(f"On ED_Monthly, select A1:D25 (Month and the three hospitals) and insert a line chart. In which month did "
             f"{T2_FACILITY} have its most ED visits? Type the month and year (for example, Jun 2025).",
             answer=t2, fmt="mmm yyyy", answer_display=_month_label(t2),
             title=f"Line chart: {T2_FACILITY}'s busiest month",
             hint="Insert → Insert Line or Area Chart → Line. Hover over the highest point of the Cedar Ridge line",
             solution="1. On **ED_Monthly**, select **A1:D25**.\n2. Choose **Insert → Insert Line or Area Chart → 2-D "
                      "Line → Line** (or **Line with Markers**). Excel recognizes the Month column as dates and builds a "
                      "date axis, one point per month.\n3. Hover over the highest point of the "
                      f"{T2_FACILITY} line, or click the line once and read the values in the tooltips.",
             live=f2,
             explanation=f"{T2_FACILITY} peaked at {cr[0][T2_FACILITY]} visits in {_month_label(t2)}. A line chart is "
                         "the standard choice for a trend over time because the slope between points shows the change. "
                         f"Notice how flat the {T2_FACILITY} and Ashby Falls lines look: Bluestone Memorial (peak "
                         f"{mem_peak['Bluestone Memorial']} in {_month_label(mem_peak['Month'])}) sets the scale, so the "
                         "smaller hospitals' swings shrink. When a small series matters, give it its own chart. "
                         + cross(f2)),
        Task("On Readmits, sort the table by ReadmitRate from largest to smallest. Then select ServiceLine and ReadmitRate "
             "(A1:A9, then Ctrl+click D1:D9, or ⌘+click on a Mac) and insert a clustered bar chart. Before you change anything else, which "
             "service line's bar is at the TOP of the chart? (Afterwards, fix the order so the highest rate is on top.)",
             answer=t3, accept=[t3.replace("&", "and")], title="Bar chart: which bar Excel puts on top",
             hint="Read the chart, not the table. Guide section 6 explains the bar order and how to fix it",
             solution="1. Click any ReadmitRate cell and choose **Data → Sort Largest to Smallest** (Mac: **Data → "
                      "Sort**, or the column's filter button).\n2. Select **A1:A9**, hold **Ctrl** (Mac: **⌘**), and "
                      "select **D1:D9**.\n3. Choose **Insert → Insert Column or Bar Chart → 2-D Bar → Clustered Bar**.\n"
                      f"4. Read the top bar: it's {t3}, the lowest rate.\n5. To fix the order, double-click the vertical "
                      "(category) axis to open **Format Axis**, then under **Axis Options** tick **Categories in reverse "
                      "order**. Set **Horizontal axis crosses** to **At maximum category** so the value axis stays at "
                      "the bottom.",
             live=f3,
             explanation="A bar chart plots the first category next to the origin, which is the bottom of a bar chart. "
                         "So a table sorted from highest to lowest produces a chart with the highest bar at the bottom and "
                         f"the lowest ({t3}, {lo_sl['ReadmitRate']:.1%}) at the top, which is upside down for a ranking. "
                         "**Categories in reverse order** flips it without re-sorting the data. Bars suit this data "
                         "better than columns because the service-line names are long and read easily on the left. "
                         + cross(f3)),
        Task(f"The Makeover sheet has a colleague's column chart of the same readmission rates. Its vertical axis starts at "
             f"{AXIS_MIN:.0%}, not 0%. Measured from that axis, how many times taller is the {hi_sl['ServiceLine']} bar "
             f"than the {lo_sl['ServiceLine']} bar? Round to 1 decimal place. Then fix the chart in place, using the "
             "checklist in Guide section 16, and add alt text.",
             answer=t4, fmt="0.0", tol=0.051, title="Makeover: how much a truncated axis exaggerates",
             hint=f"A bar's drawn height is its value minus the axis minimum ({AXIS_MIN:.0%})",
             solution=f"1. Each bar is drawn up from the axis minimum, so its height is its rate minus {AXIS_MIN:.0%}.\n"
                      f"2. ({hi_sl['ReadmitRate']:.4f} − {AXIS_MIN}) ÷ ({lo_sl['ReadmitRate']:.4f} − {AXIS_MIN}) = "
                      f"**{t4_raw:.1f}**.\n3. Fix the chart: double-click the vertical axis, and in **Format Axis → Axis "
                      "Options → Bounds** set **Minimum** to `0` (or click **Reset** so Excel chooses 0). Turn off "
                      "**Vary colors by point** (**Format Data Series → Fill & Line → Fill**), delete the legend, add data "
                      "labels, and give the chart a title that states the finding.\n4. Right-click the chart → **Edit Alt "
                      "Text** and describe it, for example: *Column chart of 2025 30-day readmission rates by service line. "
                      f"{hi_sl['ServiceLine']} is highest at {hi_sl['ReadmitRate']:.1%} and {lo_sl['ServiceLine']} lowest "
                      f"at {lo_sl['ReadmitRate']:.1%}.*",
             live=f4,
             explanation=f"The real ratio is {hi_sl['ReadmitRate']:.1%} ÷ {lo_sl['ReadmitRate']:.1%} = {t4_true:.1f}, but "
                         f"the truncated chart draws the {hi_sl['ServiceLine']} bar {t4_raw:.1f} times as tall. A bar's "
                         "length is how readers judge its value, so every bar and column chart needs a value axis that "
                         "starts at zero. Line charts are different: their message is in the slope, so a line chart's axis "
                         "may start above zero as long as it's labeled. The other problems on the Makeover chart (rainbow "
                         "colors and a legend that repeats the axis labels) add color without adding information. "
                         + cross(f4)),
        Task("On PayerMix, insert a pie chart of 2025 encounters by PayerType. Add data labels that show the Percentage "
             "(not the Value), formatted with 1 decimal place. What does the Commercial label show?",
             answer=t5, fmt="0.0%", title="Pie chart: Commercial share",
             hint="Format Data Labels → Label Options: tick Percentage, untick Value. Then Number → Percentage, 1 decimal",
             solution="1. Click any cell in tblPayerMix and choose **Insert → Insert Pie or Doughnut Chart → 2-D Pie**.\n"
                      "2. Add labels: click the **Chart Elements** button (**+**) → **Data Labels → More Options…** (Mac: "
                      "**Chart Design → Add Chart Element → Data Labels → More Data Label Options**).\n"
                      "3. In **Label Options**, tick **Percentage** and **Category Name** and untick **Value**. Under "
                      "**Number**, choose **Percentage** with **1** decimal place.\n4. Read the Commercial slice.",
             live=f5,
             explanation=f"Excel computes each slice's percentage itself: {commercial:,} ÷ {pm_total:,} = {t5:.1%}. "
                         "The default label format has no decimals, so it would round to "
                         f"{round(t5 * 100):.0f}%. A pie works here because there are only five parts of one whole and "
                         "two of them dominate. "
                         f"{smallest_payer['PayerType']} ({smallest_payer['Encounters'] / pm_total:.1%}) is barely a sliver, "
                         "which is the usual limit of a pie: with more or smaller slices, a sorted bar chart compares "
                         "parts far more accurately. " + cross(f5)),
        Task(f"On Stays, insert a histogram of LOSDays (all {SAMPLE_N} stays). Format the horizontal axis with Bin width "
             f"{1}, Overflow bin {HIST_OVER}, and Underflow bin {HIST_UNDER}. How many stays fall in the overflow bin "
             f"(longer than {HIST_OVER} days)?",
             answer=t6, title="Histogram: stays in the overflow bin",
             hint="Select D1:D301, then Insert → Insert Statistic Chart → Histogram. Double-click the horizontal axis",
             solution="1. On **Stays**, select **D1:D301** (the LOSDays column with its header).\n"
                      "2. Choose **Insert → Insert Statistic Chart → Histogram**. On a Mac, the same Statistic Chart "
                      "button is on the Insert tab.\n3. Double-click the horizontal axis. Under **Axis Options → Bins**, choose **Bin "
                      f"width** `1`, tick **Overflow bin** and type `{HIST_OVER}`, and tick **Underflow bin** and type "
                      f"`{HIST_UNDER}`.\n4. Hover over the last column, labeled **>{HIST_OVER}**.",
             live=f6,
             explanation=f"A histogram answers \"how are the values distributed?\" by counting values in equal-width bins. "
                         f"Here the tallest bin is {tallest_bin[0]} days with {tallest_bin[1]} stays, and the long right "
                         f"tail ends in the >{HIST_OVER} overflow bin with {t6} stays. That skew is why the median length "
                         f"of stay ({los_median:.2f} days) sits below the mean ({los_mean:.2f} days). The few long stays "
                         "pull the mean up but barely move the median. A bin label such as (3, 4] means "
                         "\"more than 3, up to and including 4\". The overflow and underflow bins stop a few extreme "
                         "stays from stretching the axis. " + cross(f6)),
        Task("On Stays, select D1:E301 (LOSDays and TotalCharges) and insert an XY scatter chart, so LOSDays is on the "
             "horizontal axis. Add a linear trendline and display its equation. By how many dollars do charges rise for "
             "each extra day in the hospital? Confirm with SLOPE and round to the nearest dollar.",
             answer=t7, fmt="#,##0", tol=0.5, answer_display=f"{t7:,.0f} (dollars per extra day)", title="Scatter + trendline: dollars per extra day",
             hint="Insert → Insert Scatter (X, Y) or Bubble Chart → Scatter. Right-click a point → Add Trendline, then "
                  "tick Display Equation on chart. SLOPE takes the Y range first",
             solution="1. Select **D1:E301**. The left column (LOSDays) becomes X, the right one (TotalCharges) becomes Y.\n"
                      "2. Choose **Insert → Insert Scatter (X, Y) or Bubble Chart → Scatter**.\n"
                      "3. Right-click any point → **Add Trendline…**. In **Format Trendline**, keep **Linear** and tick "
                      "**Display Equation on chart** and **Display R-squared value on chart**.\n"
                      f"4. The equation reads about *y = {t7:.0f}x + {reg_ch.intercept:.0f}*. The slope is the number "
                      f"in front of x.\n5. Confirm in the yellow cell: `=SLOPE(Stays!E2:E301,Stays!D2:D301)`.",
             live=f7,
             explanation=f"A linear trendline is the least-squares line through the points, the same line that SLOPE and "
                         f"INTERCEPT calculate. Its slope says that each extra day adds about ${t7:,.0f} in charges on "
                         f"average. R² is {r2_ch:.2f}, so length of stay explains about {r2_ch:.0%} of the variation in "
                         "charges. The rest comes from what happened during the stay, such as surgery, ICU days, and "
                         "imaging. SLOPE takes the Y range first, which is easy to get backwards. " + cross(f7)),
        Task("Make a second scatter chart on Stays with AgeAtAdmit on the horizontal axis and LOSDays on the vertical axis "
             "(C1:D301). Add a linear trendline and display the R-squared value. What is R²? Confirm with RSQ and enter "
             "it to 4 decimal places.",
             answer=t8, fmt="0.0000", tol=0.00006, title="Scatter + trendline: R² for age vs. length of stay",
             hint="RSQ(known_y's, known_x's)",
             solution="1. Select **C1:D301** and choose **Insert → Insert Scatter (X, Y) or Bubble Chart → Scatter**.\n"
                      "2. Right-click a point → **Add Trendline…** → **Linear**, and tick **Display R-squared value on "
                      f"chart**.\n3. The label reads *R² = {t8:.4f}*.\n"
                      "4. Confirm in the yellow cell: `=RSQ(Stays!D2:D301,Stays!C2:C301)`.",
             live=f8,
             explanation=f"R² = {t8:.4f} means age explains under {math.ceil(t8 * 100)}% of the variation in length of "
                         f"stay. The trendline still slopes upward (about {reg_age.slope:.3f} days per "
                         f"year of age, or {reg_age.slope * 10:.2f} days per decade), but the points scatter widely around "
                         "it. Always show R² next to a trendline, because a line can be drawn through any cloud of points, "
                         "related or not. Even a high R² shows association, not cause. " + cross(f8)),
        Task("Back on ED_Hourly, select A1:C25 and insert a combo chart with Arrivals as clustered columns and "
             "AvgDoorToProviderMin as a line on the secondary axis. Which arrival hour has the longest average wait to see "
             "a provider? Type the hour as a number from 0 to 23.",
             answer=t9, accept=[t9 / 24], title="Combo chart: hour with the longest wait",
             hint="Insert → Insert Combo Chart → Clustered Column – Line on Secondary Axis",
             solution="1. Select **A1:C25** on **ED_Hourly**.\n2. Choose **Insert → Insert Combo Chart → Clustered Column "
                      "– Line on Secondary Axis**. (Or insert any chart, then **Chart Design → Change Chart Type → "
                      "Combo**, set AvgDoorToProviderMin to **Line** and tick its **Secondary Axis** box.)\n"
                      "3. Add axis titles with **Chart Elements (+) → Axis Titles** (Mac: **Chart Design → Add Chart "
                      "Element → Axis Titles**): *Arrivals* on the left and *Avg minutes to provider* on the right.\n"
                      "4. Hover over the highest point of the line.",
             live=f9,
             explanation=f"Arrivals peak at {hour_txt(t1)}, and the average wait peaks at {hour_txt(t9)} "
                         f"({by_wait[0]['AvgDoorToProviderMin']} minutes, against {short_wait['AvgDoorToProviderMin']} at "
                         f"{hour_txt(short_wait['Hour'])}). The two series use different units (visits and minutes) and "
                         "very different sizes, so the line needs its own axis. Otherwise it would be squashed flat. "
                         f"The two shapes match closely (correlation {corr_hour:.2f}), so waits rise and fall with "
                         "arrivals. That suggests staffing doesn't keep pace in the busy hours, but the chart alone can't "
                         "prove the cause. Always title both axes on a dual-axis chart so nobody reads minutes off the "
                         "arrivals scale. " + cross(f9)),
        Task("On Denials, select A1:B8 (DenialReason and Claims) and insert a Pareto chart. What cumulative percentage does "
             "the line reach at the second bar (the top two reasons together)? Enter it as a percentage to 1 decimal place.",
             answer=t10, fmt="0.0%", title="Pareto: cumulative share of the top two denial reasons",
             hint="Insert → Insert Statistic Chart → Pareto. Confirm with LARGE and SUM",
             solution="1. Select **A1:B8** on **Denials**.\n2. Choose **Insert → Insert Statistic Chart → Pareto**. Excel "
                      "sorts the reasons from most to fewest claims and adds a cumulative-percentage line on a secondary "
                      "axis that runs to 100%.\n3. Hover over the line at the second bar.\n4. To get the exact value, "
                      "type `=(LARGE(Denials!B2:B8,1)+LARGE(Denials!B2:B8,2))/SUM(Denials!B2:B8)` and format it as a "
                      "percentage with 1 decimal place.",
             live=f10,
             explanation=f"{dn_by_claims[0]['DenialReason']} ({dn_by_claims[0]['Claims']} claims) and "
                         f"{dn_by_claims[1]['DenialReason']} ({dn_by_claims[1]['Claims']}) make up "
                         f"{dn_by_claims[0]['Claims'] + dn_by_claims[1]['Claims']} of {dn_total:,} denied claims. A Pareto "
                         "chart ranks causes so a team can see the \"vital few\" worth fixing first. Here the first two "
                         f"reasons cover {t10:.1%} of denials and the first four cover {cum_claims[3]:.1%}. The built-in "
                         "Pareto sorts the data for you, so the table itself can stay in any order. LARGE(range,1) and "
                         "LARGE(range,2) return the largest and second-largest claim counts, which are the first two "
                         "bars. The shorter cross-check formula that follows hands LARGE the array constant {1,2} (Lesson "
                         "2.5) so it returns both at once, and SUM adds them. " + cross(f10)),
        Task("On Budget, select A4:B14 and insert a waterfall chart. Set the first and last bars as totals. Which step is "
             "the largest drop (the longest downward bar)? Type the Step name as it appears in the table.",
             answer=t11, accept=["Revenue", "Net patient service revenue", "Net revenue", "Patient revenue"],
             title="Waterfall: the largest unfavorable variance",
             hint="Insert → Insert Waterfall, Funnel, Stock, Surface, or Radar Chart → Waterfall. Right-click a bar → Set as Total",
             solution="1. Select **A4:B14** on **Budget**.\n2. Choose **Insert → Insert Waterfall, Funnel, Stock, Surface, "
                      "or Radar Chart → Waterfall** (Mac: **Insert → Waterfall**).\n3. Click the first bar once to "
                      "select the series and once more to select only that bar. Right-click it → **Set as Total**. Do the "
                      "same for the last bar.\n4. Find the longest bar in the **Decrease** color (the legend shows which color "
                      "that is). Hover over it to read its Step name.",
             live=f11,
             explanation=f"A waterfall shows how a starting total becomes an ending total through a series of increases "
                         f"and decreases. 4 West budgeted a margin of ${b_margin:,} and earned ${a_margin:,}. Revenue came "
                         f"in ${-t11_row['Amount']:,} under budget, the biggest drop, and "
                         f"{biggest_up['Step'].lower()} saved ${biggest_up['Amount']:,}, the biggest rise, so expense "
                         "control offset most of the revenue shortfall. Without **Set as Total**, Excel treats the last "
                         "row as one more increase and floats it on top of the running total. " + cross(f11)),
        Task("On ED_Monthly, insert line sparklines in the yellow cells B27:D27 (Data Range B2:D25), one per hospital. "
             "Turn on First Point and Last Point markers. For how many of the three hospitals is the last point (Dec 2025) "
             "higher than the first point (Jan 2024)?",
             answer=t12, title="Sparklines: hospitals that ended higher than they started",
             hint="Insert → Sparklines → Line. Then the Sparkline tab → Show → First Point, Last Point",
             solution="1. Select **B27:D27** on **ED_Monthly**.\n2. Choose **Insert → Sparklines → Line**. In **Data "
                      "Range** type `B2:D25`. **Location Range** already shows `$B$27:$D$27`. Click **OK**. Excel draws "
                      "one sparkline per column.\n3. On the **Sparkline** tab, tick **First Point** and **Last Point** "
                      "(and **High Point** if you like).\n4. Compare each sparkline's two end markers. To check, compare "
                      f"row {mlast} with row {mfirst}.",
             live=f12,
             explanation=f"{', '.join(up[:-1]) + ' and ' + up[-1] if len(up) > 1 else up[0]} ended higher than they "
                         f"started. {', '.join(down)} ended lower ({first_m[down[0]]} visits in Jan 2024, "
                         f"{last_m[down[0]]} in Dec 2025). A sparkline is a word-sized chart in a cell, so it shows each "
                         "series' shape next to the numbers. By default every sparkline gets its own vertical scale, which "
                         "is right for comparing a series with itself but wrong for comparing hospitals. When heights should "
                         "be comparable, open **Sparkline → Axis** and set both the minimum and the maximum to **Same for "
                         "All Sparklines**. " + cross(f12)),
        Task("Write a formula in the yellow cell that builds this chart title from tblEDMonthly: ED visits by facility, "
             "[first month] to [last month] ([total visits] visits). Show each month as a three-letter month and year, and "
             "the total with a thousands separator. Example of the pattern: ED visits by facility, Mar 2023 to Feb 2024 "
             "(9,876 visits). Then link the title of your Task 2 line chart to this cell.",
             answer=t13, accept=t13_accept, solution=f13, title="Dynamic chart title",
             hint='TEXT(MIN(…),"mmm yyyy") and TEXT(SUM(…),"#,##0"), joined with &',
             live=True,
             explanation="Build the title in a cell, because a chart title can show text or one cell reference, but not "
                         "a formula. To link it, click the chart title on ED_Monthly, type `=` in the formula bar, click "
                         "the **Practice** sheet tab, click this yellow cell, and press **Enter**. The formula bar then "
                         "shows a reference such as `=Practice!$D$18`. "
                         "TEXT turns the dates and the total into formatted text. Without it, `&` would join the raw "
                         f"serial number {excel_serial(monthly[0]['Month']):.0f} instead of Jan 2024. Because the formula "
                         "uses tblEDMonthly, adding January 2026 as a new row updates the chart and its title together."),
    ]

    # ------------------------------------------------------------------ bonus
    L.bonus_title = "Bonus: Charts for the operations review"
    L.bonus_scenario = (
        "Two directors bring questions to Bluestone's monthly operations review. The ED medical director believes patients "
        "leave without being seen (LWBS) mainly in busy months, and the ED's target is an LWBS rate of "
        f"{LWBS_TARGET:.0%} or less. The revenue-cycle director wants to know which denial reasons to work first, ranked by "
        "dollars at risk (DeniedCharges) rather than by claim counts. Build a combo chart and a scatter chart for the "
        "first question and a Pareto chart for the second, then check the numbers behind them.")
    rate_col = mo.col("LWBSRate")
    L.bonus = [
        Task(f"On ED_Monthly, fill the yellow LWBSRate column with LWBS ÷ Total for each month. The column is already "
             f"formatted as a percentage. The gray cell counts the months above the {LWBS_TARGET:.0%} target. How many of "
             "the 24 months missed the target?",
             answer=b1, title="LWBSRate column (months above the 2% target)",
             solution="=[@LWBS]/[@Total]",
             summary=f'=IF(COUNT({lwbs_rng})=0,"",COUNTIF({lwbs_rng},">{LWBS_TARGET}"))',
             fill={"range": f"ED_Monthly!{rate_col}{mfirst}:{rate_col}{mlast}",
                   "formula": f"={mo.col('LWBS')}{mfirst}/{mo.col('Total')}{mfirst}"},
             live=fb1_live,
             hint=f"Write one formula in {rate_col}{mfirst} with [@Column] references (Lesson 3.1). The Table fills the "
                  "other 23 rows",
             explanation=f"Type the formula in {rate_col}{mfirst} (`=F{mfirst}/E{mfirst}` works too). [@LWBS] means "
                         "\"the LWBS value in this row\", and the Table fills the column for you. Rates make months comparable even though volume swings by "
                         f"{busiest_month['Total'] - quietest_month['Total']} visits between the busiest and quietest "
                         f"months. The gray cell runs `COUNTIF({rate_col}{mfirst}:{rate_col}{mlast},\">{LWBS_TARGET}\")` on "
                         f"your column, and {b1} of the 24 months were above {LWBS_TARGET:.0%}."),
        Task("Build a combo chart on ED_Monthly with Total as clustered columns and LWBSRate as a line on the secondary "
             "axis (select Month, then Ctrl+click or ⌘+click Total and LWBSRate). Which month had the highest LWBS rate? Type the "
             "month and year.",
             answer=b2, fmt="mmm yyyy", answer_display=_month_label(b2), title="Combo chart: month with the highest LWBS rate",
             hint="Select A1:A25, Ctrl+click E1:E25 and G1:G25, then Insert → Insert Combo Chart",
             solution="1. Select **A1:A25**, then hold **Ctrl** (Mac: **⌘**) and select **E1:E25** and **G1:G25**.\n"
                      "2. Choose **Insert → Insert Combo Chart → Clustered Column – Line on Secondary Axis**.\n"
                      "3. Title both axes (*ED visits* and *LWBS rate*). If the secondary axis shows too many decimals, set "
                      "its number format to a percentage with 1 decimal place (**Format Axis → Number**).\n4. Hover over the highest point of the line.",
             live=fb2,
             explanation=f"{_month_label(b2)} had {b2_row['LWBS']} LWBS out of {b2_row['Total']} visits ({b2_rate:.2%}), "
                         f"yet it was {b2_note}. The busiest month, {_month_label(busiest_month['Month'])} "
                         f"({busiest_month['Total']} visits), had an LWBS rate of only {busiest_rate:.2%}. On the combo "
                         "chart, the line's peaks don't line up with the tallest columns, which is the first hint that "
                         "volume isn't the whole story. " + cross(fb2)),
        Task("Test the director's theory with a scatter chart of Total (horizontal axis) against LWBSRate (vertical axis), "
             "with a linear trendline and its R². What is R²? Enter it to 3 decimal places.",
             answer=b3, fmt="0.000", tol=0.0006, title="Scatter: how much volume explains LWBS",
             hint="Select E1:E25, Ctrl+click G1:G25, then Insert → Scatter. RSQ(known_y's, known_x's) confirms it",
             solution="1. Select **E1:E25**, then **Ctrl+click** (Mac: **⌘+click**) **G1:G25**. The left column (Total) "
                      "becomes X.\n2. Choose **Insert → Insert Scatter (X, Y) or Bubble Chart → Scatter**.\n"
                      "3. Right-click a point → **Add Trendline…** → **Linear**, tick **Display R-squared value on "
                      "chart**.\n4. Confirm: `=RSQ(ED_Monthly!G2:G25,ED_Monthly!E2:E25)`.",
             live=fb3,
             explanation=f"R² = {b3:.3f}. The trendline slopes upward (about {reg_lwbs.slope * 100 * 100:.2f} percentage "
                         "points of LWBS per 100 extra visits), so busier months do run slightly higher, but volume explains "
                         f"only about {b3:.0%} of the month-to-month variation. The director should look at staffing, "
                         "boarding, and triage flow as well. Twenty-four points is a small sample, so treat any pattern as "
                         "a lead to investigate rather than proof. " + cross(fb3)),
        Task("On Denials, build a Pareto chart by DeniedCharges by hand: sort the table by DeniedCharges (largest first), "
             "fill CumulativePct with each row's running share of total DeniedCharges, then insert a combo chart with "
             "DeniedCharges as columns and CumulativePct as a line on the secondary axis (fix that axis at 0% to 100%). "
             f"How many reasons does it take to reach at least {PARETO_CUT:.0%} of denied charges?",
             answer=b4, title="Manual Pareto: reasons needed to reach 80% of denied charges",
             hint="Follow 'A Pareto by hand' in Guide section 12. DeniedCharges is column C. Sort before you chart",
             solution="1. Click a DeniedCharges cell → **Data → Sort Largest to Smallest**.\n"
                      "2. In **D2** type `=SUM($C$2:C2)/SUM($C$2:$C$8)` and press **Enter**. The Table fills it down, and "
                      "the column is already formatted as a percentage.\n"
                      "3. Select **A1:A8**, then **Ctrl+click** (Mac: **⌘+click**) **C1:D8**, and choose **Insert → Insert "
                      "Combo Chart → Clustered Column – Line on Secondary Axis**.\n"
                      "4. Double-click the secondary axis and set **Minimum** `0` and **Maximum** `1`. Optionally set the "
                      "column **Gap Width** to about 10%.\n"
                      f"5. Find the first bar where the line reaches {PARETO_CUT:.0%}.",
             live=fb4,
             explanation=f"Cumulative shares by dollars: "
                         + ", ".join(f"{r['DenialReason']} {c:.1%}" for r, c in zip(dn_by_dollars, cum_dollars))
                         + f". The line first passes {PARETO_CUT:.0%} at reason {b4}, so {b4} of the "
                           f"{len(denials)} reasons hold {cum_dollars[b4 - 1]:.1%} of the ${dollars_total:,.0f} at risk. "
                           "The expanding range `$C$2:C2` (Lesson 3.3) has an absolute start and a relative end "
                           "(Lesson 1.5), so its start stays fixed while its end moves down one row at a time. Building a Pareto by hand takes longer than the built-in chart, but it "
                           "works in every Excel version and lets you add an 80% reference line or label the cut-off. "
                           "Cross-check (counts the reasons whose running share is still below 80%, then adds one): "
                           f"`{fb4}`"),
        Task("Compare this Pareto with the one you built by claim count in Task 10. Exactly one reason ranks higher by "
             "denied charges than by number of claims. Which one? Type it as it appears in the table.",
             answer=b5, title="Which reason climbs when you rank by dollars",
             hint="Compare the order of the bars in the two Pareto charts",
             solution=f"Read the bar order in both charts. By claims, {b5} is number {rank_claims[b5]}, and by denied charges "
                      f"it's number {rank_dollars[b5]}, so it swaps places with {swapped}.",
             live=fb5,
             explanation=f"{b5} denials are fewer ({b5_row['Claims']} claims against {swapped_row['Claims']} for {swapped}) "
                         f"but larger: about ${avg_claim[b5]:,.0f} per claim against ${avg_claim[swapped]:,.0f}. That's "
                         "why the director asked for dollars. A count Pareto ranks the work queue by volume, and a dollar "
                         "Pareto ranks it by money at risk. Put the two charts side by side, at the same size, so the "
                         "committee sees the swap at a glance. Cross-check (COUNTIF(range,\">\"&range) gives each "
                         "reason's rank minus 1, so MATCH finds the reason whose dollar rank beats its claim rank. In "
                         "Excel 2019 or earlier, confirm it with Ctrl + Shift + Enter): `" + fb5 + "`"),
    ]

    # ------------------------------------------------------------------ customize: Makeover, sparkline cells, Chart Key
    steps_wf = []
    run_total = 0
    for i, r in enumerate(budget):
        if i == 0 or i == len(budget) - 1:
            steps_wf.append((r["Step"], 0, 0, 0, r["Amount"]))
            run_total = r["Amount"]
        else:
            new = run_total + r["Amount"]
            steps_wf.append((r["Step"], min(run_total, new), max(r["Amount"], 0), max(-r["Amount"], 0), 0))
            run_total = new
    assert min(s[1] for s in steps_wf) >= 0, "stacked-column waterfall needs a positive running total"

    @L.customize
    def _custom(wb, lesson, selftest):
        title_font = Font(bold=True, size=14, color=NAVY)
        note_font = Font(italic=True, color="595959")
        assert lesson.tasks[12].answer_cell == "D18", "Task 13's explanation quotes =Practice!$D$18"
        wb.loaded_theme = _office_2013_theme()

        def no_delete(chart, *axes):
            for ax in axes or (chart.x_axis, chart.y_axis):
                ax.delete = False

        # ---- sparkline cells under the ED_Monthly table
        ws_m = wb["ED_Monthly"]
        spark_row = mlast + 2
        ws_m.cell(row=spark_row, column=1, value="Trend").font = Font(bold=True)
        for c in range(2, 5):
            cell = ws_m.cell(row=spark_row, column=c)
            cell.fill = INPUT_FILL
            cell.border = INPUT_BORDER
        ws_m.row_dimensions[spark_row].height = 30
        ws_m.cell(row=spark_row + 1, column=1, value="Task 12: put line sparklines in the yellow cells B27:D27.").font = note_font
        assert spark_row == 27

        # ---- Makeover sheet: a deliberately misleading chart
        mk = wb.create_sheet("Makeover")
        mk.sheet_properties.tabColor = "BF9000"
        mk["A1"] = "Chart makeover (Task 4)"
        mk["A1"].font = title_font
        mk["A2"] = ("A colleague made this chart of 2025 readmission rates from the Readmits sheet. It has several design "
                    "problems. Answer Task 4, then fix the chart in place.")
        mk["A2"].font = note_font
        ws_r = wb["Readmits"]
        bad = BarChart()
        bad.type = "col"
        bad.title = "Readmission Rates"
        bad.varyColors = True
        bad.add_data(Reference(ws_r, min_col=4, min_row=ra.first_row - 1, max_row=ra.last_row), titles_from_data=True)
        bad.set_categories(Reference(ws_r, min_col=1, min_row=ra.first_row, max_row=ra.last_row))
        bad.y_axis.scaling.min = AXIS_MIN
        bad.y_axis.number_format = "0%"
        bad.y_axis.majorUnit = 0.02
        no_delete(bad)
        bad.legend.position = "r"
        bad.height, bad.width = 9, 18
        mk.add_chart(bad, "A4")
        mk.column_dimensions["A"].width = 12

        # ---- Chart Key (hidden): reference charts
        ws = wb.create_sheet("Chart Key")
        ws.sheet_properties.tabColor = "C00000"
        ws["A1"] = "🔑 Chart Key — Lesson 3.5"
        ws["A1"].font = Font(bold=True, size=16, color="7B2C2C")
        ws["A2"] = ("Reference versions of every chart in the lesson. Helper tables on the left feed the charts that need "
                    "sorted or pre-binned data. Your colors, sizes, and fonts can differ, so compare the shapes and the numbers.")
        ws["A2"].font = note_font
        ws.column_dimensions["A"].width = 30
        for letter in "BCDEF":
            ws.column_dimensions[letter].width = 14
        row = [4]

        def block(title, headers, body, fmts, height=20):
            """Write a titled helper table at column A; return (first data row, last data row, anchor cell)."""
            r = row[0]
            ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12, color=NAVY)
            hdr = r + 1
            for j, h in enumerate(headers, 1):
                c = ws.cell(row=hdr, column=j, value=h)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = HEADER_FILL
                c.border = BOX
                c.alignment = Alignment(horizontal="left" if j == 1 else "right", wrap_text=True)
            for i, vals in enumerate(body):
                for j, v in enumerate(vals, 1):
                    c = ws.cell(row=hdr + 1 + i, column=j, value=v)
                    if fmts[j - 1]:
                        c.number_format = fmts[j - 1]
            row[0] = r + max(height, len(body) + 4)
            return hdr, hdr + len(body), f"H{r + 1}"

        def size(chart, h=7.5, w=16):
            chart.height, chart.width = h, w

        ws_h, ws_s, ws_p = wb["ED_Hourly"], wb["Stays"], wb["PayerMix"]

        # Task 1 column
        _, _, anchor = block("Task 1 · Column chart: 2025 ED arrivals by hour of day", ["Data"],
                             [["ED_Hourly!A1:B25"]], [None])
        c = BarChart()
        c.type = "col"
        c.title = f"ED arrivals peak in the afternoon ({hour_txt(t1)} busiest, 2025)"
        c.add_data(Reference(ws_h, min_col=2, min_row=1, max_row=25), titles_from_data=True)
        c.set_categories(Reference(ws_h, min_col=1, min_row=2, max_row=25))
        c.legend = None
        c.gapWidth = 50
        c.y_axis.title = "Arrivals"
        c.y_axis.majorGridlines = None
        c.series[0].graphicalProperties.solidFill = BLUE
        no_delete(c)
        size(c)
        ws.add_chart(c, anchor)

        # Tasks 2 + 13 line chart with the dynamic title text
        _, _, anchor = block("Tasks 2 & 13 · Line chart: monthly ED visits by hospital (title linked to Task 13's cell)",
                             ["Data", "Title text"], [["ED_Monthly!A1:D25", t13]], [None, None])
        c = LineChart()
        c.title = t13
        c.add_data(Reference(ws_m, min_col=2, max_col=4, min_row=1, max_row=mlast), titles_from_data=True)
        c.set_categories(Reference(ws_m, min_col=1, min_row=mfirst, max_row=mlast))
        c.y_axis.crossAx = 500
        c.x_axis = DateAxis(crossAx=100)
        c.x_axis.number_format = "mmm yy"
        c.x_axis.majorTimeUnit = "months"
        c.x_axis.majorUnit = 3
        c.y_axis.title = "ED visits"
        c.legend.position = "b"
        for s, color in zip(c.series, (BLUE, ORANGE, SKY)):
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 22000
            s.smooth = False
        no_delete(c)
        size(c, 8, 18)
        ws.add_chart(c, anchor)

        # Task 3 bar chart (sorted, reversed categories)
        first, last, anchor = block("Task 3 · Bar chart: readmission rate by service line (sorted, categories reversed)",
                                    ["ServiceLine", "ReadmitRate"],
                                    [[r["ServiceLine"], r["ReadmitRate"]] for r in ra_sorted], [None, "0.0%"])
        c = BarChart()
        c.type = "bar"
        c.title = f"{hi_sl['ServiceLine']} has the highest 30-day readmission rate (2025)"
        c.add_data(Reference(ws, min_col=2, min_row=first, max_row=last), titles_from_data=True)
        c.set_categories(Reference(ws, min_col=1, min_row=first + 1, max_row=last))
        c.x_axis.scaling.orientation = "maxMin"
        c.y_axis.crosses = "max"
        c.y_axis.number_format = "0%"
        c.y_axis.scaling.min = 0
        c.y_axis.majorGridlines = None
        c.legend = None
        c.gapWidth = 60
        c.series[0].graphicalProperties.solidFill = BLUE
        c.dataLabels = _Labels("0.0%", val=True)
        no_delete(c)
        size(c)
        ws.add_chart(c, anchor)

        # Task 4 fixed makeover chart (axis from zero, one color, direct labels)
        first, last, anchor = block("Task 4 · Makeover fixed: axis from 0%, one color, labels instead of a legend",
                                    ["ServiceLine", "ReadmitRate"],
                                    [[r["ServiceLine"], r["ReadmitRate"]] for r in ra_sorted], [None, "0.0%"])
        c = BarChart()
        c.type = "col"
        c.title = f"Readmission rates range from {lo_sl['ReadmitRate']:.0%} to {hi_sl['ReadmitRate']:.0%} (2025)"
        c.add_data(Reference(ws, min_col=2, min_row=first, max_row=last), titles_from_data=True)
        c.set_categories(Reference(ws, min_col=1, min_row=first + 1, max_row=last))
        c.y_axis.scaling.min = 0
        c.y_axis.number_format = "0%"
        c.y_axis.majorGridlines = None
        c.legend = None
        c.series[0].graphicalProperties.solidFill = BLUE
        c.dataLabels = _Labels("0.0%", val=True)
        no_delete(c)
        size(c)
        ws.add_chart(c, anchor)

        # Task 5 pie
        _, _, anchor = block("Task 5 · Pie chart: 2025 encounters by payer type (Percentage labels, 1 decimal)",
                             ["Data"], [["PayerMix!A1:B6"]], [None])
        c = PieChart()
        c.title = "Government and Commercial payers cover most 2025 encounters"
        c.add_data(Reference(ws_p, min_col=2, min_row=1, max_row=pm.last_row), titles_from_data=True)
        c.set_categories(Reference(ws_p, min_col=1, min_row=2, max_row=pm.last_row))
        c.dataLabels = _Labels("0.0%", pct=True, cat=True)
        c.dataLabels.showLeaderLines = True
        c.legend = None
        size(c, 10, 16)
        ws.add_chart(c, anchor)

        # Task 6 histogram (pre-binned)
        first, last, anchor = block(f"Task 6 · Histogram of LOSDays: bin width 1, underflow {HIST_UNDER}, overflow "
                                    f"{HIST_OVER} (classic column version)", ["Bin (days)", "Stays"],
                                    [[b, n] for b, n in hist_bins], [None, "0"])
        c = BarChart()
        c.type = "col"
        mid = sum(n for b, n in hist_bins if b in ("(2, 3]", "(3, 4]", "(4, 5]"))
        assert mid > SAMPLE_N / 2, "the chart title claims most stays last 2-5 days"
        c.title = f"Most stays last 2–5 days, and {t6} run past {HIST_OVER} days"
        c.add_data(Reference(ws, min_col=2, min_row=first, max_row=last), titles_from_data=True)
        c.set_categories(Reference(ws, min_col=1, min_row=first + 1, max_row=last))
        c.gapWidth = 6
        c.legend = None
        c.y_axis.title = "Stays"
        c.x_axis.title = "Length of stay (days)"
        c.y_axis.majorGridlines = None
        c.series[0].graphicalProperties.solidFill = BLUE
        c.series[0].graphicalProperties.line.solidFill = "FFFFFF"
        c.dataLabels = _Labels(val=True)
        no_delete(c)
        size(c)
        ws.add_chart(c, anchor)

        def scatter(xcol, ycol, title, xt, yt, yfmt=None):
            sc = ScatterChart()
            sc.title = title
            sc.style = 13
            xs = Reference(ws_s, min_col=xcol, min_row=sy.first_row, max_row=sy.last_row)
            ys = Reference(ws_s, min_col=ycol, min_row=sy.first_row - 1, max_row=sy.last_row)
            s = Series(ys, xs, title_from_data=True)
            s.marker = Marker(symbol="circle", size=5)
            s.marker.graphicalProperties = GraphicalProperties(solidFill=SKY)
            s.marker.graphicalProperties.line.solidFill = BLUE
            s.graphicalProperties.line.noFill = True
            s.trendline = Trendline(trendlineType="linear", dispEq=True, dispRSqr=True, trendlineLbl=TrendlineLabel(
                numFmt=NumFmt(formatCode="#,##0.0000", sourceLinked=False)))
            sc.series.append(s)
            sc.x_axis.title = xt
            sc.y_axis.title = yt
            if yfmt:
                sc.y_axis.number_format = yfmt
            sc.legend = None
            no_delete(sc)
            size(sc, 8, 16)
            return sc

        _, _, anchor = block("Task 7 · Scatter: LOSDays (x) vs TotalCharges (y), linear trendline", ["SLOPE", "INTERCEPT", "R²"],
                             [[t7, reg_ch.intercept, r2_ch]], ["#,##0.00", "#,##0.00", "0.0000"])
        ws.add_chart(scatter(sy.headers.index("LOSDays") + 1, sy.headers.index("TotalCharges") + 1,
                             f"Each extra day adds about ${t7:,.0f} in charges", "Length of stay (days)",
                             "Total charges ($)", "#,##0"), anchor)
        _, _, anchor = block("Task 8 · Scatter: AgeAtAdmit (x) vs LOSDays (y), linear trendline", ["SLOPE", "INTERCEPT", "R²"],
                             [[reg_age.slope, reg_age.intercept, t8]], ["0.0000", "0.0000", "0.0000"])
        ws.add_chart(scatter(sy.headers.index("AgeAtAdmit") + 1, sy.headers.index("LOSDays") + 1,
                             f"Age explains almost none of the variation in LOS (R² = {t8:.3f})", "Age at admission",
                             "Length of stay (days)"), anchor)

        def combo(cats, bars, line, title, bar_t, line_t, line_fmt, line_max=None, gap=60):
            b = BarChart()
            b.type = "col"
            b.title = title
            b.add_data(bars, titles_from_data=True)
            b.set_categories(cats)
            b.y_axis.title = bar_t
            b.y_axis.majorGridlines = None
            b.gapWidth = gap
            b.series[0].graphicalProperties.solidFill = SKY
            ln = LineChart()
            ln.add_data(line, titles_from_data=True)
            ln.y_axis.axId = 200
            ln.y_axis.title = line_t
            ln.y_axis.number_format = line_fmt
            ln.y_axis.crosses = "max"
            ln.y_axis.majorGridlines = None
            if line_max is not None:
                ln.y_axis.scaling.min = 0
                ln.y_axis.scaling.max = line_max
            ln.series[0].graphicalProperties.line.solidFill = VERMILLION
            ln.series[0].graphicalProperties.line.width = 28000
            ln.series[0].smooth = False
            no_delete(b)
            no_delete(ln, ln.y_axis)
            b.legend.position = "b"
            b += ln
            size(b, 8, 18)
            return b

        _, _, anchor = block("Task 9 · Combo: arrivals (columns) + average door-to-provider minutes (line, secondary axis)",
                             ["Data"], [["ED_Hourly!A1:C25"]], [None])
        ws.add_chart(combo(Reference(ws_h, min_col=1, min_row=2, max_row=25),
                           Reference(ws_h, min_col=2, min_row=1, max_row=25),
                           Reference(ws_h, min_col=3, min_row=1, max_row=25),
                           "Waits rise and fall with arrivals (2025)", "Arrivals", "Avg minutes to provider", "0"), anchor)

        first, last, anchor = block("Task 10 · Pareto by claim count (sorted columns + cumulative % line)",
                                    ["DenialReason", "Claims", "Cumulative %"],
                                    [[r["DenialReason"], r["Claims"], cm] for r, cm in zip(dn_by_claims, cum_claims)],
                                    [None, "#,##0", "0.0%"])
        ws.add_chart(combo(Reference(ws, min_col=1, min_row=first + 1, max_row=last),
                           Reference(ws, min_col=2, min_row=first, max_row=last),
                           Reference(ws, min_col=3, min_row=first, max_row=last),
                           f"Two reasons account for {t10:.0%} of 2025 denials", "Denied claims", "Cumulative %", "0%",
                           line_max=1, gap=10), anchor)

        first, last, anchor = block("Task 11 · Waterfall (classic stacked-column version; Base series has no fill)",
                                    ["Step", "Base", "Increase", "Decrease", "Total"],
                                    [list(s) for s in steps_wf], [None, "#,##0", "#,##0", "#,##0", "#,##0"])
        c = BarChart()
        c.type = "col"
        c.grouping = "stacked"
        c.overlap = 100
        c.gapWidth = 40
        c.title = "Revenue shortfall, mostly offset by labor savings (4 West, FY2025)"
        c.add_data(Reference(ws, min_col=2, max_col=5, min_row=first, max_row=last), titles_from_data=True)
        c.set_categories(Reference(ws, min_col=1, min_row=first + 1, max_row=last))
        c.series[0].graphicalProperties.noFill = True
        c.series[0].graphicalProperties.line.noFill = True
        for s, color in zip(c.series[1:], (BLUE, ORANGE, GREY)):
            s.graphicalProperties.solidFill = color
        c.y_axis.number_format = "#,##0"
        c.y_axis.majorGridlines = None
        c.legend.position = "b"
        no_delete(c)
        size(c, 11, 20)
        ws.add_chart(c, anchor)

        _, _, anchor = block("Task 12 · Sparklines", ["Expected result"],
                             [[f"{h}: {'ends higher' if h in up else 'ends lower'} (Jan 2024 {first_m[h]} → Dec 2025 {last_m[h]})"]
                              for h in FAC_COLS.values()], [None], height=6)
        ws.cell(row=row[0] - 2, column=1,
                value="This sheet can't show sparklines, so compare your sparklines with the values above.").font = note_font

        # Bonus charts
        first, last, anchor = block("Bonus B2 · Combo: monthly ED visits (columns) + LWBS rate (line, secondary axis)",
                                    ["Month", "Total", "LWBS rate"],
                                    [[r["Month"], r["Total"], r["LWBS"] / r["Total"]] for r in monthly],
                                    ["mmm yy", "#,##0", "0.00%"], height=30)
        ws.add_chart(combo(Reference(ws, min_col=1, min_row=first + 1, max_row=last),
                           Reference(ws, min_col=2, min_row=first, max_row=last),
                           Reference(ws, min_col=3, min_row=first, max_row=last),
                           f"LWBS spikes don't follow volume: {_month_label(b2)} peaked at {b2_rate:.1%}",
                           "ED visits", "LWBS rate", "0.0%"), anchor)
        sc = ScatterChart()
        sc.title = f"Volume explains only {b3:.0%} of LWBS variation (R² = {b3:.3f})"
        sc.style = 13
        s = Series(Reference(ws, min_col=3, min_row=first, max_row=last),
                   Reference(ws, min_col=2, min_row=first + 1, max_row=last), title_from_data=True)
        s.marker = Marker(symbol="circle", size=7)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=SKY)
        s.marker.graphicalProperties.line.solidFill = BLUE
        s.graphicalProperties.line.noFill = True
        s.trendline = Trendline(trendlineType="linear", dispEq=False, dispRSqr=True, trendlineLbl=TrendlineLabel(
            numFmt=NumFmt(formatCode="0.000", sourceLinked=False)))
        sc.series.append(s)
        sc.x_axis.title = "ED visits in the month"
        sc.y_axis.title = "LWBS rate"
        sc.y_axis.number_format = "0.0%"
        sc.legend = None
        no_delete(sc)
        size(sc, 8, 16)
        ws.add_chart(sc, f"T{first - 1}")

        first, last, anchor = block("Bonus B4–B5 · Manual Pareto by DeniedCharges", ["DenialReason", "DeniedCharges",
                                                                                     "CumulativePct"],
                                    [[r["DenialReason"], r["DeniedCharges"], cm] for r, cm in zip(dn_by_dollars, cum_dollars)],
                                    [None, "#,##0", "0.0%"])
        ws.add_chart(combo(Reference(ws, min_col=1, min_row=first + 1, max_row=last),
                           Reference(ws, min_col=2, min_row=first, max_row=last),
                           Reference(ws, min_col=3, min_row=first, max_row=last),
                           f"{b4} reasons hold {cum_dollars[b4 - 1]:.0%} of denied charges", "Denied charges ($)",
                           "Cumulative %", "0%", line_max=1, gap=10), anchor)
        ws.sheet_state = "hidden"

        # Answer keys: show the Markdown steps and explanations as plain text in Excel.
        for key_name in (lesson.key_sheet, lesson.bonus_key_sheet):
            ks = wb[key_name]
            for r in range(5, ks.max_row + 1):
                for cc in (4, 6):  # sample solution and explanation columns
                    v = ks.cell(row=r, column=cc).value
                    if isinstance(v, str) and not v.startswith("="):
                        ks.cell(row=r, column=cc).value = re.sub(r"\*([^*\n]+)\*", r"\1", _plain(v))
        for sheet in (ws, mk):
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True

    L.sheet_order = ["Start Here", "Practice", "ED_Hourly", "ED_Monthly", "Readmits", "Makeover", "PayerMix", "Stays",
                     "Denials", "Budget", "Bonus", "Answer Key", "Bonus Key", "Chart Key"]
    return L
