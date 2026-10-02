"""Lesson 3.1 · Excel Tables, Structured References & Named Ranges.

The supply inventory ships as a PLAIN RANGE on purpose: task 1 has the learner convert it to a Table
named tblInventory. Because that Table (and the names ReportDate, ExpiringWindowDays, CycleCountDays,
DiscountMin, DiscountPct, and the Table tblVendors) only exist after the learner creates them, nothing in
the pristine workbook may refer to them directly (Excel can refuse to open a file whose formulas point at a
missing Table). So:

* Practice/Bonus summaries use plain A1 ranges, or INDIRECT("tblInventory") / INDIRECT("ReportDate"), which
  return #REF! (wrapped in IFERROR → "") until the learner creates the Table or name.
* Answer Key live formulas use A1 ranges (or live=False when the formula only makes sense with the
  learner's Tables).
* In self-test mode the ``_simulate_learner`` hook builds exactly what the learner would: the Table with
  its Total Row and the bonus OrderCost column, the vendor Table with its calculated column, and the five
  defined names. Then the sample solutions (written with structured references) are typed into the
  answer cells and must show ✔.
"""
from __future__ import annotations

from datetime import date, timedelta

from openpyxl.styles import Alignment, Font
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

from xlcourse import Lesson, Task, data
from xlcourse.xlfn import to_file_formula

CODE = "3.1"
TBL = "tblInventory"
VTBL = "tblVendors"

REPORT_DATE = data.AS_OF          # 12/31/2025
EXPIRING_WINDOW_DAYS = 90
CYCLE_COUNT_DAYS = 30             # named constant (task 12)
DISCOUNT_MIN = 20000              # named constant (bonus)
DISCOUNT_PCT = 0.02               # named constant (bonus)
SLICER_FACILITY = "Cedar Ridge Medical Center"
SLICER_CATEGORY = "Surgical"
SHARE_CATEGORY = "Orthopedic Implants"


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="03-data-analysis", slug="01-tables-named-ranges",
        title="Excel Tables, Structured References & Named Ranges", level="Intermediate", minutes=50,
        objectives=[
            "Convert ranges to Excel Tables and name them",
            "Write structured references like tblInventory[UnitCost] and [@QtyOnHand]",
            "Use calculated columns, the Total Row, and slicers on Tables",
            "Create, manage, and use named ranges and named constants",
        ],
    )

    # ------------------------------------------------------------------ data
    deps = data.index(data.load("departments"), "DeptID")
    facs = data.index(data.load("facilities"), "FacilityID")
    inv = data.load("supply_inventory")
    inv.sort(key=lambda r: r["StockID"])
    for r in inv:  # pre-join readable location names (this lesson isn't about lookups)
        d = deps[r["LocationDeptID"]]
        r["Facility"] = facs[d["FacilityID"]]["FacilityName"]
        r["Location"] = d["DeptName"]
    n_rooms = len({r["LocationDeptID"] for r in inv})
    n_facs = len({r["Facility"] for r in inv})
    words = {2: "two", 3: "three", 4: "four"}
    L.data_note = (f"Supply inventory snapshot for {n_rooms} storerooms across {words.get(n_facs, n_facs)} Bluestone "
                   f"hospitals ({len(inv)} stock rows, as of {REPORT_DATE:%m/%d/%Y}), a Settings sheet with the report "
                   f"date and expiry window, and Purchasing's vendor list.")

    cols = ["StockID", "SKU", "ItemDescription", "Category", "Vendor", "UnitOfMeasure", "UnitCost", "Facility",
            "Location", "QtyOnHand", "ParLevel", "ReorderPoint", "ReorderQty", "LeadTimeDays", "LastCountDate",
            "ExpirationDate", "LotNumber"]
    sd = L.add_table_sheet(
        "Inventory", inv, columns=cols, as_table=False, extra_cols=["ExtendedValue", "NeedsReorder"],
        formats={"UnitCost": "#,##0.00", "ExtendedValue": "#,##0.00"},
        widths={"ItemDescription": 40, "Facility": 30, "Location": 26, "ExtendedValue": 15, "NeedsReorder": 15,
                "LastCountDate": 14, "ExpirationDate": 15},
    )
    settings_rows = [
        {"Setting": "ReportDate", "Value": REPORT_DATE,
         "Notes": "The 'as of' date of this inventory snapshot. Use it instead of TODAY() so answers never change."},
        {"Setting": "ExpiringWindowDays", "Value": EXPIRING_WINDOW_DAYS,
         "Notes": "Perishable stock expiring within this many days after ReportDate gets flagged for use-first."},
    ]
    st = L.add_table_sheet("Settings", settings_rows, columns=["Setting", "Value", "Notes"], as_table=False,
                           widths={"Setting": 22, "Value": 13, "Notes": 70})
    vendor_names = sorted({r["Vendor"] for r in inv})
    vendors = [{"Vendor": v, "AccountNo": f"BHS-V{1001 + i}"} for i, v in enumerate(vendor_names)]
    vd = L.add_table_sheet("Vendors", vendors, columns=["Vendor", "AccountNo"], as_table=False,
                           extra_cols=["ReorderCost"], formats={"ReorderCost": "#,##0.00"},
                           widths={"Vendor": 30, "AccountNo": 13, "ReorderCost": 15})

    first, last = sd.first_row, sd.last_row
    n = len(inv)

    def a1(col: str) -> str:
        """Plain A1 range for a data column, e.g. Inventory!$J$2:$J$258 (used in summaries and the key)."""
        return sd.rng(col)

    def col_range(letter: str) -> str:
        return f"Inventory!${letter}${first}:${letter}${last}"

    ev_rng, nr_rng = a1("ExtendedValue"), a1("NeedsReorder")
    oc_letter = "T"  # the learner adds OrderCost here (first empty column right of the range) in the bonus
    assert sd.col("NeedsReorder") == "S", "OrderCost must sit immediately right of the range"
    oc_rng = col_range(oc_letter)
    report_cell = f"Settings!${st.col('Value')}${st.first_row}"
    window_cell = f"Settings!${st.col('Value')}${st.first_row + 1}"
    qty, cost, rp = a1("QtyOnHand"), a1("UnitCost"), a1("ReorderPoint")

    # ------------------------------------------------------------------ answers (computed in Python)
    for r in inv:
        r["ExtendedValue"] = r["QtyOnHand"] * r["UnitCost"]
        r["NeedsReorder"] = r["QtyOnHand"] <= r["ReorderPoint"]
        r["OrderCost"] = r["ReorderQty"] * r["UnitCost"] if r["NeedsReorder"] else 0.0

    stockouts = sum(1 for r in inv if r["QtyOnHand"] == 0)
    total_value = sum(r["ExtendedValue"] for r in inv)
    share_value = sum(r["ExtendedValue"] for r in inv if r["Category"] == SHARE_CATEGORY)
    share = share_value / total_value
    needs_reorder = sum(1 for r in inv if r["NeedsReorder"])
    avg_lead = sum(r["LeadTimeDays"] for r in inv) / n
    all_rows_with_totals = 1 + n + 1           # header + data + Total Row
    slicer_rows = [r for r in inv if r["Facility"] == SLICER_FACILITY and r["Category"] == SLICER_CATEGORY]
    slicer_sum = sum(r["ExtendedValue"] for r in slicer_rows)
    window_end = REPORT_DATE + timedelta(days=EXPIRING_WINDOW_DAYS)
    expiring = sum(1 for r in inv if r["ExpirationDate"] and REPORT_DATE < r["ExpirationDate"] <= window_end)
    expired_value = sum(r["ExtendedValue"] for r in inv if r["ExpirationDate"] and r["ExpirationDate"] < REPORT_DATE)
    expired_rows = sum(1 for r in inv if r["ExpirationDate"] and r["ExpirationDate"] < REPORT_DATE)
    overdue = sum(1 for r in inv if (REPORT_DATE - r["LastCountDate"]).days > CYCLE_COUNT_DAYS)
    # sanity: the "later than ReportDate" vs "on or after" wording can't change task 10's answer
    assert not any(r["ExpirationDate"] == REPORT_DATE for r in inv)

    # bonus
    order_total = sum(r["OrderCost"] for r in inv)
    by_vendor = {v: sum(r["OrderCost"] for r in inv if r["Vendor"] == v) for v in vendor_names}
    top_vendor = max(by_vendor, key=by_vendor.get)
    top_amount = by_vendor[top_vendor]
    assert sorted(by_vendor.values())[-1] > sorted(by_vendor.values())[-2] + 1, "top vendor must be unambiguous"
    discount_base = sum(v for v in by_vendor.values() if v >= DISCOUNT_MIN)
    discount = discount_base * DISCOUNT_PCT
    vendors_discounted = sum(1 for v in by_vendor.values() if v >= DISCOUNT_MIN)

    L.practice_intro = (
        f"Every task uses the Inventory sheet: supply stock in {n_rooms} storerooms across {words.get(n_facs, n_facs)} "
        f"Bluestone hospitals, counted on {REPORT_DATE:%m/%d/%Y}. Do the tasks in order, because later tasks use the Table, columns, Total Row, and names you "
        "create in earlier ones. After task 1, write formulas with structured references such as "
        "tblInventory[QtyOnHand] instead of cell ranges.")
    L.start_notes = [
        "The Inventory, Settings, and Vendors sheets are plain ranges on purpose. You turn them into Tables and "
        "named cells as part of the practice.",
    ]

    L.tasks = [
        Task("On the Inventory sheet, click any cell in the data and convert the range to an Excel Table "
             "(Ctrl + T; Mac: Control + T). Then rename the Table tblInventory in the Table Name box on the Table Design "
             "tab. The gray cell finds a Table with that exact name and counts its data rows. It stays blank until "
             "the Table exists.",
             answer=n, title="Convert the range to a Table named tblInventory (data rows)",
             summary=f'=IFERROR(ROWS(INDIRECT("{TBL}")),"")',
             # The _simulate_learner hook creates the table in self-test mode; this no-op fill (rewriting the A1
             # header with its own text) just tells the verifier to run the self-test for this summary task.
             fill={"range": "Inventory!A1:A1", "values": ["StockID"]},
             live=f"=COUNTA({a1('StockID')})",
             solution="1. Click any cell inside the data, for example **A2**.\n"
                      "2. Press **Ctrl + T** (Mac: **Control + T**), or choose **Insert → Table**.\n"
                      f"3. Check that the range shows **$A$1:$S${last}** and that **My table has headers** is ticked, "
                      "then click **OK**.\n"
                      "4. On the **Table Design** tab (Mac: **Table** tab), click in the **Table Name** box at the far "
                      "left, type `tblInventory`, and press **Enter**.",
             hint="Make sure 'My table has headers' is ticked",
             explanation="Ctrl + T selects the whole block of connected cells around the active cell, so you don't need "
                         "to select the data first. Excel names new Tables Table1, Table2, and so on. Renaming the Table "
                         "is what makes formulas readable: `tblInventory[UnitCost]` tells the next reader what the formula "
                         "uses, while `Table1[UnitCost]` doesn't. The gray cell uses `INDIRECT(\"tblInventory\")`, "
                         "which looks the Table up by name and returns its data rows (headers and Total Row excluded)."),
        Task("How many stock rows are completely out of stock (QtyOnHand = 0)? Use a structured reference to the "
             "QtyOnHand column instead of a cell range.",
             answer=stockouts, solution=f"=COUNTIF({TBL}[QtyOnHand],0)",
             live=f"=COUNTIF({qty},0)",
             hint="COUNTIF. Type tblInventory[ and pick the column from the list",
             explanation="`tblInventory[QtyOnHand]` means \"every data cell in the QtyOnHand column.\" It excludes the "
                         "header and the Total Row, and it grows when rows are added. After you type `tblInventory[`, "
                         "Excel lists the columns. Pick one with the arrow keys and press **Tab**. These rows are "
                         "*stockouts*: a nurse who reaches for the item finds an empty bin."),
        Task(f"Fill the yellow ExtendedValue column (column R) with a calculated column: each row's QtyOnHand × "
             f"UnitCost. Type the formula once in R{first} and press Enter, and the Table fills every row. The gray "
             f"cell totals the column: what's the value of all stock on hand across the system?",
             answer=round(total_value, 2), fmt="#,##0.00", title="ExtendedValue calculated column (total stock value)",
             solution="=[@QtyOnHand]*[@UnitCost]",
             summary=f'=IF(COUNT({ev_rng})=0,"",SUM({ev_rng}))',
             fill={"range": f"Inventory!R{first}:R{last}", "formula": "=[@QtyOnHand]*[@UnitCost]"},
             live=f"=SUMPRODUCT({qty},{cost})", table=TBL,
             hint="Click the QtyOnHand cell in the same row while typing; Excel writes [@QtyOnHand]",
             explanation=f"When you enter a formula in an empty Table column, Excel copies it to every row of that column. "
                         f"That's a *calculated column*. `[@QtyOnHand]` means \"QtyOnHand in this row,\" so the formula "
                         f"reads the same on all {n} rows and says what it multiplies. The A1 version, "
                         f"`={sd.col('QtyOnHand')}{first}*{sd.col('UnitCost')}{first}`, changes on every row and tells the "
                         f"next reader nothing about the columns it uses."),
        Task(f"What share of the system's total inventory value (ExtendedValue) is {SHARE_CATEGORY} stock? Divide "
             f"the category's value by the total of all rows and enter it as a percentage, to 1 decimal place.",
             answer=share, fmt="0.0%",
             solution=f'=SUMIFS({TBL}[ExtendedValue],{TBL}[Category],"{SHARE_CATEGORY}")/SUM({TBL}[ExtendedValue])',
             live=f'=SUMPRODUCT(({a1("Category")}="{SHARE_CATEGORY}")*{qty}*{cost})/SUMPRODUCT({qty},{cost})',
             hint="SUMIFS(…)/SUM(…), both with structured references",
             explanation=f"Every argument in SUMIFS is a whole Table column, so the ranges are automatically the same "
                         f"size, which SUMIFS requires. Only {sum(1 for r in inv if r['Category'] == SHARE_CATEGORY)} of "
                         f"the {n} rows are implants, yet they hold {share:.0%} of the value on the shelves. "
                         f"High-cost, low-volume items like these are the ones a supply chain team counts most carefully."),
        Task(f"Fill the yellow NeedsReorder column (column S) with a calculated column that returns TRUE when "
             f"QtyOnHand is at or below ReorderPoint, and FALSE otherwise. The gray cell counts the TRUE rows: how many "
             f"stock rows need reordering?",
             answer=needs_reorder, title="NeedsReorder calculated column (rows to reorder)",
             solution="=[@QtyOnHand]<=[@ReorderPoint]",
             summary=f'=IF(COUNTA({nr_rng})=0,"",COUNTIF({nr_rng},TRUE))',
             fill={"range": f"Inventory!S{first}:S{last}", "formula": "=[@QtyOnHand]<=[@ReorderPoint]"},
             live=f"=SUMPRODUCT(--({qty}<={rp}))", table=TBL,
             hint="A comparison already returns TRUE or FALSE, so you don't need IF",
             explanation="A comparison such as `[@QtyOnHand]<=[@ReorderPoint]` evaluates to TRUE or FALSE by itself "
                         "(Lesson 2.1). Storing the result as a column turns a row-by-row test into something you can "
                         "filter, count with `COUNTIF(tblInventory[NeedsReorder],TRUE)`, and reuse in other formulas, "
                         "as the bonus does. Use \"at or below\" (`<=`): an item sitting exactly at its reorder point "
                         "is due."),
        Task("Turn on the Table's Total Row and set the LeadTimeDays total to Average. Then, in the yellow cell, type = "
             "and click that Total Row cell, so Excel writes a [#Totals] reference. What's the average vendor lead "
             "time in days, to 2 decimal places?",
             answer=avg_lead, fmt="0.00", solution=f"={TBL}[[#Totals],[LeadTimeDays]]",
             live=f"=AVERAGE({a1('LeadTimeDays')})",
             hint="Table Design → Total Row (Windows: Ctrl + Shift + T), then use the dropdown in the Total Row cell",
             explanation="The Total Row's dropdown writes `=SUBTOTAL(101,[LeadTimeDays])`. Function number 101 means "
                         "AVERAGE of the visible rows only, ignoring rows hidden by a filter or hidden by hand "
                         "(109 = SUM, 103 = COUNTA, 104 = MAX). "
                         "`tblInventory[[#Totals],[LeadTimeDays]]` points at that total cell by name, so it still works "
                         "if rows are added. If you turn the Total Row off, the reference returns #REF!."),
        Task("With the Total Row still on, what does =ROWS(tblInventory[#All]) return? Predict it first, then type the "
             "formula to check.",
             answer=all_rows_with_totals, solution=f"=ROWS({TBL}[#All])", live=False,
             hint="#All is everything: which rows does it include that tblInventory[QtyOnHand] doesn't?",
             explanation=f"`[#All]` covers the header row, the {n} data rows, and the Total Row: 1 + {n} + 1 = "
                         f"{all_rows_with_totals}. `tblInventory[#Data]` (or just `tblInventory`) is the {n} "
                         f"data rows, `[#Headers]` is the header row alone, and `[#Totals]` is the Total Row alone."),
        Task(f"Insert slicers for Facility and Category (Table Design → Insert Slicer). Select {SLICER_FACILITY} in "
             f"the Facility slicer and {SLICER_CATEGORY} in the Category slicer. Set the Total Row's ExtendedValue cell "
             f"to Sum. What value does it show? Type the number, then clear both slicers, because while they filter the "
             f"Table the Total Row (and task 6) only summarizes the visible rows.",
             answer=round(slicer_sum, 2), fmt="#,##0.00", title=f"Slicers: {SLICER_CATEGORY} stock at {SLICER_FACILITY}",
             solution="1. Click inside the Table, then **Table Design → Insert Slicer** (Mac: **Table → Insert "
                      "Slicer**). Tick **Facility** and **Category** and click **OK**.\n"
                      f"2. In the Facility slicer click **{SLICER_FACILITY}**. In the Category slicer click "
                      f"**{SLICER_CATEGORY}**. The Table now shows {len(slicer_rows)} rows.\n"
                      "3. In the Total Row, click the ExtendedValue cell, open its dropdown, and choose **Sum**.\n"
                      f"4. Read the total and type it in the answer cell: **{slicer_sum:,.2f}**.\n"
                      "5. Click the **Clear Filter** button (funnel with a red X) at the top right of each slicer.",
             live=f'=SUMPRODUCT(({a1("Facility")}="{SLICER_FACILITY}")*({a1("Category")}="{SLICER_CATEGORY}")*{qty}*{cost})',
             hint="Slicers filter the Table, and SUBTOTAL ignores rows the filter hides",
             explanation=f"Slicers are buttons that apply the Table's AutoFilter for you. Choices in different slicers "
                         f"combine with AND, so only rows that match both stay visible. The Total Row uses SUBTOTAL(109,…), "
                         f"which adds only the visible rows. That's why it changes as you click. You can check the "
                         f"number without slicers using "
                         f"`=SUMIFS(tblInventory[ExtendedValue],tblInventory[Facility],\"{SLICER_FACILITY}\","
                         f"tblInventory[Category],\"{SLICER_CATEGORY}\")`."),
        Task("On the Settings sheet, select A2:B3 and use Formulas → Create from Selection (tick only Left column) to "
             "name cell B2 ReportDate and cell B3 ExpiringWindowDays. The gray cell looks up both names and adds them: "
             "what date does it show?",
             answer=window_end, fmt="mm/dd/yyyy", title="Name the Settings cells with Create from Selection",
             summary='=IFERROR(INDIRECT("ReportDate")+INDIRECT("ExpiringWindowDays"),"")',
             # names are created by the _simulate_learner hook in self-test mode; no-op fill opts in to the self-test
             fill={"range": "Settings!A2:A2", "values": ["ReportDate"]},
             live=f"={report_cell}+{window_cell}",
             solution="1. Go to the **Settings** sheet and select **A2:B3** (labels and values together).\n"
                      "2. Choose **Formulas → Create from Selection** (Windows shortcut: **Ctrl + Shift + F3**).\n"
                      "3. Tick **Left column** only, untick **Top row**, and click **OK**.\n"
                      "4. Check the result: click B2 and the Name Box shows **ReportDate**. Open **Formulas → Name "
                      "Manager** to see both names and what they refer to.",
             hint="The labels in column A become the names of the cells beside them",
             explanation=f"Create from Selection turns each label in the left column into a name for the cell to its "
                         f"right: ReportDate → `=Settings!$B$2` and ExpiringWindowDays → `=Settings!$B$3`. Names are "
                         f"absolute references, so they never shift when you copy a formula. "
                         f"{REPORT_DATE:%m/%d/%Y} + {EXPIRING_WINDOW_DAYS} days = {window_end:%m/%d/%Y}, because dates "
                         f"are day counts (Lesson 2.3). If the gray cell stays blank, check the spelling of each name in "
                         f"Name Manager."),
        Task("How many stock rows expire inside the expiry window? Count rows whose ExpirationDate is later than "
             "ReportDate and no later than ReportDate + ExpiringWindowDays. Use the two names in a COUNTIFS (rows with "
             "no ExpirationDate don't count).",
             answer=expiring,
             solution=f'=COUNTIFS({TBL}[ExpirationDate],">"&ReportDate,{TBL}[ExpirationDate],"<="&ReportDate+ExpiringWindowDays)',
             live=f'=COUNTIFS({a1("ExpirationDate")},">"&{report_cell},{a1("ExpirationDate")},"<="&{report_cell}+{window_cell})',
             hint="Join an operator to a name the way you join it to a cell: \">\"&ReportDate",
             explanation="Criteria are text, so join the operator to the name with `&`, exactly as you would with a "
                         "cell reference (Lesson 2.5). Blank ExpirationDate cells never match a date criterion, so "
                         "non-perishables drop out automatically. The formula reads like the policy, and when the "
                         "next snapshot arrives you only change the date in Settings!B2."),
        Task("Expired stock must be pulled from the shelves and written off. What's the total ExtendedValue of rows whose "
             "ExpirationDate is before ReportDate?",
             answer=round(expired_value, 2), fmt="#,##0.00",
             solution=f'=SUMIFS({TBL}[ExtendedValue],{TBL}[ExpirationDate],"<"&ReportDate)',
             live=f'=SUMPRODUCT(({a1("ExpirationDate")}<{report_cell})*({a1("ExpirationDate")}<>"")*{qty}*{cost})',
             hint="SUMIFS with a \"<\"&ReportDate criterion",
             explanation=f"In this snapshot, {expired_rows} rows are past their expiration date. That's a patient-safety issue as well as "
                         f"a cost: surveyors such as The Joint Commission look for expired supplies in clinical areas. "
                         f"The live formula in the key guards against blanks with `<>\"\"`, but SUMIFS doesn't need "
                         f"that guard, because an empty cell never matches \"<\"&ReportDate."),
        Task(f"Policy says every stock row must be cycle-counted at least every {CYCLE_COUNT_DAYS} days. Create a named "
             f"constant CycleCountDays that refers to ={CYCLE_COUNT_DAYS} (Formulas → Define Name; it doesn't live in "
             f"any cell). How many rows are overdue, meaning ReportDate − LastCountDate is greater than CycleCountDays?",
             answer=overdue,
             solution=f'=COUNTIF({TBL}[LastCountDate],"<"&ReportDate-CycleCountDays)',
             live=f'=COUNTIF({a1("LastCountDate")},"<"&{report_cell}-{CYCLE_COUNT_DAYS})',
             hint="Rearrange: overdue means LastCountDate < ReportDate − CycleCountDays",
             explanation=f"In the New Name dialog, type `CycleCountDays` as the Name and `={CYCLE_COUNT_DAYS}` in "
                         f"**Refers to**. ReportDate − CycleCountDays is {REPORT_DATE - timedelta(days=CYCLE_COUNT_DAYS):%m/%d/%Y}, "
                         f"so any row last counted before that date is more than {CYCLE_COUNT_DAYS} days old. A named "
                         f"constant suits a policy value that shouldn't be typed over by accident. A named cell, like "
                         f"ReportDate, suits an input that people change often."),
    ]

    # ------------------------------------------------------------------ bonus
    L.bonus_title = "Bonus: This week's reorder report"
    L.bonus_scenario = (
        "Every Monday, Bluestone's purchasing team turns the inventory snapshot into purchase orders. Each stock row at "
        "or below its reorder point gets an order for its standard ReorderQty, and all the lines for one vendor go on a "
        "single purchase order. Build the report with Tables and names, so next week it updates itself when the new "
        "snapshot is pasted in. Finish practice tasks 1–12 first, because this uses tblInventory, NeedsReorder, and "
        "the names.")
    L.bonus = [
        Task(f"Add a new column to tblInventory by typing OrderCost in cell {oc_letter}1, just right of the NeedsReorder "
             f"header. The Table expands to include it. Make it a calculated column that returns ReorderQty × UnitCost "
             f"for rows where NeedsReorder is TRUE, and 0 for every other row. The gray cell totals the column: what's "
             f"the total cost of this week's orders?",
             answer=round(order_total, 2), fmt="#,##0.00", title="OrderCost calculated column (total of all orders)",
             solution="=IF([@NeedsReorder],[@ReorderQty]*[@UnitCost],0)",
             summary=f'=IF(COUNT({oc_rng})=0,"",SUM({oc_rng}))',
             fill={"range": f"Inventory!{oc_letter}{first}:{oc_letter}{last}",
                   "formula": "=IF([@NeedsReorder],[@ReorderQty]*[@UnitCost],0)"},
             live=f"=SUMPRODUCT(({qty}<={rp})*{a1('ReorderQty')}*{cost})", table=TBL,
             hint="IF can test a TRUE/FALSE column directly: IF([@NeedsReorder], …, 0)",
             explanation=f"Typing in the first empty column next to a Table adds a column to it (Excel's "
                         f"*AutoExpansion*). `IF([@NeedsReorder],…)` reuses the calculated column you built in task 5 "
                         f"instead of repeating the comparison. The {needs_reorder} rows that need an order return their "
                         f"cost and the rest return 0, so the column total is the cost of the whole report."),
        Task(f"On the Vendors sheet, convert the vendor list to a Table named {VTBL}. Fill its yellow ReorderCost column "
             f"with a calculated column that adds up tblInventory[OrderCost] for that row's vendor. Which vendor gets the "
             f"largest purchase order? (Type the vendor's name, or return it with a formula.)",
             answer=top_vendor,
             solution=f"=INDEX({VTBL}[Vendor],MATCH(MAX({VTBL}[ReorderCost]),{VTBL}[ReorderCost],0))",
             live=False, hint="In tblVendors, SUMIFS can add up another Table's column, with [@Vendor] as the criterion",
             explanation=f"The calculated column in {VTBL} is "
                         f"`=SUMIFS(tblInventory[OrderCost],tblInventory[Vendor],[@Vendor])`. It mixes whole columns "
                         f"from *another* Table with `[@Vendor]` from this row, which is how Tables talk to each other. "
                         f"INDEX/MATCH (Lesson 2.6) then returns the vendor on the row with the largest total. "
                         f"`=XLOOKUP(MAX({VTBL}[ReorderCost]),{VTBL}[ReorderCost],{VTBL}[Vendor])` works too in Excel "
                         f"2021 or Microsoft 365. Orthopedic implants cost thousands of dollars each, so even a few "
                         f"reorder lines make the largest order."),
        Task("What's the value of that vendor's purchase order?",
             answer=round(top_amount, 2), fmt="#,##0.00", solution=f"=MAX({VTBL}[ReorderCost])",
             live=f'=SUMPRODUCT(({a1("Vendor")}="{top_vendor}")*({qty}<={rp})*{a1("ReorderQty")}*{cost})',
             hint="The largest value in tblVendors[ReorderCost]",
             explanation=f"That's {top_amount / order_total:.0%} of the week's spend on a single purchase order. In "
                         f"practice, a buyer would confirm implant orders with the OR schedule before sending them."),
        Task(f"The group purchasing contract gives a {DISCOUNT_PCT:.0%} discount on any single vendor order of "
             f"${DISCOUNT_MIN:,} or more. Define two named constants, DiscountMin (={DISCOUNT_MIN}) and DiscountPct "
             f"(={DISCOUNT_PCT}), then calculate the total discount on this week's orders.",
             answer=round(discount, 2), fmt="#,##0.00",
             solution=f'=SUMIFS({VTBL}[ReorderCost],{VTBL}[ReorderCost],">="&DiscountMin)*DiscountPct',
             live=False,
             hint="SUMIFS can use the same column as the sum range and the criteria range",
             explanation=f"Orders from {vendors_discounted} vendors reach ${DISCOUNT_MIN:,}. Together they total "
                         f"${discount_base:,.2f}, and {DISCOUNT_PCT:.0%} of that is the discount. With the thresholds in "
                         f"named constants, Purchasing can test a new contract (say 3% at $15,000) by editing two names in "
                         f"Name Manager, and every formula that uses them updates."),
    ]

    # ------------------------------------------------------------------ custom content
    @L.customize
    def _sheet_polish(wb, lesson, selftest):
        ws = wb["Settings"]
        for row in ws.iter_rows(min_row=st.first_row, max_row=st.last_row):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=c.column == 3)
        ws[f"B{st.first_row + 1}"].number_format = "0"
        note_row = st.last_row + 2
        ws.cell(row=note_row, column=1,
                value="Task 9 names B2 and B3 from the labels in column A. Later tasks and the bonus add named "
                      "constants (CycleCountDays, DiscountMin, DiscountPct) in Name Manager; they don't live in cells.")
        ws.cell(row=note_row, column=1).font = Font(italic=True, color="595959")
        vws = wb["Vendors"]
        vnote = vd.last_row + 2
        vws.cell(row=vnote, column=1, value="Bonus: convert this list to a Table named tblVendors, then fill ReorderCost "
                                           "with a calculated column.")
        vws.cell(row=vnote, column=1).font = Font(italic=True, color="595959")

    @L.customize
    def _simulate_learner(wb, lesson, selftest):
        """Self-test only: build the Tables, Total Row, OrderCost column, and names the learner creates."""
        if not selftest:
            return
        ws = wb["Inventory"]
        ws[f"{oc_letter}1"] = "OrderCost"
        headers = sd.headers + ["OrderCost"]
        total_row = last + 1
        ws.cell(row=total_row, column=1, value="Total")
        ws[f"{sd.col('LeadTimeDays')}{total_row}"] = f"=SUBTOTAL(101,{TBL}[LeadTimeDays])"
        ws[f"{sd.col('ExtendedValue')}{total_row}"] = f"=SUBTOTAL(109,{TBL}[ExtendedValue])"
        tcols = []
        for j, h in enumerate(headers, 1):
            tc = TableColumn(id=j, name=h)
            if j == 1:
                tc.totalsRowLabel = "Total"
            elif h == "LeadTimeDays":
                tc.totalsRowFunction = "average"
            elif h == "ExtendedValue":
                tc.totalsRowFunction = "sum"
            tcols.append(tc)
        t = Table(displayName=TBL, ref=f"A1:{oc_letter}{total_row}", totalsRowCount=1)
        t.tableColumns = tcols
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)

        vws = wb["Vendors"]
        for i in range(vd.first_row, vd.last_row + 1):
            vws[f"{vd.col('ReorderCost')}{i}"] = to_file_formula(
                "=SUMIFS(tblInventory[OrderCost],tblInventory[Vendor],[@Vendor])", VTBL)
        vt = Table(displayName=VTBL, ref=f"A1:{vd.col('ReorderCost')}{vd.last_row}")
        vt.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        vws.add_table(vt)

        names = {
            "ReportDate": f"Settings!${st.col('Value')}${st.first_row}",
            "ExpiringWindowDays": f"Settings!${st.col('Value')}${st.first_row + 1}",
            "CycleCountDays": str(CYCLE_COUNT_DAYS),
            "DiscountMin": str(DISCOUNT_MIN),
            "DiscountPct": str(DISCOUNT_PCT),
        }
        for name, ref in names.items():
            wb.defined_names[name] = DefinedName(name, attr_text=ref)

    return L
