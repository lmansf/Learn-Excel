"""Lesson 5.4 · VBA: Custom Functions, Dictionaries & Error Handling.

Data (Cedar Ridge Medical Center, F03, 2025):
  Patients    adults (18+ on 12/31/2025) with any 2025 Cedar Ridge encounter: DOB, height, weight
  Encounters  every 2025 inpatient and observation stay at Cedar Ridge, with the patient's DOB joined in.
              Four stays get deliberate discharge typos (three with the year keyed as 2024, one AM/PM slip)
              so LOSDAYS has something to reject.
  Claims      every claim with a 2025 service date for a Cedar Ridge encounter. Five rows get a blank or
              text BilledAmount ("N/A", "VOID") so the bonus macro has rows to skip.
  Payers      the 8 payers.

Files written next to the workbook (ASCII with CRLF line endings, because the VBE imports Windows-style text):
  starter/HealthUDFs.bas                    TODO stubs for the UDF tasks 1-6 (helpers finished)
  starter/ClaimDictionaries.bas             TODO stubs for the Collection/Dictionary tasks 8-10
  starter/Snippets.bas                      the four predict-the-output snippets (tasks 7, 11, 12, 13)
  starter/PayerSummary.bas                  bonus skeleton (helpers finished)
  solutions/HealthUDFs_Solution.bas         spoilers
  solutions/ClaimDictionaries_Solution.bas  spoilers (Windows, Scripting.Dictionary)
  solutions/ClaimCollectionsMac_Solution.bas  spoilers (Mac-friendly: Collections only)
  solutions/PayerSummary_Solution.bas       bonus spoilers (Windows, Scripting.Dictionary)
  solutions/PayerSummaryMac_Solution.bas    bonus spoilers (Mac-friendly: Collections only)

Smoke test (LibreOffice VBA-compatibility mode). LibreOffice can't create Scripting.Dictionary, so the smoke
test runs the UDFs and the Collection-only (Mac) versions of the macros, which share the Dictionary versions' logic:
  cd tools && python3 -c "import lessons.lesson_5_4 as m; m.smoke_test()"
"""
from __future__ import annotations

from collections import Counter, OrderedDict
from datetime import date, timedelta
from pathlib import Path

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from xlcourse import Lesson, Task, data

CODE = "5.4"
FAC = "F03"                    # Cedar Ridge Medical Center
AS_OF = date(2025, 12, 31)
SUMMARY_SHEET = "PayerSummary"
OUT = {"reasons": "B4", "patients": "B5", "top_name": "B6", "top_count": "B7"}

NAVY = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
MACRO_FILL = PatternFill("solid", fgColor="DDEBF7")   # blue, like the Output sheets of 5.2 and 5.3
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------------------
# VBA source (ASCII only). One source of truth for the .bas files, the Snippets
# sheet, the README answer key, and the LibreOffice smoke test.
# ---------------------------------------------------------------------------
UDF_HELPERS = '''' ---------------------------------------------------------------------
' Helpers. Private = code in other modules can't call them, and they
' stay out of AutoComplete and the Insert Function list.
' ---------------------------------------------------------------------

' A cell reference arrives as a Range object; a typed number or a formula
' result arrives as a plain value. Return the value either way.
Private Function ArgValue(arg As Variant) As Variant
    If IsObject(arg) Then
        ArgValue = arg.Cells(1, 1).Value
    Else
        ArgValue = arg
    End If
End Function

' True for real numbers and dates (dates are numbers in Excel).
' False for blanks, text, and error values. IsNumeric is too forgiving:
' IsNumeric(Empty) and IsNumeric("12") are both True.
Private Function IsNumberValue(v As Variant) As Boolean
    Select Case VarType(v)
        Case vbInteger, vbLong, vbSingle, vbDouble, vbCurrency, vbDecimal, vbDate, vbByte
            IsNumberValue = True
        Case Else
            IsNumberValue = False
    End Select
End Function'''

SOL_BMI = '''' Tasks 1-2: body-mass index = 703 x weight (lb) / height (in) ^ 2
Public Function BMI(weightLb As Variant, heightIn As Variant) As Variant
    Dim w As Variant, h As Variant
    w = ArgValue(weightLb)
    h = ArgValue(heightIn)
    If Not IsNumberValue(w) Or Not IsNumberValue(h) Then
        BMI = CVErr(xlErrValue)            ' blank or text input: #VALUE!
    ElseIf w <= 0 Or h <= 0 Then
        BMI = CVErr(xlErrNum)              ' impossible measurement: #NUM!
    Else
        BMI = 703 * w / h ^ 2
    End If
End Function'''

SOL_AGEAT = '''' Task 3: age in completed years on asOf. Leave asOf out to use today's date.
Public Function AGEAT(dob As Variant, Optional asOf As Variant) As Variant
    Dim birthValue As Variant, refValue As Variant
    Dim birth As Date, refDate As Date, years As Long

    birthValue = ArgValue(dob)
    If IsMissing(asOf) Then
        refValue = Date                    ' second argument left out: today
        Application.Volatile               ' today changes, so recalculate every time
    Else
        refValue = ArgValue(asOf)
    End If
    If Not IsNumberValue(birthValue) Or Not IsNumberValue(refValue) Then
        AGEAT = CVErr(xlErrValue)
        Exit Function
    End If

    birth = CDate(birthValue)
    refDate = CDate(refValue)
    If refDate < birth Then
        AGEAT = CVErr(xlErrNum)            ' asOf is before the date of birth
        Exit Function
    End If

    years = Year(refDate) - Year(birth)
    ' Birthday still to come in the asOf year? Then the person is a year younger.
    If DateSerial(Year(refDate), Month(birth), Day(birth)) > refDate Then years = years - 1
    AGEAT = years
End Function'''

SOL_LOSDAYS = '''' Tasks 4-5: length of stay = midnights between admission and discharge.
' #VALUE! when an input isn't a date or the discharge is before the admission.
Public Function LOSDAYS(admitDateTime As Variant, dischargeDateTime As Variant) As Variant
    Dim a As Variant, d As Variant
    a = ArgValue(admitDateTime)
    d = ArgValue(dischargeDateTime)
    If Not IsNumberValue(a) Or Not IsNumberValue(d) Then
        LOSDAYS = CVErr(xlErrValue)
    ElseIf CDate(d) < CDate(a) Then
        LOSDAYS = CVErr(xlErrValue)        ' discharge before admission: a data-entry error
    Else
        LOSDAYS = DateDiff("d", CDate(a), CDate(d))   ' "d" counts the midnights crossed
    End If
End Function'''

SOL_DENIALRATE = '''' Task 6: share of non-blank cells in statusRange that equal statusToCount.
Public Function DENIALRATE(statusRange As Range, Optional statusToCount As String = "Denied") As Variant
    Dim cell As Range, total As Long, hits As Long
    For Each cell In statusRange.Cells
        If Not IsEmpty(cell.Value) And Not IsError(cell.Value) Then
            total = total + 1
            If StrComp(Trim$(CStr(cell.Value)), statusToCount, vbTextCompare) = 0 Then hits = hits + 1
        End If
    Next cell
    If total = 0 Then
        DENIALRATE = CVErr(xlErrDiv0)      ' nothing to divide by: #DIV/0!
    Else
        DENIALRATE = hits / total
    End If
End Function'''


def _udf_solution_module() -> str:
    return f'''Attribute VB_Name = "HealthUDFsSolution"
Option Explicit

' =====================================================================
' Lesson 5.4 - reference solutions for the UDF tasks 1-6 (SPOILERS)
'
' Import this into a spare copy of the workbook, or remove your own
' HealthUDFs module first (right-click it > Remove HealthUDFs), so two
' modules don't define the same function names.
'
' Worksheet functions in this module:
'   =BMI(weightLb, heightIn)                    body-mass index
'   =AGEAT(dob, [asOf])                         age in completed years
'   =LOSDAYS(admitDateTime, dischargeDateTime)  length of stay (midnights)
'   =DENIALRATE(statusRange, [statusToCount])   share of claims with a status
' UDFs must live in a standard module (Insert > Module), not in a sheet
' module or ThisWorkbook; otherwise the cell shows #NAME?.
' =====================================================================

{SOL_BMI}

{SOL_AGEAT}

{SOL_LOSDAYS}

{SOL_DENIALRATE}

{UDF_HELPERS}
'''


def _udf_starter_module() -> str:
    return f'''Attribute VB_Name = "HealthUDFs"
Option Explicit

' =====================================================================
' Lesson 5.4 - starter module for the UDF tasks 1-6
'
' 1. Save the workbook as .xlsm.  2. In the VBE: File > Import File...
' 3. Replace each TODO. 4. Use the functions in cells, e.g. =BMI(G2,F2).
' Public functions in a standard module like this one can be called from
' worksheet cells. The two helpers at the bottom are finished for you.
' =====================================================================

' Tasks 1-2: body-mass index = 703 x weight (lb) / height (in) ^ 2
Public Function BMI(weightLb As Variant, heightIn As Variant) As Variant
    Dim w As Variant, h As Variant
    w = ArgValue(weightLb)
    h = ArgValue(heightIn)
    ' TODO 1: if w or h is not a number (use IsNumberValue), return a #VALUE! error
    ' TODO 2: if w or h is zero or negative, return a #NUM! error
    ' TODO 3: otherwise return the BMI
End Function

' Task 3: age in completed years on asOf. Leave asOf out to use today's date.
Public Function AGEAT(dob As Variant, Optional asOf As Variant) As Variant
    ' TODO 1: read dob with ArgValue. If asOf is missing use Date, otherwise ArgValue(asOf)
    ' TODO 2: return #VALUE! if either value is not a number or date
    ' TODO 3: work out the age in completed years (careful: a birthday that is still
    '         to come in the asOf year doesn't count yet)
End Function

' Tasks 4-5: length of stay = midnights between admission and discharge
Public Function LOSDAYS(admitDateTime As Variant, dischargeDateTime As Variant) As Variant
    ' TODO 1: read both arguments with ArgValue; #VALUE! if either is not a date
    ' TODO 2: #VALUE! if the discharge date-time is earlier than the admission
    ' TODO 3: otherwise return the number of midnights between the two (DateDiff)
End Function

' Task 6: share of non-blank cells in statusRange that equal statusToCount
Public Function DENIALRATE(statusRange As Range, Optional statusToCount As String = "Denied") As Variant
    ' TODO 1: loop over statusRange.Cells: count the non-blank cells (total) and the
    '         cells whose text equals statusToCount, ignoring case (hits)
    ' TODO 2: return #DIV/0! if total = 0, otherwise hits / total
End Function

{UDF_HELPERS}
'''


ADD_IF_NEW = '''' Adds keyText to a Collection once. Returns False if it was already there.
Private Function AddIfNew(col As Collection, ByVal keyText As String) As Boolean
    On Error Resume Next
    col.Add Item:=keyText, Key:=keyText    ' error 457 if the key already exists
    AddIfNew = (Err.Number = 0)
    On Error GoTo 0
End Function'''

HEADER_COLUMN = '''' Column number of a header in row 1. Raises a clear error if it is missing.
Private Function HeaderColumn(ws As Worksheet, ByVal headerText As String) As Long
    Dim pos As Variant
    pos = Application.Match(headerText, ws.Rows(1), 0)   ' an error VALUE, not a crash
    If IsError(pos) Then
        Err.Raise vbObjectError + 513, "HeaderColumn", _
                  "Column '" & headerText & "' was not found on sheet '" & ws.Name & "'."
    End If
    HeaderColumn = CLng(pos)
End Function'''

DICT_HELPERS = f'''' ---------------------------------------------------------------------
' Helpers
' ---------------------------------------------------------------------

{ADD_IF_NEW}

{HEADER_COLUMN}'''

SOL_REASONS = '''' Task 8: how many different denial reasons? A Collection refuses a duplicate
' key, so adding every reason with Key:=reason keeps exactly one of each.
Public Sub CountDenialReasons()
    Dim ws As Worksheet, claims As Variant
    Dim reasonCol As Long, r As Long, reason As String
    Dim reasons As Collection

    Set ws = ThisWorkbook.Worksheets("Claims")
    claims = ws.Range("A1").CurrentRegion.Value        ' header + data in one 2-D array
    reasonCol = HeaderColumn(ws, "DenialReason")
    Set reasons = New Collection

    For r = 2 To UBound(claims, 1)
        reason = Trim$(CStr(claims(r, reasonCol)))     ' a blank cell becomes ""
        If Len(reason) > 0 Then AddIfNew reasons, reason
    Next r

    ThisWorkbook.Worksheets("Output").Range("B4").Value = reasons.Count
End Sub'''

SOL_PATIENTS = '''' Task 9: how many different patients have a claim? Dictionary keys are unique.
Public Sub CountDistinctPatients()
    Dim ws As Worksheet, claims As Variant
    Dim patientCol As Long, r As Long
    Dim patients As Object                             ' late-bound Scripting.Dictionary

    Set ws = ThisWorkbook.Worksheets("Claims")
    claims = ws.Range("A1").CurrentRegion.Value
    patientCol = HeaderColumn(ws, "PatientID")
    Set patients = CreateObject("Scripting.Dictionary")

    For r = 2 To UBound(claims, 1)
        If Not patients.Exists(claims(r, patientCol)) Then
            patients.Add claims(r, patientCol), 1
        End If
    Next r

    ThisWorkbook.Worksheets("Output").Range("B5").Value = patients.Count
End Sub'''

SOL_TOPPAYER = '''' Task 10: which payer has the most Denied claims? Two dictionaries:
' one to look up payer names, one to count.
Public Sub TopDeniedPayer()
    Dim wsClaims As Worksheet, claims As Variant, payers As Variant
    Dim payerCol As Long, statusCol As Long, r As Long
    Dim payerNames As Object, deniedCount As Object
    Dim payerID As Variant, bestID As String, bestCount As Long

    ' 1. Lookup dictionary: PayerID -> PayerName
    Set payerNames = CreateObject("Scripting.Dictionary")
    payers = ThisWorkbook.Worksheets("Payers").Range("A1").CurrentRegion.Value
    For r = 2 To UBound(payers, 1)
        payerNames(payers(r, 1)) = payers(r, 2)        ' assigning to a new key adds it
    Next r

    ' 2. Counting dictionary: PayerID -> number of Denied claims
    Set wsClaims = ThisWorkbook.Worksheets("Claims")
    claims = wsClaims.Range("A1").CurrentRegion.Value
    payerCol = HeaderColumn(wsClaims, "PayerID")
    statusCol = HeaderColumn(wsClaims, "ClaimStatus")
    Set deniedCount = CreateObject("Scripting.Dictionary")
    For r = 2 To UBound(claims, 1)
        If claims(r, statusCol) = "Denied" Then
            ' Reading a key that isn't there returns Empty (and adds it); Empty + 1 = 1
            deniedCount(claims(r, payerCol)) = deniedCount(claims(r, payerCol)) + 1
        End If
    Next r

    ' 3. Keep the largest count (on a tie, the payer seen first wins)
    For Each payerID In deniedCount.Keys
        If deniedCount(payerID) > bestCount Then
            bestCount = deniedCount(payerID)
            bestID = payerID
        End If
    Next payerID

    With ThisWorkbook.Worksheets("Output")
        .Range("B6").Value = payerNames(bestID)
        .Range("B7").Value = bestCount
    End With
End Sub'''


def _dict_solution_module() -> str:
    return f'''Attribute VB_Name = "ClaimDictionariesSolution"
Option Explicit

' =====================================================================
' Lesson 5.4 - reference solutions for tasks 8-10 (SPOILERS)
'
' Windows: uses Scripting.Dictionary through late binding, so you don't
' need a Tools > References setting. Excel for Mac has no
' Scripting.Dictionary: use ClaimCollectionsMac_Solution.bas instead.
' Remove your own ClaimDictionaries module first (or import into a spare
' copy) so two modules don't define the same macro names.
'
' Claims sheet: A ClaimID, B PatientID, C PayerID, D ServiceDate,
'   E BilledAmount, F AllowedAmount, G PaidAmount, H ClaimStatus,
'   I DenialReason.    Payers sheet: A PayerID, B PayerName, C PayerType.
' Results go to the Output sheet: B4 (task 8), B5 (task 9), B6:B7 (task 10).
' =====================================================================

{SOL_REASONS}

{SOL_PATIENTS}

{SOL_TOPPAYER}

{DICT_HELPERS}
'''


def _dict_starter_module() -> str:
    helpers = DICT_HELPERS.replace('''    On Error Resume Next
    col.Add Item:=keyText, Key:=keyText    ' error 457 if the key already exists
    AddIfNew = (Err.Number = 0)
    On Error GoTo 0''', '''    ' TODO: trap the error a duplicate key raises (On Error Resume Next), add the
    '       key, set AddIfNew from Err.Number, then switch trapping off again''')
    assert "TODO" in helpers
    return f'''Attribute VB_Name = "ClaimDictionaries"
Option Explicit

' =====================================================================
' Lesson 5.4 - starter module for tasks 8-10 (Collections & Dictionaries)
'
' Each macro writes its answer to the Output sheet; the gray cells on the
' Practice sheet read it from there.
'   Windows: Set d = CreateObject("Scripting.Dictionary")
'   Mac:     Scripting.Dictionary doesn't exist. Use a Collection
'            (see "Collections and dictionaries on a Mac" in the guide).
' Claims sheet: A ClaimID, B PatientID, C PayerID, D ServiceDate,
'   E BilledAmount, F AllowedAmount, G PaidAmount, H ClaimStatus,
'   I DenialReason.    Payers sheet: A PayerID, B PayerName, C PayerType.
' =====================================================================

' Task 8 -> Output!B4: how many different (non-blank) denial reasons?
Public Sub CountDenialReasons()
    ' TODO: read the Claims sheet into an array (CurrentRegion.Value), add every
    '       non-blank DenialReason to a Collection with AddIfNew, then write
    '       the Collection's Count to Output!B4
End Sub

' Task 9 -> Output!B5: how many different patients have a claim?
Public Sub CountDistinctPatients()
    ' TODO: add every PatientID to a Dictionary (Mac: a Collection) once, then
    '       write the Count to Output!B5
End Sub

' Task 10 -> Output!B6 (PayerName) and Output!B7 (number of Denied claims)
Public Sub TopDeniedPayer()
    ' TODO 1: lookup dictionary PayerID -> PayerName from the Payers sheet
    ' TODO 2: counting dictionary PayerID -> number of rows with ClaimStatus "Denied"
    ' TODO 3: loop over the keys to find the largest count, then write the payer's
    '         name to Output!B6 and the count to Output!B7
End Sub

{helpers}
'''


MAC_SLOT_HELPERS = '''' Returns the slot number stored under keyText. The first time a key is seen it
' gets the next free slot (Count + 1). A Collection can't change an item, so
' the numbers you want to update live in arrays indexed by this slot.
Private Function SlotFor(slots As Collection, ByVal keyText As String) As Long
    On Error Resume Next
    SlotFor = slots(keyText)               ' error 5 if the key isn't there yet
    If Err.Number <> 0 Then
        slots.Add Item:=slots.Count + 1, Key:=keyText
        SlotFor = slots.Count
    End If
    On Error GoTo 0
End Function

' Collection lookup that returns a default instead of raising error 5.
Private Function LookupOr(col As Collection, ByVal keyText As String, ByVal fallback As String) As String
    LookupOr = fallback
    On Error Resume Next
    LookupOr = col(keyText)                ' leaves the fallback in place if the key is missing
    On Error GoTo 0
End Function'''


def _mac_dict_solution_module() -> str:
    reasons = SOL_REASONS.replace("' Task 8:", "' Task 8 (same on Windows and Mac):")
    return f'''Attribute VB_Name = "ClaimCollectionsMacSolution"
Option Explicit

' =====================================================================
' Lesson 5.4 - Mac-friendly solutions for tasks 8-10 (SPOILERS)
'
' Excel for Mac has no Scripting.Dictionary (CreateObject fails with
' run-time error 429), so these versions use only the built-in Collection.
' They also run on Windows. Remove your own ClaimDictionaries module first
' (or import into a spare copy) so two modules don't share macro names.
' =====================================================================

{reasons}

' Task 9 (Mac): a keyed Collection keeps one entry per patient.
Public Sub CountDistinctPatients()
    Dim ws As Worksheet, claims As Variant
    Dim patientCol As Long, r As Long
    Dim patients As Collection

    Set ws = ThisWorkbook.Worksheets("Claims")
    claims = ws.Range("A1").CurrentRegion.Value
    patientCol = HeaderColumn(ws, "PatientID")
    Set patients = New Collection

    For r = 2 To UBound(claims, 1)
        AddIfNew patients, CStr(claims(r, patientCol))
    Next r

    ThisWorkbook.Worksheets("Output").Range("B5").Value = patients.Count
End Sub

' Task 10 (Mac): a Collection maps PayerID -> slot; an array holds the counts.
Public Sub TopDeniedPayer()
    Dim wsClaims As Worksheet, claims As Variant, payers As Variant
    Dim payerCol As Long, statusCol As Long, r As Long, i As Long
    Dim payerNames As Collection, slots As Collection
    Dim ids() As String, counts() As Long
    Dim bestID As String, bestCount As Long

    ' 1. Lookup Collection: Item = PayerName, Key = PayerID
    Set payerNames = New Collection
    payers = ThisWorkbook.Worksheets("Payers").Range("A1").CurrentRegion.Value
    For r = 2 To UBound(payers, 1)
        payerNames.Add Item:=payers(r, 2), Key:=CStr(payers(r, 1))
    Next r

    ' 2. Count Denied claims per payer in slot arrays
    Set wsClaims = ThisWorkbook.Worksheets("Claims")
    claims = wsClaims.Range("A1").CurrentRegion.Value
    payerCol = HeaderColumn(wsClaims, "PayerID")
    statusCol = HeaderColumn(wsClaims, "ClaimStatus")
    Set slots = New Collection
    ReDim ids(1 To UBound(claims, 1))
    ReDim counts(1 To UBound(claims, 1))
    For r = 2 To UBound(claims, 1)
        If claims(r, statusCol) = "Denied" Then
            i = SlotFor(slots, CStr(claims(r, payerCol)))
            ids(i) = CStr(claims(r, payerCol))
            counts(i) = counts(i) + 1
        End If
    Next r

    ' 3. Keep the largest count (on a tie, the payer seen first wins)
    For i = 1 To slots.Count
        If counts(i) > bestCount Then
            bestCount = counts(i)
            bestID = ids(i)
        End If
    Next i

    With ThisWorkbook.Worksheets("Output")
        .Range("B6").Value = LookupOr(payerNames, bestID, "(unknown payer " & bestID & ")")
        .Range("B7").Value = bestCount
    End With
End Sub

{DICT_HELPERS}

{MAC_SLOT_HELPERS}
'''


SNIPPETS = OrderedDict()
SNIPPETS["A"] = dict(task=7, name="SnippetA", title="ByVal vs ByRef", code='''Sub AddFeeByVal(ByVal amount As Double)
    amount = amount + 25
End Sub

Sub AddFeeByRef(ByRef amount As Double)
    amount = amount + 25
End Sub

Sub SnippetA()
    Dim charge As Double
    charge = 100
    AddFeeByVal charge
    AddFeeByRef charge
    AddFeeByRef (charge)
    Debug.Print charge
End Sub''')
SNIPPETS["B"] = dict(task=11, name="SnippetB", title="an error handler inside a loop", code='''Sub SnippetB()
    Dim readings As Variant, i As Long, total As Long
    readings = Array("120", "abc", "95", "", "88")
    On Error GoTo BadReading
    For i = LBound(readings) To UBound(readings)
        total = total + CLng(readings(i))
    Next i
    Debug.Print "Total:"; total
    Exit Sub
BadReading:
    Debug.Print "Skipped item"; i; "- error"; Err.Number
    Resume Next
End Sub''')
SNIPPETS["C"] = dict(task=12, name="SnippetC", title="On Error Resume Next", code='''Sub SnippetC()
    Dim readings As Variant, i As Long, reading As Long, total As Long
    readings = Array("120", "abc", "95", "", "88")
    On Error Resume Next
    For i = LBound(readings) To UBound(readings)
        reading = CLng(readings(i))
        total = total + reading
    Next i
    On Error GoTo 0
    Debug.Print "Total:"; total
End Sub''')
SNIPPETS["D"] = dict(task=13, name="SnippetD", title="two ways to call MATCH from VBA", code='''Sub SnippetD()
    Dim payerList As Range, pos As Variant
    Set payerList = ThisWorkbook.Worksheets("Payers").Range("A2:A9")

    pos = Application.Match("PY09", payerList, 0)
    If IsError(pos) Then
        Debug.Print "Application.Match: not found"
    Else
        Debug.Print "Application.Match: row"; pos
    End If

    pos = Application.WorksheetFunction.Match("PY09", payerList, 0)
    Debug.Print "WorksheetFunction.Match: row"; pos
End Sub''')


def _snippets_module() -> str:
    parts = ['''Attribute VB_Name = "Snippets"
Option Explicit

' =====================================================================
' Lesson 5.4 - predict-the-output snippets (tasks 7, 11, 12 and 13)
'
' Read a snippet and type your prediction on the Practice sheet FIRST.
' Then click inside the snippet's Sub and press F5 to run it (Mac:
' Run > Run Sub/UserForm). Debug.Print writes to the Immediate window
' (View > Immediate Window; Ctrl + G on Windows). Snippet D stops with a
' run-time error on purpose: note the number, then click End.
' =====================================================================''']
    for key, s in SNIPPETS.items():
        parts.append(f"' ---- Snippet {key} (task {s['task']}): {s['title']} ----\n{s['code']}")
    return "\n\n".join(parts) + "\n"


BONUS_HELPERS = '''' ---------------------------------------------------------------------
' Helpers
' ---------------------------------------------------------------------

' True if this workbook has a sheet with that name.
Private Function SheetExists(ByVal sheetName As String) As Boolean
    Dim ws As Worksheet
    On Error Resume Next
    Set ws = ThisWorkbook.Worksheets(sheetName)    ' error 9 if there is no such sheet
    On Error GoTo 0
    SheetExists = Not ws Is Nothing
End Function

' Column number of a header in row 1. Raises a clear error if it is missing.
Private Function HeaderColumn(ws As Worksheet, ByVal headerText As String) As Long
    Dim pos As Variant
    pos = Application.Match(headerText, ws.Rows(1), 0)
    If IsError(pos) Then
        Err.Raise vbObjectError + 513, "HeaderColumn", _
                  "Column '" & headerText & "' was not found on sheet '" & ws.Name & "'."
    End If
    HeaderColumn = CLng(pos)
End Function

' True for real numbers (a blank cell or text is not a number).
Private Function IsNumberValue(v As Variant) As Boolean
    Select Case VarType(v)
        Case vbInteger, vbLong, vbSingle, vbDouble, vbCurrency, vbDecimal, vbByte
            IsNumberValue = True
        Case Else
            IsNumberValue = False
    End Select
End Function'''

BONUS_WRITE = '''    ' 6. A fresh output sheet: clear it if it exists, add it if it doesn't
    If SheetExists(OUT_SHEET) Then
        Set wsOut = ThisWorkbook.Worksheets(OUT_SHEET)
        wsOut.Cells.Clear
    Else
        Set wsOut = ThisWorkbook.Worksheets.Add(After:=ThisWorkbook.Worksheets("Payers"))
        wsOut.Name = OUT_SHEET
    End If

    ' 7. Write the block in one go, sort the payer rows by Billed (largest
    '    first), then add the TOTAL row and the skipped-rows line
    With wsOut
        .Range("A1").Resize(n + 1, 7).Value = outData
        .Range("A1").Resize(n + 1, 7).Sort Key1:=.Range("D1"), Order1:=xlDescending, Header:=xlYes
        totalRow = n + 2
        .Cells(totalRow, 1).Value = "TOTAL"
        .Cells(totalRow, 3).Value = totClaims
        .Cells(totalRow, 4).Value = totBilled
        .Cells(totalRow, 5).Value = totPaid
        .Cells(totalRow, 6).Value = totDenied
        .Cells(totalRow, 7).Value = totDenied / totClaims
        .Cells(totalRow + 2, 1).Value = "Rows skipped"
        .Cells(totalRow + 2, 2).Value = skipped
        .Range("D:E").NumberFormat = "#,##0.00"
        .Range("G:G").NumberFormat = "0.0%"
        .Rows(1).Font.Bold = True
        .Rows(totalRow).Font.Bold = True
        .Columns("A:G").AutoFit
    End With

CleanExit:
    On Error Resume Next                   ' cleanup must never raise an error of its own
    Application.Calculation = prevCalc
    Application.ScreenUpdating = True
    Exit Sub

Fail:
    MsgBox "BuildPayerSummary stopped: " & Err.Description & " (error " & Err.Number & ")", _
           vbCritical, "Payer summary"
    Resume CleanExit
End Sub'''

BONUS_START = '''    ' 1. Check first: a friendly message beats a crash
    If Not SheetExists("Claims") Or Not SheetExists("Payers") Then
        MsgBox "BuildPayerSummary needs a sheet named 'Claims' and a sheet named 'Payers'." & vbNewLine & _
               "One of them is missing or has been renamed. Fix that, then run the macro again.", _
               vbExclamation, "Payer summary"
        Exit Sub
    End If

    ' 2. From here on, any run-time error jumps to Fail, which restores Excel's settings
    prevCalc = Application.Calculation
    On Error GoTo Fail
    Application.ScreenUpdating = False
    Application.Calculation = xlCalculationManual

    Set wsClaims = ThisWorkbook.Worksheets("Claims")
    claims = wsClaims.Range("A1").CurrentRegion.Value2     ' Value2: plain numbers, no Currency/Date types
    colPayer = HeaderColumn(wsClaims, "PayerID")
    colBilled = HeaderColumn(wsClaims, "BilledAmount")
    colPaid = HeaderColumn(wsClaims, "PaidAmount")
    colStatus = HeaderColumn(wsClaims, "ClaimStatus")
    payers = ThisWorkbook.Worksheets("Payers").Range("A1").CurrentRegion.Value2'''

SOL_BONUS = f'''Public Sub BuildPayerSummary()
    Dim wsClaims As Worksheet, wsOut As Worksheet
    Dim claims As Variant, payers As Variant, outData() As Variant
    Dim payerNames As Object, stats As Object
    Dim colPayer As Long, colBilled As Long, colPaid As Long, colStatus As Long
    Dim r As Long, n As Long, totalRow As Long, skipped As Long
    Dim payerID As Variant, s As Variant
    Dim totClaims As Long, totDenied As Long, totBilled As Double, totPaid As Double
    Dim prevCalc As Long

{BONUS_START}

    ' 3. Lookup dictionary: PayerID -> PayerName
    Set payerNames = CreateObject("Scripting.Dictionary")
    For r = 2 To UBound(payers, 1)
        payerNames(payers(r, 1)) = payers(r, 2)
    Next r

    ' 4. One pass over the claims: PayerID -> Array(claims, billed, paid, denied)
    Set stats = CreateObject("Scripting.Dictionary")
    For r = 2 To UBound(claims, 1)
        If IsNumberValue(claims(r, colBilled)) And IsNumberValue(claims(r, colPaid)) Then
            payerID = claims(r, colPayer)
            If Not stats.Exists(payerID) Then stats.Add payerID, Array(0, 0, 0, 0)
            s = stats(payerID)                 ' copy the array out of the dictionary...
            s(0) = s(0) + 1
            s(1) = s(1) + claims(r, colBilled)
            s(2) = s(2) + claims(r, colPaid)
            If claims(r, colStatus) = "Denied" Then s(3) = s(3) + 1
            stats(payerID) = s                 ' ...and put it back: stats(payerID)(0) = 1 would not stick
        Else
            skipped = skipped + 1              ' blank or text amount
        End If
    Next r
    n = stats.Count
    If n = 0 Then Err.Raise vbObjectError + 514, "BuildPayerSummary", "No claim rows with numeric amounts were found."

    ' 5. Header + one row per payer, built in memory, plus the running totals
    ReDim outData(1 To n + 1, 1 To 7)
    outData(1, 1) = "PayerID": outData(1, 2) = "PayerName": outData(1, 3) = "Claims"
    outData(1, 4) = "Billed": outData(1, 5) = "Paid": outData(1, 6) = "Denied": outData(1, 7) = "DenialRate"
    r = 1
    For Each payerID In stats.Keys
        r = r + 1
        s = stats(payerID)
        outData(r, 1) = payerID
        If payerNames.Exists(payerID) Then
            outData(r, 2) = payerNames(payerID)
        Else
            outData(r, 2) = "(unknown payer)"
        End If
        outData(r, 3) = s(0): outData(r, 4) = s(1): outData(r, 5) = s(2): outData(r, 6) = s(3)
        outData(r, 7) = s(3) / s(0)
        totClaims = totClaims + s(0): totBilled = totBilled + s(1)
        totPaid = totPaid + s(2): totDenied = totDenied + s(3)
    Next payerID

{BONUS_WRITE}'''

SOL_BONUS_MAC = f'''Public Sub BuildPayerSummary()
    Dim wsClaims As Worksheet, wsOut As Worksheet
    Dim claims As Variant, payers As Variant, outData() As Variant
    Dim payerNames As Collection, slots As Collection
    Dim ids() As String, cnt() As Long, denied() As Long, billed() As Double, paid() As Double
    Dim colPayer As Long, colBilled As Long, colPaid As Long, colStatus As Long
    Dim r As Long, i As Long, n As Long, totalRow As Long, skipped As Long
    Dim totClaims As Long, totDenied As Long, totBilled As Double, totPaid As Double
    Dim prevCalc As Long

{BONUS_START}

    ' 3. Lookup Collection: Item = PayerName, Key = PayerID
    Set payerNames = New Collection
    For r = 2 To UBound(payers, 1)
        payerNames.Add Item:=payers(r, 2), Key:=CStr(payers(r, 1))
    Next r

    ' 4. One pass over the claims. A Collection maps PayerID -> slot number;
    '    parallel arrays hold each slot's claims, billed, paid and denied.
    Set slots = New Collection
    ReDim ids(1 To UBound(claims, 1)): ReDim cnt(1 To UBound(claims, 1))
    ReDim denied(1 To UBound(claims, 1)): ReDim billed(1 To UBound(claims, 1))
    ReDim paid(1 To UBound(claims, 1))
    For r = 2 To UBound(claims, 1)
        If IsNumberValue(claims(r, colBilled)) And IsNumberValue(claims(r, colPaid)) Then
            i = SlotFor(slots, CStr(claims(r, colPayer)))
            ids(i) = CStr(claims(r, colPayer))
            cnt(i) = cnt(i) + 1
            billed(i) = billed(i) + claims(r, colBilled)
            paid(i) = paid(i) + claims(r, colPaid)
            If claims(r, colStatus) = "Denied" Then denied(i) = denied(i) + 1
        Else
            skipped = skipped + 1              ' blank or text amount
        End If
    Next r
    n = slots.Count
    If n = 0 Then Err.Raise vbObjectError + 514, "BuildPayerSummary", "No claim rows with numeric amounts were found."

    ' 5. Header + one row per payer, built in memory, plus the running totals
    ReDim outData(1 To n + 1, 1 To 7)
    outData(1, 1) = "PayerID": outData(1, 2) = "PayerName": outData(1, 3) = "Claims"
    outData(1, 4) = "Billed": outData(1, 5) = "Paid": outData(1, 6) = "Denied": outData(1, 7) = "DenialRate"
    For i = 1 To n
        outData(i + 1, 1) = ids(i)
        outData(i + 1, 2) = LookupOr(payerNames, ids(i), "(unknown payer)")
        outData(i + 1, 3) = cnt(i): outData(i + 1, 4) = billed(i): outData(i + 1, 5) = paid(i)
        outData(i + 1, 6) = denied(i): outData(i + 1, 7) = denied(i) / cnt(i)
        totClaims = totClaims + cnt(i): totBilled = totBilled + billed(i)
        totPaid = totPaid + paid(i): totDenied = totDenied + denied(i)
    Next i

{BONUS_WRITE}'''


def _bonus_header(mac: bool) -> str:
    name = "PayerSummaryMacSolution" if mac else "PayerSummarySolution"
    platform = ("Mac-friendly: Collections only (it also runs on Windows)." if mac else
                "Windows: Scripting.Dictionary. Excel for Mac: use PayerSummaryMac_Solution.bas.")
    return f'''Attribute VB_Name = "{name}"
Option Explicit

' =====================================================================
' Lesson 5.4 - reference solution for the bonus (SPOILERS)
' {platform}
'
' BuildPayerSummary writes a payer scorecard to the PayerSummary sheet:
'   row 1        PayerID | PayerName | Claims | Billed | Paid | Denied | DenialRate
'   rows 2..n+1  one row per payer, sorted by Billed (largest first)
'   row n+2      TOTAL
'   row n+4      "Rows skipped" and how many claim rows had a blank/text amount
' Remove your own PayerSummary module first (or import into a spare copy).
' =====================================================================

Private Const OUT_SHEET As String = "PayerSummary"
'''


def _bonus_solution_module(mac: bool) -> str:
    body = SOL_BONUS_MAC if mac else SOL_BONUS
    helpers = BONUS_HELPERS + ("\n\n" + MAC_SLOT_HELPERS if mac else "")
    return _bonus_header(mac) + "\n" + body + "\n\n" + helpers + "\n"


def _bonus_starter_module() -> str:
    return f'''Attribute VB_Name = "PayerSummary"
Option Explicit

' =====================================================================
' Lesson 5.4 - bonus starter: BuildPayerSummary
' The full spec is on the Bonus sheet. The helpers at the bottom are
' finished for you. Windows: use Scripting.Dictionary. Mac: use a
' Collection plus arrays (see the guide).
' =====================================================================

Private Const OUT_SHEET As String = "PayerSummary"

Public Sub BuildPayerSummary()
    ' Declare your variables here (Option Explicit is on)

    ' 1. If the Claims or Payers sheet is missing: show a friendly MsgBox and Exit Sub

    ' 2. Remember Application.Calculation, then On Error GoTo Fail, then turn off
    '    ScreenUpdating and switch to manual calculation

    ' 3. Read Claims and Payers into arrays (CurrentRegion.Value2) and find the
    '    columns you need with HeaderColumn

    ' 4. Lookup dictionary PayerID -> PayerName

    ' 5. One pass over the claims: skip (and count) rows whose BilledAmount or
    '    PaidAmount is not a number; otherwise update the payer's claims, billed,
    '    paid and denied totals

    ' 6. Build the output array: header row + one row per payer, and the totals

    ' 7. Clear or create the PayerSummary sheet, write the array, sort the payer
    '    rows by Billed (largest first), then add the TOTAL row and, two rows
    '    below it, "Rows skipped" with the count

CleanExit:
    ' 8. Restore ScreenUpdating and Calculation. This block runs after success
    '    AND after an error.
    Exit Sub

Fail:
    ' 9. Tell the user what went wrong (Err.Description, Err.Number), then
    '    Resume CleanExit
End Sub

{BONUS_HELPERS}
'''


def _crlf(text: str) -> str:
    assert text.isascii(), "VBA modules must be pure ASCII (the VBE imports them as ANSI)"
    return text.replace("\r\n", "\n").replace("\n", "\r\n")


def _count_word(n: int) -> str:
    return {1: "One", 2: "Two", 3: "Three", 4: "Four"}.get(n, str(n))


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------
def _age(dob: date, on) -> int:
    on = on.date() if hasattr(on, "date") else on
    return on.year - dob.year - ((on.month, on.day) < (dob.month, dob.day))


def _num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _datasets():
    pats = data.index(data.load("patients"), "PatientID")
    encounters = data.load("encounters")
    enc_fac = {e["EncounterID"]: e["FacilityID"] for e in encounters}
    f03 = [e for e in encounters if e["FacilityID"] == FAC and e["AdmitDateTime"].year == 2025]

    # Patients: adults (18+ on the as-of date) with any 2025 Cedar Ridge encounter
    pids = sorted({e["PatientID"] for e in f03})
    patients = [pats[p] for p in pids if _age(pats[p]["DOB"], AS_OF) >= 18]

    # Encounters: 2025 inpatient + observation stays, DOB joined in
    stays = sorted((e for e in f03 if e["EncounterType"] in ("Inpatient", "Observation")),
                   key=lambda e: (e["AdmitDateTime"], e["EncounterID"]))
    stays = [dict(e, DOB=pats[e["PatientID"]]["DOB"]) for e in stays]
    typos = []
    # Three discharges keyed with the wrong year (2024 instead of 2025)
    for i in (45, 205, 365):
        s = stays[i]
        s["DischargeDateTime"] = s["DischargeDateTime"].replace(year=2024)
        typos.append(s["EncounterID"])
    # One AM/PM slip: a same-day afternoon discharge keyed 12 hours early, which lands before the admission
    for s in stays:
        a, d = s["AdmitDateTime"], s["DischargeDateTime"]
        if s["EncounterID"] not in typos and a.date() == d.date() and d.hour >= 12 and d - timedelta(hours=12) < a:
            s["DischargeDateTime"] = d - timedelta(hours=12)
            typos.append(s["EncounterID"])
            break
    assert len(typos) == 4

    # Claims: every claim with a 2025 service date for a Cedar Ridge encounter
    claims = sorted((c for c in data.load("claims") if c["ServiceDate"].year == 2025 and enc_fac[c["EncounterID"]] == FAC),
                    key=lambda c: c["ClaimID"])
    # Five interface glitches in BilledAmount: blanks and text. One of them is a Silverline (PY02) denial,
    # so a macro that treats a blank as 0 instead of skipping it gets a different denial rate.
    py02_denied = next(i for i, c in enumerate(claims) if i >= 300 and c["PayerID"] == "PY02" and c["ClaimStatus"] == "Denied")
    glitches = {100: "N/A", py02_denied: None, 450: "VOID", 700: None, 950: "N/A"}
    assert len(glitches) == 5
    for i, v in glitches.items():
        claims[i]["BilledAmount"] = v

    payers = data.load("payers")
    return patients, stays, claims, payers


def _check_guide_examples(patients, stays):
    """The README guide's worked examples quote these rows and numbers in prose. Fail the build if the data drifts."""
    p = patients[1]                                   # "Albert Romero (PT10005, row 3 of the Patients sheet)"
    assert (p["PatientID"], p["FirstName"], p["LastName"], p["WeightLb"], p["HeightIn"]) == \
        ("PT10005", "Albert", "Romero", 197.9, 70.4), p
    assert repr(703 * p["WeightLb"] / p["HeightIn"] ** 2) == "28.0708653473657"   # what ? BMI(197.9, 70.4) prints
    s = stays[12]                                     # "Encounter ENC111079 (row 14 of the Encounters sheet)"
    assert (s["EncounterID"], s["DOB"], s["AdmitDateTime"].date()) == ("ENC111079", date(1961, 10, 9), date(2025, 1, 10))
    assert _age(s["DOB"], s["AdmitDateTime"]) == 63
    s = stays[24]                                     # "Encounter ENC111251 (row 26 of the Encounters sheet)"
    assert (s["EncounterID"], s["EncounterType"]) == ("ENC111251", "Observation")
    assert (s["AdmitDateTime"].strftime("%m/%d/%Y %H:%M"), s["DischargeDateTime"].strftime("%m/%d/%Y %H:%M")) == \
        ("01/15/2025 21:16", "01/16/2025 10:27")
    assert (s["DischargeDateTime"].date() - s["AdmitDateTime"].date()).days == 1
    by_type = Counter(x["EncounterType"] for x in stays)   # CountStaysByType prints "Inpatient 408" then "Observation 97"
    assert list(by_type.items()) == [("Inpatient", 408), ("Observation", 97)], by_type


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------
def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="05-automation-vba", slug="04-vba-functions-error-handling",
        title="VBA: Custom Functions, Dictionaries & Error Handling", level="Expert", minutes=65,
        objectives=[
            "Write user-defined functions (UDFs) you can call from cells",
            "Pass arguments ByVal/ByRef, use Optional arguments, and return errors with CVErr",
            "Count and group with Collections and Scripting.Dictionary",
            "Handle errors gracefully with On Error, the Err object, and cleanup code",
        ],
        data_note="Cedar Ridge Medical Center, 2025: adult patients with height and weight, inpatient and observation "
                  "stays, every claim with its payer and status, and the 8 payers. A few rows contain deliberate "
                  "data-entry errors for your code to catch.",
    )
    patients, stays, claims, payers = _datasets()
    _check_guide_examples(patients, stays)
    payer_name = {p["PayerID"]: p["PayerName"] for p in payers}

    pat = L.add_table_sheet(
        "Patients", patients, table="tblPatients",
        columns=["PatientID", "FirstName", "LastName", "Sex", "DOB", "HeightIn", "WeightLb"],
        extra_cols=["BMI"], formats={"BMI": "0.0", "HeightIn": "0.0", "WeightLb": "0.0"}, widths={"BMI": 10},
    )
    enc = L.add_table_sheet(
        "Encounters", stays, table="tblEncounters",
        columns=["EncounterID", "PatientID", "EncounterType", "DOB", "AdmitDateTime", "DischargeDateTime"],
        extra_cols=["AgeAtAdmit", "LOSDays"], formats={"AgeAtAdmit": "0", "LOSDays": "0"},
        widths={"AdmitDateTime": 18, "DischargeDateTime": 19, "AgeAtAdmit": 12, "LOSDays": 10},
    )
    clm = L.add_table_sheet(
        "Claims", claims, table="tblClaims",
        columns=["ClaimID", "PatientID", "PayerID", "ServiceDate", "BilledAmount", "AllowedAmount", "PaidAmount",
                 "ClaimStatus", "DenialReason"],
        formats={"BilledAmount": "#,##0.00", "AllowedAmount": "#,##0.00", "PaidAmount": "#,##0.00"},
        widths={"DenialReason": 26, "BilledAmount": 14},
    )
    pay = L.add_table_sheet("Payers", payers, table="tblPayers", columns=["PayerID", "PayerName", "PayerType"])
    assert [r["PayerID"] for r in payers] == [f"PY0{i}" for i in range(1, 9)]
    # The VBA reads these sheets by fixed position (A1 CurrentRegion, Payers columns A:B, Snippet D's A2:A9)
    assert pat.first_row == enc.first_row == clm.first_row == pay.first_row == 2 and pay.last_row == 9
    assert clm.headers[:9] == ["ClaimID", "PatientID", "PayerID", "ServiceDate", "BilledAmount", "AllowedAmount",
                               "PaidAmount", "ClaimStatus", "DenialReason"]

    def rng(sd, col):
        return sd.rng(col)

    # ------------------------------------------------------------------ answers (computed in Python)
    p0 = patients[0]
    bmi_vals = [703 * p["WeightLb"] / p["HeightIn"] ** 2 for p in patients]
    assert all(abs(b - 30) > 1e-6 for b in bmi_vals)
    bmi_first = bmi_vals[0]
    obese = sum(1 for b in bmi_vals if b >= 30)

    ages = [_age(s["DOB"], s["AdmitDateTime"]) for s in stays]
    avg_age = sum(ages) / len(ages)
    naive_ages = [s["AdmitDateTime"].year - s["DOB"].year for s in stays]          # DateDiff("yyyy", dob, admit)
    naive_avg = sum(naive_ages) / len(naive_ages)
    n_bday_ahead = sum(1 for a, b in zip(ages, naive_ages) if a != b)

    los = []
    for s in stays:
        a, d = s["AdmitDateTime"], s["DischargeDateTime"]
        los.append("#VALUE!" if d < a else (d.date() - a.date()).days)
    los_total = sum(x for x in los if x != "#VALUE!")
    los_bad = sum(1 for x in los if x == "#VALUE!")
    # A version that only compares dates (not times) misses the AM/PM slip
    los_bad_dates_only = sum(1 for s in stays if s["DischargeDateTime"].date() < s["AdmitDateTime"].date())

    statuses = [c["ClaimStatus"] for c in claims]
    denial_rate = statuses.count("Denied") / len(statuses)

    reasons = OrderedDict()
    for c in claims:
        r = (c["DenialReason"] or "").strip()
        if r:
            reasons[r] = True
    n_reasons = len(reasons)
    n_patients = len({c["PatientID"] for c in claims})
    repeat_patients = sum(1 for v in Counter(c["PatientID"] for c in claims).values() if v > 1)

    denied_by_payer: "OrderedDict[str, int]" = OrderedDict()
    for c in claims:
        if c["ClaimStatus"] == "Denied":
            denied_by_payer[c["PayerID"]] = denied_by_payer.get(c["PayerID"], 0) + 1
    top_id = max(denied_by_payer, key=lambda k: denied_by_payer[k])   # first maximum, like the VBA's ">" test
    top_count = denied_by_payer[top_id]
    claims_by_payer = Counter(c["PayerID"] for c in claims)
    most_claims_id = claims_by_payer.most_common(1)[0][0]
    assert sorted(denied_by_payer.values())[-1] > sorted(denied_by_payer.values())[-2]   # no tie for the top spot

    # Predict-the-output snippets (simulate the VBA semantics)
    charge = 100                             # AddFeeByVal charge: works on a copy, so no change
    charge += 25                             # AddFeeByRef charge: changes the caller's variable
    snippet_a = charge                       # AddFeeByRef (charge): the parentheses pass a temporary copy

    readings = ["120", "abc", "95", "", "88"]
    snippet_b = sum(int(x) for x in readings if x.strip().isdigit())          # bad items skipped by Resume Next
    reading, snippet_c = 0, 0
    for x in readings:                                                        # On Error Resume Next keeps the old value
        if x.strip().isdigit():
            reading = int(x)
        snippet_c += reading
    snippet_d = 1004   # Excel raises run-time error 1004 when WorksheetFunction.Match finds nothing

    # Bonus: payer scorecard
    kept = [c for c in claims if _num(c["BilledAmount"]) and _num(c["PaidAmount"])]
    skipped = len(claims) - len(kept)
    stats: "OrderedDict[str, list]" = OrderedDict()
    for c in kept:
        s = stats.setdefault(c["PayerID"], [0, 0.0, 0.0, 0])
        s[0] += 1
        s[1] += c["BilledAmount"]
        s[2] += c["PaidAmount"]
        s[3] += c["ClaimStatus"] == "Denied"
    ranked = sorted(stats.items(), key=lambda kv: -kv[1][1])
    billed_sorted = [v[1] for _, v in ranked]
    assert all(a > b for a, b in zip(billed_sorted, billed_sorted[1:]))      # no ties in the sort
    n_pay = len(ranked)
    third_id = ranked[2][0]
    third_name = payer_name[third_id]
    silver_id = next(k for k, v in payer_name.items() if v == "Silverline Medicare Advantage")
    silver = stats[silver_id]
    silver_rate = silver[3] / silver[0]
    silver_row = 2 + [k for k, _ in ranked].index(silver_id)          # its row on PayerSummary
    silver_all = [c for c in claims if c["PayerID"] == silver_id]
    silver_blank_denials = sum(1 for c in silver_all if c["BilledAmount"] is None and c["ClaimStatus"] == "Denied")
    n_blank = sum(1 for c in claims if c["BilledAmount"] is None)
    text_vals = sorted({c["BilledAmount"] for c in claims if isinstance(c["BilledAmount"], str)})
    n_text = sum(1 for c in claims if isinstance(c["BilledAmount"], str))
    silver_rate_blank_as_zero = sum(c["ClaimStatus"] == "Denied" for c in silver_all if c["BilledAmount"] is None or _num(c["BilledAmount"])) / \
        sum(1 for c in silver_all if c["BilledAmount"] is None or _num(c["BilledAmount"]))
    tot_claims = sum(v[0] for _, v in ranked)
    tot_billed = sum(v[1] for _, v in ranked)
    tot_paid = sum(v[2] for _, v in ranked)
    tot_denied = sum(v[3] for _, v in ranked)
    total_row = n_pay + 2
    summary_rows = [["PayerID", "PayerName", "Claims", "Billed", "Paid", "Denied", "DenialRate"]]
    summary_rows += [[k, payer_name[k], v[0], v[1], v[2], v[3], v[3] / v[0]] for k, v in ranked]
    summary_rows += [["TOTAL", None, tot_claims, tot_billed, tot_paid, tot_denied, tot_denied / tot_claims], [],
                     ["Rows skipped", skipped]]
    assert len(summary_rows) == total_row + 2

    # ------------------------------------------------------------------ helpers for formulas
    P_W, P_H, P_BMI = rng(pat, "WeightLb"), rng(pat, "HeightIn"), rng(pat, "BMI")
    E_DOB, E_ADM, E_DIS = rng(enc, "DOB"), rng(enc, "AdmitDateTime"), rng(enc, "DischargeDateTime")
    E_AGE, E_LOS = rng(enc, "AgeAtAdmit"), rng(enc, "LOSDays")
    C_PID, C_PAY, C_BILL = rng(clm, "PatientID"), rng(clm, "PayerID"), rng(clm, "BilledAmount")
    C_PAID, C_STAT, C_REAS = rng(clm, "PaidAmount"), rng(clm, "ClaimStatus"), rng(clm, "DenialReason")
    PAY_ID, PAY_NAME = rng(pay, "PayerID"), rng(pay, "PayerName")
    los_started = f"COUNT({E_LOS})+SUMPRODUCT(--ISERROR({E_LOS}))=0"

    def out_summary(cell):
        return f'=IF(Output!{cell}="","",Output!{cell})'

    def out_fill(cell, *values):
        col, row = cell[0], int(cell[1:])
        return {"range": f"Output!{cell}:{col}{row + len(values) - 1}", "values": list(values)}

    def ps(cell_range):
        return f"INDIRECT(\"'{SUMMARY_SHEET}'!{cell_range}\")"

    def ps_lookup(value_col, label_col, label):
        return (f'=IFERROR(INDEX({ps(f"{value_col}1:{value_col}40")},MATCH("{label}",'
                f'{ps(f"{label_col}1:{label_col}40")},0)),"")')

    snip_note = "Predict first, then run it to check."

    L.start_notes = [
        "Macros can't be saved in an .xlsx file. Before you write any code, choose File → Save As and pick "
        "'Excel Macro-Enabled Workbook (*.xlsm)'.",
        "Download the four .bas files in the lesson's starter/ folder (on GitHub, open each file and click Download raw "
        "file). Open the Visual Basic Editor with Alt + F11 (Mac: Option + F11, or Developer → Visual Basic), then choose "
        "File → Import File… for each one. HealthUDFs.bas is for tasks 1–6, ClaimDictionaries.bas for tasks 8–10, "
        "Snippets.bas for tasks 7 and 11–13, and PayerSummary.bas for the bonus.",
        "Scripting.Dictionary works only in Excel for Windows. On a Mac, use the Collection techniques from the guide. "
        "The solutions/ folder has Mac versions of every macro.",
        "The bonus macro creates a PayerSummary sheet. The gray cells on the Bonus sheet read it, so they stay blank until "
        "the macro has run.",
    ]
    # Output and Snippets are made in the customize hook, so list them under "Sheets in this workbook" here.
    # PayerSummary isn't listed: it doesn't exist until the learner's bonus macro creates it.
    L.sheet_notes = [
        ("Output", "Your macros for tasks 8–10 write their results here. The gray cells on Practice read them."),
        ("Snippets", "The code for the predict-the-output tasks 7 and 11–13 (the same code as starter/Snippets.bas)."),
    ]
    L.bonus_where = ("Write the macro in the VBE. It creates the **PayerSummary** sheet, and the gray cells on the "
                     "workbook's **Bonus** sheet read it, so there's nothing to type there.")
    # The library's generic how-to lines say "type a formula or value into each yellow cell". Here 7 of the 13 Practice
    # answers and every auto-checked Bonus answer are gray cells fed by UDF columns or macros, so say that instead.
    L.practice_how = ("Go to the 'Practice' sheet. Type your answers for tasks 1, 6, 7, and 11–13 in the yellow cells. For "
                      "tasks 2–5, fill the yellow UDF columns on the data sheets. For tasks 8–10, run your macros. The gray "
                      "cells for those tasks fill in by themselves.")
    L.practice_instructions = (
        "Type your answers for tasks 1, 6, 7, and 11–13 in the yellow cells. The gray cells for tasks 2–5 and 8–10 fill in "
        "by themselves from your UDF columns and your macros. The Check column turns green when your answer matches. "
        f"Stuck? Read the hint, then the lesson guide. Answers: right-click a sheet tab → Unhide… → '{L.key_sheet}'.")
    L.bonus_instructions = (
        "There's nothing to type on this sheet. The gray cells for B1–B4 fill in once BuildPayerSummary has run, and the "
        "Check column turns green when they match. B5 is a test you run by hand, so compare your results with the key. "
        f"Stuck? Read the hint, then the lesson guide. Answers: right-click a sheet tab → Unhide… → '{L.bonus_key_sheet}'.")
    L.practice_intro = (
        "Save the workbook as .xlsm and import the starter modules first (see the Start Here sheet). In tasks 1–6 you write "
        "user-defined functions and use them in cells: tasks 1 and 6 in the yellow cell, and tasks 2–5 in yellow columns on "
        "the data sheets, which the gray cells summarize. Tasks 7 and 11–13 ask you to predict what a snippet on the "
        "Snippets sheet does. In tasks 8–10 you run macros, and the gray cells read what they wrote to the Output sheet.")

    pid_cell = f"Patients!{pat.col('WeightLb')}{pat.first_row}"
    hid_cell = f"Patients!{pat.col('HeightIn')}{pat.first_row}"
    L.tasks = [
        # ---------------------------------------------------------- UDFs
        Task(f"Write BMI(weightLb, heightIn) in a standard module (start from starter/HealthUDFs.bas). BMI = 703 × weight "
             f"in pounds ÷ (height in inches)². It must return #VALUE! for a blank or text input and #NUM! for a zero or "
             f"negative one. In the yellow cell, call it for the first patient on the Patients sheet (row 2, {p0['PatientID']}). "
             f"Keep full precision. The check accepts 2 decimal places.",
             answer=bmi_first, fmt="0.00", title="BMI for the first patient",
             solution=f"' In the yellow cell: =BMI({pid_cell},{hid_cell})\n\n" + SOL_BMI, solution_lang="vba",
             live=f"=703*{pid_cell}/{hid_cell}^2",
             hint="Assign the result to the function's own name. Return errors with CVErr",
             explanation=f"A **Function** hands back a value by assigning it to its own name (`BMI = …`). Because it's Public and "
                         f"sits in a standard module, Excel lists it with the built-in functions, so `=BMI(` works in any cell. "
                         f"The return type is Variant so the same function can return a number or an error value made by "
                         f"`CVErr`. The two helpers come from the starter module: `ArgValue` turns a cell reference into its "
                         f"value, and `IsNumberValue` rejects blanks and text. {p0['FirstName']} {p0['LastName']} weighs "
                         f"{p0['WeightLb']} lb at {p0['HeightIn']} in, so BMI = 703 × {p0['WeightLb']} ÷ {p0['HeightIn']}² "
                         f"≈ {bmi_first:.2f}. If the cell shows #NAME?, the function isn't in a standard module of *this* "
                         f"workbook, or macros are disabled."),
        Task("Fill the yellow BMI column on the Patients sheet with your BMI function for every patient. Don't round inside "
             "the function: the column's number format already shows 1 decimal place. The gray cell counts the patients "
             "whose unrounded BMI is 30 or more (the usual adult obesity threshold).",
             answer=obese, title="BMI column (patients with BMI ≥ 30)",
             solution="=BMI([@WeightLb],[@HeightIn])",
             summary=f'=IF(COUNT({P_BMI})=0,"",COUNTIF({P_BMI},">=30"))',
             fill={"range": f"Patients!{pat.col('BMI')}{pat.first_row}:{pat.col('BMI')}{pat.last_row}", "values": bmi_vals},
             live=f"=SUMPRODUCT(--(703*{P_W}/{P_H}^2>=30))",
             hint="Type it once in the first BMI cell. The Table fills the rest",
             explanation=f"Type the formula in {pat.cell('BMI', 0, sheet=False)} (or `=BMI({pat.col('WeightLb')}2,"
                         f"{pat.col('HeightIn')}2)`, which does the same). "
                         f"A UDF behaves like any other function: write it once in the first row and the Excel Table copies it "
                         f"down as a calculated column, so each of the {len(patients)} rows recalculates whenever its height or "
                         f"weight changes. That's the payoff of a UDF over a macro, because a macro would write values once and "
                         f"go stale. The key's live formula does the same math with SUMPRODUCT, which shows a UDF is a "
                         f"convenience: anything it does with simple arithmetic you could also do with a formula."),
        Task("Write AGEAT(dob, [asOf]): a person's age in completed years on the asOf date, or on today's date when asOf is "
             "left out (use IsMissing). Then fill the yellow AgeAtAdmit column on the Encounters sheet with each patient's age "
             "on the admission date. The gray cell averages your column, and the check compares that average to 2 decimal "
             "places.",
             answer=avg_age, fmt="0.00", title="AgeAtAdmit column (average age at admission)",
             solution=f"' In the AgeAtAdmit column ({enc.cell('AgeAtAdmit', 0, sheet=False)}): =AGEAT([@DOB],[@AdmitDateTime])\n\n"
                      + SOL_AGEAT, solution_lang="vba",
             summary=f'=IF(COUNT({E_AGE})=0,"",AVERAGE({E_AGE}))',
             fill={"range": f"Encounters!{enc.col('AgeAtAdmit')}{enc.first_row}:{enc.col('AgeAtAdmit')}{enc.last_row}",
                   "values": ages},
             live=f"=SUMPRODUCT(YEAR({E_ADM})-YEAR({E_DOB})-(MONTH({E_ADM})*100+DAY({E_ADM})<MONTH({E_DOB})*100+DAY({E_DOB})))"
                  f"/ROWS({E_DOB})",
             hint="Optional asOf As Variant + IsMissing. Then check whether this year's birthday has happened",
             explanation=f"`IsMissing` works only on an **Optional Variant** argument, which is why `asOf` is declared "
                         f"`As Variant`. The tricky part is the birthday: `Year(asOf) - Year(dob)` counts calendar years, "
                         f"so it's one too high whenever the birthday falls later in the year than the admission. "
                         f"`DateSerial(Year(asOf), Month(dob), Day(dob))` builds this year's birthday. If that's after "
                         f"asOf, subtract 1. `DateDiff(\"yyyy\", dob, asOf)` has the same flaw because it counts year "
                         f"boundaries, not birthdays. On this data it is wrong for {n_bday_ahead} of {len(stays)} stays "
                         f"and gives an average of {naive_avg:.2f} instead of {avg_age:.2f}."),
        Task("Write LOSDAYS(admitDateTime, dischargeDateTime): the number of midnights between admission and discharge, or "
             "#VALUE! when an input isn't a date or the discharge date-time is earlier than the admission date-time. Fill the "
             "yellow LOSDays column on the Encounters sheet. The gray cell adds up your column and skips error cells. "
             "What is the total?",
             answer=los_total, title="LOSDays column (total of the valid stays)",
             solution=f"' In the LOSDays column ({enc.cell('LOSDays', 0, sheet=False)}): "
                      f"=LOSDAYS([@AdmitDateTime],[@DischargeDateTime])\n\n" + SOL_LOSDAYS, solution_lang="vba",
             summary=f'=IF({los_started},"",AGGREGATE(9,6,{E_LOS}))',
             fill={"range": f"Encounters!{enc.col('LOSDays')}{enc.first_row}:{enc.col('LOSDays')}{enc.last_row}", "values": los},
             live=f"=SUMPRODUCT(({E_DIS}>={E_ADM})*(INT({E_DIS})-INT({E_ADM})))",
             hint="CVErr(xlErrValue) for bad rows. DateDiff(\"d\", …) counts midnights",
             explanation="`DateDiff(\"d\", admit, discharge)` counts the calendar-day boundaries crossed, so a patient admitted "
                         "at 23:00 and discharged at 01:00 the next morning has a stay of 1 midnight, which is how inpatient "
                         "days are counted. Returning `CVErr(xlErrValue)` makes the bad rows impossible to miss, and it keeps "
                         "them out of totals: the gray cell uses `AGGREGATE(9,6,…)` (9 = SUM, 6 = ignore error values). If "
                         "you had returned 0 or a negative number instead, it would have silently dragged the total down."),
        Task("How many rows does your LOSDays column flag with #VALUE!? The gray cell counts them. (Each one is a data-entry "
             "error to send back to Health Information Management.)",
             answer=los_bad, title="LOSDays rows flagged #VALUE!",
             solution='=COUNTIF(tblEncounters[LOSDays],"#VALUE!")',
             summary=f'=IF({los_started},"",COUNTIF({E_LOS},"#VALUE!"))',
             fill={"range": f"Encounters!{enc.col('LOSDays')}{enc.first_row}:{enc.col('LOSDays')}{enc.last_row}", "values": los},
             live=f"=SUMPRODUCT(--({E_DIS}<{E_ADM}))",
             hint="Filter the LOSDays column to see the error rows",
             explanation=f"Nothing new to type here: the gray cell already holds this formula. COUNTIF can count one "
                         f"specific error value when you give the error's text as the criteria. "
                         f"Three discharges were keyed with the year 2024, and one same-day stay has its discharge time keyed "
                         f"12 hours early (an AM/PM slip). Compare full date-times, not just dates: a check like "
                         f"`DateValue(d) < DateValue(a)` finds only {los_bad_dates_only} of the {los_bad} problems, because "
                         f"the AM/PM slip has the right date. The check counts #VALUE! only, so a function that returns "
                         f"#N/A or #NUM! for these rows shows ✘."),
        Task("Write DENIALRATE(statusRange, [statusToCount]): the share of non-blank cells in statusRange whose text equals "
             "statusToCount (ignoring case). statusToCount is Optional and defaults to \"Denied\". Return #DIV/0! if the range "
             "has no non-blank cells. In the yellow cell, call DENIALRATE on the ClaimStatus column of the Claims sheet and "
             "leave out the second argument. The cell is already formatted as a percentage.",
             answer=denial_rate, fmt="0.0%", title="DENIALRATE on the ClaimStatus column",
             solution=f"' In the yellow cell: =DENIALRATE(tblClaims[ClaimStatus])\n"
                      f"' (or =DENIALRATE(Claims!{clm.col('ClaimStatus')}2:{clm.col('ClaimStatus')}{clm.last_row}))\n\n"
                      + SOL_DENIALRATE, solution_lang="vba",
             live=f'=COUNTIF({C_STAT},"Denied")/COUNTA({C_STAT})',
             hint="Optional statusToCount As String = \"Denied\", then For Each cell In statusRange.Cells",
             explanation=f"A typed Optional argument gets its default in the declaration (`= \"Denied\"`), so you don't need "
                         f"IsMissing. Declaring `statusRange As Range` lets you loop over its cells, and it means Excel "
                         f"returns #VALUE! by itself if someone passes a plain number. {statuses.count('Denied')} of "
                         f"{len(claims):,} claims are Denied. Because statusToCount is a parameter, the same function "
                         f"answers other questions too: `=DENIALRATE(tblClaims[ClaimStatus],\"Pending\")` gives the pending "
                         f"rate. The key's live cell uses COUNTIF/COUNTA because your UDF isn't in the downloaded .xlsx."),
        Task(f"Snippet A on the Snippets sheet passes a $100 charge to AddFeeByVal and AddFeeByRef. What number does "
             f"Debug.Print charge show? {snip_note}",
             answer=snippet_a, title="Snippet A: ByVal vs ByRef",
             solution="Trace it: `AddFeeByVal charge` adds 25 to a **copy** (charge stays 100). `AddFeeByRef charge` adds 25 to "
                      "the caller's variable itself (125). `AddFeeByRef (charge)` looks the same, but the parentheses turn "
                      "`charge` into an expression, so VBA passes a temporary copy and charge stays **125**.",
             hint="ByRef shares the caller's variable. What do parentheses around an argument do?", live=False,
             explanation="VBA passes arguments **ByRef by default**, so a procedure can change the caller's variable. Write "
                         "`ByVal` when a procedure must not do that. The parentheses trap catches experienced developers: "
                         "when you call a Sub without `Call`, wrapping one argument in parentheses evaluates it first and "
                         "passes the result, which silently turns ByRef into ByVal. Call Subs without parentheses, or use "
                         "`Call AddFeeByRef(charge)`."),
        # ---------------------------------------------------------- Collections & Dictionaries
        Task("Complete CountDenialReasons in starter/ClaimDictionaries.bas: add every non-blank DenialReason on the Claims "
             "sheet to a Collection, using the reason itself as the key, so each reason is kept once. The macro writes the "
             "Collection's Count to Output!B4, and the gray cell reads it.",
             answer=n_reasons, title="CountDenialReasons (Collection)",
             solution=SOL_REASONS + "\n\n" + ADD_IF_NEW, solution_lang="vba",
             summary=out_summary(OUT["reasons"]), fill=out_fill(OUT["reasons"], n_reasons),
             live=f'=ROWS(UNIQUE(FILTER({C_REAS},{C_REAS}<>"")))',
             hint="col.Add Item:=reason, Key:=reason raises an error for a key it already has",
             explanation=f"A Collection key must be unique, so adding a duplicate key raises run-time error 457. `AddIfNew` "
                         f"traps that error with `On Error Resume Next`, reads `Err.Number`, and switches trapping off "
                         f"again straight away with `On Error GoTo 0`. The {n_reasons} reasons are: "
                         f"{', '.join(sorted(reasons))}. Skipping blanks matters: most claims have no denial reason, "
                         f"and counting the empty string as a reason would give {n_reasons + 1}."),
        Task("Complete CountDistinctPatients: add every PatientID on the Claims sheet to a Scripting.Dictionary (Mac: a "
             "Collection) once, then write the number of keys to Output!B5. How many different patients had a claim?",
             answer=n_patients, title="CountDistinctPatients (Dictionary)",
             solution=SOL_PATIENTS, solution_lang="vba",
             summary=out_summary(OUT["patients"]), fill=out_fill(OUT["patients"], n_patients),
             live=f"=ROWS(UNIQUE({C_PID}))",
             hint="CreateObject(\"Scripting.Dictionary\"), then .Exists and .Add",
             explanation=f"`.Exists` asks whether a key is already in the dictionary, so each PatientID is added once and "
                         f"`.Count` is the number of different patients. {len(claims):,} claims come from {n_patients} "
                         f"patients because {repeat_patients} patients had more than one claim. `Dim patients As Object` with "
                         f"`CreateObject` is **late binding**: it needs no reference to the Scripting Runtime library, so "
                         f"the file works on any Windows PC. In Microsoft 365 and Excel 2021 or later, the key's "
                         f"`=ROWS(UNIQUE(…))` gives the same answer, but the dictionary pattern scales to jobs a formula "
                         f"can't do, like the bonus."),
        Task("Complete TopDeniedPayer: build one dictionary that maps PayerID → PayerName (from the Payers sheet) and another "
             "that counts the claims with ClaimStatus \"Denied\" for each PayerID (Mac: a Collection for the names, and "
             "SlotFor plus an array for the counts). Write the PayerName with the most denied claims to Output!B6 and its "
             "count to Output!B7. The gray cell reads B6.",
             answer=payer_name[top_id], accept=[top_id], title="TopDeniedPayer (two dictionaries)",
             answer_display=f"{payer_name[top_id]} ({top_count} denied claims)",
             solution=SOL_TOPPAYER, solution_lang="vba",
             summary=out_summary(OUT["top_name"]), fill=out_fill(OUT["top_name"], payer_name[top_id], top_count),
             live=(f'=INDEX({PAY_NAME},MATCH(MAX(COUNTIFS({C_PAY},{PAY_ID},{C_STAT},"Denied")),'
                   f'COUNTIFS({C_PAY},{PAY_ID},{C_STAT},"Denied"),0))'),
             hint="d(key) = d(key) + 1 counts. Then loop over .Keys to find the largest",
             explanation=f"`deniedCount(id) = deniedCount(id) + 1` is the **counting pattern**: reading a key that isn't there "
                         f"yet returns Empty (and quietly adds the key), and Empty + 1 = 1. Then one loop over `.Keys` keeps the "
                         f"largest count. The lookup dictionary replaces a VLOOKUP inside the loop. Note that the payer with the "
                         f"most claims overall is {payer_name[most_claims_id]} ({claims_by_payer[most_claims_id]} claims), but "
                         f"the most *denials* come from {payer_name[top_id]}: {top_count} of its "
                         f"{claims_by_payer[top_id]} claims."),
        # ---------------------------------------------------------- error handling
        Task(f"Snippet B reads five blood-glucose readings typed as text and adds them up, with an error handler that skips "
             f"bad entries. What total does the last line print? {snip_note}",
             answer=snippet_b, title="Snippet B: On Error GoTo + Resume Next",
             solution="CLng fails (run-time error 13, Type mismatch) on \"abc\" and on the empty string. Each time, VBA jumps to "
                      "BadReading, prints a line, and `Resume Next` continues with the statement **after** the one that "
                      "failed, which is `Next i`. So the bad items add nothing: 120 + 95 + 88 = **303**.",
             hint="Resume Next continues after the line that failed", live=False,
             explanation="`On Error GoTo label` sends any run-time error to the handler, and the `Exit Sub` before the label "
                         "keeps normal runs out of it. Inside the handler, `Err.Number` says what went wrong (13 = Type "
                         "mismatch). `Resume Next` skips the failing statement, `Resume` retries it (an endless loop here, "
                         "because \"abc\" never becomes a number), and `Resume SomeLabel` continues at a label."),
        Task(f"Snippet C reads the same five values with On Error Resume Next instead of a handler. What total does it print? "
             f"{snip_note}",
             answer=snippet_c, title="Snippet C: the On Error Resume Next trap",
             solution="When `CLng(\"abc\")` fails, the assignment never happens, so `reading` still holds 120 from the line "
                      "before, and it's added again. The same thing happens for \"\" (reading is still 95). "
                      "120 + 120 + 95 + 95 + 88 = **518**.",
             hint="When an assignment fails, what's left in the variable?", live=False,
             explanation=f"`On Error Resume Next` doesn't fix an error. It hides it. The variable keeps its previous value and "
                         f"the code carries on with wrong data: here {snippet_c - snippet_b} phantom units of glucose. Keep "
                         f"Resume Next to the one line you expect might fail, check `Err.Number` right after it, and switch "
                         f"back with `On Error GoTo 0`. AddIfNew in task 8 follows that pattern."),
        Task(f"Snippet D looks up a payer ID that doesn't exist (PY09), first with Application.Match and then with "
             f"Application.WorksheetFunction.Match. The first prints a message. The second stops the macro with a run-time "
             f"error. Type the error number. {snip_note}",
             answer=snippet_d, answer_display=str(snippet_d), title="Snippet D: Application.Match vs WorksheetFunction.Match",
             solution="`Application.Match` returns the error **value** Error 2042 (#N/A), which `IsError` detects, so the code "
                      "prints \"Application.Match: not found\". `WorksheetFunction.Match` raises a run-time **error** instead: "
                      "**1004**, *Unable to get the Match property of the WorksheetFunction class*.",
             hint="The 'application-defined or object-defined' error number", live=False,
             explanation="Both call Excel's MATCH, but they report failure differently. Use `Application.Match` (result in a "
                         "Variant, then `IsError`) when \"not found\" is a normal outcome, as in the HeaderColumn helper. "
                         "`WorksheetFunction.Match` only makes sense when a miss is a real error that your handler should "
                         "catch. The same split applies to VLOOKUP, INDEX, and the other lookup functions."),
    ]

    # ------------------------------------------------------------------ bonus
    L.bonus_title = "Bonus: a one-click payer scorecard"
    L.bonus_scenario = (
        "Cedar Ridge's revenue-cycle director wants a payer scorecard she can rebuild with one click every month. Write "
        "BuildPayerSummary (start from starter/PayerSummary.bas) so that it does five things:\n"
        "1. If the Claims or Payers sheet is missing, it shows a friendly message and stops without a run-time error.\n"
        "2. It reads the Claims table into an array and uses a Scripting.Dictionary keyed by PayerID (Mac: a Collection plus "
        "arrays) to work out four numbers for each payer: Claims (number of rows), Billed (sum of BilledAmount), Paid (sum of "
        "PaidAmount), and Denied (rows with ClaimStatus \"Denied\").\n"
        "3. It skips, and counts, every row whose BilledAmount or PaidAmount is not a number. A blank cell is not a number, "
        "and a skipped row is left out of every column.\n"
        "4. It writes a sheet named PayerSummary, clearing it first if it already exists. Row 1 holds the headers PayerID, "
        "PayerName, Claims, Billed, Paid, Denied, DenialRate (Denied ÷ Claims). Below it go one row per payer, sorted by "
        "Billed from largest to smallest, and then a TOTAL row: the word TOTAL in column A, the totals in C:F, and the overall "
        "rate in G. Two rows below TOTAL, column A holds the text Rows skipped and column B holds the count.\n"
        "5. It restores ScreenUpdating and Calculation even when an error stops the macro.")

    bonus_sol = SOL_BONUS + "\n\n" + BONUS_HELPERS
    L.bonus = [
        Task("Run BuildPayerSummary. Which payer has the third-highest billed charges? (It's the PayerName in cell B4 of "
             "PayerSummary, and the gray cell reads it.)",
             answer=third_name, accept=[third_id], title="Third-highest billed payer",
             solution=bonus_sol, solution_lang="vba",
             summary=f'=IFERROR(IF({ps("B4")}="","",{ps("B4")}),"")',
             fill={"range": f"{SUMMARY_SHEET}!B4:B4", "values": [third_name]},
             live=f"=INDEX(SORTBY({PAY_NAME},SUMIFS({C_BILL},{C_PAY},{PAY_ID}),-1),3)",
             hint="Range.Sort with Key1:=the Billed header cell, Order1:=xlDescending, Header:=xlYes",
             explanation=f"The dictionary loop collects the totals in whatever order payers first appear. Sorting the written "
                         f"block, header excluded, puts the biggest payers on top: {', '.join(payer_name[k] for k, _ in ranked[:3])}. "
                         f"Third place is close ({third_name} billed {ranked[2][1][1]:,.2f} against {ranked[3][1][1]:,.2f} "
                         f"for {payer_name[ranked[3][0]]}), so an unsorted or wrongly sorted table shows ✘. The macro "
                         f"writes the whole block with one `.Value = outData` assignment, which is much faster than writing "
                         f"cell by cell. **Mac:** the Collection version is in solutions/PayerSummaryMac_Solution.bas."),
        Task("What DenialRate does your PayerSummary show for Silverline Medicare Advantage? (The gray cell finds "
             "Silverline's row on PayerSummary and reads column G.)",
             answer=silver_rate, fmt="0.0%", title="Silverline Medicare Advantage denial rate",
             solution="Read it from PayerSummary column G. Formula cross-check (blank and text amounts excluded): "
                      f"`=COUNTIFS(tblClaims[PayerID],\"{silver_id}\",tblClaims[ClaimStatus],\"Denied\",tblClaims[BilledAmount],\">=0\")"
                      f"/COUNTIFS(tblClaims[PayerID],\"{silver_id}\",tblClaims[BilledAmount],\">=0\")`",
             summary=ps_lookup("G", "B", "Silverline Medicare Advantage"),
             fill={"range": f"{SUMMARY_SHEET}!G{silver_row}:G{silver_row}", "values": [silver_rate]},
             live=(f'=COUNTIFS({C_PAY},"{silver_id}",{C_STAT},"Denied",{C_BILL},">=0")'
                   f'/COUNTIFS({C_PAY},"{silver_id}",{C_BILL},">=0")'),
             hint="The array-in-a-dictionary pattern: copy it out, change it, put it back",
             explanation=f"Silverline has {silver[3]} denied claims out of {silver[0]} kept rows. "
                         f"{_count_word(silver_blank_denials)} Silverline denial"
                         f"{'s have' if silver_blank_denials != 1 else ' has'} a blank BilledAmount. "
                         f"`IsNumeric(Empty)` returns True, so a macro that tests amounts with IsNumeric "
                         f"keeps that row and reports {silver_rate_blank_as_zero:.1%} instead of {silver_rate:.1%}. Test the "
                         f"type with `VarType` instead, as IsNumberValue does. Also watch the array trap: "
                         f"`stats(id)(0) = stats(id)(0) + 1` changes a temporary copy and the dictionary never sees it. "
                         f"Copy the array out, change it, and assign it back."),
        Task("What total Paid does the TOTAL row of PayerSummary show? (The gray cell reads column E of the TOTAL row, and "
             "the check compares it to the cent.)",
             answer=round(tot_paid, 2), fmt="#,##0.00", title="Total Paid on the TOTAL row",
             solution="Read it from the TOTAL row, column E. Formula cross-check: "
                      "`=SUMIFS(tblClaims[PaidAmount],tblClaims[BilledAmount],\">=0\")`",
             summary=ps_lookup("E", "A", "TOTAL"),
             fill={"range": f"{SUMMARY_SHEET}!E{total_row}:E{total_row}", "values": [tot_paid]},
             live=f'=SUMIFS({C_PAID},{C_BILL},">=0")',
             hint="Accumulate the totals in the same loop that fills the output array",
             explanation=f"The macro adds each payer's figures to running totals while it builds the output array, then writes "
                         f"the TOTAL row after sorting, so the sort can't move it. It covers {tot_claims:,} claims "
                         f"and ${tot_billed:,.2f} billed. Writing the TOTAL row before sorting is a classic bug: Range.Sort "
                         f"would treat it as a payer and sort it to the top."),
        Task("How many claim rows did your macro skip because BilledAmount or PaidAmount was not a number? (The gray cell reads "
             "the number next to 'Rows skipped'.)",
             answer=skipped, title="Rows skipped",
             solution="Read it from the cell next to *Rows skipped*. Formula cross-check: "
                      "`=ROWS(tblClaims[BilledAmount])-COUNT(tblClaims[BilledAmount])`",
             summary=ps_lookup("B", "A", "Rows skipped"),
             fill={"range": f"{SUMMARY_SHEET}!B{total_row + 2}:B{total_row + 2}", "values": [skipped]},
             live=f"=ROWS({C_BILL})-COUNT({C_BILL})",
             hint="A blank cell read into an array is Empty, and VarType(Empty) is vbEmpty",
             explanation=f"Filter the Claims sheet's BilledAmount column to find them: {_count_word(n_blank).lower()} are "
                         f"blank and {_count_word(n_text).lower()} hold text ({', '.join(text_vals)}). Without the check, `s(1) + \"N/A\"` raises run-time error 13 (Type mismatch), the "
                         f"handler stops the macro, and nobody gets a scorecard. Counting and reporting skipped rows is "
                         f"better than silently dropping them, because the director can see that "
                         f"{skipped} claims need fixing at the source."),
        Task("Test the error handling. (1) Rename the Claims sheet to Claims_old and run BuildPayerSummary: you should get your "
             "friendly message, not a run-time error. (2) Rename it back. (3) Temporarily change the header text in the Claims "
             "sheet's BilledAmount cell (E1) to Billed and run the macro again: the Fail handler should report the missing "
             "column, and **Formulas → Calculation Options** should still show Automatic afterwards. (4) Put the header back and "
             "run the macro once more.",
             answer=None, check="manual", title="Robustness test",
             solution="1. With Claims renamed, `SheetExists(\"Claims\")` is False, so the macro shows the friendly MsgBox and "
                      "exits **before** it changes any setting.\n"
                      "2. With the header renamed, `HeaderColumn` raises error vbObjectError + 513. HeaderColumn has no "
                      "handler of its own, so the error travels up to BuildPayerSummary's `Fail` handler, which shows "
                      "*Column 'BilledAmount' was not found on sheet 'Claims'*. Then `Resume CleanExit` restores "
                      "Calculation and ScreenUpdating.\n"
                      "3. If Calculation Options shows Manual after the error, your error path skipped the cleanup block. "
                      "Set it back with **Formulas → Calculation Options → Automatic**, then route the handler through "
                      "`Resume CleanExit`. (Excel usually switches ScreenUpdating back on by itself when a macro ends, but the "
                      "calculation mode stays manual for every open workbook until someone changes it.)",
             hint="Check before you change settings, and clean up in one place",
             explanation="Good error handling has two layers. **Validate** the things you can predict (a missing sheet) and "
                         "explain them in plain language. **Trap** everything else in one handler that tells the user what "
                         "happened and always runs the cleanup code. Because errors travel up the call stack until they meet "
                         "a handler, the custom error raised in a helper function is reported by the main macro's handler."),
    ]

    # ------------------------------------------------------------------ files next to the workbook
    L.extra_files = {
        "starter/HealthUDFs.bas": _crlf(_udf_starter_module()),
        "starter/ClaimDictionaries.bas": _crlf(_dict_starter_module()),
        "starter/Snippets.bas": _crlf(_snippets_module()),
        "starter/PayerSummary.bas": _crlf(_bonus_starter_module()),
        "solutions/HealthUDFs_Solution.bas": _crlf(_udf_solution_module()),
        "solutions/ClaimDictionaries_Solution.bas": _crlf(_dict_solution_module()),
        "solutions/ClaimCollectionsMac_Solution.bas": _crlf(_mac_dict_solution_module()),
        "solutions/PayerSummary_Solution.bas": _crlf(_bonus_solution_module(mac=False)),
        "solutions/PayerSummaryMac_Solution.bas": _crlf(_bonus_solution_module(mac=True)),
    }

    # ------------------------------------------------------------------ extra sheets
    @L.customize
    def _extras(wb, lesson, selftest):
        _output_sheet(wb)
        _snippets_sheet(wb)
        if selftest:
            # Simulate what BuildPayerSummary writes, so the bonus summaries can be self-tested
            ws = wb.create_sheet(SUMMARY_SHEET)
            for row in summary_rows:
                ws.append(row)

    L.sheet_order = ["Start Here", "Practice", "Patients", "Encounters", "Claims", "Payers", "Output", "Snippets",
                     "Bonus", "Answer Key", "Bonus Key"]
    return L


def _output_sheet(wb):
    ws = wb.create_sheet("Output")
    ws.sheet_properties.tabColor = "2E75B6"
    ws["A1"] = "Output: your macros write their results here"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = ("Macros write into the blue cells. Don't type in column B: run the macro named in column C. The gray cells "
                "on the Practice sheet read these cells.")
    ws["A2"].font = Font(italic=True, color="595959")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:C2")
    ws.row_dimensions[2].height = 32
    rows = [
        ("Result", "Value", "Written by (task)"),
        ("Number of different denial reasons (blanks ignored)", None, "CountDenialReasons (task 8)"),
        ("Number of different patients on the Claims sheet", None, "CountDistinctPatients (task 9)"),
        ("Payer with the most Denied claims (PayerName)", None, "TopDeniedPayer (task 10)"),
        ("…its number of Denied claims", None, "TopDeniedPayer (task 10)"),
    ]
    for i, (a, b, c) in enumerate(rows):
        r = 3 + i
        for j, v in enumerate((a, b, c), 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.border = BOX
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if i == 0:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = HEADER_FILL
            elif j == 2:
                cell.fill = MACRO_FILL
    assert [f"B{3 + i}" for i in range(1, 5)] == [OUT["reasons"], OUT["patients"], OUT["top_name"], OUT["top_count"]]
    for col, w in zip("ABC", (52, 34, 32)):
        ws.column_dimensions[col].width = w
    _fit_page(ws)


def _fit_page(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _snippets_sheet(wb):
    ws = wb.create_sheet("Snippets")
    ws.sheet_properties.tabColor = "7030A0"
    ws.column_dimensions["A"].width = 100
    ws["A1"] = "Snippets for the predict-the-output tasks"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = ("Read the code and type your prediction on the Practice sheet first. Then import starter/Snippets.bas, click "
                "inside the Sub, and press F5 to run it (Mac: Run → Run Sub/UserForm). Debug.Print writes to the Immediate "
                "window: press Ctrl + G on Windows, or choose View → Immediate Window on a Mac. Snippet D stops with a "
                "run-time error on purpose: note the number, then click End.")
    ws["A2"].font = Font(italic=True, color="595959")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 62
    r = 4
    code_font = Font(name="Consolas", size=10)
    for key, s in SNIPPETS.items():
        c = ws.cell(row=r, column=1, value=f"Snippet {key} (task {s['task']}): {s['title']}. Run: {s['name']}")
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        r += 1
        for line in s["code"].split("\n"):
            c = ws.cell(row=r, column=1, value=line or None)
            if line:
                c.data_type = "s"           # never let openpyxl read "=..." as a formula
            c.font = code_font
            c.fill = PatternFill("solid", fgColor="F3F3F3")
            r += 1
        r += 1
    _fit_page(ws)


# ---------------------------------------------------------------------------
# Smoke test in LibreOffice (best effort; see module docstring)
# ---------------------------------------------------------------------------
def smoke_test():
    import tempfile

    from xlcourse import vba

    L = build()
    tmp = Path(tempfile.mkdtemp(prefix="l54_"))
    wb_path = L.save(tmp / L.workbook_name)
    patients, stays, claims, payers = _datasets()
    n_pat, n_enc = len(patients), len(stays)
    mods = {}
    # LibreOffice Basic can't compile "Debug.Print a; b", so the smoke test writes each snippet's result to a cell
    snip_lo = "Option Explicit\n" + "\n\n".join(s["code"] for k, s in SNIPPETS.items() if k != "D")
    snip_lo = (snip_lo.replace("Debug.Print charge", 'Worksheets("Output").Range("F1").Value = charge')
               .replace('    Debug.Print "Skipped item"; i; "- error"; Err.Number\n', "")
               .replace('Debug.Print "Total:"; total', 'Worksheets("Output").Range("F2").Value = total', 1)
               .replace('Debug.Print "Total:"; total', 'Worksheets("Output").Range("F3").Value = total', 1))
    assert "Debug.Print" not in snip_lo
    for name, text in [("udf", _udf_solution_module()), ("mac", _mac_dict_solution_module()),
                       ("bonus", _bonus_solution_module(mac=True)), ("snip", snip_lo)]:
        p = tmp / f"{name}.bas"
        p.write_text(text, encoding="utf-8")
        mods[name] = p
    # UDFs called from VBA (Range arguments) and from cells (values), written to spare columns
    test = f'''Option Explicit
Sub XlcUdfs()
    Dim r As Long, ws As Worksheet
    Set ws = Worksheets("Patients")
    For r = 2 To {n_pat + 1}
        ws.Cells(r, 8).Value = BMI(ws.Cells(r, 7), ws.Cells(r, 6))
    Next r
    Dim bad As Long
    Set ws = Worksheets("Encounters")
    For r = 2 To {n_enc + 1}
        ws.Cells(r, 7).Value = AGEAT(ws.Cells(r, 4), ws.Cells(r, 5))
        ws.Cells(r, 8).Value = LOSDAYS(ws.Cells(r, 5), ws.Cells(r, 6))
        If IsError(LOSDAYS(ws.Cells(r, 5), ws.Cells(r, 6))) Then bad = bad + 1
    Next r
    Worksheets("Output").Range("E8").Value = bad
    With Worksheets("Output")
        .Range("E1").Value = DENIALRATE(Worksheets("Claims").Range("H2:H{len(claims) + 1}"))
        .Range("E2").Value = DENIALRATE(Worksheets("Claims").Range("H2:H{len(claims) + 1}"), "Pending")
        .Range("E3").Formula = "=BMI(Patients!G2,Patients!F2)"
        .Range("E4").Formula = "=AGEAT(DATE(1960,12,31),DATE(2025,12,30))"
        .Range("E5").Value = IsError(BMI("abc", 60)) And IsError(BMI(150, 0)) And Not IsError(BMI(150, 60))
        .Range("E6").Value = IsError(LOSDAYS(#3/2/2025#, #3/1/2025#)) And IsError(LOSDAYS("x", #3/1/2025#))
        .Range("E7").Value = AGEAT(#12/31/1960#)
    End With
End Sub
'''
    tpath = tmp / "test.bas"
    tpath.write_text(test, encoding="utf-8")
    macros = ["XlcUdfs", "CountDenialReasons", "CountDistinctPatients", "TopDeniedPayer", "BuildPayerSummary",
              "SnippetA", "SnippetB", "SnippetC"]
    read = [("Output", c) for c in ("B4", "B5", "B6", "B7", "E1", "E2", "E3", "E4", "E5", "E6", "E7", "E8", "F1", "F2", "F3")]
    read += [(SUMMARY_SHEET, f"{c}{r}") for r in range(1, 13) for c in "ABCDEFG"]
    read += [("Patients", f"H{r}") for r in range(2, n_pat + 2)]
    read += [("Encounters", f"{c}{r}") for r in range(2, n_enc + 2) for c in "GH"]
    # Mac/Collection module and the Windows module define the same macro names, so test the Mac versions here
    res = vba.run(wb_path, [mods["udf"], mods["mac"], mods["bonus"], mods["snip"], tpath], macros, read=read)
    cells = res["cells"]
    print("errors:", res["errors"])
    print("sheets:", res["sheets"])
    exp = {t.title: t.answer for t in L.tasks + L.bonus}
    print("Output B4..B7:", [cells[("Output", c)] for c in ("B4", "B5", "B6", "B7")])
    print("DENIALRATE / Pending / cell BMI / AGEAT trap / IsError x2 / AGEAT(today) / LOS errors:",
          [cells[("Output", c)] for c in ("E1", "E2", "E3", "E4", "E5", "E6", "E7", "E8")])
    # LibreOffice passes "AddFeeByRef (charge)" by reference (150); real VBA passes a temporary copy (125)
    print("Snippets A (LO-specific), B, C:", [cells[("Output", c)] for c in ("F1", "F2", "F3")])
    bmi = [cells[("Patients", f"H{r}")] for r in range(2, n_pat + 2)]
    print("BMI >= 30:", sum(1 for b in bmi if isinstance(b, float) and b >= 30), "expected",
          exp["BMI column (patients with BMI ≥ 30)"])
    ages = [cells[("Encounters", f"G{r}")] for r in range(2, n_enc + 2)]
    print("avg age:", sum(ages) / len(ages), "expected", exp["AgeAtAdmit column (average age at admission)"])
    los = [cells[("Encounters", f"H{r}")] for r in range(2, n_enc + 2)]
    print("LOS total:", sum(x for x in los if isinstance(x, float)), "non-numeric:", [x for x in los if not isinstance(x, float)],
          "expected", exp["LOSDays column (total of the valid stays)"], exp["LOSDays rows flagged #VALUE!"])
    for r in range(1, 13):
        print("  ", [cells[(SUMMARY_SHEET, f"{c}{r}")] for c in "ABCDEFG"])
    print("bonus expected:", [exp[t.title] for t in L.bonus if t.answer is not None])
    print("task expected:", {k: v for k, v in exp.items() if k.startswith(("Count", "Top", "DENIAL", "BMI for"))})
    return res
