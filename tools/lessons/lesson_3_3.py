"""Lesson 3.3 · Cleaning Messy Data.

The data is the deliberately messy registration export in data/messy/patient_registrations_raw.csv (650 rows). The answer
truth file tools/_truth/patient_registrations_truth.csv maps each RecordID to its real PatientID; it is used here only to
prove that the cleaning formulas recover the right patients (it never goes into the workbook).
"""
from __future__ import annotations

import csv
from collections import Counter, defaultdict
from datetime import date, datetime

from openpyxl.styles import Alignment, Font

from xlcourse import Lesson, Task, data
from xlcourse.lesson import INPUT_FILL, ROOT, estimate_lines

CODE = "3.3"
RAW_COLS = ["RecordID", "PatientName", "DOB", "Sex", "Phone", "Email", "CityStateZip", "Insurance", "MRN", "RegisteredOn"]
CLEAN_COLS = ["NameClean", "DOBClean", "SexClean", "PhoneClean", "ZIP5", "PayerID", "MRNClean", "RegisteredClean", "Keep"]
MONTHS = "JanFebMarAprMayJunJulAugSepOctNovDec"
PAIR = ("R1017", "R1409")   # same patient, different name layout, phone layout, and MRN (zero lost)
KEEP_MRN = "01276260"          # bonus B3: three rows, stripped MRN on two of them, latest row is a 12-hour PM time
REPORT_ICU = ("Cedar Ridge Medical Center", "Intensive Care Unit")
CENSUS_WEEK = (date(2025, 12, 15), date(2025, 12, 21))
CENSUS_UNITS = ["D120", "D130", "D230", "D210", "D320", "D330", "D310"]   # facility blocks F01, F02, F03


# ---------------------------------------------------------------------------
# Python models of the Excel formulas (every answer is computed the way Excel computes it)
# ---------------------------------------------------------------------------
def xl_trim(s: str) -> str:
    """TRIM: strip ASCII spaces at both ends and collapse inner runs to one space."""
    return " ".join(p for p in s.split(" ") if p)


def xl_proper(s: str) -> str:
    """PROPER: uppercase a letter that follows a non-letter (or starts the text); lowercase every other letter."""
    out, prev = [], False
    for ch in s:
        if ch.isalpha():
            out.append(ch.lower() if prev else ch.upper())
            prev = True
        else:
            out.append(ch)
            prev = False
    return "".join(out)


def name_clean(s: str) -> str:
    """=PROPER(IF(ISNUMBER(FIND(",",B2)),TRIM(MID(B2,FIND(",",B2)+1,LEN(B2)))&" "&TRIM(LEFT(B2,FIND(",",B2)-1)),TRIM(B2)))"""
    if "," in s:
        i = s.find(",")
        return xl_proper(xl_trim(s[i + 1:]) + " " + xl_trim(s[:i]))
    return xl_proper(xl_trim(s))


def sex_clean(s: str) -> str:
    """=UPPER(LEFT(TRIM(D2),1))"""
    return xl_trim(s)[:1].upper()


def phone_clean(s: str) -> str:
    """=RIGHT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(E2,"(",""),")",""),"-",""),".","")," ",""),10)"""
    for ch in "()-. ":
        s = s.replace(ch, "")
    return s[-10:]


def mrn_clean(s: str) -> str:
    """=TEXT(VALUE(SUBSTITUTE(I2,"MRN-","")),"00000000")"""
    return f"{int(s.replace('MRN-', '')):08d}"


def dob_parse(s: str) -> date:
    """The locale-proof DATE() formula: ISO yyyy-mm-dd, dd-Mon-yyyy, and m/d/yyyy with / or . separators."""
    if s[4] == "-":
        return date(int(s[:4]), int(s[5:7]), int(s[-2:]))
    if s[2] == "-":
        return date(int(s[-4:]), (MONTHS.lower().find(s[3:6].lower()) + 1 + 2) // 3, int(s[:2]))
    p = s.replace(".", "/").find("/") + 1           # FIND is 1-based
    return date(int(s[-4:]), int(s[:p - 1]), int(s[p:p + len(s) - p - 5]))


def dob_us_datevalue(s: str) -> date:
    """What =DATEVALUE(SUBSTITUTE(C2,".","/")) returns with US (month-first) date settings."""
    t = s.replace(".", "/")
    for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    raise ValueError(s)


def reg_parse(s: str) -> datetime:
    """=IF(MID(J2,5,1)="-",DATE(LEFT(J2,4),MID(J2,6,2),MID(J2,9,2)),DATE(MID(J2,7,4),LEFT(J2,2),MID(J2,4,2)))+TIMEVALUE(MID(J2,12,8))"""
    if s[4] == "-":
        return datetime.strptime(s, "%Y-%m-%d %H:%M:%S")
    return datetime.strptime(s, "%m/%d/%Y %I:%M %p")


def zip5(s: str) -> str:
    """Flash Fill result (and =LEFT(TEXTAFTER(G2," ",-1),5)): the first five digits of the ZIP."""
    return s.rsplit(" ", 1)[1][:5]


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="03-data-analysis", slug="03-data-cleaning",
        title="Cleaning Messy Data", level="Intermediate", minutes=60,
        objectives=[
            "Spot common data problems: stray spaces, inconsistent case and categories, text dates, duplicates",
            "Fix them with TRIM, PROPER, SUBSTITUTE, VALUE, DATEVALUE, and lookup mapping tables",
            "Use Text to Columns, Flash Fill, Remove Duplicates, and Go To Special",
            "Document a repeatable cleaning process",
        ],
    )

    # ------------------------------------------------------------------ data: the raw registration export
    raw = data.load_raw("messy/patient_registrations_raw")
    raw_rows = [{c: (r[c] if r[c] != "" else None) for c in RAW_COLS} for r in raw]   # truly empty cells, not ""
    n = len(raw)
    truth_path = ROOT / "tools" / "_truth" / "patient_registrations_truth.csv"
    with truth_path.open(newline="", encoding="utf-8") as f:
        truth = {r["RecordID"]: r["PatientID"] for r in csv.DictReader(f)}
    patients = data.index(data.load("patients"), "PatientID")
    payers = data.index(data.load("payers"), "PayerID")

    widths = {"RecordID": 10, "PatientName": 26, "DOB": 12, "Sex": 8, "Phone": 17, "Email": 34, "CityStateZip": 28,
              "Insurance": 30, "MRN": 14, "RegisteredOn": 21}
    rw = L.add_table_sheet("Raw", raw_rows, columns=RAW_COLS, as_table=False, widths=widths, tab_color="7F7F7F")
    cl = L.add_table_sheet(
        "Clean", raw_rows, columns=RAW_COLS, table="tblClean", extra_cols=CLEAN_COLS,
        formats={"DOBClean": "mm/dd/yyyy", "RegisteredClean": "mm/dd/yyyy hh:mm"},
        widths={**widths, "NameClean": 22, "DOBClean": 12, "SexClean": 10, "PhoneClean": 13, "ZIP5": 8, "PayerID": 10,
                "MRNClean": 11, "RegisteredClean": 17, "Keep": 8},
    )
    assert rw.first_row == cl.first_row == 2 and rw.last_row == cl.last_row == n + 1

    # ------------------------------------------------------------------ data: payer mapping table
    variant_map = {
        "MEDICARE": "PY01", "MEDICARE PART A": "PY01",
        "SILVERLINE MEDICARE ADVANTAGE": "PY02", "SILVERLINE MA": "PY02", "SILVERLINE MEDICARE ADV.": "PY02",
        "STATE MEDICAID": "PY03", "MEDICAID": "PY03",
        "KEYSTONE HEALTH PARTNERS": "PY04", "KEYSTONE": "PY04", "KEYSTONE HEALTH PTNRS": "PY04",
        "EVERGREEN MUTUAL INSURANCE": "PY05", "EVERGREEN MUTUAL": "PY05", "EVERGREEN MUTUAL INS": "PY05",
        "SUMMIT CHOICE PPO": "PY06", "SUMMIT PPO": "PY06", "SUMMIT CHOICE": "PY06",
        "SELF-PAY": "PY07", "SELF PAY": "PY07", "UNINSURED": "PY07",
    }
    map_rows = [{"Variant": v, "PayerID": pid, "StandardName": payers[pid]["PayerName"],
                 "Note": "Business rule: no coverage on file = Self-Pay" if v == "UNINSURED" else None}
                for v, pid in sorted(variant_map.items(), key=lambda kv: (kv[1], kv[0]))]
    pm = L.add_table_sheet("PayerMap", map_rows, columns=["Variant", "PayerID", "StandardName", "Note"], table="tblPayerMap",
                           widths={"Variant": 32, "PayerID": 10, "StandardName": 30, "Note": 44})
    # every raw spelling must map after TRIM (XLOOKUP ignores case), and to the patient's real payer
    for r in raw:
        pid = variant_map[xl_trim(r["Insurance"]).upper()]
        real = patients[truth[r["RecordID"]]]["PrimaryPayerID"]
        assert pid == real, (r["RecordID"], r["Insurance"], pid, real)

    # ------------------------------------------------------------------ data: bed-board census export (Go To Special)
    deps = data.index(data.load("departments"), "DeptID")
    facs = data.index(data.load("facilities"), "FacilityID")
    census = [c for c in data.load("daily_census")
              if c["DeptID"] in CENSUS_UNITS and CENSUS_WEEK[0] <= c["CensusDate"] <= CENSUS_WEEK[1]]
    census.sort(key=lambda c: (CENSUS_UNITS.index(c["DeptID"]), c["CensusDate"]))
    assert len(census) == 7 * len(CENSUS_UNITS)
    export_rows, filled = [], []
    prev_fac = prev_unit = None
    for c in census:
        fac = facs[c["FacilityID"]]["FacilityName"]
        unit = deps[c["DeptID"]]["DeptName"]
        filled.append((fac, unit))
        export_rows.append({"Facility": fac if fac != prev_fac else None,
                            "Unit": unit if (fac, unit) != (prev_fac, prev_unit) else None,
                            "CensusDate": c["CensusDate"], "StaffedBeds": c["StaffedBeds"],
                            "MidnightCensus": c["MidnightCensus"]})
        prev_fac, prev_unit = fac, unit
    ce = L.add_table_sheet("CensusExport", export_rows,
                           columns=["Facility", "Unit", "CensusDate", "StaffedBeds", "MidnightCensus"], as_table=False,
                           widths={"Facility": 32, "Unit": 24, "CensusDate": 12, "StaffedBeds": 12, "MidnightCensus": 15})
    icu_days = sum(c["MidnightCensus"] for c, (fa, un) in zip(census, filled) if (fa, un) == REPORT_ICU)
    first_block_only = sum(c["MidnightCensus"] for c, r in zip(census, export_rows)
                           if r["Facility"] == REPORT_ICU[0] and r["Unit"] == REPORT_ICU[1])
    n_blank_cells = sum((r["Facility"] is None) + (r["Unit"] is None) for r in export_rows)
    units_per_name = Counter(un for _, un in set(filled))
    assert units_per_name[REPORT_ICU[1]] == 3, "the ICU name must repeat across facilities so both columns matter"

    # ------------------------------------------------------------------ data: pipe-delimited lab interface feed (Text to Columns)
    enc = data.index(data.load("encounters"), "EncounterID")
    labs = [lb for lb in data.load("lab_results")
            if lb["CollectedDateTime"].year == 2025 and lb["CollectedDateTime"].month == 12 and lb["Priority"] == "STAT"
            and enc[lb["EncounterID"]]["FacilityID"] == "F03" and lb["TestCode"] in ("K", "NA", "GLU", "CREAT", "BUN")]
    labs.sort(key=lambda lb: (lb["CollectedDateTime"], lb["LabResultID"]))
    feed_header = "MRN|TestCode|Result|Units|CollectedDateTime"
    feed = []
    for lb in labs:
        mrn = patients[lb["PatientID"]]["MRN"]
        feed.append({"line": f"{mrn}|{lb['TestCode']}|{lb['ResultValue']:g}|{lb['Units']}|"
                             f"{lb['CollectedDateTime']:%m/%d/%Y %H:%M}",
                     "parts": [mrn, lb["TestCode"], lb["ResultValue"], lb["Units"], lb["CollectedDateTime"]]})
    lf = L.add_table_sheet("LabFeed", [{feed_header: x["line"]} for x in feed], columns=[feed_header], as_table=False,
                           widths={feed_header: 58})
    k_vals = [x["parts"][2] for x in feed if x["parts"][1] == "K"]
    avg_k = sum(k_vals) / len(k_vals)
    lost_zero = sum(1 for x in feed if x["parts"][0].startswith("0"))
    assert lost_zero == len(feed), "every MRN in the feed starts with 0, so General format visibly breaks them all"

    # ------------------------------------------------------------------ data: cleaning log template
    log_rows = [
        {"Step": 1, "Column": "(all)", "ProblemFound": "Export must stay exactly as received, for audit and re-runs",
         "FixApplied": "Raw copied to Clean. Every fix is a formula in a new yellow column", "RowsChanged": 0,
         "Notes": "Never type over Raw"},
        {"Step": 2, "Column": "Sex", "ProblemFound": "10 spellings of 2 values: F, f, ' F', Female, FEMALE, M, m, 'M ', Male, MALE",
         "FixApplied": None, "RowsChanged": None, "Notes": "Practice tasks 2 and 12"},   # FixApplied left for the learner (it
         # would otherwise print task 2's solution formula on a visible sheet)
        {"Step": 3, "Column": "PatientName", "ProblemFound": "Two layouts (Last, First and First Last), stray spaces, ALL CAPS and lowercase"},
        {"Step": 4, "Column": "Phone", "ProblemFound": "Six layouts, some with a +1 country code, and some blanks"},
        {"Step": 5, "Column": "MRN", "ProblemFound": "Leading zeros lost on some rows, and an MRN- prefix on others"},
        {"Step": 6, "Column": "DOB", "ProblemFound": "Dates stored as text in five layouts"},
        {"Step": 7, "Column": "Insurance", "ProblemFound": "Many spellings of 7 payers, some with trailing spaces"},
        {"Step": 8, "Column": "CityStateZip", "ProblemFound": "ZIP+4 mixed with 5-digit ZIPs"},
        {"Step": 9, "Column": "(rows)", "ProblemFound": "Exact copies and re-registrations of the same patient"},
    ]
    lg = L.add_table_sheet("CleaningLog", log_rows,
                           columns=["Step", "Column", "ProblemFound", "FixApplied", "RowsChanged", "Notes"], table="tblLog",
                           widths={"Step": 7, "Column": 14, "ProblemFound": 62, "FixApplied": 52, "RowsChanged": 14, "Notes": 22})
    sex_log_cell = f"CleaningLog!{lg.col('RowsChanged')}{lg.first_row + 1}"

    L.sheet_order = ["Start Here", "Practice", "Raw", "Clean", "PayerMap", "CensusExport", "LabFeed", "CleaningLog", "Bonus"]

    # ------------------------------------------------------------------ helpers
    F, LR = rw.first_row, rw.last_row

    def R(col: str) -> str:                       # Raw!B2:B651
        return f"Raw!{rw.col(col)}{F}:{rw.col(col)}{LR}"

    def C(col: str) -> str:                       # Clean!K2:K651
        return f"Clean!{cl.col(col)}{F}:{cl.col(col)}{LR}"

    def c1(col: str) -> str:                      # first-row cell on Clean, no sheet: K2
        return f"{cl.col(col)}{F}"

    def row_of(rid: str) -> int:
        return F + next(i for i, r in enumerate(raw) if r["RecordID"] == rid)

    # ------------------------------------------------------------------ the README guide quotes these rows; fail loudly if they move
    guide_refs = {
        ("R1001", "DOB"): "2002-03-31", ("R1002", "DOB"): "03-Apr-1957", ("R1003", "DOB"): "08/06/1989",
        ("R1019", "DOB"): "3/22/1961", ("R1029", "DOB"): "01.11.2003",
        ("R1001", "Phone"): "555.468.7891", ("R1002", "Phone"): "(555) 476-7432", ("R1006", "Phone"): "555-722-8468",
        ("R1024", "Phone"): "5553869767", ("R1019", "Phone"): "+1 555 597 3811", ("R1011", "Phone"): "",
        ("R1004", "MRN"): "MRN-05927600", ("R1010", "MRN"): "1955231", ("R1015", "MRN"): "488484",
        ("R1002", "Insurance"): "Medicare ", ("R1008", "Insurance"): "SILVERLINE MEDICARE ADV.",
        ("R1018", "Insurance"): "medicaid", ("R1012", "Insurance"): "Silverline MA",
        ("R1009", "PatientName"): "  Alexander ,  Gary ", ("R1004", "PatientName"): "MURRAY, MARGARET",
        ("R1017", "PatientName"): "Amber White", ("R1007", "PatientName"): "perry, janet",
        ("R1001", "PatientName"): "Haddad, Jonathan",
        ("R1007", "Email"): "none", ("R1011", "Email"): "N/A", ("R1030", "Email"): " carlos.santos91@example.com ",
        ("R1010", "Sex"): " F", ("R1012", "Sex"): "M ", ("R1004", "RegisteredOn"): "05/30/2022 05:42 PM",
        ("R1001", "RegisteredOn"): "2015-11-06 18:33:00", ("R1001", "CityStateZip"): "Bluestone OH 45501",
        ("R1010", "CityStateZip"): "Bluestone, OH 45501-8106", ("R1002", "CityStateZip"): "CEDAR RIDGE, OH 45720",
    }
    by_rid = {r["RecordID"]: r for r in raw}
    for (rid, col), want in guide_refs.items():
        assert by_rid[rid][col] == want, f"README guide example {rid}.{col} changed: {by_rid[rid][col]!r} != {want!r}"
    assert row_of("R1009") == 10 and row_of("R1029") == 30
    assert next(r["RecordID"] for r in raw if "-" in r["CityStateZip"]) == "R1010"   # Flash Fill steps name row 11
    a, b = by_rid[PAIR[0]], by_rid[PAIR[1]]
    assert (a["PatientName"], b["PatientName"]) == ("Amber White", "WHITE, AMBER") and truth[PAIR[0]] == truth[PAIR[1]]
    # README section 14 says Email is the ONLY raw column the pair shares character for character
    assert [c for c in RAW_COLS[1:] if a[c] == b[c]] == ["Email"], "README section 14 duplicate-pair table changed"
    assert (a["CityStateZip"], b["CityStateZip"]) == ("Cedar Ridge OH 45720", "Cedar Ridge, OH 45720")

    # ------------------------------------------------------------------ cleaning results (computed in Python)
    names = [name_clean(r["PatientName"]) for r in raw]
    sexes = [sex_clean(r["Sex"]) for r in raw]
    phones = [phone_clean(r["Phone"]) for r in raw]
    mrns = [mrn_clean(r["MRN"]) for r in raw]
    dobs = [dob_parse(r["DOB"]) for r in raw]
    regs = [reg_parse(r["RegisteredOn"]) for r in raw]
    zips = [zip5(r["CityStateZip"]) for r in raw]
    payer_ids = [variant_map[xl_trim(r["Insurance"]).upper()] for r in raw]
    pids = [truth[r["RecordID"]] for r in raw]

    # Cleaning must reproduce the master patient file (proves the formulas, not just the arithmetic)
    for i, r in enumerate(raw):
        p = patients[pids[i]]
        assert names[i] == xl_proper(f"{p['FirstName']} {p['LastName']}"), r
        assert sexes[i] == p["Sex"] and mrns[i] == p["MRN"] and dobs[i] == p["DOB"], r
        assert dobs[i] == dob_us_datevalue(r["DOB"]), r   # the short US-settings formula agrees everywhere
        assert phones[i] == ("".join(ch for ch in p["Phone"] if ch.isdigit()) if r["Phone"] else ""), r
        assert regs[i].date() == p["RegistrationDate"], r
        assert zips[i] == p["ZIP"], r

    # 1 · names with extra spaces
    extra_spaces = sum(1 for r in raw if len(r["PatientName"]) != len(xl_trim(r["PatientName"])))
    no_comma = sum(1 for r in raw if "," not in r["PatientName"])
    # 2 · SexClean
    n_f = sum(1 for s in sexes if s == "F")
    sex_spellings = Counter(r["Sex"] for r in raw)
    # 3 · exact duplicates (all columns but RecordID); Remove Duplicates and UNIQUE ignore case, so check both ways
    keys = [tuple(r[c] for c in RAW_COLS[1:]) for r in raw]
    exact_dups = n - len(set(keys))
    assert exact_dups == n - len({tuple(x.lower() for x in k) for k in keys})
    # 4 · census fill-down
    # 5 · ZIP5 distinct count; how many if ZIP+4 is left in
    distinct_zips = len(set(zips))
    distinct_zip_raw = len({r["CityStateZip"].rsplit(" ", 1)[1] for r in raw})
    zip4_rows = sum(1 for r in raw if "-" in r["CityStateZip"])
    # 6 · phones
    n_phone10 = sum(1 for p in phones if len(p) == 10 and p.isdigit())
    n_phone_blank = sum(1 for r in raw if not r["Phone"])
    n_plus1 = sum(1 for r in raw if r["Phone"].startswith("+1"))
    # 7 · MRN
    distinct_mrn = len(set(mrns))
    assert distinct_mrn == len(set(pids)) == 560
    mrn_prefix = sum(1 for r in raw if r["MRN"].startswith("MRN-"))
    mrn_short = sum(1 for r in raw if not r["MRN"].startswith("MRN-") and len(r["MRN"]) < 8)
    # 8 · payer
    n_py02 = sum(1 for p in payer_ids if p == "PY02")
    n_py02_untrimmed = sum(1 for r in raw if r["Insurance"].upper() in variant_map and variant_map[r["Insurance"].upper()] == "PY02")
    assert n_py02_untrimmed < n_py02
    insurance_spellings = len(set(r["Insurance"] for r in raw))
    # 9 · lab feed: average potassium
    # 10 · names: distinct clean names (fewer than patients: some different patients share a name)
    distinct_names = len({x.lower() for x in names})
    distinct_names_noflip = len({xl_proper(xl_trim(r["PatientName"].replace(",", ""))).lower() for r in raw})
    distinct_names_trim_only = len({xl_proper(xl_trim(r["PatientName"])).lower() for r in raw})   # =PROPER(TRIM(B2))
    assert distinct_names < distinct_names_noflip < distinct_names_trim_only
    name_twins = sorted({nm for nm in names if len({pids[i] for i, x in enumerate(names) if x == nm}) > 1})
    mc_rows = [r["RecordID"] for r, nm in zip(raw, names) if "Mcd" in nm]
    # 11 · DOB: rows whose day is 13+ (impossible if month and day were swapped)
    day13 = sum(1 for d in dobs if d.day > 12)
    day13_iso_mon = sum(1 for r, d in zip(raw, dobs) if d.day > 12 and "-" in r["DOB"])
    dob_layouts = Counter(("iso" if r["DOB"][4] == "-" else "mon" if r["DOB"][2] == "-" else "dot" if "." in r["DOB"]
                           else "slash") for r in raw)
    # 12 · log: Sex values that the cleaning changed (case-sensitive comparison)
    sex_changed = sum(1 for r, s in zip(raw, sexes) if r["Sex"] != s)

    # bonus
    evening = sum(1 for d in regs if d.hour >= 17)
    evening_if_pm_ignored = sum(1 for r, d in zip(raw, regs) if d.hour >= 17 and r["RegisteredOn"][4] == "-")
    groups: dict[str, list[int]] = defaultdict(list)
    for i, m in enumerate(mrns):
        groups[m].append(i)
    keep = [False] * n
    for m, idx in groups.items():
        latest = max(regs[i] for i in idx)
        keep[next(i for i in idx if regs[i] == latest)] = True     # first row among ties
    n_keep = sum(keep)
    assert n_keep == distinct_mrn
    kept_reg = next(regs[i] for i in groups[KEEP_MRN] if keep[i])
    first_reg = regs[groups[KEEP_MRN][0]]
    assert kept_reg == datetime(2019, 12, 21, 14, 6) and first_reg < kept_reg and len(groups[KEEP_MRN]) == 3
    assert any(raw[i]["RegisteredOn"][4] != "-" and keep[i] for i in groups[KEEP_MRN])
    not_first = sum(1 for m, idx in groups.items() for i in idx if keep[i] and i != idx[0])
    lost_phone = sum(1 for m, idx in groups.items() for i in idx
                     if keep[i] and phones[i] == "" and any(phones[j] for j in idx))
    assert 0 < lost_phone < 10
    rows_per_patient = Counter(len(v) for v in groups.values())

    # ------------------------------------------------------------------ data note
    L.data_note = (f"A {n}-row patient registration export from Bluestone's legacy registration system (names, DOBs, phones, "
                   "insurance, and MRNs typed every which way, plus duplicates), a payer mapping table, a one-week bed-board census "
                   f"export, and a {len(feed)}-line lab interface feed from Cedar Ridge Medical Center.")

    # ------------------------------------------------------------------ practice tasks
    L.practice_intro = (
        "Raw is the untouched export: never edit it. Clean is your working copy of the same 650 rows, with yellow columns to "
        "fill. For a column task, type the formula in row 2 of the yellow column. Clean is an Excel Table, so the formula fills "
        "down by itself. A gray cell here then checks the whole column. Other tasks use the CensusExport, LabFeed, and "
        "CleaningLog sheets.")
    L.start_notes = [
        "Clean's last two yellow columns (RegisteredClean and Keep) belong to the Bonus.",
        "Task 3 deletes rows, so do it on a copy of Raw (right-click the Raw tab → Move or Copy → Create a copy). "
        "Tasks 4 and 9 change the CensusExport and LabFeed sheets in place. If something goes wrong, press Ctrl + Z "
        "(Mac: ⌘ + Z) or download a fresh workbook.",
        "Version notes: XLOOKUP needs Microsoft 365 or Excel 2021+ (use VLOOKUP with FALSE in older versions). MAXIFS needs "
        "Excel 2019+. Flash Fill needs Excel 2013+ on Windows, or Excel 2019 / Microsoft 365 on a Mac.",
    ]

    sex_formula = f"=UPPER(LEFT(TRIM(D{F}),1))"
    name_formula = (f'=PROPER(IF(ISNUMBER(FIND(",",B{F})),TRIM(MID(B{F},FIND(",",B{F})+1,LEN(B{F})))&" "&'
                    f'TRIM(LEFT(B{F},FIND(",",B{F})-1)),TRIM(B{F})))')
    phone_formula = (f'=RIGHT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(E{F},"(",""),")",""),"-",""),".","")," ",""),10)')
    mrn_formula = f'=TEXT(VALUE(SUBSTITUTE(I{F},"MRN-","")),"00000000")'
    map_var = f"PayerMap!${pm.col('Variant')}${pm.first_row}:${pm.col('Variant')}${pm.last_row}"
    map_id = f"PayerMap!${pm.col('PayerID')}${pm.first_row}:${pm.col('PayerID')}${pm.last_row}"
    payer_formula = f'=XLOOKUP(TRIM(H{F}),{map_var},{map_id},"UNMAPPED")'          # A1 form for the self-test fill
    payer_solution = f'=XLOOKUP(TRIM(H{F}),tblPayerMap[Variant],tblPayerMap[PayerID],"UNMAPPED")'
    dob_formula = (f'=IF(MID(C{F},5,1)="-",DATE(LEFT(C{F},4),MID(C{F},6,2),RIGHT(C{F},2)),'
                   f'IF(MID(C{F},3,1)="-",DATE(RIGHT(C{F},4),(SEARCH(MID(C{F},4,3),"{MONTHS}")+2)/3,LEFT(C{F},2)),'
                   f'DATE(RIGHT(C{F},4),LEFT(C{F},FIND("/",SUBSTITUTE(C{F},".","/"))-1),'
                   f'MID(C{F},FIND("/",SUBSTITUTE(C{F},".","/"))+1,LEN(C{F})-FIND("/",SUBSTITUTE(C{F},".","/"))-5))))')
    reg_formula = (f'=IF(MID(J{F},5,1)="-",DATE(LEFT(J{F},4),MID(J{F},6,2),MID(J{F},9,2)),'
                   f'DATE(MID(J{F},7,4),LEFT(J{F},2),MID(J{F},4,2)))+TIMEVALUE(MID(J{F},12,8))')
    Q, Rg = cl.col("MRNClean"), cl.col("RegisteredClean")
    keep_formula = (f"=AND({Rg}{F}=MAXIFS(${Rg}${F}:${Rg}${LR},${Q}${F}:${Q}${LR},{Q}{F}),"
                    f"COUNTIFS(${Q}${F}:{Q}{F},{Q}{F},${Rg}${F}:{Rg}{F},{Rg}{F})=1)")

    def fill(col: str, formula: str | None = None, values: list | None = None) -> dict:
        d = {"range": f"Clean!{cl.col(col)}{F}:{cl.col(col)}{LR}"}
        if values is not None:
            d["values"] = values
        else:
            d["formula"] = formula
        return d

    raw_mrn_text = f'TEXT(VALUE(SUBSTITUTE({R("MRN")},"MRN-","")),"00000000")'
    # Region-proof array versions of the DOB and RegisteredOn formulas, run on Raw, for the key's Live result column
    # (DATEVALUE and --text follow the computer's regional settings, so they would show #VALUE! in a day-first region).
    raw_dob_dates = dob_formula[1:].replace(f"C{F}", R("DOB"))
    raw_reg_times = reg_formula[1:].replace(f"J{F}", R("RegisteredOn"))
    assert raw_dob_dates.count("Raw!") == dob_formula.count(f"C{F}") > 0
    assert raw_reg_times.count("Raw!") == reg_formula.count(f"J{F}") > 0
    raw_phone = (f'RIGHT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({R("Phone")},"(",""),")",""),"-",""),".",""),'
                 f'" ",""),10)')
    ce_rng = lambda col: f"CensusExport!{ce.col(col)}{ce.first_row}:{ce.col(col)}{ce.last_row}"  # noqa: E731
    lf_rng = lambda col: f"LabFeed!{col}{lf.first_row}:{col}{lf.last_row}"  # noqa: E731
    feed_rng = lf_rng("A")
    p1 = f'FIND("|",{feed_rng})'
    p2 = f'FIND("|",{feed_rng},{p1}+1)'
    p3 = f'FIND("|",{feed_rng},{p2}+1)'

    L.tasks = [
        # ---------------------------------------------------------------- profile
        Task(f"Profile first. On the Raw sheet, how many PatientName values contain extra spaces (leading, trailing, or "
             f"doubled)? Compare each name's length with the length of its TRIMmed version.",
             answer=extra_spaces, title="Names with extra spaces (profiling)",
             solution=f"=SUMPRODUCT(--(LEN({R('PatientName')})<>LEN(TRIM({R('PatientName')}))))",
             hint="LEN(x)<>LEN(TRIM(x)) is TRUE when TRIM would remove something. SUMPRODUCT(--(…)) counts the TRUEs",
             explanation="TRIM removes leading and trailing spaces and shrinks inner runs to one space, so any name it shortens "
                         "had extra spaces. Comparing lengths tests all 650 rows at once: the comparison gives TRUE/FALSE for "
                         "each row, the double minus turns those into 1/0, and SUMPRODUCT adds them. You can't see a trailing "
                         "space by looking at a cell, so a count like this is how you find out a problem exists before you fix "
                         f"it. In Microsoft 365 or Excel 2021+, `=SUM(--(LEN(…)<>LEN(TRIM(…))))` works too."),
        # ---------------------------------------------------------------- standardize a category with a formula
        Task(f"On the Clean sheet, fill the yellow SexClean column ({cl.col('SexClean')}) with a single capital letter, F or M. "
             f"The raw Sex column has {len(sex_spellings)} spellings, including Female, MALE, lowercase f, and values with a "
             f"stray space. The gray cell counts cells that are exactly F (case-sensitive).",
             answer=n_f, title="SexClean column (rows that are exactly F)",
             solution=sex_formula,
             summary=f'=IF(COUNTA({C("SexClean")})=0,"",SUMPRODUCT(--EXACT({C("SexClean")},"F")))',
             fill=fill("SexClean", sex_formula),
             live=f'=SUMPRODUCT(--EXACT(UPPER(LEFT(TRIM({R("Sex")}),1)),"F"))',
             hint="TRIM first, then take the first letter, then capitalize it",
             explanation="Every spelling starts with the right letter once the stray space is gone, so the fix is TRIM, then "
                         "`LEFT(…,1)`, then UPPER. The order matters: `LEFT(\" F\",1)` is a space. The check uses EXACT because COUNTIF ignores "
                         "case and would count a lowercase f as F. This shortcut works only because each value is identified by its "
                         "first letter. Medicare and Medicaid share a first letter, so Insurance needs a mapping table instead "
                         "(task 8)."),
        # ---------------------------------------------------------------- Remove Duplicates (on a copy)
        Task("Make a copy of the Raw sheet (right-click its tab → **Move or Copy…** → tick **Create a copy**). On the copy, run "
             f"**Data → Remove Duplicates** on all the data (A1:{rw.col('RegisteredOn')}{LR}) with every column ticked except "
             "RecordID (each row has its own "
             "RecordID, so leaving it ticked finds nothing). How many duplicate rows does Excel remove?",
             answer=exact_dups, title="Remove Duplicates on the raw columns",
             solution=("1. Right-click the **Raw** tab → **Move or Copy…** → tick **Create a copy** → **OK**. Work on "
                       "**Raw (2)**.\n"
                       "2. Click any cell in the data, then **Data → Remove Duplicates**.\n"
                       "3. Keep **My data has headers** ticked. Untick **RecordID** and leave the other nine columns ticked. "
                       "Click **OK**.\n"
                       f"4. Excel reports *{exact_dups} duplicate values found and removed; {n - exact_dups} unique values "
                       "remain.*\n\n"
                       f"Formula check without deleting anything (Microsoft 365 or Excel 2021+): "
                       f"`=ROWS(Raw!B{F}:J{LR})-ROWS(UNIQUE(Raw!B{F}:J{LR}))`"),
             live=f"=ROWS(Raw!B{F}:J{LR})-ROWS(UNIQUE(Raw!B{F}:J{LR}))",
             hint="Untick RecordID in the Remove Duplicates dialog. Excel's message tells you the count",
             explanation=f"Remove Duplicates only removes rows that match in **every** ticked column, so it finds the {exact_dups} "
                         f"rows that were exported twice, character for character. That is not all of them, as task 7 will show "
                         "when you count the real patients. Many more rows are re-registrations typed differently. For example, row "
                         f"{row_of(PAIR[0])} (\"{by_rid[PAIR[0]]['PatientName']}\", {by_rid[PAIR[0]]['Phone']}) and row "
                         f"{row_of(PAIR[1])} (\"{by_rid[PAIR[1]]['PatientName']}\", {by_rid[PAIR[1]]['Phone']}) are the same "
                         "patient, but Remove Duplicates sees two different rows. "
                         "Standardize first, then remove duplicates on the cleaned key column. UNIQUE on all nine columns "
                         "counts the same thing without deleting anything. Both ignore case."),
        # ---------------------------------------------------------------- Go To Special: blanks, fill down
        Task(f"Switch to CensusExport, a bed-board report that prints each Facility and Unit only on the first row of its block. "
             f"Fill every blank cell in {ce.col('Facility')}{ce.first_row}:{ce.col('Unit')}{ce.last_row} with the value above it: "
             f"select that range, use **Go To Special → Blanks**, type = and press the Up arrow, then press Ctrl + Enter "
             f"(Mac: ⌘ + Return). The gray cell stays blank until every gap is filled, then totals the week's midnight census "
             f"for the Intensive Care Unit at Cedar Ridge Medical Center (ICU patient days).",
             answer=icu_days, title="Go To Special → Blanks fill-down (Cedar Ridge ICU patient days)",
             solution=(f"1. Select **{ce.col('Facility')}{ce.first_row}:{ce.col('Unit')}{ce.last_row}** on CensusExport (not "
                       "the whole columns, and not the header row).\n"
                       "2. **Home → Find & Select → Go To Special…** (or press **F5**, Mac: **Control + G**, then **Special…**). "
                       "Choose **Blanks** → **OK**. Only the empty cells stay selected.\n"
                       f"3. Without clicking anywhere, type **=** and press **↑**. The active cell is the first "
                       f"blank, {ce.col('Facility')}{ce.first_row + 1}, so the formula reads "
                       f"`={ce.col('Facility')}{ce.first_row}`.\n"
                       "4. Press **Ctrl + Enter** (Mac: **⌘ + Return**) to put that formula in every selected blank. Each copy "
                       "points to the cell above itself.\n"
                       f"5. Turn the formulas into values: select {ce.col('Facility')}{ce.first_row}:{ce.col('Unit')}"
                       f"{ce.last_row}, copy, then **Paste Special → Values**.\n\n"
                       f"Check with `=SUMIFS({ce_rng('MidnightCensus')},{ce_rng('Facility')},\"{REPORT_ICU[0]}\","
                       f"{ce_rng('Unit')},\"{REPORT_ICU[1]}\")`"),
             summary=(f'=IF(COUNTBLANK(CensusExport!{ce.col("Facility")}{ce.first_row}:{ce.col("Unit")}{ce.last_row})>0,"",'
                      f'SUMIFS({ce_rng("MidnightCensus")},{ce_rng("Facility")},"{REPORT_ICU[0]}",{ce_rng("Unit")},"{REPORT_ICU[1]}"))'),
             fill={"range": f"CensusExport!{ce.col('Facility')}{ce.first_row}:{ce.col('Unit')}{ce.last_row}",
                   "values": [v for pair in filled for v in pair]},
             live=False,
             hint="Go To Special selects only the blanks, and Ctrl + Enter fills them all with one relative formula",
             explanation=f"After Go To Special, Excel has selected only the {n_blank_cells} empty cells, and the formula you type "
                         "goes into all of them at once. Because the reference is relative, each blank points to the cell "
                         "directly above it, and that cell either holds the label or points further up. The unit name "
                         f"\"{REPORT_ICU[1]}\" appears at all three hospitals, so the SUMIFS needs both columns filled. With "
                         f"the blanks left in, it returns {first_block_only}: no Cedar Ridge ICU row carries both labels, because "
                         "the facility name was printed only once, on the facility's first unit. Paste the result as values "
                         "before you sort, because a formula that points "
                         "\"one row up\" points somewhere else after a sort."),
        # ---------------------------------------------------------------- Flash Fill ZIP5
        Task(f"Back on Clean, use Flash Fill to fill the yellow ZIP5 column ({cl.col('ZIP5')}) with the 5-digit ZIP from "
             f"CityStateZip. {zip4_rows} rows carry a ZIP+4 such as 45501-8106, and those must become 45501. Type three "
             f"examples yourself: rows {F} and {F + 1}, plus the first ZIP+4 row (row {row_of('R1010')}). Then select the "
             f"first empty cell and press Ctrl + E (Mac: **Data → Flash Fill**). The gray cell counts the distinct ZIP codes "
             f"in your column.",
             answer=distinct_zips, title="Flash Fill ZIP5 (distinct ZIP codes)",
             solution=(f"1. In **{c1('ZIP5')}** type **{zips[0]}** (from *{raw[0]['CityStateZip']}*), and in "
                       f"**{cl.col('ZIP5')}{F + 1}** type **{zips[1]}** (from *{raw[1]['CityStateZip']}*).\n"
                       f"2. In the first ZIP+4 row, **{cl.col('ZIP5')}{row_of('R1010')}** "
                       f"(*{by_rid['R1010']['CityStateZip']}*), type **{zip5(by_rid['R1010']['CityStateZip'])}**.\n"
                       f"3. Select **{cl.col('ZIP5')}{F + 2}**, the first empty cell, and press **Ctrl + E** (or "
                       "**Data → Flash Fill**).\n"
                       f"4. Check the result: `=SUMPRODUCT(--(LEN({C('ZIP5')})<>5))` should return 0. If some rows are "
                       "wrong, press Ctrl + Z (Mac: ⌘ + Z), type the correct ZIP on one of the wrong rows as an extra "
                       "example, and run Flash Fill again.\n\n"
                       f"Formula alternative (Microsoft 365 or Excel 2024): `=LEFT(TEXTAFTER({c1('CityStateZip')},\" \",-1),5)`"),
             summary=f'=IF(COUNTA({C("ZIP5")})=0,"",SUMPRODUCT(({C("ZIP5")}<>"")/COUNTIF({C("ZIP5")},{C("ZIP5")}&"")))',
             fill=fill("ZIP5", values=zips),
             live=f'=ROWS(UNIQUE(MID({R("CityStateZip")},FIND(" OH ",{R("CityStateZip")})+4,5)))',
             hint="Give Flash Fill an example from each layout, especially a ZIP+4 row",
             explanation=f"Flash Fill looks for a rule that explains every example you typed. Two plain ZIPs fit several rules "
                         "(the first number, the last number, the last five characters), and on 45501-8106 those give different "
                         "answers. The ZIP+4 example leaves only \"the first number\". If the +4 survives, the distinct count "
                         f"jumps to {distinct_zip_raw}. Flash Fill writes typed-in values, not formulas, so it won't update when "
                         "next month's export arrives. The Lesson 2.2 bonus did the same job with TEXTAFTER. The formula is the "
                         "repeatable choice, and Flash Fill is the fast one for a one-off."),
        # ---------------------------------------------------------------- SUBSTITUTE phones
        Task(f"Fill the yellow PhoneClean column ({cl.col('PhoneClean')}) with each phone as exactly 10 digits and nothing else, "
             f"or an empty result when Phone is blank. Phones arrive as (555) 476-7432, 555-722-8468, 555.468.7891, "
             f"5553869767, and +1 555 597 3811. The gray cell counts rows that end up with exactly 10 digits.",
             answer=n_phone10, title="PhoneClean column (rows with exactly 10 digits)",
             solution=phone_formula,
             summary=(f'=IF(COUNTA({C("PhoneClean")})=0,"",SUMPRODUCT((LEN({C("PhoneClean")})=10)'
                      f'*ISNUMBER(--{C("PhoneClean")})))'),
             fill=fill("PhoneClean", phone_formula),
             live=f"=SUMPRODUCT((LEN({raw_phone})=10)*ISNUMBER(--{raw_phone}))",
             hint="Nest one SUBSTITUTE per unwanted character, then keep the RIGHT 10 characters",
             explanation=f"Each SUBSTITUTE deletes one kind of character: parentheses, dashes, dots, and spaces. That leaves "
                         f"5554767432 for most rows, but +15555973811 for the {n_plus1} numbers with a country code, so "
                         "`RIGHT(…,10)` keeps the last ten digits, which is always the U.S. number. A blank phone needs no "
                         "special handling, because SUBSTITUTE and RIGHT on an empty cell return an empty result. The "
                         f"{n_phone_blank} blanks are why the count is {n_phone10}, not {n}. Keep phones as text, because they are "
                         "identifiers, not quantities. In Microsoft 365, `=RIGHT(REGEXREPLACE(E2,\"[^0-9]\",\"\"),10)` removes "
                         "every non-digit in one step."),
        # ---------------------------------------------------------------- MRN: numbers stored as text, leading zeros
        Task(f"Fill the yellow MRNClean column ({cl.col('MRNClean')}) with every MRN as 8-digit text: 05927600, not "
             f"MRN-05927600, and 01955231, not 1955231. The gray cell first checks that every value is 8 characters long, "
             f"then counts the distinct MRNs, which is the number of real patients in the file.",
             answer=distinct_mrn, title="MRNClean column (distinct patients)",
             solution=mrn_formula,
             summary=(f'=IF(COUNTA({C("MRNClean")})=0,"",IF(SUMPRODUCT(--(LEN({C("MRNClean")})<>8))>0,"MRN not 8 chars",'
                      f'SUMPRODUCT(({C("MRNClean")}<>"")/COUNTIF({C("MRNClean")},{C("MRNClean")}&""))))'),
             fill=fill("MRNClean", mrn_formula),
             live=f"=ROWS(UNIQUE({raw_mrn_text}))",
             hint="Remove the prefix with SUBSTITUTE, make it a number with VALUE, then pad it back with TEXT(…,\"00000000\")",
             explanation=f"Two problems hide in this column. {mrn_prefix} rows carry an MRN- prefix, which SUBSTITUTE removes. "
                         f"{mrn_short} rows lost their leading zeros somewhere upstream, because a system treated the MRN as a "
                         "number. VALUE turns the remaining text into a number, and `TEXT(…,\"00000000\")` writes it back as text "
                         "with exactly 8 digits, adding the zeros it needs. `=RIGHT(\"0000000\"&SUBSTITUTE(I2,\"MRN-\",\"\"),8)` "
                         f"does the same without VALUE. The result: {n} rows but only {distinct_mrn} patients, so "
                         f"{n - distinct_mrn} rows are duplicates."),
        # ---------------------------------------------------------------- mapping table
        Task(f"Fill the yellow PayerID column ({cl.col('PayerID')}) by looking up each Insurance value in the PayerMap "
             f"table, so all {insurance_spellings} spellings become one of seven PayerIDs. Some values carry a trailing "
             f"space. Return UNMAPPED for anything missing from the map. The gray cell shows a warning if any row is "
             f"unmapped. Otherwise it counts PY02 (Silverline Medicare Advantage) rows.",
             answer=n_py02, title="PayerID column via mapping table (PY02 rows)",
             solution=payer_solution,
             summary=(f'=IF(COUNTA({C("PayerID")})=0,"",IF(COUNTIF({C("PayerID")},"PY*")<ROWS({C("PayerID")}),'
                      f'"Unmapped rows",COUNTIF({C("PayerID")},"PY02")))'),
             fill=fill("PayerID", payer_formula),
             live=f'=SUMPRODUCT(--(XLOOKUP(TRIM({R("Insurance")}),{map_var},{map_id},"UNMAPPED")="PY02"))',
             hint="XLOOKUP the TRIMmed Insurance value in tblPayerMap, and use XLOOKUP's if_not_found argument for UNMAPPED",
             explanation=f"A mapping table turns a messy category into a standard code with one lookup, and it documents your "
                         "decisions where everyone can see them, such as the rule that Uninsured means PY07 (Self-Pay). XLOOKUP ignores case, so the "
                         "map needs only one row per spelling, not one per capitalization. It does not ignore spaces, so TRIM "
                         f"comes first. Without TRIM, the {n_py02 - n_py02_untrimmed} rows typed \"Silverline Medicare "
                         "Advantage \" with a trailing space come back UNMAPPED. The \"UNMAPPED\" result makes new spellings "
                         "easy to find next month: filter for it, add the spelling as a new row of the map, and every formula picks it "
                         "up. The Table references (`tblPayerMap[Variant]`) grow with the map. A fixed range such as "
                         f"`{map_var}` works today but would miss a row added below it. Without XLOOKUP, use "
                         "`=IFNA(VLOOKUP(TRIM(H2),tblPayerMap,2,FALSE),\"UNMAPPED\")`."),
        # ---------------------------------------------------------------- Text to Columns
        Task(f"Switch to LabFeed: {len(feed)} STAT results from a lab interface, each crammed into one cell as "
             f"MRN|TestCode|Result|Units|CollectedDateTime. Split A1:A{lf.last_row} into five columns with **Data → Text to "
             f"Columns** (Delimited, Other: |). In step 3 of the wizard, set the MRN column's format to Text so its leading "
             f"zeros survive. The gray cell checks the MRNs, then shows the average potassium (TestCode K) result in mmol/L.",
             answer=avg_k, fmt="0.00", title="Text to Columns on the lab feed (average potassium)",
             solution=(f"1. On **LabFeed**, select **A1:A{lf.last_row}** (header included).\n"
                       "2. **Data → Text to Columns**. Choose **Delimited** → **Next**.\n"
                       "3. Untick **Tab**, tick **Other**, and type a vertical bar **|** in the box. The preview splits into "
                       "five columns. Click **Next**.\n"
                       "4. Click the **MRN** column in the preview and choose **Text** as its column data format. Optionally "
                       "click the last column and choose **Date: MDY**. Click **Finish**.\n\n"
                       f"Check with `=AVERAGEIF({lf_rng('B')},\"K\",{lf_rng('C')})`"),
             summary=(f'=IF(COUNTA({lf_rng("C")})=0,"",IF(SUMPRODUCT(--(LEN({lf_rng("A")})<>8))>0,"MRN lost zeros",'
                      f'AVERAGEIF({lf_rng("B")},"K",{lf_rng("C")})))'),
             fill={"range": f"LabFeed!A{lf.first_row}:E{lf.last_row}", "values": [v for x in feed for v in x["parts"]]},
             live=(f'=AVERAGE(IF(MID({feed_rng},{p1}+1,{p2}-{p1}-1)="K",--MID({feed_rng},{p2}+1,{p3}-{p2}-1)))'),
             hint="Text to Columns step 3: click the first column in the preview, then choose Text",
             explanation=f"Text to Columns splits each cell at every |, and step 3 decides what each piece becomes. Under "
                         f"**General**, Excel turns {feed[0]['parts'][0]} (row 2) into the number {int(feed[0]['parts'][0])}. All "
                         f"{len(feed)} MRNs in this feed start with 0, so every one would break and stop matching the "
                         "registration file. Choosing **Text** keeps the digits exactly as sent. The results (for example "
                         f"{feed[0]['parts'][2]:g} and {feed[1]['parts'][2]:g} in rows 2 and 3) should stay General so they "
                         "become real numbers you can average. If you get it wrong, press Ctrl + Z (Mac: ⌘ + Z) to undo the "
                         "split and run it again."),
        # ---------------------------------------------------------------- names: two layouts
        Task(f"Fill the yellow NameClean column ({cl.col('NameClean')}) with every name as First Last in Proper Case, with "
             f"single spaces. For example, `\"Haddad, Jonathan\"` becomes Jonathan Haddad, `\"  Alexander ,  Gary \"` becomes "
             f"Gary Alexander, and `\"Amber White\"` stays Amber White. {no_comma} rows are already First Last (no comma). The gray cell warns if "
             f"any name still has extra spaces or is ALL CAPS or all lowercase. Otherwise it counts distinct names.",
             answer=distinct_names, title="NameClean column (distinct names)",
             solution=name_formula,
             summary=(f'=IF(COUNTA({C("NameClean")})=0,"",IF(SUMPRODUCT(EXACT({C("NameClean")},UPPER({C("NameClean")}))'
                      f'+EXACT({C("NameClean")},LOWER({C("NameClean")}))+NOT(EXACT({C("NameClean")},TRIM({C("NameClean")}))))>0,'
                      f'"Fix spaces/case",SUMPRODUCT(({C("NameClean")}<>"")/COUNTIF({C("NameClean")},{C("NameClean")}&""))))'),
             fill=fill("NameClean", name_formula),
             live=False,   # LibreOffice mis-evaluates the array form of this formula; the self-test covers it
             hint="IF(ISNUMBER(FIND(\",\",B2)), flip the two parts, just TRIM), then wrap everything in PROPER",
             explanation=f"`FIND(\",\",B2)` returns a position when there is a comma and #VALUE! when there isn't, so "
                         "`ISNUMBER(FIND(…))` tells the two layouts apart. For Last, First rows, MID takes everything after the "
                         "comma (the first name) and LEFT takes everything before it (the last name). TRIM each piece before "
                         "joining, because the spaces sit around the comma. PROPER on the outside fixes the case of both "
                         f"layouts at once. Two checks on the result. First, the flip matters: `=PROPER(TRIM(B2))` alone gives "
                         f"{distinct_names_trim_only} distinct names, and deleting the comma without flipping still gives "
                         f"{distinct_names_noflip}, because the same patient then appears as both \"Amber White\" (row "
                         f"{row_of(PAIR[0])}) and \"White Amber\" (row {row_of(PAIR[1])}). Second, the count is {distinct_names}, "
                         f"not {distinct_mrn}, because "
                         f"{len(name_twins)} names belong to two different patients ({', '.join(name_twins)}). That's why "
                         "you deduplicate on MRN, never on name. PROPER also turns McDonald into Mcdonald, so fix known exceptions "
                         "by hand and note them in the cleaning log."),
        # ---------------------------------------------------------------- text dates
        Task(f"Fill the yellow DOBClean column ({cl.col('DOBClean')}) with real dates. DOB is text in five layouts: "
             f"2002-03-31, 03-Apr-1957, 08/06/1989, 3/22/1961, and 01.11.2003. The export is from a U.S. system, so the "
             f"slash and dot layouts are month first. The gray cell checks that every row is a real date, then counts DOBs "
             f"that fall on the 13th or later of their month (a quick test that month and day weren't swapped).",
             answer=day13, title="DOBClean column (DOBs on day 13 or later)",
             solution=dob_formula,
             summary=(f'=IF(COUNTA({C("DOBClean")})=0,"",IF(COUNT({C("DOBClean")})<ROWS({C("DOBClean")}),"Not all dates yet",'
                      f'SUMPRODUCT(--(DAY({C("DOBClean")})>12))))'),
             fill=fill("DOBClean", dob_formula),
             live=f'=SUMPRODUCT(--(DAY({raw_dob_dates})>12))',
             hint="DATEVALUE can't read 01.11.2003 until the dots become slashes. For a version that works with any "
                  "regional setting, build DATE(year, month, day) from the pieces",
             explanation="This solution works on any computer: it reads the layout and builds DATE(year, month, day) from the "
                         "pieces. ISO dates have a dash in position 5. In 03-Apr-1957 the dash is in position 3, and the month comes "
                         f"from where Apr sits in \"{MONTHS}\" (position 10, and (10+2)/3 = 4). Everything else is month/day/year "
                         "with a / or . after the month. If your Windows or Mac region is United States, the much shorter "
                         "`=DATEVALUE(SUBSTITUTE(C2,\".\",\"/\"))` gives identical results, because U.S. settings read the ISO "
                         "and 03-Apr-1957 layouts and read slashes month first. In a day-first region it reads "
                         "08/06/1989 as 8 June and fails on 3/22/1961, and in a language whose month names differ from English "
                         "it can fail on 03-Apr-1957 too. The day-13 test catches a swap: a month number is "
                         "never above 12, so if the slash and dot rows had month and day swapped, none of them could show a "
                         f"day above 12 and the count would fall to {day13_iso_mon} instead of {day13}."),
        # ---------------------------------------------------------------- document
        Task(f"Document your work. On the CleaningLog sheet, put a formula in the yellow RowsChanged cell for step 2 "
             f"({sex_log_cell.split('!')[1]}) that counts how many rows your SexClean column actually changed: rows where the "
             f"raw Sex value is not exactly the same as SexClean (case-sensitive). The gray cell here reads your log entry.",
             answer=sex_changed, title="Cleaning log: rows changed by the Sex step",
             solution=f"=SUMPRODUCT(--NOT(EXACT({C('Sex')},{C('SexClean')})))",
             summary=f'=IF({sex_log_cell}="","",{sex_log_cell})',
             fill={"range": f"{sex_log_cell}:{sex_log_cell.split('!')[1]}", "formula": f"=SUMPRODUCT(--NOT(EXACT({C('Sex')},{C('SexClean')})))"},
             live=f'=SUMPRODUCT(--NOT(EXACT({R("Sex")},UPPER(LEFT(TRIM({R("Sex")}),1)))))',
             hint="EXACT compares case-sensitively, NOT flips TRUE and FALSE, and SUMPRODUCT(--…) counts the TRUEs",
             explanation=f"A cleaning log records what you changed and how much, so the next person (or you, next month) can "
                         f"repeat it and audit it. \"Rows changed\" is the most useful number in it: here {sex_changed} of {n} "
                         f"rows changed and {n - sex_changed} were already F or M. EXACT matters again, because with `=` the "
                         "values \"f\" and \"F\" count as equal and those rows would look unchanged. A count that is "
                         "suspiciously high or low is often the first sign that a cleaning formula is wrong."),
    ]

    # ------------------------------------------------------------------ bonus: master patient index
    L.bonus_title = "Bonus: Build the master patient list for the EHR migration"
    L.bonus_scenario = (
        "Bluestone is moving to a new EHR, and the vendor needs a master patient load file with exactly one row per patient. "
        "The rule from Patient Access is: keep each patient's most recent registration (latest RegisteredOn). Fill the Clean "
        "sheet's last two yellow columns (RegisteredClean and Keep), using your MRNClean and PhoneClean columns from the "
        "practice tasks. RegisteredOn has two "
        "layouts: 2015-11-06 18:33:00 (24-hour clock) and 05/30/2022 05:42 PM (12-hour clock). If you want helper columns, "
        "put them on the Clean sheet in column V or further right, outside the Table.")
    L.bonus = [
        Task(f"Fill the yellow RegisteredClean column ({cl.col('RegisteredClean')}) with real date-times. Make sure 05:42 PM "
             f"becomes 17:42. The gray cell checks every row, then counts registrations made at 5:00 PM or later (the evening "
             f"Patient Access shift).",
             answer=evening, title="RegisteredClean column (evening registrations)",
             solution=reg_formula,
             summary=(f'=IF(COUNTA({C("RegisteredClean")})=0,"",IF(COUNT({C("RegisteredClean")})<ROWS({C("RegisteredClean")}),'
                      f'"Not all date-times yet",SUMPRODUCT(--(HOUR({C("RegisteredClean")})>=17))))'),
             fill=fill("RegisteredClean", reg_formula),
             live=f"=SUMPRODUCT(--(HOUR({raw_reg_times})>=17))",
             hint="In both layouts the time starts at character 12, and TIMEVALUE understands both 18:33:00 and 05:42 PM. "
                  "Build the date part with DATE, then add the time to it",
             explanation="Both layouts put the time at character 12, so `TIMEVALUE(MID(J2,12,8))` reads \"18:33:00\" and "
                         "\"05:42 PM\" alike, and AM/PM is handled for you. The date part differs, so the IF builds it with DATE "
                         "from the right pieces. In the ISO layout the day is `MID(J2,9,2)`, not `RIGHT(J2,2)`, because the time "
                         "follows it. A date-time is just date + time (a whole number plus a fraction of a day). "
                         f"If PM were ignored, the evening count would drop to {evening_if_pm_ignored}. With U.S. regional "
                         "settings, `=--J2` (or `=VALUE(J2)`) converts both layouts in one step."),
        Task(f"Fill the yellow Keep column ({cl.col('Keep')}) with TRUE on exactly one row per patient (MRNClean): the row "
             f"with that patient's latest RegisteredClean. If the latest time appears on two identical rows, keep only the "
             f"first of them. Every other row is FALSE. The gray cell counts the TRUEs.",
             answer=n_keep, title="Keep column (rows in the master list)",
             solution=keep_formula,
             summary=f'=IF(COUNTA({C("Keep")})=0,"",COUNTIF({C("Keep")},TRUE))',
             fill=fill("Keep", keep_formula),
             live=f"=ROWS(UNIQUE({raw_mrn_text}))",
             hint="MAXIFS (Lesson 2.5) finds the patient's latest time. A COUNTIFS over an expanding range "
                  "(`$Q$2:Q2`) breaks ties",
             explanation="The first test, RegisteredClean = MAXIFS(all RegisteredClean, all MRNClean, this MRN), is TRUE on the "
                         "patient's latest row. Exact copies share that time, so the second test counts how many rows so far "
                         "(the expanding range `$Q$2:Q2` grows as the formula goes down) have this MRN and this time. It "
                         f"equals 1 only on the first of them. The total must equal the {distinct_mrn} distinct MRNs from practice task 7, which "
                         "is a good cross-check. Avoid comparing date-times with `\">\"&R2` inside COUNTIFS, because that turns "
                         "the time into text with 15 digits and can miss by a rounding error."),
        Task(f"Which RegisteredOn did you keep for MRN {KEEP_MRN}? Enter it as a date and time.",
             answer=kept_reg, fmt="mm/dd/yyyy hh:mm", title=f"Registration kept for MRN {KEEP_MRN}",
             solution=f'=MAXIFS({C("RegisteredClean")},{C("MRNClean")},"{KEEP_MRN}",{C("Keep")},TRUE)',
             live=(f'=MAX(IF({raw_mrn_text}="{KEEP_MRN}",{raw_reg_times}))'),
             hint="MAXIFS with two conditions: MRNClean is this MRN, and Keep is TRUE. Or filter Clean on those two "
                  "columns and read RegisteredClean",
             explanation=f"This patient has three rows. Two of them say {by_rid[raw[groups[KEEP_MRN][0]]['RecordID']]['MRN']} (the "
                         f"MRN lost its leading zero) with the time {first_reg:%H:%M}, and the third says {KEEP_MRN} with "
                         f"{raw[next(i for i in groups[KEEP_MRN] if keep[i])]['RegisteredOn']}. You only find all three after "
                         f"cleaning the MRN, and you only pick the right one if 02:06 PM became 14:06: read as 02:06 AM it "
                         f"would look older than {first_reg:%H:%M}."),
        Task("Remove Duplicates always keeps the first row it meets. For how many patients is the row you kept NOT that "
             "patient's first row in the file? These are the patients that Remove Duplicates on MRNClean would have gotten wrong.",
             answer=not_first, title="Kept rows that are not the first occurrence",
             solution=(f"=SUMPRODUCT({C('Keep')}*(MATCH({C('MRNClean')},{C('MRNClean')},0)"
                       f"<>ROW({C('MRNClean')})-{F - 1}))"),
             live=False,
             hint="A row is a patient's first occurrence when MATCH(its MRN, the MRN column, 0) returns its own position. "
                  "Or add a helper column with `=COUNTIF($Q$2:Q2,Q2)=1`",
             explanation="MATCH(MRN, MRN column, 0) returns the position of the first row with that MRN. `ROW(…)-1` is each "
                         "row's own position in the column (row 2 is position 1). Where they differ, the row is a repeat. "
                         "Multiplying by Keep counts kept rows that are repeats. "
                         f"For these {not_first} patients a newer registration sits further down the file. To make Remove "
                         "Duplicates keep the newest row, sort by RegisteredClean (newest to oldest) first, then remove "
                         "duplicates on MRNClean: the first row it meets is then the newest."),
        Task("The rule \"keep the latest row\" decides which registration survives into the master list, and it can throw "
             "away good data. How many kept rows have an empty PhoneClean even though another registration for the same "
             "patient has a phone number?",
             answer=lost_phone, title="Kept rows that lose a phone number",
             solution=(f'=SUMPRODUCT({C("Keep")}*({C("PhoneClean")}="")'
                       f'*(COUNTIFS({C("MRNClean")},{C("MRNClean")},{C("PhoneClean")},"?*")>0))'),
             live=False,
             hint="For each row, COUNTIFS(MRN column, this MRN, PhoneClean column, `\"?*\"`) counts that patient's rows "
                  "that have a phone",
             explanation="When its criteria argument is a range instead of a single value, COUNTIFS returns one count per row: "
                         "how many rows share this row's MRN and have a phone. The criteria `\"?*\"` means at least one "
                         "character, which works because PhoneClean is text. Multiply by Keep and by an empty PhoneClean "
                         f"to find the {lost_phone} patients who would arrive in the new EHR with no phone, although an older "
                         "registration has one. Real master patient index (MPI) loads prevent this with **survivorship rules**, "
                         "which decide field by field which value survives. A common rule takes each field from the most "
                         "recent row that has it filled in, instead of taking every field from the most recent row."),
    ]

    # Only tasks 1 and 3 (and B3–B5) are typed on the Practice/Bonus sheets. The other parts are worked on Clean,
    # CensusExport, LabFeed, or CleaningLog and read by gray cells, so Start Here, the two how-to lines, and the README's
    # bonus line say where the work happens.
    typed = [i for i, t in enumerate(L.tasks, 1) if not t.summary]
    assert typed == [1, 3], typed
    assert [i for i, t in enumerate(L.bonus, 1) if t.summary] == [1, 2]
    L.practice_how = ("Go to the 'Practice' sheet. Type your answers to tasks 1 and 3 in its yellow cells. For tasks 2 and "
                      f"4–{len(L.tasks)}, do the work on the sheet the task names (Clean, CensusExport, LabFeed, or "
                      "CleaningLog), and the task's gray cell on Practice reads it.")
    L.practice_instructions = (
        "For tasks 1 and 3, type a formula or value in the yellow cell. For the other tasks, do the work on the sheet the "
        "task names, and the gray cell here reads it. The Check column turns green when your answer matches. Stuck? Read "
        f"the hint, then the lesson guide. Answers: right-click a sheet tab → Unhide… → '{L.key_sheet}'.")
    L.bonus_instructions = (
        "For B1 and B2, fill the yellow RegisteredClean and Keep columns on the Clean sheet, and the gray cell here reads "
        "them. For B3–B5, type a formula or value in the yellow cell. The Check column turns green when your answer matches. "
        f"Stuck? Read the hint, then the lesson guide. Answers: right-click a sheet tab → Unhide… → '{L.bonus_key_sheet}'.")
    L.bonus_where = ("Fill the RegisteredClean and Keep columns on the **Clean** sheet, and type your answers to B3–B5 in "
                     "the yellow cells on the **Bonus** sheet.")

    # Start Here describes each data sheet only by its size, so the hook below adds what each one is for.
    practice_cols, bonus_cols = [cl.col(c) for c in CLEAN_COLS[:7]], [cl.col(c) for c in CLEAN_COLS[7:]]
    assert "".join(practice_cols + bonus_cols) == "KLMNOPQRS"
    sheet_roles = {
        "Raw": "The registration export exactly as received. Read it, but never edit it.",
        "Clean": (f"Your working copy of Raw. Fill its yellow columns: {practice_cols[0]}–{practice_cols[-1]} in the "
                  f"practice tasks, and {bonus_cols[0]}–{bonus_cols[-1]} in the bonus."),
        "PayerMap": "The mapping table that turns each insurance spelling into a PayerID (task 8).",
        "CensusExport": "A one-week bed-board census export with blank Facility and Unit cells to fill (task 4).",
        "LabFeed": "A pipe-delimited lab interface feed to split with Text to Columns (task 9).",
        "CleaningLog": "Your cleaning log. Fill in its yellow cells as you go. Task 12 checks step 2's RowsChanged cell.",
    }

    # ------------------------------------------------------------------ workbook polish
    log_fill_cols = ("FixApplied", "RowsChanged")

    @L.customize
    def _polish(wb, lesson, selftest):
        # Start Here: append each data sheet's role to the library's "Data: N rows" line (each sheet is listed once).
        sh = wb["Start Here"]
        hdr = next(r for r in range(1, sh.max_row + 1) if sh.cell(row=r, column=2).value == "Sheets in this workbook")
        done = []
        for r in range(hdr + 1, sh.max_row + 1):
            name = sh.cell(row=r, column=2).value
            if name in sheet_roles:
                c = sh.cell(row=r, column=3)
                c.value = f"{c.value}. {sheet_roles[name]}"
                sh.row_dimensions[r].height = 15 * estimate_lines(c.value, 100) + 3
                done.append(name)
        assert sorted(done) == sorted(sheet_roles), done
        # CleaningLog: learner cells in yellow; wrap long text
        lg_ws = wb["CleaningLog"]
        for r in range(lg.first_row, lg.last_row + 1):
            for col in lg.headers:
                c = lg_ws[f"{lg.col(col)}{r}"]
                c.alignment = Alignment(vertical="top", wrap_text=col in ("ProblemFound", "FixApplied", "Notes"))
                if col in log_fill_cols and c.value is None:
                    c.fill = INPUT_FILL
        note = lg.last_row + 2
        lg_ws.cell(row=note, column=1, value="Fill in FixApplied and RowsChanged as you finish each practice task. "
                                             "Practice task 12 checks the RowsChanged cell for step 2.")
        lg_ws.cell(row=note, column=1).font = Font(italic=True, color="595959")
        # CensusExport and LabFeed: short notes to the right of the data
        ce_ws = wb["CensusExport"]
        ce_ws.cell(row=1, column=7, value="Bed-board export: Facility and Unit print only on the first row of each block.")
        ce_ws.cell(row=2, column=7, value=f"Practice task 4: fill the blanks in {ce.col('Facility')}{ce.first_row}:"
                                          f"{ce.col('Unit')}{ce.last_row} with Go To Special → Blanks.")
        for rr in (1, 2):
            ce_ws.cell(row=rr, column=7).font = Font(italic=True, color="595959")
        lf_ws = wb["LabFeed"]
        lf_ws.cell(row=1, column=8, value="Lab interface feed: one result per line, fields separated by | (vertical bar).")
        lf_ws.cell(row=2, column=8, value="Practice task 9: split column A with Data → Text to Columns. Keep MRN as Text.")
        for rr in (1, 2):
            lf_ws.cell(row=rr, column=8).font = Font(italic=True, color="595959")
        pm_ws = wb["PayerMap"]
        pm_ws.cell(row=pm.last_row + 2, column=1,
                   value="Variants are trimmed and in capitals. XLOOKUP and VLOOKUP ignore case, so one row per spelling is enough.")
        pm_ws.cell(row=pm.last_row + 2, column=1).font = Font(italic=True, color="595959")

    # stash numbers the README quotes, for the author's convenience (python -c "…build().readme_numbers")
    L.readme_numbers = dict(
        extra_spaces=extra_spaces, no_comma=no_comma, sex_spellings=dict(sex_spellings), n_f=n_f, exact_dups=exact_dups,
        icu_days=icu_days, n_blank_cells=n_blank_cells, distinct_zips=distinct_zips, zip4_rows=zip4_rows,
        n_phone10=n_phone10, n_phone_blank=n_phone_blank, n_plus1=n_plus1, distinct_mrn=distinct_mrn, mrn_prefix=mrn_prefix,
        mrn_short=mrn_short, n_py02=n_py02, insurance_spellings=insurance_spellings, avg_k=avg_k, feed=len(feed),
        distinct_names=distinct_names, name_twins=name_twins, mc_rows=mc_rows, day13=day13, dob_layouts=dict(dob_layouts),
        sex_changed=sex_changed, evening=evening, n_keep=n_keep, not_first=not_first, lost_phone=lost_phone,
        rows_per_patient=dict(rows_per_patient), placeholders=Counter(r["Email"] for r in raw if r["Email"] in ("N/A", "none")),
        email_blank=sum(1 for r in raw if not r["Email"]),
        email_space=sum(1 for r in raw if r["Email"] and r["Email"] != r["Email"].strip()),
        email_upper=sum(1 for r in raw if r["Email"] and r["Email"] not in ("N/A", "none") and r["Email"].strip().isupper()),
        phone_layouts=Counter(("blank" if not r["Phone"] else "paren" if r["Phone"].startswith("(") else "plus1"
                               if r["Phone"].startswith("+") else "dots" if "." in r["Phone"] else "dash"
                               if "-" in r["Phone"] else "digits") for r in raw),
        reg_layouts=Counter("iso" if r["RegisteredOn"][4] == "-" else "us" for r in raw),
        all_caps=sum(1 for r in raw if r["PatientName"] == r["PatientName"].upper()),
        ins_spaces=sum(1 for r in raw if r["Insurance"] != xl_trim(r["Insurance"])),
        countif_medicare=sum(1 for r in raw if r["Insurance"].lower() == "medicare"),
        py01=sum(1 for x in payer_ids if x == "PY01"),
        map_rows=len(map_rows), unique_trim=len({xl_trim(r["Insurance"]).lower() for r in raw}),
        email_changed=sum(1 for r in raw if (r["Email"] or "") != xl_trim(r["Email"] or "").lower()),
        csz_upper=sum(1 for r in raw if r["CityStateZip"].split(",")[0] == r["CityStateZip"].split(",")[0].upper()),
        csz_nocomma=sum(1 for r in raw if "," not in r["CityStateZip"]),
        find_comma_r1009=by_rid["R1009"]["PatientName"].find(",") + 1,
        dob_count_text=n, first_feed=feed[0]["line"],
        all_lower=sum(1 for r in raw if r["PatientName"] == r["PatientName"].lower()),
    )
    return L
