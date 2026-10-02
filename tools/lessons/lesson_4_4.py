"""Lesson 4.4 · Data Model, Power Pivot & DAX.

The workbook ships a star schema as nine Excel Tables (two fact tables, seven dimension tables).
Learners load them into the Data Model, relate them, write DAX measures, and type the value a Data Model
PivotTable shows. LibreOffice can't evaluate DAX, so the sample solutions are DAX (solution_lang='dax') or build
steps, the self-test types the Python answer, and each key row carries a LIVE worksheet formula (COUNTIFS, SUMIFS,
SUMPRODUCT, XLOOKUP, UNIQUE) that recomputes the same number from the same tables without the Data Model. The
verifier checks those live formulas against the Python answers, so the DAX filter logic is cross-checked twice.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from xlcourse import Lesson, Task, data
from xlcourse.lesson import HEADER_FILL, NAVY, WRAP_TOP

CODE = "4.4"

AS_OF = date(2025, 12, 31)
FACT_STYLE, DIM_STYLE = "TableStyleMedium2", "TableStyleMedium7"
FACT_TAB, DIM_TAB, DATE_TAB = "1F4E79", "548235", "BF9000"
AGE_BANDS = [(0, 17, "0-17"), (18, 44, "18-44"), (45, 64, "45-64"), (65, 74, "65-74"), (75, 200, "75+")]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def _age(dob: date, ref: date = AS_OF) -> int:
    return ref.year - dob.year - ((ref.month, ref.day) < (dob.month, dob.day))


def _age_group(age: int) -> str:
    return next(label for lo, hi, label in AGE_BANDS if lo <= age <= hi)


def _dr(col: str, d1: date, d2: date) -> str:
    """COUNTIFS/SUMIFS criteria pair for a date column between d1 and d2 (inclusive)."""
    return (f'{col},">="&DATE({d1.year},{d1.month},{d1.day}),'
            f'{col},"<="&DATE({d2.year},{d2.month},{d2.day})')


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="04-advanced-analysis", slug="04-power-pivot-dax",
        title="Data Model, Power Pivot & DAX", level="Advanced", minutes=75,
        objectives=[
            "Design a star schema and load tables into the Data Model",
            "Create relationships and a proper date table",
            "Write DAX measures with SUM, COUNTROWS, DISTINCTCOUNT, DIVIDE, and CALCULATE",
            "Use time intelligence (TOTALYTD, SAMEPERIODLASTYEAR) and iterators (SUMX, AVERAGEX)",
        ],
        data_note="A star schema of the whole Bluestone Health System for 2024–2025: every encounter (21,857 rows) and every "
                  "claim (21,857 rows) as fact tables, plus dimension tables for dates, facilities, departments, providers, "
                  "patients, payers, and diagnoses.",
    )

    # ================================================================== source data
    enc_src = data.load("encounters")
    claims_src = data.load("claims")
    facilities = sorted(data.load("facilities"), key=lambda r: r["FacilityID"])
    departments = sorted(data.load("departments"), key=lambda r: r["DeptID"])
    providers = sorted(data.load("providers"), key=lambda r: r["ProviderID"])
    patients = sorted(data.load("patients"), key=lambda r: r["PatientID"])
    payers = sorted(data.load("payers"), key=lambda r: r["PayerID"])
    diagnoses = sorted(data.load("diagnoses"), key=lambda r: r["DxCode"])

    FAC = {f["FacilityID"]: f for f in facilities}
    FAC_ID = {f["FacilityName"]: f["FacilityID"] for f in facilities}
    DEPT = {d["DeptID"]: d for d in departments}
    PAY = {p["PayerID"]: p for p in payers}
    DX = {d["DxCode"]: d for d in diagnoses}

    # ------------------------------------------------------------------ FactEncounters (one row per encounter)
    def los_days(e) -> int:
        nights = (e["DischargeDateTime"].date() - e["AdmitDateTime"].date()).days
        return max(1, nights) if e["EncounterType"] == "Inpatient" else nights

    fact_enc = []
    for e in sorted(enc_src, key=lambda r: (r["AdmitDateTime"], r["EncounterID"])):
        fact_enc.append({
            "EncounterID": e["EncounterID"], "PatientID": e["PatientID"],
            "AdmitDate": e["AdmitDateTime"].date(), "DischargeDate": e["DischargeDateTime"].date(),
            "FacilityID": e["FacilityID"], "DeptID": e["DeptID"], "AttendingProviderID": e["AttendingProviderID"],
            "PayerID": e["PayerID"], "PrimaryDxCode": e["PrimaryDxCode"], "EncounterType": e["EncounterType"],
            "LOSDays": los_days(e), "TotalCharges": e["TotalCharges"],
            "Readmit30": None if e["Readmit30"] is None else (1 if e["Readmit30"] == "Y" else 0),
        })
    enc_by_id = {r["EncounterID"]: r for r in fact_enc}

    # ------------------------------------------------------------------ FactClaims (one row per claim)
    fact_claims = []
    for c in sorted(claims_src, key=lambda r: (r["ServiceDate"], r["ClaimID"])):
        fact_claims.append({k: c[k] for k in ("ClaimID", "EncounterID", "PayerID", "ServiceDate", "BilledAmount",
                                              "AllowedAmount", "PaidAmount", "ClaimStatus", "DenialReason")})

    # ------------------------------------------------------------------ dimensions
    first_day, last_day = date(2024, 1, 1), date(2025, 12, 31)
    dim_date = []
    d = first_day
    while d <= last_day:
        dim_date.append({"Date": d, "Year": d.year, "Quarter": f"Q{(d.month - 1) // 3 + 1}", "MonthNum": d.month,
                         "MonthName": MONTHS[d.month - 1], "YearMonth": f"{d.year}-{d.month:02d}",
                         "DayName": DAYS[d.weekday()]})
        d = date.fromordinal(d.toordinal() + 1)
    dim_fac = [{k: f[k] for k in ("FacilityID", "FacilityName", "City", "FacilityType", "LicensedBeds")} for f in facilities]
    dim_dept = [{k: x[k] for k in ("DeptID", "DeptName", "ServiceLine", "UnitType", "StaffedBeds")} for x in departments]
    dim_prov = [{"ProviderID": p["ProviderID"], "ProviderName": f"{p['FirstName']} {p['LastName']}",
                 "Credential": p["Credential"], "Specialty": p["Specialty"]} for p in providers]
    dim_pat = [{"PatientID": p["PatientID"], "Sex": p["Sex"], "AgeGroup": _age_group(_age(p["DOB"])), "City": p["City"],
                "PreferredLanguage": p["PreferredLanguage"]} for p in patients]
    dim_pay = [{k: p[k] for k in ("PayerID", "PayerName", "PayerType", "AvgAllowedPctOfCharges")} for p in payers]
    dim_dx = [{k: x[k] for k in ("DxCode", "DxDescription", "DxCategory", "ChronicCondition", "ExpectedLOS")} for x in diagnoses]
    PROV_NAME = {p["ProviderID"]: p["ProviderName"] for p in dim_prov}
    AGEGRP = {p["PatientID"]: p["AgeGroup"] for p in dim_pat}

    # ------------------------------------------------------------------ model integrity (what Power Pivot will enforce)
    for rows, key in ((dim_date, "Date"), (dim_fac, "FacilityID"), (dim_dept, "DeptID"), (dim_prov, "ProviderID"),
                      (dim_pat, "PatientID"), (dim_pay, "PayerID"), (dim_dx, "DxCode")):
        keys = [r[key] for r in rows]
        assert len(keys) == len(set(keys)), f"dimension key {key} must be unique"
    assert len({p["ProviderName"] for p in dim_prov}) == len(dim_prov), "provider names must be unique for RANKX by name"
    date_keys = {r["Date"] for r in dim_date}
    for r in fact_enc:  # every foreign key finds its dimension row, so no '(blank)' rows appear in pivots
        assert r["AdmitDate"] in date_keys and r["DischargeDate"] in date_keys
        assert r["FacilityID"] in FAC and r["DeptID"] in DEPT and r["AttendingProviderID"] in PROV_NAME
        assert r["PatientID"] in AGEGRP and r["PayerID"] in PAY and r["PrimaryDxCode"] in DX
    for c in fact_claims:
        assert c["ServiceDate"] in date_keys and c["PayerID"] in PAY
    assert len({r["EncounterID"] for r in fact_enc}) == len(fact_enc)

    # ================================================================== sheets
    def notes(title: str, text: str) -> list[str]:
        return [title, text]

    L.add_table_sheet(
        "FactEncounters", fact_enc, table="FactEncounters", style=FACT_STYLE, tab_color=FACT_TAB, start_row=3,
        notes=notes("FactEncounters · fact table · one row per encounter (visit or stay), 2024–2025",
                    "Keys point to the Dim tables. LOSDays = discharge date − admit date (inpatient minimum 1). "
                    "Readmit30: 1 = readmitted within 30 days, 0 = not, blank = not an inpatient stay."),
        formats={"LOSDays": "0", "Readmit30": "0"}, widths={"AttendingProviderID": 20, "TotalCharges": 13},
    )
    L.add_table_sheet(
        "FactClaims", fact_claims, table="FactClaims", style=FACT_STYLE, tab_color=FACT_TAB, start_row=3,
        notes=notes("FactClaims · fact table · one row per claim (one claim per encounter), 2024–2025 service dates",
                    "Status as of 12/31/2025. Relate it to DimDate (ServiceDate) and DimPayer only. Never relate two fact "
                    "tables to each other."),
        widths={"DenialReason": 24},
    )
    L.add_table_sheet(
        "DimDate", dim_date, table="DimDate", style=DIM_STYLE, tab_color=DATE_TAB, start_row=3,
        notes=notes("DimDate · date dimension · one row per day, 01/01/2024–12/31/2025 (731 days, no gaps)",
                    "In the Power Pivot window, mark it as the date table (Design → Mark as Date Table) and sort MonthName by "
                    "MonthNum (Home → Sort by Column)."),
    )
    L.add_table_sheet("DimFacility", dim_fac, table="DimFacility", style=DIM_STYLE, tab_color=DIM_TAB, start_row=3,
                      notes=notes("DimFacility · one row per facility", "Key: FacilityID."))
    L.add_table_sheet("DimDepartment", dim_dept, table="DimDepartment", style=DIM_STYLE, tab_color=DIM_TAB, start_row=3,
                      notes=notes("DimDepartment · one row per department (cost center)",
                                  "Key: DeptID. Department names repeat across facilities (every hospital has an "
                                  "'Emergency Department'), so pair DeptName with DimFacility[FacilityName] in pivots."))
    L.add_table_sheet("DimProvider", dim_prov, table="DimProvider", style=DIM_STYLE, tab_color=DIM_TAB, start_row=3,
                      notes=notes("DimProvider · one row per provider", "Key: ProviderID. FactEncounters stores it as "
                                                                       "AttendingProviderID."))
    L.add_table_sheet("DimPatient", dim_pat, table="DimPatient", style=DIM_STYLE, tab_color=DIM_TAB, start_row=3,
                      notes=notes("DimPatient · one row per registered patient (4,000)",
                                  "Key: PatientID. AgeGroup is the patient's age on 12/31/2025 (the course's as-of date)."))
    L.add_table_sheet("DimPayer", dim_pay, table="DimPayer", style=DIM_STYLE, tab_color=DIM_TAB, start_row=3,
                      formats={"AvgAllowedPctOfCharges": "0%"}, widths={"AvgAllowedPctOfCharges": 24},
                      notes=notes("DimPayer · one row per payer",
                                  "Key: PayerID. AvgAllowedPctOfCharges = the share of billed charges the payer's contract "
                                  "typically allows."))
    L.add_table_sheet("DimDiagnosis", dim_dx, table="DimDiagnosis", style=DIM_STYLE, tab_color=DIM_TAB, start_row=3,
                      formats={"ExpectedLOS": "0.0"}, widths={"DxDescription": 50},
                      notes=notes("DimDiagnosis · one row per ICD-10-CM code",
                                  "Key: DxCode (FactEncounters stores it as PrimaryDxCode). ExpectedLOS = benchmark "
                                  "inpatient length of stay in days."))

    # ================================================================== answers (computed in Python)
    def yr(r, y=2025):
        return r["AdmitDate"].year == y

    def is_type(r, t):
        return r["EncounterType"] == t

    F01, F02, F03, F04 = (FAC_ID[n] for n in ("Bluestone Memorial Hospital", "Ashby Falls Community Hospital",
                                              "Cedar Ridge Medical Center", "Bluestone Outpatient Pavilion"))
    ED_DEPTS = {k for k, v in DEPT.items() if v["DeptName"] == "Emergency Department"}
    MEDICARE = next(p["PayerID"] for p in payers if p["PayerName"] == "Medicare")

    # T1 relationship check: encounters per facility via DimFacility (all dates)
    t1 = sum(1 for r in fact_enc if r["FacilityID"] == F03)
    t1_total = len(fact_enc)

    # T2 Total Charges, 2025
    t2_exact = sum(r["TotalCharges"] for r in fact_enc if yr(r))
    t2 = int(round(t2_exact))

    # T3 Cedar Ridge ED encounters, Q3 2025 (DeptName filter x FacilityName filter x Year/Quarter)
    t3 = sum(1 for r in fact_enc if r["FacilityID"] == F03 and r["DeptID"] in ED_DEPTS and yr(r)
             and 7 <= r["AdmitDate"].month <= 9)

    # T4 distinct patients, 2025 (and why it isn't additive across facilities)
    pts25 = {r["PatientID"] for r in fact_enc if yr(r)}
    t4 = len(pts25)
    t4_enc = sum(1 for r in fact_enc if yr(r))
    t4_sum_of_facilities = sum(len({r["PatientID"] for r in fact_enc if yr(r) and r["FacilityID"] == f}) for f in FAC)
    assert t4_sum_of_facilities > t4

    # T5 average inpatient LOS, Bluestone Memorial, 2025
    los_f01 = [r["LOSDays"] for r in fact_enc if yr(r) and r["FacilityID"] == F01 and is_type(r, "Inpatient")]
    t5 = round(sum(los_f01) / len(los_f01), 2)
    los_f01_all = [r["LOSDays"] for r in fact_enc if yr(r) and r["FacilityID"] == F01]
    t5_unfiltered = sum(los_f01_all) / len(los_f01_all)

    # T6 denial rate, Medicare Advantage claims with 2025 service dates (Denied or Appealed = denied)
    ma = [c for c in fact_claims if c["ServiceDate"].year == 2025 and PAY[c["PayerID"]]["PayerType"] == "Medicare Advantage"]
    ma_denied = sum(1 for c in ma if c["ClaimStatus"] in ("Denied", "Appealed"))
    t6 = ma_denied / len(ma)
    t6_denied_only = sum(1 for c in ma if c["ClaimStatus"] == "Denied") / len(ma)

    # T7 readmission rate, Medicare (PayerName) inpatient stays admitted in 2025
    mc = [r for r in fact_enc if yr(r) and is_type(r, "Inpatient") and r["PayerID"] == MEDICARE]
    t7 = sum(r["Readmit30"] for r in mc) / len(mc)

    # T8 Ashby Falls share of system ED visits, 2025 (REMOVEFILTERS on DimFacility)
    ed25 = [r for r in fact_enc if yr(r) and is_type(r, "Emergency")]
    t8 = sum(1 for r in ed25 if r["FacilityID"] == F02) / len(ed25)

    # T9 FILTER + RELATED: Cedar Ridge inpatient stays (2025) with LOSDays > their diagnosis's ExpectedLOS
    cr_ip = [r for r in fact_enc if yr(r) and is_type(r, "Inpatient") and r["FacilityID"] == F03]
    t9 = sum(1 for r in cr_ip if r["LOSDays"] > DX[r["PrimaryDxCode"]]["ExpectedLOS"])

    # T10 SUMX + RELATED: expected allowed = charges x payer allowed %, Bluestone Memorial, 2025
    t10_exact = sum(r["TotalCharges"] * PAY[r["PayerID"]]["AvgAllowedPctOfCharges"]
                    for r in fact_enc if yr(r) and r["FacilityID"] == F01)
    t10 = int(round(t10_exact))
    t10_avg_pct = t10_exact / sum(r["TotalCharges"] for r in fact_enc if yr(r) and r["FacilityID"] == F01)

    # T11 AVERAGEX over patients: average 2025 charges per patient, AgeGroup 75+
    per_pt = defaultdict(float)
    for r in fact_enc:
        if yr(r) and AGEGRP[r["PatientID"]] == "75+":
            per_pt[r["PatientID"]] += r["TotalCharges"]
    t11_exact = sum(per_pt.values()) / len(per_pt)
    t11 = int(round(t11_exact))
    enc75 = [r["TotalCharges"] for r in fact_enc if yr(r) and AGEGRP[r["PatientID"]] == "75+"]
    t11_per_encounter = sum(enc75) / len(enc75)

    # T12 TOTALYTD: Bluestone Outpatient Pavilion charges, Jan 1 to Sep 30 2025
    t12_exact = sum(r["TotalCharges"] for r in fact_enc if r["FacilityID"] == F04
                    and date(2025, 1, 1) <= r["AdmitDate"] <= date(2025, 9, 30))
    t12 = int(round(t12_exact))
    t12_month_only = sum(r["TotalCharges"] for r in fact_enc if r["FacilityID"] == F04
                        and r["AdmitDate"].year == 2025 and r["AdmitDate"].month == 9)

    # T13 SAMEPERIODLASTYEAR: Bluestone Memorial ED visits, 2025 vs 2024
    ed_f01 = {y: sum(1 for r in fact_enc if yr(r, y) and is_type(r, "Emergency") and r["FacilityID"] == F01)
              for y in (2024, 2025)}
    t13 = (ed_f01[2025] - ed_f01[2024]) / ed_f01[2024]

    # ------------------------------------------------------------------ bonus
    def ed_month(y, m, fac=None):
        return sum(1 for r in fact_enc if is_type(r, "Emergency") and r["AdmitDate"].year == y and r["AdmitDate"].month == m
                   and (fac is None or r["FacilityID"] == fac))

    b1_months = [ed_month(2025, m) for m in (10, 11, 12)]
    b1_exact = sum(b1_months) / 3
    b1 = round(b1_exact, 1)
    b2_now, b2_ly = ed_month(2025, 12, F02), ed_month(2024, 12, F02)
    b2 = (b2_now - b2_ly) / b2_ly
    ip25_by_prov = Counter(r["AttendingProviderID"] for r in fact_enc if yr(r) and is_type(r, "Inpatient"))
    (b3_id, b3_n), (_, b3_second) = ip25_by_prov.most_common(2)
    assert b3_n > b3_second, "the #1 provider must be unique"
    b3 = PROV_NAME[b3_id]
    b3_last_first = ", ".join(reversed(b3.split(" ", 1)))
    b4 = sum(1 for r in fact_enc if is_type(r, "Inpatient") and r["DischargeDate"].year == 2025
             and r["DischargeDate"].month == 12)
    b4_admits = sum(1 for r in fact_enc if is_type(r, "Inpatient") and yr(r) and r["AdmitDate"].month == 12)
    assert b4 != b4_admits
    dec1 = date(2025, 12, 1)
    b4_carry_in = sum(1 for r in fact_enc if is_type(r, "Inpatient") and r["AdmitDate"] < dec1
                      and dec1 <= r["DischargeDate"] <= date(2025, 12, 31))
    b4_still_in = sum(1 for r in fact_enc if is_type(r, "Inpatient") and r["AdmitDate"] >= dec1
                      and r["DischargeDate"] > date(2025, 12, 31))
    assert b4 == b4_admits - b4_still_in + b4_carry_in
    if b4_still_in == 0:
        b4_why = (f"The difference is the {b4_carry_in} stays admitted before December and discharged in December. "
                  f"All {b4_admits} December admissions in this extract went home by 12/31. In a live system, some "
                  "December admissions would still be in the hospital at midnight on 12/31, which would pull the "
                  "two numbers apart the other way.")
    else:
        b4_why = (f"{b4_carry_in} stays admitted before December were discharged in December, and {b4_still_in} "
                  "December admissions were still in the hospital at midnight on 12/31.")

    # ================================================================== live cross-check formulas (worksheet only)
    E, C = "FactEncounters", "FactClaims"
    ADM = f"{E}[AdmitDate]"
    Y25 = _dr(ADM, date(2025, 1, 1), date(2025, 12, 31))
    TYPE, FACID, CHG = f"{E}[EncounterType]", f"{E}[FacilityID]", f"{E}[TotalCharges]"
    x_t1 = f'=COUNTIFS({FACID},"{F03}")'
    x_t2 = f"=SUMIFS({CHG},{Y25})"
    x_t3 = (f'=SUMPRODUCT(COUNTIFS({FACID},"{F03}",{E}[DeptID],DimDepartment[DeptID],{_dr(ADM, date(2025, 7, 1), date(2025, 9, 30))})'
            f'*(DimDepartment[DeptName]="Emergency Department"))')
    x_t4 = f"=ROWS(UNIQUE(FILTER({E}[PatientID],YEAR({ADM})=2025)))"
    x_t5 = f'=AVERAGEIFS({E}[LOSDays],{TYPE},"Inpatient",{FACID},"{F01}",{Y25})'
    x_t6 = (f'=LET(ptype,XLOOKUP({C}[PayerID],DimPayer[PayerID],DimPayer[PayerType]),'
            f'keep,(ptype="Medicare Advantage")*(YEAR({C}[ServiceDate])=2025),'
            f'SUMPRODUCT(keep*(({C}[ClaimStatus]="Denied")+({C}[ClaimStatus]="Appealed")))/SUM(keep))')
    x_t7 = (f'=SUMIFS({E}[Readmit30],{E}[PayerID],"{MEDICARE}",{TYPE},"Inpatient",{Y25})'
            f'/COUNTIFS({E}[PayerID],"{MEDICARE}",{TYPE},"Inpatient",{Y25})')
    x_t8 = f'=COUNTIFS({FACID},"{F02}",{TYPE},"Emergency",{Y25})/COUNTIFS({TYPE},"Emergency",{Y25})'
    x_t9 = (f'=SUMPRODUCT(({FACID}="{F03}")*({TYPE}="Inpatient")*(YEAR({ADM})=2025)'
            f'*({E}[LOSDays]>XLOOKUP({E}[PrimaryDxCode],DimDiagnosis[DxCode],DimDiagnosis[ExpectedLOS])))')
    x_t10 = (f'=SUMPRODUCT({CHG}*XLOOKUP({E}[PayerID],DimPayer[PayerID],DimPayer[AvgAllowedPctOfCharges])'
             f'*({FACID}="{F01}")*(YEAR({ADM})=2025))')
    x_t11 = (f'=LET(age,XLOOKUP({E}[PatientID],DimPatient[PatientID],DimPatient[AgeGroup]),'
             f'keep,(age="75+")*(YEAR({ADM})=2025),'
             f'SUMPRODUCT(keep*{CHG})/ROWS(UNIQUE(FILTER({E}[PatientID],keep=1))))')
    x_t12 = f'=SUMIFS({CHG},{FACID},"{F04}",{_dr(ADM, date(2025, 1, 1), date(2025, 9, 30))})'
    x_t13 = (f'=COUNTIFS({FACID},"{F01}",{TYPE},"Emergency",{Y25})'
             f'/COUNTIFS({FACID},"{F01}",{TYPE},"Emergency",{_dr(ADM, date(2024, 1, 1), date(2024, 12, 31))})-1')
    x_b1 = f'=COUNTIFS({TYPE},"Emergency",{_dr(ADM, date(2025, 10, 1), date(2025, 12, 31))})/3'
    x_b2 = (f'=COUNTIFS({FACID},"{F02}",{TYPE},"Emergency",{_dr(ADM, date(2025, 12, 1), date(2025, 12, 31))})'
            f'/COUNTIFS({FACID},"{F02}",{TYPE},"Emergency",{_dr(ADM, date(2024, 12, 1), date(2024, 12, 31))})-1')
    x_b3 = (f'=LET(n,COUNTIFS({E}[AttendingProviderID],DimProvider[ProviderID],{TYPE},"Inpatient",{Y25}),'
            f'INDEX(DimProvider[ProviderName],MATCH(MAX(n),n,0)))')
    x_b4 = f'=COUNTIFS({TYPE},"Inpatient",{_dr(E + "[DischargeDate]", date(2025, 12, 1), date(2025, 12, 31))})'

    # ================================================================== tasks
    L.practice_intro = (
        "Start by building the model (Guide sections 3–5): add all nine tables to the Data Model, create the ten relationships "
        "on the Model Map sheet, and mark DimDate as the date table. Task 1 checks that the model works. For every other task, "
        "create the measure, show it in a PivotTable built from the Data Model, and type the number the PivotTable shows into "
        "the yellow cell (or link to it, as Guide section 15 shows). \"In 2025\" means DimDate[Year] = 2025 through the active "
        "relationships, so encounters are dated by AdmitDate and claims by ServiceDate. \"ED visits\" means encounters with "
        "EncounterType = Emergency.")

    def dax(*lines: str) -> str:
        return "\n".join(lines)

    L.tasks = [
        Task("Load all nine tables into the Data Model and create the relationships listed on the Model Map sheet. Then insert "
             "a PivotTable from the Data Model with DimFacility[FacilityName] in Rows and FactEncounters[EncounterID] in "
             "Values (Excel names it Count of EncounterID). How many encounters (2024 and 2025 together) does Cedar Ridge "
             "Medical Center show?",
             answer=t1, live=x_t1, title="Relationship check: Cedar Ridge encounters",
             hint="If every facility shows the same number, a relationship is missing",
             solution=(
                 "1. Click any cell in the FactEncounters table, then **Power Pivot → Add to Data Model**. Repeat for the "
                 "other eight tables.\n"
                 "2. In the Power Pivot window, click **Home → Diagram View**. Drag **FactEncounters[FacilityID]** onto "
                 "**DimFacility[FacilityID]**, and create the other relationships on the Model Map the same way.\n"
                 "3. Click **Home → PivotTable → PivotTable**, choose **New Worksheet**, and click OK.\n"
                 "4. Tick **DimFacility → FacilityName** so it goes to Rows. Then drag **FactEncounters → EncounterID** "
                 "into the **Values** area. (Ticking a text field such as EncounterID would put it in Rows instead.)\n"
                 f"5. Read the Cedar Ridge Medical Center row. The Grand Total is {t1_total:,}."),
             explanation=(
                 "The facility names live in DimFacility and the encounters live in FactEncounters. The PivotTable can split "
                 "the count by facility only because the relationship carries each facility's filter from DimFacility[FacilityID] "
                 f"to the matching rows of FactEncounters. Without it, every row shows the grand total ({t1_total:,}) and the "
                 "field list displays *Relationships between tables may be needed*. Dragging a field into Values creates an "
                 "**implicit measure** (Count of EncounterID). It works, but you can't reuse it in other formulas, so from "
                 "Task 2 on you write explicit measures.")),
        Task("Create the measure Total Charges = the sum of FactEncounters[TotalCharges]. Show it in a PivotTable with "
             "DimDate[Year] in Rows. What are the total charges for 2025? Enter the amount rounded to the nearest dollar.",
             answer=t2, fmt="#,##0", tol=0.5, live=x_t2, solution_lang="dax", title="Total Charges, 2025",
             hint="Power Pivot → Measures → New Measure…, then SUM",
             solution=dax("Total Charges := SUM(FactEncounters[TotalCharges])",
                          "-- PivotTable: DimDate[Year] in Rows, [Total Charges] in Values"),
             explanation=(
                 "A **measure** is a named formula that the PivotTable evaluates once for every cell. In the 2025 row, the "
                 "filter Year = 2025 travels from DimDate to FactEncounters through the AdmitDate relationship, so SUM adds "
                 "only the 2025 rows. The same measure gives the 2024 row and the grand total without any change. Set the "
                 "measure's format to Currency in the Measure dialog so every PivotTable that uses it shows dollars.")),
        Task("Create the measure Encounters = the number of rows in FactEncounters. How many encounters did the Emergency "
             "Department at Cedar Ridge Medical Center have in Q3 2025? Filter DimFacility[FacilityName] = Cedar Ridge "
             "Medical Center, DimDepartment[DeptName] = Emergency Department, DimDate[Year] = 2025, and "
             "DimDate[Quarter] = Q3.",
             answer=t3, live=x_t3, solution_lang="dax", title="Cedar Ridge ED encounters, Q3 2025",
             hint="COUNTROWS counts the rows of a table",
             solution=dax("Encounters := COUNTROWS(FactEncounters)",
                          "-- PivotTable: DimDate[Year] and DimDate[Quarter] in Rows,",
                          "--   DimFacility[FacilityName] and DimDepartment[DeptName] in Filters (or slicers)"),
             explanation=(
                 "COUNTROWS(FactEncounters) counts whatever rows survive the filters on the cell. Here four filters arrive "
                 "from three different dimension tables, and each one narrows FactEncounters through its own relationship. "
                 "Filtering on DeptName alone would count the Emergency Departments of all three hospitals together, because "
                 "the name repeats, so the FacilityName filter is what isolates Cedar Ridge. COUNTROWS is the usual way to "
                 "count a fact table, because it doesn't depend on any column being filled in.")),
        Task("Create the measure Patients = the number of distinct PatientID values in FactEncounters. How many different "
             "patients had at least one encounter in 2025, system-wide? (Use the Grand Total, not a sum of facility rows.)",
             answer=t4, live=x_t4, solution_lang="dax", title="Distinct patients, 2025",
             hint="DISTINCTCOUNT",
             solution=dax("Patients := DISTINCTCOUNT(FactEncounters[PatientID])",
                          "-- PivotTable: DimDate[Year] in Rows (2025 row), optionally DimFacility[FacilityName] in Columns"),
             explanation=(
                 f"In 2025 there were {t4_enc:,} encounters but only {t4:,} different patients, because many patients came back. "
                 "DISTINCTCOUNT counts each PatientID once in the current filter context. Distinct counts are **not "
                 "additive**: if you put facilities in Columns, the facility values add up to "
                 f"{t4_sum_of_facilities:,}, which is more than the {t4:,} in the Grand Total, because a patient seen at two "
                 "facilities counts once per facility but only once system-wide. A measure recomputes the total from the rows. "
                 "It never adds up the visible cells.")),
        Task("Create the measure Avg LOS (IP) = the average of FactEncounters[LOSDays] for Inpatient encounters only, using "
             "CALCULATE so that the measure applies the filter itself. What is it for Bluestone Memorial Hospital in 2025? "
             "Enter it to 2 decimal places.",
             answer=t5, fmt="0.00", live=x_t5, solution_lang="dax", title="Average inpatient LOS, Bluestone Memorial, 2025",
             hint="CALCULATE(expression, Table[Column] = \"value\")",
             solution=dax("Avg LOS (IP) :=",
                          "CALCULATE(",
                          "    AVERAGE(FactEncounters[LOSDays]),",
                          "    FactEncounters[EncounterType] = \"Inpatient\"",
                          ")",
                          "-- PivotTable: DimFacility[FacilityName] in Rows, DimDate[Year] = 2025 in Filters"),
             explanation=(
                 "CALCULATE evaluates its first argument after adding the filters you list. The PivotTable supplies the "
                 "facility and the year, and CALCULATE adds EncounterType = Inpatient, so the average covers only inpatient "
                 "stays at Bluestone Memorial in 2025. Without the filter, the same cell also averages ED visits and observation "
                 f"stays (most ED visits have 0 LOS days) and drops to {t5_unfiltered:.2f}. Building the filter into the measure means "
                 "nobody has to remember to set an EncounterType slicer.")),
        Task("FactClaims shares DimDate and DimPayer with FactEncounters. Create Claims (rows of FactClaims), Denied Claims "
             "(claims whose ClaimStatus is Denied or Appealed, because an appealed claim was denied first), and Denial Rate = "
             "Denied Claims ÷ Claims, using DIVIDE. What is the denial rate for DimPayer[PayerType] = Medicare Advantage in "
             "2025? Enter it as a percentage with 1 decimal place.",
             answer=t6, fmt="0.0%", live=x_t6, solution_lang="dax", title="Denial rate, Medicare Advantage, 2025",
             hint="Two conditions on the same column can be joined with || inside CALCULATE",
             solution=dax("Claims := COUNTROWS(FactClaims)",
                          "Denied Claims :=",
                          "CALCULATE(",
                          "    [Claims],",
                          "    FactClaims[ClaimStatus] = \"Denied\" || FactClaims[ClaimStatus] = \"Appealed\"",
                          ")",
                          "Denial Rate := DIVIDE([Denied Claims], [Claims])",
                          "-- PivotTable: DimPayer[PayerType] in Rows, DimDate[Year] = 2025 in Filters"),
             explanation=(
                 "The PayerType filter reaches FactClaims through FactClaims[PayerID] and the year reaches it through "
                 "ServiceDate. Those are the claims table's own relationships to the shared (conformed) dimensions, so one "
                 "PivotTable can show encounter measures and claim measures side by side. The || operator means OR, and "
                 "because both conditions test the same column, CALCULATE accepts it as a single filter. In Microsoft 365 you "
                 "can also write FactClaims[ClaimStatus] IN {\"Denied\", \"Appealed\"}. DIVIDE returns a blank whenever "
                 "the denominator is 0 or blank, where the / operator would return Infinity or NaN (DAX never shows "
                 "#DIV/0!). Counting only the Denied status would give "
                 f"{t6_denied_only:.1%}.")),
        Task("Create Inpatient Stays = Encounters for EncounterType = Inpatient (build it on your Encounters measure), and "
             "Readmission Rate = the sum of FactEncounters[Readmit30] ÷ Inpatient Stays. What is the readmission rate for "
             "inpatient stays billed to the payer named Medicare (DimPayer[PayerName]) in 2025? Enter it as a percentage with "
             "1 decimal place.",
             answer=t7, fmt="0.0%", live=x_t7, solution_lang="dax", title="Readmission rate, Medicare, 2025",
             hint="A measure can use another measure: CALCULATE([Encounters], …)",
             solution=dax("Inpatient Stays := CALCULATE([Encounters], FactEncounters[EncounterType] = \"Inpatient\")",
                          "Readmissions := SUM(FactEncounters[Readmit30])",
                          "Readmission Rate := DIVIDE([Readmissions], [Inpatient Stays])",
                          "-- PivotTable: DimPayer[PayerName] in Rows, DimDate[Year] = 2025 in Filters"),
             explanation=(
                 "Readmit30 holds 1 or 0 on inpatient rows and is blank on every other row, so SUM counts the readmitted "
                 "stays. The denominator must count inpatient stays only, which is why Inpatient Stays wraps Encounters in "
                 "CALCULATE. Building measures on measures keeps each definition in one place: if the definition of an "
                 "inpatient stay ever changes, you fix one measure and every rate that uses it updates. Medicare's "
                 "readmission rate matters because CMS reduces payments to hospitals with excess readmissions.")),
        Task("Create ED Visits = Encounters for EncounterType = Emergency, and % of System ED = ED Visits ÷ ED Visits with "
             "every DimFacility filter removed. In a PivotTable with DimFacility[FacilityName] in Rows and DimDate[Year] = "
             "2025, what share of the system's ED visits did Ashby Falls Community Hospital handle? Enter it as a percentage "
             "with 1 decimal place.",
             answer=t8, fmt="0.0%", live=x_t8, solution_lang="dax", title="Ashby Falls share of system ED visits, 2025",
             hint="CALCULATE([ED Visits], ALL(DimFacility)), or REMOVEFILTERS(DimFacility) in Microsoft 365",
             solution=dax("ED Visits := CALCULATE([Encounters], FactEncounters[EncounterType] = \"Emergency\")",
                          "% of System ED :=",
                          "DIVIDE(",
                          "    [ED Visits],",
                          "    CALCULATE([ED Visits], ALL(DimFacility))",
                          ")",
                          "-- Microsoft 365 can also use REMOVEFILTERS(DimFacility) in place of ALL(DimFacility)",
                          "-- PivotTable: DimFacility[FacilityName] in Rows, DimDate[Year] = 2025 in Filters"),
             explanation=(
                 "In the Ashby Falls row, the numerator sees two filters: the facility and the year. The denominator uses "
                 "ALL(DimFacility) to clear the facility filter but keeps the year, so it returns all 2025 ED visits "
                 "in the system. REMOVEFILTERS(DimFacility) does the same job and reads more clearly, but only Excel for "
                 "Microsoft 365 recognizes it. ALL works in every version. This is the DAX "
                 "version of Show Values As → % of Column Total, with one big advantage: it's a real measure that you can "
                 "reuse in other measures, KPIs, and CUBEVALUE formulas. It only works when the PivotTable filters facilities "
                 "through DimFacility. A filter on FactEncounters[FacilityID] would survive ALL(DimFacility).")),
        Task("Create Stays Over Expected = the number of Inpatient encounters whose LOSDays is greater than the ExpectedLOS of "
             "their primary diagnosis in DimDiagnosis. Use FILTER over FactEncounters and RELATED. How many inpatient stays at "
             "Cedar Ridge Medical Center in 2025 ran longer than expected?",
             answer=t9, live=x_t9, solution_lang="dax", title="Cedar Ridge stays over expected LOS, 2025",
             hint="FILTER(FactEncounters, … && FactEncounters[LOSDays] > RELATED(DimDiagnosis[ExpectedLOS]))",
             solution=dax("Stays Over Expected :=",
                          "COUNTROWS(",
                          "    FILTER(",
                          "        FactEncounters,",
                          "        FactEncounters[EncounterType] = \"Inpatient\"",
                          "            && FactEncounters[LOSDays] > RELATED(DimDiagnosis[ExpectedLOS])",
                          "    )",
                          ")",
                          "-- PivotTable: DimFacility[FacilityName] in Rows, DimDate[Year] = 2025 in Filters"),
             explanation=(
                 "This condition compares a column in the fact table with a column in a dimension table, row by row. A simple "
                 "CALCULATE filter can't do that, because it tests one column against fixed values. FILTER walks through every "
                 "FactEncounters row visible in the cell (Cedar Ridge, 2025) and keeps the rows where the test is TRUE. "
                 "Because FILTER works one row at a time (a **row context**), RELATED can follow that row's relationship to "
                 f"DimDiagnosis and fetch its ExpectedLOS. {t9} of Cedar Ridge's {len(cr_ip)} inpatient stays in 2025 ran "
                 "long. The calculated-column alternative is in Guide section 10.")),
        Task("Expected reimbursement: create Expected Allowed = the sum, row by row, of FactEncounters[TotalCharges] × the "
             "payer's AvgAllowedPctOfCharges from DimPayer. Use SUMX and RELATED. What is Expected Allowed for Bluestone "
             "Memorial Hospital in 2025? Enter it rounded to the nearest dollar.",
             answer=t10, fmt="#,##0", tol=0.5, live=x_t10, solution_lang="dax",
             title="Expected Allowed (SUMX), Bluestone Memorial, 2025",
             hint="SUMX(table, expression evaluated on each row)",
             solution=dax("Expected Allowed :=",
                          "SUMX(",
                          "    FactEncounters,",
                          "    FactEncounters[TotalCharges] * RELATED(DimPayer[AvgAllowedPctOfCharges])",
                          ")",
                          "-- PivotTable: DimFacility[FacilityName] in Rows, DimDate[Year] = 2025 in Filters"),
             explanation=(
                 "Each payer pays a different share of charges, so you must multiply on every row before you add. "
                 "SUM(TotalCharges) × AVERAGE(AvgAllowedPctOfCharges) would weight every payer equally and give the wrong "
                 "answer. SUMX is an **iterator**: it evaluates the expression once per visible row (here, Bluestone "
                 "Memorial's 2025 encounters) and then adds the results. RELATED fetches each encounter's payer percentage "
                 f"through the PayerID relationship. The result works out to {t10_avg_pct:.1%} of charges, the blended rate "
                 "that finance uses to estimate net revenue before the claims are paid.")),
        Task("Create Avg Charges per Patient = the average, over the patients in the current filter context, of each "
             "patient's Total Charges. Use AVERAGEX over VALUES(FactEncounters[PatientID]) with your Total Charges measure. "
             "What is it for patients in DimPatient[AgeGroup] = 75+ in 2025? Enter it rounded to the nearest dollar.",
             answer=t11, fmt="#,##0", tol=0.5, live=x_t11, solution_lang="dax", title="Avg charges per patient (AVERAGEX), 75+, 2025",
             hint="AVERAGEX(VALUES(…), [Total Charges])",
             solution=dax("Avg Charges per Patient :=",
                          "AVERAGEX(",
                          "    VALUES(FactEncounters[PatientID]),",
                          "    [Total Charges]",
                          ")",
                          "-- PivotTable: DimPatient[AgeGroup] in Rows, DimDate[Year] = 2025 in Filters"),
             explanation=(
                 "VALUES returns the list of distinct patients visible in the cell (aged 75+, with 2025 encounters). AVERAGEX "
                 "evaluates [Total Charges] once for each patient and averages the results. Calling a measure inside an "
                 "iterator triggers **context transition**: the current patient becomes a filter, so [Total Charges] returns "
                 "that one patient's charges. The result is the same as DIVIDE([Total Charges], [Patients]), and it is a "
                 f"different question from AVERAGE(TotalCharges), which averages per encounter ({t11_per_encounter:,.0f} "
                 "here). Always decide what one unit of the average is: an encounter, a patient, or a month.")),
        Task("Create Charges YTD = Total Charges accumulated from January 1 to the last date in the current filter context, "
             "using TOTALYTD and DimDate[Date]. Put DimDate[Year] and DimDate[MonthName] in Rows and filter Bluestone "
             "Outpatient Pavilion. What does Charges YTD show for September 2025? Enter it rounded to the nearest dollar.",
             answer=t12, fmt="#,##0", tol=0.5, live=x_t12, solution_lang="dax", title="Charges YTD at September 2025, Outpatient Pavilion",
             hint="TOTALYTD(expression, DimDate[Date])",
             solution=dax("Charges YTD := TOTALYTD([Total Charges], DimDate[Date])",
                          "-- PivotTable: DimDate[Year], DimDate[MonthName] in Rows, DimFacility[FacilityName] in Filters"),
             explanation=(
                 "In the September 2025 row, the filter context holds the dates September 1 to September 30, 2025. TOTALYTD "
                 "replaces them with every date from January 1 through September 30, 2025 and evaluates Total Charges over "
                 f"that range. September on its own was {t12_month_only:,.0f}. Time-intelligence functions need a proper date table: one row per day with "
                 "no gaps, marked as the date table, and related to the fact table on a date column. If MonthName sorts "
                 "alphabetically (Apr, Aug, Dec…), set Sort by Column to MonthNum in Power Pivot.")),
        Task("Create ED Visits LY = ED Visits for the same period one year earlier (SAMEPERIODLASTYEAR) and ED YoY % = (ED "
             "Visits − ED Visits LY) ÷ ED Visits LY. What is the year-over-year change in ED visits at Bluestone Memorial "
             "Hospital for 2025 compared with 2024? Enter it as a percentage with 1 decimal place (negative if visits fell).",
             answer=t13, fmt="0.0%", live=x_t13, solution_lang="dax", title="ED visits YoY %, Bluestone Memorial, 2025 vs 2024",
             hint="CALCULATE([ED Visits], SAMEPERIODLASTYEAR(DimDate[Date]))",
             solution=dax("ED Visits LY := CALCULATE([ED Visits], SAMEPERIODLASTYEAR(DimDate[Date]))",
                          "ED YoY % := DIVIDE([ED Visits] - [ED Visits LY], [ED Visits LY])",
                          "-- PivotTable: DimDate[Year] in Rows, DimFacility[FacilityName] in Filters"),
             explanation=(
                 f"In the 2025 row, SAMEPERIODLASTYEAR shifts the year's dates back one year, so ED Visits LY returns the 2024 "
                 f"count ({ed_f01[2024]:,}) next to the 2025 count ({ed_f01[2025]:,}). The 2024 row has no prior year in the "
                 "data, so its LY value is blank and DIVIDE returns a blank. The / operator would show Infinity there. "
                 "The same measures work by "
                 "quarter or by month without any change, because they shift whatever dates the cell contains. "
                 "CALCULATE([ED Visits], DATEADD(DimDate[Date], -1, YEAR)) gives the same result.")),
    ]

    # ================================================================== bonus
    L.bonus_title = "Bonus: the December 2025 operations briefing"
    L.bonus_scenario = (
        "The Chief Medical Officer wants a one-page December 2025 briefing built from the Data Model, so that next month it "
        "refreshes instead of being rebuilt. It needs a rolling ED trend, a same-month comparison for the smallest hospital, "
        "the busiest inpatient attending, and discharges counted by discharge date. Use your measures from the practice "
        "tasks and add new ones as needed. ED visits are encounters with EncounterType = Emergency.")
    L.bonus = [
        Task("Create ED Visits 3M Avg = the average monthly ED visits over the three months ending with the last date in the "
             "current filter context (use DATESINPERIOD). What does it show for December 2025, system-wide? Enter it to 1 "
             "decimal place.",
             answer=b1, fmt="0.0", tol=0.051, live=x_b1, solution_lang="dax", title="Rolling 3-month average ED visits, Dec 2025",
             hint="DATESINPERIOD(DimDate[Date], MAX(DimDate[Date]), -3, MONTH)",
             solution=dax("ED Visits 3M Avg :=",
                          "CALCULATE(",
                          "    AVERAGEX(VALUES(DimDate[YearMonth]), [ED Visits]),",
                          "    DATESINPERIOD(DimDate[Date], MAX(DimDate[Date]), -3, MONTH)",
                          ")",
                          "-- PivotTable: DimDate[Year], DimDate[MonthName] in Rows"),
             explanation=(
                 "In the December 2025 row, MAX(DimDate[Date]) is 12/31/2025, and DATESINPERIOD returns the three months of "
                 "dates ending there (October 1 to December 31). CALCULATE swaps those dates in for December's, then AVERAGEX "
                 "evaluates ED Visits once per month (another context transition) and averages "
                 f"{b1_months[0]}, {b1_months[1]}, and {b1_months[2]}. Dividing the three-month total by 3 gives the same "
                 "number here, but AVERAGEX is more honest at the start of the data: in January 2024 there is only one month "
                 "to average, and /3 would understate it.")),
        Task("Ashby Falls Community Hospital: what is the percentage change in ED visits for December 2025 compared with "
             "December 2024? Enter it as a percentage with 1 decimal place (negative if visits fell).",
             answer=b2, fmt="0.0%", live=x_b2, solution_lang="dax", title="Ashby Falls ED visits, Dec 2025 vs Dec 2024",
             hint="Your ED YoY % measure from Task 13 works at month level too",
             solution=dax("-- Reuse Task 13's measures in a month-level PivotTable:",
                          "ED Visits LY := CALCULATE([ED Visits], SAMEPERIODLASTYEAR(DimDate[Date]))",
                          "ED YoY % := DIVIDE([ED Visits] - [ED Visits LY], [ED Visits LY])",
                          "-- PivotTable: DimDate[Year], DimDate[MonthName] in Rows, FacilityName = Ashby Falls in Filters"),
             explanation=(
                 f"Ashby Falls had {b2_now} ED visits in December 2025 against {b2_ly} in December 2024. You don't need a new "
                 "measure: in the December 2025 cell, SAMEPERIODLASTYEAR shifts December 2025's dates to December 2024. This "
                 "is the payoff of measures. You define the logic once, and it adapts to whatever year, month, or facility the "
                 "cell represents. Small hospitals have small monthly counts, so a large percentage swing can come from a "
                 "difference of a couple of dozen visits. Show the counts next to the percentage in the briefing.")),
        Task("Create Provider Rank = the rank of each attending provider by Inpatient Stays (1 = most stays), using RANKX over "
             "ALL(DimProvider[ProviderName]) and returning a blank on the Grand Total row. With DimProvider[ProviderName] in "
             "Rows and DimDate[Year] = 2025, which provider is ranked 1? Enter the name exactly as it appears in "
             "DimProvider[ProviderName].",
             answer=b3, accept=[b3_id, b3_last_first], live=x_b3, solution_lang="dax", title="Top inpatient attending (RANKX), 2025",
             hint="RANKX(ALL(…), [measure]) ranks against every provider. HASONEVALUE is TRUE only on a single-provider row",
             solution=dax("Provider Rank :=",
                          "IF(",
                          "    HASONEVALUE(DimProvider[ProviderName]),",
                          "    RANKX(ALL(DimProvider[ProviderName]), [Inpatient Stays])",
                          ")",
                          "-- PivotTable: DimProvider[ProviderName] in Rows, DimDate[Year] = 2025 in Filters"),
             explanation=(
                 "RANKX is an iterator. ALL(DimProvider[ProviderName]) gives it every provider name, even though the current "
                 "row shows just one, and RANKX evaluates [Inpatient Stays] for each name through context transition. It then "
                 "reports where the current row's value falls in that list. Without ALL, the list would contain only the "
                 "current provider and every row would rank 1. HASONEVALUE blanks the Grand Total, where a rank means "
                 f"nothing. {b3} attended {b3_n} inpatient stays in 2025, and second place had {b3_second}.")),
        Task("The briefing must count inpatient discharges by discharge date, not admit date. Create IP Discharges = Inpatient "
             "Stays evaluated through the inactive relationship FactEncounters[DischargeDate] → DimDate[Date]. How many "
             "inpatient discharges did the system have in December 2025?",
             answer=b4, live=x_b4, solution_lang="dax", title="Inpatient discharges, Dec 2025 (USERELATIONSHIP)",
             hint="USERELATIONSHIP(many-side column, one-side column) goes inside CALCULATE as a filter argument",
             solution=dax("IP Discharges :=",
                          "CALCULATE(",
                          "    [Inpatient Stays],",
                          "    USERELATIONSHIP(FactEncounters[DischargeDate], DimDate[Date])",
                          ")",
                          "-- PivotTable: DimDate[Year], DimDate[MonthName] in Rows"),
             explanation=(
                 "Only one relationship between two tables can be active, so the DischargeDate relationship is inactive (a "
                 "dashed line in Diagram View) and normally does nothing. USERELATIONSHIP switches it on for this one "
                 "calculation, so the December 2025 filter from DimDate now selects stays that were discharged in December. "
                 f"In the same row, Inpatient Stays shows {b4_admits} admissions and IP Discharges shows {b4}. {b4_why} "
                 "If the inactive relationship doesn't exist in your model, "
                 "USERELATIONSHIP returns an error, so create it first (see the Model Map).")),
    ]

    # ================================================================== start notes & model map
    L.start_notes = [
        "You need Excel for Windows (Microsoft 365, or Excel 2016 or later) with the Power Pivot add-in turned on: File → "
        "Options → Add-ins → Manage: COM Add-ins → Go… → tick Microsoft Power Pivot for Excel. Excel for Mac and Excel for "
        "the web can't create a Data Model, relationships, or measures, so on those you can read the guide and the answer "
        "key but not build the model.",
        "The Model Map sheet lists the nine tables, their keys, and the ten relationships to create (nine active, one "
        "inactive).",
        "Each answer is a number (or a name) that you read from a Data Model PivotTable and type into the yellow cell. You "
        "can also link to the PivotTable cell, which writes a GETPIVOTDATA formula, or use CUBEVALUE (Guide section 15).",
        "The hidden Answer Key has a worksheet formula next to each answer that recomputes it from the same tables without "
        "the Data Model, so you can compare DAX filter context with COUNTIFS and SUMIFS criteria.",
        "The workbook is large (about 44,000 fact rows). Adding the tables to the Data Model makes the file bigger again, "
        "so save it after you build the model.",
    ]
    L.sheet_order = ["Start Here", "Model Map", "Practice", "Bonus", "FactEncounters", "FactClaims", "DimDate",
                     "DimFacility", "DimDepartment", "DimProvider", "DimPatient", "DimPayer", "DimDiagnosis"]

    tables_info = [
        ("FactEncounters", "Fact", "one encounter (visit or stay)", "EncounterID", len(fact_enc)),
        ("FactClaims", "Fact", "one claim", "ClaimID", len(fact_claims)),
        ("DimDate", "Dimension (date table)", "one calendar day", "Date", len(dim_date)),
        ("DimFacility", "Dimension", "one facility", "FacilityID", len(dim_fac)),
        ("DimDepartment", "Dimension", "one department", "DeptID", len(dim_dept)),
        ("DimProvider", "Dimension", "one provider", "ProviderID", len(dim_prov)),
        ("DimPatient", "Dimension", "one patient", "PatientID", len(dim_pat)),
        ("DimPayer", "Dimension", "one payer", "PayerID", len(dim_pay)),
        ("DimDiagnosis", "Dimension", "one diagnosis code", "DxCode", len(dim_dx)),
    ]
    assert len(tables_info) == 9
    rels = [
        ("FactEncounters[AdmitDate]", "DimDate[Date]", "Active", "Dates encounters (\"in 2025\", months, YTD, last year)"),
        ("FactEncounters[DischargeDate]", "DimDate[Date]", "Inactive", "Discharges by discharge date (USERELATIONSHIP, bonus B4)"),
        ("FactEncounters[FacilityID]", "DimFacility[FacilityID]", "Active", "Facility names, types, beds"),
        ("FactEncounters[DeptID]", "DimDepartment[DeptID]", "Active", "Department names and service lines"),
        ("FactEncounters[AttendingProviderID]", "DimProvider[ProviderID]", "Active", "Attending provider names and specialties"),
        ("FactEncounters[PatientID]", "DimPatient[PatientID]", "Active", "Age group, sex, city, language"),
        ("FactEncounters[PayerID]", "DimPayer[PayerID]", "Active", "Payer names, payer types, allowed %"),
        ("FactEncounters[PrimaryDxCode]", "DimDiagnosis[DxCode]", "Active", "Diagnosis descriptions, categories, expected LOS"),
        ("FactClaims[ServiceDate]", "DimDate[Date]", "Active", "Dates claims"),
        ("FactClaims[PayerID]", "DimPayer[PayerID]", "Active", "Payer attributes for claims"),
    ]

    @L.customize
    def _model_map(wb, lesson, selftest):
        ws = wb.create_sheet("Model Map")
        ws.sheet_properties.tabColor = NAVY
        thin = Side(style="thin", color="BFBFBF")
        box = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws["A1"] = "Model Map · the Bluestone star schema"
        ws["A1"].font = Font(bold=True, size=16, color=NAVY)
        ws["A2"] = ("Two fact tables in the middle, seven dimension tables around them. Every relationship runs from a key "
                    "column in a fact table (the many side) to the unique key of a dimension table (the one side). "
                    "Filters flow from the dimension to the fact.")
        ws["A2"].font = Font(italic=True, color="404040")
        ws["A2"].alignment = WRAP_TOP
        ws.merge_cells("A2:E2")
        ws.row_dimensions[2].height = 32

        def header(row, labels):
            for j, h in enumerate(labels, 1):
                c = ws.cell(row=row, column=j, value=h)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = HEADER_FILL
                c.border = box

        r = 4
        ws.cell(row=r, column=1, value="Tables").font = Font(bold=True, size=13, color=NAVY)
        r += 1
        header(r, ["Table", "Role", "Key", "Rows", "Grain (one row per…)"])
        for name, role, grain, key, nrows in tables_info:
            r += 1
            for j, v in enumerate((name, role, key, nrows, grain), 1):
                c = ws.cell(row=r, column=j, value=v)
                c.border = box
                if j == 4:
                    c.number_format = "#,##0"
            fill = "DDEBF7" if role == "Fact" else ("FFF2CC" if "date" in role else "E2EFDA")
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=fill)
            ws.cell(row=r, column=1).font = Font(bold=True)
        r += 2
        ws.cell(row=r, column=1, value="Relationships to create").font = Font(bold=True, size=13, color=NAVY)
        r += 1
        header(r, ["#", "Many side (fact column)", "One side (dimension key)", "Active?", "What it lets you slice by"])
        for i, (frm, to, act, why) in enumerate(rels, 1):
            r += 1
            vals = [i, frm, to, act, why]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.border = box
                c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if j == 1 else None)
            if act == "Inactive":
                for j in range(1, 6):
                    ws.cell(row=r, column=j).font = Font(italic=True, color="7F7F7F")
        r += 2
        tips = [
            "Do not relate FactClaims to FactEncounters. Two fact tables share dimensions instead (DimDate and DimPayer here).",
            "Power Pivot makes the second FactEncounters → DimDate relationship inactive automatically (dashed line). That is "
            "what you want.",
            "Mark DimDate as the date table (Design → Mark as Date Table → Date) and sort MonthName by MonthNum (Home → Sort "
            "by Column).",
            "Optional polish: hide the fact tables' foreign-key columns (FacilityID, DeptID, PayerID…) from client tools, "
            "so the field list shows only the dimension columns to slice by.",
        ]
        ws.cell(row=r, column=1, value="Notes").font = Font(bold=True, size=13, color=NAVY)
        for tip in tips:
            r += 1
            ws.cell(row=r, column=1, value="•").alignment = Alignment(horizontal="right", vertical="top")
            c = ws.cell(row=r, column=2, value=tip)
            c.alignment = WRAP_TOP
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            ws.row_dimensions[r].height = 30
        for col, w in zip("ABCDE", (18, 36, 30, 12, 52)):
            ws.column_dimensions[col].width = w
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        note = ("Spoiler alert: try every task before reading this sheet. Column C is the value the Check column compares "
                "against. Column D is the DAX (or the steps) that produces it in a Data Model PivotTable. Column E is NOT the "
                "DAX: it is a worksheet formula over the same tables (COUNTIFS, SUMIFS, SUMPRODUCT, XLOOKUP, UNIQUE) that "
                "recomputes the answer without the Data Model. Click a cell in column E to read it.")
        for name in (lesson.key_sheet, lesson.bonus_key_sheet):
            if name in wb.sheetnames:
                wb[name]["A2"] = note
                wb[name].row_dimensions[2].height = 48

    return L
