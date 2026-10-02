"""Lesson 3.2 · Data Validation & Conditional Formatting.

Practice, Part A (tasks 1-6): the learner adds data-validation rules to a Patient Access intake log that already holds
some bad entries, then counts what Circle Invalid Data flags. Part B (tasks 7-13): the learner builds conditional
formatting rules on ICU STAT labs, supply inventory, and a unit census, then answers a checkable question about what
each rule highlights. Every answer is computed here in Python; each solution also shows an equivalent formula, which
runs live in the hidden key.

Bonus: a system bed-huddle board (one row per inpatient unit) with a color scale, an icon set, two prioritized row
rules, and a dependent Facility -> Unit selector that drives a pre-built unit card.
"""
from __future__ import annotations

import statistics
from datetime import date, timedelta

from openpyxl.formatting.rule import ColorScaleRule, FormatObject, FormulaRule, IconSet, Rule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from xlcourse import Lesson, Task, data
from xlcourse.lesson import BOX, HEADER_FILL, INPUT_BORDER, INPUT_FILL, NAVY, PREFILL_FILL, _plain

CODE = "3.2"
REPORT_DATE = date(2025, 12, 31)
EXPIRING_WINDOW = 90

RED_FILL = PatternFill("solid", fgColor="FFC7CE")
AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")


# ---------------------------------------------------------------------------- data builders
def _facilities():
    return {f["FacilityID"]: f["FacilityName"] for f in data.load("facilities")}


def _registration_departments():
    """Departments a patient can be registered to, by facility (Pharmacy excluded)."""
    fac = _facilities()
    out: dict[str, list[str]] = {name: [] for name in fac.values()}
    for d in sorted(data.load("departments"), key=lambda d: (d["FacilityID"], d["DeptID"])):
        if d["DeptName"] == "Pharmacy":
            continue
        out[fac[d["FacilityID"]]].append(d["DeptName"])
    return out


def _intake_log():
    """A December 2025 Patient Access intake log built from real encounters, with deliberate entry errors."""
    fac = _facilities()
    depts = data.index(data.load("departments"), "DeptID")
    pats = data.index(data.load("patients"), "PatientID")
    payers = data.index(data.load("payers"), "PayerID")
    dec = [e for e in data.load("encounters")
           if e["AdmitDateTime"].year == 2025 and e["AdmitDateTime"].month == 12]
    dec.sort(key=lambda e: (e["AdmitDateTime"], e["EncounterID"]))
    step = len(dec) // 50
    picked = [dec[i * step] for i in range(50)]
    rows = []
    for i, e in enumerate(picked):
        p = pats[e["PatientID"]]
        py = payers[e["PayerID"]]
        etype, ptype = e["EncounterType"], py["PayerType"]
        if etype == "Emergency":
            copay = {"Commercial": 100.0, "Medicare Advantage": 75.0}.get(ptype, 0.0)
        elif etype == "Outpatient":
            copay = {"Commercial": 25.0 if i % 2 else 40.0, "Medicare Advantage": 20.0, "Self-Pay": 50.0,
                     "Government": 0.0 if py["PayerName"] == "Medicare" else 3.0}.get(ptype, 0.0)
        else:
            copay = 0.0
        rows.append({
            "IntakeID": f"INT-{25001 + i}",
            "VisitDate": e["AdmitDateTime"].date(),
            "MRN": p["MRN"],
            "PatientName": f"{p['LastName']}, {p['FirstName']}",
            "ZIP": p["ZIP"],
            "Payer": py["PayerName"],
            "CopayAmt": copay,
            "Facility": fac[e["FacilityID"]],
            "Department": depts[e["DeptID"]]["DeptName"],
        })

    # ---- deliberate errors (deterministic positions) ----
    # payer typos are near misses of the patient's real payer (plus one payer that isn't contracted at all)
    used: set[int] = set()
    for target, real, bad in [(4, "Self-Pay", "Self Pay"), (13, "State Medicaid", "Medicaid"),
                              (22, "Evergreen Mutual Insurance", "BCBS of Ohio"), (31, "Summit Choice PPO", "Summit Choice"),
                              (44, "Keystone Health Partners", "Keystone Health")]:
        j = min((k for k, r in enumerate(rows) if r["Payer"] == real and k not in used), key=lambda k: abs(k - target))
        used.add(j)
        rows[j]["Payer"] = bad
    rows[7]["VisitDate"] = "12/32/2025"          # text: no such day
    rows[19]["VisitDate"] = date(2026, 12, rows[19]["VisitDate"].day)   # year typo
    d33 = rows[33]["VisitDate"]
    rows[33]["VisitDate"] = f"{d33.day:02d}/{d33.month:02d}/{d33.year}"   # text: the same date imported day-first
    rows[46]["VisitDate"] = date(2015, 12, rows[46]["VisitDate"].day)   # year typo
    rows[10]["CopayAmt"] = -20.0                 # refund typed as a copay
    rows[27]["CopayAmt"] = 2500.0                # $25.00 typed without the decimal point
    rows[38]["CopayAmt"] = 350.0                 # above the $100 registration limit
    rows[3]["ZIP"] = rows[3]["ZIP"][:4]          # 4 characters
    rows[24]["ZIP"] = rows[24]["ZIP"] + "-1234"  # ZIP+4, 10 characters
    rows[40]["ZIP"] = rows[40]["ZIP"] + "2"      # 6 characters
    rows[15]["ZIP"] = int(rows[15]["ZIP"])       # stored as a number but still 5 characters (valid)
    zero_rows = [i for i, r in enumerate(rows) if r["MRN"].startswith("0")]
    mrn_err = {6: "drop", 17: "extra", zero_rows[5]: "number", zero_rows[11]: "letterO"}
    for i, kind in mrn_err.items():
        m = rows[i]["MRN"]
        rows[i]["MRN"] = {"drop": m[:7], "extra": m + "1", "number": int(m), "letterO": "O" + m[1:]}[kind]
    # departments that exist in the system but not at the row's facility
    wrong = {"Bluestone Memorial Hospital": "Pediatric Clinic",
             "Ashby Falls Community Hospital": "Cardiac Step-Down",
             "Cedar Ridge Medical Center": "Labor & Delivery",
             "Bluestone Outpatient Pavilion": "Emergency Department"}
    targets = [9, 21, 36, 48]
    used = set()
    for t, (facname, dept) in zip(targets, wrong.items()):
        j = min((k for k, r in enumerate(rows) if r["Facility"] == facname and k not in used), key=lambda k: abs(k - t))
        used.add(j)
        rows[j]["Department"] = dept
    return rows


def _labs():
    """STAT labs for the three ICUs, Jul-Dec 2025, plus a few results the interface re-sent (duplicate IDs)."""
    fac = _facilities()
    enc = data.index(data.load("encounters"), "EncounterID")
    rows = []
    for r in data.load("lab_results"):
        e = enc[r["EncounterID"]]
        if r["Priority"] == "STAT" and e["DeptID"] in ("D130", "D230", "D330") \
                and date(2025, 7, 1) <= r["CollectedDateTime"].date() <= date(2025, 12, 31):
            r["Facility"] = fac[e["FacilityID"]]
            r["TATMin"] = int(round((r["ResultedDateTime"] - r["CollectedDateTime"]).total_seconds() / 60))
            rows.append(r)
    rows.sort(key=lambda r: (r["CollectedDateTime"], r["LabResultID"]))
    # interface re-sends: five results sent twice, one sent three times (all normal results)
    plan = {}
    for target, copies in [(30, 1), (85, 1), (140, 1), (190, 1), (245, 1), (280, 2)]:
        j = next(k for k in range(target, len(rows)) if rows[k]["AbnormalFlag"] == "N" and k not in plan)
        plan[j] = copies
    out = []
    for k, r in enumerate(rows):
        out.append(r)
        for _ in range(plan.get(k, 0)):
            out.append(dict(r))
    return out


def _supplies():
    fac = _facilities()
    depts = data.index(data.load("departments"), "DeptID")
    rows = sorted(data.load("supply_inventory"), key=lambda r: r["StockID"])
    for r in rows:
        d = depts[r["LocationDeptID"]]
        r["Facility"] = fac[d["FacilityID"]]
        r["Location"] = d["DeptName"]
        r["PctOfPar"] = r["QtyOnHand"] / r["ParLevel"]
    return rows


def _unit_census():
    rows = [r for r in data.load("daily_census")
            if r["DeptID"] == "D111" and date(2025, 11, 1) <= r["CensusDate"] <= date(2025, 12, 31)]
    rows.sort(key=lambda r: r["CensusDate"])
    for r in rows:
        r["Occupancy"] = r["MidnightCensus"] / r["StaffedBeds"]
    return rows


def _huddle_data():
    """Board rows (one per inpatient unit, census 11/24 vs 11/25/2025) and Nov 1-25 daily rows for the card."""
    fac = _facilities()
    depts = data.index(data.load("departments"), "DeptID")
    cen = data.load("daily_census")
    y = {r["DeptID"]: r for r in cen if r["CensusDate"] == date(2025, 11, 24)}
    t = {r["DeptID"]: r for r in cen if r["CensusDate"] == date(2025, 11, 25)}
    order = sorted(t, key=lambda k: (t[k]["FacilityID"], k))
    board = []
    for k in order:
        board.append({"Facility": fac[t[k]["FacilityID"]], "Unit": depts[k]["DeptName"], "DeptID": k,
                      "Beds": t[k]["StaffedBeds"], "Yesterday": y[k]["MidnightCensus"], "Today": t[k]["MidnightCensus"]})
    nov = []
    for k in order:
        for r in sorted((r for r in cen if r["DeptID"] == k and date(2025, 11, 1) <= r["CensusDate"] <= date(2025, 11, 25)),
                        key=lambda r: r["CensusDate"]):
            nov.append({"CensusDate": r["CensusDate"], "Facility": fac[r["FacilityID"]], "Unit": depts[k]["DeptName"],
                        "StaffedBeds": r["StaffedBeds"], "MidnightCensus": r["MidnightCensus"],
                        "Occupancy": r["MidnightCensus"] / r["StaffedBeds"]})
    return board, nov


# ---------------------------------------------------------------------------- lesson
def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="03-data-analysis", slug="02-data-validation-conditional-formatting",
        title="Data Validation & Conditional Formatting", level="Intermediate", minutes=55,
        objectives=[
            "Restrict entries with list, number, date, length, and custom-formula validation",
            "Build dependent drop-down lists",
            "Highlight what matters with conditional formatting rules, data bars, color scales, and icon sets",
            "Write formula-based rules that format entire rows",
        ],
        data_note="A December 2025 Patient Access intake log with entry errors, STAT lab results from the three ICUs "
                  "(Jul–Dec 2025), the system supply inventory (12/31/2025 snapshot), and Medical-Surgical 5 East's daily "
                  "census (Nov–Dec 2025). The bonus adds a bed-huddle board for every inpatient unit.",
    )

    # ------------------------------------------------------------------ data sheets
    intake = _intake_log()
    itk = L.add_table_sheet(
        "IntakeLog", intake, table="tblIntake",
        columns=["IntakeID", "VisitDate", "MRN", "PatientName", "ZIP", "Payer", "CopayAmt", "Facility", "Department"],
        formats={"CopayAmt": "#,##0.00", "VisitDate": "mm/dd/yyyy"},
        widths={"PatientName": 24, "Payer": 30, "Facility": 32, "Department": 32, "MRN": 12, "ZIP": 12, "VisitDate": 12},
    )
    labs = _labs()
    lab = L.add_table_sheet(
        "Labs", labs, table="tblLabs",
        columns=["LabResultID", "CollectedDateTime", "Facility", "PatientID", "TestCode", "TestName", "ResultValue",
                 "Units", "RefLow", "RefHigh", "AbnormalFlag", "TATMin"],
        formats={"ResultValue": "General", "RefLow": "General", "RefHigh": "General"},
        widths={"Facility": 30, "TestName": 30, "CollectedDateTime": 17, "AbnormalFlag": 13},
    )
    supplies = _supplies()
    sup = L.add_table_sheet(
        "Supplies", supplies, table="tblSupplies",
        columns=["StockID", "SKU", "ItemDescription", "Category", "Facility", "Location", "UnitOfMeasure", "QtyOnHand",
                 "ParLevel", "ReorderPoint", "PctOfPar", "ExpirationDate", "LotNumber"],
        formats={"PctOfPar": "0%", "ExpirationDate": "mm/dd/yyyy"},
        widths={"ItemDescription": 40, "Facility": 30, "Location": 28, "PctOfPar": 11, "ExpirationDate": 15},
    )
    census = _unit_census()
    cen = L.add_table_sheet(
        "Census", census, table="tblCensus",
        columns=["CensusDate", "StaffedBeds", "Admissions", "Discharges", "MidnightCensus", "Occupancy"],
        formats={"Occupancy": "0.0%"}, widths={"Occupancy": 12},
    )
    board, nov = _huddle_data()
    novsd = L.add_table_sheet(
        "NovCensus", nov, table="tblNovCensus",
        columns=["CensusDate", "Facility", "Unit", "StaffedBeds", "MidnightCensus", "Occupancy"],
        formats={"Occupancy": "0.0%"}, widths={"Facility": 32, "Unit": 26},
    )
    L.sheet_order = ["Start Here", "Practice", "IntakeLog", "Lists", "Labs", "Supplies", "Census", "Bonus", "Huddle",
                     "NovCensus", "Answer Key", "Bonus Key", "Huddle Key"]

    # ------------------------------------------------------------------ Lists sheet layout (built in customize)
    payer_names = [p["PayerName"] for p in sorted(data.load("payers"), key=lambda p: p["PayerID"])]
    fac_names = [f["FacilityName"] for f in sorted(data.load("facilities"), key=lambda f: f["FacilityID"])]
    reg = _registration_departments()
    block_cols = ["E", "F", "G", "H"]
    pairs = [(f, d) for f in fac_names for d in reg[f]]
    pay_rng = f"Lists!$A$2:$A${1 + len(payer_names)}"
    fac_rng = f"Lists!$C$2:$C${1 + len(fac_names)}"
    map_f = f"Lists!$J$2:$J${1 + len(pairs)}"
    map_d = f"Lists!$K$2:$K${1 + len(pairs)}"
    block_ranges = [f"{c}1:{c}{1 + len(reg[f])}" for c, f in zip(block_cols, fac_names)]

    # ------------------------------------------------------------------ answers (computed in Python)
    def rr(sd, col):
        """Sheet-qualified relative range without quotes, e.g. IntakeLog!F2:F51 (all sheet names here are one word)."""
        c = sd.col(col)
        return f"{sd.name}!{c}{sd.first_row}:{c}{sd.last_row}"

    ir = lambda col: rr(itk, col)
    ic = lambda col: f"{itk.col(col)}{itk.first_row}:{itk.col(col)}{itk.last_row}"
    n_int = len(intake)

    bad_payer = sum(1 for r in intake if r["Payer"] not in payer_names)
    bad_date = sum(1 for r in intake if not (isinstance(r["VisitDate"], date)
                                             and date(2025, 12, 1) <= r["VisitDate"] <= date(2025, 12, 31)))
    bad_copay = sum(1 for r in intake if not (0 <= r["CopayAmt"] <= 100))
    bad_zip = sum(1 for r in intake if len(str(r["ZIP"])) != 5)
    bad_mrn = sum(1 for r in intake if not (len(str(r["MRN"])) == 8 and str(r["MRN"]).isdigit()))
    bad_dept = sum(1 for r in intake if r["Department"] not in reg[r["Facility"]])
    text_dates = [r["VisitDate"] for r in intake if isinstance(r["VisitDate"], str)]
    day_first = next(t for t in text_dates if t != "12/32/2025")
    letter_o = next(str(r["MRN"]) for r in intake if str(r["MRN"]).startswith("O"))
    zip_num = next(r["ZIP"] for r in intake if isinstance(r["ZIP"], int))
    has_zero_copay = any(r["CopayAmt"] == 0 for r in intake)
    has_100_copay = any(r["CopayAmt"] == 100 for r in intake)
    assert has_zero_copay and has_100_copay
    assert all(r["Facility"] in fac_names for r in intake)

    # Labs
    lr = lambda col: rr(lab, col)
    lc = lambda col: f"{lab.col(col)}{lab.first_row}:{lab.col(col)}{lab.last_row}"
    ids = [r["LabResultID"] for r in labs]
    dup_cells = sum(1 for x in ids if ids.count(x) > 1)
    dup_ids = len({x for x in ids if ids.count(x) > 1})
    extra_rows = len(ids) - len(set(ids))
    tats = [r["TATMin"] for r in labs]
    tenth = sorted(tats, reverse=True)[9]
    top_cells = sum(1 for x in tats if x >= tenth)
    yellow = sum(1 for x in tats if 45 <= x < 60)
    red_icons = sum(1 for x in tats if x >= 60)
    crit_rows = sum(1 for r in labs if r["AbnormalFlag"] in ("HH", "LL"))
    n_hh = sum(1 for r in labs if r["AbnormalFlag"] == "HH")
    n_ll = sum(1 for r in labs if r["AbnormalFlag"] == "LL")
    last_lab_col = lab.col("TATMin")

    # Supplies
    sr = lambda col: rr(sup, col)
    full_bars = sum(1 for r in supplies if r["PctOfPar"] >= 1)
    max_par = max(r["PctOfPar"] for r in supplies)
    exps = [r["ExpirationDate"] for r in supplies if r["ExpirationDate"]]
    expired = sum(1 for d in exps if d <= REPORT_DATE)
    window_end = REPORT_DATE + timedelta(days=EXPIRING_WINDOW)
    expiring_all = sum(1 for d in exps if d <= window_end)
    amber = sum(1 for d in exps if REPORT_DATE < d <= window_end)
    blanks = sum(1 for r in supplies if not r["ExpirationDate"])
    assert all(d != REPORT_DATE and d != window_end for d in exps), "boundary dates would make the rule ambiguous"

    # Census color scale
    occ = [r["Occupancy"] for r in census]
    occ_median = statistics.median(occ)
    median_days = sum(1 for x in occ if abs(x - occ_median) < 1e-12)

    # ------------------------------------------------------------------ tasks
    L.practice_intro = (
        "Part A (tasks 1–6) uses the IntakeLog and Lists sheets. Add each validation rule, then click the arrow on "
        "Data → Data Validation → Circle Invalid Data and count the red circles. Part B (tasks 7–13) uses the Labs, "
        "Supplies, and Census sheets. Create each conditional formatting rule, then answer the question about what it "
        "highlights. Type the number you count, or write a formula that calculates it.")

    circle = ("Click the arrow on **Data → Data Validation**, then **Circle Invalid Data**. "
              "(**Clear Validation Circles** removes them.)")

    L.tasks = [
        # ---------------- Part A: data validation ----------------
        Task(f"IntakeLog, Payer column ({ic('Payer')}): add a List validation whose Source is the payer list on the Lists sheet "
             f"({pay_rng}). Then circle invalid data. How many Payer cells are circled?",
             answer=bad_payer, title="Payer drop-down list (circled entries)",
             hint="Data → Data Validation → Allow: List",
             solution=(f"1. On **IntakeLog**, select **{ic('Payer')}**.\n"
                       "2. **Data → Data Validation** (Windows: Alt, A, V, V). On the **Settings** tab set **Allow** to **List**.\n"
                       f"3. Click in **Source**, switch to the **Lists** sheet, and select **A2:A{1 + len(payer_names)}**. "
                       f"The box reads `={pay_rng}`. Click **OK**.\n"
                       f"4. {circle} You see **{bad_payer}** circles.\n\n"
                       f"Equivalent formula: `=SUMPRODUCT(--(COUNTIF({pay_rng},{ir('Payer')})=0))`"),
             live=f"=SUMPRODUCT(--(COUNTIF({pay_rng},{ir('Payer')})=0))",
             explanation="A validation rule only checks values as someone types them, so the bad payers already in the log "
                         "stay until you circle them. Each one is a near miss of a real payer: \"Self Pay\" instead of "
                         "\"Self-Pay\", \"Medicaid\" instead of \"State Medicaid\". A drop-down list prevents exactly this. "
                         "The formula counts entries whose COUNTIF in the payer list is 0, which means they aren't on it."),
        Task(f"IntakeLog, VisitDate ({ic('VisitDate')}): the log covers December 2025 only. Allow a Date between 12/01/2025 "
             "and 12/31/2025 (type =DATE(2025,12,1) and =DATE(2025,12,31) in the boxes so it works in any regional setting). "
             "How many VisitDate cells are circled?",
             answer=bad_date, title="December-only VisitDate rule (circled entries)",
             hint="Allow: Date, Data: between",
             solution=(f"1. Select **{ic('VisitDate')}** on IntakeLog.\n"
                       "2. **Data → Data Validation → Settings**: **Allow** = **Date**, **Data** = **between**, "
                       "**Start date** = `=DATE(2025,12,1)`, **End date** = `=DATE(2025,12,31)`. Click **OK**.\n"
                       f"3. {circle} You see **{bad_date}** circles.\n\n"
                       f"Equivalent formula: `=ROWS({ir('VisitDate')})-COUNTIFS({ir('VisitDate')},\">=\"&DATE(2025,12,1),"
                       f"{ir('VisitDate')},\"<=\"&DATE(2025,12,31))`"),
             live=f"=ROWS({ir('VisitDate')})-COUNTIFS({ir('VisitDate')},\">=\"&DATE(2025,12,1),{ir('VisitDate')},\"<=\"&DATE(2025,12,31))",
             explanation=f"Two entries are text that only looks like a date: 12/32/2025 doesn't exist, and {day_first} was "
                         "imported day-first. Both sit at the left of the cell, which is the clue that they're text. A Date "
                         "rule rejects text automatically. The other two are real dates with the wrong year (2026 and 2015). "
                         "COUNTIFS skips text, so \"all rows minus rows with a December 2025 date\" counts every bad entry."),
        Task(f"IntakeLog, CopayAmt ({ic('CopayAmt')}): allow a Decimal between 0 and 100. On the Input Message tab, add the "
             "title Copay and the message \"Enter the amount collected, $0 to $100.\" Leave the Error Alert style as Stop. "
             "How many CopayAmt cells are circled?",
             answer=bad_copay, title="Copay amount rule with an input message (circled entries)",
             hint="Allow: Decimal. Between includes both limits",
             solution=(f"1. Select **{ic('CopayAmt')}**.\n"
                       "2. **Data → Data Validation → Settings**: **Allow** = **Decimal**, **Data** = **between**, "
                       "**Minimum** = 0, **Maximum** = 100.\n"
                       "3. **Input Message** tab: Title `Copay`, message `Enter the amount collected, $0 to $100.`\n"
                       "4. **Error Alert** tab: keep **Style** = **Stop**. Click **OK**.\n"
                       f"5. {circle} You see **{bad_copay}** circles.\n\n"
                       f"Equivalent formula: `=COUNTIF({ir('CopayAmt')},\"<0\")+COUNTIF({ir('CopayAmt')},\">100\")`"),
             live=f"=COUNTIF({ir('CopayAmt')},\"<0\")+COUNTIF({ir('CopayAmt')},\">100\")",
             explanation="\"Between\" includes both limits, so the $0.00 and $100.00 copays in the log are valid. The circled "
                         "values are a negative amount (a refund typed into the wrong field), 2,500 (almost certainly $25.00 "
                         "typed without the decimal point), and 350. Click any cell in the column to see your input message."),
        Task(f"IntakeLog, ZIP ({ic('ZIP')}): allow Text length equal to 5. How many ZIP cells are circled?",
             answer=bad_zip, title="ZIP code length rule (circled entries)",
             hint="Allow: Text length, Data: equal to",
             solution=(f"1. Select **{ic('ZIP')}**.\n"
                       "2. **Data → Data Validation → Settings**: **Allow** = **Text length**, **Data** = **equal to**, "
                       "**Length** = 5. Click **OK**.\n"
                       f"3. {circle} You see **{bad_zip}** circles.\n\n"
                       f"Equivalent formula: `=SUMPRODUCT(--(LEN({ir('ZIP')})<>5))`"),
             live=f"=SUMPRODUCT(--(LEN({ir('ZIP')})<>5))",
             explanation="Text length counts characters, so the 4-digit ZIP, the 6-digit ZIP, and the ZIP+4 entry "
                         "(10 characters including the hyphen) fail. The rule also works on numbers. One ZIP in the log was "
                         f"stored as the number {zip_num} (it's right-aligned), and it passes because it has 5 characters."),
        Task(f"IntakeLog, MRN ({ic('MRN')}): an MRN must be exactly 8 characters, and all of them digits. Select the column "
             "with C2 as the active cell and add a Custom validation rule that is TRUE only for a valid MRN. "
             "How many MRN cells are circled?",
             answer=bad_mrn, title="Custom MRN rule (circled entries)",
             hint="Combine a LEN test with an ISNUMBER(--C2) test inside AND, written for the active cell C2",
             solution=(f"1. Type **{ic('MRN')}** in the Name Box and press **Enter**. The range is selected and **C2** is the active cell.\n"
                       "2. **Data → Data Validation → Settings**: **Allow** = **Custom**, **Formula** = "
                       "`=AND(LEN(C2)=8,ISNUMBER(--C2))`. Click **OK**.\n"
                       f"3. {circle} You see **{bad_mrn}** circles.\n\n"
                       f"Equivalent formula: `=ROWS({ir('MRN')})-SUMPRODUCT((LEN({ir('MRN')})=8)*ISNUMBER(--{ir('MRN')}))`"),
             live=f"=ROWS({ir('MRN')})-SUMPRODUCT((LEN({ir('MRN')})=8)*ISNUMBER(--{ir('MRN')}))",
             explanation=f"LEN catches the 7- and 9-character MRNs and the MRN that was stored as a number (Excel dropped its "
                         f"leading zero, leaving 7 digits). ISNUMBER(--C2) catches {letter_o}, where the letter O "
                         "replaced a zero: the double minus tries to turn the text into a number and fails. Each test misses "
                         "something the other catches, so both go inside AND. You write the formula for C2 only, and Excel "
                         "adjusts it for C3, C4, and so on."),
        Task("Dependent drop-down: on the Lists sheet, name each department column after its facility "
             f"(Formulas → Create from Selection → Top row, one column at a time: {', '.join(block_ranges)}). "
             f"Give IntakeLog Facility ({ic('Facility')}) a List validation from {fac_rng}. Then give Department "
             f"({ic('Department')}) a List validation whose Source is =INDIRECT(SUBSTITUTE($H2,\" \",\"_\")). "
             "How many Department cells are circled?",
             answer=bad_dept, title="Dependent Facility → Department list (circled entries)",
             hint="Names can't contain spaces, so Create from Selection uses underscores",
             solution=("1. On **Lists**, select " + ", then ".join(f"**{b}**" for b in block_ranges[:1]) +
                       " and choose **Formulas → Create from Selection**, tick only **Top row**, and click **OK**. "
                       "Repeat for " + ", ".join(f"**{b}**" for b in block_ranges[1:]) + ". "
                       "**Formulas → Name Manager** (Ctrl + F3) now lists four names such as `Cedar_Ridge_Medical_Center`.\n"
                       f"2. Select **{ic('Facility')}** on IntakeLog → **Data Validation** → **List**, Source `={fac_rng}`.\n"
                       f"3. Select **{ic('Department')}** with **I2** active → **Data Validation** → **List**, Source "
                       "`=INDIRECT(SUBSTITUTE($H2,\" \",\"_\"))`. Click **OK**.\n"
                       f"4. {circle} You see **{bad_dept}** circles.\n\n"
                       f"Equivalent formula (uses the facility–department map in Lists!J:K): "
                       f"`=SUMPRODUCT(--(COUNTIFS({map_f},{ir('Facility')},{map_d},{ir('Department')})=0))`"),
             live=f"=SUMPRODUCT(--(COUNTIFS({map_f},{ir('Facility')},{map_d},{ir('Department')})=0))",
             explanation="Create from Selection turns the header \"Cedar Ridge Medical Center\" into the name "
                         "Cedar_Ridge_Medical_Center, because names can't contain spaces. SUBSTITUTE makes the same change to "
                         "the facility in column H, and INDIRECT turns that text into a reference to the named range. Because "
                         "$H2 has no $ before the row, every row builds its list from its own facility. The circled entries are "
                         "real departments at the wrong facility: Pediatric Clinic, for example, is at the Outpatient Pavilion, "
                         "not at Bluestone Memorial Hospital."),

        # ---------------- Part B: conditional formatting ----------------
        Task(f"Labs sheet: the lab interface re-sent some results, so a few LabResultIDs appear more than once. Select "
             f"{lc('LabResultID')} and apply Highlight Cells Rules → Duplicate Values. How many cells are highlighted?",
             answer=dup_cells, title="Duplicate LabResultIDs (highlighted cells)",
             hint="Home → Conditional Formatting → Highlight Cells Rules",
             solution=(f"1. On **Labs**, select **{lc('LabResultID')}**.\n"
                       "2. **Home → Conditional Formatting → Highlight Cells Rules → Duplicate Values**. Keep **Duplicate** "
                       "and the light red fill, then click **OK**.\n"
                       f"3. Count the highlighted cells (or filter the column by color and read the status bar): **{dup_cells}**.\n\n"
                       f"Equivalent formula: `=SUMPRODUCT(--(COUNTIF({lr('LabResultID')},{lr('LabResultID')})>1))`"),
             live=f"=SUMPRODUCT(--(COUNTIF({lr('LabResultID')},{lr('LabResultID')})>1))",
             explanation=f"Duplicate Values highlights every copy of a repeated value, including the first one. There are "
                         f"{dup_ids} repeated IDs: five appear twice and one appears three times, so {dup_cells} cells light up even though only "
                         f"{extra_rows} rows are extra. When you clean the feed, you delete {extra_rows} rows, not {dup_cells}."),
        Task(f"Labs, TATMin ({lc('TATMin')}): turnaround minutes from specimen collection to result. Apply Top/Bottom Rules → "
             "Top 10 Items. What is the smallest TAT that your rule highlights?",
             answer=tenth, title="Top 10 turnaround times (smallest highlighted value)",
             hint="Ties with the 10th value are highlighted too. LARGE gives the k-th largest",
             solution=(f"1. Select **{lc('TATMin')}**.\n"
                       "2. **Home → Conditional Formatting → Top/Bottom Rules → Top 10 Items**. Keep **10**, click **OK**.\n"
                       "3. Sort or filter by color to read the smallest highlighted value.\n\n"
                       f"Equivalent formula: `=LARGE({lr('TATMin')},10)`"),
             live=f"=LARGE({lr('TATMin')},10)",
             explanation=("Top 10 Items highlights the 10 largest values plus any cell tied with the 10th, so a tie at the "
                          "boundary can light up 11 or 12 cells. "
                          + (f"Here {top_cells} cells are highlighted because several results share the boundary value of "
                             f"{tenth} minutes. " if top_cells > 10 else
                             f"Here there's no tie at the boundary, so exactly 10 cells are highlighted. ")
                          + "Either way, the smallest highlighted value is LARGE(range,10).")),
        Task(f"Supplies, PctOfPar ({sup.rng('PctOfPar', absolute=False, sheet=False)}) = QtyOnHand ÷ ParLevel. Add Data Bars, "
             "then edit the rule so Minimum is Number 0 and Maximum is Number 1. A full bar now means \"stocked to par.\" "
             "How many items show a completely full bar?",
             answer=full_bars, title="Data bars on % of par (full bars)",
             hint="Values at or above the Maximum get a full bar",
             solution=(f"1. On **Supplies**, select **{sup.rng('PctOfPar', absolute=False, sheet=False)}**.\n"
                       "2. **Home → Conditional Formatting → Data Bars** and pick any fill.\n"
                       "3. **Conditional Formatting → Manage Rules → Edit Rule**. Set **Minimum** Type = **Number**, Value = 0, "
                       "and **Maximum** Type = **Number**, Value = 1. Click **OK** twice.\n"
                       f"4. Count the full bars: **{full_bars}**.\n\n"
                       f"Equivalent formula: `=COUNTIF({sr('PctOfPar')},\">=1\")`"),
             live=f"=COUNTIF({sr('PctOfPar')},\">=1\")",
             explanation=f"With the default Automatic settings, the longest bar belongs to the largest value (about "
                         f"{max_par:.0%} of par), so an item stocked exactly to par looks only about three-quarters full. Fixing "
                         "the Maximum at 1 gives bar length a meaning: every item at or above 100% of par gets a full bar, and "
                         "an item at 50% gets a half bar."),
        Task(f"Census sheet (Medical-Surgical 5 East, Nov–Dec 2025): apply the Red - Yellow - Green Color Scale to Occupancy "
             f"({cen.rng('Occupancy', absolute=False, sheet=False)}) so the fullest days are red. In Manage Rules → Edit Rule "
             "you'll see the midpoint is the 50th percentile. Which occupancy gets the pure yellow midpoint color? Enter it as "
             "a percentage to 1 decimal place.",
             answer=occ_median, fmt="0.0%", title="Color scale midpoint (50th percentile)",
             hint="The 50th percentile has a more common name",
             solution=(f"1. On **Census**, select **{cen.rng('Occupancy', absolute=False, sheet=False)}**.\n"
                       "2. **Home → Conditional Formatting → Color Scales → Red - Yellow - Green Color Scale**. The first color "
                       "in the name goes to the highest values, so high occupancy is red.\n"
                       "3. **Conditional Formatting → Manage Rules → Edit Rule** shows Minimum = Lowest Value, "
                       "Midpoint = Percentile 50, Maximum = Highest Value.\n\n"
                       f"Equivalent formula: `=MEDIAN({rr(cen, 'Occupancy')})`"),
             live=f"=MEDIAN({rr(cen, 'Occupancy')})",
             explanation="The 50th percentile is the median, the middle value when you sort the days by occupancy. "
                         + (f"Here {median_days} days have exactly that occupancy, so they show pure yellow. "
                            if median_days > 1 else "Exactly one day has that value and shows pure yellow. ")
                         + "So the default color scale's middle color marks a typical day for this unit, not a target. "
                           "To color against a target such as 85%, change the midpoint Type to Number (you'll do that in the bonus)."),
        Task(f"Labs, TATMin: apply Icon Sets → 3 Traffic Lights (Unrimmed). Edit the rule: click Reverse Icon Order, set both "
             "Types to Number, and make red show when the value is >= 60 and yellow when it is >= 45 (green below 45). "
             "How many cells show a yellow light?",
             answer=yellow, title="Traffic-light icons on turnaround time (yellow count)",
             hint="Each icon's test is \">=\". The default Type is Percent, not Number",
             solution=(f"1. Select **{lc('TATMin')}** → **Home → Conditional Formatting → Icon Sets → 3 Traffic Lights (Unrimmed)**.\n"
                       "2. **Manage Rules → Edit Rule**. Click **Reverse Icon Order** so red is on top.\n"
                       "3. Red: **>=**, Value **60**, Type **Number**. Yellow: **>=**, Value **45**, Type **Number**. "
                       "Green covers everything below 45. Click **OK** twice.\n\n"
                       f"Equivalent formula: `=COUNTIFS({lr('TATMin')},\">=45\",{lr('TATMin')},\"<60\")`"),
             live=f"=COUNTIFS({lr('TATMin')},\">=45\",{lr('TATMin')},\"<60\")",
             explanation="The default thresholds are Percent 67 and 33. They split the span between the lowest and highest TAT "
                         "into thirds, so they don't mean minutes at all. Switching the Type to Number makes the thresholds "
                         "literal. Reverse Icon Order puts red on the high (slow) values. Yellow covers 45 ≤ TAT < 60 because "
                         f"each icon's test is \">=\" and the red test is checked first. ({red_icons} results are red.)"),
        Task(f"Labs: highlight the entire row of every critical result (AbnormalFlag HH or LL). Select A2:{last_lab_col}{lab.last_row} "
             "with A2 active, choose New Rule → \"Use a formula to determine which cells to format,\" and set a red fill. "
             "How many rows are highlighted?",
             answer=crit_rows, title="Critical results: whole-row formula rule",
             hint="Lock the column with $, not the row",
             solution=(f"1. On **Labs**, type **A2:{last_lab_col}{lab.last_row}** in the Name Box and press **Enter**. "
                       "The rows are selected and A2 is the active cell.\n"
                       "2. **Home → Conditional Formatting → New Rule → Use a formula to determine which cells to format**.\n"
                       "3. Formula: `=OR($K2=\"HH\",$K2=\"LL\")`. Click **Format → Fill**, pick red, then **OK** twice.\n\n"
                       f"Equivalent formula: `=COUNTIF({lr('AbnormalFlag')},\"HH\")+COUNTIF({lr('AbnormalFlag')},\"LL\")`"),
             live=f"=COUNTIF({lr('AbnormalFlag')},\"HH\")+COUNTIF({lr('AbnormalFlag')},\"LL\")",
             explanation=f"You write the rule for the active cell's row (row 2), and Excel shifts it for every other cell in the "
                         "Applies-to range. $K keeps every cell in the row looking at column K, and the unlocked 2 lets each row "
                         f"check its own flag. Without the $, cell B2 would test L2 and the rule breaks. There are {n_hh} HH and "
                         f"{n_ll} LL results, so {crit_rows} rows turn red."),
        Task(f"Supplies: add two formula rules to A2:M{sup.last_row}. Rule 1 (red fill): the item has expired, meaning ExpirationDate is not "
             "blank and is on or before ReportDate. Rule 2 (amber fill): ExpirationDate is not blank and is on or before "
             "ReportDate + ExpiringWindowDays. ReportDate (12/31/2025) and ExpiringWindowDays (90) are named cells on the Lists "
             "sheet. Excel puts each new rule at the top of the list, so open Manage Rules and move the red rule above the "
             "amber rule. How many rows are amber?",
             answer=amber, title="Expired vs expiring supplies: two rules and rule order",
             hint="A blank cell counts as 0, which is \"on or before\" any date. The rule higher in Manage Rules wins the fill",
             solution=(f"1. On **Supplies**, select **A2:M{sup.last_row}** with A2 active.\n"
                       "2. **New Rule → Use a formula**: `=AND($L2<>\"\",$L2<=ReportDate)` with a red fill → **OK**.\n"
                       "3. With the same range selected, add a second rule: `=AND($L2<>\"\",$L2<=ReportDate+ExpiringWindowDays)` "
                       "with an amber (light orange) fill → **OK**.\n"
                       "4. **Conditional Formatting → Manage Rules**, set **Show formatting rules for: This Worksheet**. The amber "
                       "rule is on top because Excel adds each new rule at the top of the list. Select the red rule and click "
                       "**▲** (Move Up) so it sits above the amber rule, then click **OK**.\n\n"
                       f"Equivalent formula: `=COUNTIFS({sr('ExpirationDate')},\">\"&ReportDate,{sr('ExpirationDate')},"
                       f"\"<=\"&ReportDate+ExpiringWindowDays)`"),
             live=(f"=COUNTIFS({sr('ExpirationDate')},\">\"&ReportDate,{sr('ExpirationDate')},"
                   f"\"<=\"&ReportDate+ExpiringWindowDays)"),
             explanation=f"Rule 2 is TRUE for {expiring_all} rows, because every expired item also expires before ReportDate + 90. "
                         f"The red rule sits above it and both rules set a fill, so red wins on the {expired} expired rows and "
                         f"{amber} rows stay amber. If you skip the move, the newer amber rule stays on top and you'd see "
                         f"{expiring_all} amber rows and no red. The $L2<>\"\" "
                         f"guard matters too: an empty cell counts as 0 (the date 1/0/1900), which is \"on or before ReportDate,\" "
                         f"so without the guard all {blanks} non-perishable items would turn red."),
    ]

    # ------------------------------------------------------------------ bonus: bed huddle board
    H_FIRST = 12
    H_LAST = H_FIRST + len(board) - 1
    hb = lambda col: f"Huddle!{col}{H_FIRST}:{col}{H_LAST}"
    occ_b = [r["Today"] / r["Beds"] for r in board]
    above85 = sum(1 for x in occ_b if x > 0.85)
    up = sum(1 for r in board if r["Today"] - r["Yesterday"] >= 1)
    flat = sum(1 for r in board if r["Today"] - r["Yesterday"] == 0)
    down = sum(1 for r in board if r["Today"] - r["Yesterday"] < 0)
    overflow = sum(1 for r in board if r["Today"] > r["Beds"])
    amber_b = sum(1 for r, o in zip(board, occ_b) if o >= 0.9 and r["Today"] <= r["Beds"])
    at90 = sum(1 for o in occ_b if abs(o - 0.9) < 1e-12)
    pick_fac, pick_unit = "Cedar Ridge Medical Center", "Intensive Care Unit"
    sel = [r for r in nov if r["Facility"] == pick_fac and r["Unit"] == pick_unit]
    card_occ = sum(r["MidnightCensus"] for r in sel) / sum(r["StaffedBeds"] for r in sel)
    nr = lambda col: f"NovCensus!${novsd.col(col)}${novsd.first_row}:${novsd.col(col)}${novsd.last_row}"
    unit_src = f"=OFFSET($B${H_FIRST - 1},MATCH($B$4,$A${H_FIRST}:$A${H_LAST},0),0,COUNTIF($A${H_FIRST}:$A${H_LAST},$B$4),1)"

    L.bonus_title = "Bonus: The 7 a.m. bed huddle board"
    L.bonus_scenario = (
        "It's 7 a.m. on Wednesday, November 26, 2025, the day before Thanksgiving. The house supervisor runs the system bed "
        "huddle from the Huddle sheet: one row per inpatient unit, with the midnight census for 11/24 and 11/25, the "
        "overnight change, and occupancy. Make the board readable at a glance, then build the unit selector that drives the "
        f"gray unit card (B6:B9). The card reads the NovCensus sheet (every unit, Nov 1–25). The board's data is in rows "
        f"{H_FIRST}–{H_LAST}.")
    L.bonus = [
        Task(f"Apply a 3-color scale to Occupancy ({hb('G').replace('Huddle!', '')}): Minimum = Lowest Value (green), Midpoint = "
             "Number 0.85 (yellow, the planning target), Maximum = Highest Value (red). How many units are shaded on the red "
             "side of yellow (occupancy above 85%)?",
             answer=above85, title="Color scale with an 85% target midpoint",
             hint="New Rule → Format all cells based on their values → 3-Color Scale",
             solution=(f"1. Select **{hb('G').replace('Huddle!', '')}** on Huddle → **Conditional Formatting → New Rule → "
                       "Format all cells based on their values**.\n"
                       "2. **Format Style** = **3-Color Scale**. Minimum: Lowest Value, green. Midpoint: Type **Number**, "
                       "Value **0.85**, yellow. Maximum: Highest Value, red. Click **OK**.\n\n"
                       f"Equivalent formula: `=COUNTIF({hb('G')},\">0.85\")`"),
             live=f"=COUNTIF({hb('G')},\">0.85\")",
             explanation="With a Number midpoint, yellow means \"exactly at target\" instead of \"a typical unit.\" Every unit "
                         "above 85% shades from yellow toward red, and the reddest cell is simply the fullest unit."),
        Task(f"Apply Icon Sets → 3 Arrows (Colored) to Change ({hb('F').replace('Huddle!', '')}). Edit the rule so both Types are "
             "Number: up arrow when the value is >= 1, sideways arrow when it is >= 0, down arrow otherwise. How many units show "
             "an up arrow?",
             answer=up, title="Arrows on the overnight census change",
             hint="Type = Number, not the default Percent",
             solution=(f"1. Select **{hb('F').replace('Huddle!', '')}** → **Conditional Formatting → Icon Sets → 3 Arrows (Colored)**.\n"
                       "2. **Manage Rules → Edit Rule**: green up arrow **>=** 1 **Number**, yellow sideways arrow **>=** 0 "
                       "**Number**. The red down arrow covers negative changes. Click **OK** twice.\n\n"
                       f"Equivalent formula: `=COUNTIF({hb('F')},\">=1\")`"),
             live=f"=COUNTIF({hb('F')},\">=1\")",
             explanation=f"Census rose overnight on {up} units, held steady on {flat}, and fell on {down}. Number thresholds keep "
                         "the arrows honest on any day. The default Percent thresholds depend on the day's smallest and largest "
                         "change, so a unit that gained one patient could show a sideways arrow on a busier day."),
        Task(f"Add two formula rules to the whole board (A{H_FIRST}:G{H_LAST}): a red fill when the 11/25 census is above "
             "staffed beds (overflow), and an amber fill when occupancy is at least 90%. In Manage Rules, put the overflow rule "
             "above the amber rule. How many rows end up amber?",
             answer=amber_b, title="Overflow (red) above near-capacity (amber) row rules",
             hint="Both rules are TRUE for an overflow unit, and the rule on top wins the fill",
             solution=(f"1. Select **A{H_FIRST}:G{H_LAST}** with A{H_FIRST} active.\n"
                       f"2. **New Rule → Use a formula**: `=$E{H_FIRST}>$C{H_FIRST}`, red fill.\n"
                       f"3. **New Rule → Use a formula**: `=$G{H_FIRST}>=0.9`, amber fill.\n"
                       "4. **Manage Rules** (This Worksheet): Excel added the amber rule at the top because it's newer. Select the "
                       "overflow rule and click **▲** so it's on top.\n\n"
                       f"Equivalent formula: `=SUMPRODUCT(({hb('G')}>=0.9)*({hb('E')}<={hb('C')}))`"),
             live=f"=SUMPRODUCT(({hb('G')}>=0.9)*({hb('E')}<={hb('C')}))",
             explanation=f"There are {overflow} overflow units, and both rules are TRUE for them. The overflow rule is on top, so "
                         f"those rows stay red, and {amber_b} rows are amber. Two of those amber units sit at exactly 90.0%, "
                         f"so writing > 0.9 instead of >= 0.9 would lose them. With the rules in the wrong order every "
                         f"overflow row would turn amber too."
                         if at90 >= 2 else
                         f"There are {overflow} overflow units, and both rules are TRUE for them. The overflow rule is on top, so "
                         f"those rows stay red, and {amber_b} rows are amber."),
        Task("Build the unit selector. Give B4 a List validation from the Hospitals list (I12:I14). Give B5 a dependent List "
             "built from the board's own Facility and Unit columns, with no named ranges. Then choose Cedar Ridge Medical "
             "Center → Intensive Care Unit. The gray cell shows the card's Nov 1–25 occupancy. What is it?",
             answer=card_occ, fmt="0.0%", title="Dependent unit selector driving the card",
             hint="OFFSET(start, rows down, 0, height, 1) with MATCH for the first row and COUNTIF for the height",
             solution=("1. Select **B4** on Huddle → **Data Validation → List**, Source `=$I$12:$I$14`.\n"
                       f"2. Select **B5** → **Data Validation → List**, Source:\n\n"
                       f"   `{unit_src}`\n\n"
                       "   MATCH finds the facility's first row on the board, and COUNTIF counts its units. The board is sorted "
                       "by facility, so OFFSET returns exactly that facility's block of unit names. If B4 is still empty, "
                       "Excel warns that the source currently evaluates to an error. Click **Yes**, or pick a facility in B4 "
                       "first.\n"
                       "3. Pick **Cedar Ridge Medical Center** in B4, then **Intensive Care Unit** in B5. Read the card.\n\n"
                       "If you prefer INDIRECT, copy each hospital's units into its own column under the hospital's name (like "
                       "the Lists sheet in task 6), name the columns with Create from Selection, and use "
                       "`=INDIRECT(SUBSTITUTE($B$4,\" \",\"_\"))`. The result is the same."),
             summary='=IF(Huddle!$B$8="","",Huddle!$B$8)',
             fill={"range": "Huddle!B4:B5", "values": [pick_fac, pick_unit]},
             live=(f"=SUMIFS({nr('MidnightCensus')},{nr('Facility')},\"{pick_fac}\",{nr('Unit')},\"{pick_unit}\")"
                   f"/SUMIFS({nr('StaffedBeds')},{nr('Facility')},\"{pick_fac}\",{nr('Unit')},\"{pick_unit}\")"),
             explanation="Three hospitals each have an Intensive Care Unit, so the unit name alone is ambiguous. The dependent list "
                         "forces a facility first and then offers only that facility's units, and the card's SUMIFS uses both "
                         "cells. The OFFSET version needs no named ranges, so it keeps working when you add a unit to the board "
                         "(as long as the board stays sorted by facility)."),
    ]

    L.start_notes = [
        "Part A of the practice uses IntakeLog and Lists. Part B uses Labs, Supplies, and Census. The bonus uses Huddle and NovCensus.",
        "ReportDate (12/31/2025) and ExpiringWindowDays (90) are named cells on the Lists sheet. Use the names in your rules.",
        "Circle Invalid Data circles disappear when you save or close the file. They are a checking tool, not formatting.",
    ]

    # ------------------------------------------------------------------ customize: Lists, Huddle, Huddle Key
    @L.customize
    def _custom(wb, lesson, selftest):
        bold_white = Font(bold=True, color="FFFFFF")

        # ---------------- Lists ----------------
        ws = wb.create_sheet("Lists")
        ws.sheet_properties.tabColor = "7F7F7F"

        def header(cell, text):
            c = ws[cell]
            c.value = text
            c.font = bold_white
            c.fill = HEADER_FILL
            c.border = BOX
            c.alignment = Alignment(wrap_text=False)

        header("A1", "Payers")
        for i, p in enumerate(payer_names, 2):
            ws.cell(row=i, column=1, value=p).border = BOX
        header("C1", "Facilities")
        for i, f in enumerate(fac_names, 2):
            ws.cell(row=i, column=3, value=f).border = BOX
        for col, f in zip(block_cols, fac_names):
            header(f"{col}1", f)
            for i, d in enumerate(reg[f], 2):
                ws[f"{col}{i}"] = d
                ws[f"{col}{i}"].border = BOX
        header("J1", "Facility")
        header("K1", "Department")
        for i, (f, d) in enumerate(pairs, 2):
            ws[f"J{i}"] = f
            ws[f"K{i}"] = d
        header("M1", "Setting")
        header("N1", "Value")
        ws["M2"], ws["N2"] = "ReportDate", REPORT_DATE
        ws["N2"].number_format = "mm/dd/yyyy"
        ws["M3"], ws["N3"] = "ExpiringWindowDays", EXPIRING_WINDOW
        for c in ("M2", "N2", "M3", "N3"):
            ws[c].border = BOX
        ws["M5"] = "Named cells: ReportDate = Lists!$N$2, ExpiringWindowDays = Lists!$N$3"
        ws["M5"].font = Font(italic=True, color="595959")
        ws["E19"] = ("Departments by facility, one column per facility. Task 6: name each column with Formulas → Create from "
                     "Selection (Top row).")
        ws["E19"].font = Font(italic=True, color="595959")
        ws["J34"] = "The same pairs in long form (used by the answer key's check formula)."
        ws["J34"].font = Font(italic=True, color="595959")
        for col, w in {"A": 30, "B": 3, "C": 32, "D": 3, "E": 30, "F": 32, "G": 30, "H": 32, "I": 3, "J": 32, "K": 32,
                       "L": 3, "M": 22, "N": 12}.items():
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        wb.defined_names["ReportDate"] = DefinedName("ReportDate", attr_text="Lists!$N$2")
        wb.defined_names["ExpiringWindowDays"] = DefinedName("ExpiringWindowDays", attr_text="Lists!$N$3")

        # ---------------- Huddle (learner) and Huddle Key (finished, hidden) ----------------
        def huddle(name, finished):
            hs = wb.create_sheet(name)
            hs.sheet_properties.tabColor = "C00000" if finished else "BF9000"
            hs["A1"] = "7 a.m. Bed Huddle · Wednesday, November 26, 2025" + (" (finished key)" if finished else "")
            hs["A1"].font = Font(bold=True, size=16, color=NAVY)
            hs["A2"] = ("Midnight census for Tuesday 11/25 vs Monday 11/24, every inpatient unit in the system. "
                        "Yellow cells are yours, and gray cells hold formulas.")
            hs["A2"].font = Font(italic=True, color="595959")
            hs["A3"] = "UNIT SELECTOR"
            hs["A3"].font = bold_white
            hs["A3"].fill = HEADER_FILL
            hs["B3"].fill = HEADER_FILL
            labels = ["Facility", "Unit", "Staffed beds", "Census 11/25", "Occupancy, Nov 1–25", "Overflow days, Nov 1–25"]
            for i, lab_ in enumerate(labels, 4):
                hs.cell(row=i, column=1, value=lab_).font = Font(bold=True)
                hs.cell(row=i, column=1).border = BOX
                c = hs.cell(row=i, column=2)
                c.border = INPUT_BORDER if i <= 5 else BOX
                c.fill = INPUT_FILL if i <= 5 else PREFILL_FILL
                c.alignment = Alignment(horizontal="left")
            sel_ = 'OR($B$4="",$B$5="")'
            ok_ = f"COUNTIFS($A${H_FIRST}:$A${H_LAST},$B$4,$B${H_FIRST}:$B${H_LAST},$B$5)=0"
            crit = f"$A${H_FIRST}:$A${H_LAST},$B$4,$B${H_FIRST}:$B${H_LAST},$B$5"
            ncrit = f"{nr('Facility')},$B$4,{nr('Unit')},$B$5"
            lesson.set_formula(hs, "B6", f'=IF({sel_},"",IF({ok_},"n/a",SUMIFS($C${H_FIRST}:$C${H_LAST},{crit})))', dynamic=False)
            lesson.set_formula(hs, "B7", f'=IF({sel_},"",IF({ok_},"n/a",SUMIFS($E${H_FIRST}:$E${H_LAST},{crit})))', dynamic=False)
            lesson.set_formula(hs, "B8", f'=IF({sel_},"",IFERROR(SUMIFS({nr("MidnightCensus")},{ncrit})/SUMIFS({nr("StaffedBeds")},{ncrit}),"n/a"))',
                               dynamic=False)
            lesson.set_formula(hs, "B9", f'=IF({sel_},"",IF(COUNTIFS({ncrit})=0,"n/a",COUNTIFS({ncrit},{nr("Occupancy")},">1")))',
                               dynamic=False)
            hs["B8"].number_format = "0.0%"
            for i, h in enumerate(["Facility", "Unit", "Staffed beds", "Census 11/24", "Census 11/25", "Change", "Occupancy"], 1):
                c = hs.cell(row=H_FIRST - 1, column=i, value=h)
                c.font = bold_white
                c.fill = HEADER_FILL
                c.border = BOX
                c.alignment = Alignment(horizontal="center", wrap_text=True)
            for k, r in enumerate(board):
                row = H_FIRST + k
                vals = [r["Facility"], r["Unit"], r["Beds"], r["Yesterday"], r["Today"]]
                for j, v in enumerate(vals, 1):
                    hs.cell(row=row, column=j, value=v).border = BOX
                hs[f"F{row}"] = f"=E{row}-D{row}"
                hs[f"G{row}"] = f"=E{row}/C{row}"
                hs[f"G{row}"].number_format = "0.0%"
                hs[f"F{row}"].number_format = "+0;-0;0"
                for col in "FG":
                    hs[f"{col}{row}"].border = BOX
            hs[f"I{H_FIRST - 1}"] = "Hospitals"
            hs[f"I{H_FIRST - 1}"].font = bold_white
            hs[f"I{H_FIRST - 1}"].fill = HEADER_FILL
            hospitals = list(dict.fromkeys(r["Facility"] for r in board))
            for k, hname in enumerate(hospitals):
                hs[f"I{H_FIRST + k}"] = hname
                hs[f"I{H_FIRST + k}"].border = BOX
            for col, w in {"A": 32, "B": 30, "C": 9, "D": 9, "E": 9, "F": 9, "G": 11, "H": 3, "I": 32}.items():
                hs.column_dimensions[col].width = w
            hs.row_dimensions[H_FIRST - 1].height = 30
            hs.freeze_panes = hs[f"A{H_FIRST}"]
            if finished:
                rng_g, rng_f, rng_all = f"G{H_FIRST}:G{H_LAST}", f"F{H_FIRST}:F{H_LAST}", f"A{H_FIRST}:G{H_LAST}"
                hs.conditional_formatting.add(rng_all, FormulaRule(formula=[f"$E{H_FIRST}>$C{H_FIRST}"], fill=RED_FILL, stopIfTrue=False))
                hs.conditional_formatting.add(rng_all, FormulaRule(formula=[f"$G{H_FIRST}>=0.9"], fill=AMBER_FILL))
                hs.conditional_formatting.add(rng_g, ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="num",
                                                                    mid_value=0.85, mid_color="FFEB84", end_type="max",
                                                                    end_color="F8696B"))
                icon = IconSet(iconSet="3Arrows", cfvo=[FormatObject(type="percent", val=0), FormatObject(type="num", val=0),
                                                        FormatObject(type="num", val=1)], showValue=None, reverse=None)
                hs.conditional_formatting.add(rng_f, Rule(type="iconSet", iconSet=icon))
                dv1 = DataValidation(type="list", formula1=f"$I${H_FIRST}:$I${H_FIRST + len(hospitals) - 1}", allow_blank=True)
                dv2 = DataValidation(type="list", formula1=unit_src[1:], allow_blank=True)
                hs.add_data_validation(dv1)
                hs.add_data_validation(dv2)
                dv1.add("B4")
                dv2.add("B5")
                hs["B4"], hs["B5"] = pick_fac, pick_unit
                hs.sheet_state = "hidden"
            return hs

        huddle("Huddle", finished=False)
        huddle("Huddle Key", finished=True)

        # ---------------- Answer keys: show Markdown solutions as plain text in Excel ----------------
        for key_name in (lesson.key_sheet, lesson.bonus_key_sheet):
            ks = wb[key_name]
            for row in range(5, ks.max_row + 1):
                v = ks.cell(row=row, column=4).value
                if isinstance(v, str) and not v.startswith("="):
                    ks.cell(row=row, column=4).value = _plain(v)

    return L
