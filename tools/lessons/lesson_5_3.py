"""Lesson 5.3 · VBA: Ranges, Worksheets & Workbooks.

Data: 2,000 encounters admitted in 2025 at all four Bluestone facilities (every encounter type), picked evenly
across the year from the 2025 extract (sorted by AdmitDateTime, then every ~5.6th row). The Encounters sheet is a
PLAIN RANGE on purpose, like an EHR export, so the End(xlUp) / CurrentRegion / AutoFilter idioms behave exactly
as they do on real exports. A tiny Facilities sheet is an Excel Table (tblFacilities) for the ListObject basics.

Files written next to the workbook (ASCII, CRLF so the VBE imports them cleanly):
  starter/EncounterMacros.bas            LastRow + DeleteSheetIfExists helpers, TODO stubs for tasks 7-12 and the bonus
  starter/Snippets.bas                   the five predict-the-output snippets (tasks 1-6)
  solutions/EncounterMacros_Solution.bas reference macros for tasks 7-12 (spoilers)
  solutions/MonthlyPackets_Solution.bas  reference macro for the bonus (spoilers)

Checks read what the macros create:
  * sheets a macro adds (TypeSummary, F01-F04, 2025-01..2025-12) are read through INDIRECT and guarded with
    ISREF(INDIRECT("'Name'!A1")), so the pristine file never references a missing sheet and the gray cells stay blank
    until the sheet exists (COUNTA(#REF!) would count the error as 1 in Excel, so the ISREF guard matters);
  * Output!B5 / B7 / B9 hold values the macros write;
  * the LOSDays column on Encounters is summed.
The self-test simulates the macros: a customize hook builds the sheets the macros would create (from the same Python
computations as the answers) and `fill` writes the Output cells / LOSDays column.

Smoke test of the reference macros in LibreOffice (run from tools/):
  python3 -c "import lessons.lesson_5_3 as m; m.smoke_test()"
"""
from __future__ import annotations

from collections import OrderedDict, defaultdict
from datetime import datetime

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

from xlcourse import Lesson, Task, data

CODE = "5.3"
N_SAMPLE = 2000
ENC_COLS = ["EncounterID", "PatientID", "EncounterType", "FacilityID", "DeptID", "AdmitSource", "AdmitDateTime",
            "DischargeDateTime", "PrimaryDxCode", "PayerID", "TotalCharges"]
LOS_COL = "LOSDays"
FAC_COLS = ["FacilityID", "FacilityName", "City", "FacilityType"]
TYPES = ["Emergency", "Inpatient", "Observation", "Outpatient"]
SUMMARY_SHEET = "TypeSummary"
EXPORT_FAC = "F02"
EXPORT_FILE = "F02_encounters_2025.xlsx"
SORT_CHECK_FAC = "F03"
OUT = {"rows": "B4", "charges": "B5", "file": "B6", "export": "B7", "best": "B9"}
PACKET_HEADER_ROW, PACKET_FIRST = 6, 7
PACKET_CHECK = {"count": 3, "total": 7, "top": 11}     # months checked by bonus parts B1-B3

DARK = "1F4E79"
HDR_FILL = PatternFill("solid", fgColor=DARK)
OUT_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------------------
# VBA source (ASCII only). One source of truth for the .bas files, the
# Snippets sheet, the README answer key and the LibreOffice smoke test.
# ---------------------------------------------------------------------------
SNIPPETS = OrderedDict()
SNIPPETS["A"] = dict(tasks="task 1", code='''Sub SnippetA_OffsetResize()
    Dim ws As Worksheet
    Set ws = ThisWorkbook.Worksheets("Encounters")
    Debug.Print ws.Range("C2").Offset(3, 4).Resize(5).Address
End Sub''')
SNIPPETS["B"] = dict(tasks="task 2", code='''Sub SnippetB_DataBody()
    Dim ws As Worksheet, block As Range
    Set ws = ThisWorkbook.Worksheets("Encounters")
    Set block = ws.Range("E50").CurrentRegion            ' any cell inside the data works
    Set block = block.Offset(1).Resize(block.Rows.Count - 1)
    Debug.Print block.Address(False, False)
End Sub''')
SNIPPETS["C"] = dict(tasks="tasks 3 and 4", code='''Sub SnippetC_LastRow()
    Dim ws As Worksheet
    Set ws = ThisWorkbook.Worksheets("Encounters")
    Debug.Print ws.Cells(ws.Rows.Count, "A").End(xlUp).Row     ' line 1: EncounterID
    Debug.Print ws.Cells(ws.Rows.Count, "F").End(xlUp).Row     ' line 2: AdmitSource
    Debug.Print ws.Range("F1").End(xlDown).Row                 ' line 3: AdmitSource again
End Sub''')
SNIPPETS["D"] = dict(tasks="task 5", code='''Sub SnippetD_WhichSheet()
    Dim ws As Worksheet
    Set ws = ThisWorkbook.Worksheets("Encounters")
    ThisWorkbook.Worksheets("Practice").Activate               ' Practice is now the active sheet
    Debug.Print ws.Range(Cells(2, 1), Cells(2, 11)).Address    ' row 2, columns A to K
End Sub''')
SNIPPETS["E"] = dict(tasks="task 6", code='''Sub SnippetE_DeleteOld()
    Application.DisplayAlerts = False
    ThisWorkbook.Worksheets("TypeSummary_2024").Delete         ' last year's summary sheet
    Application.DisplayAlerts = True
    Debug.Print "Old summary deleted"
End Sub''')

COLUMN_MAP = """' Encounters sheet (row 1 = headers, data from row 2, no blank rows):
'    1 A EncounterID      5 E DeptID              9 I PrimaryDxCode
'    2 B PatientID        6 F AdmitSource        10 J PayerID
'    3 C EncounterType    7 G AdmitDateTime      11 K TotalCharges
'    4 D FacilityID       8 H DischargeDateTime  12 L LOSDays (task 12 fills it)
'
' Facilities sheet: Excel Table tblFacilities (FacilityID, FacilityName, City, FacilityType)
' Output sheet: B4 and B5 SortAndReconcile, B6 and B7 ExportFacility,
'   B9 BuildMonthlyPackets (bonus, formatted as Text)."""

LAST_ROW = '''Function LastRow(ws As Worksheet, Optional col As Variant = 1) As Long
    ' The last filled row in one column: start at the bottom of the sheet and
    ' jump up, like Ctrl + Up arrow. col can be a letter ("A") or a number (1).
    LastRow = ws.Cells(ws.Rows.Count, col).End(xlUp).Row
End Function'''

DELETE_IF = '''Sub DeleteSheetIfExists(ByVal sheetName As String)
    ' Deletes the sheet if it's there, without Excel's "are you sure?" prompt.
    ' ByVal lets you pass any text: a String, a Variant, or a cell's value.
    If SheetExists(sheetName) Then
        Application.DisplayAlerts = False
        ThisWorkbook.Worksheets(sheetName).Delete
        Application.DisplayAlerts = True
    End If
End Sub'''

SOL = OrderedDict()
SOL["SheetExists"] = '''Function SheetExists(ByVal sheetName As String) As Boolean
    ' True if THIS workbook has a worksheet with that name (ignoring case,
    ' because Excel treats "typesummary" and "TypeSummary" as the same name).
    Dim ws As Worksheet
    For Each ws In ThisWorkbook.Worksheets
        If StrComp(ws.Name, sheetName, vbTextCompare) = 0 Then
            SheetExists = True
            Exit Function
        End If
    Next ws
End Function'''
SOL["BuildTypeSummary"] = '''Sub BuildTypeSummary()
    Dim wsEnc As Worksheet, wsSum As Worksheet
    Dim rngType As Range, rngCharges As Range
    Dim types As Variant
    Dim lastR As Long, i As Long

    Set wsEnc = ThisWorkbook.Worksheets("Encounters")
    lastR = LastRow(wsEnc, "A")
    Set rngType = wsEnc.Range("C2:C" & lastR)          ' EncounterType
    Set rngCharges = wsEnc.Range("K2:K" & lastR)       ' TotalCharges

    ' Delete-and-recreate: the macro can run again and again
    DeleteSheetIfExists "TypeSummary"
    Set wsSum = ThisWorkbook.Worksheets.Add(After:=ThisWorkbook.Worksheets(ThisWorkbook.Worksheets.Count))
    wsSum.Name = "TypeSummary"

    wsSum.Range("A1:C1").Value = Array("EncounterType", "Encounters", "TotalCharges")
    types = Array("Emergency", "Inpatient", "Observation", "Outpatient")
    For i = 0 To UBound(types)                          ' Array() numbers its items from 0
        With wsSum.Cells(i + 2, 1)                      ' rows 2 to 5, column A
            .Value = types(i)
            .Offset(0, 1).Value = Application.WorksheetFunction.CountIf(rngType, types(i))
            .Offset(0, 2).Value = Application.WorksheetFunction.SumIf(rngType, types(i), rngCharges)
        End With
    Next i

    wsSum.Range("A1:C1").Font.Bold = True
    wsSum.Range("C2:C5").NumberFormat = "#,##0.00"
    wsSum.Columns("A:C").AutoFit
End Sub'''
SOL["SplitByFacility"] = '''Sub SplitByFacility()
    Dim wsEnc As Worksheet, wsNew As Worksheet
    Dim lo As ListObject
    Dim idCell As Range, block As Range
    Dim facID As String

    Set wsEnc = ThisWorkbook.Worksheets("Encounters")
    Set lo = ThisWorkbook.Worksheets("Facilities").ListObjects("tblFacilities")
    If wsEnc.AutoFilterMode Then wsEnc.AutoFilterMode = False   ' start with no filter
    Set block = wsEnc.Range("A1").CurrentRegion                 ' header + every data row

    Application.ScreenUpdating = False
    On Error GoTo CleanUp                                       ' always restore settings

    For Each idCell In lo.ListColumns("FacilityID").DataBodyRange
        facID = idCell.Value
        DeleteSheetIfExists facID
        Set wsNew = ThisWorkbook.Worksheets.Add(After:=ThisWorkbook.Worksheets(ThisWorkbook.Worksheets.Count))
        wsNew.Name = facID

        block.AutoFilter Field:=4, Criteria1:=facID               ' column D = FacilityID
        block.SpecialCells(xlCellTypeVisible).Copy Destination:=wsNew.Range("A1")
        wsNew.Range("A1").CurrentRegion.Columns.AutoFit
    Next idCell

CleanUp:
    wsEnc.AutoFilterMode = False                                ' leave the export unfiltered
    Application.ScreenUpdating = True
    If Err.Number <> 0 Then MsgBox "SplitByFacility stopped: " & Err.Description, vbExclamation
End Sub'''
SOL["SortAndReconcile"] = '''Sub SortAndReconcile()
    Dim ws As Worksheet
    Dim lastR As Long, rowTotal As Long, sheetCount As Long
    Dim chargeTotal As Double

    For Each ws In ThisWorkbook.Worksheets
        If ws.Name Like "F0#" Then                  ' F01 to F09 (# = any one digit)
            ' Task 9: largest charge first. One cell is enough to name the sort column.
            ws.Range("A1").CurrentRegion.Sort Key1:=ws.Range("K1"), Order1:=xlDescending, Header:=xlYes

            ' Task 10: add this sheet to the control totals
            lastR = LastRow(ws, "A")
            rowTotal = rowTotal + (lastR - 1)       ' minus the header row
            chargeTotal = chargeTotal + Application.WorksheetFunction.Sum(ws.Range("K2:K" & lastR))
            sheetCount = sheetCount + 1
        End If
    Next ws

    With ThisWorkbook.Worksheets("Output")
        .Range("B4").Value = rowTotal
        .Range("B5").Value = chargeTotal
    End With
    MsgBox sheetCount & " facility sheets: " & rowTotal & " rows, " & _
           Format(chargeTotal, "#,##0.00") & " in charges.", vbInformation, "Reconcile the split"
End Sub'''
SOL["ExportFacility"] = '''Sub ExportFacility()
    Const FAC_ID As String = "F02"
    Dim wbNew As Workbook, wbCheck As Workbook
    Dim filePath As String, savedName As String
    Dim lastR As Long
    Dim total As Double

    If Not SheetExists(FAC_ID) Then
        MsgBox "There's no " & FAC_ID & " sheet. Run SplitByFacility first.", vbExclamation
        Exit Sub
    End If
    If ThisWorkbook.Path = "" Then
        MsgBox "Save this workbook first, so the export has a folder to go in.", vbExclamation
        Exit Sub
    End If
    filePath = ThisWorkbook.Path & Application.PathSeparator & FAC_ID & "_encounters_2025.xlsx"

    ' 1. Copy the sheet into a brand-new workbook and save that as .xlsx
    ThisWorkbook.Worksheets(FAC_ID).Copy            ' no Before/After: Excel makes a new workbook
    Set wbNew = ActiveWorkbook                      ' right after Copy, the new workbook is active
    Application.DisplayAlerts = False               ' replace last run's file without asking
    wbNew.SaveAs Filename:=filePath, FileFormat:=xlOpenXMLWorkbook
    Application.DisplayAlerts = True
    savedName = wbNew.Name                          ' read it BEFORE closing
    wbNew.Close SaveChanges:=False

    ' 2. Reopen the saved file, read it, and close it again
    Set wbCheck = Workbooks.Open(Filename:=filePath, ReadOnly:=True)
    With wbCheck.Worksheets(1)
        lastR = .Cells(.Rows.Count, "A").End(xlUp).Row
        total = Application.WorksheetFunction.Sum(.Range("K2:K" & lastR))
    End With
    wbCheck.Close SaveChanges:=False

    ' 3. Report in THIS workbook, whichever workbook is active now
    With ThisWorkbook.Worksheets("Output")
        .Range("B6").Value = savedName
        .Range("B7").Value = total
    End With
End Sub'''
SOL["FillLOSDays"] = '''Sub FillLOSDays()
    Dim ws As Worksheet
    Dim lastR As Long, i As Long
    Dim stay As Variant          ' will hold an n x 2 array: admit, discharge
    Dim los() As Variant         ' n x 1 array for the results

    Set ws = ThisWorkbook.Worksheets("Encounters")
    lastR = LastRow(ws, "A")

    stay = ws.Range("G2:H" & lastR).Value2          ' ONE read. Value2 gives dates as serial numbers
    ReDim los(1 To UBound(stay, 1), 1 To 1)

    For i = 1 To UBound(stay, 1)
        ' whole days between the two dates = midnights in the hospital
        los(i, 1) = Int(stay(i, 2)) - Int(stay(i, 1))
    Next i

    ws.Range("L2").Resize(UBound(los, 1), 1).Value = los   ' ONE write
End Sub'''

BONUS_SOL = '''Sub BuildMonthlyPackets()
    Const NCOLS As Long = 11                     ' Encounters columns A:K
    Dim wsEnc As Worksheet, wsTpl As Worksheet, wsPk As Worksheet
    Dim src As Variant, buf() As Variant
    Dim lastR As Long, i As Long, j As Long, n As Long, m As Long
    Dim monthStart As Date, nextMonth As Date
    Dim sheetName As String, bestName As String
    Dim total As Double, bestTotal As Double

    Set wsEnc = ThisWorkbook.Worksheets("Encounters")
    Set wsTpl = ThisWorkbook.Worksheets("PacketTemplate")
    lastR = LastRow(wsEnc, "A")
    src = wsEnc.Range("A2").Resize(lastR - 1, NCOLS).Value   ' read the export ONCE

    Application.ScreenUpdating = False
    On Error GoTo CleanUp

    For m = 1 To 12
        monthStart = DateSerial(2025, m, 1)
        nextMonth = DateSerial(2025, m + 1, 1)   ' month 13 rolls over to January 2026
        sheetName = Format(monthStart, "yyyy-mm")

        ' 1. Collect this month's inpatient rows in a buffer as big as the export
        ReDim buf(1 To UBound(src, 1), 1 To NCOLS)
        n = 0
        total = 0
        For i = 1 To UBound(src, 1)
            If src(i, 3) = "Inpatient" And src(i, 7) >= monthStart And src(i, 7) < nextMonth Then
                n = n + 1
                For j = 1 To NCOLS
                    buf(n, j) = src(i, j)
                Next j
                total = total + src(i, 11)
            End If
        Next i

        ' 2. A fresh copy of the template, named for the month
        DeleteSheetIfExists sheetName
        wsTpl.Copy After:=ThisWorkbook.Worksheets(ThisWorkbook.Worksheets.Count)
        Set wsPk = ActiveSheet                   ' a copied sheet is always the active one
        wsPk.Name = sheetName

        ' 3. Header block, data, sort
        wsPk.Range("B2").Value = monthStart
        wsPk.Range("B3").Value = n
        wsPk.Range("B4").Value = total
        If n > 0 Then
            ' A range smaller than the array takes just the top-left n rows of it
            wsPk.Range("A7").Resize(n, NCOLS).Value = buf
            wsPk.Range("A6").CurrentRegion.Sort Key1:=wsPk.Range("K6"), Order1:=xlDescending, Header:=xlYes
        End If

        ' 4. Remember the biggest month so far
        If total > bestTotal Then
            bestTotal = total
            bestName = sheetName
        End If
    Next m

    ThisWorkbook.Worksheets("Output").Range("B9").Value = bestName

CleanUp:
    Application.ScreenUpdating = True
    If Err.Number <> 0 Then
        MsgBox "BuildMonthlyPackets stopped: " & Err.Description, vbExclamation
    Else
        MsgBox "12 packets built. Largest month: " & bestName, vbInformation
    End If
End Sub'''


def _stub(header: str, signature: str, lines: list[str], end: str = "End Sub") -> str:
    body = "\n".join(f"    ' {ln}" if ln else "    '" for ln in lines)
    return f"{header}\n{signature}\n{body}\n{end}"


def _starter_module() -> str:
    parts = [
        'Attribute VB_Name = "EncounterMacros"',
        "Option Explicit",
        "",
        "' =====================================================================",
        "' Lesson 5.3 - VBA: Ranges, Worksheets & Workbooks - starter module",
        "' (Bluestone Health System, fictional data)",
        "'",
        "' Import it: in the Visual Basic Editor choose File > Import File... and",
        "' pick this file. Complete each TODO, then run the macro with F5 (cursor",
        "' inside the Sub) or from Developer > Macros. Save before every run:",
        "' Undo can't undo a macro.",
        "'",
        COLUMN_MAP,
        "' =====================================================================",
        "",
        "' Helper (complete). Use it in your macros: lastR = LastRow(ws, \"A\")",
        LAST_ROW,
        "",
        _stub("' Task 7, part 1", "Function SheetExists(ByVal sheetName As String) As Boolean", [
            "TODO: loop over ThisWorkbook.Worksheets with For Each.",
            "When a sheet's name matches, ignoring case:",
            "    If StrComp(ws.Name, sheetName, vbTextCompare) = 0 Then",
            "set SheetExists = True and leave with Exit Function.",
            "(If the loop ends without a match, the function returns False.)",
        ], end="End Function"),
        "",
        "' Helper (complete). It calls your SheetExists, so it only works once",
        "' SheetExists does.",
        DELETE_IF,
        "",
        _stub("' Task 7, part 2", "Sub BuildTypeSummary()", [
            "TODO:",
            "1. DeleteSheetIfExists \"TypeSummary\"",
            "2. Add a sheet at the end of the workbook and name it TypeSummary:",
            "     Set wsSum = ThisWorkbook.Worksheets.Add( _",
            "         After:=ThisWorkbook.Worksheets(ThisWorkbook.Worksheets.Count))",
            "3. Headers in A1:C1: EncounterType, Encounters, TotalCharges",
            "4. For each type in Array(\"Emergency\", \"Inpatient\", \"Observation\", \"Outpatient\")",
            "   write the type, its count, and its total charges in rows 2-5.",
            "   (Array() numbers its items from 0: see guide section 10.)",
            "   Application.WorksheetFunction.CountIf(rngType, t) counts,",
            "   Application.WorksheetFunction.SumIf(rngType, t, rngCharges) totals.",
        ]),
        "",
        _stub("' Task 8", "Sub SplitByFacility()", [
            "TODO: for each FacilityID cell in the Facilities table:",
            "    For Each idCell In ThisWorkbook.Worksheets(\"Facilities\") _",
            "        .ListObjects(\"tblFacilities\").ListColumns(\"FacilityID\").DataBodyRange",
            "1. DeleteSheetIfExists, then add a sheet at the end named after the ID.",
            "2. Filter the Encounters block (Range(\"A1\").CurrentRegion) on",
            "   column 4 (FacilityID) with .AutoFilter Field:=4, Criteria1:=...",
            "3. Copy the visible cells, header included, to A1 of the new sheet:",
            "   .SpecialCells(xlCellTypeVisible).Copy Destination:=...",
            "After the loop, turn the filter off: wsEnc.AutoFilterMode = False",
            "Bonus points: ScreenUpdating off while it runs, and back on at the end.",
        ]),
        "",
        _stub("' Tasks 9 and 10", "Sub SortAndReconcile()", [
            "TODO: loop over every worksheet with For Each ws In ThisWorkbook.Worksheets.",
            "For each sheet whose name is Like \"F0#\":",
            "  Task 9:  sort its block (Range(\"A1\").CurrentRegion) by TotalCharges",
            "           (column K), largest first, with Header:=xlYes.",
            "  Task 10: add its data rows and its TotalCharges to two running totals.",
            "After the loop, write the row total to Output!B4 and the charge total",
            "to Output!B5.",
        ]),
        "",
        _stub("' Task 11", "Sub ExportFacility()", [
            "TODO:",
            "1. Build the path: ThisWorkbook.Path & Application.PathSeparator & _",
            "                   \"F02_encounters_2025.xlsx\"",
            "2. ThisWorkbook.Worksheets(\"F02\").Copy   (no arguments: a new workbook)",
            "   Set wbNew = ActiveWorkbook",
            "3. wbNew.SaveAs Filename:=..., FileFormat:=xlOpenXMLWorkbook",
            "   (DisplayAlerts False around it, so a second run can overwrite),",
            "   then wbNew.Close SaveChanges:=False",
            "4. Set wbCheck = Workbooks.Open(Filename:=..., ReadOnly:=True),",
            "   add up its TotalCharges (column K), and close it without saving.",
            "5. Write the file name to Output!B6 and the total to Output!B7.",
        ]),
        "",
        _stub("' Task 12", "Sub FillLOSDays()", [
            "TODO:",
            "1. Read G2:H<lastRow> (AdmitDateTime, DischargeDateTime) into a Variant",
            "   with ONE .Value2 read. Value2 gives each date as a serial number.",
            "2. ReDim a results array: ReDim los(1 To UBound(stay, 1), 1 To 1)",
            "3. For each row: los(i, 1) = Int(discharge) - Int(admit)",
            "4. Write the whole array to L2 with ONE assignment:",
            "   ws.Range(\"L2\").Resize(UBound(los, 1), 1).Value = los",
        ]),
        "",
        _stub("' Bonus - see the Bonus sheet for the full brief", "Sub BuildMonthlyPackets()", [
            "TODO (bonus): for m = 1 To 12",
            "1. DeleteSheetIfExists Format(DateSerial(2025, m, 1), \"yyyy-mm\"), then",
            "   copy PacketTemplate to the end and rename the copy to that name.",
            "2. Collect the month's Inpatient rows (AdmitDateTime in the month) from",
            "   an array you read ONCE before the loop, and write them from A7.",
            "3. Sort the block at A6 by TotalCharges (K), largest first.",
            "4. B2 = first day of the month, B3 = encounters, B4 = total charges.",
            "After the loop, write the name of the sheet with the largest total",
            "charges to Output!B9. ScreenUpdating off, and always back on.",
        ]),
        "",
    ]
    return "\n".join(parts)


def _snippets_module() -> str:
    parts = [
        'Attribute VB_Name = "Snippets"',
        "Option Explicit",
        "",
        "' =====================================================================",
        "' Lesson 5.3 - predict-the-output snippets (Practice tasks 1-6)",
        "'",
        "' For each snippet: predict what it prints (or which error it stops",
        "' with), type your prediction on the Practice sheet, THEN click inside",
        "' the Sub and press F5 to check. Debug.Print writes to the Immediate",
        "' window (Ctrl + G; Mac: View > Immediate Window).",
        "' =====================================================================",
        "",
    ]
    for key, s in SNIPPETS.items():
        parts += [f"' Snippet {key} ({s['tasks']})", s["code"], ""]
    return "\n".join(parts)


def _private(code: str) -> str:
    """Make a helper Private in the solution modules, so they never clash with the learner's Public copies."""
    return code.replace("Function ", "Private Function ", 1) if code.startswith("Function ") else \
        code.replace("Sub ", "Private Sub ", 1)


def _solution_module() -> str:
    parts = [
        'Attribute VB_Name = "EncounterMacrosSolution"',
        "Option Explicit",
        "",
        "' =====================================================================",
        "' Lesson 5.3 - reference solutions for Practice tasks 7-12 (SPOILERS)",
        "'",
        "' Import this into a spare copy of the workbook (or remove your own",
        "' EncounterMacros module first) so the macro names don't clash.",
        "' The helpers are Private here, so they never clash with yours.",
        "'",
        COLUMN_MAP,
        "' =====================================================================",
        "",
        _private(LAST_ROW),
        "",
        "' Task 7, part 1",
        _private(SOL["SheetExists"]),
        "",
        _private(DELETE_IF),
        "",
    ]
    numbers = {"BuildTypeSummary": "Task 7, part 2", "SplitByFacility": "Task 8", "SortAndReconcile": "Tasks 9 and 10",
               "ExportFacility": "Task 11", "FillLOSDays": "Task 12"}
    for name, label in numbers.items():
        parts += [f"' {label}", SOL[name], ""]
    return "\n".join(parts)


def _bonus_module() -> str:
    return "\n".join([
        'Attribute VB_Name = "MonthlyPacketsSolution"',
        "Option Explicit",
        "",
        "' =====================================================================",
        "' Lesson 5.3 - reference solution for the bonus (SPOILERS)",
        "' Import into a spare copy of the workbook, or remove your own",
        "' BuildMonthlyPackets first. The helpers below are Private copies.",
        "' =====================================================================",
        "",
        _private(LAST_ROW),
        "",
        _private(SOL["SheetExists"]),
        "",
        _private(DELETE_IF),
        "",
        BONUS_SOL,
        "",
    ])


def _lines(text: str, width: int) -> int:
    return max(1, -(-len(text) // width))


def _describe_sheets(wb, lesson, n_rows: int = N_SAMPLE):
    """Start Here lists Practice, the data sheets, the sheet_notes entries (Snippets, Output, PacketTemplate) and Bonus.
    The library describes Practice, Encounters and Bonus generically ("Data: 2,000 rows"), so say what they hold here."""
    st = wb["Start Here"]
    top = next(r for r in range(1, st.max_row + 1) if st.cell(row=r, column=2).value == "Sheets in this workbook")
    better = {
        lesson.practice_sheet: f"Tasks 1–{len(lesson.tasks)}. Yellow cells take your predictions (tasks 1–6), and gray cells "
                               "read what your macros create (tasks 7–12).",
        "Encounters": f"Data: {n_rows:,} rows in columns A:K, a plain range like a raw EHR export. Task 12 fills the yellow "
                      "LOSDays column (L).",
        lesson.bonus_sheet: f"The bonus challenge ({len(lesson.bonus)} parts): one inpatient packet per month. Its gray cells "
                            "read the sheets BuildMonthlyPackets creates.",
    }
    seen = set()
    for r in range(top + 1, st.max_row + 1):
        label = st.cell(row=r, column=2).value
        if not label:
            continue
        assert label not in seen, f"Start Here lists {label} twice"
        seen.add(label)
        text = better.get(label)
        if text:
            st.cell(row=r, column=3).value = text
            st.row_dimensions[r].height = 15 * _lines(text, 100) + 3
    # The keys are hidden after the customize hooks run, so leave them out by name.
    listed = {ws.title for ws in wb.worksheets if ws.sheet_state == "visible"} - {"Start Here", lesson.key_sheet,
                                                                                 lesson.bonus_key_sheet}
    assert listed <= seen, f"Start Here doesn't list {sorted(listed - seen)}"


def _crlf(text: str) -> bytes:
    text.encode("ascii")  # raises if a non-ASCII character slipped in (the VBE reads ANSI, not UTF-8)
    return text.replace("\r\n", "\n").replace("\n", "\r\n").encode("ascii")


# ---------------------------------------------------------------------------
# Small simulators for the predict-the-output answers (Excel semantics)
# ---------------------------------------------------------------------------
def _parse(addr: str) -> tuple[int, int]:
    col = "".join(ch for ch in addr if ch.isalpha())
    return int(addr[len(col):]), column_index_from_string(col)


def _address(r1, c1, r2, c2, absolute=True) -> str:
    d = "$" if absolute else ""
    a = f"{d}{get_column_letter(c1)}{d}{r1}"
    if (r1, c1) == (r2, c2):
        return a
    return f"{a}:{d}{get_column_letter(c2)}{d}{r2}"


def _offset_resize(top_left: str, dr: int, dc: int, rows: int, cols: int = 1) -> tuple[int, int, int, int]:
    """Range(top_left).Offset(dr, dc).Resize(rows, cols) for a single-cell start."""
    r, c = _parse(top_left)
    r, c = r + dr, c + dc
    return r, c, r + rows - 1, c + cols - 1


def _end_up(filled: list[bool], first_row: int, max_row: int = 1048576) -> int:
    """Cells(Rows.Count, col).End(xlUp).Row: the last non-empty cell (header in row first_row - 1 is filled)."""
    rows = [first_row - 1] + [first_row + i for i, f in enumerate(filled) if f]
    return max(rows)


def _end_down(filled: list[bool], first_row: int, max_row: int = 1048576) -> int:
    """Range(header).End(xlDown).Row starting from a filled header cell in row first_row - 1."""
    cells = {first_row - 1: True}
    cells.update({first_row + i: f for i, f in enumerate(filled)})

    def full(r):
        return cells.get(r, False)

    r = first_row - 1
    if full(r + 1):                  # inside a block: go to its last cell
        while full(r + 1):
            r += 1
        return r
    r += 1                            # at the edge of a block: jump to the next filled cell
    while r < max_row and not full(r):
        r += 1
    return r


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------
def _sample() -> list[dict]:
    enc = [e for e in data.load("encounters") if e["AdmitDateTime"].year == 2025]
    enc.sort(key=lambda e: (e["AdmitDateTime"], e["EncounterID"]))
    n = len(enc)
    return [enc[i * n // N_SAMPLE] for i in range(N_SAMPLE)]   # evenly spaced across the year


def _los(e) -> int:
    return (e["DischargeDateTime"].date() - e["AdmitDateTime"].date()).days


def _month_rows(rows, m):
    out = [e for e in rows if e["EncounterType"] == "Inpatient" and e["AdmitDateTime"].month == m]
    return sorted(out, key=lambda e: -e["TotalCharges"])     # the packet's order after the sort


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="05-automation-vba", slug="03-vba-ranges-worksheets",
        title="VBA: Ranges, Worksheets & Workbooks", level="Expert", minutes=65,
        objectives=[
            "Navigate the object model: Application, Workbooks, Worksheets, Range",
            "Find the last row and work with CurrentRegion, Offset, and Resize",
            "Loop through sheets; add, copy, rename, and delete them",
            "Process data fast with arrays, AutoFilter, and Sort — without Select",
        ],
        data_note="2,000 encounters admitted in 2025 at all four Bluestone facilities (inpatient, observation, emergency, "
                  "and outpatient), exported from the EHR as a plain range, plus the facility list as an Excel Table.",
    )

    rows = _sample()
    facilities = sorted(data.load("facilities"), key=lambda f: f["FacilityID"])
    n = len(rows)
    enc = L.add_table_sheet(
        "Encounters", rows, columns=ENC_COLS, extra_cols=[LOS_COL], as_table=False,
        formats={LOS_COL: "0"}, widths={"AdmitSource": 21, "AdmitDateTime": 17, "DischargeDateTime": 18, LOS_COL: 10},
    )
    fac = L.add_table_sheet("Facilities", facilities, columns=FAC_COLS, table="tblFacilities",
                            widths={"FacilityName": 32, "FacilityType": 24})
    first, last = enc.first_row, enc.last_row
    assert [enc.col(c) for c in ("EncounterType", "FacilityID", "AdmitSource", "AdmitDateTime", "DischargeDateTime",
                                 "TotalCharges", LOS_COL)] == list("CDFGHKL")
    assert (first, last) == (2, n + 1)

    def rng(col):
        return f"Encounters!{enc.col(col)}{first}:{enc.col(col)}{last}"

    # ------------------------------------------------------------------ snippet answers (Excel semantics, simulated)
    a = _offset_resize("C2", 3, 4, rows=5)                       # Snippet A: Range("C2").Offset(3, 4).Resize(5)
    ans_a = _address(*a)
    n_cols = len(enc.headers)                                    # CurrentRegion spans the header row's width
    region = (1, 1, last, n_cols)                                # Range("E50").CurrentRegion
    height = region[2] - region[0] + 1
    body = (region[0] + 1, region[1], region[0] + 1 + (height - 1) - 1, region[3])   # .Offset(1).Resize(height - 1)
    ans_b = _address(*body, absolute=False)
    src_filled = [bool(e["AdmitSource"]) for e in rows]
    line1 = _end_up([True] * n, first)
    ans_c2 = _end_up(src_filled, first)
    ans_c3 = _end_down(src_filled, first)
    trailing_blank = line1 - ans_c2
    assert {e["EncounterType"] for e in rows[ans_c2 - 1:]} <= {"Emergency", "Outpatient"}   # rows below ans_c2
    first_blank_row = first + src_filled.index(False)
    ans_d = "1004"     # Run-time error 1004: Method 'Range' of object '_Worksheet' failed (cells from another sheet)
    ans_e = "9"        # Run-time error 9: Subscript out of range (no sheet with that name)

    # ------------------------------------------------------------------ macro answers
    type_count = {t: sum(1 for e in rows if e["EncounterType"] == t) for t in TYPES}
    type_total = {t: round(sum(e["TotalCharges"] for e in rows if e["EncounterType"] == t), 2) for t in TYPES}
    assert set(e["EncounterType"] for e in rows) == set(TYPES)
    fac_ids = [f["FacilityID"] for f in facilities]
    fac_rows = {f: sorted([e for e in rows if e["FacilityID"] == f], key=lambda e: -e["TotalCharges"]) for f in fac_ids}
    fac_counts = " / ".join(str(len(fac_rows[f])) for f in fac_ids)
    top = fac_rows[SORT_CHECK_FAC]
    assert top[0]["TotalCharges"] > top[1]["TotalCharges"], "tie at the top of the sort check"
    top_id = top[0]["EncounterID"]
    grand_total = round(sum(e["TotalCharges"] for e in rows), 2)
    export_total = round(sum(e["TotalCharges"] for e in fac_rows[EXPORT_FAC]), 2)
    los = [_los(e) for e in rows]
    los_total = sum(los)
    ip = [e for e in rows if e["EncounterType"] == "Inpatient"]
    ip_los = sum(_los(e) for e in ip)
    assert all(x >= 0 for x in los)

    # ------------------------------------------------------------------ bonus answers
    packets = {m: _month_rows(rows, m) for m in range(1, 13)}
    pk_total = {m: round(sum(e["TotalCharges"] for e in packets[m]), 2) for m in packets}
    best_m = max(pk_total, key=pk_total.get)
    assert sorted(pk_total.values())[-1] > sorted(pk_total.values())[-2]
    best_name = f"2025-{best_m:02d}"
    mc, mt, mtop = PACKET_CHECK["count"], PACKET_CHECK["total"], PACKET_CHECK["top"]
    pk_count_ans = len(packets[mc])
    pk_total_ans = pk_total[mt]
    pk_top = packets[mtop]
    assert pk_top[0]["TotalCharges"] > pk_top[1]["TotalCharges"]
    pk_top_ans = pk_top[0]["TotalCharges"]
    smallest_m = min(pk_total, key=pk_total.get)

    def pk(m):
        return f"2025-{m:02d}"

    def exists(sheet):
        return f'ISREF(INDIRECT("\'{sheet}\'!A1"))'

    def out_summary(cell):
        return f'=IF(Output!{cell}="","",Output!{cell})'

    def out_fill(cell, value):
        return {"range": f"Output!{cell}:{cell}", "values": [value]}

    def touch(sheet, value):
        # The customize hook builds the macro-created sheet in the self-test; this fill just exercises it.
        return {"range": f"'{sheet}'!A1:A1", "values": [value]}

    # ------------------------------------------------------------------ Practice
    L.start_notes = [
        "Macros can't be saved in an .xlsx file. Before you write any code, choose File → Save As and pick "
        "'Excel Macro-Enabled Workbook (*.xlsm)'.",
        "Import starter/EncounterMacros.bas and starter/Snippets.bas from the lesson folder: open the Visual Basic Editor with "
        "Alt + F11 (Mac: Option + F11, or Developer → Visual Basic), then choose File → Import File…",
        "Snippets shows the code for tasks 1–6. Your macros add sheets (TypeSummary, F01–F04, 2025-01…) and write to the "
        "Output sheet. The gray cells on Practice and Bonus read those, so they stay blank until the macro has run.",
    ]
    L.practice_intro = (
        "Save the workbook as .xlsm, then import starter/EncounterMacros.bas and starter/Snippets.bas (**File → Import File…** "
        "in the VBE). Tasks 1–6 ask what the code on the Snippets sheet prints: type your prediction, then run the snippet "
        "to check it. In tasks 7–12 you complete macros, and the gray cells read the sheets and cells your macros create.")
    # Tasks 1-6 are yellow predictions, while tasks 7-12 and the whole bonus are gray cells that read macro output, so the
    # generic "type a formula or value in each yellow cell" lines don't fit.
    L.practice_how = ("Go to the 'Practice' sheet. Type your predictions for tasks 1–6 in the yellow cells. For tasks 7–12, "
                      "complete and run your macros, and the gray cells fill in by themselves.")
    L.practice_instructions = (
        "Type your predictions for tasks 1–6 in the yellow cells. The gray cells for tasks 7–12 fill in when your macros run. "
        "The Check column turns green when your answer matches. "
        f"Stuck? Read the hint, then the lesson guide. Answers: right-click a sheet tab → Unhide… → '{L.key_sheet}'.")
    L.bonus_instructions = (
        "There's nothing to type on this sheet. The gray cells fill in once BuildMonthlyPackets has run, and the Check column "
        "turns green when they match. "
        f"Stuck? Read the hint, then the lesson guide. Answers: right-click a sheet tab → Unhide… → '{L.bonus_key_sheet}'.")
    L.bonus_where = ("Complete BuildMonthlyPackets in the EncounterMacros module and run it. It adds the monthly packet sheets "
                     "and writes to the **Output** sheet, and the gray cells on the **Bonus** sheet read your results.")
    # Start Here lists Practice, Encounters, Facilities, these (the customize hook creates them), then Bonus.
    L.sheet_notes = [
        ("Snippets", "The code for the predict-the-output tasks 1–6. The same code is in starter/Snippets.bas."),
        ("Output", "The blue cells your macros write results to. The Practice and Bonus sheets read them."),
        ("PacketTemplate", "The bonus template that BuildMonthlyPackets copies once per month. Leave it as it is."),
    ]

    sum_type = (f'=IF({exists(SUMMARY_SHEET)},IFERROR(INDEX(INDIRECT("\'{SUMMARY_SHEET}\'!C1:C50"),'
                f'MATCH("Inpatient",INDIRECT("\'{SUMMARY_SHEET}\'!A1:A50"),0)),"Inpatient row not found"),"")')

    def fac_part(f):
        return f'IF({exists(f)},COUNTA(INDIRECT("\'{f}\'!A:A"))-1,"missing")'
    sum_split = (f'=IF(OR({",".join(exists(f) for f in fac_ids)}),'
                 + '&" / "&'.join(fac_part(f) for f in fac_ids) + ',"")')
    live_split = "=" + '&" / "&'.join(f'COUNTIF({rng("FacilityID")},"{f}")' for f in fac_ids)

    L.tasks = [
        # ---------------------------------------------------------- 1-6 predict
        Task("Snippet A: what address does the Immediate window show? Type it as printed (the $ signs are optional).",
             answer=ans_a, accept=[ans_a.replace("$", "")], live=False,
             solution=f"`Range(\"C2\").Offset(3, 4)` moves 3 rows down and 4 columns right, to **{ans_a.split(':')[0].replace('$', '')}**. "
                      f"`Resize(5)` makes the range 5 rows tall from that corner and keeps its width of 1 column: **{ans_a}**.",
             hint="Offset moves the top-left corner, and Resize sets the size from that corner", title="Snippet A: Offset and Resize",
             explanation="Offset never changes a range's size, and Resize never moves its top-left corner, so you can read a chain "
                         "of them left to right. Leaving out Resize's second argument keeps the column count as it is. `.Address` "
                         "returns absolute references ($G$5) unless you ask for `.Address(False, False)`."),
        Task("Snippet B: what address does the Immediate window show?",
             answer=ans_b, accept=[_address(*body)], live=False,
             solution=f"`Range(\"E50\").CurrentRegion` grows from E50 to the whole block bounded by empty rows and columns: "
                      f"A1:{get_column_letter(n_cols)}{last} (the LOSDays header in {enc.col(LOS_COL)}1 makes the block "
                      f"{n_cols} columns wide). `Offset(1)` shifts the block down one row, and `Resize(Rows.Count - 1)` trims the "
                      f"row that fell off the bottom: **{ans_b}**.",
             hint="CurrentRegion is the whole block around E50 (**Go To Special → Current region** shows it). The last two steps drop "
                  "the header row",
             title="Snippet B: the data-body idiom",
             explanation="This **data-body idiom** gives you the data without its header, however many rows the export has. "
                         "It works from any cell inside the block, which is why the snippet starts in E50. CurrentRegion stops at the "
                         "first completely empty row and column, so a blank row in the middle of an export cuts it short."),
        Task(f"Snippet C: line 1 prints {line1}, the last row of column A. What number does line 2 print (the same idiom on "
             "column F, AdmitSource)?",
             answer=ans_c2, answer_display=str(ans_c2), title="Snippet C, line 2: the last row of AdmitSource", live=f'=SUMPRODUCT(MAX(({rng("AdmitSource")}<>"")*ROW({rng("AdmitSource")})))',
             solution=f"End(xlUp) from the bottom of column F stops at the last cell with something in it: row **{ans_c2}**. "
                      f"The last {trailing_blank} encounters are emergency or outpatient visits, which have no AdmitSource.",
             hint="Which encounter types leave AdmitSource blank? Look at the bottom of column F",
             explanation="The last-row idiom finds the last filled cell **in the column you give it**. AdmitSource is blank for "
                         "emergency and outpatient visits, so a loop that stopped at this row would silently skip the last "
                         f"{trailing_blank} encounters. Always use a column that is filled on every row, such as the ID column."),
        Task("Snippet C: what number does line 3 print?",
             answer=ans_c3, answer_display=str(ans_c3), title="Snippet C, line 3: End(xlDown)", live=False,
             solution=f"F1 and F2 are both filled and F{first_blank_row} is empty, so End(xlDown) stops at the end of that first "
                      f"block: row **{ans_c3}**.",
             hint="End(xlDown) works like Ctrl + ↓ (Mac: ⌘ + ↓): it stops at the last filled cell before a gap",
             explanation="`Range(\"F1\").End(xlDown)` is the top-down version of the idiom, and it fails on the first blank cell. "
                         "On a column with no gaps it gives the same answer as End(xlUp), which is why it seems to work until "
                         "the day a blank appears. Start from the bottom and go up with `Cells(Rows.Count, col).End(xlUp)`."),
        Task("Snippet D: running SnippetD_WhichSheet stops with a run-time error. Type the error number.",
             answer=ans_d, accept=["error 1004", "run-time error 1004", "runtime error 1004"], live=False,
             title="Snippet D: the unqualified Cells error",
             solution="Run-time error **1004**: *Method 'Range' of object '_Worksheet' failed*. The two bare `Cells(...)` belong to "
                      "the active sheet (Practice), and `ws.Range(...)` can't build an Encounters range from Practice cells.",
             hint="Which sheet does a Cells(…) with nothing in front of it belong to?",
             explanation="In a standard module, `Range` and `Cells` with no object in front mean *the active sheet*. Qualify every "
                         "one: `ws.Range(ws.Cells(2, 1), ws.Cells(2, 11))`, or use a With block and a leading dot: "
                         "`.Range(.Cells(2, 1), .Cells(2, 11))`. The same macro works when Encounters happens to be active, "
                         "which is what makes this bug hard to spot."),
        Task("Snippet E: running SnippetE_DeleteOld stops with a run-time error. Type the error number.",
             answer=ans_e, accept=["subscript out of range", "error 9", "run-time error 9", "runtime error 9"], live=False,
             title="Snippet E: deleting a sheet that isn't there",
             solution="Run-time error **9**: *Subscript out of range*. There is no sheet named TypeSummary_2024, so "
                      "`Worksheets(\"TypeSummary_2024\")` fails before `.Delete` even runs.",
             hint="What does a collection do when you ask for an item it doesn't have?",
             explanation="`DisplayAlerts = False` only silences Excel's *confirmation* prompt. It doesn't make a missing sheet "
                         "exist. That's why delete-and-recreate macros check first with a function such as **SheetExists**. "
                         "The macro also stopped before its `DisplayAlerts = True` line. Excel resets DisplayAlerts by itself "
                         "when code finishes, but it doesn't do that for Calculation or EnableEvents, which is why restore "
                         "lines belong in cleanup code that always runs (guide section 15)."),
        # ---------------------------------------------------------- 7 sheets: delete & recreate
        Task(f"Complete SheetExists and BuildTypeSummary in the starter module. BuildTypeSummary deletes any old {SUMMARY_SHEET} "
             f"sheet (with DeleteSheetIfExists, which calls your SheetExists), adds a new sheet named {SUMMARY_SHEET} at the end of "
             "the workbook, and writes one row per EncounterType (Emergency, Inpatient, Observation, Outpatient) with the columns "
             "EncounterType (A), Encounters (B, the count), and TotalCharges (C), headers in row 1. Run it twice: the second run "
             "must not stop with an error or ask a question. The gray cell looks up the Inpatient TotalCharges on your sheet.",
             answer=type_total["Inpatient"], fmt="#,##0.00", title=f"SheetExists + BuildTypeSummary (Inpatient TotalCharges)",
             solution=SOL["SheetExists"] + "\n\n" + SOL["BuildTypeSummary"], solution_lang="vba",
             summary=sum_type, fill=touch(SUMMARY_SHEET, "EncounterType"),
             live=f'=SUMIFS({rng("TotalCharges")},{rng("EncounterType")},"Inpatient")',
             hint="SheetExists: For Each ws In ThisWorkbook.Worksheets … StrComp(ws.Name, sheetName, vbTextCompare) = 0",
             explanation="**Delete-and-recreate** makes a macro safe to rerun: without the delete, the second run stops at "
                         "`wsSum.Name = \"TypeSummary\"` with run-time error 1004 (*That name is already taken*), because sheet "
                         "names must be unique. `Worksheets.Add` returns the new sheet, so `Set wsSum = …` gives you a variable "
                         "for it and you never need ActiveSheet. SheetExists compares with `vbTextCompare` because Excel sheet "
                         "names ignore case. The finished sheet: "
                         + ", ".join(f"{t} {type_count[t]:,} encounters / ${type_total[t]:,.2f}" for t in TYPES)
                         + f". Inpatient stays are {type_count['Inpatient'] / n:.0%} of the encounters but "
                           f"{type_total['Inpatient'] / sum(type_total.values()):.0%} of the charges."),
        # ---------------------------------------------------------- 8 AutoFilter split
        Task("Complete SplitByFacility: for each FacilityID in the Facilities table (tblFacilities), delete any old sheet with that "
             "name, add a new sheet named after the ID (F01, F02, F03, F04), AutoFilter the Encounters block on FacilityID, and copy "
             "the visible cells (header included) to A1 of the new sheet. Turn the filter off at the end. The gray cell lists the "
             "number of data rows on F01, F02, F03, and F04.",
             answer=fac_counts, title="SplitByFacility (data rows per facility sheet)",
             solution=SOL["SplitByFacility"], solution_lang="vba",
             summary=sum_split, fill=touch("F01", "EncounterID"), live=live_split,
             hint="block.AutoFilter Field:=4, Criteria1:=facID, then block.SpecialCells(xlCellTypeVisible).Copy",
             explanation="AutoFilter hides the rows that don't match, and `SpecialCells(xlCellTypeVisible)` picks only the rows "
                         "still showing, so the copy brings the header and that facility's rows and nothing else. Looping over the "
                         "Table's FacilityID column means a fifth facility would get its own sheet with no code change. The "
                         "`On Error GoTo CleanUp` block turns the filter off and ScreenUpdating back on even if something fails. "
                         "If a count is one too high or too low, check whether your copy included the header row, because the "
                         "gray cell subtracts 1 for it."),
        # ---------------------------------------------------------- 9-10 loop through sheets
        Task(f"Complete SortAndReconcile, part 1: loop through every worksheet and, for each sheet whose name is Like \"F0#\", sort its "
             f"block by TotalCharges (column K), largest first, with Range.Sort and a header row. Which EncounterID is now in A2 of "
             f"sheet {SORT_CHECK_FAC}? (The gray cell reads {SORT_CHECK_FAC}!A2.)",
             answer=top_id, title=f"SortAndReconcile: top EncounterID on {SORT_CHECK_FAC}",
             solution=SOL["SortAndReconcile"], solution_lang="vba",
             summary=f'=IF({exists(SORT_CHECK_FAC)},INDIRECT("\'{SORT_CHECK_FAC}\'!A2")&"","")',
             fill=touch(SORT_CHECK_FAC, "EncounterID"),
             live=(f'=INDEX({rng("EncounterID")},MATCH(1,({rng("FacilityID")}="{SORT_CHECK_FAC}")*'
                   f'({rng("TotalCharges")}=MAX(IF({rng("FacilityID")}="{SORT_CHECK_FAC}",{rng("TotalCharges")}))),0))'),
             hint="Test ws.Name Like \"F0#\" inside the loop, then sort ws.Range(\"A1\").CurrentRegion with Key1 in column K "
                  "and Order1:=xlDescending (guide section 12)",
             explanation=f"`For Each ws In ThisWorkbook.Worksheets` visits every worksheet, including the hidden key sheets, so the "
                         f"`Like \"F0#\"` test (# means any one digit) picks out the facility sheets. `Header:=xlYes` keeps row 1 "
                         f"in place as a header. Leave it out and Range.Sort treats row 1 as data, so an ascending sort would bury "
                         f"the header among the rows. {top_id} is Cedar Ridge's most "
                         f"expensive 2025 encounter in the sample (${top[0]['TotalCharges']:,.2f}, "
                         f"{top[0]['EncounterType'].lower()})."),
        Task("SortAndReconcile, part 2: in the same loop, add each facility sheet's data rows and TotalCharges to two running "
             "totals, then write the row total to Output!B4 and the charge total to Output!B5. What is the total charge on the "
             "facility sheets? (The gray cell reads Output!B5.)",
             answer=grand_total, fmt="#,##0.00", title="SortAndReconcile: control total of charges",
             solution="The second half of the loop in task 9's macro: `rowTotal = rowTotal + (lastR - 1)` and "
                      "`chargeTotal = chargeTotal + Application.WorksheetFunction.Sum(ws.Range(\"K2:K\" & lastR))`, "
                      "then both are written to Output after `Next ws`.",
             summary=out_summary(OUT["charges"]), fill=out_fill(OUT["charges"], grand_total),
             live=f"=SUM({rng('TotalCharges')})",
             hint="Find each sheet's last row with LastRow(ws, \"A\"), then WorksheetFunction.Sum the K cells",
             explanation=f"This is a **reconciliation**: the facility sheets together must hold exactly the {n:,} rows and "
                         f"${grand_total:,.2f} of the export (compare with `=SUM({rng('TotalCharges')})`). If the totals differ, "
                         "the split lost or duplicated rows, and you know before anyone reads the sheets. The same control-total "
                         "habit applies to any macro that moves data around."),
        # ---------------------------------------------------------- 11 workbooks
        Task(f"Complete ExportFacility: copy sheet {EXPORT_FAC} into a new workbook (Worksheet.Copy with no arguments), save it in "
             f"the same folder as this workbook as {EXPORT_FILE} with SaveAs and FileFormat:=xlOpenXMLWorkbook, and close it. "
             "Then reopen the file with Workbooks.Open, add up its TotalCharges column, write the total to Output!B7 (and the "
             "file name to Output!B6), and close it without saving. The gray cell reads Output!B7.",
             answer=export_total, fmt="#,##0.00", title=f"ExportFacility ({EXPORT_FAC} total read back from the saved file)",
             solution=SOL["ExportFacility"], solution_lang="vba",
             summary=out_summary(OUT["export"]), fill=out_fill(OUT["export"], export_total),
             live=f'=SUMIFS({rng("TotalCharges")},{rng("FacilityID")},"{EXPORT_FAC}")',
             hint="After Worksheet.Copy the new workbook is the ActiveWorkbook, so write results through ThisWorkbook",
             explanation="Once a second workbook is open, **ActiveWorkbook** and **ThisWorkbook** are different files. ThisWorkbook "
                         "is always the file that holds the code, so it's the safe way back to Output. `FileFormat:=xlOpenXMLWorkbook` "
                         "(51) must match the .xlsx extension, or Excel refuses or saves a file that won't open. DisplayAlerts False "
                         "lets a second run overwrite the old file. Opening with `ReadOnly:=True` and closing with "
                         "`SaveChanges:=False` guarantees the check never changes the export."),
        # ---------------------------------------------------------- 12 arrays
        Task("Complete FillLOSDays: read AdmitDateTime and DischargeDateTime (Encounters G2:H2001) into a Variant array with one "
             ".Value2 read, calculate each row's length of stay in days with Int(discharge) - Int(admit) (the number of midnights, "
             "ignoring the times), and write all the results to the yellow LOSDays column (L) with one assignment. The gray cell "
             "adds up your column.",
             answer=los_total, title="FillLOSDays (sum of the LOSDays column)",
             solution=SOL["FillLOSDays"], solution_lang="vba",
             summary=f'=IF(COUNT({rng(LOS_COL)})=0,"",SUM({rng(LOS_COL)}))',
             fill={"range": f"Encounters!{enc.col(LOS_COL)}{first}:{enc.col(LOS_COL)}{last}", "values": los},
             live=f"=SUMPRODUCT(INT({rng('DischargeDateTime')})-INT({rng('AdmitDateTime')}))",
             hint="stay = ws.Range(\"G2:H\" & lastR).Value2, then ReDim los(1 To UBound(stay, 1), 1 To 1)",
             explanation=f"Every read or write of a cell crosses from VBA into Excel, and that crossing is the slow part. The array "
                         f"version crosses twice (one read, one write) instead of {3 * n:,} times ({2 * n:,} reads and {n:,} writes) for "
                         f"{n:,} rows. A multi-cell range's "
                         f"`.Value2` is always a **2-D** array numbered from 1, even for one column, so the results array must be "
                         f"2-D too: `los(1 To n, 1 To 1)`. Value2 hands back dates as serial numbers, so `Int` strips the time. "
                         f"Inpatient stays account for {ip_los:,} of the {los_total:,} days. Emergency and observation visits "
                         f"that crossed midnight count 1 or 2."),
    ]

    # ------------------------------------------------------------------ bonus
    L.bonus_title = "Bonus: the monthly inpatient packet"
    L.bonus_scenario = (
        "Finance wants a monthly packet of inpatient encounters, one sheet per month of 2025, each built from the PacketTemplate "
        "sheet. Complete BuildMonthlyPackets so that, for each month 1 to 12, it does four things:\n"
        "1. Deletes any old sheet named for the month in yyyy-mm form (for example 2025-03), copies PacketTemplate to the end of "
        "the workbook, and renames the copy to that name.\n"
        f"2. Writes that month's Inpatient encounters (AdmitDateTime in the month) to the packet: all {len(ENC_COLS)} columns "
        f"A:K, in the same order as Encounters. The template's headers are in row {PACKET_HEADER_ROW}, so data starts in "
        f"A{PACKET_FIRST}. Use arrays: read the Encounters data once, before the month loop, collect each month's rows in a "
        "second array, and write them with one assignment.\n"
        f"3. Sorts the packet's rows by TotalCharges, largest first.\n"
        "4. Fills the header block: B2 the first day of the month, B3 the number of encounters, and B4 their total charges.\n\n"
        "After the loop, write the name of the packet sheet with the largest total charges to Output!B9 (it's formatted as Text, "
        "so Excel keeps 2025-xx as text and doesn't turn it into a date). Turn ScreenUpdating off while the macro runs, and make "
        "sure it comes back on even if something fails.")
    pk_exist = {m: exists(pk(m)) for m in (mc, mt, mtop)}
    best_check = (f'OR(LOWER(TRIM({{cell}}&""))="{best_name}",AND(ISNUMBER({{cell}}),'
                  f'IFERROR(YEAR({{cell}})=2025,FALSE),IFERROR(MONTH({{cell}})={best_m},FALSE)))')
    L.bonus = [
        Task(f"How many inpatient encounters are in the {pk(mc)} packet? (The gray cell counts the data rows from row {PACKET_FIRST} down.)",
             answer=pk_count_ans, title=f"Encounters in the {pk(mc)} packet",
             solution=BONUS_SOL, solution_lang="vba",
             summary=f'=IF({pk_exist[mc]},COUNTA(INDIRECT("\'{pk(mc)}\'!A{PACKET_FIRST}:A2000")),"")',
             fill=touch(pk(mc), "Monthly inpatient packet"),
             live=(f'=COUNTIFS({rng("EncounterType")},"Inpatient",{rng("AdmitDateTime")},">="&DATE(2025,{mc},1),'
                   f'{rng("AdmitDateTime")},"<"&DATE(2025,{mc + 1},1))'),
             hint="Test src(i, 3) = \"Inpatient\" And src(i, 7) >= monthStart And src(i, 7) < nextMonth",
             explanation="Comparing against the **first day of the next month** with `<` catches every admission on the last day of "
                         "the month, whatever the time. `<= DateSerial(2025, 3, 31)` would miss a patient admitted at 3/31 14:20, "
                         "because that datetime is larger than midnight on 3/31. `DateSerial(2025, m + 1, 1)` works for December "
                         "too: month 13 rolls over to January 2026. The `ReDim buf(1 To UBound(src, 1), …)` buffer is as big as the "
                         "whole export, and `Range(\"A7\").Resize(n, 11).Value = buf` writes only its top-left n rows."),
        Task(f"What total charges does the {pk(mt)} packet show in its header (cell B4)?",
             answer=pk_total_ans, fmt="#,##0.00", title=f"Header total on the {pk(mt)} packet",
             solution="Add each matching row's TotalCharges (`src(i, 11)`) to `total` inside the collecting loop, then write "
                      "`wsPk.Range(\"B4\").Value = total` after copying the template.",
             summary=f'=IF({pk_exist[mt]},INDIRECT("\'{pk(mt)}\'!B4"),"")', fill=touch(pk(mt), "Monthly inpatient packet"),
             live=(f'=SUMIFS({rng("TotalCharges")},{rng("EncounterType")},"Inpatient",{rng("AdmitDateTime")},">="&DATE(2025,{mt},1),'
                   f'{rng("AdmitDateTime")},"<"&DATE(2025,{mt + 1},1))'),
             hint="Reset total = 0 at the start of every month",
             explanation="If every month after January shows a bigger total than the one before, `total` wasn't reset inside the "
                         "month loop, so each packet carries the previous months' charges. The same goes for `n`. Accumulators "
                         "that belong to one pass of an outer loop must be reset at the top of that pass."),
        Task(f"What is the top charge in the {pk(mtop)} packet (TotalCharges in K{PACKET_FIRST}, after sorting)?",
             answer=pk_top_ans, fmt="#,##0.00", title=f"Top charge on the {pk(mtop)} packet",
             solution=f"Sort the block that starts at the header row: `wsPk.Range(\"A{PACKET_HEADER_ROW}\").CurrentRegion.Sort "
                      f"Key1:=wsPk.Range(\"K{PACKET_HEADER_ROW}\"), Order1:=xlDescending, Header:=xlYes`.",
             summary=f'=IF({pk_exist[mtop]},INDIRECT("\'{pk(mtop)}\'!K{PACKET_FIRST}"),"")', fill=touch(pk(mtop), "Monthly inpatient packet"),
             live=(f'=MAX(IF(({rng("EncounterType")}="Inpatient")*({rng("AdmitDateTime")}>=DATE(2025,{mtop},1))*'
                   f'({rng("AdmitDateTime")}<DATE(2025,{mtop + 1},1)),{rng("TotalCharges")}))'),
             hint=f"Row {PACKET_HEADER_ROW - 1} of the template is empty, so CurrentRegion from A{PACKET_HEADER_ROW} is just the table",
             explanation=f"The template leaves row {PACKET_HEADER_ROW - 1} empty on purpose. CurrentRegion stops at an empty row, so "
                         f"`Range(\"A{PACKET_HEADER_ROW}\").CurrentRegion` is the header plus the data and never includes the "
                         f"labels in A2:B4. Without that empty row, the region would reach up to the title and labels in rows 1–"
                         f"{PACKET_HEADER_ROW - 2}, and the sort would shuffle them in among the encounters. "
                         f"{pk_top[0]['EncounterID']} tops the {pk(mtop)} packet at ${pk_top_ans:,.2f}."),
        Task("Which month's packet has the largest total charges? (The gray cell reads Output!B9.)",
             answer=best_name, check="custom", custom_check=best_check, live=False,
             title="Packet with the largest total charges",
             solution="Track a running maximum in the month loop. After each month, `If total > bestTotal Then` store `total` in "
                      "`bestTotal` and `sheetName` in `bestName`. After `Next m`, write `bestName` to Output!B9.",
             summary=out_summary(OUT["best"]), fill=out_fill(OUT["best"], best_name),
             hint="The running-maximum pattern from Lesson 5.2, one level up: once per month, not once per row",
             explanation=(f"**{best_name}**: ${pk_total[best_m]:,.2f} from {len(packets[best_m])} stays. The quietest month was "
                          f"{pk(smallest_m)} (${pk_total[smallest_m]:,.2f}, {len(packets[smallest_m])} stays). Output!B9 is "
                          "formatted as Text because Excel may read a value like 2025-01 written into a General cell as a date (January 2025). "
                          "The check accepts either form. In the solution, the `On Error GoTo CleanUp` line sends any error to the "
                          "cleanup block, which turns ScreenUpdating back on before reporting the problem.")),
    ]

    # ------------------------------------------------------------------ extra files
    L.extra_files = {
        "starter/EncounterMacros.bas": _crlf(_starter_module()),
        "starter/Snippets.bas": _crlf(_snippets_module()),
        "solutions/EncounterMacros_Solution.bas": _crlf(_solution_module()),
        "solutions/MonthlyPackets_Solution.bas": _crlf(_bonus_module()),
    }

    # ------------------------------------------------------------------ custom sheets + self-test simulation
    enc_headers = list(enc.headers)

    def write_block(ws, top_row, recs, fmt_from=None):
        for j, h in enumerate(ENC_COLS, 1):
            ws.cell(row=top_row, column=j, value=h)
        for i, e in enumerate(recs, top_row + 1):
            for j, h in enumerate(ENC_COLS, 1):
                c = ws.cell(row=i, column=j, value=e[h])
                if h in ("AdmitDateTime", "DischargeDateTime"):
                    c.number_format = "mm/dd/yyyy hh:mm"
                elif h == "TotalCharges":
                    c.number_format = "#,##0.00"

    @L.customize
    def _sheets(wb, lesson, selftest):
        # ---- Output
        ws = wb.create_sheet("Output")
        ws.sheet_properties.tabColor = "2E75B6"
        ws["A1"] = "Output: your macros write their results here"
        ws["A1"].font = Font(bold=True, size=14, color=DARK)
        ws["A2"] = ("Macros write into the blue cells. The Practice and Bonus sheets read them, so keep them where they are and "
                    "don't type in them by hand.")
        ws["A2"].font = Font(italic=True, color="595959")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells("A2:C2")
        ws.row_dimensions[2].height = 32
        for coord, text in {"A3": "Result", "B3": "Value", "C3": "Written by (task)"}.items():
            c = ws[coord]
            c.value = text
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HDR_FILL
            c.border = BOX
        items = [
            ("rows", "Data rows on the facility sheets", "SortAndReconcile (10)", "#,##0"),
            ("charges", "TotalCharges on the facility sheets", "SortAndReconcile (10)", "#,##0.00"),
            ("file", "Exported file name", "ExportFacility (11)", None),
            ("export", "TotalCharges read back from the exported file", "ExportFacility (11)", "#,##0.00"),
            ("best", "Bonus: packet with the largest total charges", "BuildMonthlyPackets (bonus)", "@"),
        ]
        for key, label, who, nf in items:
            r = int(OUT[key][1:])
            ws.cell(row=r, column=1, value=label).border = BOX
            v = ws.cell(row=r, column=2)
            v.fill = OUT_FILL
            v.border = BOX
            if nf:
                v.number_format = nf
            w = ws.cell(row=r, column=3, value=who)
            w.font = Font(italic=True, color="7F7F7F")
            w.border = BOX
        for col, wdt in zip("ABC", (46, 26, 28)):
            ws.column_dimensions[col].width = wdt

        # ---- PacketTemplate (bonus)
        tp = wb.create_sheet("PacketTemplate")
        tp.sheet_properties.tabColor = "BF9000"
        tp["A1"] = "Monthly inpatient packet"
        tp["A1"].font = Font(bold=True, size=14, color=DARK)
        tp["D1"] = "Bluestone Health System · Finance"
        tp["D1"].font = Font(italic=True, color="7F7F7F")
        labels = [("A2", "Month", "mmmm yyyy"), ("A3", "Inpatient encounters", "#,##0"), ("A4", "Total charges", "#,##0.00")]
        for coord, text, nf in labels:
            tp[coord] = text
            tp[coord].font = Font(bold=True)
            val = tp.cell(row=tp[coord].row, column=2)
            val.number_format = nf
            val.fill = OUT_FILL
            val.border = BOX
            val.alignment = Alignment(horizontal="left")
        for j, h in enumerate(ENC_COLS, 1):
            c = tp.cell(row=PACKET_HEADER_ROW, column=j, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HDR_FILL
        for r in range(PACKET_FIRST, PACKET_FIRST + 150):          # pre-formatted rows: the macro only writes values
            for j, h in enumerate(ENC_COLS, 1):
                if h in ("AdmitDateTime", "DischargeDateTime"):
                    tp.cell(row=r, column=j).number_format = "mm/dd/yyyy hh:mm"
                elif h == "TotalCharges":
                    tp.cell(row=r, column=j).number_format = "#,##0.00"
        for j, h in enumerate(ENC_COLS, 1):
            tp.column_dimensions[get_column_letter(j)].width = {"AdmitSource": 21, "AdmitDateTime": 17,
                                                                "DischargeDateTime": 18, "EncounterType": 14,
                                                                "PrimaryDxCode": 15, "TotalCharges": 14}.get(h, 13)
        tp.column_dimensions["A"].width = 20
        tp.column_dimensions["B"].width = 16      # B2 shows the month as "September 2025"; narrower shows ####
        tp.freeze_panes = f"A{PACKET_FIRST}"

        # ---- Snippets
        sn = wb.create_sheet("Snippets")
        sn.sheet_properties.tabColor = "7030A0"
        sn.column_dimensions["A"].width = 100
        sn["A1"] = "Snippets for Practice tasks 1–6"
        sn["A1"].font = Font(bold=True, size=14, color=DARK)
        sn["A2"] = ("Predict what each snippet prints (or which error stops it), type your prediction on the Practice sheet, then "
                    "run it from the Snippets module (starter/Snippets.bas) to check. Debug.Print writes to the Immediate window, "
                    "which Ctrl + G opens (Mac: View → Immediate Window).")
        sn["A2"].font = Font(italic=True, color="595959")
        sn["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        sn.row_dimensions[2].height = 48
        code_fill = PatternFill("solid", fgColor="F2F2F2")
        r = 4
        for key, s in SNIPPETS.items():
            c = sn.cell(row=r, column=1, value=f"Snippet {key}  ·  Practice {s['tasks']}")
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HDR_FILL
            r += 1
            for line in s["code"].split("\n"):
                c = sn.cell(row=r, column=1, value=line)
                c.font = Font(name="Consolas", size=10)
                c.fill = code_fill
                c.data_type = "s"
                r += 1
            r += 1

        # Print each helper sheet one page wide, in landscape like Practice (on screen nothing changes).
        for sheet in (sn, ws):
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0

        _describe_sheets(wb, lesson, n)

        if not selftest:
            return
        # ---- self-test: build what the macros would build, from the same Python data as the answers
        s = wb.create_sheet(SUMMARY_SHEET)
        s.append(["EncounterType", "Encounters", "TotalCharges"])
        for t in TYPES:
            s.append([t, type_count[t], type_total[t]])
        for f in fac_ids:
            write_block(wb.create_sheet(f), 1, fac_rows[f])
        for m in range(1, 13):
            p = wb.create_sheet(pk(m))
            p["A1"] = "Monthly inpatient packet"
            p["B2"] = datetime(2025, m, 1)
            p["B3"] = len(packets[m])
            p["B4"] = pk_total[m]
            write_block(p, PACKET_HEADER_ROW, packets[m])

    # The hidden keys sit BEFORE PacketTemplate, so Worksheets(Worksheets.Count) is a visible sheet and the macros'
    # "add/copy after the last sheet" never targets a hidden sheet. The visible tab order is unchanged.
    L.sheet_order = ["Start Here", "Practice", "Snippets", "Encounters", "Facilities", "Output", "Bonus", "Answer Key",
                     "Bonus Key", "PacketTemplate"]

    # Keep the README's snippet listings in sync with the code above.
    readme = L.dir / "README.md"
    if readme.exists():
        text = readme.read_text(encoding="utf-8")
        missing = [k for k, sn in SNIPPETS.items() if "TODO" not in text and sn["code"] not in text]
        if missing:
            print(f"  warning: README is missing the current code for snippet(s) {', '.join(missing)}")
    return L


# ---------------------------------------------------------------------------
# Smoke test in LibreOffice (not part of the build)
# ---------------------------------------------------------------------------
def smoke_test(save_dir: str | None = None):
    """Run the reference macros in LibreOffice (VBA-compatibility mode) and compare them with the Python answers.

    LibreOffice Basic can't run three Excel members this lesson uses: ListObjects, copying a multi-area range
    (SpecialCells(xlCellTypeVisible).Copy) and clearing a filter with AutoFilterMode = False. So SplitByFacility and
    ExportFacility (which needs a facility sheet and Worksheet.Copy into a new workbook) can't be smoke-tested here. The test
    pre-builds F01-F04 in export order, exactly as SplitByFacility leaves them, and runs everything else on that copy.
    """
    import tempfile
    from pathlib import Path

    from openpyxl import load_workbook

    from xlcourse import vba

    L = build()
    out_dir = Path(save_dir or tempfile.mkdtemp(prefix="xlc_5_3_"))
    out_dir.mkdir(parents=True, exist_ok=True)
    sol = L.dir / "solutions" / "EncounterMacros_Solution.bas"
    bon = L.dir / "solutions" / "MonthlyPackets_Solution.bas"

    rows = _sample()
    wb = load_workbook(L.dir / L.workbook_name)
    for f in ("F01", "F02", "F03", "F04"):
        ws = wb.create_sheet(f)
        ws.append(ENC_COLS + [LOS_COL])
        for e in rows:
            if e["FacilityID"] == f:
                ws.append([e[h] for h in ENC_COLS])
    start = out_dir / "smoke_start.xlsx"
    wb.save(start)

    read = ([("TypeSummary", f"C{r}") for r in range(2, 6)] + [("F03", "A2"), ("Output", "B4"), ("Output", "B5"),
            ("Output", "B9"), ("Encounters", "L2001")]
            + [(f"2025-{m:02d}", c) for m in (3, 7, 11) for c in ("B3", "B4", "K7")])
    res = vba.run(
        workbook=start, modules=[sol, bon],
        macros=["BuildTypeSummary", "BuildTypeSummary", "SortAndReconcile", "FillLOSDays", "BuildMonthlyPackets"],
        read=read, save_as=out_dir / "smoke.xlsx",
    )
    cells = res["cells"]
    tasks = {str(i): t for i, t in enumerate(L.tasks, 1)}
    tasks.update({f"B{i}": t for i, t in enumerate(L.bonus, 1)})
    los_last = _los(rows[-1])
    expect = [
        ("7  TypeSummary Inpatient", cells[("TypeSummary", "C3")], tasks["7"].answer),
        ("9  F03!A2 after sort", cells[("F03", "A2")], tasks["9"].answer),
        ("10 Output!B5", cells[("Output", "B5")], tasks["10"].answer),
        ("10 Output!B4", cells[("Output", "B4")], len(rows)),
        ("12 LOSDays last row", cells[("Encounters", "L2001")], los_last),
        ("B1 2025-03 B3", cells[("2025-03", "B3")], tasks["B1"].answer),
        ("B2 2025-07 B4", cells[("2025-07", "B4")], tasks["B2"].answer),
        ("B3 2025-11 K7", cells[("2025-11", "K7")], tasks["B3"].answer),
        ("B4 Output!B9", cells[("Output", "B9")], tasks["B4"].answer),
    ]
    ok = True
    for label, got, want in expect:
        good = (abs(float(got) - float(want)) < 0.006) if isinstance(want, (int, float)) and got not in (None, "") \
            else str(got) == str(want)
        ok &= good
        print(f"  {'OK  ' if good else 'FAIL'} {label}: got {got!r}, expected {want!r}")
    print("  errors:", res["errors"] or "none")
    print("  smoke test:", "PASS" if ok and not res["errors"] else "CHECK")
    return res
