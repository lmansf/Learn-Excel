"""Lesson 4.2 · Advanced Formulas: LET, LAMBDA & Array Logic."""
from __future__ import annotations

from collections import defaultdict
from datetime import date, timedelta

from openpyxl.styles import Alignment, Font, PatternFill

from xlcourse import Lesson, Task, data
from xlcourse.lesson import INPUT_FILL, NAVY

CODE = "4.2"

EXCLUDED = ("Expired", "Transfer to Another Hospital", "Left AMA")
HRRP = ("I21.4", "I50.9", "J18.9", "J44.1")
FACS = {"F01": "Bluestone Memorial Hospital", "F02": "Ashby Falls Community Hospital", "F03": "Cedar Ridge Medical Center"}


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="04-advanced-analysis", slug="02-advanced-formulas-let-lambda",
        title="Advanced Formulas: LET, LAMBDA & Array Logic", level="Advanced", minutes=65,
        objectives=[
            "Write multi-condition array logic with SUMPRODUCT and boolean math",
            "Make complex formulas readable and fast with LET",
            "Create reusable custom functions with LAMBDA and the Name Manager",
            "Use MAP, REDUCE, SCAN, BYROW, and audit formulas with Evaluate Formula",
        ],
        data_note="All 5,586 inpatient stays at Bluestone's three hospitals, Jan 2024–Dec 2025 (with the Readmit30 flag "
                  "removed so you can rebuild it), the diagnosis lookup with benchmark LOS, 1,018 lab results from Ashby Falls "
                  "stays discharged Jul–Dec 2025, and the Bluestone Memorial ICU's daily census for 2024–2025.",
    )
    L.start_notes = [
        "Tasks that use LAMBDA, MAP, BYROW, SCAN, or REDUCE need Microsoft 365 or Excel 2024. In older versions they show #NAME?, "
        "so read those solutions in the lesson README instead. LET, XLOOKUP, and FILTER need Excel 2021 or later.",
        "Answer cells must hold ONE value. Wrap array results in SUM, MAX, or INDEX, and use the Sandbox sheet to watch arrays spill.",
        "The Stays sheet has 5,586 rows, and some of these formulas compare every row with every other row. "
        "Excel may pause for a second or two while it recalculates.",
        "Two extra sheets: 'Audit' holds a colleague's broken formula for Task 13, and 'Sandbox' is blank space for "
        "watching array formulas spill.",
    ]

    # ------------------------------------------------------------------ data
    enc = data.load("encounters")
    stays = [e for e in enc if e["EncounterType"] == "Inpatient"]
    stays.sort(key=lambda e: (e["AdmitDateTime"], e["EncounterID"]))
    st = L.add_table_sheet(
        "Stays", stays, table="tblStays",
        columns=["EncounterID", "PatientID", "FacilityID", "PrimaryDxCode", "AdmitDateTime", "DischargeDateTime",
                 "DischargeDisposition", "TotalCharges"],
        extra_cols=["Readmit30"], formats={"Readmit30": "0"},
        widths={"AdmitDateTime": 17, "DischargeDateTime": 18, "DischargeDisposition": 27, "Readmit30": 11},
    )
    dx_rows = sorted(data.load("diagnoses"), key=lambda d: d["DxCode"])
    dxs = L.add_table_sheet("Diagnoses", dx_rows, table="tblDx",
                            columns=["DxCode", "DxDescription", "DxCategory", "ExpectedLOS"],
                            formats={"ExpectedLOS": "0.0"}, widths={"DxDescription": 48})
    dx = {d["DxCode"]: d for d in dx_rows}

    ip_by_id = {e["EncounterID"]: e for e in stays}

    def lab_keep(l):
        e = ip_by_id.get(l["EncounterID"])
        return (e is not None and e["FacilityID"] == "F02" and e["DischargeDateTime"].year == 2025
                and e["DischargeDateTime"].month >= 7)
    labs = [l for l in data.load("lab_results") if lab_keep(l)]
    labs.sort(key=lambda l: (l["CollectedDateTime"], l["LabResultID"]))
    lb = L.add_table_sheet(
        "Labs", labs, table="tblLabs",
        columns=["LabResultID", "EncounterID", "PatientID", "TestCode", "TestName", "ResultValue", "RefLow", "RefHigh",
                 "Units", "Priority", "CollectedDateTime", "ResultedDateTime"],
        widths={"TestName": 24, "CollectedDateTime": 18, "ResultedDateTime": 18},
    )
    census = [r for r in data.load("daily_census") if r["DeptID"] == "D130"]
    census.sort(key=lambda r: r["CensusDate"])
    cen = L.add_table_sheet("Census", census, table="tblCensus",
                            columns=["CensusDate", "StaffedBeds", "Admissions", "Discharges", "MidnightCensus"])

    # ------------------------------------------------------------------ shared Python helpers
    def nights(e):  # midnights crossed: discharge date - admit date
        return (e["DischargeDateTime"].date() - e["AdmitDateTime"].date()).days

    def los(e):  # LOS days: midnights, minimum 1 (same-day stays count as 1 day)
        return max(1, nights(e))

    d25 = [e for e in stays if e["DischargeDateTime"].year == 2025]

    by_patient = defaultdict(list)
    for e in stays:
        by_patient[e["PatientID"]].append(e)

    def readmitted(e):  # dataset rule: later inpatient admit, 0-30 days after discharge (by date), any facility
        d = e["DischargeDateTime"]
        return any(x["AdmitDateTime"] > d and (x["AdmitDateTime"].date() - d.date()).days <= 30
                   for x in by_patient[e["PatientID"]])

    flags = [1 if readmitted(e) else 0 for e in stays]
    assert sum(flags) == sum(1 for e in stays if e["Readmit30"] == "Y"), "readmission rule must match the dataset flag"
    assert all((f == 1) == (e["Readmit30"] == "Y") for f, e in zip(flags, stays))
    # what a datetime-based window (admit <= discharge + 30 days, to the minute) would give instead
    naive_flags = sum(1 for e in stays if any(e["DischargeDateTime"] < x["AdmitDateTime"] <= e["DischargeDateTime"] + timedelta(days=30)
                                              for x in by_patient[e["PatientID"]]))

    S = "tblStays"
    FAC, DX, ADM, DIS, DISPO, CHG = (f"{S}[FacilityID]", f"{S}[PrimaryDxCode]", f"{S}[AdmitDateTime]",
                                     f"{S}[DischargeDateTime]", f"{S}[DischargeDisposition]", f"{S}[TotalCharges]")
    PID = f"{S}[PatientID]"
    first, last = st.first_row, st.last_row

    # T1 weekend admissions in 2025
    t1 = sum(1 for e in stays if e["AdmitDateTime"].year == 2025 and e["AdmitDateTime"].isoweekday() >= 6)

    # T2 OR logic: LOS >= 7 OR discharged to SNF, 2025 discharges, total charges (each stay once)
    review = [e for e in d25 if los(e) >= 7 or e["DischargeDisposition"] == "Skilled Nursing Facility"]
    t2 = round(sum(e["TotalCharges"] for e in review), 2)
    t2_double = round(sum(e["TotalCharges"] for e in d25 if los(e) >= 7)
                      + sum(e["TotalCharges"] for e in d25 if e["DischargeDisposition"] == "Skilled Nursing Facility"), 2)
    t2_both = sum(1 for e in d25 if los(e) >= 7 and e["DischargeDisposition"] == "Skilled Nursing Facility")

    # T3 multi-criteria XLOOKUP: most recent CREAT for PT11882 (Labs sorted oldest -> newest)
    t3_pt, t3_test = "PT11882", "CREAT"
    t3_matches = [l for l in labs if l["PatientID"] == t3_pt and l["TestCode"] == t3_test]
    assert len(t3_matches) >= 3 and t3_matches[0]["ResultValue"] != t3_matches[-1]["ResultValue"]
    t3 = t3_matches[-1]["ResultValue"]
    t3_first = t3_matches[0]["ResultValue"]

    # T4 FREQUENCY of LOS days, 2025 discharges, bins {1,2,3,5,7,14} -> 7th count (> 14 days)
    bins = [1, 2, 3, 5, 7, 14]
    freq = [0] * (len(bins) + 1)
    for e in d25:
        v = nights(e)  # same bin as LOS days: 0 and 1 both land in the "<= 1" bin
        freq[next((i for i, b in enumerate(bins) if v <= b), len(bins))] += 1
    t4 = freq[-1]
    assert t4 == sum(1 for e in d25 if los(e) > 14)

    # T5 readmission flag column
    t5 = sum(flags)

    # T6 LET average LOS: sepsis at F01, 2025 discharges
    t6_set = [los(e) for e in d25 if e["FacilityID"] == "F01" and e["PrimaryDxCode"] == "A41.9"]
    t6 = round(sum(t6_set) / len(t6_set), 2)

    # T7 O/E LOS index: F03, 2025 discharges
    t7_set = [e for e in d25 if e["FacilityID"] == "F03"]
    t7 = sum(los(e) for e in t7_set) / sum(dx[e["PrimaryDxCode"]]["ExpectedLOS"] for e in t7_set)

    # T8 LOSDAYS total: F02, 2025 discharges
    t8_set = [e for e in d25 if e["FacilityID"] == "F02"]
    t8 = sum(los(e) for e in t8_set)
    t8_maxbug = max(los(e) for e in stays) * len(t8_set)  # what MAX(1, …) on whole columns would return

    # T9 MAP: F01 2025 discharges spanning fewer than 2 midnights
    t9 = sum(1 for e in d25 if e["FacilityID"] == "F01" and nights(e) < 2)

    # T10 BYROW: labs outside reference range
    t10 = sum(1 for l in labs if l["ResultValue"] < l["RefLow"] or l["ResultValue"] > l["RefHigh"])
    assert t10 == sum(1 for l in labs if l["AbnormalFlag"] != "N")

    # T11 SCAN: longest run of consecutive ICU days at >= 90% occupancy
    run = t11 = 0
    for r in census:
        run = run + 1 if r["MidnightCensus"] / r["StaffedBeds"] >= 0.9 else 0
        t11 = max(t11, run)

    # T12 REDUCE: HRRP cohort at F01, 2025 discharges
    t12 = sum(1 for e in d25 if e["FacilityID"] == "F01" and e["PrimaryDxCode"] in HRRP)

    # T13 audit fix: F02, 2025 discharges to Home Health
    t13 = sum(1 for e in d25 if e["FacilityID"] == "F02" and e["DischargeDisposition"] == "Home Health")
    t13_f02 = sum(1 for e in stays if e["FacilityID"] == "F02")  # each condition counted alone (audit technique)
    t13_hh = sum(1 for e in stays if e["DischargeDisposition"] == "Home Health")

    # ------------------------------------------------------------------ formulas
    los_let = (f"=LET(\n  adm, {ADM},\n  dis, {DIS},\n  nights, INT(dis)-INT(adm),\n  los, IF(nights<1, 1, nights),\n")
    readmit_fill = (f'=IF(COUNTIFS($B${first}:$B${last},B{first},$E${first}:$E${last},">"&F{first},'
                    f'$E${first}:$E${last},"<"&(INT(F{first})+31))>0,1,0)')
    readmit_struct = (f'=IF(COUNTIFS({PID},[@PatientID],{ADM},">"&[@DischargeDateTime],'
                      f'{ADM},"<"&(INT([@DischargeDateTime])+31))>0,1,0)')
    readmit_live = (f'=SUM(--(COUNTIFS({PID},{PID},{ADM},">"&{DIS},{ADM},"<"&(INT({DIS})+31))>0))')
    rcol = f"Stays!$I${first}:$I${last}"
    assert st.col("Readmit30") == "I" and st.col("PatientID") == "B" and st.col("AdmitDateTime") == "E" \
        and st.col("DischargeDateTime") == "F"

    audit_broken = (f'=SUMPRODUCT(({FAC}="F02")*(YEAR({DIS})="2025")*({DISPO}="Home Health"))')
    audit_fixed = (f'=SUMPRODUCT(({FAC}="F02")*(YEAR({DIS})=2025)*({DISPO}="Home Health"))')

    L.practice_intro = (
        "Every task works on the Excel Tables in this workbook. Table references such as tblStays[FacilityID] are the easiest to "
        f"read, and A1 ranges such as Stays!$C${first}:$C${last} work too. LOS days = discharge date − admit date (midnights), "
        "with a minimum of 1 day. Each answer cell must return ONE value, so try array formulas on the Sandbox sheet first.")

    L.tasks = [
        Task("How many inpatient stays were admitted on a Saturday or Sunday in 2025? Use AdmitDateTime for both the weekday and the year.",
             answer=t1, title="Weekend admissions in 2025",
             solution=f"=SUMPRODUCT((WEEKDAY({ADM},2)>5)*(YEAR({ADM})=2025))",
             hint="WEEKDAY(date, 2) numbers Monday as 1 and Sunday as 7. Multiply two TRUE/FALSE arrays inside SUMPRODUCT.",
             explanation="Each comparison returns an array of 5,586 TRUE/FALSE values. Multiplying the two arrays turns TRUE into 1 and "
                         "FALSE into 0, so a row is 1 only when both conditions hold (AND logic). SUMPRODUCT adds the 1s. COUNTIFS "
                         "could handle the year with two date criteria, but it can't apply WEEKDAY to the range first, so the "
                         "weekend test needs either a helper column or array logic."),
        Task("Case management reviews every 2025 discharge that had LOS days of 7 or more OR went to a Skilled Nursing Facility. "
             "What were the total charges (TotalCharges) of the stays on that review list? Count each stay once, and enter "
             "dollars and cents.",
             answer=t2, fmt="#,##0.00", title="Charges on the case-management review list (OR logic)",
             solution=(f'=SUMPRODUCT((((INT({DIS})-INT({ADM})>=7)+({DISPO}="Skilled Nursing Facility"))>0)'
                       f"*(YEAR({DIS})=2025),{CHG})"),
             hint="Adding two conditions gives OR logic, but a stay that meets both scores 2. Wrap the OR in (…>0).",
             explanation=f"(A)+(B) is 0, 1, or 2. The {t2_both:,} long SNF stays score 2, so multiplying by charges without the "
                         f"`>0` test counts their charges twice and gives ${t2_double:,.2f}. `((A)+(B))>0` turns the sum back into "
                         "TRUE/FALSE, so every stay counts once. LOS days of 7 or more only depends on the midnight count, so "
                         "`INT(dis)-INT(adm)>=7` is enough here (the 1-day minimum only changes 0-night stays)."),
        Task(f"The Labs sheet is sorted from oldest to newest CollectedDateTime. What was patient {t3_pt}'s most recent creatinine "
             f"(TestCode {t3_test}) result?",
             answer=t3, fmt="0.00", title=f"Most recent creatinine for {t3_pt} (multi-criteria XLOOKUP)",
             solution=f'=XLOOKUP(1,(tblLabs[PatientID]="{t3_pt}")*(tblLabs[TestCode]="{t3_test}"),tblLabs[ResultValue],"Not found",0,-1)',
             hint="XLOOKUP can look for 1 in an array of 1s and 0s that you build from two conditions. Its 6th argument chooses the search direction.",
             explanation=f"Multiplying the two conditions builds an array that is 1 only on this patient's creatinine rows. XLOOKUP "
                         f"finds the value 1 in that array. search_mode -1 searches from the bottom up, and because the sheet is sorted "
                         f"oldest to newest, the bottom match is the latest result. A top-down search returns the first result "
                         f"({t3_first}), which is how creatinine looked months earlier. Rising creatinine can signal kidney injury."),
        Task("Use FREQUENCY with the bins {1,2,3,5,7,14} on the LOS days of every stay discharged in 2025. FREQUENCY returns 7 counts. "
             "Enter the 7th: the number of 2025 discharges with LOS days greater than 14.",
             answer=t4, title="FREQUENCY of LOS days: stays longer than 14 days",
             solution=(f"=INDEX(FREQUENCY(FILTER(INT({DIS})-INT({ADM}),YEAR({DIS})=2025),{{1,2,3,5,7,14}}),7)"),
             hint="FILTER the LOS values to 2025 first, then pick one count out of FREQUENCY's result with INDEX.",
             explanation="FREQUENCY counts values ≤ 1, then >1 to 2, >2 to 3, >3 to 5, >5 to 7, >7 to 14, and finally everything "
                         "above the last bin. That extra overflow count is why 6 bins return 7 numbers. INDEX(…, 7) returns just the "
                         "overflow count, so the answer cell holds one value. The solution skips the 1-day minimum because a 0-night "
                         "stay lands in the ≤ 1 bin either way, so applying it gives the same counts. "
                         "`FREQUENCY(IF(YEAR(dis)=2025, INT(dis)-INT(adm)), bins)` also works, because FREQUENCY ignores the FALSE "
                         "values that IF returns for other years."),
        Task("On the Stays sheet, fill the yellow Readmit30 column: 1 if the same patient has another stay whose AdmitDateTime is "
             "after this stay's DischargeDateTime and whose admit DATE is 0–30 days after this stay's discharge DATE, otherwise 0. "
             "The gray cell totals your column. How many stays were followed by a 30-day readmission?",
             answer=t5, title="Readmit30 flag column (stays followed by a readmission)",
             solution=readmit_struct,
             summary=f'=IF(COUNT({rcol})=0,"",SUM({rcol}))',
             fill={"range": f"Stays!I{first}:I{last}", "formula": readmit_fill},
             live=readmit_live,
             hint="COUNTIFS over the whole table, using this row's PatientID and DischargeDateTime as criteria. Compare dates with INT().",
             explanation=f"For each row, COUNTIFS counts the stays that have the same PatientID, an AdmitDateTime after this "
                         f"discharge, and an AdmitDateTime earlier than INT(discharge)+31, which is midnight at the start of the 31st "
                         f"day after the discharge date. That last test is the same as an admit DATE no more than 30 days later. Any "
                         f"count above 0 means a readmission. In A1 style the first row is "
                         f"`{readmit_fill}`. Your total should be {t5}, the same as the official Readmit30 = Y count in the full "
                         f"dataset. A window measured to the minute (admit ≤ discharge + 30) finds only {naive_flags} because it misses "
                         f"readmissions on day 30 that happen later in the day than the discharge did. The key's live formula does "
                         f"all rows in one cell by giving COUNTIFS whole columns as criteria."),
        Task("Write ONE LET formula that names the admit and discharge columns, computes LOS days (minimum 1), and returns the average "
             "LOS days for sepsis stays (PrimaryDxCode A41.9) at Bluestone Memorial (F01) discharged in 2025. Round to 2 decimal "
             "places with ROUND.",
             answer=t6, fmt="0.00", title="Average LOS days for F01 sepsis stays (LET)",
             solution=(los_let + f'  keep, ({FAC}="F01")*({DX}="A41.9")*(YEAR(dis)=2025),\n'
                       "  ROUND(SUM(los*keep)/SUM(keep), 2)\n)"),
             hint="Name each step (nights, then los, then a 1/0 keep array). A conditional average is SUM(los*keep)/SUM(keep).",
             explanation=f"Each name is calculated once and reused, so INT(dis)-INT(adm) isn't repeated. `los*keep` zeroes out the "
                         f"rows you don't want, and SUM(keep) counts the rows you do ({len(t6_set)} stays), which gives a conditional "
                         "average without a helper column. `AVERAGE(FILTER(los, keep))` is an equally good last step."),
        Task("The LOS index (observed ÷ expected) compares actual LOS days with the benchmark ExpectedLOS of each stay's diagnosis "
             "(Diagnoses sheet). Calculate it for Cedar Ridge (F03) stays discharged in 2025: total LOS days ÷ total ExpectedLOS. "
             "Enter the ratio to at least 2 decimal places.",
             answer=t7, fmt="0.00", title="LOS index (O/E) for Cedar Ridge in 2025 (LET + XLOOKUP)",
             solution=(los_let + f"  explos, XLOOKUP({DX}, tblDx[DxCode], tblDx[ExpectedLOS]),\n"
                       f'  keep, ({FAC}="F03")*(YEAR(dis)=2025),\n  SUM(los*keep)/SUM(explos*keep)\n)'),
             hint="Inside LET, XLOOKUP the whole PrimaryDxCode column. You get one ExpectedLOS for every stay.",
             explanation="Given a whole column as its lookup value, XLOOKUP returns one benchmark per stay. Naming that array "
                         "`explos` keeps the final line readable: observed days ÷ expected days for the kept rows. An index above 1.00 "
                         "means patients stayed longer than the benchmark. Don't average the per-stay ratios. A ratio of totals "
                         "weights long and short stays correctly."),
        Task("Create a named function LOSDAYS(admit, discharge) in the Name Manager that returns LOS days (discharge date − admit "
             "date, minimum 1). Then call it on whole columns. What is the total of LOS days for Ashby Falls (F02) stays "
             "discharged in 2025?",
             answer=t8, title="LOSDAYS named LAMBDA: total LOS days for F02 in 2025",
             solution=f'=SUM(LOSDAYS({ADM},{DIS})*({FAC}="F02")*(YEAR({DIS})=2025))',
             live=False, self_test=False,
             hint="Formulas → Name Manager → New. Test the LAMBDA in a cell first, for example =LAMBDA(…)(Stays!E2,Stays!F2).",
             explanation=f"Define LOSDAYS as `=LAMBDA(admit, discharge, LET(nights, INT(discharge)-INT(admit), IF(nights<1, 1, nights)))`. "
                         f"Called with two whole columns, it returns 5,586 LOS values, and the two conditions keep the "
                         f"{len(t8_set)} F02 stays from 2025. Use IF for the minimum, not MAX. `MAX(1, INT(discharge)-INT(admit))` "
                         f"collapses the whole column into one number (the longest stay, {max(los(e) for e in stays)} days), and the "
                         f"total becomes {t8_maxbug:,}."),
        Task("Utilization review audits inpatient stays that spanned fewer than 2 midnights (discharge date − admit date < 2). "
             "Using MAP with a LAMBDA that uses AND, count Bluestone Memorial (F01) stays discharged in 2025 that spanned fewer "
             "than 2 midnights.",
             answer=t9, title="Short stays at F01 in 2025 (MAP)",
             solution=(f"=SUM(MAP({FAC}, {ADM}, {DIS},\n"
                       '  LAMBDA(fac, adm, dis, IF(AND(fac="F01", YEAR(dis)=2025, INT(dis)-INT(adm)<2), 1, 0))))'),
             live=False, self_test=False,
             hint="MAP can take several same-size arrays and passes one value from each to your LAMBDA. Wrap the result in SUM.",
             explanation="MAP calls the LAMBDA once per stay, so `fac`, `adm`, and `dis` are single values and AND works. Outside MAP, "
                         "AND(...) on whole columns would return one TRUE/FALSE for all 5,586 rows. The boolean-math version "
                         "`=SUMPRODUCT((tblStays[FacilityID]=\"F01\")*(YEAR(tblStays[DischargeDateTime])=2025)"
                         "*(INT(tblStays[DischargeDateTime])-INT(tblStays[AdmitDateTime])<2))` gives the same count. MAP is "
                         "worth it when the per-row logic reads better with AND, OR, or MAX."),
        Task("On the Labs sheet, count the results outside their reference range (ResultValue below RefLow or above RefHigh). "
             "Use BYROW over the three adjacent columns ResultValue:RefHigh with a LAMBDA that uses OR.",
             answer=t10, title="Lab results outside the reference range (BYROW)",
             solution=("=SUM(BYROW(tblLabs[[ResultValue]:[RefHigh]],\n"
                       "  LAMBDA(r, IF(OR(INDEX(r,1)<INDEX(r,2), INDEX(r,1)>INDEX(r,3)), 1, 0))))"),
             live=False, self_test=False,
             hint="Inside BYROW, the LAMBDA gets one row of 3 cells: INDEX(r,1) is the value, INDEX(r,2) the low, INDEX(r,3) the high.",
             explanation="BYROW hands the LAMBDA one 1×3 row at a time, so OR tests just that row and returns one TRUE/FALSE. "
                         "BYROW stacks the results into a column of 1s and 0s for SUM. `tblLabs[[ResultValue]:[RefHigh]]` is the "
                         "structured reference for the three adjacent columns. Cross-check: the lab's own abnormal flags (not "
                         f"included here) mark the same {t10} results."),
        Task("The Census sheet holds the Bluestone Memorial ICU's daily census for 2024–2025. A day is 'strained' when "
             "MidnightCensus ÷ StaffedBeds is 90% or more. Using SCAN, find the longest run of consecutive strained days.",
             answer=t11, title="Longest run of strained ICU days (SCAN)",
             solution=("=MAX(SCAN(0, tblCensus[MidnightCensus]/tblCensus[StaffedBeds]>=0.9,\n"
                       "  LAMBDA(run, strained, IF(strained, run+1, 0))))"),
             live=False, self_test=False,
             hint="Keep a running count that adds 1 on a strained day and resets to 0 otherwise. Then take the MAX.",
             explanation="SCAN walks down the 731 TRUE/FALSE values and keeps a running number (the accumulator `run`). On a strained "
                         "day it adds 1, and on any other day it resets to 0. The result is one running count per day, so its MAX is "
                         "the longest streak. A streak depends on the previous row's result, which SUMPRODUCT and COUNTIFS can't "
                         "express without a helper column."),
        Task("Medicare's Hospital Readmissions Reduction Program (HRRP) tracks readmissions after heart attack (AMI, I21.4), heart "
             "failure (I50.9), pneumonia (J18.9), and COPD (chronic obstructive pulmonary disease, J44.1), among other conditions. "
             'Using REDUCE to loop over the array {"I21.4","I50.9","J18.9","J44.1"}, count Bluestone Memorial (F01) stays '
             "discharged in 2025 with any of these primary diagnoses.",
             answer=t12, title="HRRP-condition stays at F01 in 2025 (REDUCE)",
             solution=('=REDUCE(0, {"I21.4","I50.9","J18.9","J44.1"},\n'
                       f'  LAMBDA(total, dxcode, total + COUNTIFS({DX}, dxcode, {FAC}, "F01",\n'
                       f'    {DIS}, ">="&DATE(2025,1,1), {DIS}, "<"&DATE(2026,1,1))))'),
             live=False, self_test=False,
             hint="REDUCE(0, codes, LAMBDA(total, dxcode, total + …)) starts at 0 and adds one COUNTIFS per code.",
             explanation="REDUCE starts the accumulator `total` at 0, then calls the LAMBDA once per code, adding that code's "
                         "COUNTIFS each time. It returns only the final total. The codes can't overlap (a stay has one primary "
                         "diagnosis), so adding is safe. `SUM(COUNTIFS(…, {codes}, …))` gives the same answer. REDUCE earns its keep "
                         "when each step needs more logic than one function call."),
        Task("The Audit sheet has a colleague's formula that should count Ashby Falls (F02) stays discharged in 2025 to Home Health, "
             "but it returns 0. Step through it with Evaluate Formula (on a Mac, test its pieces on the Sandbox sheet instead), "
             "fix the bug, and enter the correct count.",
             answer=t13, title="Audit: fix the colleague's formula",
             solution=audit_fixed,
             hint="Check the data type on each side of every comparison. Text shows in quotes. Count each condition's TRUE values with SUM(--(…)).",
             explanation=f"YEAR returns the number 2025, but the formula compares it with the text \"2025\" (in quotes). In Excel a "
                         f"number never equals text, so that comparison is FALSE on every row and every product is 0. In Evaluate "
                         f"Formula the clue is the data types: the YEAR step shows plain numbers such as 2024, and they're compared "
                         f"with \"2025\" in quotes. Counting each condition on the Sandbox sheet proves it: "
                         f"`=SUM(--(tblStays[FacilityID]=\"F02\"))` returns {t13_f02}, the Home Health test returns {t13_hh}, but "
                         f"`=SUM(--(YEAR(tblStays[DischargeDateTime])=\"2025\"))` returns 0. Without the quotes it returns "
                         f"{len(d25):,}. COUNTIFS would have accepted \"2025\" because it converts criteria strings, and that habit is "
                         f"how this bug usually gets in."),
    ]

    # ------------------------------------------------------------------ bonus
    idx = [e for e in stays if date(2025, 1, 1) <= e["DischargeDateTime"].date() <= date(2025, 11, 30)
           and e["DischargeDisposition"] not in EXCLUDED]
    b1 = len(idx)

    def rate(rows):
        return sum(1 for e in rows if readmitted(e)) / len(rows)

    rates = {f: rate([e for e in idx if e["FacilityID"] == f]) for f in FACS}
    best = max(rates, key=rates.get)
    assert best != "F01" and sorted(rates.values())[-1] - sorted(rates.values())[-2] > 0.001
    hrrp_idx = [e for e in idx if e["PrimaryDxCode"] in HRRP]
    b5 = rate(hrrp_idx)
    raw_2025 = [e for e in stays if date(2025, 1, 1) <= e["DischargeDateTime"].date() <= date(2025, 11, 30)]

    excl_arr = '{"Expired","Transfer to Another Hospital","Left AMA"}'
    readmit_line = f'  readmit, COUNTIFS({PID}, {PID}, adm, ">"&dis, adm, "<"&(INT(dis)+31))>0,\n'
    eligible_line = (f"  eligible, (dis>=DATE(2025,1,1))*(dis<DATE(2025,12,1))\n"
                     f"            *ISNA(MATCH({DISPO}, {excl_arr}, 0)),\n")

    def rate_let(fac: str) -> str:
        return (f"=LET(\n  adm, {ADM},\n  dis, {DIS},\n" + readmit_line + eligible_line
                + f'  keep, eligible*({FAC}="{fac}"),\n  SUM(keep*readmit)/SUM(keep)\n)')

    hrrp_arr = "{" + ",".join(f'"{c}"' for c in HRRP) + "}"
    L.bonus_title = "Bonus: CMS-style readmission rates in one formula"
    L.bonus_scenario = (
        "Bluestone's quality committee wants 30-day all-cause readmission rates by hospital, built with simplified versions of the "
        "rules used by CMS (the Centers for Medicare & Medicaid Services, which also risk-adjusts its rates and ignores planned "
        "readmissions). Index stays are inpatient stays discharged from 01/01/2025 through 11/30/2025, so every stay has a full 30 days of "
        "follow-up in the data. Exclude stays that ended in death (Expired), a transfer (Transfer to Another Hospital), or a "
        "discharge against medical advice (Left AMA). A readmission is any later inpatient admission of the same patient, at any "
        "Bluestone hospital, with an admit date 0–30 days after the index discharge date (the Task 5 rule). "
        "Rate = index stays followed by a readmission ÷ index stays. Build each answer as one LET formula over tblStays.")
    L.bonus = [
        Task("How many index stays are there system-wide?", answer=b1, title="Eligible index stays (system-wide)",
             solution=(f"=LET(\n  dis, {DIS},\n" + eligible_line + "  SUM(eligible)\n)"),
             hint="ISNA(MATCH(disposition, {list}, 0)) is TRUE when a disposition is NOT on the list.",
             explanation=f"MATCH returns #N/A for every disposition that isn't one of the three exclusions, so ISNA turns \"not "
                         f"excluded\" into TRUE. The two date tests keep discharges from 01/01/2025 up to (not including) "
                         f"12/01/2025. Without the exclusions there would be {len(raw_2025):,} stays in the window."),
        Task("Write one LET formula that returns the readmission rate for Bluestone Memorial (F01). Enter it as a percentage "
             "(1 decimal place is enough).", answer=rates["F01"], fmt="0.0%", title="Readmission rate for F01",
             solution=rate_let("F01"),
             hint="Name readmit (the Task 5 test for every row at once, as in guide section 4c), eligible, and keep. Then divide "
                  "SUM(keep*readmit) by SUM(keep).",
             explanation=f"`readmit` hands COUNTIFS whole columns as criteria, so it returns one count per stay (5,586 of them), and "
                         f"`>0` turns those counts into TRUE/FALSE. LET calculates that expensive array once. `keep` narrows the "
                         f"index stays to F01 ({sum(1 for e in idx if e['FacilityID'] == 'F01'):,} stays). If you finished Task 5, "
                         f"`SUM(keep*{S}[Readmit30])/SUM(keep)` is a shorter last step that reuses your column. The self-contained "
                         f"version keeps working even if someone deletes that column."),
        Task("Which hospital has the highest rate? Enter its FacilityID.", answer=best, accept=[FACS[best]],
             title="Hospital with the highest readmission rate",
             solution=("=LET(\n  facs, {\"F01\";\"F02\";\"F03\"},\n  rates, MAP(facs, LAMBDA(f, READMITRATE(f))),\n"
                       "  XLOOKUP(MAX(rates), rates, facs)\n)"),
             live=False, self_test=False,
             hint='Wrap your B2 formula in LAMBDA(facility, …), save it as READMITRATE, and MAP it over {"F01";"F02";"F03"}. '
                  "Or edit the facility ID three times.",
             explanation=(f"Define READMITRATE in the Name Manager as your B2 formula wrapped in `=LAMBDA(facility, LET(…))`, with "
                          f"\"F01\" replaced by `facility`. MAP passes each ID in turn to `LAMBDA(f, READMITRATE(f))`, and XLOOKUP "
                          f"returns the ID next to the largest rate. The rates are " + ", ".join(f"{f} {rates[f]:.1%}" for f in FACS) + f". {FACS[best]} has the highest "
                          f"rate even though it's the smallest hospital, which is why rates, not counts, are compared.")),
        Task("What is that hospital's rate? Enter it as a percentage (1 decimal place is enough).", answer=rates[best], fmt="0.0%",
             title="Highest hospital readmission rate", solution=rate_let(best),
             hint="Reuse B2 with the other facility ID, or call your READMITRATE function.",
             explanation=f"Same formula as B2 with \"{best}\". With the named LAMBDA it is simply `=READMITRATE(\"{best}\")`. "
                         f"Turning a long LET into a LAMBDA means one tested definition serves every hospital, so a fix to the "
                         f"rule only has to be made once."),
        Task("CMS also publishes readmission rates by condition. What is the system-wide rate for index stays whose primary diagnosis "
             "is one of the four HRRP conditions from Task 12? Enter it as a percentage (1 decimal place is enough).",
             answer=b5, fmt="0.0%", title="System-wide readmission rate for HRRP conditions",
             solution=(f"=LET(\n  adm, {ADM},\n  dis, {DIS},\n" + readmit_line + eligible_line
                       + f"  keep, eligible*ISNUMBER(MATCH({DX}, {hrrp_arr}, 0)),\n  SUM(keep*readmit)/SUM(keep)\n)"),
             hint="Swap the facility test for a diagnosis-list test: ISNUMBER(MATCH(dx, {list}, 0)).",
             explanation=f"ISNUMBER(MATCH(…)) is the mirror image of the exclusion test: TRUE when the diagnosis IS on the list. "
                         f"The {len(hrrp_idx)} HRRP index stays are readmitted more often than index stays overall "
                         f"({rate(idx):.1%}), which is why CMS targets these conditions."),
    ]

    # ------------------------------------------------------------------ extra sheets
    @L.customize
    def extra_sheets(wb, lesson, selftest):
        ws = wb.create_sheet("Audit")
        ws.sheet_properties.tabColor = "C55A11"
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 110
        ws["A1"] = "Audit: a formula that needs fixing"
        ws["A1"].font = Font(bold=True, size=14, color=NAVY)
        ws["A3"] = "Request"
        ws["B3"] = "Count the Ashby Falls (F02) inpatient stays discharged in 2025 that went to Home Health."
        ws["A4"] = "Colleague's note"
        ws["B4"] = "\"My formula returns 0, but I know Ashby Falls sends patients home with home health every week. Can you find the bug?\""
        ws["A6"] = "The formula"
        lesson.set_formula(ws, "B6", audit_broken)
        ws["B6"].fill = INPUT_FILL
        ws["B6"].alignment = Alignment(horizontal="left")
        ws["A7"] = "Formula text"
        ws["B7"] = audit_broken
        ws["B7"].data_type = "s"
        ws["B7"].font = Font(name="Consolas", size=10)
        ws["B7"].alignment = Alignment(wrap_text=True, vertical="top")
        ws["B7"]._style.quotePrefix = 1
        ws.row_dimensions[7].height = 32
        steps = [
            "1. Select B6, then choose Formulas → Evaluate Formula (Formula Auditing group). Evaluate Formula is in Excel for "
            "Windows only.",
            "2. Click Evaluate repeatedly. Watch the underlined part of the formula turn into its result.",
            "3. At each comparison, check the data type on both sides. Text appears in quotes, such as \"F02\", and numbers "
            "appear without quotes. Only the first values of an array fit on screen, and the first rows are 2024 stays, so "
            "FALSE at the start of the 2025 test is normal.",
            "4. To test one condition on all 5,586 rows, count its TRUE values on the Sandbox sheet, for example "
            "=SUM(--(tblStays[FacilityID]=\"F02\")). A condition that counts 0 is never TRUE.",
            "5. Fix the formula in B6 and check that the result is believable.",
            "6. Enter the corrected count in Task 13 on the Practice sheet.",
            "On a Mac (no Evaluate Formula)? Skip steps 1–3 and use step 4 on each of the three conditions in the formula.",
        ]
        ws["A9"] = "How to audit"
        for i, s in enumerate(steps):
            c = ws.cell(row=9 + i, column=2, value=s)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[9 + i].height = 15 * max(1, -(-len(s) // 100))
        for r in (3, 4, 6, 7, 9):
            ws.cell(row=r, column=1).font = Font(bold=True)
            ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
        for r in (3, 4):
            ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[4].height = 30

        sb = wb.create_sheet("Sandbox")
        sb.sheet_properties.tabColor = "7F7F7F"
        sb.column_dimensions["A"].width = 22  # the notes overflow into the empty cells to the right
        sb["A1"] = "Sandbox: try array formulas here"
        sb["A1"].font = Font(bold=True, size=14, color=NAVY)
        notes = [
            "Answer cells on the Practice and Bonus sheets must return ONE value, so build and inspect your array formulas here first.",
            "Type a formula in A10 (or any cell with empty cells below it) and watch it spill. Examples to try:",
            "   =INT(tblStays[DischargeDateTime])-INT(tblStays[AdmitDateTime])       (midnights for every stay)",
            "   =SCAN(0, tblCensus[Admissions], LAMBDA(total, x, total + x))       (running total of ICU admissions)",
            "When the spill looks right, wrap it in SUM, MAX, or INDEX and copy the formula to the Practice sheet.",
        ]
        for i, s in enumerate(notes, 2):
            sb.cell(row=i, column=1, value=s).font = Font(italic=i > 2, color="404040")
        sb["A9"] = "Your workspace ↓"
        sb["A9"].font = Font(bold=True, color=NAVY)

    L.sheet_order = ["Start Here", "Practice", "Stays", "Diagnoses", "Labs", "Census", "Audit", "Sandbox", "Bonus",
                     "Answer Key", "Bonus Key"]
    return L
