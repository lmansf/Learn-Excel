"""Lesson 3.4 · PivotTables.

Learners build PivotTables from tblEncounters (every encounter that began in 2025) and type the numbers the
pivots show. A pivot can't be checked by a formula until the learner has built it, so:

* Each pivot task's ``solution`` is Markdown build steps, and the answer is computed in Python.
* Each task's ``live`` formula is an independent cross-check written with COUNTIFS/SUMIFS/AVERAGEIFS on the
  Table, so the verifier proves the Python answer in LibreOffice and learners see the "formula twin" of every
  pivot in the Answer Key.
* The self-test types the expected value into each answer cell (the default for non-formula solutions).
* Task 13 is a GETPIVOTDATA formula that points at the learner's own pivot sheet, so its solution is Markdown
  (steps plus the formula) rather than a formula string; nothing in the workbook refers to a sheet that doesn't
  exist yet.
* A hidden "Pivot Key" sheet (built in a customize hook from Python values) shows each finished pivot in full,
  so learners can compare their whole layout, not just one number.
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from xlcourse import Lesson, Task, data
from xlcourse.lesson import BOX, HEADER_FILL, NAVY, _plain

CODE = "3.4"
TBL = "tblEncounters"
YEAR = 2025

# Slices used by the tasks (kept here so prompts, answers and cross-check formulas can't drift apart).
T1_FACILITY = "Cedar Ridge Medical Center"
T2_SERVICE_LINE = "Cardiovascular"
T4_PAYER_TYPE = "Self-Pay"
T5_SERVICE_LINE = "Medicine"
T6_PAYER = "State Medicaid"
T10_BAND = (70, 79)
T11_PAYER_TYPES = ("Government", "Medicare Advantage")
T11_FACILITY = "Cedar Ridge Medical Center"
T11_QUARTER = 4
T12_FACILITY = "Ashby Falls Community Hospital"
T13_FACILITY = "Cedar Ridge Medical Center"
T13_TYPE = "Observation"
B_MIN_STAYS = 30
B_TOP_N = 3

WORDS = {2: "two", 3: "three", 4: "four", 5: "five", 6: "six"}
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTHS_LONG = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
               "November", "December"]

COLUMNS = ["EncounterID", "PatientID", "EncounterType", "AdmitDate", "DischargeDate", "LOSDays", "FacilityName",
           "DeptName", "ServiceLine", "PayerName", "PayerType", "DxCode", "DxDescription", "DxCategory", "AgeAtAdmit",
           "TotalCharges", "Readmit30", "ReadmitFlag"]

DICTIONARY = [
    ("EncounterID", "Unique ID of the encounter. One row = one encounter, so Count of EncounterID counts encounters.", "ENC112345"),
    ("PatientID", "The patient. One patient can have several encounters, so counting PatientID still counts encounters, "
                  "not people.", "PT10493"),
    ("EncounterType", "Inpatient, Observation, Emergency (an ED visit that didn't become an inpatient or observation stay: "
                      "discharged home, left without being seen, left against advice, or transferred out), or Outpatient "
                      "(clinic visit).", "Inpatient"),
    ("AdmitDate", "Date the encounter began. For patients admitted through the ED, the date of the admit decision. "
                  "Every row is in 2025.", "01/03/2025"),
    ("DischargeDate", "Date the encounter ended.", "01/07/2025"),
    ("LOSDays", "Length of stay in days = DischargeDate − AdmitDate (midnights). 0 for same-day visits.", "4"),
    ("FacilityName", "Hospital or outpatient site.", "Bluestone Memorial Hospital"),
    ("DeptName", "Unit or clinic. Names repeat across facilities (each hospital has an Emergency Department), so pair "
                 "it with FacilityName.", "Cardiac Step-Down"),
    ("ServiceLine", "Clinical service line of the department.", "Cardiovascular"),
    ("PayerName", "Insurance payer billed for the encounter.", "State Medicaid"),
    ("PayerType", "Government (Medicare, State Medicaid), Medicare Advantage, Commercial, Self-Pay, or Workers' Comp.",
     "Government"),
    ("DxCode", "Primary diagnosis, ICD-10-CM code.", "I50.9"),
    ("DxDescription", "Primary diagnosis description.", "Heart failure, unspecified"),
    ("DxCategory", "Body-system category of the diagnosis.", "Circulatory"),
    ("AgeAtAdmit", "Patient's age in completed years on AdmitDate.", "72"),
    ("TotalCharges", "Gross billed charges in dollars (not what the payer pays).", "20,407.87"),
    ("Readmit30", "Inpatient stays only: Y if the same patient had another inpatient admission 0–30 days after this "
                  "discharge, otherwise N. Blank for other encounter types. Stays ending in death are always N.", "N"),
    ("ReadmitFlag", "Readmit30 as a number: 1 = Y, 0 = N, blank for non-inpatient rows. Sum = readmissions, "
                    "Count = index stays, Average = readmission rate.", "0"),
]


def _age(dob: date, on: date) -> int:
    return on.year - dob.year - ((on.month, on.day) < (dob.month, dob.day))


def encounter_rows() -> list[dict]:
    """Every encounter that began in 2025, with names pre-joined (this lesson isn't about lookups)."""
    fac = data.index(data.load("facilities"), "FacilityID")
    dep = data.index(data.load("departments"), "DeptID")
    pay = data.index(data.load("payers"), "PayerID")
    dx = data.index(data.load("diagnoses"), "DxCode")
    pat = data.index(data.load("patients"), "PatientID")
    enc = [e for e in data.load("encounters") if e["AdmitDateTime"].year == YEAR]
    enc.sort(key=lambda e: (e["AdmitDateTime"], e["EncounterID"]))
    rows = []
    for e in enc:
        admit, disch = e["AdmitDateTime"].date(), e["DischargeDateTime"].date()
        flag = None if e["Readmit30"] is None else (1 if e["Readmit30"] == "Y" else 0)
        rows.append({
            "EncounterID": e["EncounterID"], "PatientID": e["PatientID"], "EncounterType": e["EncounterType"],
            "AdmitDate": admit, "DischargeDate": disch, "LOSDays": (disch - admit).days,
            "FacilityName": fac[e["FacilityID"]]["FacilityName"], "DeptName": dep[e["DeptID"]]["DeptName"],
            "ServiceLine": dep[e["DeptID"]]["ServiceLine"], "PayerName": pay[e["PayerID"]]["PayerName"],
            "PayerType": pay[e["PayerID"]]["PayerType"], "DxCode": e["PrimaryDxCode"],
            "DxDescription": dx[e["PrimaryDxCode"]]["DxDescription"], "DxCategory": dx[e["PrimaryDxCode"]]["DxCategory"],
            "AgeAtAdmit": _age(pat[e["PatientID"]]["DOB"], admit), "TotalCharges": e["TotalCharges"],
            "Readmit30": e["Readmit30"], "ReadmitFlag": flag,
        })
    return rows


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="03-data-analysis", slug="04-pivottables",
        title="PivotTables", level="Intermediate", minutes=60,
        objectives=[
            "Build PivotTables from a Table and arrange rows, columns, values, and filters",
            "Change summaries (Sum, Count, Average) and Show Values As (% of total, difference, running total)",
            "Group dates and numbers; filter with slicers and timelines",
            "Add calculated fields and use GETPIVOTDATA",
        ],
    )

    # ------------------------------------------------------------------ data
    rows = encounter_rows()
    n = len(rows)
    assert all(r["AdmitDate"].year == YEAR for r in rows)
    L.data_note = (f"All {n:,} encounters that began in 2025 at Bluestone's four facilities (one row per encounter), with "
                   "facility, department, service line, payer, and diagnosis names already joined in, plus age, length of "
                   "stay, charges, and a 30-day readmission flag.")
    sd = L.add_table_sheet(
        "Encounters", rows, columns=COLUMNS, table=TBL,
        formats={"AdmitDate": "mm/dd/yyyy", "DischargeDate": "mm/dd/yyyy", "LOSDays": "0", "AgeAtAdmit": "0",
                 "TotalCharges": "#,##0.00", "ReadmitFlag": "0"},
        widths={"EncounterID": 13, "PatientID": 11, "EncounterType": 14, "AdmitDate": 12, "DischargeDate": 14,
                "LOSDays": 9, "FacilityName": 31, "DeptName": 31, "ServiceLine": 18, "PayerName": 29, "PayerType": 19,
                "DxCode": 9, "DxDescription": 44, "DxCategory": 25, "AgeAtAdmit": 12, "TotalCharges": 14,
                "Readmit30": 11, "ReadmitFlag": 12},
    )
    L.add_table_sheet("Data Dictionary", [{"Column": c, "Meaning": m, "Example": x} for c, m, x in DICTIONARY],
                      columns=["Column", "Meaning", "Example"], as_table=False,
                      widths={"Column": 16, "Meaning": 110, "Example": 28})

    def col(name: str) -> str:
        """Structured reference to a Table column (used by the cross-check formulas)."""
        assert name in COLUMNS, name
        return f"{TBL}[{name}]"

    def q(s: str) -> str:
        return '"' + s.replace('"', '""') + '"'

    def cross(formula: str) -> str:
        return f"Cross-check without a pivot: `{formula}`"

    by_type = defaultdict(list)
    for r in rows:
        by_type[r["EncounterType"]].append(r)
    ip, ed = by_type["Inpatient"], by_type["Emergency"]
    assert all(r["ReadmitFlag"] is not None for r in ip)
    assert all(r["ReadmitFlag"] is None for r in rows if r["EncounterType"] != "Inpatient")
    facilities = sorted({r["FacilityName"] for r in rows})
    payer_names = sorted({r["PayerName"] for r in rows})
    types = ["Emergency", "Inpatient", "Observation", "Outpatient"]

    def month_counts(subset):
        c = Counter(r["AdmitDate"].month for r in subset)
        return [c.get(m, 0) for m in range(1, 13)]

    # ------------------------------------------------------------------ answers (computed in Python)
    # T1 count by facility
    t1 = sum(1 for r in rows if r["FacilityName"] == T1_FACILITY)

    # T2 sum of inpatient charges for one service line
    t2_rows = [r for r in ip if r["ServiceLine"] == T2_SERVICE_LINE]
    t2 = round(sum(r["TotalCharges"] for r in t2_rows), 2)

    # T3 drill-down: most expensive stay in that service line (must be unique, and its charge unique in the table)
    t3_top = max(t2_rows, key=lambda r: r["TotalCharges"])
    assert sorted(r["TotalCharges"] for r in t2_rows)[-2] < t3_top["TotalCharges"]
    assert sum(1 for r in rows if r["TotalCharges"] == t3_top["TotalCharges"]) == 1
    t3 = t3_top["EncounterID"]

    # T4 average ED charges for one payer type
    t4_vals = [r["TotalCharges"] for r in ed if r["PayerType"] == T4_PAYER_TYPE]
    t4 = sum(t4_vals) / len(t4_vals)

    # T5 inpatient readmission rate for one service line (average of a 1/0 flag)
    t5_vals = [r["ReadmitFlag"] for r in ip if r["ServiceLine"] == T5_SERVICE_LINE]
    t5 = sum(t5_vals) / len(t5_vals)

    # T6 % of column total: share of ED encounters billed to one payer
    t6 = sum(1 for r in ed if r["PayerName"] == T6_PAYER) / len(ed)

    # T7 month with the fewest ED visits (must be unique)
    ed_by_month = month_counts(ed)
    t7_idx = ed_by_month.index(min(ed_by_month))
    assert sorted(ed_by_month)[0] < sorted(ed_by_month)[1], "fewest-ED month must be unique"
    t7 = MONTHS[t7_idx]

    # T8 Difference From (previous): December minus November ED visits
    t8 = ed_by_month[11] - ed_by_month[10]

    # T9 Running Total In months: inpatient stays that began January–June
    ip_by_month = month_counts(ip)
    ip_running = [sum(ip_by_month[:i + 1]) for i in range(12)]
    t9 = ip_running[5]

    # T10 age band count (inpatient)
    lo_age, hi_age = T10_BAND
    t10 = sum(1 for r in ip if lo_age <= r["AgeAtAdmit"] <= hi_age)
    max_age = max(r["AgeAtAdmit"] for r in rows)
    assert min(r["AgeAtAdmit"] for r in rows) == 0 and max_age == 99, "grouping prompt assumes ages 0–99"

    # T11 slicer (two payer types) + timeline (Q4) → one facility's ED visits
    q_start = date(YEAR, 3 * (T11_QUARTER - 1) + 1, 1)
    t11_rows = [r for r in ed if r["PayerType"] in T11_PAYER_TYPES and r["FacilityName"] == T11_FACILITY
                and r["AdmitDate"] >= q_start]
    t11 = len(t11_rows)

    # T12 calculated field ChargesPerDay = TotalCharges / LOSDays (sums first, then divides)
    def cpd(subset):
        return sum(r["TotalCharges"] for r in subset) / sum(r["LOSDays"] for r in subset)
    t12_rows = [r for r in ip if r["FacilityName"] == T12_FACILITY]
    t12 = cpd(t12_rows)
    t12_naive_rows = [r for r in t12_rows if r["LOSDays"] > 0]
    t12_naive = sum(r["TotalCharges"] / r["LOSDays"] for r in t12_naive_rows) / len(t12_naive_rows)
    t12_zero = sum(1 for r in t12_rows if r["LOSDays"] == 0)
    zero_txt = "its one same-day stay would return" if t12_zero == 1 else f"its {t12_zero} same-day stays would return"

    # T13 GETPIVOTDATA: Sum of TotalCharges for one facility × encounter type
    t13 = round(sum(r["TotalCharges"] for r in rows if r["FacilityName"] == T13_FACILITY and r["EncounterType"] == T13_TYPE), 2)

    # ------------------------------------------------------------------ bonus answers
    combos = defaultdict(list)
    for r in ip:
        combos[(r["ServiceLine"], r["PayerType"])].append(r["ReadmitFlag"])
    stats = {k: (len(v), sum(v), sum(v) / len(v)) for k, v in combos.items()}
    kept = {k: s for k, s in stats.items() if s[0] >= B_MIN_STAYS}
    b1 = len(kept)
    ranked = sorted(kept.items(), key=lambda kv: -kv[1][2])
    assert ranked[0][1][2] > ranked[1][1][2] + 1e-9, "top combination must be unique"
    (b2_sl, b2_pt), (b2_n, b2_ra, b2_rate) = ranked[0]
    b2 = f"{b2_sl}, {b2_pt}"
    small_rivals = sorted(((s[2], s[0], s[1], k) for k, s in stats.items() if s[0] < B_MIN_STAYS), reverse=True)
    dx_readmits = Counter(r["DxDescription"] for r in ip if r["ReadmitFlag"] == 1)
    dx_ranked = dx_readmits.most_common()
    assert dx_ranked[B_TOP_N - 1][1] > dx_ranked[B_TOP_N][1], "top-3 cut must not tie"
    assert dx_ranked[B_TOP_N - 2][1] > dx_ranked[B_TOP_N - 1][1], "3rd place must be unique"
    assert {d for d, _ in dx_ranked[:B_TOP_N]} == {  # the B5 explanation names these three conditions
        "Pneumonia, unspecified organism", "Chronic obstructive pulmonary disease with (acute) exacerbation",
        "Heart failure, unspecified"}
    b4 = dx_ranked[B_TOP_N - 1][0]
    b4_code = next(r["DxCode"] for r in ip if r["DxDescription"] == b4)
    total_readmits = sum(r["ReadmitFlag"] for r in ip)
    top_sum = sum(c for _, c in dx_ranked[:B_TOP_N])
    b5 = top_sum / total_readmits

    # ------------------------------------------------------------------ cross-check formulas
    ET, FN, SL, PT, PN = col("EncounterType"), col("FacilityName"), col("ServiceLine"), col("PayerType"), col("PayerName")
    AD, TC, RF, AGE, LOS = col("AdmitDate"), col("TotalCharges"), col("ReadmitFlag"), col("AgeAtAdmit"), col("LOSDays")
    f1 = f"=COUNTIFS({FN},{q(T1_FACILITY)})"
    f2 = f'=SUMIFS({TC},{ET},"Inpatient",{SL},{q(T2_SERVICE_LINE)})'
    f3 = (f'=INDEX({col("EncounterID")},MATCH(MAXIFS({TC},{ET},"Inpatient",{SL},{q(T2_SERVICE_LINE)}),{TC},0))')
    f4 = f'=AVERAGEIFS({TC},{ET},"Emergency",{PT},{q(T4_PAYER_TYPE)})'
    f5 = f'=AVERAGEIFS({RF},{ET},"Inpatient",{SL},{q(T5_SERVICE_LINE)})'
    f6 = f'=COUNTIFS({ET},"Emergency",{PN},{q(T6_PAYER)})/COUNTIFS({ET},"Emergency")'
    months_const = "{" + ",".join(str(m) for m in range(1, 13)) + "}"
    ed_month_counts = (f'COUNTIFS({ET},"Emergency",{AD},">="&DATE({YEAR},{months_const},1),'
                       f'{AD},"<"&DATE({YEAR},{months_const}+1,1))')
    f7 = f'=TEXT(DATE({YEAR},MATCH(MIN({ed_month_counts}),{ed_month_counts},0),1),"mmm")'
    f7_one = (f'=COUNTIFS({ET},"Emergency",{AD},">="&DATE({YEAR},{t7_idx + 1},1),'
              f'{AD},"<"&DATE({YEAR},{t7_idx + 2},1))')
    f8 = (f'=COUNTIFS({ET},"Emergency",{AD},">="&DATE({YEAR},12,1))'
          f'-COUNTIFS({ET},"Emergency",{AD},">="&DATE({YEAR},11,1),{AD},"<"&DATE({YEAR},12,1))')
    f9 = f'=COUNTIFS({ET},"Inpatient",{AD},"<"&DATE({YEAR},7,1))'
    f10 = f'=COUNTIFS({ET},"Inpatient",{AGE},">={lo_age}",{AGE},"<={hi_age}")'
    f11 = (f'=SUM(COUNTIFS({ET},"Emergency",{FN},{q(T11_FACILITY)},{PT},'
           + "{" + ",".join(q(p) for p in T11_PAYER_TYPES) + "}"
           + f',{AD},">="&DATE({YEAR},{q_start.month},1)))')
    f12 = (f'=SUMIFS({TC},{ET},"Inpatient",{FN},{q(T12_FACILITY)})'
           f'/SUMIFS({LOS},{ET},"Inpatient",{FN},{q(T12_FACILITY)})')
    f13 = f'=SUMIFS({TC},{FN},{q(T13_FACILITY)},{ET},{q(T13_TYPE)})'
    fb1 = (f'=LET(s,UNIQUE({SL}),p,TRANSPOSE(UNIQUE({PT})),'
           f'n,COUNTIFS({ET},"Inpatient",{SL},s,{PT},p),SUM(--(n>={B_MIN_STAYS})))')
    fb3 = f'=AVERAGEIFS({RF},{ET},"Inpatient",{SL},{q(b2_sl)},{PT},{q(b2_pt)})'
    dx_counts = (f'LET(d,UNIQUE(FILTER({col("DxDescription")},{RF}=1)),'
                 f'n,COUNTIFS({col("DxDescription")},d,{RF},1),')
    fb4 = f"={dx_counts}INDEX(SORTBY(d,n,-1),{B_TOP_N}))"
    fb5 = f"={dx_counts}SUM(LARGE(n,SEQUENCE({B_TOP_N})))/SUM({RF}))"

    # ------------------------------------------------------------------ practice tasks
    L.practice_intro = (
        f"Every task uses {TBL} on the Encounters sheet: all {n:,} encounters that began in 2025. Build each PivotTable on a "
        "new worksheet (**Insert → PivotTable**, then **New Worksheet**), or rearrange the one you already have. When you reuse a pivot, "
        "make it match the layout the task lists: remove leftover fields and clear filters the task doesn't mention. Then "
        "type the answer the pivot shows into the yellow cell. Type the value itself rather than a reference to a pivot "
        "cell, because a reference like =B7 points somewhere else as soon as you rearrange the pivot. Task 13 is the "
        "exception.")
    L.start_notes = [
        "You'll create several PivotTable sheets as you work. That's expected. Delete the ones you no longer need, or keep "
        "rearranging one pivot.",
        "The hidden 'Pivot Key' sheet shows each finished pivot in full, so you can compare your whole layout, not just one "
        "number. Unhide it the same way as the Answer Key.",
    ]
    L.practice_how = ("Go to the 'Practice' sheet. Build each PivotTable on a new worksheet (or rearrange one you already "
                      "have), then type what it shows into the task's yellow cell. Task 13 asks for a GETPIVOTDATA formula.")
    L.bonus_where = "Build the pivots on new sheets, and type your answers in the yellow cells on the **Bonus** sheet."

    steps_new = ("Click any cell in tblEncounters on the Encounters sheet, then choose **Insert → PivotTable** "
                 "(Microsoft 365 for Windows: **Insert → PivotTable → From Table/Range**). Check that Table/Range says "
                 f"`{TBL}`, choose **New Worksheet**, and click **OK**.")
    fewest = sorted(ed_by_month)
    L.tasks = [
        Task(f"Create a PivotTable from {TBL} on a new worksheet. Put FacilityName in Rows and EncounterID in Values. "
             f"How many 2025 encounters did {T1_FACILITY} have?",
             answer=t1, hint="**Insert → PivotTable**, then drag fields into the four areas",
             solution=f"1. {steps_new}\n2. In the PivotTable Fields pane, drag **FacilityName** to **Rows** and **EncounterID** "
                      f"to **Values**. Excel names the value field *Count of EncounterID*.\n3. Read the {T1_FACILITY} row.",
             live=f1,
             explanation="EncounterID is text, so Excel summarizes it with Count, which adds 1 for every non-empty cell. Each "
                         f"row is one encounter, so the Grand Total ({n:,}) equals the number of rows in the Table. That's a "
                         "quick check that the pivot sees all of the data. " + cross(f1)),
        Task(f"Rearrange the pivot. Remove FacilityName, put EncounterType in Filters and select Inpatient, put ServiceLine in "
             f"Rows, and put TotalCharges in Values in place of EncounterID. What were the total charges for "
             f"{T2_SERVICE_LINE} inpatient stays? Enter the amount to the cent.",
             answer=t2, fmt="#,##0.00", hint="The filter button for the Filters area appears above the pivot",
             solution="1. Uncheck **FacilityName** in the field list (or drag it out of Rows).\n"
                      "2. Drag **EncounterType** to **Filters**. Open the filter button that appears in B1, pick **Inpatient**, "
                      "and click **OK**.\n3. Drag **ServiceLine** to **Rows** and **TotalCharges** to **Values**. Remove "
                      f"*Count of EncounterID* from Values.\n4. Read the {T2_SERVICE_LINE} row of *Sum of TotalCharges*.",
             live=f2,
             explanation="TotalCharges contains only numbers, so Excel sums it by default. The Filters area filters the whole "
                         "pivot without adding rows or columns, which keeps the layout simple when you only need one slice. "
                         + cross(f2)),
        Task(f"Drill down: in that pivot, double-click the Sum of TotalCharges value on the {T2_SERVICE_LINE} row. Excel lists "
             "the stays behind the number on a new sheet. Sort that list by TotalCharges, largest first. What is the "
             f"EncounterID of the most expensive {T2_SERVICE_LINE} inpatient stay?",
             answer=t3, hint="Double-click a value cell (**Show Details**)",
             solution=f"1. Double-click the *Sum of TotalCharges* cell for {T2_SERVICE_LINE}. Excel inserts a new sheet with a "
                      f"Table of the {len(t2_rows)} matching rows.\n2. Right-click any TotalCharges value in that Table and "
                      "choose **Sort → Sort Largest to Smallest** (or use the TotalCharges filter button).\n3. Read the "
                      "EncounterID in the first row.",
             live=f3,
             explanation=f"Drilling down answers \"which records make up this number?\" The detail sheet holds "
                         f"{len(t2_rows)} rows, exactly the stays in the total, and the top one is a "
                         f"{t3_top['LOSDays']}-day stay for {t3_top['DxDescription']} with charges of "
                         f"${t3_top['TotalCharges']:,.2f}. The detail sheet is a static copy: it doesn't update when the data "
                         "changes, so delete it when you're done. " + cross(f3)),
        Task("Build a pivot of average ED charges by payer type: EncounterType = Emergency in Filters, PayerType in Rows, and "
             "TotalCharges in Values. Change the summary from Sum to Average. What was the average charge for a "
             f"{T4_PAYER_TYPE} ED visit? Round to 2 decimal places.",
             answer=t4, fmt="#,##0.00", hint="Right-click a value → **Summarize Values By**, or **Value Field Settings**",
             solution="1. Set the **EncounterType** filter to **Emergency** and put **PayerType** in **Rows** and "
                      "**TotalCharges** in **Values**.\n2. Right-click any number in the pivot and choose **Summarize Values "
                      "By → Average** (or **Value Field Settings → Summarize Values By → Average**).\n3. Click **Number "
                      "Format** in Value Field Settings and choose Number with 2 decimals, or Currency.\n"
                      f"4. Read the {T4_PAYER_TYPE} row.",
             live=f4,
             explanation=f"Average divides the sum of the charges by the number of visits in each row ({len(t4_vals)} Self-Pay "
                         "ED visits). The Grand Total row is the average over every ED visit, not the average of the "
                         "payer-type averages, so it weights each payer type by its volume. " + cross(f4)),
        Task("Build a readmission pivot: EncounterType = Inpatient in Filters, ServiceLine in Rows, and ReadmitFlag in Values. "
             "Notice which summary Excel picks, then change it to Average and format it as a percentage. What was the "
             f"30-day readmission rate for the {T5_SERVICE_LINE} service line? Enter it as a percentage to 1 decimal place.",
             answer=t5, fmt="0.0%", hint="Average of a 1/0 column is the share of 1s",
             solution="1. Set the **EncounterType** filter to **Inpatient**, put **ServiceLine** in **Rows**, and drag "
                      "**ReadmitFlag** to **Values**. Excel shows *Count of ReadmitFlag*.\n2. Open **Value Field Settings**, "
                      "choose **Average**, click **Number Format**, and pick **Percentage** with 1 decimal place.\n"
                      f"3. Read the {T5_SERVICE_LINE} row.",
             live=f5,
             explanation="Excel picks **Count** because ReadmitFlag has blank cells (every non-inpatient row), and a pivot only "
                         "defaults to Sum when a column is 100% numbers. Count of ReadmitFlag is the number of index stays, Sum "
                         "is the number of readmissions, and Average (Sum ÷ Count) is the readmission rate: "
                         f"{sum(t5_vals)} ÷ {len(t5_vals):,} for {T5_SERVICE_LINE}. Average ignores blanks, so the "
                         "non-inpatient rows can't dilute the rate. " + cross(f5)),
        Task("Build a payer-mix pivot on a new sheet, with no filters: PayerName in Rows, EncounterType in Columns, and "
             "EncounterID in Values. Show the values as % of Column Total. What percentage of ED visits (the Emergency "
             f"column) were billed to {T6_PAYER}? Enter it to 1 decimal place.",
             answer=t6, fmt="0.0%", hint="**Value Field Settings → Show Values As**",
             solution="1. Start a new pivot. If you reuse the last one instead, clear the EncounterType filter first: a field "
                      "moved from Filters to Columns can keep its old selection and hide the Emergency column. Put "
                      "**PayerName** in **Rows**, **EncounterType** in **Columns**, and **EncounterID** in **Values**.\n"
                      "2. Right-click a number → "
                      "**Show Values As → % of Column Total**.\n"
                      f"3. Read the {T6_PAYER} row in the Emergency column.",
             live=f6,
             explanation="% of Column Total divides each cell by its column's total, so every column adds up to 100% and you "
                         "read each encounter type's payer mix down the column. % of Row Total would answer a different "
                         "question: what share of that payer's encounters were ED visits. Filtering EncounterType to "
                         "Emergency and choosing % of Grand Total gives the same answer. " + cross(f6)),
        Task("Group dates by month: EncounterType = Emergency in Filters, AdmitDate in Rows, and EncounterID in Values. Group "
             "AdmitDate by Months. Which month of 2025 had the fewest ED visits? Type the month's three-letter name as the "
             "pivot shows it (for example, Mar).",
             answer=t7, accept=[MONTHS_LONG[t7_idx], str(t7_idx + 1)],
             hint="Right-click a date → **Group…**, then sort by the count",
             solution="1. Set the **EncounterType** filter to **Emergency**, put **AdmitDate** in **Rows**, and put "
                      "**EncounterID** in **Values**.\n2. If you see individual dates, right-click one → **Group…**, select "
                      "**Months** only, and click **OK**. (Excel 2016 and later may group by month automatically.)\n"
                      "3. Right-click a count → **Sort → Sort Smallest to Largest**. The first month is the answer.",
             live=f7,
             explanation=f"{t7} had {fewest[0]} ED visits, just below {MONTHS[ed_by_month.index(fewest[1])]} with "
                         f"{fewest[1]}. Grouping sorts {len(ed):,} ED visits into 12 monthly buckets without a helper "
                         "column. Sorting by value reorders the months, so remember to sort back (**Sort A to Z** on the "
                         "month labels keeps calendar order) before you compare one month with the next. "
                         + cross(f7_one) + f" returns {fewest[0]}. Repeat it for each month, or see the key's live "
                         "formula, which checks all twelve at once."),
        Task("In the same pivot, first sort the months back into calendar order (Jan at the top). Then add EncounterID to "
             "Values a second time and show it as Difference From the (previous) month. By how many visits did December's "
             "ED volume differ from November's? Type a negative number if December was lower.",
             answer=t8, hint="Right-click a month → **Sort → Sort A to Z**. Then **Show Values As → Difference From**, Base "
                             "item (previous)",
             solution="1. Sort the months back into calendar order (right-click a month → **Sort → Sort A to Z**).\n"
                      "2. Drag **EncounterID** into **Values** again. Right-click one of the new numbers → **Show Values As → "
                      "Difference From…**\n3. Base field: the field that shows the months (**AdmitDate**, or "
                      "**Months (AdmitDate)** if Excel created it). Base item: **(previous)**. Click **OK**.\n"
                      "4. Read the Dec row of the new column.",
             live=f8,
             explanation=f"Difference From (previous) subtracts the month above: {ed_by_month[11]} − {ed_by_month[10]} = "
                         f"{t8}. January is blank because it has no previous month in the data. % Difference From would "
                         f"show the same change as a percentage ({t8 / ed_by_month[10]:.1%}). " + cross(f8)),
        Task("Change the EncounterType filter to Inpatient, and change the second value field to Running Total In the months "
             "field. How many inpatient stays began from January 1 through June 30, 2025 (the running total on the Jun row)?",
             answer=t9, hint="**Show Values As → Running Total In**",
             solution="1. Set the **EncounterType** filter to **Inpatient**.\n2. Right-click a number in the second value "
                      "column → **Show Values As → Running Total In…** → Base field: the months field → **OK**.\n"
                      "3. Read the Jun row.",
             live=f9,
             explanation="Running Total In adds each month to everything above it, so the Jun row is the year-to-date total "
                         f"at the end of June, and the Dec row equals the Grand Total ({len(ip):,}). The running total follows "
                         "the order of the months, which is another reason to keep them in calendar order. " + cross(f9)),
        Task("Group numbers into bands: EncounterType = Inpatient in Filters, AgeAtAdmit in Rows, and EncounterID in Values. "
             f"Group AgeAtAdmit starting at 0, ending at {max_age}, by 10. How many inpatient stays were for patients aged "
             f"{lo_age}–{hi_age}?",
             answer=t10, hint="Right-click an age → **Group…** (Starting at, Ending at, By)",
             solution="1. Set the **EncounterType** filter to **Inpatient**, put **AgeAtAdmit** in **Rows**, and "
                      "**EncounterID** in **Values**.\n2. Right-click any age → **Group…** → Starting at **0**, Ending at "
                      f"**{max_age}**, By **10** → **OK**.\n3. Read the **{lo_age}-{hi_age}** row.",
             live=f10,
             explanation=f"Grouping turns {len({r['AgeAtAdmit'] for r in ip})} distinct ages into ten bands labeled "
                         f"0-9, 10-19, … {max_age - 9}-{max_age}. Each band includes both ends, so {lo_age}-{hi_age} means "
                         f"ages {lo_age} through {hi_age}. It's the busiest band for inpatient care. " + cross(f10)),
        Task("Build a new pivot with FacilityName in Rows, EncounterType in Columns, and EncounterID in Values. Insert a slicer "
             "for PayerType and a timeline for AdmitDate. In the slicer, select both "
             f"{T11_PAYER_TYPES[0]} and {T11_PAYER_TYPES[1]}. In the timeline, switch to QUARTERS and select "
             f"Q{T11_QUARTER} {YEAR}. How many Emergency encounters does {T11_FACILITY} show?",
             answer=t11,
             hint="**PivotTable Analyze → Insert Slicer** and **Insert Timeline**. Ctrl + click (Mac: ⌘ + click) picks a "
                  "second button",
             solution="1. Build the pivot on a new sheet: **FacilityName** in **Rows**, **EncounterType** in **Columns**, "
                      "**EncounterID** in **Values**.\n2. **PivotTable Analyze → Insert Slicer** → tick **PayerType** → "
                      f"**OK**. Click **{T11_PAYER_TYPES[0]}**, then Ctrl + click (Mac: ⌘ + click) **{T11_PAYER_TYPES[1]}**.\n"
                      "3. **PivotTable Analyze → Insert Timeline** → tick **AdmitDate** → **OK**. Change the time level "
                      f"(top right of the timeline) to **QUARTERS** and click **Q{T11_QUARTER}**.\n"
                      f"4. Read the {T11_FACILITY} row in the Emergency column.",
             live=f11,
             explanation="Buttons selected in one slicer combine with OR (Government or Medicare Advantage), and separate "
                         "filters combine with AND (those payer types and Q4). A timeline is a slicer built for dates: it "
                         "selects a continuous period at the level you choose. "
                         + cross(f11)),
        Task("Add a calculated field named ChargesPerDay with the formula =TotalCharges/LOSDays. Use a pivot with "
             "EncounterType = Inpatient in Filters, FacilityName in Rows, and ChargesPerDay in Values. What is ChargesPerDay "
             f"for {T12_FACILITY}? Round to 2 decimal places.",
             answer=t12, fmt="#,##0.00", hint="**PivotTable Analyze → Fields, Items, & Sets → Calculated Field**",
             solution="1. Click inside a pivot built from tblEncounters that no slicer or timeline is filtering (or build a "
                      "new one). A leftover slicer selection from task 11 would change the result.\n2. **PivotTable Analyze → "
                      "Fields, Items, & Sets → Calculated Field…**\n3. Name: `ChargesPerDay`. Formula: `=TotalCharges/LOSDays` "
                      "(double-click the fields in the list to insert them). Click **Add**, then **OK**.\n4. Set the "
                      "**EncounterType** filter to **Inpatient**, put **FacilityName** in **Rows**, and keep only "
                      f"*Sum of ChargesPerDay* in **Values**.\n5. Read the {T12_FACILITY} row.",
             live=f12,
             explanation="A calculated field adds up each field first and then applies the formula: Sum of TotalCharges ÷ Sum "
                         f"of LOSDays for the {len(t12_rows)} {T12_FACILITY} stays. That's charges per patient day. It is not "
                         "the average of each stay's charges ÷ days. That would be "
                         f"${t12_naive:,.2f}, and {zero_txt} #DIV/0!. Excel labels the field *Sum of ChargesPerDay* even "
                         "though it's a ratio. "
                         + cross(f12)),
        Task("On a new sheet, build a pivot with FacilityName in Rows, EncounterType in Columns, and TotalCharges in Values. "
             "Then click this task's yellow cell, type =, switch to the pivot sheet, click the "
             f"{T13_FACILITY} × {T13_TYPE} cell, and press Enter. Excel writes a GETPIVOTDATA formula. What does it return? "
             "Leave the formula in the cell.",
             answer=t13, fmt="#,##0.00", hint="If you get a plain reference such as =Sheet7!D8 instead, turn **Generate GetPivotData** back on",
             solution="1. Build the pivot on a new sheet (Excel names it something like *Sheet7*).\n2. Click the yellow answer "
                      "cell on the Practice sheet and type `=`.\n"
                      f"3. Switch to the pivot sheet, click the {T13_FACILITY} × {T13_TYPE} cell, and press **Enter**. "
                      "Excel writes a formula like:\n\n"
                      f'```\n=GETPIVOTDATA("TotalCharges",Sheet7!$A$3,"FacilityName","{T13_FACILITY}",'
                      f'"EncounterType","{T13_TYPE}")\n```',
             live=f13,
             explanation="GETPIVOTDATA looks a value up by its field and item names instead of by its cell address, so it "
                         "keeps returning the right number when the pivot is sorted, filtered, or rearranged, as long as the "
                         "item stays visible. If someone filters Cedar Ridge out of the pivot, the formula returns #REF!. "
                         + cross(f13)),
    ]

    # ------------------------------------------------------------------ bonus
    combo_variants = []
    for sep in (", ", ",", " / ", "/", " - ", " – ", " & ", " and ", " ", " x ", " × "):
        combo_variants += [f"{b2_sl}{sep}{b2_pt}", f"{b2_pt}{sep}{b2_sl}"]
    combo_variants = [v for v in dict.fromkeys(combo_variants) if v.lower() != b2.lower()]
    rivals = [x for x in small_rivals if x[0] >= small_rivals[0][0] - 1e-9]  # every small group tied at the top
    rival_parts = [f"{k[0]} / {k[1]} ({ra} of {cnt})" for rate, cnt, ra, k in rivals]
    rival_txt = ", ".join(rival_parts[:-1]) + ", and " + rival_parts[-1]
    rival_rate = small_rivals[0][0]
    assert rival_rate < b2_rate and len(rivals) >= 2
    assert all((ra + 1) / cnt > b2_rate for _, cnt, ra, _ in rivals), "'one more readmission' claim must hold"

    L.bonus_title = "Bonus: Readmissions deep-dive"
    L.bonus_scenario = (
        "Bluestone's Quality Committee is preparing for its annual readmissions review and wants to know where 30-day "
        "readmissions concentrate. An index stay is any inpatient stay that could be followed by a readmission: every row "
        "where ReadmitFlag is 1 or 0. Small groups produce extreme rates by chance (1 readmission among 4 stays is 25%), so "
        f"the committee only reviews groups with at least {B_MIN_STAYS} index stays. Build the pivots on new sheets and "
        "answer the committee's questions.")
    L.bonus = [
        Task("Build a pivot with EncounterType = Inpatient in Filters, ServiceLine and then PayerType in Rows, and ReadmitFlag "
             "in Values twice: once as Count (rename it Index stays) and once as Average (rename it Readmit rate, formatted "
             f"as a percentage). Add a value filter on PayerType that keeps only rows with at least {B_MIN_STAYS} index "
             "stays. How many service line × payer type combinations remain?",
             answer=b1,
             hint="PayerType's filter menu → **Value Filters → Greater Than Or Equal To**. Tabular Form makes rows easy "
                  "to count",
             solution="1. New pivot: **EncounterType** in **Filters** (select **Inpatient**); **ServiceLine** then "
                      "**PayerType** in **Rows**.\n2. Drag **ReadmitFlag** to **Values** twice. In **Value Field Settings**, "
                      "set the first to **Count** with Custom Name `Index stays`, and the second to **Average** with Custom "
                      "Name `Readmit rate` and a percentage number format.\n3. **Design → Report Layout → Show in Tabular "
                      "Form** and **Design → Subtotals → Do Not Show Subtotals** make each combination one row.\n"
                      "4. Open the filter button on the **PayerType** header (in Compact Form, the Row Labels button, then "
                      "choose **PayerType** under *Select field*) → **Value Filters → Greater Than Or Equal To…** → "
                      f"*Index stays* · `{B_MIN_STAYS}` → **OK**.\n5. Count the remaining PayerType rows (select the Readmit "
                      "rate cells above the Grand Total and read **Count** on the status bar).",
             live=fb1,
             explanation="A value filter on the inner row field is applied within each service line, so it keeps or hides "
                         f"each service line × payer type cell separately. {len(stats)} combinations exist in the data, and "
                         f"{len(stats) - b1} of them have fewer than {B_MIN_STAYS} index stays. Custom names keep the two "
                         "ReadmitFlag fields apart and make the filter dialog readable. The key's cross-check uses dynamic-array "
                         "functions (Microsoft 365 and Excel 2021 or later) that you'll meet in Lessons 4.1 and 4.2: " + f"`{fb1}`"),
        Task("Among the remaining combinations, which has the highest readmission rate? Type it as ServiceLine, PayerType "
             "(for example: Medicine, Commercial).",
             answer=b2, accept=combo_variants, hint="Sort the Readmit rate column, or scan it with a color scale",
             solution="1. Right-click a *Readmit rate* value on a payer-type row → **Sort → Sort Largest to Smallest**. In a "
                      "two-level pivot this sorts the payer types *within* each service line, so the highest rate in each "
                      "service line moves to the top of its group.\n2. Compare those top rows, one per service line, and "
                      "pick the highest. (Sorting the service lines as well orders them by their overall rate, which doesn't "
                      "guarantee that the best single combination ends up first.)\n3. Optional: **Home → Conditional Formatting → Color Scales** on "
                      "the Readmit rate cells makes the highest rate stand out.",
             live=False,
             explanation=f"{b2_sl} stays paid by {b2_pt} plans had {b2_ra} readmissions in {b2_n} index stays. "
                         f"{WORDS[len(rivals)].capitalize()} groups too small to qualify sit right behind it at {rival_rate:.1%}: {rival_txt}. "
                         "With one more readmission, any of them would top the list, which is why the committee sets a "
                         "minimum volume before it ranks rates."),
        Task("What is that combination's readmission rate? Enter it as a percentage to 1 decimal place.",
             answer=b2_rate, fmt="0.0%", hint="Read the Readmit rate cell",
             solution=f"Read the *Readmit rate* cell on the {b2_sl} / {b2_pt} row.",
             live=fb3,
             explanation=f"{b2_ra} ÷ {b2_n} = {b2_rate:.1%}, against {total_readmits / len(ip):.1%} for all inpatient stays. "
                         + cross(fb3)),
        Task("Now rank diagnoses. In a new pivot (EncounterType = Inpatient in Filters), put DxDescription in Rows and "
             f"ReadmitFlag in Values as Sum, which is the number of readmissions. Apply a Top 10 filter that keeps the top "
             f"{B_TOP_N} items. Which diagnosis has the third-highest number of readmissions? Type the description exactly "
             "as it appears.",
             answer=b4, accept=[b4_code], hint="**Row Labels** filter → **Value Filters → Top 10…**",
             solution="1. New pivot: **EncounterType** in **Filters** (**Inpatient**), **DxDescription** in **Rows**, "
                      "**ReadmitFlag** in **Values** (change it to **Sum**).\n2. Open the **Row Labels** filter button → "
                      f"**Value Filters → Top 10…** → **Top** `{B_TOP_N}` **Items** by *Sum of ReadmitFlag* → **OK**.\n"
                      "3. Sort largest to smallest and read the third row.",
             live=fb4,
             explanation="Sum of a 1/0 flag counts the 1s, so Sum of ReadmitFlag is the number of readmissions. The Top 10 "
                         "filter keeps the top N items by any value field, not just 10. The top three are "
                         + ", ".join(f"{d} ({c})" for d, c in dx_ranked[:B_TOP_N])
                         + f", and fourth place has {dx_ranked[B_TOP_N][1]}. Cross-check (Microsoft 365 and Excel 2021 or later): "
                         f"`{fb4}`"),
        Task(f"What share of all {YEAR} inpatient readmissions do those {WORDS[B_TOP_N]} diagnoses account for together? Enter it "
             "as a percentage to 1 decimal place.",
             answer=b5, fmt="0.0%", hint="With the Top 3 filter on, the Grand Total adds only the visible rows",
             solution=f"1. Note the Grand Total with the Top {B_TOP_N} filter on ({top_sum} readmissions).\n2. Clear the "
                      "filter (**Row Labels → Clear Filter From \"DxDescription\"**) and note the new Grand Total "
                      f"({total_readmits}).\n3. Divide: {top_sum} ÷ {total_readmits}.",
             live=fb5,
             explanation="A regular PivotTable totals only the items that survive its filters. With the Top 3 filter on, the "
                         f"Grand Total is {top_sum}, so % of Grand Total would show the three diagnoses adding up to 100%. "
                         f"The true denominator is all {total_readmits} readmissions, so about {b5:.0%} of readmissions come "
                         "from just three diagnoses. Pneumonia, COPD, and heart failure are also conditions in Medicare's "
                         "Hospital Readmissions Reduction Program, so quality teams watch them closely. "
                         f"Cross-check (Microsoft 365 and Excel 2021 or later): `{fb5}`"),
    ]

    # ------------------------------------------------------------------ Pivot Key (hidden reference pivots)
    def avg(v):
        return sum(v) / len(v) if v else None

    blocks = []  # (title, headers, rows, formats, total_row)

    by_fac = Counter(r["FacilityName"] for r in rows)
    blocks.append(("Task 1 · Count of EncounterID by FacilityName", ["Row Labels", "Count of EncounterID"],
                   [[f, by_fac[f]] for f in facilities], ["@", "#,##0"], ["Grand Total", n]))
    sl_ip = defaultdict(list)
    for r in ip:
        sl_ip[r["ServiceLine"]].append(r)
    sls_ip = sorted(sl_ip)
    blocks.append(("Tasks 2–3 · EncounterType = Inpatient · Sum of TotalCharges by ServiceLine",
                   ["Row Labels", "Sum of TotalCharges"],
                   [[s, sum(x["TotalCharges"] for x in sl_ip[s])] for s in sls_ip], ["@", "#,##0.00"],
                   ["Grand Total", sum(r["TotalCharges"] for r in ip)]))
    pt_ed = defaultdict(list)
    for r in ed:
        pt_ed[r["PayerType"]].append(r["TotalCharges"])
    blocks.append(("Task 4 · EncounterType = Emergency · Average of TotalCharges by PayerType",
                   ["Row Labels", "Average of TotalCharges"], [[p, avg(pt_ed[p])] for p in sorted(pt_ed)],
                   ["@", "#,##0.00"], ["Grand Total", avg([r["TotalCharges"] for r in ed])]))
    blocks.append(("Task 5 · EncounterType = Inpatient · ReadmitFlag by ServiceLine (Count, Sum, and Average shown together)",
                   ["Row Labels", "Count of ReadmitFlag", "Sum of ReadmitFlag", "Average of ReadmitFlag"],
                   [[s, len(sl_ip[s]), sum(x["ReadmitFlag"] for x in sl_ip[s]), avg([x["ReadmitFlag"] for x in sl_ip[s]])]
                    for s in sls_ip], ["@", "#,##0", "#,##0", "0.0%"],
                   ["Grand Total", len(ip), total_readmits, total_readmits / len(ip)]))
    type_tot = Counter(r["EncounterType"] for r in rows)
    pn_type = Counter((r["PayerName"], r["EncounterType"]) for r in rows)
    pn_tot = Counter(r["PayerName"] for r in rows)
    blocks.append(("Task 6 · Count of EncounterID shown as % of Column Total (PayerName × EncounterType)",
                   ["Row Labels"] + types + ["Grand Total"],
                   [[p] + [(pn_type[(p, t)] / type_tot[t]) if pn_type[(p, t)] else None for t in types] + [pn_tot[p] / n]
                    for p in payer_names],
                   ["@"] + ["0.0%"] * (len(types) + 1), ["Grand Total"] + [1.0] * (len(types) + 1)))
    ed_diff = [None] + [ed_by_month[i] - ed_by_month[i - 1] for i in range(1, 12)]
    blocks.append(("Tasks 7–8 · EncounterType = Emergency · AdmitDate grouped by Months (calendar order)",
                   ["Row Labels", "Count of EncounterID", "Difference From (previous)"],
                   [[MONTHS[i], ed_by_month[i], ed_diff[i]] for i in range(12)], ["@", "#,##0", "+#,##0;-#,##0;0"],
                   ["Grand Total", len(ed), None]))
    blocks.append(("Task 9 · EncounterType = Inpatient · AdmitDate grouped by Months",
                   ["Row Labels", "Count of EncounterID", "Running Total In AdmitDate"],
                   [[MONTHS[i], ip_by_month[i], ip_running[i]] for i in range(12)], ["@", "#,##0", "#,##0"],
                   ["Grand Total", len(ip), None]))
    band = Counter(r["AgeAtAdmit"] // 10 for r in ip)
    blocks.append(("Task 10 · EncounterType = Inpatient · AgeAtAdmit grouped 0 to 99 by 10",
                   ["Row Labels", "Count of EncounterID"],
                   [[f"{b * 10}-{b * 10 + 9}", band[b]] for b in sorted(band)], ["@", "#,##0"], ["Grand Total", len(ip)]))
    sl_rows = [r for r in rows if r["PayerType"] in T11_PAYER_TYPES and r["AdmitDate"] >= q_start]
    sl_c = Counter((r["FacilityName"], r["EncounterType"]) for r in sl_rows)
    sl_types = [t for t in types if any(r["EncounterType"] == t for r in sl_rows)]
    sl_facs = sorted({r["FacilityName"] for r in sl_rows})
    blocks.append((f"Task 11 · Slicer PayerType = {' + '.join(T11_PAYER_TYPES)} · Timeline Q{T11_QUARTER} {YEAR} · "
                   "Count of EncounterID",
                   ["Row Labels"] + sl_types + ["Grand Total"],
                   [[f] + [sl_c[(f, t)] or None for t in sl_types] + [sum(sl_c[(f, t)] for t in sl_types)] for f in sl_facs],
                   ["@"] + ["#,##0"] * (len(sl_types) + 1),
                   ["Grand Total"] + [sum(sl_c[(f, t)] for f in sl_facs) for t in sl_types] + [len(sl_rows)]))
    fac_ip = defaultdict(list)
    for r in ip:
        fac_ip[r["FacilityName"]].append(r)
    blocks.append(("Task 12 · EncounterType = Inpatient · calculated field ChargesPerDay = TotalCharges/LOSDays",
                   ["Row Labels", "Sum of TotalCharges", "Sum of LOSDays", "Sum of ChargesPerDay"],
                   [[f, sum(x["TotalCharges"] for x in fac_ip[f]), sum(x["LOSDays"] for x in fac_ip[f]), cpd(fac_ip[f])]
                    for f in sorted(fac_ip)], ["@", "#,##0.00", "#,##0", "#,##0.00"],
                   ["Grand Total", sum(r["TotalCharges"] for r in ip), sum(r["LOSDays"] for r in ip), cpd(ip)]))
    ft = defaultdict(float)
    for r in rows:
        ft[(r["FacilityName"], r["EncounterType"])] += r["TotalCharges"]
    blocks.append(("Task 13 · Sum of TotalCharges (FacilityName × EncounterType); blank = no encounters",
                   ["Row Labels"] + types + ["Grand Total"],
                   [[f] + [ft.get((f, t)) for t in types] + [sum(v for (ff, _), v in ft.items() if ff == f)]
                    for f in facilities],
                   ["@"] + ["#,##0.00"] * (len(types) + 1),
                   ["Grand Total"] + [sum(v for (_, tt), v in ft.items() if tt == t) for t in types]
                   + [sum(r["TotalCharges"] for r in rows)]))
    blocks.append((f"Bonus B1–B3 · EncounterType = Inpatient · Tabular Form · value filter Index stays ≥ {B_MIN_STAYS} "
                   "(sorted by Readmit rate within each service line)",
                   ["ServiceLine", "PayerType", "Index stays", "Readmit rate"],
                   [[k[0], k[1], s[0], s[2]] for k, s in sorted(kept.items(), key=lambda kv: (kv[0][0], -kv[1][2]))],
                   ["@", "@", "#,##0", "0.0%"],
                   ["Grand Total", None, sum(s[0] for s in kept.values()),
                    sum(s[1] for s in kept.values()) / sum(s[0] for s in kept.values())]))
    blocks.append((f"Bonus B4–B5 · EncounterType = Inpatient · Sum of ReadmitFlag by DxDescription · Top {B_TOP_N} filter",
                   ["Row Labels", "Sum of ReadmitFlag", "Share of all readmissions"],
                   [[d, c, c / total_readmits] for d, c in dx_ranked[:B_TOP_N]], ["@", "#,##0", "0.0%"],
                   ["Grand Total (visible rows only)", top_sum, b5]))

    @L.customize
    def _custom(wb, lesson, selftest):
        ws = wb.create_sheet("Pivot Key")
        ws.sheet_properties.tabColor = "C00000"
        ws["A1"] = "🔑 Pivot Key — Lesson 3.4"
        ws["A1"].font = Font(bold=True, size=16, color="7B2C2C")
        ws["A2"] = ("Each block is a finished PivotTable from the practice, typed out as values so you can compare your whole "
                    "layout with it.")
        ws["A3"] = ("Your pivot may list rows in a different order, use other number formats, and name a second copy of a "
                    "value field differently (Excel calls it Count of EncounterID2 until you rename it).")
        for note in ("A2", "A3"):
            ws[note].font = Font(italic=True, color="595959")
        r = 5
        title_font = Font(bold=True, size=12, color=NAVY)
        hdr_font = Font(bold=True, color="FFFFFF")
        total_font = Font(bold=True)
        total_border = Border(top=Side(style="thin", color="1F4E79"))
        light = PatternFill("solid", fgColor="DDEBF7")
        for title, headers, body, fmts, total in blocks:
            ws.cell(row=r, column=1, value=title).font = title_font
            r += 1
            for j, h in enumerate(headers, 1):
                c = ws.cell(row=r, column=j, value=h)
                c.font = hdr_font
                c.fill = HEADER_FILL
                c.border = BOX
                c.alignment = Alignment(horizontal="left" if j == 1 else "right", wrap_text=True)
            r += 1
            for row_vals in body:
                for j, v in enumerate(row_vals, 1):
                    c = ws.cell(row=r, column=j, value=v)
                    if fmts[j - 1] != "@":
                        c.number_format = fmts[j - 1]
                r += 1
            if total:
                for j, v in enumerate(total, 1):
                    c = ws.cell(row=r, column=j, value=v)
                    c.font = total_font
                    c.fill = light
                    c.border = total_border
                    if fmts[j - 1] != "@" and isinstance(v, (int, float)):
                        c.number_format = fmts[j - 1]
                r += 1
            r += 1
        ws.column_dimensions["A"].width = 52
        for letter in "BCDEF":
            ws.column_dimensions[letter].width = 19
        ws.sheet_state = "hidden"

        # Data Dictionary: wrap the long descriptions; both reference sheets print one page wide.
        dd = wb["Data Dictionary"]
        for row in dd.iter_rows(min_row=2, max_row=dd.max_row):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for sheet in (ws, dd):
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True

        # Start Here lists every data sheet as "Data: N rows". Say what the dictionary holds instead.
        start = wb["Start Here"]
        hdr = next(r for r in range(1, start.max_row + 1) if start.cell(row=r, column=2).value == "Sheets in this workbook")
        dd_row = next(r for r in range(hdr + 1, start.max_row + 1) if start.cell(row=r, column=2).value == "Data Dictionary")
        start.cell(row=dd_row, column=3).value = (f"What each of the {len(DICTIONARY)} columns in {TBL} means, with an "
                                                  "example value.")

        # Answer keys: show the Markdown steps and explanations as plain text in Excel.
        for key_name in (lesson.key_sheet, lesson.bonus_key_sheet):
            ks = wb[key_name]
            for row in range(5, ks.max_row + 1):
                for c in (4, 6):  # sample solution and explanation columns
                    v = ks.cell(row=row, column=c).value
                    if isinstance(v, str) and not v.startswith("="):
                        ks.cell(row=row, column=c).value = re.sub(r"\*([^*\n]+)\*", r"\1", _plain(v))

    L.sheet_order = ["Start Here", "Practice", "Encounters", "Data Dictionary", "Bonus", "Answer Key", "Bonus Key",
                     "Pivot Key"]
    return L
