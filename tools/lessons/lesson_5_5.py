"""Lesson 5.5 · Events, UserForms & Automated Reports.

The workbook holds the sheets the learner's VBA works on (BedBoard, Intake, Log, Combined, Settings, Lists).
The macros themselves live in the lesson's starter/ and solutions/ folders (hand-written; the builder reads the
solution files so the README answer key always shows the same code). The 12 monthly census CSVs the Dir task
combines are generated here from data/daily_census.csv and written to data/census_monthly/ (+ a .zip).

Checks read what the macros write: gray summary formulas count Log rows, tblIntake rows and Combined rows, and the
bonus reads the Report_yyyy-mm sheets the learner's macro creates through INDIRECT (a missing sheet → "").
The self-test simulates the macros' output (Log rows, three intake rows, the combined CSV rows, the report sheets).
"""
from __future__ import annotations

import calendar
import csv
import io
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from xlcourse import Lesson, Task, data

CODE = "5.5"
LESSON_DIR = Path(__file__).resolve().parents[2] / "05-automation-vba" / "05-events-userforms-automation"
NAVY = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
NOTE_FONT = Font(italic=True, color="595959")

CSV_COLS = ["CensusDate", "FacilityID", "DeptID", "UnitName", "StaffedBeds", "Admissions", "Discharges",
            "MidnightCensus"]


# ---------------------------------------------------------------------------- helpers
def _read(rel: str) -> str:
    return (LESSON_DIR / rel).read_text(encoding="utf-8")


def _body(rel: str) -> str:
    """A paste-in file (.cls/.vba) without its header comment block (everything up to the dashed line)."""
    lines = _read(rel).splitlines()
    cut = next((i for i, l in enumerate(lines) if l.startswith("'-----")), -1)
    return "\n".join(lines[cut + 1:]).strip("\n")


def _proc(rel: str, name: str) -> str:
    """One Sub/Function from a VBA file, including the comment lines directly above it."""
    lines = _read(rel).splitlines()
    pat = re.compile(rf"^\s*(Public |Private )?(Sub|Function) {re.escape(name)}\b")
    start = next(i for i, l in enumerate(lines) if pat.match(l))
    end = next(i for i in range(start, len(lines)) if re.match(r"^\s*End (Sub|Function)\b", lines[i]))
    while start > 0 and lines[start - 1].lstrip().startswith("'") and not lines[start - 1].lstrip().startswith("'---") \
            and not lines[start - 1].lstrip().startswith("'==="):
        start -= 1
    return "\n".join(lines[start:end + 1])


def _excel_like_hhnn(dt: datetime) -> str:
    """VBA Format(dt, "yyyy-mm-dd_hhnn"): 4-digit year, 2-digit month/day, 24-hour hour, nn = minutes."""
    return f"{dt.year:04d}-{dt.month:02d}-{dt.day:02d}_{dt.hour:02d}{dt.minute:02d}"


def _csv_text(rows: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(CSV_COLS)
    for r in rows:
        w.writerow([r["CensusDate"].isoformat()] + [r[c] for c in CSV_COLS[1:]])
    return buf.getvalue()


def _zip_bytes(files: dict[str, str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name in sorted(files):
            info = zipfile.ZipInfo(name, date_time=(2025, 12, 31, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            z.writestr(info, files[name].encode("utf-8"))
    return buf.getvalue()


def _unit_summary(rows: list[dict], month: date) -> tuple[list[tuple], tuple]:
    """Report rows for one month, sorted by occupancy desc: (DeptID, Unit, Facility, PD, BD, ADC, Occ) + total row."""
    ndays = calendar.monthrange(month.year, month.month)[1]
    pd, bd, meta = defaultdict(int), defaultdict(int), {}
    for r in rows:
        if r["CensusDate"].year == month.year and r["CensusDate"].month == month.month:
            pd[r["DeptID"]] += r["MidnightCensus"]
            bd[r["DeptID"]] += r["StaffedBeds"]
            meta[r["DeptID"]] = (r["UnitName"], r["FacilityID"])
    out = [(d, meta[d][0], meta[d][1], pd[d], bd[d], pd[d] / ndays, pd[d] / bd[d]) for d in pd]
    out.sort(key=lambda t: -t[6])
    tp, tb = sum(pd.values()), sum(bd.values())
    total = ("Total", "All units", "", tp, tb, tp / ndays, tp / tb)
    return out, total


# ---------------------------------------------------------------------------- build
def build() -> Lesson:
    deps = data.index(data.load("departments"), "DeptID")
    pats = data.index(data.load("patients"), "PatientID")
    facs = data.index(data.load("facilities"), "FacilityID")

    # ------------------------------------------------------------------ census CSVs (Dir task + bonus)
    census = [r for r in data.load("daily_census") if r["CensusDate"].year == 2025]
    census.sort(key=lambda r: (r["CensusDate"], r["FacilityID"], r["DeptID"]))
    n_units = len({r["DeptID"] for r in census})
    n_beds_4w = 36                       # 4 West (D110) bed board below: rooms 401-418, beds A/B
    assert {r["StaffedBeds"] for r in census if r["DeptID"] == "D110" and r["CensusDate"] == date(2025, 12, 31)} \
        == {n_beds_4w}

    L = Lesson(
        code=CODE, module_dir="05-automation-vba", slug="05-events-userforms-automation",
        title="Events, UserForms & Automated Reports", level="Expert", minutes=75,
        objectives=[
            "Respond to workbook and worksheet events (Open, Change, BeforeSave)",
            "Build a validated data-entry UserForm that writes to a Table",
            "Automate a report: combine files with Dir, export to PDF, save timestamped copies",
            "Know the safe limits of automation and when to use Office Scripts or Power Automate",
        ],
        data_note=f"Bluestone Memorial Hospital: the 4 West bed board ({n_beds_4w} beds, 12/31/2025) and the ED intake "
                  f"log (Dec 25–31, 2025), plus 12 monthly census CSV files covering all {n_units} inpatient units in the "
                  f"health system in 2025 ({len(census):,} rows).",
    )
    for r in census:
        r["UnitName"] = deps[r["DeptID"]]["DeptName"]
    readme_txt = ("Bluestone Health System - bed management export\n"
                  "Midnight census by inpatient unit, 2025. One CSV file per month (census_yyyy-mm.csv).\n"
                  "Columns: " + ", ".join(CSV_COLS) + "\n"
                  "Dates are ISO yyyy-mm-dd. Synthetic data for the Learn Excel course.\n")
    folder_files = {"census_monthly/README.txt": readme_txt}
    for m in range(1, 13):
        month_rows = [r for r in census if r["CensusDate"].month == m]
        folder_files[f"census_monthly/census_2025-{m:02d}.csv"] = _csv_text(month_rows)
    for name, text in folder_files.items():
        L.extra_files[f"data/{name}"] = text
    L.extra_files["data/census_monthly.zip"] = _zip_bytes(folder_files)

    # ------------------------------------------------------------------ BedBoard (4 West, 12/31/2025 14:00)
    bed_ids = [f"4W-{room}{ab}" for room in range(401, 419) for ab in "AB"]
    assert len(bed_ids) == n_beds_4w
    status0 = {b: "Occupied" for b in bed_ids}
    status0.update({"4W-410A": "Dirty", "4W-410B": "Dirty", "4W-411A": "Dirty",
                    "4W-417B": "Clean", "4W-418B": "Clean", "4W-418A": "Blocked"})
    notes0 = {"4W-402A": "Fall risk", "4W-405B": "Discharge order written; waiting for ride",
              "4W-409B": "Telemetry", "4W-414A": "NPO after midnight",
              "4W-418A": "Contact isolation; room needs terminal clean"}
    beds = []
    for i, b in enumerate(bed_ids):
        st = status0[b]
        if st == "Occupied":
            t = datetime(2025, 12, 26, 8, 0) + timedelta(minutes=(i * 397) % (5 * 24 * 60))
        elif st == "Dirty":
            t = datetime(2025, 12, 31, 10, 5) + timedelta(minutes=17 * (i % 5))
        elif st == "Clean":
            t = datetime(2025, 12, 31, 9, 20) + timedelta(minutes=11 * (i % 4))
        else:
            t = datetime(2025, 12, 30, 18, 0)
        beds.append({"BedID": b, "Room": int(b[3:6]), "Status": st, "StatusTime": t, "Notes": notes0.get(b)})
    bed_sd = L.add_table_sheet("BedBoard", beds, table="tblBeds",
                               columns=["BedID", "Room", "Status", "StatusTime", "Notes"],
                               formats={"StatusTime": "mm/dd/yyyy hh:mm"},
                               widths={"BedID": 11, "Room": 8, "Status": 12, "StatusTime": 18, "Notes": 44})

    # ------------------------------------------------------------------ Intake (ED, Bluestone Memorial)
    cutoff = datetime(2025, 12, 31, 11, 40)
    ed = [r for r in data.load("ed_visits")
          if r["FacilityID"] == "F01" and r["ArrivalDateTime"] >= datetime(2025, 12, 25)]
    ed.sort(key=lambda r: r["ArrivalDateTime"])
    before = [r for r in ed if r["ArrivalDateTime"] < cutoff]
    new3 = [r for r in ed if r["ArrivalDateTime"] >= cutoff][:3]

    def intake_row(i, r):
        p = pats[r["PatientID"]]
        t = r["ArrivalDateTime"]
        return {"IntakeID": 1001 + i, "ArrivalDateTime": t, "MRN": p["MRN"], "LastName": p["LastName"],
                "FirstName": p["FirstName"], "ArrivalMode": r["ArrivalMode"], "ESILevel": r["ESILevel"],
                "ChiefComplaint": r["ChiefComplaint"], "Interpreter": p["PreferredLanguage"] != "English",
                "EnteredBy": "D. Okafor" if 7 <= t.hour < 19 else "M. Reyes"}
    intake = [intake_row(i, r) for i, r in enumerate(before)]
    added = [intake_row(len(before) + i, r) for i, r in enumerate(new3)]
    for a in added:
        a["EnteredBy"] = "(you)"
    n0 = len(intake)
    in_sd = L.add_table_sheet(
        "Intake", intake, table="tblIntake",
        columns=["IntakeID", "ArrivalDateTime", "MRN", "LastName", "FirstName", "ArrivalMode", "ESILevel",
                 "ChiefComplaint", "Interpreter", "EnteredBy"],
        formats={"ArrivalDateTime": "mm/dd/yyyy hh:mm"},
        widths={"ArrivalDateTime": 17, "MRN": 11, "ChiefComplaint": 30, "Interpreter": 12, "EnteredBy": 12})

    arrival_modes = ["Walk-In", "Ambulance", "Police", "Air Transport"]
    complaints = sorted({r["ChiefComplaint"] for r in data.load("ed_visits")})
    for a in added:
        assert a["ArrivalMode"] in arrival_modes and a["ChiefComplaint"] in complaints

    # ------------------------------------------------------------------ answers (computed in Python)
    # 1 · Intersect(Target, Status column) for a selection that overhangs the table
    st_col = bed_sd.col("Status")
    st_first, st_last = bed_sd.first_row, bed_sd.last_row
    sel_c1, sel_c2, sel_r1, sel_r2 = "A", "E", 30, 40
    overlap_rows = max(0, min(sel_r2, st_last) - max(sel_r1, st_first) + 1)
    overlap_cols = 1 if sel_c1 <= st_col <= sel_c2 else 0
    intersect_count = overlap_rows * overlap_cols
    status_addr = f"{st_col}{st_first}:{st_col}{st_last}"
    sel_addr = f"{sel_c1}{sel_r1}:{sel_c2}{sel_r2}"

    # 3 · Format(Now, "yyyy-mm-dd_hhnn") at 17:07 on 12/31/2025
    run_at = datetime(2025, 12, 31, 17, 7)
    archive_name = "CensusReport_" + _excel_like_hhnn(run_at) + ".xlsm"

    # 7 · bed-board edits: each changed Status cell → one "Bed status" log row; Notes edits → none
    edits = [(["4W-405B"], "Status", "Dirty"),
             (["4W-410A", "4W-410B", "4W-411A"], "Status", "Clean"),   # one Ctrl+Enter action, Target = 3 cells
             (["4W-415A"], "Notes", "Bed alarm on")]
    for cells, col, _ in edits:
        if col == "Status" and len(cells) > 1:     # Ctrl+Enter needs adjacent cells
            rows_ = [bed_ids.index(c) for c in cells]
            assert rows_ == list(range(rows_[0], rows_[0] + len(rows_)))
    bed_log_rows = sum(len(cells) for cells, col, _ in edits if col == "Status")

    # 8–10 · intake form
    final_intake = intake + added
    rows_after = len(final_intake)
    text_mrns = sum(1 for r in final_intake if isinstance(r["MRN"], str) and len(r["MRN"]) == 8)
    esi12 = sum(1 for r in final_intake if r["ESILevel"] <= 2)
    esi12_new = [a for a in added if a["ESILevel"] <= 2]
    bad_mrn = added[1]["MRN"].lstrip("0")          # the 'dropped leading zeros' MRN the form must reject
    assert len(bad_mrn) != 8

    # 11–12 · combine + analyze
    combined_rows = len(census)
    d130 = [r for r in census if r["DeptID"] == "D130"]
    d130_occ = sum(r["MidnightCensus"] for r in d130) / sum(r["StaffedBeds"] for r in d130)

    # bonus
    dec_units, dec_total = _unit_summary(census, date(2025, 12, 1))
    jan_units, jan_total = _unit_summary(census, date(2025, 1, 1))
    dec_d130_adc = next(u[5] for u in dec_units if u[0] == "D130")
    assert dec_units[0][6] > dec_units[1][6] and jan_units[0][6] > jan_units[1][6]   # no tie at the top
    # B3 explanation: the smallest unit vs 4 West (D110), beds = December bed days ÷ 31
    dec_days = calendar.monthrange(2025, 12)[1]
    small = min(dec_units, key=lambda u: u[4])
    small_beds, w4_beds = small[4] / dec_days, next(u[4] for u in dec_units if u[0] == "D110") / dec_days
    assert small_beds == int(small_beds) and w4_beds == n_beds_4w
    small_desc = (f"the {int(small_beds)}-bed {small[1]} at "
                  f"{facs[small[2]]['FacilityName'].replace(' Community Hospital', '').replace(' Medical Center', '')}")

    # ------------------------------------------------------------------ tasks
    log_rng = "Log!B2:B1000"
    comb_a = "Combined!A2:A20000"

    def arrival_line(a):
        return (f"• {a['ArrivalDateTime']:%H:%M} · MRN {a['MRN']} · {a['LastName']}, {a['FirstName']} · "
                f"{a['ArrivalMode']} · ESI {a['ESILevel']} · {a['ChiefComplaint']} · "
                f"{'interpreter needed' if a['Interpreter'] else 'no interpreter'}")

    L.practice_intro = (
        "Tasks 1–5 are quick concept checks you can answer from the guide. Tasks 6–12 are build tasks: save this "
        "workbook as an .xlsm first, import the starter modules, and work in order. The gray cells read what your "
        "macros write, so they stay blank until your code has run.")
    L.start_notes = [
        "Save this file as an Excel Macro-Enabled Workbook (.xlsm) before you write any code: File → Save As → "
        "Excel Macro-Enabled Workbook.",
        "Put the census_monthly folder (from the lesson's data folder or census_monthly.zip) in the same folder as "
        "your .xlsm. CombineCensusFiles looks for it there.",
        "Starter code: the lesson's starter folder. Import the .bas files with File → Import File… in the VBE. Paste "
        "the .cls and .vba files into the module named at the top of each file.",
        "Other sheets: Log (WriteLog appends a row here), Combined (CombineCensusFiles fills it from the CSV files), "
        "Settings (cells the report macros read and write), Lists (drop-down lists for the intake form).",
    ]

    L.tasks = [
        Task(f"On the BedBoard sheet, the Status column of tblBeds is {status_addr}. Suppose you select {sel_addr} "
             f"and press Delete. Worksheet_Change runs once with Target = {sel_addr}. How many cells are in "
             f"Intersect(Target, Range(\"{status_addr}\"))?",
             answer=intersect_count, title="Cells in Intersect(Target, Status column)",
             solution=f"Draw both rectangles. Target covers columns {sel_c1}–{sel_c2}, rows {sel_r1}–{sel_r2}. "
                      f"The Status column covers column {st_col}, rows {st_first}–{st_last}. They share "
                      f"{st_col}{max(sel_r1, st_first)}:{st_col}{min(sel_r2, st_last)}.",
             hint="Intersect keeps only the cells that are in both ranges: shared columns AND shared rows",
             explanation=f"Rows {st_last + 1}–{sel_r2} are below the table and columns other than {st_col} aren't "
                         f"Status cells, so they drop out, leaving {intersect_count} cells. This is why a handler loops "
                         "over Intersect(Target, …) instead of over Target: it touches only the Status cells that "
                         "changed, however big the selection was."),
        Task("A run-time error stopped your Worksheet_Change handler right after it set Application.EnableEvents = "
             "False, and now nothing happens when you edit a Status cell. Type the exact statement you would run "
             "in the Immediate window (Ctrl + G) to switch events back on.",
             answer="Application.EnableEvents = True",
             accept=["Application.EnableEvents=True", "Application.EnableEvents = True"],
             title="Turn events back on from the Immediate window",
             solution="`Application.EnableEvents = True` (type it in the Immediate window and press Enter)",
             hint="It's the same property your handler switched off",
             explanation="EnableEvents belongs to the whole Excel application, not to one workbook or sheet. It stays "
                         "False after the macro stops, for every open workbook, until something sets it back. That's "
                         "why every handler that turns it off needs an error handler that always turns it back on. "
                         "The EventsOn macro in modLog does the same job from Alt + F8 (Mac: Option + F8, or Developer → Macros)."),
        Task("Your archive macro runs at 5:07 PM on 12/31/2025 and calls ThisWorkbook.SaveCopyAs folder & "
             "\"CensusReport_\" & Format(Now, \"yyyy-mm-dd_hhnn\") & \".xlsm\". What file name does it create? "
             "Type the name only, without the folder.",
             answer=archive_name, title="Timestamped file name from Format",
             solution=f"`{archive_name}`",
             hint="Look up each code in the Format table in Guide §9 (hh is a 24-hour clock unless you add AM/PM)",
             explanation="yyyy-mm-dd gives 2025-12-31. hh gives the hour on a 24-hour clock (17) because the format "
                         "has no AM/PM code, and nn gives the minutes (07). Year-month-day order makes the archive "
                         "files sort by date when you sort them by name. VBA would also read mm right after hh as "
                         "minutes, but nn is never ambiguous, so prefer it."),
        Task("After ThisWorkbook.SaveCopyAs runs, which workbook are you working in? Type A, B, or C.\n"
             "A = the original file (the copy was written to disk but isn't open)\n"
             "B = the new copy (Excel switched to it, the way Save As does)\n"
             "C = both files are open in Excel",
             answer="A", title="Where you are after SaveCopyAs",
             solution="**A.** SaveCopyAs writes a snapshot to a new file and leaves the open workbook alone.",
             hint="Compare Save, SaveAs, and SaveCopyAs in Guide §9",
             explanation="SaveAs renames the open workbook, so afterwards you're editing the new file and the "
                         "original file is no longer open. SaveCopyAs is the right tool for backups because nothing about the open "
                         "workbook changes: same name, same folder, and the copy keeps the original's file format, "
                         "so give it the same extension (.xlsm)."),
        Task("The house supervisor wants the census summary rebuilt at 06:00 every morning and posted to a Microsoft "
             "Teams channel, even on days when nobody has Excel open. Which tool fits? Type A, B, or C.\n"
             "A = Application.OnTime in your .xlsm\n"
             "B = a Workbook_Open macro\n"
             "C = an Office Script run by a scheduled Power Automate flow",
             answer="C", title="Unattended daily automation",
             solution="**C.** A scheduled (Recurrence) Power Automate flow runs an Office Script in the cloud.",
             hint="Which option doesn't need desktop Excel to be running?",
             explanation="OnTime and Workbook_Open both need desktop Excel running on someone's PC with the workbook "
                         "open (or being opened). Neither can fire at 06:00 if that PC is off, nobody is logged on, or "
                         "Excel was closed overnight. A Power "
                         "Automate flow with a Recurrence trigger runs in Microsoft's cloud, calls an Office Script on "
                         "the workbook stored in OneDrive or SharePoint, and can post the result to Teams. VBA can't "
                         "run in that setting at all."),
        Task("Save the workbook as .xlsm and import modLog. Paste starter/ThisWorkbook.cls into the ThisWorkbook "
             "module and finish Workbook_Open and Workbook_BeforeSave so they call WriteLog \"Open\", … and "
             "WriteLog \"Save\", …. Save, close, reopen with macros enabled, then save again. The gray cell shows TRUE when the Log sheet has at least "
             "one Open row and one Save row.",
             answer=True, title="Workbook_Open and Workbook_BeforeSave write to the Log",
             summary=f'=IF(COUNTA({log_rng})=0,"",AND(COUNTIF({log_rng},"Open")>0,COUNTIF({log_rng},"Save")>0))',
             fill={"range": "Log!B2:B3", "values": ["Open", "Save"]}, live=False,
             solution=_body("solutions/ThisWorkbook.cls"), solution_lang="vba",
             hint="Choose Workbook in the left drop-down of the ThisWorkbook code window, then the event on the right",
             explanation="Workbook events run only from the ThisWorkbook module, and their names and arguments must "
                         "match exactly, which is why picking them from the drop-downs beats typing. BeforeSave runs "
                         "before Excel writes the file, so the Save row ends up inside the saved file. Setting "
                         "Cancel = True stops the save. The solution uses that to refuse saving while a required "
                         "Intake field is blank, and it logs the attempt before it exits."),
        Task("Paste starter/BedBoard.cls into the BedBoard sheet's module and finish its Worksheet_Change handler. "
             "For each changed Status cell it should "
             "stamp StatusTime with Now and call WriteLog \"Bed status\", … once. Delete any \"Bed status\" rows "
             "that your testing left on the Log sheet, then make exactly these edits:\n"
             "• Change 4W-405B's Status to Dirty.\n"
             "• Select the Status cells of 4W-410A, 4W-410B and 4W-411A, type Clean, and press Ctrl + Enter "
             "(Mac: ⌘ + Return) to fill all three at once.\n"
             "• In 4W-415A's Notes cell, type Bed alarm on.\n"
             "The gray cell counts the Log's \"Bed status\" rows.",
             answer=bed_log_rows, title="Bed board change log (Worksheet_Change + Intersect)",
             summary=f'=IF(COUNTIF({log_rng},"Bed status")=0,"",COUNTIF({log_rng},"Bed status"))',
             fill={"range": f"Log!B4:B{3 + bed_log_rows}", "values": ["Bed status"] * bed_log_rows}, live=False,
             solution=_body("solutions/BedBoard.cls"), solution_lang="vba",
             hint="Intersect Target with the Status column, then loop For Each over the result",
             explanation="Ctrl + Enter changes three cells in one action, so the event fires once with a 3-cell "
                         "Target. A handler that reads only Target.Value or Target.Row logs 1 row instead of 3. The "
                         "Notes edit fires the event too, but its Intersect with the Status column is Nothing, so the "
                         f"handler exits. Total: 1 + 3 = {bed_log_rows}. EnableEvents = False stops the handler's own "
                         "writes (the tidied status and the StatusTime stamp) from firing it again, and the CleanUp "
                         "label turns events back on even after an error."),
        Task("Build frmIntake (Guide §6–7) and use it to add these three ED arrivals on 12/31/2025. Type each "
             "arrival as yyyy-mm-dd hh:mm, for example 2025-12-31 11:41.\n"
             + "\n".join(arrival_line(a) for a in added) +
             f"\nThen try to save a fourth record with MRN {bad_mrn}. Your form must refuse it. "
             "The gray cell shows how many rows tblIntake has now.",
             answer=rows_after, title="Intake form: rows in tblIntake",
             summary=f'=IF(ROWS(tblIntake[IntakeID])<={n0},"",ROWS(tblIntake[IntakeID]))',
             fill={"range": f"Intake!A{in_sd.first_row + n0}:A{in_sd.first_row + n0 + 2}",
                   "values": [a["IntakeID"] for a in added]}, live=False,
             solution=_proc("solutions/frmIntake.vba", "cmdSave_Click") + "\n\n"
                      + _proc("solutions/frmIntake.vba", "AddIntakeRow") + "\n\n"
                      + _proc("solutions/frmIntake.vba", "SetField"),
             solution_lang="vba",
             hint="ListRows.Add, then write each field into the new row",
             explanation=f"The table started with {n0} arrivals, and three valid saves make {rows_after}. If you see "
                         f"{rows_after + 1}, the {len(bad_mrn)}-digit MRN got through, so check your "
                         "Like \"########\" test (# matches exactly one digit). An extra, half-empty row instead "
                         "means a run-time error stopped AddIntakeRow after ListRows.Add: delete that row "
                         "(right-click → Delete → Table Rows) and fix the error. ListRows.Add grows the Table itself, "
                         "so formats, formulas, and anything that refers to tblIntake pick up the new row "
                         "automatically. The full form code is in solutions/frmIntake.vba."),
        Task("How many MRNs in tblIntake are stored as 8-character text? The gray cell counts them with ISTEXT and "
             "LEN. If the count is lower than your row count in task 8, your form wrote some MRNs as numbers. Delete "
             "those rows, fix AddIntakeRow, and enter the patients again.",
             answer=text_mrns, title="Intake form: MRNs kept as text",
             summary=f'=IF(ROWS(tblIntake[IntakeID])<={n0},"",SUMPRODUCT(--ISTEXT(tblIntake[MRN]),--(LEN(tblIntake[MRN])=8)))',
             fill={"range": f"Intake!C{in_sd.first_row + n0}:C{in_sd.first_row + n0 + 2}",
                   "values": [a["MRN"] for a in added]}, live=False,
             solution="newRow.Range.Cells(1, lo.ListColumns(\"MRN\").Index).NumberFormat = \"@\"   ' Text format first...\n"
                      "SetField lo, newRow, \"MRN\", Trim$(txtMRN.Value)                     ' ...so leading zeros survive",
             solution_lang="vba",
             hint="A TextBox gives you text, but Excel converts number-like text that lands in a General cell",
             explanation="txtMRN.Value is the text \"00787672\". When VBA writes number-like text into a "
                         "General-format cell, Excel converts it just as if you had typed it, so the cell stores "
                         "787672 and the leading zeros are gone. Formatting the cell as Text (\"@\") before writing "
                         "keeps it as text, and so does writing \"'\" & mrn. Don't rely on the column's existing "
                         "format: new Table rows copy whatever format the row above has."),
        Task("How many arrivals in tblIntake are ESI level 1 or 2 now? The gray cell uses "
             "COUNTIFS(tblIntake[ESILevel],\"<=2\"), which counts only real numbers.",
             answer=esi12, title="Intake form: ESI 1–2 arrivals",
             summary=f'=IF(ROWS(tblIntake[IntakeID])<={n0},"",COUNTIFS(tblIntake[ESILevel],"<=2"))',
             fill={"range": f"Intake!G{in_sd.first_row + n0}:G{in_sd.first_row + n0 + 2}",
                   "values": [a["ESILevel"] for a in added]}, live=False,
             solution=_proc("solutions/frmIntake.vba", "SelectedESI"), solution_lang="vba",
             hint="Loop over the five option buttons with Me.Controls(\"optESI\" & i)",
             explanation="Option buttons don't share a single value, so SelectedESI asks optESI1 to optESI5 in turn "
                         "and returns the number of the one that's selected (0 if none, which validation rejects). "
                         "Writing that Long keeps ESILevel numeric, so COUNTIFS, AVERAGE, and PivotTables all work. "
                         f"{esi12 - len(esi12_new)} of the original {n0} arrivals were ESI 1–2, and "
                         f"{esi12_new[0]['FirstName']} {esi12_new[0]['LastName']} (ESI {esi12_new[0]['ESILevel']}) "
                         f"brings the total to {esi12}."),
        Task("Finish CombineCensusFiles in modReports so it loops through every .csv file in the census_monthly "
             "folder with Dir, opens each one, copies its data rows (not its header) to the next empty row of the "
             "Combined sheet, and closes it without saving. Run it. The gray cell counts the data rows on Combined, or "
             "shows a message if column A holds text or blank rows.",
             answer=combined_rows, title="CombineCensusFiles: rows on Combined",
             summary=(f'=IF(COUNTA({comb_a})=0,"",IF(COUNT({comb_a})<COUNTA({comb_a}),"Text rows found",'
                      f'IF(COUNTBLANK(Combined!A2:INDEX({comb_a},COUNTA({comb_a})))>0,"Blank rows found",'
                      f'COUNTA({comb_a}))))'),
             fill={"range": f"Combined!A2:H{1 + combined_rows}",
                   "values": [r[c] for r in census for c in CSV_COLS]}, live=False,
             solution=_proc("solutions/modReports.bas", "CombineCensusFiles"), solution_lang="vba",
             hint="Dir(folder) gives the first file and Dir() each next one. CurrentRegion.Offset(1).Resize(…) drops the header",
             explanation=f"12 files × {n_units} units × the days in each month = {combined_rows:,} rows. The gray cell "
                         "also inspects column A (CensusDate), where every cell should be a real date. \"Text rows "
                         "found\" means header rows or README.txt lines were copied in: drop each file's header with "
                         "Offset(1, 0).Resize(rows − 1), and skip other files with LCase$(fileName) Like \"*.csv\". "
                         "\"Blank rows found\" means each file left an empty row behind. That happens with Offset(1, 0) "
                         "alone, because the shifted block keeps the header's row in its count and so ends one row "
                         "below the data. Because the macro clears row 2 down first, running it twice gives the same "
                         "count instead of doubling it."),
        Task("Now analyze what your macro imported. In the yellow cell, write one formula that reads the Combined "
             "sheet: what was the 2025 occupancy of Bluestone Memorial's Intensive Care Unit (DeptID D130)? Use patient days (sum of "
             "MidnightCensus) ÷ bed days (sum of StaffedBeds), as a percentage to 1 decimal place.",
             answer=d130_occ, fmt="0.0%", live=False, title="D130 occupancy from the combined data",
             solution='=SUMIFS(Combined!H:H,Combined!C:C,"D130")/SUMIFS(Combined!E:E,Combined!C:C,"D130")',
             hint="SUMIFS ÷ SUMIFS on the Combined columns",
             explanation="Occupancy for a period is total patient days ÷ total bed days, the same total-over-total "
                         "rule as in Lesson 1.4. Whole-column references such as Combined!H:H suit a sheet that a "
                         "macro refills, because the formula keeps working however many rows arrive next month."),
    ]

    # ------------------------------------------------------------------ bonus
    rep_dec, rep_jan = "Report_2025-12", "Report_2025-01"
    rep_max = 60

    def rep(sheet, rng):
        return f'INDIRECT("\'{sheet}\'!{rng}")'

    L.bonus_title = "Bonus: one-click monthly census report"
    L.bonus_scenario = (
        "The Chief Nursing Officer wants a monthly census packet without anyone copying and pasting. Write "
        "BuildCensusReport in modReports. It reads the month from Settings!B3 (ReportMonth) and then:\n"
        "1. Runs CombineCensusFiles.\n"
        "2. Deletes and re-creates a sheet named Report_yyyy-mm (for example Report_2025-12).\n"
        "3. Puts the headers DeptID, Unit, Facility, PatientDays, BedDays, ADC, Occupancy in A3:G3 and one row per "
        "unit from row 4, sorted by Occupancy (highest first). PatientDays = sum of MidnightCensus, BedDays = sum of "
        "StaffedBeds, ADC = PatientDays ÷ days in the month, and Occupancy = PatientDays ÷ BedDays. Store full "
        "precision, and format ADC as 0.0 and Occupancy as 0.0%.\n"
        "4. Adds a row under the units with Total in column A and the system-wide values in D:G.\n"
        "5. Exports the report sheet to CensusReport_yyyy-mm.pdf in the Reports folder next to the workbook (the "
        "folder name is in Settings!B5), writes the PDF's full path into Settings!B6, and saves a timestamped "
        "backup with SaveCopyAs.\n\n"
        "Run it for December 2025 (the month already in Settings!B3), then set Settings!B3 to 01/01/2025 and run it "
        "again. The gray cells read your report sheets.")
    L.bonus = [
        Task("What was the December 2025 average daily census (ADC) of Bluestone Memorial's Intensive Care Unit "
             "(D130)? The gray cell looks it up on Report_2025-12.",
             answer=dec_d130_adc, fmt="0.0", tol=0.051, title="December ADC for D130",
             summary=f'=IFERROR(INDEX({rep(rep_dec, f"F4:F{rep_max}")},MATCH("D130",{rep(rep_dec, f"A4:A{rep_max}")},0)),"")',
             fill={"range": f"'{rep_dec}'!F4:F{3 + len(dec_units)}", "values": [u[5] for u in dec_units]},
             live=False,
             solution=_proc("solutions/modReports.bas", "BuildCensusReport") + "\n\n"
                      + _proc("solutions/modReports.bas", "UnitIndex"),
             solution_lang="vba", hint="ADC = patient days ÷ 31 for December",
             explanation="The macro reads Combined into an array once, then loops in memory. That's far faster than "
                         "reading cells one by one. A Collection maps each DeptID to its position in the parallel "
                         "arrays (UnitIndex returns 0 for a new unit), which works on Windows and Mac. "
                         "Scripting.Dictionary (Lesson 5.4) would also work, but it's Windows-only. "
                         "DateSerial(y, m + 1, 0) returns the last day of the month, so the same code handles 28-, "
                         "30-, and 31-day months."),
        Task("Which unit (DeptID) had the highest occupancy in December 2025? The gray cell reads the first unit "
             "row of Report_2025-12, so your sort must be right.",
             answer=dec_units[0][0], title="December's fullest unit",
             summary=f'=IFERROR({rep(rep_dec, "A4")}&"","")',
             fill={"range": f"'{rep_dec}'!A4:A4", "values": [dec_units[0][0]]}, live=False,
             solution=".Range(\"A\" & FIRST_ROW).Resize(n, 7).Sort Key1:=.Range(\"G\" & FIRST_ROW), _\n"
                      "    Order1:=xlDescending, Header:=xlNo",
             solution_lang="vba", hint="Range.Sort on the Occupancy column, descending, before you add the Total row",
             explanation=f"{dec_units[0][0]} ({dec_units[0][1]}, {dec_units[0][2]}) ran at "
                         f"{dec_units[0][6]:.1%}, just ahead of {dec_units[1][0]} at {dec_units[1][6]:.1%}. Sort "
                         "before you write the Total row. Otherwise the Total row gets sorted in with the units."),
        Task("What was the system-wide inpatient occupancy in December 2025? The gray cell reads it from the Total "
             "row of Report_2025-12.",
             answer=dec_total[6], fmt="0.0%", title="December system occupancy (Total row)",
             summary=f'=IFERROR(INDEX({rep(rep_dec, f"G4:G{rep_max}")},MATCH("Total",{rep(rep_dec, f"A4:A{rep_max}")},0)),"")',
             fill={"range": f"'{rep_dec}'!G{4 + len(dec_units)}:G{4 + len(dec_units)}", "values": [dec_total[6]]},
             live=False,
             solution="totPD / totBD   ' total patient days / total bed days across all units",
             solution_lang="vba", hint="Total patient days ÷ total bed days, not the average of the unit percentages",
             explanation=f"{dec_total[3]:,} patient days ÷ {dec_total[4]:,} bed days = {dec_total[6]:.1%}. "
                         f"Averaging the {len(dec_units)} unit percentages would give every unit equal weight, so "
                         f"{small_desc} would count as much as the {n_beds_4w}-bed 4 West."),
        Task("After rerunning for January 2025, which DeptID tops Report_2025-01?",
             answer=jan_units[0][0], title="January's fullest unit (parameterized rerun)",
             summary=f'=IFERROR({rep(rep_jan, "A4")}&"","")',
             fill={"range": f"'{rep_jan}'!A4:A4", "values": [jan_units[0][0]]}, live=False,
             solution="reportMonth = DateSerial(Year(wsSet.Range(\"B3\").Value), Month(wsSet.Range(\"B3\").Value), 1)\n"
                      "sheetName = \"Report_\" & Format(reportMonth, \"yyyy-mm\")",
             solution_lang="vba", hint="Change Settings!B3, run the same macro again",
             explanation=f"January's leader is {jan_units[0][0]} ({jan_units[0][1]}) at {jan_units[0][6]:.1%}. "
                         "Reading the month from a Settings cell instead of hard-coding it is what turns a one-off "
                         "macro into a reusable report: next month, someone changes one cell and clicks once."),
        Task("What is the file name (without the folder) of the last PDF your macro exported? The gray cell pulls "
             "it from the full path in Settings!B6.",
             answer="CensusReport_2025-01.pdf", accept=["CensusReport_2025-12.pdf"],
             title="PDF file name recorded in Settings!B6",
             summary='=IF(Settings!B6="","",TRIM(RIGHT(SUBSTITUTE(SUBSTITUTE(Settings!B6,"\\","/"),"/",REPT(" ",200)),200)))',
             fill={"range": "Settings!B6:B6",
                   "values": ["C:\\Users\\you\\Documents\\Learn-Excel\\Reports\\CensusReport_2025-01.pdf"]},
             live=False,
             solution="pdfPath = folder & sep & \"CensusReport_\" & Format(reportMonth, \"yyyy-mm\") & \".pdf\"\n"
                      "wsRep.ExportAsFixedFormat Type:=xlTypePDF, Filename:=pdfPath, Quality:=xlQualityStandard, _\n"
                      "    IncludeDocProperties:=True, IgnorePrintAreas:=False, OpenAfterPublish:=False\n"
                      "wsSet.Range(\"B6\").Value = pdfPath",
             solution_lang="vba", hint="Format(reportMonth, \"yyyy-mm\") inside the file name",
             explanation="After the January rerun the answer is CensusReport_2025-01.pdf (the check also accepts the "
                         "December name if you ran December last). Recording the output path in a cell leaves an "
                         "audit trail and lets the next step, such as an email or a Power Automate flow, find the "
                         "file. The gray formula turns every \\ into /, then keeps the text after the last /, so it "
                         "works with Windows and Mac paths."),
    ]

    # ------------------------------------------------------------------ extra sheets & self-test simulation
    @L.customize
    def extras(wb, lesson, selftest):
        bold_white = Font(bold=True, color="FFFFFF")

        def header(ws, cells):
            for c, text in cells:
                ws[c] = text
                ws[c].font = bold_white
                ws[c].fill = HEADER_FILL

        # Log: plain range that WriteLog appends to
        log = wb.create_sheet("Log")
        header(log, [("A1", "Timestamp"), ("B1", "Event"), ("C1", "Detail"), ("D1", "User")])
        for col, w in zip("ABCD", (20, 14, 52, 20)):
            log.column_dimensions[col].width = w
        log.freeze_panes = "A2"
        log.sheet_properties.tabColor = "7F7F7F"

        # Combined: headers only; CombineCensusFiles fills rows 2+
        comb = wb.create_sheet("Combined")
        header(comb, [(f"{chr(65 + i)}1", h) for i, h in enumerate(CSV_COLS)])
        for col, w in zip("ABCDEFGH", (12, 11, 9, 24, 12, 11, 11, 15)):
            comb.column_dimensions[col].width = w
        comb.freeze_panes = "A2"

        # Settings: values the report macros read
        st = wb.create_sheet("Settings")
        st["A1"] = "Report settings"
        st["A1"].font = Font(bold=True, size=13, color=NAVY)
        st["A2"] = "Macros read these cells. Change ReportMonth and run BuildCensusReport again to report on another month."
        st["A2"].font = NOTE_FONT
        rows = [("ReportMonth", date(2025, 12, 1), "First day of the month to report on (bonus)"),
                ("CsvFolder", "census_monthly", "Folder of monthly CSV files, next to this workbook"),
                ("ReportFolder", "Reports", "Folder for PDFs and backup copies, next to this workbook (created if missing)"),
                ("LastExport", None, "Full path of the last PDF exported (written by BuildCensusReport)")]
        for i, (k, v, note) in enumerate(rows, 3):
            st.cell(row=i, column=1, value=k).font = Font(bold=True)
            c = st.cell(row=i, column=2, value=v)
            if k != "LastExport":                       # LastExport is written by the macro, not typed
                c.fill = PatternFill("solid", fgColor="FFF2CC")
            st.cell(row=i, column=3, value=note).font = NOTE_FONT
        st["B3"].number_format = "mmmm yyyy"
        st["B3"].alignment = Alignment(horizontal="left")
        for col, w in zip("ABC", (16, 46, 70)):
            st.column_dimensions[col].width = w

        # Lists: sources for the form's combo boxes
        ls = wb.create_sheet("Lists")
        header(ls, [("A1", "ArrivalMode"), ("C1", "ChiefComplaint")])
        for i, v in enumerate(arrival_modes, 2):
            ls.cell(row=i, column=1, value=v)
        for i, v in enumerate(complaints, 2):
            ls.cell(row=i, column=3, value=v)
        ls.column_dimensions["A"].width = 16
        ls.column_dimensions["C"].width = 32

        # BedBoard: Status dropdown
        bb = wb["BedBoard"]
        dv = DataValidation(type="list", formula1='"Occupied,Clean,Dirty,Blocked"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Bed status",
                            error="Choose Occupied, Clean, Dirty, or Blocked.")
        bb.add_data_validation(dv)
        dv.add(f"{st_col}{st_first}:{st_col}{st_last}")

        sh = wb["Start Here"]
        for row in sh.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("Go to the 'Practice' sheet."):
                    cell.value = ("Go to the 'Practice' sheet. Type answers in the yellow cells. The gray cells fill "
                                  "in by themselves once your macros have run.")

        lesson.sheet_order = ["Start Here", "Practice", "BedBoard", "Intake", "Log", "Combined", "Settings", "Lists",
                              "Bonus"]

        if not selftest:
            return
        # --- simulate the learner's macros (fills then overwrite the checked cells with identical values)
        ws = wb["Intake"]
        for i, a in enumerate(added):
            r = in_sd.first_row + n0 + i
            for j, h in enumerate(in_sd.headers, 1):
                c = ws.cell(row=r, column=j, value=a[h])
                if h == "ArrivalDateTime":
                    c.number_format = "mm/dd/yyyy hh:mm"
                if h == "MRN":
                    c.number_format = "@"
        tbl = ws.tables["tblIntake"]
        new_ref = re.sub(r"\d+$", str(in_sd.first_row + n0 + len(added) - 1), tbl.ref)
        tbl.ref = new_ref
        if tbl.autoFilter is not None:
            tbl.autoFilter.ref = new_ref
        for name, units, total in ((rep_dec, dec_units, dec_total), (rep_jan, jan_units, jan_total)):
            rs = wb.create_sheet(name)
            rs["A1"] = "Bluestone Health System - Inpatient Census Report"
            for j, h in enumerate(["DeptID", "Unit", "Facility", "PatientDays", "BedDays", "ADC", "Occupancy"], 1):
                rs.cell(row=3, column=j, value=h)
            for i, row in enumerate(list(units) + [total]):
                for j, v in enumerate(row, 1):
                    rs.cell(row=4 + i, column=j, value=v)

    return L
