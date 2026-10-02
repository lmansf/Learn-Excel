"""Lesson 5.2 · VBA Fundamentals.

Data: 505 lab results from Bluestone's three ICUs (collected Jan–Apr 2025), sorted by result time.
Files written next to the workbook:
  starter/LabMacros.bas            TODO stubs for the practice macros (+ a complete Warmup and a buggy macro to debug)
  starter/Snippets.bas             the six predict-the-output snippets (no Option Explicit on purpose)
  solutions/LabMacros_Solution.bas reference solutions for the practice macros (spoilers)
  solutions/LabSnapshot_Solution.bas reference solution for the bonus (spoilers)
The .bas files are ASCII with CRLF line endings because the VBE imports Windows-style text files.

Smoke test (LibreOffice VBA-compatibility mode) — run from tools/:
  python3 -c "import lessons.lesson_5_2 as m; print(m.smoke_test())"
In LibreOffice, Debug.Print raises "error 91" (it has no Immediate window), so Warmup, AveragePotassium,
FindFirstCritical and FlagSlowStat report that error after writing correct results. Excel is unaffected.
LibreOffice also compares text with numbers more leniently than Excel, so the starter BuggyGlucoseCount
writes 0 there instead of stopping with run-time error 13 as it does in Excel.
"""
from __future__ import annotations

import re
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from xlcourse import Lesson, Task, data

CODE = "5.2"
ICU_DEPTS = ("D130", "D230", "D330")  # ICUs at Bluestone Memorial, Ashby Falls, Cedar Ridge

# Labs sheet column map (1-based numbers are what the VBA uses with Cells(r, c))
LAB_COLS = ["LabResultID", "EncounterID", "TestCode", "TestName", "ResultValue", "Units", "RefLow", "RefHigh",
            "AbnormalFlag", "Priority", "CollectedDateTime", "ResultedDateTime"]
FLAG_COL = "TATFlag"

# Output sheet cells the macros write to
OUT = {
    "rows": "B5", "crit": "B6", "kavg": "B7", "first": "B8", "above": "B9", "glu": "B10",
}
SUMMARY_FIRST, SUMMARY_LAST = 5, 40   # bonus summary table rows on Output (E:H)
TOP_CELL = "J5"

AMBER = (255, 235, 156)
RED = (255, 199, 206)
SLOW_LIMIT = 60          # minutes, STAT collect-to-result target
ABOVE_TEST, ABOVE_LIMIT = "CREAT", 2
GLU_LIMIT = 180


# ---------------------------------------------------------------------------
# VBA source (ASCII only). One source of truth for the .bas files, the
# Snippets sheet, the README answer key and the LibreOffice smoke test.
# ---------------------------------------------------------------------------
def _snippets(tat_minutes: int) -> "OrderedDict[str, dict]":
    s = OrderedDict()
    s["A"] = dict(name="SnippetA_Typo", code='''Sub SnippetA_Typo()
    Dim drawCount As Long
    drawCount = 5
    drawCuont = drawCount + 1       ' add one more blood draw
    Debug.Print drawCount
End Sub''')
    s["B"] = dict(name="SnippetB_Overflow", code='''Sub SnippetB_Overflow()
    Dim labRows As Integer
    labRows = 51527                 ' rows in the full two-year lab file
    Debug.Print labRows
End Sub''')
    s["C"] = dict(name="SnippetC_Turnaround", code=f'''Sub SnippetC_Turnaround()
    Dim tatMinutes As Long
    tatMinutes = {tat_minutes}                ' slowest STAT result on the Labs sheet
    Debug.Print tatMinutes \\ 60 & " h " & tatMinutes Mod 60 & " min"
End Sub''')
    s["D"] = dict(name="SnippetD_VitalSigns", code='''Sub SnippetD_VitalSigns()
    Dim hr As Long, checks As Long
    For hr = 0 To 23 Step 4         ' vital signs every 4 hours (q4h)
        checks = checks + 1
    Next hr
    Debug.Print hr
End Sub''')
    s["E"] = dict(name="SnippetE_Lactate", code='''Sub SnippetE_Lactate()
    Dim lactate As Double, category As String
    lactate = 4.6                   ' mmol/L
    Select Case lactate
        Case Is < 0.5
            category = "Low"
        Case Is <= 2
            category = "Normal"
        Case Is > 2
            category = "Elevated"
        Case Is > 4
            category = "Critical"
        Case Else
            category = "Check value"
    End Select
    Debug.Print category
End Sub''')
    s["F"] = dict(name="SnippetF_HalfLife", code='''Sub SnippetF_HalfLife()
    Dim level As Double, hours As Long
    level = 400                     ' ng/mL right after the dose
    Do While level > 50
        level = level / 2           ' the drug's half-life is 6 hours
        hours = hours + 6
    Loop
    Debug.Print hours
End Sub''')
    return s


HEADER_COMMENT = """' Labs sheet columns (row 1 = headers, data starts in row 2):
'    1 A LabResultID      5 E ResultValue      9 I AbnormalFlag
'    2 B EncounterID      6 F Units           10 J Priority
'    3 C TestCode         7 G RefLow          11 K CollectedDateTime
'    4 D TestName         8 H RefHigh         12 L ResultedDateTime
'                                             13 M TATFlag (task 11 fills it)
'
' Output sheet cells: B5 Warmup, B6 CountCriticals, B7 AveragePotassium,
'   B8 FindFirstCritical, B9 WriteCountAbove, B10 BuggyGlucoseCount.
'   Bonus: summary table E5:H40, top test in J5."""

WARMUP = '''Sub Warmup()
    Dim ws As Worksheet
    Dim lastRow As Long

    Set ws = ThisWorkbook.Worksheets("Labs")
    lastRow = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row      ' last filled row in column A
    ThisWorkbook.Worksheets("Output").Range("B5").Value = lastRow - 1   ' minus the header row
    Debug.Print "Warmup ran: " & (lastRow - 1) & " lab rows"
End Sub'''

SOL = OrderedDict()
SOL["CountCriticals"] = '''Sub CountCriticals()
    Dim ws As Worksheet
    Dim lastRow As Long
    Dim cell As Range
    Dim criticalCount As Long

    Set ws = ThisWorkbook.Worksheets("Labs")
    lastRow = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row

    For Each cell In ws.Range("I2:I" & lastRow)          ' the AbnormalFlag column
        If cell.Value = "HH" Or cell.Value = "LL" Then
            criticalCount = criticalCount + 1
        End If
    Next cell

    ThisWorkbook.Worksheets("Output").Range("B6").Value = criticalCount
    MsgBox criticalCount & " critical results", vbInformation, "Count criticals"
End Sub'''
SOL["AveragePotassium"] = '''Sub AveragePotassium()
    Dim ws As Worksheet
    Dim r As Long, lastRow As Long
    Dim total As Double, n As Long

    Set ws = ThisWorkbook.Worksheets("Labs")
    lastRow = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row

    For r = 2 To lastRow
        If ws.Cells(r, 3).Value = "K" Then           ' column C = TestCode
            total = total + ws.Cells(r, 5).Value     ' column E = ResultValue
            n = n + 1
        End If
    Next r

    If n > 0 Then
        ThisWorkbook.Worksheets("Output").Range("B7").Value = total / n
        Debug.Print n & " potassium results, average " & Format(total / n, "0.00")
    Else
        MsgBox "No potassium results found.", vbExclamation
    End If
End Sub'''
SOL["FindFirstCritical"] = '''Sub FindFirstCritical()
    Dim ws As Worksheet
    Dim r As Long, lastRow As Long
    Dim firstID As String

    Set ws = ThisWorkbook.Worksheets("Labs")
    lastRow = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row

    For r = 2 To lastRow
        Select Case ws.Cells(r, 9).Value             ' column I = AbnormalFlag
            Case "HH", "LL"
                firstID = ws.Cells(r, 1).Value       ' column A = LabResultID
                Exit For                             ' found it: stop looping
        End Select
    Next r

    If firstID = "" Then firstID = "None found"
    ThisWorkbook.Worksheets("Output").Range("B8").Value = firstID
    Debug.Print "First critical: " & firstID & " (row " & r & ")"
End Sub'''
SOL["FlagSlowStat"] = '''Sub FlagSlowStat()
    Dim ws As Worksheet
    Dim r As Long, lastRow As Long
    Dim minutes As Long, slowCount As Long

    Set ws = ThisWorkbook.Worksheets("Labs")
    lastRow = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row

    For r = 2 To lastRow
        ws.Cells(r, 13).ClearContents                ' start clean so a re-run is safe
        If ws.Cells(r, 10).Value = "STAT" Then       ' column J = Priority
            ' "n" = minutes (not "m", which is months)
            minutes = DateDiff("n", ws.Cells(r, 11).Value, ws.Cells(r, 12).Value)
            If minutes > 60 Then
                ws.Cells(r, 13).Value = "SLOW"       ' column M = TATFlag
                slowCount = slowCount + 1
            End If
        End If
    Next r

    Debug.Print slowCount & " slow STAT results"
End Sub'''
SOL["CountAbove"] = '''Sub CountAbove()
    Dim testCode As String
    Dim limitInput As Variant

    testCode = UCase(Trim(InputBox("Which TestCode? (for example CREAT)", "Count results above a limit")))
    If testCode = "" Then Exit Sub                   ' Cancel (or nothing typed)

    limitInput = Application.InputBox("Count results above what value?", "Count results above a limit", Type:=1)
    If VarType(limitInput) = vbBoolean Then Exit Sub ' Cancel returns False

    WriteCountAbove testCode, CDbl(limitInput)
End Sub

Sub WriteCountAbove(testCode As String, limit As Double)
    Dim ws As Worksheet
    Dim r As Long, lastRow As Long, n As Long

    Set ws = ThisWorkbook.Worksheets("Labs")
    lastRow = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row

    For r = 2 To lastRow
        If ws.Cells(r, 3).Value = testCode Then
            If ws.Cells(r, 5).Value > limit Then n = n + 1
        End If
    Next r

    ThisWorkbook.Worksheets("Output").Range("B9").Value = n
    MsgBox n & " " & testCode & " results above " & limit, vbInformation, "Count results above a limit"
End Sub'''
SOL["BuggyGlucoseCount"] = '''Sub BuggyGlucoseCount()
    Dim ws As Worksheet
    Dim r As Long, lastRow As Long
    Dim highCount As Long

    Set ws = ThisWorkbook.Worksheets("Labs")
    lastRow = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row

    For r = 2 To lastRow                             ' FIX 1: start below the header row
        If ws.Cells(r, 3).Value = "GLU" And ws.Cells(r, 5).Value > 180 Then   ' FIX 2: "GLU", not "Glu"
            highCount = highCount + 1
        End If
    Next r

    ThisWorkbook.Worksheets("Output").Range("B10").Value = highCount
End Sub'''

BUGGY = '''Sub BuggyGlucoseCount()
    Dim ws As Worksheet
    Dim r As Long, lastRow As Long
    Dim highCount As Long

    Set ws = ThisWorkbook.Worksheets("Labs")
    lastRow = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row

    For r = 1 To lastRow
        If ws.Cells(r, 3).Value = "Glu" And ws.Cells(r, 5).Value > 180 Then
            highCount = highCount + 1
        End If
    Next r

    ThisWorkbook.Worksheets("Output").Range("B10").Value = highCount
End Sub'''

BONUS_SOL = '''Sub LabSnapshot()
    Dim wsLab As Worksheet, wsOut As Worksheet
    Dim r As Long, lastRow As Long
    Dim k As Long, nextRow As Long
    Dim flag As String, code As String
    Dim found As Boolean
    Dim pct As Double, topPct As Double, topCode As String

    Set wsLab = ThisWorkbook.Worksheets("Labs")
    Set wsOut = ThisWorkbook.Worksheets("Output")
    lastRow = wsLab.Cells(wsLab.Rows.Count, 1).End(xlUp).Row

    ' 1. Reset: remove old fills and the old summary so the macro can be re-run
    wsLab.Range("A2:L" & lastRow).Interior.ColorIndex = xlNone
    wsOut.Range("E5:H40").ClearContents
    wsOut.Range("J5").ClearContents
    nextRow = 5                                      ' first empty row of the summary table

    For r = 2 To lastRow
        flag = wsLab.Cells(r, 9).Value               ' column I = AbnormalFlag
        code = wsLab.Cells(r, 3).Value               ' column C = TestCode

        ' 2. Color the row (columns A:L) by its flag
        Select Case flag
            Case "HH", "LL"
                wsLab.Range("A" & r & ":L" & r).Interior.Color = RGB(255, 199, 206)
            Case "H", "L"
                wsLab.Range("A" & r & ":L" & r).Interior.Color = RGB(255, 235, 156)
        End Select

        ' 3. Find this test in the summary table; add a new row if it isn't there yet
        found = False
        For k = 5 To nextRow - 1
            If wsOut.Cells(k, 5).Value = code Then
                found = True
                Exit For                             ' k now points at the test's row
            End If
        Next k
        If Not found Then
            k = nextRow
            wsOut.Cells(k, 5).Value = code
            wsOut.Cells(k, 6).Value = 0
            wsOut.Cells(k, 7).Value = 0
            nextRow = nextRow + 1
        End If

        ' 4. Update the counts
        wsOut.Cells(k, 6).Value = wsOut.Cells(k, 6).Value + 1          ' Results
        If flag <> "N" Then
            wsOut.Cells(k, 7).Value = wsOut.Cells(k, 7).Value + 1      ' Abnormal
        End If
    Next r

    ' 5. Percent abnormal for each test, tracking the highest
    For k = 5 To nextRow - 1
        pct = wsOut.Cells(k, 7).Value / wsOut.Cells(k, 6).Value
        wsOut.Cells(k, 8).Value = pct
        If pct > topPct Then
            topPct = pct
            topCode = wsOut.Cells(k, 5).Value
        End If
    Next k
    wsOut.Range("H5:H40").NumberFormat = "0.0%"
    wsOut.Range("J5").Value = topCode

    MsgBox "Tests summarized: " & (nextRow - 5) & vbNewLine & _
           "Highest % abnormal: " & topCode & " (" & Format(topPct, "0.0%") & ")", _
           vbInformation, "Lab snapshot"
End Sub'''


def _stub(task_no: int, name: str, lines: list[str], args: str = "") -> str:
    body = "\n".join(f"    ' {ln}" if ln else "    '" for ln in lines)
    return f"""' Task {task_no}
Sub {name}({args})
{body}
End Sub"""


def _starter_module() -> str:
    parts = [
        'Attribute VB_Name = "LabMacros"',
        "Option Explicit",
        "",
        "' =====================================================================",
        "' Lesson 5.2 - VBA Fundamentals - starter module (Bluestone Health, fictional)",
        "'",
        "' Import it: in the Visual Basic Editor choose File > Import File... and",
        "' pick this file. Write your code between each Sub and End Sub, then run",
        "' it with F5 (cursor inside the Sub) or from Developer > Macros.",
        "'",
        "' A compile error in one Sub (a red line, or an undeclared variable) can",
        "' stop every macro in this module from running, even Warmup. If that",
        "' happens, choose Debug > Compile VBAProject to jump to the problem line.",
        "'",
        HEADER_COMMENT,
        "' =====================================================================",
        "",
        "' Task 1 - already complete. Run it to check that macros work in your copy.",
        WARMUP,
        "",
        _stub(8, "CountCriticals", [
            "TODO: count the results whose AbnormalFlag (column I) is HH or LL.",
            "Use For Each over the range I2:I<lastRow>, then write the count",
            "to Output!B6. (Copy the Dim, Set ws, and lastRow lines from Warmup.)",
        ]),
        "",
        _stub(9, "AveragePotassium", [
            "TODO: with For r = 2 To lastRow, add up ResultValue (column E) for",
            "every row whose TestCode (column C) is K, and count those rows.",
            "Write total / count to Output!B7.",
        ]),
        "",
        _stub(10, "FindFirstCritical", [
            "TODO: find the FIRST row (from the top) whose AbnormalFlag is HH or LL.",
            "Write its LabResultID (column A) to Output!B8 and leave the loop",
            "with Exit For.",
        ]),
        "",
        _stub(11, "FlagSlowStat", [
            "TODO: for every STAT row (Priority, column J), work out the",
            "turnaround in minutes: DateDiff(\"n\", Collected (K), Resulted (L)).",
            "If it is more than 60, write SLOW in column M (TATFlag) of that row.",
            "Leave every other row's column M empty.",
        ]),
        "",
        _stub(12, "CountAbove", [
            "TODO: ask for a TestCode with InputBox and for a limit with",
            "Application.InputBox(..., Type:=1). Stop if the user cancels.",
            "Then call WriteCountAbove with the two answers. Pass the Variant",
            "limit as CDbl(yourVariable): a Variant passed to an argument declared",
            "As Double stops with Compile error: ByRef argument type mismatch.",
        ]),
        "",
        "Sub WriteCountAbove(testCode As String, limit As Double)",
        "    ' TODO: count the rows whose TestCode = testCode and ResultValue > limit.",
        "    ' Write the count to Output!B9 and show it in a MsgBox.",
        "End Sub",
        "",
        "' Task 13 - this macro should count glucose (GLU) results above 180 mg/dL",
        "' and write the count to Output!B10. It has TWO bugs. Debug it and fix it.",
        BUGGY,
        "",
        "' Bonus - see the Bonus sheet for the full brief.",
        "Sub LabSnapshot()",
        "    ' TODO (bonus):",
        "    ' 1. Clear old fills on Labs A2:L<lastRow>, Output!E5:H40 and Output!J5.",
        "    ' 2. Color each row A:L: HH/LL RGB(255, 199, 206), H/L RGB(255, 235, 156).",
        "    ' 3. Build the summary table at Output!E5 down: TestCode, Results,",
        "    '    Abnormal (flag is not N), PctAbnormal (Abnormal / Results).",
        "    ' 4. Write the TestCode with the highest PctAbnormal to Output!J5.",
        "End Sub",
        "",
    ]
    return "\n".join(parts)


def _snippets_module(snips) -> str:
    parts = [
        'Attribute VB_Name = "Snippets"',
        "' =====================================================================",
        "' Lesson 5.2 - predict-the-output snippets (Practice tasks 2-7)",
        "'",
        "' This module deliberately has NO Option Explicit line, so that",
        "' Snippet A can run. Never leave it out of your own modules.",
        "'",
        "' For each snippet: predict what it prints, type your prediction on the",
        "' Practice sheet, THEN click inside the Sub and press F5 to check.",
        "' Debug.Print writes to the Immediate window (View > Immediate Window).",
        "' =====================================================================",
        "",
    ]
    for key, s in snips.items():
        parts += [f"' Snippet {key}", s["code"], ""]
    return "\n".join(parts)


def _solution_module() -> str:
    parts = [
        'Attribute VB_Name = "LabMacrosSolution"',
        "Option Explicit",
        "",
        "' =====================================================================",
        "' Lesson 5.2 - reference solutions for Practice tasks 8-13 (SPOILERS)",
        "'",
        "' Import this into a spare copy of the workbook, or delete your own",
        "' LabMacros module first, so the two modules don't fight over names.",
        "'",
        HEADER_COMMENT,
        "' =====================================================================",
        "",
        "' Task 1",
        WARMUP,
        "",
    ]
    numbers = {"CountCriticals": 8, "AveragePotassium": 9, "FindFirstCritical": 10, "FlagSlowStat": 11,
               "CountAbove": 12, "BuggyGlucoseCount": 13}
    for name, code in SOL.items():
        parts += [f"' Task {numbers[name]}", code, ""]
    return "\n".join(parts)


def _bonus_module() -> str:
    return "\n".join([
        'Attribute VB_Name = "LabSnapshotSolution"',
        "Option Explicit",
        "",
        "' =====================================================================",
        "' Lesson 5.2 - reference solution for the bonus (SPOILERS)",
        "' Import into a spare copy of the workbook, or remove your own",
        "' LabSnapshot first.",
        "' =====================================================================",
        "",
        BONUS_SOL,
        "",
    ])


def _cell_text(md: str) -> str:
    """Markdown → plain text for an Excel cell: drop **bold**, `code` and *italic* markers."""
    s = (md or "").replace("**", "").replace("`", "")
    s = re.sub(r"(?<![\w*])\*(?=\S)([^*\n]+?)(?<=\S)\*(?![\w*])", r"\1", s)
    return s.strip()


def _fit_page(ws, landscape: bool):
    """Print/PDF setup: fit every column on one page width."""
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _crlf(text: str) -> bytes:
    text.encode("ascii")  # raises if a non-ASCII character slipped in (the VBE reads ANSI, not UTF-8)
    return text.replace("\r\n", "\n").replace("\n", "\r\n").encode("ascii")


# ---------------------------------------------------------------------------
# Lesson
# ---------------------------------------------------------------------------
def _lab_rows() -> list[dict]:
    enc = data.index(data.load("encounters"), "EncounterID")
    rows = [r for r in data.load("lab_results")
            if enc[r["EncounterID"]]["DeptID"] in ICU_DEPTS
            and datetime(2025, 1, 1) <= r["CollectedDateTime"] < datetime(2025, 5, 1)]
    rows.sort(key=lambda r: (r["ResultedDateTime"], r["LabResultID"]))
    return rows


def _tat(r) -> int:
    return round((r["ResultedDateTime"] - r["CollectedDateTime"]).total_seconds() / 60)


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="05-automation-vba", slug="02-vba-fundamentals",
        title="VBA Fundamentals", level="Expert", minutes=60,
        objectives=[
            "Navigate the Visual Basic Editor and organize code in modules",
            "Declare variables with the right data types and Option Explicit",
            "Control flow with If, Select Case, For, For Each, and Do loops",
            "Debug with breakpoints, stepping, the Immediate window, and Debug.Print",
        ],
        data_note="505 lab results from Bluestone's three intensive care units (Bluestone Memorial, Ashby Falls, "
                  "Cedar Ridge), collected January–April 2025 and listed in the order the results were released.",
    )

    labs = _lab_rows()
    n = len(labs)
    L.data_note = (f"{n} lab results from Bluestone's three intensive care units (Bluestone Memorial, Ashby Falls, "
                   "Cedar Ridge), collected January–April 2025 and listed in the order the results were released.")
    lab = L.add_table_sheet(
        "Labs", labs, table="tblLabs", columns=LAB_COLS, extra_cols=[FLAG_COL],
        formats={"ResultValue": "General", "RefLow": "General", "RefHigh": "General"},
        widths={"TestName": 24, "CollectedDateTime": 18, "ResultedDateTime": 18, FLAG_COL: 10},
    )
    first, last = lab.first_row, lab.last_row
    assert [lab.col(c) for c in ("TestCode", "ResultValue", "AbnormalFlag", "Priority", FLAG_COL)] == list("CEIJM")

    def rng(col):
        return lab.rng(col, absolute=False)

    # ------------------------------------------------------------------ snippet answers (simulated in Python)
    stat_tats = [_tat(r) for r in labs if r["Priority"] == "STAT"]
    tat_max = max(stat_tats)
    snips = _snippets(tat_max)

    # A: the typo creates a second (Variant) variable; drawCount itself never changes
    draw_count = 5
    _draw_cuont = draw_count + 1
    ans_a = draw_count
    # B: 51,527 doesn't fit in an Integer (-32,768 to 32,767) -> run-time error 6 (Overflow)
    ans_b = "6" if not (-32768 <= 51527 <= 32767) else "51527"
    # C: integer division and Mod
    ans_c = f"{tat_max // 60} h {tat_max % 60} min"
    # D: For hr = 0 To 23 Step 4 -> counter after the loop
    hr, checks = 0, 0
    while hr <= 23:
        checks += 1
        hr += 4
    ans_d = hr
    # E: Select Case takes the FIRST matching Case
    lactate = 4.6
    for cond, label in [(lambda v: v < 0.5, "Low"), (lambda v: v <= 2, "Normal"), (lambda v: v > 2, "Elevated"),
                        (lambda v: v > 4, "Critical")]:
        if cond(lactate):
            ans_e = label
            break
    else:
        ans_e = "Check value"
    # F: Do While with a pre-test
    level, hours = 400.0, 0
    while level > 50:
        level /= 2
        hours += 6
    ans_f = hours

    # ------------------------------------------------------------------ macro answers
    crit = [r for r in labs if r["AbnormalFlag"] in ("HH", "LL")]
    k_vals = [r["ResultValue"] for r in labs if r["TestCode"] == "K"]
    k_avg = sum(k_vals) / len(k_vals)
    first_crit = crit[0]["LabResultID"]
    first_crit_row = first + labs.index(crit[0])
    slow_flags = ["SLOW" if (r["Priority"] == "STAT" and _tat(r) > SLOW_LIMIT) else None for r in labs]
    slow_count = sum(1 for f in slow_flags if f)
    above_count = sum(1 for r in labs if r["TestCode"] == ABOVE_TEST and r["ResultValue"] > ABOVE_LIMIT)
    glu_count = sum(1 for r in labs if r["TestCode"] == "GLU" and r["ResultValue"] > GLU_LIMIT)
    glu_row = first + next(i for i, r in enumerate(labs) if r["TestCode"] == "GLU")   # first glucose row (hint)

    # ------------------------------------------------------------------ bonus answers
    amber = sum(1 for r in labs if r["AbnormalFlag"] in ("H", "L"))
    summary: "OrderedDict[str, list]" = OrderedDict()
    for r in labs:
        s = summary.setdefault(r["TestCode"], [0, 0])
        s[0] += 1
        if r["AbnormalFlag"] != "N":
            s[1] += 1
    summary_rows = [(c, a, b, b / a) for c, (a, b) in summary.items()]
    top_code = max(summary_rows, key=lambda x: x[3])[0]  # first one wins a tie, like the VBA's ">" test
    creat_abn = summary["CREAT"][1]
    lact_pct = summary["LACT"][1] / summary["LACT"][0]
    # The B5 explanation names the top two tests and the test with the most critical (HH/LL) results.
    ranked = sorted(summary_rows, key=lambda x: -x[3])
    assert (ranked[0][0], ranked[1][0]) == ("GLU", "LACT"), ranked[:2]
    crit_by_test = {}
    for r in crit:
        crit_by_test[r["TestCode"]] = crit_by_test.get(r["TestCode"], 0) + 1
    assert max(crit_by_test, key=crit_by_test.get) == "LACT", crit_by_test
    summary_fill = [v for row in summary_rows for v in row]
    summary_range = f"Output!E{SUMMARY_FIRST}:H{SUMMARY_FIRST + len(summary_rows) - 1}"
    n_tests = len(summary_rows)

    def out_summary(cell):
        return f'=IF(Output!{cell}="","",Output!{cell})'

    def out_fill(cell, value):
        return {"range": f"Output!{cell}:{cell}", "values": [value]}

    L.start_notes = [
        "Macros can't be saved in an .xlsx file. Before you write any code, choose File → Save As and pick "
        "'Excel Macro-Enabled Workbook (*.xlsm)'.",
        "Import starter/LabMacros.bas and starter/Snippets.bas from the lesson folder: open the Visual Basic Editor "
        "(Alt+F11; Mac: Option+F11 or Developer → Visual Basic), then File → Import File…",
        "Debug.Print writes to the Immediate window. Open it in the Visual Basic Editor with Ctrl+G (Mac: View → Immediate Window) "
        "before you run the snippets.",
        "The Snippets sheet shows the code for the predict-the-output tasks. The Output sheet is where your macros write their results.",
    ]
    L.practice_intro = (
        "Save this workbook as .xlsm, then import starter/LabMacros.bas and starter/Snippets.bas (VBE → File → Import File…). "
        "Tasks 2–7 are predict-the-output questions about the code on the Snippets sheet: type your prediction first, then run the "
        "snippet to check it. In tasks 1 and 8–13 you run macros, and the gray cells read what your macros wrote to the Output sheet "
        "(or to the TATFlag column on the Labs sheet).")

    L.tasks = [
        # ---------------------------------------------------------- 1 warm-up
        Task("Import the starter module, then run the Warmup macro (it's already written). It writes the number of lab rows "
             "to Output!B5, and the gray cell reads it from there.",
             answer=n, title="Warmup: run your first macro",
             solution=WARMUP, solution_lang="vba",
             summary=out_summary(OUT["rows"]), fill=out_fill(OUT["rows"], n),
             live=f"=COUNTA({rng('LabResultID')})",
             hint="Click inside Warmup and press F5, or in Excel press Alt+F8 (Mac: Option+F8), pick Warmup, and click Run",
             explanation="Warmup finds the last filled row in column A, subtracts 1 for the header row, and writes the result "
                         "into a cell. If the gray cell stays empty, check that you saved as .xlsm, enabled macros, and ran "
                         "Warmup in **this** workbook. If you made a Personal Macro Workbook in Lesson 5.1, also check that "
                         "LabMacros was imported into this workbook's project and not into PERSONAL.XLSB. Open the "
                         "Immediate window (Ctrl + G; Mac: View → Immediate Window) to see the line Debug.Print wrote."),
        # ---------------------------------------------------------- 2-7 predict the output
        Task("Snippet A: what number does the Immediate window show? Predict first, then run SnippetA_Typo to check.",
             answer=ans_a, solution="Trace it: `drawCount` is set to 5. The next line assigns to `drawCuont`, which is a "
                                    "*different* variable, so `drawCount` is still 5 when it's printed.",
             hint="Which variable does Debug.Print actually read?", live=False,
             explanation="Without **Option Explicit**, VBA silently creates a new Variant for any misspelled name, so the typo "
                         "doesn't raise an error and the result is quietly wrong. Add `Option Explicit` at the top of the "
                         "module and the same code stops with *Compile error: Variable not defined*, highlighting `drawCuont`."),
        Task("Snippet B: running SnippetB_Overflow stops with a run-time error. Type the error number.",
             answer=ans_b, accept=["overflow", "error 6", "run-time error 6", "runtime error 6"],
             solution="Run-time error **6: Overflow**. An `Integer` holds only −32,768 to 32,767, and 51,527 is bigger.",
             hint="Check the Integer row of the data-types table", live=False,
             explanation="`Integer` is a 16-bit type left over from early versions of VBA. Use `Long` (up to about 2.1 billion) "
                         "for counts and row numbers. A worksheet has 1,048,576 rows, so an `Integer` row counter fails on "
                         "any large sheet. `Long` is also no slower on modern computers."),
        Task("Snippet C: what exactly does the Immediate window show? Type the whole line in the same pattern, for example 1 h 5 min.",
             answer=ans_c, accept=[ans_c.replace(" h ", "h "), ans_c.replace(" min", "min"),
                                   ans_c.replace(" h ", "h ").replace(" min", "min"),
                                   ans_c.replace(" ", "")],
             solution=f"`{tat_max} \\ 60` is {tat_max // 60} (whole hours) and `{tat_max} Mod 60` is {tat_max % 60} "
                      f"(leftover minutes). `&` joins everything into one string: **{ans_c}**.",
             hint="\\ keeps the whole part of a division, and Mod keeps the remainder", live=False,
             explanation="Arithmetic operators run before `&`, so each piece is calculated first and then joined. Integer "
                         "division `\\` and `Mod` are the standard way to split minutes into hours and minutes. "
                         f"({tat_max} minutes was the slowest STAT turnaround on the Labs sheet, against a "
                         f"{SLOW_LIMIT}-minute target.)"),
        Task("Snippet D: what number does Debug.Print hr show after the loop finishes?",
             answer=ans_d, solution=f"The loop body runs for hr = 0, 4, 8, 12, 16, 20 ({checks} checks). `Next` then adds the "
                                    f"Step again (20 + 4 = {ans_d}), which is past 23, so the loop ends with hr = **{ans_d}**.",
             hint="List every value hr takes, then apply Step one more time", live=False,
             explanation="After a For…Next loop finishes normally, the counter holds the first value that *failed* the test, "
                         "not the last value used. That's why code that needs \"the last row processed\" should save it in "
                         "its own variable or leave the loop with Exit For."),
        Task("Snippet E: which category does the Immediate window show? (Type the word.)",
             answer=ans_e, solution="4.6 is not < 0.5 and not ≤ 2, but it **is** > 2, so `Case Is > 2` matches and sets "
                                    f"**{ans_e}**. VBA then jumps to End Select, so `Case Is > 4` is never tested.",
             hint="Select Case runs only the first Case that matches", live=False,
             explanation="Order matters in Select Case. Put the most specific (most extreme) test first: `Case Is > 4` before "
                         "`Case Is > 2`. Here a critical lactate of 4.6 mmol/L is mislabeled \"Elevated\", which is exactly "
                         "the kind of silent bug that testing with real values catches."),
        Task("Snippet F: how many hours does the Immediate window show?",
             answer=ans_f, solution="level 400 → 200 (6 h) → 100 (12 h) → 50 (18 h). Before the next pass VBA tests "
                                    f"`50 > 50`, which is False, so the loop stops: **{ans_f}**.",
             hint="Write down level and hours after each pass, and test the condition before every pass", live=False,
             explanation="`Do While … Loop` tests the condition *before* each pass, so the body may run zero times. Use "
                         "`Do … Loop While` when the body must run at least once. Watch the boundary: `> 50` stops at "
                         "exactly 50, while `>= 50` would run one more pass and print 24."),
        # ---------------------------------------------------------- 8-13 macros
        Task("Complete CountCriticals: use For Each to loop over the AbnormalFlag cells (Labs column I) and count the results "
             "flagged HH or LL (critical values). Write the count to Output!B6.",
             answer=len(crit), solution=SOL["CountCriticals"], solution_lang="vba",
             summary=out_summary(OUT["crit"]), fill=out_fill(OUT["crit"], len(crit)),
             live=f'=COUNTIF({rng("AbnormalFlag")},"HH")+COUNTIF({rng("AbnormalFlag")},"LL")',
             hint="For Each cell In ws.Range(\"I2:I\" & lastRow) … If … Or … Then",
             explanation="This is the **counter pattern**: start a Long at 0 and add 1 each time a condition is true. Each `Or` "
                         "side must be a full comparison: `cell.Value = \"HH\" Or cell.Value = \"LL\"`. Writing "
                         "`cell.Value = \"HH\" Or \"LL\"` gives a Type mismatch error. Cross-check with a formula: "
                         f"`=COUNTIF({rng('AbnormalFlag')},\"HH\")+COUNTIF({rng('AbnormalFlag')},\"LL\")`."),
        Task("Complete AveragePotassium: loop over the rows with For…Next, add up ResultValue for every potassium result "
             "(TestCode K) and count them, then write the average to Output!B7. (Full precision or 2 decimal places are both accepted.)",
             answer=k_avg, fmt="0.00", solution=SOL["AveragePotassium"], solution_lang="vba",
             summary=out_summary(OUT["kavg"]), fill=out_fill(OUT["kavg"], k_avg),
             live=f'=AVERAGEIF({rng("TestCode")},"K",{rng("ResultValue")})',
             hint="Two accumulators: a Double for the total and a Long for the count",
             explanation=f"The **accumulator pattern** keeps a running total and a running count, then divides once at the end. "
                         f"The `If n > 0` guard matters when no rows match: then `total` and `n` are both still 0, and `0 / 0` "
                         f"stops with run-time error 6 (*Overflow*). Any other number divided by 0 gives error 11 "
                         f"(*Division by zero*). "
                         f"The {len(k_vals)} potassium results average {k_avg:.2f} mmol/L, which is comfortably inside the "
                         f"3.5–5.1 reference range. If you round in VBA, remember that `Round` uses banker's rounding."),
        Task("Complete FindFirstCritical: find the first row (from the top) whose AbnormalFlag is HH or LL, write its "
             "LabResultID to Output!B8, and leave the loop with Exit For.",
             answer=first_crit, solution=SOL["FindFirstCritical"], solution_lang="vba",
             summary=out_summary(OUT["first"]), fill=out_fill(OUT["first"], first_crit),
             live=(f'=INDEX({rng("LabResultID")},MIN(IFERROR(MATCH("HH",{rng("AbnormalFlag")},0),99999),'
                   f'IFERROR(MATCH("LL",{rng("AbnormalFlag")},0),99999)))'),
             hint="Case \"HH\", \"LL\" matches either value, and Exit For stops the loop",
             explanation=f"`Exit For` leaves the loop the moment the first match is found (sheet row {first_crit_row}), so the "
                         "macro doesn't waste time on the remaining rows and `firstID` can't be overwritten by a later match. "
                         "A Case with a comma-separated list (`Case \"HH\", \"LL\"`) matches any of the values."),
        Task(f"Complete FlagSlowStat: for every STAT result whose turnaround (ResultedDateTime minus CollectedDateTime, in "
             f"minutes, with DateDiff) is more than {SLOW_LIMIT} minutes, write SLOW in that row's TATFlag cell (Labs column M). "
             f"Leave all other rows empty. The gray cell counts the SLOW flags.",
             answer=slow_count, title="FlagSlowStat (count of SLOW flags)",
             solution=SOL["FlagSlowStat"], solution_lang="vba",
             summary=f'=IF(COUNTA({rng(FLAG_COL)})=0,"",COUNTIF({rng(FLAG_COL)},"SLOW"))',
             fill={"range": f"Labs!{lab.col(FLAG_COL)}{first}:{lab.col(FLAG_COL)}{last}", "values": slow_flags},
             live=(f'=SUMPRODUCT(({rng("Priority")}="STAT")*'
                   f'(ROUND(({rng("ResultedDateTime")}-{rng("CollectedDateTime")})*1440,0)>{SLOW_LIMIT}))'),
             hint="Use DateDiff(\"n\", start, finish), then combine the two tests with And or a nested If",
             explanation=f"`DateDiff(\"n\", …)` returns whole minutes. The interval code is `\"n\"` because `\"m\"` means "
                         f"*months*. {slow_count} of the {len(stat_tats)} STAT results missed the {SLOW_LIMIT}-minute target. "
                         f"Clearing column M at the start of each pass makes the macro safe to run twice, because an old "
                         f"SLOW flag can't survive a re-run."),
        Task(f"Complete CountAbove and WriteCountAbove. CountAbove asks for a TestCode (InputBox) and a limit "
             f"(Application.InputBox with Type:=1), then calls WriteCountAbove, which counts the results of that test above the "
             f"limit, writes the count to Output!B9, and shows it in a MsgBox. Run CountAbove and enter {ABOVE_TEST} and {ABOVE_LIMIT}.",
             answer=above_count, title=f"CountAbove ({ABOVE_TEST} above {ABOVE_LIMIT})",
             solution=SOL["CountAbove"], solution_lang="vba",
             summary=out_summary(OUT["above"]), fill=out_fill(OUT["above"], above_count),
             live=f'=COUNTIFS({rng("TestCode")},"{ABOVE_TEST}",{rng("ResultValue")},">{ABOVE_LIMIT}")',
             hint="Call a Sub that takes arguments without parentheses: SubName arg1, arg2. Convert the Variant limit with CDbl",
             explanation=f"Splitting the work into two Subs keeps the dialog code apart from the counting code. "
                         f"You can test `WriteCountAbove` straight from the Immediate window "
                         f"(`WriteCountAbove \"{ABOVE_TEST}\", {ABOVE_LIMIT}`) without clicking through dialogs. "
                         f"`Application.InputBox` with `Type:=1` refuses non-numbers and returns **False** on Cancel, which "
                         f"is why the result goes into a Variant and is checked with `VarType`. Creatinine above "
                         f"{ABOVE_LIMIT} mg/dL is a rough screen for impaired kidney function. Formal acute kidney injury "
                         f"criteria compare each patient with their own baseline."),
        Task(f"BuggyGlucoseCount (already in the starter module) should count glucose (GLU) results above {GLU_LIMIT} mg/dL and "
             f"write the count to Output!B10. It has two bugs: the first stops it with a run-time error, and after you fix that "
             f"one it writes 0. Use F8, the Locals window, and the Immediate window to find and fix both, then run it.",
             answer=glu_count, title="Debug BuggyGlucoseCount",
             solution=SOL["BuggyGlucoseCount"], solution_lang="vba",
             summary=out_summary(OUT["glu"]), fill=out_fill(OUT["glu"], glu_count),
             live=f'=COUNTIFS({rng("TestCode")},"GLU",{rng("ResultValue")},">{GLU_LIMIT}")',
             hint=(f"When it stops, click Debug and hover over r. While it's paused, try ? ws.Cells({glu_row}, 3).Value "
                   f"(a glucose row) in the Immediate window"),
             explanation="**Bug 1:** the loop starts at row 1, the header row. `ws.Cells(1, 5).Value` is the text "
                         "\"ResultValue\", and comparing text with the number 180 raises *Run-time error 13: Type mismatch*. "
                         "VBA's `And` evaluates both sides even when the first is already False, so the TestCode test "
                         "doesn't protect you. **Bug 2:** text comparison in VBA is case-sensitive by default, so "
                         "`\"GLU\" = \"Glu\"` is False and nothing is ever counted. Fix it by matching the data exactly "
                         "(`\"GLU\"`), or compare `UCase(ws.Cells(r, 3).Value) = \"GLU\"`."),
    ]

    # ------------------------------------------------------------------ bonus
    L.bonus_title = "Bonus: one-click lab snapshot"
    L.bonus_scenario = (
        "The ICU medical director wants a one-click 'lab snapshot' for the morning huddle. Complete the LabSnapshot stub in "
        "the starter module so that it does four things:\n"
        f"1. Clears the old fills on Labs!A{first}:L{last}, the old summary in Output!E{SUMMARY_FIRST}:H{SUMMARY_LAST}, and "
        f"Output!{TOP_CELL}, so you can run it again safely.\n"
        "2. Colors each data row (columns A:L) by its AbnormalFlag: light red RGB(255, 199, 206) for HH or LL, amber "
        "RGB(255, 235, 156) for H or L, and no fill for N.\n"
        f"3. Builds a summary table on the Output sheet from row {SUMMARY_FIRST} down, with one row per TestCode in the order "
        "each code first appears. The columns are TestCode (E), Results (F), Abnormal (G, any flag other than N), and "
        "PctAbnormal (H, Abnormal ÷ Results, formatted 0.0%).\n"
        f"4. Writes the TestCode with the highest PctAbnormal to Output!{TOP_CELL}.\n\n"
        "Don't use Scripting.Dictionary, because that's Lesson 5.4. A nested loop that searches the table you're building "
        "is enough.")
    L.bonus = [
        Task("Run LabSnapshot. Then click the filter arrow on the Labs table's AbnormalFlag header and choose Filter by Color → "
             f"your amber fill. How many rows are amber? (The status bar shows 'x of {n} records found'.)",
             answer=amber, solution=BONUS_SOL, solution_lang="vba", live=f'=COUNTIF({rng("AbnormalFlag")},"H")+COUNTIF({rng("AbnormalFlag")},"L")',
             hint="Select Case flag: Case \"HH\", \"LL\" … Case \"H\", \"L\" …",
             title="Run LabSnapshot, then count the amber rows",
             explanation="A worksheet formula can't see fill colors, so Filter by Color (Lesson 1.6) is how you check the "
                         "coloring. Clearing "
                         "the old fills first matters: without it, a row that was amber yesterday would stay amber even if its "
                         "flag changed. Clear the filter afterwards (Data → Clear). The full macro is also in "
                         "solutions/LabSnapshot_Solution.bas."),
        Task("How many TestCode rows does your summary table have? (The gray cell counts the codes in Output!E5:E40.)",
             answer=n_tests,
             solution="Run LabSnapshot (full code in B1). The gray cell uses `=COUNTA(Output!E5:E40)`.",
             summary=f'=IF(COUNTA(Output!E{SUMMARY_FIRST}:E{SUMMARY_LAST})=0,"",COUNTA(Output!E{SUMMARY_FIRST}:E{SUMMARY_LAST}))',
             fill={"range": summary_range, "values": summary_fill},
             live=f"=SUMPRODUCT(1/COUNTIF({rng('TestCode')},{rng('TestCode')}))",
             hint="Search rows 5 to nextRow - 1 with an inner For loop. If the code isn't there, add a row",
             explanation="The inner loop searches only the rows already written (5 to nextRow − 1). If it finishes without a "
                         "match, the code is new, so the macro writes it in `nextRow` and moves `nextRow` down one row. "
                         f"That search-or-add pattern is what a Dictionary does for you in Lesson 5.4. The ICU slice has "
                         f"{n_tests} distinct tests."),
        Task("How many creatinine (CREAT) results are abnormal (any flag other than N)? (The gray cell looks up CREAT in your table.)",
             answer=creat_abn,
             solution="Run LabSnapshot. The gray cell uses `=INDEX(Output!G5:G40,MATCH(\"CREAT\",Output!E5:E40,0))`.",
             summary=(f'=IFERROR(INDEX(Output!G{SUMMARY_FIRST}:G{SUMMARY_LAST},'
                      f'MATCH("CREAT",Output!E{SUMMARY_FIRST}:E{SUMMARY_LAST},0)),"")'),
             fill={"range": summary_range, "values": summary_fill},
             live=f'=COUNTIFS({rng("TestCode")},"CREAT",{rng("AbnormalFlag")},"<>N")',
             hint="If flag <> \"N\" Then add 1 to column G of the test's row",
             explanation="`<> \"N\"` counts H, L, HH, and LL together, which is what *abnormal* means here. Updating the count in "
                         "the found row `k` works because `Exit For` left `k` pointing at that row."),
        Task("What percentage of lactate (LACT) results are abnormal? (The gray cell looks up LACT in your table.)",
             answer=lact_pct, fmt="0.0%",
             solution="Run LabSnapshot. The gray cell uses `=INDEX(Output!H5:H40,MATCH(\"LACT\",Output!E5:E40,0))`.",
             summary=(f'=IFERROR(INDEX(Output!H{SUMMARY_FIRST}:H{SUMMARY_LAST},'
                      f'MATCH("LACT",Output!E{SUMMARY_FIRST}:E{SUMMARY_LAST},0)),"")'),
             fill={"range": summary_range, "values": summary_fill},
             live=f'=COUNTIFS({rng("TestCode")},"LACT",{rng("AbnormalFlag")},"<>N")/COUNTIF({rng("TestCode")},"LACT")',
             hint="Compute the percentages in a second loop, after every row has been counted",
             explanation=f"Store the ratio ({lact_pct:.3f}…) and let `.NumberFormat = \"0.0%\"` display it. If you store "
                         f"{lact_pct * 100:.1f} instead, the check accepts that too, but a true fraction is easier to reuse in later formulas. Dividing in a "
                         "separate loop after counting means each test is divided once, with its final totals."),
        Task("Which TestCode has the highest percentage of abnormal results? (The gray cell reads Output!J5.)",
             answer=top_code,
             solution="Run LabSnapshot. Step 5 keeps the best value so far in `topPct` and its code in `topCode`.",
             summary=out_summary(TOP_CELL), fill=out_fill(TOP_CELL, top_code),
             live=False,
             hint="Track the best value so far: If pct > topPct Then …",
             explanation=f"The **running maximum** pattern compares each value with the best one seen so far. In this ICU "
                         f"sample, glucose ({summary['GLU'][1]} of {summary['GLU'][0]}, {ranked[0][3]:.1%}) edges out "
                         f"lactate ({ranked[1][3]:.1%}). That's plausible, because stress hyperglycemia is common in "
                         f"critically ill patients. Lactate has the most *critical* values, though ({crit_by_test['LACT']} of "
                         f"the {len(crit)} HH or LL results), so \"most often abnormal\" and \"most dangerous\" are "
                         f"different questions."),
    ]

    # ------------------------------------------------------------------ extra files
    L.extra_files = {
        "starter/LabMacros.bas": _crlf(_starter_module()),
        "starter/Snippets.bas": _crlf(_snippets_module(snips)),
        "solutions/LabMacros_Solution.bas": _crlf(_solution_module()),
        "solutions/LabSnapshot_Solution.bas": _crlf(_bonus_module()),
    }

    # ------------------------------------------------------------------ Output + Snippets sheets
    @L.customize
    def _sheets(wb, lesson, selftest):
        navy = "1F4E79"
        hdr_fill = PatternFill("solid", fgColor=navy)
        out_fill_ = PatternFill("solid", fgColor="DDEBF7")
        thin = Side(style="thin", color="BFBFBF")
        box = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws = wb.create_sheet("Output")
        ws.sheet_properties.tabColor = "2E75B6"
        ws["A1"] = "Output: your macros write their results here"
        ws["A1"].font = Font(bold=True, size=14, color=navy)
        ws["A2"] = ("Macros write into the blue cells. The Practice and Bonus sheets read these cells, so keep them where they "
                    "are and don't type in them by hand.")
        ws["A2"].font = Font(italic=True, color="595959")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells("A2:J2")
        ws.row_dimensions[2].height = 32
        ws["E3"] = "Bonus: LabSnapshot writes one row per TestCode from row 5 down."
        ws["E3"].font = Font(italic=True, color="595959")
        heads = {"A4": "Result", "B4": "Value", "C4": "Written by (task)", "E4": "TestCode", "F4": "Results",
                 "G4": "Abnormal", "H4": "PctAbnormal", "J4": "Highest % abnormal"}
        for coord, text in heads.items():
            c = ws[coord]
            c.value = text
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = hdr_fill
            c.border = box
        items = [
            ("rows", "Lab rows on the Labs sheet", "Warmup (1)"),
            ("crit", "Critical results (HH or LL)", "CountCriticals (8)"),
            ("kavg", "Average potassium, mmol/L", "AveragePotassium (9)"),
            ("first", "First critical LabResultID", "FindFirstCritical (10)"),
            ("above", "Results above your limit", "WriteCountAbove (12)"),
            ("glu", f"Glucose results above {GLU_LIMIT} mg/dL", "BuggyGlucoseCount (13)"),
        ]
        for key, label, who in items:
            row = int(OUT[key][1:])
            ws.cell(row=row, column=1, value=label).border = box
            v = ws.cell(row=row, column=2)
            v.fill = out_fill_
            v.border = box
            if key == "kavg":
                v.number_format = "0.00"
            w = ws.cell(row=row, column=3, value=who)
            w.font = Font(italic=True, color="7F7F7F")
            w.border = box
        ws[TOP_CELL].fill = out_fill_
        ws[TOP_CELL].border = box
        for r in range(SUMMARY_FIRST, SUMMARY_LAST + 1):
            ws.cell(row=r, column=8).number_format = "0.0%"
        for col, wdt in zip("ABCDEFGHIJ", (38, 14, 24, 3, 11, 10, 11, 13, 3, 20)):
            ws.column_dimensions[col].width = wdt
        ws.freeze_panes = "A5"

        sn = wb.create_sheet("Snippets")
        sn.sheet_properties.tabColor = "7030A0"
        sn.column_dimensions["A"].width = 100
        sn["A1"] = "Snippets for Practice tasks 2–7"
        sn["A1"].font = Font(bold=True, size=14, color=navy)
        sn["A2"] = ("Predict what each snippet prints, type your prediction on the Practice sheet, then run it from the Snippets "
                    "module (starter/Snippets.bas) to check. Debug.Print writes to the Immediate window (Ctrl+G; Mac: View → Immediate Window).")
        sn["A2"].font = Font(italic=True, color="595959")
        sn["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        sn.row_dimensions[2].height = 48
        code_fill = PatternFill("solid", fgColor="F2F2F2")
        r = 4
        for i, (key, s) in enumerate(snips.items()):
            c = sn.cell(row=r, column=1, value=f"Snippet {key}  ·  Practice task {i + 2}")
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = hdr_fill
            r += 1
            for line in s["code"].split("\n"):
                c = sn.cell(row=r, column=1, value=line)
                c.font = Font(name="Consolas", size=10)
                c.fill = code_fill
                c.alignment = Alignment(indent=0)
                c.data_type = "s"
                r += 1
            r += 1

        _fit_page(ws, landscape=True)
        _fit_page(sn, landscape=False)

        # The library's generic instructions say "type a formula or value". Here most answers are gray cells that fill in
        # when a macro runs, so say that instead.
        generic = ("Type a formula or value into each yellow cell.", "Type a formula or value in each yellow cell.")
        lesson_text = "Type your answer in each yellow cell, or run the task's macro so its gray cell fills in by itself."
        for name in ("Start Here", "Practice", "Bonus"):
            if name not in wb.sheetnames:
                continue
            for row in wb[name].iter_rows(max_row=40):
                for c in row:
                    if isinstance(c.value, str) and not c.value.startswith("="):
                        for g in generic:
                            if g in c.value:
                                c.value = c.value.replace(g, lesson_text)

        # The key sheets show Markdown markers (** ` *) literally. Prose solutions and explanations read better as plain
        # text, so strip them there; VBA solutions stay as code.
        for key_name, tasks in (("Answer Key", lesson.tasks), ("Bonus Key", lesson.bonus)):
            if key_name not in wb.sheetnames:
                continue
            key = wb[key_name]
            for i, t in enumerate(tasks):
                row = 5 + i
                if t.solution_lang != "vba":
                    sol = key.cell(row=row, column=4)
                    sol.value = _cell_text(t.solution)
                    sol.font = Font(size=11)
                ex = key.cell(row=row, column=6)
                if isinstance(ex.value, str):
                    ex.value = _cell_text(ex.value)

    L.sheet_order = ["Start Here", "Practice", "Snippets", "Labs", "Output", "Bonus", "Answer Key", "Bonus Key"]

    # Keep the README's snippet listings in sync with the code above.
    readme = L.dir / "README.md"
    if readme.exists():
        text = readme.read_text(encoding="utf-8")
        missing = [k for k, s in snips.items() if "TODO" not in text and s["code"] not in text]
        if missing:
            print(f"  warning: README is missing the current code for snippet(s) {', '.join(missing)}")
    return L


# ---------------------------------------------------------------------------
# Smoke test in LibreOffice (not part of the build)
# ---------------------------------------------------------------------------
def smoke_test():
    """Run the reference solutions in LibreOffice and compare the cells they write with the Python answers."""
    import tempfile

    from xlcourse import vba

    L = build()
    sol = L.dir / "solutions" / "LabMacros_Solution.bas"
    bon = L.dir / "solutions" / "LabSnapshot_Solution.bas"
    res = vba.run(
        workbook=L.dir / L.workbook_name,
        modules=[sol, bon],
        macros=["Warmup", "CountCriticals", "AveragePotassium", "FindFirstCritical", "FlagSlowStat", "XlcAbove",
                "BuggyGlucoseCount", "LabSnapshot"],
        extra_code=f'Sub XlcAbove()\n  WriteCountAbove "{ABOVE_TEST}", {ABOVE_LIMIT}\nEnd Sub',
        read=[("Output", c) for c in ("B5", "B6", "B7", "B8", "B9", "B10", "E5", "F5", "G5", "H5", "J5")]
             + [("Output", f"E{r}") for r in range(5, 17)]
             + [("Output", f"G{r}") for r in range(5, 17)]
             + [("Output", f"H{r}") for r in range(5, 17)],
        save_as=Path(tempfile.mkdtemp(prefix="xlc_5_2_")) / "smoke.xlsx",   # inspect fills / TATFlag here
    )
    return res
