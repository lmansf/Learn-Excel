"""Lesson 4.3 · Power Query: Import, Transform & Combine.

The practice data are CSV files written next to the workbook (lesson.extra_files), because the point of
the lesson is importing files with Power Query:

  data/claims_monthly/claims_2025_01.csv … _12.csv   2025 claims split by SubmitDate month (raw strings
                                                     copied from data/claims.csv, same 13 columns)
  data/new_month/claims_2026_01.csv                  claims submitted in January 2026 (bonus refresh). These
                                                     are the dataset's own late-2025 services billed in
                                                     January 2026, so their IDs join to encounters_2025.csv
  data/payers.csv                                    copy of data/payers.csv
  data/encounters_2025.csv                           encounters discharged in 2025 (10 columns)
  data/budget_2025_wide.csv                          2025 budget & actual, one row per dept × category ×
                                                     measure, one column per month plus an FY Total column
  data/4.3-power-query-data.zip                      all of the above in one download (deterministic zip)
  starter/AgedPending.m                              the M query learners paste and edit (task 11)
  solutions/*.m                                      reference queries (spoilers)

Every answer is computed by re-reading those generated CSV files as text and replaying the Power Query
steps in Python (filters, joins, groups, unpivot), then cross-checked against xlcourse.data. The
workbook itself holds a hand-maintained mapping table (tblDenialMap) with two deliberate key defects
(a trailing space and a lower-case letter) so learners see that merges match text exactly.

Solutions are M code (solution_lang="m"), so the key has no live formulas; the self-test types each
value into its answer cell.
"""
from __future__ import annotations

import csv
import io
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal

from openpyxl.styles import Alignment, Font, PatternFill

from xlcourse import Lesson, Task, data
from xlcourse.lesson import NAVY

CODE = "4.3"
AS_OF = date(2025, 12, 31)
WIN_DATA = "C:\\PQ\\data\\"          # example location used in the guide and the generated-code solutions
ZIP_NAME = "data/4.3-power-query-data.zip"

CLAIM_COLS = ["ClaimID", "EncounterID", "PatientID", "PayerID", "ServiceDate", "SubmitDate", "BilledAmount",
              "AllowedAmount", "PatientResponsibility", "PaidAmount", "ClaimStatus", "DenialReason", "PaidDate"]
ENC_COLS = ["EncounterID", "PatientID", "EncounterType", "FacilityID", "DeptID", "AdmitDateTime",
            "DischargeDateTime", "PrimaryDxCode", "PayerID", "TotalCharges"]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
BUDGET_LABELS = ["FacilityID", "Department", "LineType", "Category", "Measure"]
BUDGET_COLS = BUDGET_LABELS + MONTHS + ["FY Total"]

# Hand-maintained mapping table in the workbook. Two keys are deliberately imperfect (see module doc).
DENIAL_MAP = [
    ("Authorization Required ", "Front End", "Patient Access"),   # trailing space
    ("Eligibility / Coverage", "Front End", "Patient Access"),
    ("Medical Necessity", "Mid-Cycle", "Utilization Review"),
    ("Missing Documentation", "Mid-Cycle", "HIM & Coding"),
    ("Coding Error", "Mid-Cycle", "HIM & Coding"),
    ("Duplicate Claim", "Back End", "Business Office"),
    ("Timely filing", "Back End", "Business Office"),             # lower-case f
]


# ---------------------------------------------------------------------------
# M code (one source of truth for the README answer key, starter/ and solutions/)
# ---------------------------------------------------------------------------
def _types(cols: list[str], kinds: dict[str, str]) -> str:
    return "{" + ", ".join(f'{{"{c}", {kinds.get(c, "type text")}}}' for c in cols) + "}"


CLAIM_TYPES = _types(CLAIM_COLS, {
    "ServiceDate": "type date", "SubmitDate": "type date", "PaidDate": "type date", "BilledAmount": "type number",
    "AllowedAmount": "type number", "PatientResponsibility": "type number", "PaidAmount": "type number"})

M_JAN = f'''// Query: claims_2025_01
// Data > Get Data > From File > From Text/CSV > pick the file > Transform Data (or Load)
let
    Source = Csv.Document(File.Contents("{WIN_DATA}claims_monthly\\claims_2025_01.csv"),[Delimiter=",", Columns=13, Encoding=65001, QuoteStyle=QuoteStyle.None]),
    #"Promoted Headers" = Table.PromoteHeaders(Source, [PromoteAllScalars=true]),
    #"Changed Type" = Table.TransformColumnTypes(#"Promoted Headers",{CLAIM_TYPES})
in
    #"Changed Type"'''

M_APPEND = '''// Import claims_2025_11.csv and claims_2025_12.csv the same way as task 1, then
// Home > Append Queries > Append Queries as New > Two tables > claims_2025_11 + claims_2025_12
// Query: Claims_NovDec
let
    Source = Table.Combine({claims_2025_11, claims_2025_12})
in
    Source'''

M_PARAM = f'''// Parameter: DataFolder   (Home > Manage Parameters > New Parameter)
// Name: DataFolder   Type: Text   Current Value: the folder that holds the lesson CSV files, ending in a separator
"{WIN_DATA}" meta [IsParameterQuery=true, Type="Text", IsParameterQueryRequired=true]'''

CLAIMS_TYPES_WITH_SOURCE = _types(["Source.Name"] + CLAIM_COLS, {
    "ServiceDate": "type date", "SubmitDate": "type date", "PaidDate": "type date", "BilledAmount": "type number",
    "AllowedAmount": "type number", "PatientResponsibility": "type number", "PaidAmount": "type number"})

M_CLAIMS = f'''// Query: Claims   (Data > Get Data > From File > From Folder > pick claims_monthly > Combine & Transform Data,
// choose the first file as the sample, OK, then rename the query from claims_monthly to Claims)
// Excel also creates a "Helper Queries" group (Sample File, Parameter1, Transform Sample File, Transform File).
// The Source line below uses the DataFolder parameter (Guide section 12); the generated code has your full path.
// The helper query Sample File repeats the folder path in its own Source step, so make the same change there:
//     Source = Folder.Files(DataFolder & "claims_monthly"),
// (On a Mac, keep full paths such as "/Users/you/PQ/data/claims_monthly" instead of the parameter.)
// Step names can differ slightly between Excel versions. Optional hardening (Guide section 11): insert
//     #"CSV Only" = Table.SelectRows(Source, each Text.Lower([Extension]) = ".csv"),
// after Source (and make the next step read #"CSV Only") so stray files, such as a Mac .DS_Store file,
// are never combined.
let
    Source = Folder.Files(DataFolder & "claims_monthly"),
    #"Filtered Hidden Files1" = Table.SelectRows(Source, each [Attributes]?[Hidden]? <> true),
    #"Invoke Custom Function1" = Table.AddColumn(#"Filtered Hidden Files1", "Transform File", each #"Transform File"([Content])),
    #"Renamed Columns1" = Table.RenameColumns(#"Invoke Custom Function1", {{"Name", "Source.Name"}}),
    #"Removed Other Columns1" = Table.SelectColumns(#"Renamed Columns1", {{"Source.Name", "Transform File"}}),
    #"Expanded Table Column1" = Table.ExpandTableColumn(#"Removed Other Columns1", "Transform File", Table.ColumnNames(#"Transform File"(#"Sample File"))),
    #"Changed Type" = Table.TransformColumnTypes(#"Expanded Table Column1",{CLAIMS_TYPES_WITH_SOURCE})
in
    #"Changed Type"'''

M_TOTAL_PAID = '''// Right-click Claims > Reference, rename the new query PaidTotal,
// select the PaidAmount column > Transform > Statistics > Sum. The preview shows one number.
let
    Source = Claims,
    #"Calculated Sum" = List.Sum(Source[PaidAmount])
in
    #"Calculated Sum"
// Alternative: load Claims to a sheet and type =SUM(Claims[PaidAmount]) in any empty cell.'''

M_DENIALS_BASIC = '''// Query: DenialsByReason   (right-click Claims > Reference)
let
    Source = Claims,
    #"Filtered Rows" = Table.SelectRows(Source, each ([ClaimStatus] = "Denied")),
    #"Grouped Rows" = Table.Group(#"Filtered Rows", {"DenialReason"}, {{"Claims", each Table.RowCount(_), Int64.Type}}),
    #"Sorted Rows" = Table.Sort(#"Grouped Rows",{{"Claims", Order.Descending}})
in
    #"Sorted Rows"'''

M_DENIALS_ADV = '''// Query: DenialsByReason, Grouped Rows step edited (gear icon > Advanced > Add aggregation)
let
    Source = Claims,
    #"Filtered Rows" = Table.SelectRows(Source, each ([ClaimStatus] = "Denied")),
    #"Grouped Rows" = Table.Group(#"Filtered Rows", {"DenialReason"}, {
        {"Claims", each Table.RowCount(_), Int64.Type},
        {"BilledDenied", each List.Sum([BilledAmount]), type nullable number},
        {"AvgBilled", each List.Average([BilledAmount]), type nullable number}}),
    #"Sorted Rows" = Table.Sort(#"Grouped Rows",{{"AvgBilled", Order.Descending}})
in
    #"Sorted Rows"'''

M_PAYERS = f'''// Query: Payers   (From Text/CSV: payers.csv)
let
    Source = Csv.Document(File.Contents(DataFolder & "payers.csv"),[Delimiter=",", Columns=6, Encoding=65001, QuoteStyle=QuoteStyle.None]),
    #"Promoted Headers" = Table.PromoteHeaders(Source, [PromoteAllScalars=true]),
    #"Changed Type" = Table.TransformColumnTypes(#"Promoted Headers",{{{{"PayerID", type text}}, {{"PayerName", type text}}, {{"PayerType", type text}}, {{"AvgAllowedPctOfCharges", type number}}, {{"AvgDaysToPay", Int64.Type}}, {{"TimelyFilingDays", Int64.Type}}}})
in
    #"Changed Type"'''

M_PAID_BY_TYPE = '''// Query: PaidByPayerType   (right-click Claims > Reference; Home > Merge Queries)
let
    Source = Claims,
    #"Merged Queries" = Table.NestedJoin(Source, {"PayerID"}, Payers, {"PayerID"}, "Payers", JoinKind.LeftOuter),
    #"Expanded Payers" = Table.ExpandTableColumn(#"Merged Queries", "Payers", {"PayerType"}, {"PayerType"}),
    #"Grouped Rows" = Table.Group(#"Expanded Payers", {"PayerType"}, {
        {"Claims", each Table.RowCount(_), Int64.Type},
        {"Paid", each List.Sum([PaidAmount]), type nullable number}})
in
    #"Grouped Rows"'''

M_SUBMIT_LAG = '''// Query: SubmitLag   (right-click Claims > Reference; Add Column > Custom Column)
let
    Source = Claims,
    #"Added Custom" = Table.AddColumn(Source, "DaysToSubmit", each Duration.Days([SubmitDate] - [ServiceDate]), Int64.Type),
    #"Filtered Rows" = Table.SelectRows(#"Added Custom", each [DaysToSubmit] > 30)
in
    #"Filtered Rows"'''

M_DENIAL_MAP = '''// Query: DenialMap   (click inside tblDenialMap > Data > From Table/Range)
let
    Source = Excel.CurrentWorkbook(){[Name="tblDenialMap"]}[Content],
    #"Changed Type" = Table.TransformColumnTypes(Source,{{"DenialReason", type text}, {"RevCycleStage", type text}, {"OwnerTeam", type text}}),
    #"Trimmed Text" = Table.TransformColumns(#"Changed Type",{{"DenialReason", Text.Trim, type text}}),
    #"Capitalized Each Word" = Table.TransformColumns(#"Trimmed Text",{{"DenialReason", Text.Proper, type text}})
in
    #"Capitalized Each Word"

// Query: DenialsByOwner   (right-click Claims > Reference)
let
    Source = Claims,
    #"Filtered Rows" = Table.SelectRows(Source, each ([ClaimStatus] = "Denied")),
    #"Merged Queries" = Table.NestedJoin(#"Filtered Rows", {"DenialReason"}, DenialMap, {"DenialReason"}, "DenialMap", JoinKind.LeftOuter),
    #"Expanded DenialMap" = Table.ExpandTableColumn(#"Merged Queries", "DenialMap", {"OwnerTeam"}, {"OwnerTeam"}),
    #"Grouped Rows" = Table.Group(#"Expanded DenialMap", {"OwnerTeam"}, {
        {"Claims", each Table.RowCount(_), Int64.Type},
        {"BilledDenied", each List.Sum([BilledAmount]), type nullable number}})
in
    #"Grouped Rows"'''

ENC_TYPES = _types(ENC_COLS, {"AdmitDateTime": "type datetime", "DischargeDateTime": "type datetime",
                              "TotalCharges": "type number"})
M_ENCOUNTERS = f'''// Query: Encounters2025   (From Text/CSV: encounters_2025.csv)
let
    Source = Csv.Document(File.Contents(DataFolder & "encounters_2025.csv"),[Delimiter=",", Columns=10, Encoding=65001, QuoteStyle=QuoteStyle.None]),
    #"Promoted Headers" = Table.PromoteHeaders(Source, [PromoteAllScalars=true]),
    #"Changed Type" = Table.TransformColumnTypes(#"Promoted Headers",{ENC_TYPES})
in
    #"Changed Type"'''

M_UNBILLED = '''// Query: Unbilled   (Home > Merge Queries > Merge Queries as New:
// top = Encounters2025, bottom = Claims, click EncounterID in both, Join Kind = Left Anti)
let
    Source = Table.NestedJoin(Encounters2025, {"EncounterID"}, Claims, {"EncounterID"}, "Claims", JoinKind.LeftAnti),
    #"Removed Columns" = Table.RemoveColumns(Source, {"Claims"})
in
    #"Removed Columns"'''

M_AGED_STARTER = '''// AgedPending: Pending claims that have waited too long for payment.
// Paste into Data > Get Data > From Other Sources > Blank Query > Home > Advanced Editor.
// It reads your Claims query (task 3), so that query must exist and be named Claims.
let
    Source = Claims,
    AsOf = #date(2025, 12, 31),
    PendingOnly = Table.SelectRows(Source, each [ClaimStatus] = "Pending"),
    AddDaysPending = Table.AddColumn(PendingOnly, "DaysPending", each Duration.Days(AsOf - [SubmitDate]), Int64.Type),
    Aged = Table.SelectRows(AddDaysPending, each [DaysPending] > 60),
    KeepColumns = Table.SelectColumns(Aged, {"ClaimID", "PayerID", "SubmitDate", "BilledAmount", "DaysPending"}),
    Sorted = Table.Sort(KeepColumns, {{"DaysPending", Order.Descending}})
in
    Sorted'''
M_AGED_SOLUTION = M_AGED_STARTER.replace("[DaysPending] > 60", "[DaysPending] > 90")

BUDGET_TYPES = _types(BUDGET_COLS, {m: "Int64.Type" for m in MONTHS + ["FY Total"]})
M_BUDGET = f'''// Query: Budget2025   (From Text/CSV: budget_2025_wide.csv)
// Select FY Total > Remove Columns; select FacilityID..Measure > Transform > Unpivot Columns > Unpivot Other Columns;
// rename Attribute to Month and Value to Amount
let
    Source = Csv.Document(File.Contents(DataFolder & "budget_2025_wide.csv"),[Delimiter=",", Columns=18, Encoding=65001, QuoteStyle=QuoteStyle.None]),
    #"Promoted Headers" = Table.PromoteHeaders(Source, [PromoteAllScalars=true]),
    #"Changed Type" = Table.TransformColumnTypes(#"Promoted Headers",{BUDGET_TYPES}),
    #"Removed Columns" = Table.RemoveColumns(#"Changed Type",{{"FY Total"}}),
    #"Unpivoted Other Columns" = Table.UnpivotOtherColumns(#"Removed Columns", {{"FacilityID", "Department", "LineType", "Category", "Measure"}}, "Attribute", "Value"),
    #"Renamed Columns" = Table.RenameColumns(#"Unpivoted Other Columns",{{{{"Attribute", "Month"}}, {{"Value", "Amount"}}}})
in
    #"Renamed Columns"'''

_SPLIT_STEP = ('    #"Split Column by Delimiter" = Table.SplitColumn(#"Renamed Columns", "Department", '
               'Splitter.SplitTextByEachDelimiter({" - "}, QuoteStyle.Csv, false), {"CostCenter", "DeptName"})')
_BUDGET_TAIL = '\nin\n    #"Renamed Columns"'
assert M_BUDGET.endswith(_BUDGET_TAIL)
M_ICU = (
    '// Query: Budget2025 after task 13: the task 12 query plus one step at the end. Select Department >\n'
    '// Transform > Split Column > By Delimiter > --Custom-- " - " (space hyphen space) > Split at: Left-most delimiter.\n'
    '// The dialog names the parts Department.1 and Department.2 (and adds a Changed Type step); typing the final\n'
    '// names in the step, as here, splits and renames at once.\n'
    + M_BUDGET.split("\n", 3)[3][: -len(_BUDGET_TAIL)]
    + ",\n" + _SPLIT_STEP + '\nin\n    #"Split Column by Delimiter"'
    + '''

// Query: ICU_Q1_Salaries   (right-click Budget2025 > Reference)
let
    Source = Budget2025,
    #"Filtered Rows" = Table.SelectRows(Source, each [DeptName] = "Intensive Care Unit"
        and [Category] = "Salaries & Wages" and [Measure] = "Actual"
        and List.Contains({"Jan", "Feb", "Mar"}, [Month])),
    #"Grouped Rows" = Table.Group(#"Filtered Rows", {"DeptName"}, {{"Q1Actual", each List.Sum([Amount]), type nullable number}})
in
    #"Grouped Rows"''')

M_DASHBOARD = '''// Query: DenialDashboard   (right-click Claims > Reference)
let
    Source = Claims,
    #"Filtered Rows" = Table.SelectRows(Source, each ([ClaimStatus] = "Denied")),
    #"Merged Queries" = Table.NestedJoin(#"Filtered Rows", {"PayerID"}, Payers, {"PayerID"}, "Payers", JoinKind.LeftOuter),
    #"Expanded Payers" = Table.ExpandTableColumn(#"Merged Queries", "Payers", {"PayerType"}, {"PayerType"}),
    #"Grouped Rows" = Table.Group(#"Expanded Payers", {"PayerType", "DenialReason"}, {
        {"DeniedClaims", each Table.RowCount(_), Int64.Type},
        {"DeniedBilled", each List.Sum([BilledAmount]), type nullable number}}),
    #"Sorted Rows" = Table.Sort(#"Grouped Rows",{{"DeniedBilled", Order.Descending}})
in
    #"Sorted Rows"'''

M_RATE = '''// Query: DenialRateByPayerType   (right-click Claims > Reference)
let
    Source = Claims,
    #"Merged Queries" = Table.NestedJoin(Source, {"PayerID"}, Payers, {"PayerID"}, "Payers", JoinKind.LeftOuter),
    #"Expanded Payers" = Table.ExpandTableColumn(#"Merged Queries", "Payers", {"PayerType"}, {"PayerType"}),
    #"Added Conditional Column" = Table.AddColumn(#"Expanded Payers", "IsDenied", each if [ClaimStatus] = "Denied" then 1 else 0, Int64.Type),
    #"Grouped Rows" = Table.Group(#"Added Conditional Column", {"PayerType"}, {
        {"AllClaims", each Table.RowCount(_), Int64.Type},
        {"Denied", each List.Sum([IsDenied]), type nullable number}}),
    #"Added Custom" = Table.AddColumn(#"Grouped Rows", "DenialRate", each [Denied] / [AllClaims], Percentage.Type),
    #"Sorted Rows" = Table.Sort(#"Added Custom",{{"DenialRate", Order.Descending}})
in
    #"Sorted Rows"'''


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------
def _csv_text(header: list[str], rows: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    for r in rows:
        w.writerow(["" if r.get(c) is None else r[c] for c in header])
    text = buf.getvalue()
    assert text.isascii(), "lesson CSVs should be plain ASCII"
    return text


def _parse(text: str) -> list[dict]:
    """Read a generated CSV back the way Power Query sees it: every field is text, empty = ''."""
    return list(csv.DictReader(io.StringIO(text)))


def _budget_wide() -> list[dict]:
    depts = data.index(data.load("departments"), "DeptID")
    lines: dict[tuple, dict] = {}
    for r in data.load("budget"):
        if r["FiscalYear"] != 2025:
            continue
        key = (r["FacilityID"], r["CostCenter"], r["DeptID"], r["LineType"], r["Category"])
        lines.setdefault(key, {})[r["FiscalMonth"]] = (r["BudgetAmount"], r["ActualAmount"])
    out = []
    line_order = {"Revenue": 0, "Expense": 1}
    for key in sorted(lines, key=lambda k: (k[0], k[1], line_order[k[3]], k[4])):
        fac, cc, dept, line, cat = key
        months = lines[key]
        assert sorted(months) == list(range(1, 13)), key
        assert depts[dept]["CostCenter"] == cc
        for mi, measure in enumerate(("Budget", "Actual")):
            row = {"FacilityID": fac, "Department": f"{cc} - {depts[dept]['DeptName']}", "LineType": line,
                   "Category": cat, "Measure": measure}
            for m in range(1, 13):
                row[MONTHS[m - 1]] = months[m][mi]
            row["FY Total"] = sum(months[m][mi] for m in range(1, 13))
            out.append(row)
    return out


def build_files() -> dict[str, str]:
    """Return {relative path: CSV text} for everything in the lesson's data/ folder (except the zip)."""
    claims = data.load_raw("claims")
    files: dict[str, str] = {}
    for m in range(1, 13):
        month = [r for r in claims if r["SubmitDate"].startswith(f"2025-{m:02d}-")]
        month.sort(key=lambda r: r["ClaimID"])
        files[f"data/claims_monthly/claims_2025_{m:02d}.csv"] = _csv_text(CLAIM_COLS, month)
    jan26 = sorted((r for r in claims if r["SubmitDate"].startswith("2026-01-")), key=lambda r: r["ClaimID"])
    files["data/new_month/claims_2026_01.csv"] = _csv_text(CLAIM_COLS, jan26)
    files["data/payers.csv"] = _csv_text(data.columns("payers"), data.load_raw("payers"))
    enc = [r for r in data.load_raw("encounters") if r["DischargeDateTime"].startswith("2025-")]
    enc.sort(key=lambda r: r["EncounterID"])
    files["data/encounters_2025.csv"] = _csv_text(ENC_COLS, enc)
    files["data/budget_2025_wide.csv"] = _csv_text(BUDGET_COLS, _budget_wide())
    return files


def _zip(files: dict[str, str]) -> bytes:
    """A deterministic zip (fixed timestamps, sorted names) so rebuilding doesn't change the bytes."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name in sorted(files):
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            z.writestr(info, files[name])
    return buf.getvalue()


def _m_file(title: str, body: str) -> str:
    return f"// Lesson 4.3 · {title}\n// SPOILER: reference solution. Try the task first.\n\n{body}\n"


# ---------------------------------------------------------------------------
# Replaying the queries in Python
# ---------------------------------------------------------------------------
def _d(s: str) -> date | None:
    return date.fromisoformat(s) if s else None


def _money(values) -> float:
    return float(sum((Decimal(v) for v in values), Decimal("0")))


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="04-advanced-analysis", slug="03-power-query",
        title="Power Query: Import, Transform & Combine", level="Advanced", minutes=70,
        objectives=[
            "Import CSV files and whole folders with Get & Transform",
            "Clean and reshape data with Applied Steps: types, splits, filters, unpivot, group by",
            "Merge (join) and append queries",
            "Build refreshable, repeatable data pipelines and read the M code behind them",
        ],
        data_note="CSV exports in the lesson's data folder: twelve monthly files of claims submitted in 2025, a January "
                  "2026 claim file (simulated next-month export) for the bonus, the payer list, 11,196 encounters discharged in 2025, and a wide "
                  "2025 budget-vs-actual export. The workbook adds a hand-maintained denial-reason mapping table.",
    )

    files = build_files()
    parsed = {k: _parse(v) for k, v in files.items()}

    # ------------------------------------------------------------------ replay the queries
    monthly = [parsed[f"data/claims_monthly/claims_2025_{m:02d}.csv"] for m in range(1, 13)]
    claims = [r for month in monthly for r in month]                 # the Claims folder query
    jan26 = parsed["data/new_month/claims_2026_01.csv"]
    payers = {p["PayerID"]: p for p in parsed["data/payers.csv"]}
    enc25 = parsed["data/encounters_2025.csv"]
    budget = parsed["data/budget_2025_wide.csv"]

    # cross-check the files against the typed loaders (catches any file-writing slip)
    typed = data.load("claims")
    assert len(claims) == sum(1 for c in typed if c["SubmitDate"].year == 2025)
    assert abs(_money(r["PaidAmount"] for r in claims)
               - sum(c["PaidAmount"] for c in typed if c["SubmitDate"].year == 2025)) < 0.01
    assert len(enc25) == sum(1 for e in data.load("encounters") if e["DischargeDateTime"].year == 2025)
    assert all(len(r) == len(CLAIM_COLS) for r in claims)

    n_jan = len(monthly[0])
    n_novdec = len(monthly[10]) + len(monthly[11])
    n_claims = len(claims)
    total_paid = round(_money(r["PaidAmount"] for r in claims), 2)

    denied = [r for r in claims if r["ClaimStatus"] == "Denied"]
    by_reason = defaultdict(list)
    for r in denied:
        by_reason[r["DenialReason"]].append(Decimal(r["BilledAmount"]))
    assert "" not in by_reason, "every denied claim should carry a reason"
    auth_count = len(by_reason["Authorization Required"])
    avg_rank = sorted(by_reason, key=lambda k: -(sum(by_reason[k]) / len(by_reason[k])))
    top_avg_reason = avg_rank[0]
    top_avg = sum(by_reason[top_avg_reason]) / len(by_reason[top_avg_reason])
    second_avg = sum(by_reason[avg_rank[1]]) / len(by_reason[avg_rank[1]])
    assert top_avg - second_avg > 100, "average billed per reason: top two too close"
    top_count_reason = max(by_reason, key=lambda k: len(by_reason[k]))

    paid_by_type = defaultdict(list)
    for r in claims:
        paid_by_type[payers[r["PayerID"]]["PayerType"]].append(r["PaidAmount"])
    commercial_paid = round(_money(paid_by_type["Commercial"]), 2)

    lag = [(_d(r["SubmitDate"]) - _d(r["ServiceDate"])).days for r in claims]
    late_submits = sum(1 for x in lag if x > 30)

    # Denial map: an exact (un-cleaned) merge misses two reasons; Trim + Proper fixes both.
    raw_map = {k: team for k, _, team in DENIAL_MAP}
    clean_map = {k.strip().title(): team for k, _, team in DENIAL_MAP}
    assert all(k.title() == k for k in by_reason), "claim reasons must already be in Proper case"
    assert set(clean_map) == set(by_reason), "mapping table must cover every denial reason once cleaned"
    unmatched_raw = sum(1 for r in denied if r["DenialReason"] not in raw_map)
    assert unmatched_raw > 0
    pa_billed = round(_money(r["BilledAmount"] for r in denied if clean_map[r["DenialReason"]] == "Patient Access"), 2)
    pa_billed_raw = round(_money(r["BilledAmount"] for r in denied if raw_map.get(r["DenialReason"]) == "Patient Access"), 2)
    assert abs(pa_billed - pa_billed_raw) > 1

    billed_ids = {r["EncounterID"] for r in claims}
    unbilled = [e for e in enc25 if e["EncounterID"] not in billed_ids]
    n_unbilled = len(unbilled)

    pending = [r for r in claims if r["ClaimStatus"] == "Pending"]
    aged60 = sum(1 for r in pending if (AS_OF - _d(r["SubmitDate"])).days > 60)
    aged90 = sum(1 for r in pending if (AS_OF - _d(r["SubmitDate"])).days > 90)

    n_unpivot = len(budget) * len(MONTHS)     # no blank month cells, so the unpivot keeps every value
    assert all(r[m] != "" for r in budget for m in MONTHS)
    icu = [r for r in budget if r["Department"].split(" - ", 1)[1] == "Intensive Care Unit"
           and r["Category"] == "Salaries & Wages" and r["Measure"] == "Actual"]
    assert len({r["Department"] for r in icu}) == 3
    icu_q1 = sum(int(r[m]) for r in icu for m in ("Jan", "Feb", "Mar"))
    assert all(int(r["FY Total"]) == sum(int(r[m]) for m in MONTHS) for r in budget)
    n_dept_names = len({r["Department"].split(" - ", 1)[1] for r in budget})
    n_depts = len({r["Department"] for r in budget})
    n_commercial_payers = sum(1 for p in payers.values() if p["PayerType"] == "Commercial")

    # Bonus: dashboard before and after the January 2026 file
    def dashboard(rows):
        g = defaultdict(list)
        for r in rows:
            if r["ClaimStatus"] == "Denied":
                g[(payers[r["PayerID"]]["PayerType"], r["DenialReason"])].append(r["BilledAmount"])
        return {k: round(_money(v), 2) for k, v in g.items()}

    def rates(rows):
        n, d = Counter(), Counter()
        for r in rows:
            t = payers[r["PayerID"]]["PayerType"]
            n[t] += 1
            d[t] += r["ClaimStatus"] == "Denied"
        return {t: d[t] / n[t] for t in n}

    dash_before = dashboard(claims)
    dash_after = dashboard(claims + jan26)
    combo = ("Government", "Authorization Required")
    assert max(dash_before, key=dash_before.get) == combo
    rate_before = rates(claims)
    top_rate_type = max(rate_before, key=rate_before.get)
    assert top_rate_type == "Medicare Advantage", "the B2 explanation discusses Medicare Advantage plans"
    ranked_rates = sorted(rate_before.values(), reverse=True)
    assert ranked_rates[0] - ranked_rates[1] > 0.01
    billed_after = {r["EncounterID"] for r in claims + jan26}
    unbilled_after = sum(1 for e in enc25 if e["EncounterID"] not in billed_after)
    assert dash_after[combo] != dash_before[combo]
    denied_after = sum(1 for r in claims + jan26 if r["ClaimStatus"] == "Denied")
    # Where the still-unbilled encounters end up in the full course claims file (for the B5 explanation)
    billed_later = {r["EncounterID"]: r["SubmitDate"] for r in data.load_raw("claims")}
    still_unbilled = [e["EncounterID"] for e in enc25 if e["EncounterID"] not in billed_after]
    later_months = sorted({billed_later[i][:7] for i in still_unbilled if i in billed_later})
    assert len(later_months) >= 2 and all(i in billed_later for i in still_unbilled)
    later_span = (f"{datetime.strptime(later_months[0], '%Y-%m'):%B} and "
                  f"{datetime.strptime(later_months[-1], '%Y-%m'):%B %Y}")
    rate_top = rate_before[max(rate_before, key=rate_before.get)]

    # ------------------------------------------------------------------ workbook sheets
    file_rows = [
        {"File": "data/claims_monthly/claims_2025_01.csv … claims_2025_12.csv", "Rows": None,
         "Contents": "One file per month of claims SUBMITTED in 2025 (12 files, 13 columns each, same layout as the "
                     "course's claims.csv). Used from task 1 on."},
        {"File": "data/new_month/claims_2026_01.csv", "Rows": len(jan26),
         "Contents": "Simulated next-month export: claims submitted in January 2026 for services in late 2025 (the "
                     "course data is otherwise as of 12/31/2025). Bonus only: keep it out of claims_monthly until the "
                     "bonus tells you to copy it in."},
        {"File": "data/payers.csv", "Rows": len(payers), "Contents": "Payer list: PayerID, PayerName, PayerType, contract terms."},
        {"File": "data/encounters_2025.csv", "Rows": len(enc25),
         "Contents": "Encounters DISCHARGED in 2025 (10 columns: IDs, type, facility, department, admit/discharge, "
                     "diagnosis, payer, charges)."},
        {"File": "data/budget_2025_wide.csv", "Rows": None,
         "Contents": "2025 budget and actual by department and category: one row per Budget/Actual line, one column "
                     "per month (Jan–Dec) plus an FY Total column."},
        {"File": "data/4.3-power-query-data.zip", "Rows": None,
         "Contents": "All of the files above in one download. Extract it (for example to C:\\PQ) before you import."},
    ]
    L.add_table_sheet("Data Files", file_rows, columns=["File", "Rows", "Contents"], as_table=False,
                      formats={"Rows": "#,##0"}, widths={"File": 58, "Rows": 9, "Contents": 100}, start_row=4,
                      notes=["Lesson 4.3 data files (CSV, in the lesson's data folder)",
                             "Power Query reads these files from disk. The workbook doesn't contain them, so download "
                             "and extract the data folder first (see Guide section 2).",
                             "Tip: copy the folder to a short path such as C:\\PQ\\data (Mac: /Users/you/PQ/data)."])
    map_rows = [{"DenialReason": k, "RevCycleStage": s, "OwnerTeam": t} for k, s, t in DENIAL_MAP]
    L.add_table_sheet("DenialMap", map_rows, table="tblDenialMap", columns=["DenialReason", "RevCycleStage", "OwnerTeam"],
                      widths={"DenialReason": 28, "RevCycleStage": 16, "OwnerTeam": 22}, start_row=4,
                      notes=["Denial reason → owner team (maintained by hand by the Revenue Integrity team)",
                             "Load this Table into Power Query with Data → From Table/Range (task 9).",
                             "Each denial reason should appear once. OwnerTeam is the department that fixes that kind of denial."])

    L.sheet_order = ["Start Here", "Practice", "Data Files", "DenialMap", "Starter M", "Bonus", "Answer Key", "Bonus Key"]
    L.start_notes = [
        "This lesson's data lives in CSV files, not in this workbook. Download the lesson's data folder (or the single "
        "file data/4.3-power-query-data.zip) and extract it before you start. The 'Data Files' sheet lists every file.",
        "Build your queries in THIS workbook so tasks 9 and 11 can use the DenialMap and Starter M sheets. Power Query "
        "needs desktop Excel: Microsoft 365 or Excel 2016+ on Windows, or Microsoft 365 for Mac (see the guide's version notes).",
        "Do the practice before the bonus. The bonus adds a January 2026 file to the claims folder, which changes several "
        "practice totals.",
    ]

    @L.customize
    def tidy_files_sheet(wb, lesson, selftest):
        ws = wb["Data Files"]
        for row in ws.iter_rows(min_row=5, max_row=4 + len(file_rows)):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            longest = max(-(-len(str(row[0].value or "")) // 55), -(-len(str(row[2].value or "")) // 85))
            ws.row_dimensions[row[0].row].height = 15 * max(2, longest) + 2
        ws.column_dimensions["C"].width = 80
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    @L.customize
    def starter_sheet(wb, lesson, selftest):
        ws = wb.create_sheet("Starter M")
        ws.sheet_properties.tabColor = "7030A0"
        ws["A1"] = "Task 11 · Starter M code for the AgedPending query"
        ws["A1"].font = Font(bold=True, size=14, color=NAVY)
        ws["A2"] = ("Select A4:A{0}, copy (Ctrl + C, Mac: ⌘ + C), then paste into Home → Advanced Editor of a new Blank Query, "
                    "replacing everything there. The same code is in starter/AgedPending.m.").format(3 + len(M_AGED_STARTER.splitlines()))
        ws["A2"].font = Font(italic=True, color="595959")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 32
        code_fill = PatternFill("solid", fgColor="F2F2F2")
        for i, line in enumerate(M_AGED_STARTER.splitlines(), 4):
            c = ws.cell(row=i, column=1, value=line)
            c.data_type = "s"
            c.font = Font(name="Consolas", size=10, color="1F3864")
            c.fill = code_fill
            if line.startswith("="):
                c._style.quotePrefix = 1
        ws.column_dimensions["A"].width = 135

    # ------------------------------------------------------------------ tasks
    L.practice_intro = (
        "Build every query in this workbook from the CSV files in the lesson's data folder (Guide section 2), and do the tasks "
        "in order because later tasks reuse earlier queries. Type each result in the yellow cell as a plain number or text. "
        "From task 3 on, the answer key's M code uses a DataFolder parameter (Guide section 12). If you skip the parameter, "
        "or work on a Mac, your code shows your full folder path in its place.")
    L.tasks = [
        Task("Import claims_2025_01.csv from the claims_monthly folder with Data → Get Data → From File → From Text/CSV and "
             "load it to a new sheet. How many claims (rows, not counting the header) does the January 2025 file contain?",
             answer=n_jan, solution=M_JAN, solution_lang="m", live=False,
             hint="After loading, the Queries & Connections pane says 'N rows loaded'",
             title="Import one CSV file (January 2025)",
             explanation="From Text/CSV writes three steps for you. **Source** reads the file with `Csv.Document`, "
                         "**Promoted Headers** turns the first line into column names, and **Changed Type** sets each column's "
                         "data type from the first 200 rows. The status bar inside the editor only counts the rows in the "
                         "preview, so read the total from the Queries & Connections pane (or from the loaded Table) after "
                         "**Close & Load**."),
        Task("Import claims_2025_11.csv and claims_2025_12.csv as two more queries. Then, in the Power Query Editor, stack "
             "them with Home → Append Queries → Append Queries as New. How many rows does the appended query return?",
             answer=n_novdec, solution=M_APPEND, solution_lang="m", live=False,
             hint="Append stacks rows, and columns line up by name",
             title="Append two monthly files",
             explanation="Appending stacks the rows of one query under another (`Table.Combine`). It matches columns by "
                         "**name**, not position, so files with the same headers line up even if the column order differs. "
                         "The result has November's rows plus December's rows. Appending works for two or three files, but "
                         "importing twelve files one by one doesn't scale, which is what task 3 fixes."),
        Task("Now combine all twelve 2025 files at once. Choose Data → Get Data → From File → From Folder, select the "
             "claims_monthly folder, then Combine & Transform Data. Rename the new query Claims. How many rows does Claims return?",
             answer=n_claims, solution=M_CLAIMS, solution_lang="m", live=False,
             hint="One folder query replaces twelve imports",
             title="Combine a folder of files (Claims)",
             explanation="**From Folder** lists every file in the folder (and its subfolders), runs the same transformation "
                         "on each one through the helper function *Transform File*, and appends the results. It also adds a "
                         "**Source.Name** column with each row's file name, which is handy for tracing a row back to its "
                         "export. Next month you drop a new file into the folder and click Refresh, with no new query needed."),
        Task("What is the total PaidAmount across all 2025 claims in your Claims query? Check that PaidAmount has the "
             "Decimal Number type (1.2 icon), then enter the total to the cent.",
             answer=total_paid, fmt="#,##0.00", solution=M_TOTAL_PAID, solution_lang="m", live=False,
             hint="Reference Claims, then Transform → Statistics → Sum on PaidAmount",
             title="Total PaidAmount in Claims",
             explanation="**Reference** creates a new query whose source is the output of Claims, so you can summarize "
                         "without changing Claims itself. **Statistics → Sum** turns the query into a single number "
                         "(`List.Sum`). Never add that step to Claims, because every query built on Claims would then receive a "
                         "number instead of a table. If PaidAmount were still Text (ABC icon), Sum would be grayed out."),
        Task("Reference Claims, keep only rows whose ClaimStatus is Denied, and use Home → Group By on DenialReason with "
             "the operation Count Rows. How many denied claims list Authorization Required as the reason?",
             answer=auth_count, solution=M_DENIALS_BASIC, solution_lang="m", live=False,
             hint="Filter first, then Group By (Basic)",
             title="Group denied claims by reason",
             explanation="Group By collapses the rows into one row per DenialReason and counts the rows in each group, "
                         "like a PivotTable, but the result is a table that refreshes with the data. The filter must come "
                         "**before** the Group By step: Applied Steps run top to bottom, and each step works on the output "
                         f"of the step above it. {top_count_reason} is the most common reason."),
        Task("Edit that Group By step (gear icon → Advanced) and add two aggregations of BilledAmount: Sum and Average. "
             "Which denial reason has the highest AVERAGE BilledAmount per denied claim? Type the reason exactly as it appears.",
             answer=top_avg_reason, solution=M_DENIALS_ADV, solution_lang="m", live=False,
             hint="Group By → Advanced → Add aggregation",
             title="Group By with several aggregations",
             explanation=f"Advanced Group By returns several summaries per group in one pass. {top_count_reason} has the "
                         f"most denials, but {top_avg_reason} denials average about ${float(top_avg):,.0f} of billed charges "
                         "each, the highest of any reason. Volume and dollars tell different stories, which is why denial "
                         "reports show both counts and amounts."),
        Task("Import payers.csv as a query named Payers. Then reference Claims, merge it with Payers on PayerID (Join Kind: "
             "Left Outer), expand only PayerType, and group by PayerType with Sum of PaidAmount. What was the total "
             "PaidAmount for the Commercial payer type? Enter it to the cent.",
             answer=commercial_paid, fmt="#,##0.00", solution=f"{M_PAYERS}\n\n{M_PAID_BY_TYPE}", solution_lang="m", live=False,
             hint="Home → Merge Queries, then the expand button (two arrows) in the new column's header",
             title="Merge Claims with Payers (PaidAmount by PayerType)",
             explanation="A merge is Power Query's lookup. **Left Outer** keeps every claim and brings in the matching payer "
                         "row as a nested table, and the expand button pulls out just the columns you need. Commercial "
                         f"combines {n_commercial_payers} different payers, so you can't answer this by PayerID alone. You need the PayerType "
                         "from the lookup table and then a Group By. Uncheck *Use original column name as prefix* when you "
                         "expand, or the column is named Payers.PayerType."),
        Task("Reference Claims and add a custom column DaysToSubmit that holds the number of days from ServiceDate to "
             "SubmitDate. How many 2025 claims took MORE than 30 days to submit?",
             answer=late_submits, solution=M_SUBMIT_LAG, solution_lang="m", live=False,
             hint="Add Column → Custom Column, then Duration.Days of the date difference",
             title="Custom column: days from service to submission",
             explanation="Subtracting one date from another in M gives a **duration**, not a number. `Duration.Days` "
                         "turns it into whole days. You can also build it without typing: select SubmitDate, Ctrl-click "
                         "ServiceDate, then Add Column → Date → Subtract Days (the order you click sets which date comes "
                         "first). Claims billed more than 30 days after service delay cash and risk timely-filing denials."),
        Task("Load tblDenialMap (on the DenialMap sheet) with Data → From Table/Range and name the query DenialMap. Reference Claims, keep the Denied rows, "
             "merge them with the map on DenialReason (Left Outer), and expand OwnerTeam. Check that EVERY denied claim "
             "found a match (no null OwnerTeam), and fix the keys in the map query if some didn't. What is the total denied "
             "BilledAmount owned by Patient Access? Enter it to the cent.",
             answer=pa_billed, fmt="#,##0.00", solution=M_DENIAL_MAP, solution_lang="m", live=False,
             hint="Merges match text exactly: look at Transform → Format → Trim and Capitalize Each Word",
             title="Merge with an Excel Table (denials owned by Patient Access)",
             explanation="Power Query merges are **exact and case-sensitive**. The hand-typed map has *Authorization "
                         "Required* with a trailing space and *Timely filing* with a lower-case f, so a plain merge leaves "
                         f"{unmatched_raw} denied claims with a null OwnerTeam and undercounts Patient Access. Trim and "
                         "Capitalize Each Word in the DenialMap query fix both keys. (Fixing the cells on the DenialMap "
                         "sheet and refreshing works too.) Always check a merge by filtering the new column for null."),
        Task("Import encounters_2025.csv (one row per encounter discharged in 2025) as a query named Encounters2025. Use "
             "Merge Queries as New with Encounters2025 on top, Claims below, EncounterID in both, and Join Kind Left Anti. "
             "How many 2025 encounters have no claim in the 2025 claim files?",
             answer=n_unbilled, solution=f"{M_ENCOUNTERS}\n\n{M_UNBILLED}", solution_lang="m", live=False,
             hint="Left Anti = rows only in the first (top) table",
             title="Left Anti merge: encounters with no claim",
             explanation="A **Left Anti** join keeps only the rows in the first table that have **no** match in the "
                         "second. It's the fastest way to answer \"what's missing?\" questions. These encounters were "
                         "discharged but not yet billed (hospitals call this *discharged not final billed*, or DNFB), "
                         f"and together they carry ${float(_money(e['TotalCharges'] for e in unbilled)):,.2f} of charges. "
                         "Name the query Unbilled, because the bonus refreshes it."),
        Task("Create a blank query (Data → Get Data → From Other Sources → Blank Query), name it AgedPending, open Home → "
             "Advanced Editor, and replace its contents with the starter code on the Starter M sheet (also in "
             "starter/AgedPending.m). It lists Pending claims more than 60 days old as of 12/31/2025. Read the code, change "
             "it to more than 90 days, and report how many Pending claims are more than 90 days old.",
             answer=aged90, solution=M_AGED_SOLUTION, solution_lang="m", live=False,
             hint="Find the step that compares DaysPending with 60",
             title="Read and edit M code (aged pending claims)",
             explanation="Each line between `let` and `in` is one Applied Step: a name, an equals sign, and an "
                         "expression that usually uses the step above it. `AsOf` is a step that just holds a date, and "
                         "the **Aged** step is the filter, so changing `> 60` to `> 90` is the whole edit. With 60 days the "
                         f"query returns {aged60} claims, and with 90 it returns {aged90}. Claims pending that long need a "
                         "follow-up call to the payer."),
        Task("Import budget_2025_wide.csv as a query named Budget2025. Remove the FY Total column, select the five label "
             "columns (FacilityID, Department, LineType, Category, Measure), and choose Transform → Unpivot Columns → "
             "Unpivot Other Columns. Rename Attribute to Month and Value to Amount. How many rows does Budget2025 return?",
             answer=n_unpivot, solution=M_BUDGET, solution_lang="m", live=False,
             hint="Unpivot Other Columns keeps the selected columns and unpivots the rest",
             title="Unpivot the wide budget",
             explanation=f"The wide file has {len(budget)} rows with a column per month. Unpivoting turns each month cell "
                         f"into its own row, so you get {len(budget)} × 12 = {n_unpivot:,} rows with a Month and an Amount "
                         "column. That long shape is what PivotTables, Group By, and the Data Model want. FY Total has to go "
                         "first, or it would be unpivoted as a thirteenth 'month' and double every annual total. "
                         "**Unpivot Other Columns** records the five label columns to keep (`Table.UnpivotOtherColumns`), "
                         "so a future file with extra month columns still unpivots correctly."),
        Task("In Budget2025, split Department (for example '6130 - Intensive Care Unit') into CostCenter and DeptName with "
             "Transform → Split Column → By Delimiter, using ' - ' (space, hyphen, space) at the left-most delimiter. "
             "What was the Q1 2025 (Jan–Mar) Actual Salaries & Wages for all three departments named Intensive Care Unit "
             "combined? Enter whole dollars.",
             answer=icu_q1, solution=M_ICU, solution_lang="m", live=False,
             hint="Filter DeptName, Category, Measure, and Month, then Group By or Sum",
             title="Split a column, then filter and sum (ICU Q1 salaries)",
             explanation="Splitting at ' - ' (with the spaces) separates the cost center from the name without breaking "
                         "names that contain a plain hyphen, such as *Medical-Surgical*. After the split, all three hospitals' "
                         f"ICUs share the DeptName *Intensive Care Unit* (the file has {n_dept_names} distinct department "
                         f"names for {n_depts} departments), so one filter catches all three. The unpivoted Month column makes "
                         "\"Q1\" a simple filter on Jan, Feb, and Mar."),
    ]

    # ------------------------------------------------------------------ bonus
    L.bonus_title = "Bonus: A refreshable denial dashboard"
    L.bonus_scenario = (
        "It's the first week of February 2026 and the CFO wants a denial dashboard she can refresh every month without anyone "
        "rebuilding it. Build it on top of your Claims query. First, make the pipeline portable: if you haven't yet, create the "
        "DataFolder parameter (Guide section 12) and use it in the Source steps of Claims and its Sample File helper query "
        "(on a Mac, keep the full paths instead). Answer B1–B3 BEFORE you add the January 2026 file, then follow B4 and B5.")
    L.bonus = [
        Task("Build a query named DenialDashboard: reference Claims, keep Denied claims, merge Payers to get PayerType, and "
             "group by BOTH PayerType and DenialReason with Count Rows and Sum of BilledAmount. What is the denied BilledAmount "
             "for Government + Authorization Required? Enter it to the cent.",
             answer=dash_before[combo], fmt="#,##0.00", solution=M_DASHBOARD, solution_lang="m", live=False,
             hint="Ctrl-click two columns in the Group By dialog (Advanced)",
             title="DenialDashboard by PayerType × DenialReason",
             explanation=f"Grouping by two columns gives one row per combination ({len(dash_before)} rows here). "
                         "Government + Authorization Required is the largest pool of denied dollars, so authorization work "
                         "for Government-insured patients (Medicare and State Medicaid in this data) is where the "
                         "revenue-cycle team should start."),
        Task("Build DenialRateByPayerType from ALL claims: reference Claims, merge Payers for PayerType, add a conditional "
             "column IsDenied (1 if ClaimStatus is Denied, otherwise 0), then group by PayerType with Count Rows and Sum of "
             "IsDenied. Which payer type has the highest denial rate (denied claims ÷ all claims)?",
             answer=top_rate_type, solution=M_RATE, solution_lang="m", live=False,
             hint="Add Column → Conditional Column, then a custom column Denied / AllClaims",
             title="Denial rate by payer type",
             explanation="A rate needs two counts per group: all claims (Count Rows) and denied claims. Summing a 1/0 flag "
                         "counts the rows where the condition is true, the same trick as SUMPRODUCT with booleans. "
                         "Dividing in a custom column after the Group By gives the rate. In this data, "
                         f"{top_rate_type} denies about 1 claim in {round(1 / rate_top)}, the highest rate of any payer "
                         "type, which mirrors what many US health systems report about Medicare Advantage plans."),
        Task("What is that payer type's denial rate? Enter it as a percentage rounded to 1 decimal place.",
             answer=rate_before[top_rate_type], fmt="0.0%", live=False,
             solution=("Use the DenialRateByPayerType query from B2 and read the DenialRate value on the "
                       f"{top_rate_type} row. The Percentage type displays it as a percentage."),
             hint="Format the DenialRate column as Percentage",
             title="Highest denial rate (value)",
             explanation=f"{top_rate_type}: {rate_top * 100:.1f}% of its claims were denied. Typing "
                         f"{rate_top * 100:.1f} or {rate_top * 100:.1f}% both pass the check."),
        Task("Now copy data\\new_month\\claims_2026_01.csv into the claims_monthly folder and click Data → Refresh All. You "
             "don't edit any query. What is the denied BilledAmount for Government + Authorization Required now? Enter it to "
             "the cent.",
             answer=dash_after[combo], fmt="#,##0.00", solution=(
                 "1. In File Explorer (Mac: Finder), copy **claims_2026_01.csv** from `data\\new_month` into "
                 "`data\\claims_monthly`.\n"
                 "2. In Excel, click **Data → Refresh All** (Windows: Ctrl + Alt + F5).\n"
                 "3. Read the Government / Authorization Required row of the loaded DenialDashboard table."),
             live=False, title="Drop in the January 2026 file and Refresh All",
             hint="Every query built on Claims re-reads the folder when you refresh",
             explanation="Claims reads whatever files are in the folder, so the thirteenth file flows through every query "
                         f"that references it: Claims now returns {n_claims + len(jan26):,} rows and the dashboard counts "
                         f"{denied_after} denied claims instead of {len(denied)}. This is the payoff of a folder-based "
                         "pipeline. If Claims had a filter that kept only 2025 SubmitDates, the new file would have been "
                         "silently ignored, so keep date filters out of the base query."),
        Task("Your Unbilled query (task 10) refreshed too. How many 2025 encounters are still unbilled now that the "
             "January 2026 claims are in?",
             answer=unbilled_after, solution="Read the row count of the refreshed **Unbilled** query in the Queries & "
                                             "Connections pane (or `=ROWS(Unbilled)` if you loaded it as a Table named Unbilled).",
             live=False, title="Unbilled encounters after the refresh",
             hint="Look at the Queries & Connections pane after the refresh",
             explanation=f"The January 2026 file billed {n_unbilled - unbilled_after} of the {n_unbilled} encounters on the "
                         f"unbilled list, so {unbilled_after} remain. (In the course's full claims data, their claims go "
                         f"out between {later_span}, so later monthly files keep shrinking the list.) Because the Left "
                         "Anti merge points at Claims, the unbilled list maintains itself as new files arrive. No one has "
                         "to rerun a lookup."),
    ]

    # ------------------------------------------------------------------ files next to the workbook
    extra: dict[str, str | bytes] = dict(files)
    extra[ZIP_NAME] = _zip(files)
    extra["starter/AgedPending.m"] = M_AGED_STARTER + "\n"
    sol = {
        "DataFolder": M_PARAM, "claims_2025_01": M_JAN, "Claims_NovDec": M_APPEND, "Claims": M_CLAIMS,
        "PaidTotal": M_TOTAL_PAID, "DenialsByReason": M_DENIALS_ADV, "Payers": M_PAYERS,
        "PaidByPayerType": M_PAID_BY_TYPE, "SubmitLag": M_SUBMIT_LAG, "DenialMap_and_DenialsByOwner": M_DENIAL_MAP,
        "Encounters2025": M_ENCOUNTERS, "Unbilled": M_UNBILLED, "AgedPending": M_AGED_SOLUTION,
        "Budget2025": M_BUDGET, "ICU_Q1_Salaries": M_ICU, "DenialDashboard": M_DASHBOARD,
        "DenialRateByPayerType": M_RATE,
    }
    for name, body in sol.items():
        extra[f"solutions/{name}.m"] = _m_file(name, body)
    L.extra_files = extra
    return L
