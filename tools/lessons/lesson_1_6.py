"""Lesson 1.6 · Sorting & Filtering Data.

Data: every 2025 emergency department visit at Cedar Ridge Medical Center (F03), 937 rows, as a PLAIN RANGE
(not an Excel Table) so learners meet the Sort Warning dialog and turn AutoFilter on themselves.

Lesson-specific data choices (documented in the README and on the Start Here sheet):
  * ArrivalDay ("Sun".."Sat") and ArrivalHour (0-23) are pre-computed from ArrivalDateTime, as ED exports usually do.
  * DoorToProviderMin = minutes from arrival to first provider contact; blank for LWBS (left without being seen).
  * SystolicBP is blank for the 15 LWBS visits ("not recorded"), so the live ShockIndex formula (=HeartRate/SystolicBP)
    shows #DIV/0! on those rows. That gives AGGREGATE a realistic column of errors to ignore.
  * Rows shaded orange = triage sepsis screen positive (suspected-infection complaint + temperature, heart-rate and
    respiratory-rate criteria all met). Color is the ONLY marker, so tasks must use Filter/Sort by Color.

Task mechanics:
  * Sorting/filtering tasks are "do it, then type what you see" (Markdown solution steps). The self-test types the
    Python answer; the key's live column re-computes each answer with an independent formula (except the color task,
    which no worksheet function can see).
  * Tasks 11-12 are SUBTOTAL/AGGREGATE formulas that read the rows an ESI-3 AutoFilter leaves visible. In self-test
    mode the customize hook hides every non-ESI-3 row on EDVisits so LibreOffice evaluates the learner formula exactly
    as Excel would with that filter on. The key's live formulas use filter-free equivalents (AVERAGEIFS, AGGREGATE array).
  * The Workspace sheet has a fixed layout (WS_* constants) that the prompts, solutions, gray row-3 labels and the
    README guide all point at, so the guide's try-it formulas, task 9/10 output and the bonus criteria never collide.
"""
from __future__ import annotations

import re
from collections import Counter
from decimal import ROUND_HALF_UP, Decimal
from statistics import median

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

from xlcourse import Lesson, Task, data

CODE = "1.6"
FACILITY = "F03"                      # Cedar Ridge Medical Center
FLAG_COLOR = "F8CBAD"                 # Office theme "Orange, Accent 2, Lighter 60%"
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]          # index = date.weekday()
WEEK_LIST = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]     # Excel's built-in custom list
LEVEL_OF_CARE = ["Admitted", "Observation", "Transferred", "Discharged", "Left AMA", "LWBS"]
INFECTION_WORDS = ("Fever", "Cough", "Urination", "Redness")

COLUMNS = ["EDVisitID", "PatientID", "ArrivalDateTime", "ArrivalDay", "ArrivalHour", "ArrivalMode", "ESILevel",
           "ChiefComplaint", "HeartRate", "SystolicBP", "ShockIndex", "DoorToProviderMin", "EDDisposition"]


def xround(x: float, nd: int) -> float:
    """Excel-style ROUND (half away from zero on the decimal representation)."""
    q = Decimal(1).scaleb(-nd)
    return float(Decimal(repr(x)).quantize(q, rounding=ROUND_HALF_UP))


def absr(sheet: str, rng: str) -> str:
    """'A4:B6' -> 'Sheet!$A$4:$B$6' (how the Advanced Filter dialog shows a range)."""
    return f"{sheet}!" + re.sub(r"([A-Z]+)(\d+)", r"$\1$\2", rng)


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="01-foundations", slug="06-sorting-filtering",
        title="Sorting & Filtering Data", level="Beginner", minutes=45,
        objectives=[
            "Sort by one or several columns, including custom orders",
            "Filter text, numbers, and dates with AutoFilter (and Top 10, by color)",
            "Summarize only visible rows with SUBTOTAL and AGGREGATE",
            "Use Advanced Filter for complex AND/OR criteria and unique lists",
        ],
        data_note="All 937 emergency department visits at Cedar Ridge Medical Center in 2025: arrival time, day and hour, "
                  "arrival mode, ESI triage level, chief complaint, two triage vital signs, a shock index formula, "
                  "door-to-provider minutes, and disposition. Rows shaded orange were flagged by the triage sepsis screen.",
    )

    # ------------------------------------------------------------------ data
    visits = sorted((r for r in data.load("ed_visits")
                     if r["FacilityID"] == FACILITY and r["ArrivalDateTime"].year == 2025),
                    key=lambda r: r["EDVisitID"])
    assert len(visits) == 937
    # EDVisitID order is arrival order, so "sort by EDVisitID A to Z" restores the original order (taught in the guide).
    assert all(a["ArrivalDateTime"] <= b["ArrivalDateTime"] for a, b in zip(visits, visits[1:]))

    first_row = 2
    rows = []
    for i, v in enumerate(visits):
        r = first_row + i
        arr = v["ArrivalDateTime"]
        seen = v["ProviderSeenDateTime"]
        lwbs = seen is None
        assert lwbs == (v["EDDisposition"] == "LWBS")
        wait = None if lwbs else (seen - arr).total_seconds() / 60
        assert wait is None or wait == int(wait)
        sirs_all = (v["TempF"] > 100.4 or v["TempF"] < 96.8) and v["HeartRate"] > 90 and v["RespRate"] > 20
        infection = any(w in v["ChiefComplaint"] for w in INFECTION_WORDS)
        rows.append({
            "EDVisitID": v["EDVisitID"], "PatientID": v["PatientID"], "ArrivalDateTime": arr,
            "ArrivalDay": DAYS[arr.weekday()], "ArrivalHour": arr.hour, "ArrivalMode": v["ArrivalMode"],
            "ESILevel": v["ESILevel"], "ChiefComplaint": v["ChiefComplaint"], "HeartRate": v["HeartRate"],
            "SystolicBP": None if lwbs else v["SystolicBP"],          # BP not recorded for patients who left early
            "ShockIndex": f"=I{r}/J{r}",                                # live formula: #DIV/0! where SBP is blank
            "DoorToProviderMin": None if wait is None else int(wait),
            "EDDisposition": v["EDDisposition"],
            # python-only helpers (not written: not in COLUMNS)
            "_si": None if lwbs else v["HeartRate"] / v["SystolicBP"],
            "_flag": bool(sirs_all and infection),
            "_row": r,
        })
    ed = L.add_table_sheet(
        "EDVisits", rows, columns=COLUMNS, as_table=False,
        formats={"ShockIndex": "0.00", "ArrivalHour": "0", "DoorToProviderMin": "0", "ESILevel": "0"},
        widths={"EDVisitID": 11, "PatientID": 10, "ArrivalDateTime": 17, "ArrivalDay": 11, "ArrivalHour": 12,
                "ArrivalMode": 13, "ESILevel": 9, "ChiefComplaint": 30, "HeartRate": 10, "SystolicBP": 11,
                "ShockIndex": 11, "DoorToProviderMin": 18, "EDDisposition": 14},
    )
    first, last = ed.first_row, ed.last_row
    assert (first, last) == (2, 938)
    n = len(rows)
    col = ed.col

    def r_(h: str) -> str:                     # learner-style range, e.g. EDVisits!L2:L938 (no quotes needed)
        return f"EDVisits!{col(h)}{first}:{col(h)}{last}"

    def R(h: str) -> str:                      # absolute range for key formulas
        return ed.rng(h)

    W = "DoorToProviderMin"
    waits = [x[W] for x in rows if x[W] is not None]
    flagged = [x for x in rows if x["_flag"]]
    assert 50 <= len(flagged) <= 90

    # ------------------------------------------------------------------ answers (computed in Python)
    # 1 · one-column sort, largest to smallest (blanks always go last)
    max_wait = max(waits)
    assert waits.count(max_wait) == 1
    t1_id = next(x["EDVisitID"] for x in rows if x[W] == max_wait)

    # 2 · three-level sort: ESI asc, wait desc (blanks last), arrival asc -> row 11
    def k2(x):
        return (x["ESILevel"], x[W] is None, -(x[W] or 0), x["ArrivalDateTime"])
    s2 = sorted(rows, key=k2)
    t2_row = 11
    t2_rec = s2[t2_row - first]
    assert sum(1 for x in rows if k2(x) == k2(t2_rec)) == 1          # position fully determined by the three keys
    assert t2_rec["ESILevel"] == 1 and t2_rec[W] is not None
    # without the 3rd level the row would sit inside a tie group (that's why the task adds it)
    assert sum(1 for x in rows if (x["ESILevel"], x[W]) == (t2_rec["ESILevel"], t2_rec[W])) > 1
    t2_id = t2_rec["EDVisitID"]
    esi1_n = sum(1 for x in rows if x["ESILevel"] == 1)

    # 3 · custom list Sun..Sat -> row of the first Wednesday
    day_ct = Counter(x["ArrivalDay"] for x in rows)
    t3_row = first + sum(day_ct[d] for d in WEEK_LIST[:WEEK_LIST.index("Wed")])
    az_wed_row = first + sum(c for d, c in day_ct.items() if d < "Wed")       # where A->Z would put it
    assert t3_row != az_wed_row

    # 4 · AutoFilter, two columns (AND): Ambulance and ESI 1-2
    t4 = sum(1 for x in rows if x["ArrivalMode"] == "Ambulance" and x["ESILevel"] <= 2)

    # 5 · search box: complaint contains "fever"
    fever_items = sorted({x["ChiefComplaint"] for x in rows if "fever" in x["ChiefComplaint"].lower()})
    t5 = sum(1 for x in rows if "fever" in x["ChiefComplaint"].lower())
    fever_begins = sum(1 for x in rows if x["ChiefComplaint"].lower().startswith("fever"))

    # 6 · Top 10 items -> smallest value still shown (10th largest)
    sw = sorted(waits, reverse=True)
    t6 = sw[9]
    assert sw.count(t6) == 1 and sw[10] < t6                        # exactly 10 rows shown, no ties at the cut-off

    # 7 · date filter: November 2025, Saturday or Sunday
    t7 = sum(1 for x in rows if x["ArrivalDateTime"].month == 11 and x["ArrivalDay"] in ("Sat", "Sun"))
    nov_n = sum(1 for x in rows if x["ArrivalDateTime"].month == 11)

    # 8 · filter by color + number filter
    t8 = sum(1 for x in flagged if x[W] is not None and x[W] > 30)

    # 9 · Advanced Filter, OR across two columns
    t9 = sum(1 for x in rows if x["ESILevel"] == 1 or (x[W] is not None and x[W] > 120))
    long_waits = sum(1 for x in rows if x[W] is not None and x[W] > 120)
    both9 = sum(1 for x in rows if x["ESILevel"] == 1 and x[W] is not None and x[W] > 120)

    # 10 · Advanced Filter, unique records only
    t10 = len({x["ChiefComplaint"] for x in rows})

    # 11 · SUBTOTAL average of visible rows (ESI 3 filter)
    esi3 = [x for x in rows if x["ESILevel"] == 3]
    esi3_waits = [x[W] for x in esi3 if x[W] is not None]
    t11 = xround(sum(esi3_waits) / len(esi3_waits), 1)
    esi3_blank = sum(1 for x in esi3 if x[W] is None)

    # 12 · AGGREGATE max ShockIndex of visible rows, ignoring #DIV/0!
    esi3_si = [x["_si"] for x in esi3 if x["_si"] is not None]
    t12 = max(esi3_si)
    esi3_err = sum(1 for x in esi3 if x["_si"] is None)
    assert esi3_err > 0

    # ------------------------------------------------------------------ bonus answers
    def in_review(x):
        return x["ESILevel"] == 1 or (x["ESILevel"] == 2 and x["ArrivalMode"] == "Ambulance" and x["ArrivalHour"] >= 18)
    review = [x for x in rows if in_review(x)]
    b1 = len(review)
    rev_w = [x[W] for x in review if x[W] is not None]
    b2 = xround(sum(rev_w) / len(rev_w), 1)

    def kb3(x):
        return (LEVEL_OF_CARE.index(x["EDDisposition"]), x[W] is None, -(x[W] or 0))
    rev_sorted = sorted(review, key=kb3)
    disp_ct = Counter(x["EDDisposition"] for x in review)
    b3_row = first + disp_ct["Admitted"]                  # first Observation row on the Review sheet
    b3_rec = rev_sorted[b3_row - first]
    assert b3_rec["EDDisposition"] == "Observation"
    assert sum(1 for x in review if kb3(x) == kb3(b3_rec)) == 1
    b3_id = b3_rec["EDVisitID"]
    # A->Z or Z->A by disposition would put a different visit in that row
    az = sorted(review, key=lambda x: (x["EDDisposition"], -(x[W] or 0)))
    za = sorted(review, key=lambda x: (tuple(-ord(c) for c in x["EDDisposition"]), -(x[W] or 0)))
    assert az[b3_row - first]["EDVisitID"] != b3_id and za[b3_row - first]["EDVisitID"] != b3_id

    rev_esi2 = [x[W] for x in review if x["ESILevel"] == 2 and x[W] is not None]
    b4 = median(rev_esi2)
    b4_mean = sum(rev_esi2) / len(rev_esi2)

    # ------------------------------------------------------------------ sheets
    L.sheet_order = ["Start Here", "Practice", "EDVisits", "Workspace", "Bonus", "Answer Key", "Bonus Key"]
    # Workspace layout (labels written by the customize hook; prompts and solutions point at the same cells)
    WS_CRIT9, WS_OUT9, WS_UNIQ, WS_CRITB = "A4:B6", "D4", "S4", "U4:W6"
    WS_FREE = "A9"                               # free cells for the guide's try-it formulas

    L.start_notes = [
        "EDVisits is a plain range on purpose (not an Excel Table), so you turn AutoFilter on yourself with "
        "Ctrl+Shift+L (Mac: ⌘+Shift+F). Rows shaded orange were flagged by the triage sepsis screen.",
        "To put the rows back in their original order at any time, sort by EDVisitID A to Z (IDs follow arrival order).",
        "Use the Workspace sheet for Advanced Filter criteria ranges and results, and for trying the guide's formulas. "
        "The gray labels in its row 3 show where each task's cells go.",
    ]

    L.practice_intro = (
        f"Every task uses the EDVisits sheet: {n} visits in rows {first}–{last}, columns A–{col(COLUMNS[-1])}. "
        "For tasks 1–10, sort or filter, then type what you see in the yellow cell: an ID, a row number, a count, or a value. "
        "Before each new filter task, clear the filters left over from the task before (Data → Clear). "
        "Tasks 9 and 10 use the Workspace sheet. "
        "Tasks 11 and 12 need an ESI 3 filter left on, so do them last."
    )

    L.tasks = [
        Task("Sort the EDVisits sheet by DoorToProviderMin, Largest to Smallest. Which EDVisitID is now in row 2? "
             "(That's the visit with the longest wait from arrival to first provider contact.)",
             answer=t1_id, title="Longest door-to-provider wait (one-column sort)",
             solution=f"1. Click any cell in column {col(W)} (DoorToProviderMin), for example {col(W)}2.\n"
                      "2. Click **Data → Sort Z to A** (the Z→A button). On a number column, the **Home → Sort & Filter** "
                      "menu calls the same command **Sort Largest to Smallest**. Because you clicked a single cell, "
                      "Excel sorts the whole block of data with it.\n"
                      f"3. Read A2: **{t1_id}**, a wait of {max_wait} minutes.",
             live=f"=INDEX({R('EDVisitID')},MATCH(MAX({R(W)}),{R(W)},0))",
             hint="Click one cell in the column, then use the Z→A button",
             explanation=f"Clicking one cell (not selecting the column) lets Excel find the whole data block and keep every row "
                         f"together. Scroll to the bottom: the {n - len(waits)} LWBS visits with a blank wait sit in rows "
                         f"{first + len(waits)}–{last}. Excel always sorts blanks last, whether you sort ascending or descending."),
        Task("Use the Sort dialog (Data → Sort) to sort by three levels: ESILevel Smallest to Largest, then DoorToProviderMin "
             "Largest to Smallest, then ArrivalDateTime Oldest to Newest. "
             f"Which EDVisitID is in row {t2_row}?",
             answer=t2_id, title=f"Three-level sort (row {t2_row})",
             solution="1. Click any cell in the data and choose **Data → Sort**. Check that **My data has headers** is ticked.\n"
                      "2. Sort by **ESILevel**, Sort On **Cell Values**, Order **Smallest to Largest**.\n"
                      "3. Click **Add Level**: Then by **DoorToProviderMin**, **Largest to Smallest**.\n"
                      "4. Click **Add Level** again: Then by **ArrivalDateTime**, **Oldest to Newest**. Click **OK**.\n"
                      f"5. Read A{t2_row}: **{t2_id}**.",
             live=f"=INDEX(SORTBY({R('EDVisitID')},{R('ESILevel')},1,{R(W)},-1,{R('ArrivalDateTime')},1),{t2_row - first + 1})",
             hint="Data → Sort, then Add Level twice. The top level wins",
             explanation=f"The first level groups the {esi1_n} ESI 1 visits at the top (rows {first}–{first + esi1_n - 1}). "
                         f"Inside that group the second level puts the longest waits first. Several ESI 1 patients waited exactly "
                         f"{t2_rec[W]} minutes, so the third level (arrival time) breaks the tie. Without a tiebreaker, rows that "
                         "tie on every level stay in whatever order they happened to be in, which depends on your earlier sorts."),
        Task("Sort by ArrivalDay using the built-in custom list Sun, Mon, Tue, Wed, Thu, Fri, Sat (Sort dialog → Order → "
             "Custom List…). On which row does the first Wednesday (Wed) visit appear? Type the row number.",
             answer=t3_row, title="Custom-list sort by weekday",
             solution="1. **Data → Sort**. If levels from task 2 are still listed, select each one and click **Delete Level**.\n"
                      "2. Sort by **ArrivalDay**, Sort On **Cell Values**, Order **Custom List…**.\n"
                      "3. In the Custom Lists box pick **Sun, Mon, Tue, Wed, Thu, Fri, Sat**, then **OK** twice.\n"
                      f"4. Scroll down column {col('ArrivalDay')} until the days change from Tue to Wed: row **{t3_row}**.",
             live=f'=COUNTIF({R("ArrivalDay")},"Sun")+COUNTIF({R("ArrivalDay")},"Mon")+COUNTIF({R("ArrivalDay")},"Tue")+{first}',
             hint="Sunday, Monday and Tuesday visits come first",
             explanation=f"A custom list sorts in the order you give it, not alphabetically. Sunday ({day_ct['Sun']} visits), "
                         f"Monday ({day_ct['Mon']}) and Tuesday ({day_ct['Tue']}) fill rows {first}–{t3_row - 1}, so Wednesday "
                         f"starts in row {t3_row}. An A→Z sort would have put the days in the order Fri, Mon, Sat, Sun, Thu, Tue, "
                         f"Wed, and the first Wednesday would land in row {az_wed_row}."),
        Task("Turn on AutoFilter (Ctrl+Shift+L; Mac: ⌘+Shift+F). Show only visits with ArrivalMode = Ambulance and "
             "ESILevel 1 or 2. How many visits are visible?",
             answer=t4, title="Ambulance arrivals with ESI 1–2 (two-column filter)",
             solution="1. Click any cell in the data and press **Ctrl+Shift+L** (Mac: **⌘+Shift+F**), or choose **Data → Filter**.\n"
                      f"2. Open the **ArrivalMode** arrow, untick **(Select All)**, tick **Ambulance**, **OK**.\n"
                      "3. Open the **ESILevel** arrow, untick 3, 4 and 5 (or use **Number Filters → Less Than Or Equal To → 2**), **OK**.\n"
                      f"4. The status bar reads **{t4} of {n} records found**.",
             live=f'=COUNTIFS({R("ArrivalMode")},"Ambulance",{R("ESILevel")},"<=2")',
             hint="Filter two columns, then read the status bar",
             explanation="Filters on different columns combine with AND: a row stays visible only if it passes every column's filter. "
                         "Ticking 1 and 2 inside one column is OR within that column. If the status bar shows *Filter Mode* instead of "
                         f"a count, select A{first}:A{last} and read **Count** on the status bar."),
        Task("Clear the filters (Data → Clear). Use the Search box in the ChiefComplaint filter to show every visit whose "
             "complaint contains the word fever anywhere. How many visits are visible?",
             answer=t5, title="Chief complaints containing \"fever\" (search box)",
             solution="1. **Data → Clear** removes the filters from task 4 but keeps the filter arrows.\n"
                      "2. Open the **ChiefComplaint** arrow and type **fever** in the **Search** box.\n"
                      f"3. The list shrinks to {len(fever_items)} items ({', '.join(fever_items)}). Click **OK**.\n"
                      f"4. The status bar reads **{t5} of {n} records found**.",
             live=f'=COUNTIF({R("ChiefComplaint")},"*fever*")',
             hint="Type in the Search box, then check which items it ticks before you click OK",
             explanation="The Search box matches text anywhere in the value and ignores case, so it finds complaints that start, "
                         "end, or contain *fever*. It works the same as **Text Filters → Contains**. Always glance at the ticked items "
                         "before clicking OK, because a search can catch values you didn't intend."),
        Task("Clear the filters. Use Number Filters → Top 10 on DoorToProviderMin to show the 10 longest waits. What is the "
             "smallest DoorToProviderMin still visible (the 10th-longest wait), in minutes?",
             answer=t6, title="Top 10 longest waits (10th-longest value)",
             solution="1. **Data → Clear**.\n"
                      "2. Open the **DoorToProviderMin** arrow → **Number Filters → Top 10…**.\n"
                      "3. Leave **Top**, **10**, **Items** and click **OK**.\n"
                      f"4. Ten rows remain. Select {col(W)}{first}:{col(W)}{last} and read **Minimum** on the status bar, which "
                      "skips filtered-out rows (right-click the status bar to turn Minimum on, as in Lesson 1.1). Or sort the "
                      f"column Largest to Smallest and look at the last visible row. Either way it's **{t6}** minutes.",
             live=f"=LARGE({R(W)},10)",
             hint="After the Top 10 filter, Minimum on the status bar finds the smallest visible wait",
             explanation=f"Top 10 keeps the rows whose value is at least the 10th largest. Here the 10th-longest wait is {t6} minutes, "
                         "so exactly ten rows stay visible. If several rows had tied at the cut-off value, Excel would show all of them, "
                         "so a Top 10 filter can show more than ten rows. The same dialog does Bottom 10 and Top 10 Percent."),
        Task("Clear the filters. Show only visits that arrived in November 2025 on a Saturday or Sunday. How many are there?",
             answer=t7, title="November weekend arrivals (date filter)",
             solution="1. **Data → Clear**.\n"
                      "2. Open the **ArrivalDateTime** arrow. In the date tree, untick **(Select All)**, expand **2025**, tick "
                      "**November**, **OK**. (Or use **Date Filters → All Dates in the Period → November**.)\n"
                      f"3. The status bar shows {nov_n} November visits. Now open the **ArrivalDay** arrow, keep only **Sat** and **Sun**, **OK**.\n"
                      f"4. The status bar reads **{t7} of {n} records found**.",
             live=f'=SUMPRODUCT((MONTH({R("ArrivalDateTime")})=11)*(({R("ArrivalDay")}="Sat")+({R("ArrivalDay")}="Sun")))',
             hint="Date tree for the month, then ArrivalDay for the weekend",
             explanation="Excel groups real dates in the filter list by year, month, and day, so you can tick a whole month at once. "
                         "There's no built-in *weekend* date filter, which is why ED exports often include a day-of-week column like "
                         "ArrivalDay. Avoid **Between 11/1/2025 and 11/30/2025** here: the values include times, and 11/30/2025 means "
                         "midnight at the start of November 30, so visits later that day would be left out."),
        Task("Clear the filters. Rows shaded orange are visits that triage flagged as a positive sepsis screen. Filter to the "
             "orange rows, then also keep only those with DoorToProviderMin greater than 30. How many flagged visits waited "
             "more than 30 minutes to see a provider?",
             answer=t8, title="Sepsis-flagged visits that waited over 30 minutes (filter by color)",
             solution="1. **Data → Clear**.\n"
                      "2. Open any column's arrow (EDVisitID works) → **Filter by Color** → pick the orange swatch under "
                      f"**Filter by Cell Color**. {len(flagged)} rows remain.\n"
                      "3. Open the **DoorToProviderMin** arrow → **Number Filters → Greater Than…** → type **30** → **OK**.\n"
                      f"4. The status bar reads **{t8} of {n} records found**.",
             live=False,
             hint="Filter by Color, then Number Filters → Greater Than",
             explanation=f"Color is the only marker for the sepsis flag, so Filter by Color is the only way to isolate those {len(flagged)} rows. "
                         "No worksheet function can read a fill color, which is why this task has no live formula in the key. "
                         "Combining a color filter on one column with a number filter on another is still AND logic. "
                         "Sepsis care is time-critical, so these are the waits a quality team reviews first."),
        Task("Clear the filters. Use Advanced Filter to find visits that were ESI level 1 OR waited more than 120 minutes for a "
             f"provider. Type the criteria range in Workspace!{WS_CRIT9}. Then, with the Workspace sheet active, choose "
             f"Data → Advanced and copy the results to Workspace!{WS_OUT9}. How many visits does it copy? Don't count the header.",
             answer=t9, title="Advanced Filter: ESI 1 OR wait over 120 minutes",
             solution="1. **Data → Clear** on EDVisits.\n"
                      f"2. On **Workspace**, type the criteria range in {WS_CRIT9}. Copy the two headers from row 1 of EDVisits:\n\n"
                      "   | Row | A | B |\n   |:-:|---|---|\n   | 4 | ESILevel | DoorToProviderMin |\n   | 5 | 1 | |\n   | 6 | | >120 |\n\n"
                      f"3. Still on Workspace, click {WS_OUT9}, then choose **Data → Advanced**. Select "
                      f"**Copy to another location**. List range: "
                      f"`EDVisits!$A$1:${col(COLUMNS[-1])}${last}`. Criteria range: `{absr('Workspace', WS_CRIT9)}`. "
                      f"Copy to: `{absr('Workspace', WS_OUT9)}`. **OK**.\n"
                      f"4. The copy has a header row plus **{t9}** visit rows.",
             live=f"=SUMPRODUCT(--((({R('ESILevel')}=1)+({R(W)}>120))>0))",
             hint="Conditions on different rows of the criteria range mean OR",
             explanation=f"Each criteria row is one way to qualify. Row 5 catches the {esi1_n} ESI 1 visits and row 6 catches the "
                         f"{long_waits} waits over 120 minutes. "
                         + ("No ESI 1 patient waited that long, so the two groups don't overlap here. When they do, a visit that meets "
                            "both rows is copied once, not twice. " if both9 == 0 else
                            f"The {both9} visits that meet both rows are copied once, not twice. ")
                         + "AutoFilter can't do this, because filters on two columns always combine with AND. The criteria headers must "
                         "match the data headers exactly, so copy them from row 1 of EDVisits rather than typing them."),
        Task("Use Advanced Filter with Unique records only to copy a list of the different ChiefComplaint values to "
             f"Workspace!{WS_UNIQ}. How many different chief complaints are there? Don't count the header.",
             answer=t10, title="Unique chief complaints (Advanced Filter)",
             solution=f"1. On **Workspace**, click **{WS_UNIQ}** and choose **Data → Advanced**.\n"
                      "2. Select **Copy to another location**. List range: "
                      f"`EDVisits!${col('ChiefComplaint')}$1:${col('ChiefComplaint')}${last}` (just that one column, header included).\n"
                      f"3. Make the **Criteria range** box empty: Excel may fill in `{absr('Workspace', WS_CRIT9)}` from task 9, so delete it. "
                      f"Copy to: `{absr('Workspace', WS_UNIQ)}`. Tick **Unique records only**. **OK**.\n"
                      f"4. The list has a header plus **{t10}** complaints. Select them and read Count on the status bar.",
             live=f"=SUMPRODUCT(1/COUNTIF({R('ChiefComplaint')},{R('ChiefComplaint')}))",
             hint="List range: the ChiefComplaint column only. Empty the Criteria range box",
             explanation="With an empty criteria range every row qualifies, and **Unique records only** keeps the first copy of each "
                         "distinct value. Because the list range is a single column, the duplicates are judged on that column alone. "
                         "With the whole table as the list range, a row would only count as a duplicate if every column matched."),
        Task("Filter EDVisits to ESILevel 3 only, with no other filters on. In the yellow cell, write a formula that averages "
             "DoorToProviderMin for the visible rows only, rounded to 1 decimal place with ROUND. Leave the filter on for task 12.",
             answer=t11, fmt="0.0", tol=0.0001, title="SUBTOTAL average of the visible rows (ESI 3)",
             solution=f"=ROUND(SUBTOTAL(101,{r_(W)}),1)",
             live=f"=ROUND(AVERAGEIFS({R(W)},{R('ESILevel')},3),1)",
             hint="SUBTOTAL's AVERAGE is function_num 1 (or 101)",
             explanation=f"SUBTOTAL skips rows hidden by a filter, so it looks only at the {len(esi3)} ESI 3 visits. Function 1 would work "
                         "too, because both 1 and 101 ignore filtered-out rows (101 also ignores rows you hide by hand). "
                         f"AVERAGE-type functions skip blank cells, so the {esi3_blank} ESI 3 LWBS visits with no wait don't drag the "
                         "average toward zero. Plain AVERAGE would ignore the filter and average all "
                         f"{len(waits)} waits. The key's live cell uses AVERAGEIFS (Lesson 2.5) so it works without a filter. "
                         "If you clear the filter later, this check turns red, which shows SUBTOTAL responding to the filter. To keep "
                         "the result, copy the cell and paste it back as a value."),
        Task("Keep the ESI 3 filter on. ShockIndex (HeartRate ÷ SystolicBP) shows #DIV/0! where no blood pressure was recorded, "
             f"so =SUBTOTAL(104,{r_('ShockIndex')}) returns #DIV/0!. In the yellow cell, write an AGGREGATE formula that returns "
             "the highest ShockIndex among the visible rows while ignoring the errors. Don't round it.",
             answer=t12, fmt="0.00", title="AGGREGATE maximum, ignoring hidden rows and errors",
             solution=f"=AGGREGATE(4,7,{r_('ShockIndex')})",
             live=f"=AGGREGATE(14,6,{R('ShockIndex')}/({R('ESILevel')}=3),1)",
             hint="MAX is function 4. Pick the option that ignores hidden rows AND error values",
             explanation=f"AGGREGATE(4, 7, range) means MAX (4) while ignoring hidden rows and error values (7). The {esi3_err} visible "
                         "#DIV/0! cells come from LWBS patients whose blood pressure wasn't recorded. SUBTOTAL and MAX return an error "
                         "as soon as the range holds one, while AGGREGATE steps over it. Option 6 would ignore errors but include the hidden "
                         "rows, giving the highest shock index of the whole year. A shock index near or above 1.0 suggests the heart is "
                         "racing to keep blood pressure up, which is why EDs watch it."),
    ]

    # ------------------------------------------------------------------ bonus
    mask = (f"(({R('ESILevel')}=1)+({R('ESILevel')}=2)*({R('ArrivalMode')}=\"Ambulance\")*({R('ArrivalHour')}>=18))")
    loc_array = "{" + ",".join(f'"{d}"' for d in LEVEL_OF_CARE) + "}"
    L.bonus_title = "Bonus: the high-acuity evening review"
    L.bonus_scenario = (
        "Cedar Ridge's ED medical director is preparing a high-acuity review for the quality committee. It covers every ESI 1 "
        "visit, plus ESI 2 patients who arrived by ambulance in the evening (ArrivalHour 18 or later, which means 6:00 pm to "
        "11:59 pm). That rule is an OR across different columns, so AutoFilter can't do it in one step. "
        "Your practice task 11 and 12 formulas follow the ESI 3 filter, so they switch to ✘ when you clear it. That's expected. "
        "To keep them green, first copy each of those two answer cells and paste it back as a value. "
        "Then clear every filter on EDVisits (Data → Clear), insert a new sheet named Review, type the criteria range in "
        f"Workspace!{WS_CRITB}, and run Data → Advanced from the Review sheet to copy the matching rows to Review!A1."
    )
    L.bonus = [
        Task("How many visits does your Advanced Filter copy to the Review sheet? Don't count the header row.",
             answer=b1, title="Visits in the review list",
             solution="1. **Data → Clear** on EDVisits. Insert a sheet with the **+** button next to the sheet tabs (or press "
                      "**Shift+F11**; on a Mac laptop, **Fn+Shift+F11**), then double-click its tab and rename it **Review**.\n"
                      f"2. On **Workspace**, type this criteria range in {WS_CRITB}, to the right of your task 10 list. Copy the "
                      "headers from row 1 of EDVisits:\n\n"
                      "   | Row | U | V | W |\n   |:-:|---|---|---|\n   | 4 | ESILevel | ArrivalMode | ArrivalHour |\n"
                      "   | 5 | 1 | | |\n   | 6 | 2 | Ambulance | >=18 |\n\n"
                      "3. Click **Review!A1**, then **Data → Advanced** → **Copy to another location**. List range "
                      f"`EDVisits!$A$1:${col(COLUMNS[-1])}${last}`, Criteria range `{absr('Workspace', WS_CRITB)}`, "
                      "Copy to `Review!$A$1`. **OK**.\n"
                      f"4. Review shows a header plus **{b1}** rows (rows 2–{first + b1 - 1}).",
             live=f"=SUMPRODUCT({mask})",
             hint="Two criteria rows: ESI 1 alone, and ESI 2 + Ambulance + >=18 together",
             explanation=f"Row 5 of the criteria range (ESILevel = 1) catches all {esi1_n} ESI 1 visits. Row 6 is an AND: ESI 2 *and* "
                         f"Ambulance *and* ArrivalHour ≥ 18, which adds {b1 - esi1_n} more. The two rows together are OR. Excel only copies "
                         "filtered data to the active sheet, so you must start Data → Advanced from the Review sheet, or Excel shows an error."),
        Task("In the yellow cell, write a formula for the average DoorToProviderMin of the review visits, rounded to "
             "1 decimal place.",
             answer=b2, fmt="0.0", tol=0.0001, title="Average wait in the review list",
             solution=f"`=ROUND(AVERAGE(Review!{col(W)}2:{col(W)}{first + b1 - 1}),1)` "
                      f"(or `=ROUND(AVERAGE(Review!{col(W)}:{col(W)}),1)`, because AVERAGE skips the text header).",
             live=f"=ROUND(SUMPRODUCT({mask}*ISNUMBER({R(W)}),{R(W)})/SUMPRODUCT({mask}*ISNUMBER({R(W)})),1)",
             hint=f"AVERAGE over column {col(W)} of the Review sheet",
             explanation="The Advanced Filter result is ordinary cells, so regular functions work on it. AVERAGE ignores text and blanks, "
                         "so even a whole-column reference gives the right answer. The copy is a snapshot: if EDVisits changes, run the "
                         "Advanced Filter again."),
        Task("The director lists dispositions by level of care. Create the custom list Admitted, Observation, Transferred, "
             "Discharged, Left AMA, LWBS. Sort the Review sheet by EDDisposition with that list, then by DoorToProviderMin Largest "
             f"to Smallest. Which EDVisitID is in row {b3_row} of the Review sheet?",
             answer=b3_id, title=f"Custom-list sort of the review list (row {b3_row})",
             solution="1. On Review, click any cell in the data → **Data → Sort**. Sort by **EDDisposition**, Order **Custom List…**.\n"
                      "2. Select **NEW LIST**, type the six entries one per line (press Enter after each), click **Add**, then **OK**.\n"
                      "3. **Add Level**: Then by **DoorToProviderMin**, **Largest to Smallest**. **OK**.\n"
                      f"4. Rows 2–{b3_row - 1} hold the {disp_ct['Admitted']} Admitted visits, so row {b3_row} is the first Observation "
                      f"visit: **{b3_id}**.",
             live=(f"=INDEX(SORTBY(FILTER({R('EDVisitID')},{mask}),XMATCH(FILTER({R('EDDisposition')},{mask}),{loc_array}),1,"
                   f"FILTER({R(W)},{mask}),-1),{b3_row - first + 1})"),
             hint="Order → Custom List… → NEW LIST",
             explanation=f"The review list holds {disp_ct['Admitted']} Admitted, {disp_ct['Observation']} Observation and "
                         f"{disp_ct['Discharged']} Discharged visits, so the custom order puts Observation in rows {b3_row}–"
                         f"{b3_row + disp_ct['Observation'] - 1}. Row {b3_row} is the Observation patient who waited longest "
                         f"({b3_rec[W]} minutes). An A→Z sort would have put Discharged second, and a different visit in row {b3_row}. "
                         "Every disposition appears in the list, even ones the review doesn't contain, so the list works on any month's data. "
                         "Excel saves a custom list on your computer, so it's available in every workbook you open there."),
        Task("CMS reports ED wait times as medians, because a few very long waits pull an average up. Filter the Review sheet to "
             "ESILevel 2 only. Then, in the yellow cell, write an AGGREGATE formula that returns the median DoorToProviderMin "
             "of the visible Review rows.",
             answer=b4, fmt=None if float(b4).is_integer() else "0.0",
             title="Median wait for the ESI 2 evening ambulance patients",
             solution=f"1. On Review, press **Ctrl+Shift+L** (Mac: **⌘+Shift+F**) and filter **ESILevel** to **2**.\n"
                      f"2. In the yellow B4 cell on the Bonus sheet, type `=AGGREGATE(12,5,Review!{col(W)}2:{col(W)}{first + b1 - 1})`. "
                      "The Bonus sheet isn't filtered, so the filter can't hide your formula.\n"
                      f"3. It returns **{b4:g}**.",
             live=f"=MEDIAN(IF({mask}*({R('ESILevel')}=2)*ISNUMBER({R(W)}),{R(W)}))",
             hint="SUBTOTAL has no median, but AGGREGATE function 12 does",
             explanation=f"AGGREGATE(12, 5, range) is MEDIAN (12) that ignores hidden rows (5). SUBTOTAL can't do this, because its eleven "
                         f"functions don't include MEDIAN, LARGE, SMALL, or PERCENTILE. The {len(rev_esi2)} ESI 2 evening ambulance "
                         f"patients had a median wait of {b4:g} minutes but a mean of {b4_mean:.1f}, because a few long waits pull the mean up. "
                         "That gap is why CMS reports medians."),
    ]

    # ------------------------------------------------------------------ customize: color flags, Workspace, self-test filter
    flag_rows = [x["_row"] for x in flagged]
    non_esi3_rows = [x["_row"] for x in rows if x["ESILevel"] != 3]
    ncols = len(COLUMNS)

    @L.customize
    def _sheets(wb, lesson, selftest):
        ws = wb["EDVisits"]
        fill = PatternFill("solid", fgColor=FLAG_COLOR)
        for r in flag_rows:
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).fill = fill
        ws.row_dimensions[1].height = 18

        wsp = wb.create_sheet("Workspace")
        wsp.sheet_properties.tabColor = "70AD47"
        wsp["A1"] = "Workspace: Advanced Filter criteria and results"
        wsp["A1"].font = Font(bold=True, size=13, color="1F4E79")
        wsp["A2"] = ("Copy header names from row 1 of EDVisits so they match exactly. "
                     "Start Data → Advanced from this sheet when the results should land here.")
        wsp["A2"].font = Font(italic=True, color="595959")
        wsp["A2"].alignment = Alignment(wrap_text=False)
        # Gray labels show where each task's cells go (they match the prompts). Row 3 sits above the work areas, and
        # nothing is placed below a Copy to cell, because Advanced Filter writes over the cells under it.
        def above(ref: str) -> str:                              # cell just above the top-left of a range
            c, r = coordinate_from_string(ref.split(":")[0])
            return f"{c}{r - 1}"

        label_font = Font(italic=True, size=10, color="7F7F7F")
        for ref, text in ((WS_CRIT9, f"Task 9 criteria: {WS_CRIT9}"),
                          (WS_OUT9, f"Task 9 results: copy to {WS_OUT9}"),
                          (WS_UNIQ, f"Task 10 list: copy to {WS_UNIQ}"),
                          (WS_CRITB, f"Bonus criteria: {WS_CRITB}"),
                          (WS_FREE, f"Free cells for the guide's examples: {WS_FREE} and below")):
            wsp[above(ref)] = text
            wsp[above(ref)].font = label_font
        for i in range(1, 24):                                     # A:W
            wsp.column_dimensions[get_column_letter(i)].width = 14
        # Task 9 copies all 13 EDVisits columns to D:P, so widen the ones whose values need room.
        out_idx = column_index_from_string(coordinate_from_string(WS_OUT9)[0])
        for h, wdt in (("ArrivalDateTime", 17), ("ChiefComplaint", 30), ("DoorToProviderMin", 18)):
            wsp.column_dimensions[get_column_letter(out_idx + COLUMNS.index(h))].width = wdt
        wsp.column_dimensions["B"].width = 18                      # DoorToProviderMin criteria header
        wsp.column_dimensions[coordinate_from_string(WS_UNIQ)[0]].width = 30   # unique ChiefComplaint list
        wsp.freeze_panes = "A4"

        if selftest:
            # Simulate the learner's "ESILevel = 3" AutoFilter for tasks 11-12.
            for r in non_esi3_rows:
                ws.row_dimensions[r].hidden = True

    return L
