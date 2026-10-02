"""Lesson 5.1 · Recording Your First Macros.

Data (Cedar Ridge Medical Center, F03):
  Census_Nov / Census_Dec   raw daily census exports for the Medical-Surgical unit (D310), Nov and Dec 2025.
                            Plain ranges, plain headers, dates left in General format (they show as serial numbers),
                            exactly like a CSV export. Learners record FormatCensusReport on Census_Nov (TOTAL row
                            recorded with relative references) and replay it on Census_Dec, which has one more day.
  ED_Nov / ED_Dec           raw ED visit exports (sorted by PatientID, as the tracking system exports them) for the
                            bonus: record ExtractHighAcuity on ED_Nov, edit it, run it on ED_Dec.

Checks read cells the learner's macros write:
  * Practice summaries find the row labeled TOTAL with MATCH and read its cells.
  * Bonus summaries read the HighAcuity sheet through INDIRECT, so the pristine file never references a missing sheet.
    B5 (ED_Dec filter off) also waits until HighAcuity's A2 is a December visit, so a sheet left over from recording on
    ED_Nov can't turn it green.
  * The self-test simulates the macros: `fill` writes the TOTAL rows, and a customize hook builds HighAcuity.

Files written next to the workbook (ASCII, CRLF so the VBE imports them cleanly):
  solutions/modCensusReport.bas   FormatCensusReport as the recorder writes it + a cleaned version (spoilers)
  solutions/modHighAcuity.bas     the bonus macro as recorded on ED_Nov + the edited version for ED_Dec (spoilers)

Smoke test of the VBA in LibreOffice (run from tools/):
  python3 -c "import lessons.lesson_5_1 as m; m.smoke_test()"
Both census macros write TOTAL to row 32 on Census_Nov and row 33 on Census_Dec (patient days 809 and 863). LibreOffice has
no Worksheet.Sort object, so the bonus macro stops at its Sort block there; the smoke test also runs a copy with that block
swapped for Range.Sort, which reproduces the answer key (26 rows, ED211706 first, ED211715 in A5, ED212209 last, ED_Dec filter off).
"""
from __future__ import annotations

from openpyxl.styles import Font, PatternFill

from xlcourse import Lesson, Task, data
from xlcourse.lesson import estimate_lines

CODE = "5.1"
FACILITY = "F03"           # Cedar Ridge Medical Center
CENSUS_DEPT = "D310"       # Medical-Surgical, 30 staffed beds
CENSUS_COLS = ["CensusDate", "FacilityID", "DeptID", "StaffedBeds", "Admissions", "Discharges", "MidnightCensus"]
ED_COLS = ["EDVisitID", "PatientID", "ArrivalDateTime", "ArrivalMode", "ESILevel", "ChiefComplaint",
           "TriageDateTime", "ProviderSeenDateTime", "DepartureDateTime", "EDDisposition"]
HA_SHEET = "HighAcuity"


# ---------------------------------------------------------------------------
# VBA source (ASCII only): one source of truth for the .bas files, the README
# answer key and the LibreOffice smoke test.
# ---------------------------------------------------------------------------
def _crlf(text: str) -> bytes:
    """VBE's File > Import File expects Windows line endings and ANSI text."""
    return text.replace("\r\n", "\n").replace("\n", "\r\n").encode("ascii")


CENSUS_RECORDED = '''Sub FormatCensusReport()
'
' FormatCensusReport Macro
' Formats the raw daily census export and adds a TOTAL row.
'
' Keyboard Shortcut: Ctrl+Shift+R
'
    ' --- Use Relative References is OFF: these lines name exact cells ---
    Range("A1:H1").Select
    Selection.Font.Bold = True
    With Selection.Interior
        .Pattern = xlSolid
        .PatternColorIndex = xlAutomatic
        .ThemeColor = xlThemeColorAccent1
        .TintAndShade = 0.799981688894314
        .PatternTintAndShade = 0
    End With
    Range("H1").Select
    ActiveCell.FormulaR1C1 = "Occupancy"
    Range("H2").Select
    Columns("A:A").Select
    Selection.NumberFormat = "mm/dd/yyyy"
    Range("A1").Select
    Selection.End(xlDown).Select
    ' --- Use Relative References turned ON: moves are stored as offsets ---
    ActiveCell.Offset(1, 0).Range("A1").Select
    ActiveCell.FormulaR1C1 = "TOTAL"
    ActiveCell.Offset(0, 1).Range("A1").Select
    ActiveCell.Offset(0, 2).Range("A1:D1").Select
    Selection.FormulaR1C1 = "=SUM(R2C:R[-1]C)"
    ActiveCell.Offset(0, 4).Range("A1").Select
    Selection.FormulaR1C1 = "=RC[-1]/RC[-4]"
    Selection.Style = "Percent"
    Selection.NumberFormat = "0.0%"
    ActiveCell.Rows("1:1").EntireRow.Select
    Selection.Font.Bold = True
    ' --- Use Relative References turned OFF again ---
    Columns("A:H").Select
    Columns("A:H").EntireColumn.AutoFit
    Range("A1").Select
    With ActiveWindow
        .SplitColumn = 0
        .SplitRow = 1
    End With
    ActiveWindow.FreezePanes = True
End Sub'''

CENSUS_CLEAN = '''Sub FormatCensusReport_Clean()
'   The same report as FormatCensusReport, tidied by hand:
'   no Select/Selection pairs and no default properties.
'   Run it with a raw census export as the active sheet.
    With Range("A1:H1")
        .Font.Bold = True
        .Interior.ThemeColor = xlThemeColorAccent1
        .Interior.TintAndShade = 0.8
    End With
    Range("H1").Value = "Occupancy"
    Columns("A").NumberFormat = "mm/dd/yyyy"

    ' The TOTAL row goes one row below the last date in column A
    ' (the same cell that Ctrl + Down arrow, then Down arrow, would reach).
    With Range("A1").End(xlDown).Offset(1, 0)
        .Value = "TOTAL"
        ' Columns D:G: from row 2 (fixed) down to the row above (relative).
        .Offset(0, 3).Resize(1, 4).FormulaR1C1 = "=SUM(R2C:R[-1]C)"
        ' Column H: patient days (G) divided by staffed-bed days (D).
        .Offset(0, 7).FormulaR1C1 = "=RC[-1]/RC[-4]"
        .Offset(0, 7).NumberFormat = "0.0%"
        .EntireRow.Font.Bold = True
    End With

    Columns("A:H").AutoFit
    ' Freeze the header row (window settings always act on the active sheet).
    ActiveWindow.FreezePanes = False
    ActiveWindow.ScrollRow = 1
    With ActiveWindow
        .SplitColumn = 0
        .SplitRow = 1
        .FreezePanes = True
    End With
End Sub'''


def _ha_recorded(nov_last: int, nov_ha_last: int) -> str:
    return f'''Sub ExtractHighAcuity_AsRecorded()
'
' ExtractHighAcuity Macro
' Copies ESI 1-2 visits to a new HighAcuity sheet, sorted by ESI level and arrival.
'
' Keyboard Shortcut: Ctrl+Shift+H
'
'   What the recorder wrote on ED_Nov (the <-- notes are added). It works on
'   November. Lines marked <-- are tied to November and need editing for ED_Dec.
    Range("A1").Select
    Selection.AutoFilter
    ActiveSheet.Range("$A$1:$J${nov_last}").AutoFilter Field:=5, Criteria1:="<=2", _
        Operator:=xlAnd                                 ' <-- November's exact range
    Range(Selection, Selection.End(xlToRight)).Select
    Range(Selection, Selection.End(xlDown)).Select
    Selection.Copy
    Sheets.Add After:=ActiveSheet
    ActiveSheet.Paste
    Application.CutCopyMode = False
    Sheets("Sheet1").Select                             ' <-- the new sheet's
    Sheets("Sheet1").Name = "HighAcuity"                ' <-- temporary name
    ActiveWorkbook.Worksheets("HighAcuity").Sort.SortFields.Clear
    ActiveWorkbook.Worksheets("HighAcuity").Sort.SortFields.Add2 Key:=Range( _
        "E2:E{nov_ha_last}"), SortOn:=xlSortOnValues, Order:=xlAscending, DataOption:= _
        xlSortNormal                                    ' <-- November's row count
    ActiveWorkbook.Worksheets("HighAcuity").Sort.SortFields.Add2 Key:=Range( _
        "C2:C{nov_ha_last}"), SortOn:=xlSortOnValues, Order:=xlAscending, DataOption:= _
        xlSortNormal                                    ' <-- November's row count
    With ActiveWorkbook.Worksheets("HighAcuity").Sort
        .SetRange Range("A1:J{nov_ha_last}")                       ' <-- November's row count
        .Header = xlYes
        .MatchCase = False
        .Orientation = xlTopToBottom
        .SortMethod = xlPinYin
        .Apply
    End With
    Cells.Select
    Cells.EntireColumn.AutoFit
    Range("A1").Select
    Sheets("ED_Nov").Select                             ' <-- November's sheet
    Range("A1").Select
    Selection.AutoFilter
End Sub'''


def _ha_edited(nov_last: int, nov_ha_last: int) -> str:
    return f'''Sub ExtractHighAcuity()
'
' ExtractHighAcuity Macro
' Copies ESI 1-2 visits to a new HighAcuity sheet, sorted by ESI level and arrival.
'
' Keyboard Shortcut: Ctrl+Shift+H
'
'   Recorded on ED_Nov, then edited to run on ED_Dec. Each EDIT comment marks a change.
'   Delete any old HighAcuity sheet before you run it: a second sheet with that
'   name is not allowed (run-time error 1004).
    Sheets("ED_Dec").Select                             ' EDIT: added, start on December
    Range("A1").Select
    Selection.AutoFilter
    Range("A1").CurrentRegion.AutoFilter Field:=5, Criteria1:="<=2", _
        Operator:=xlAnd                                 ' EDIT: was ActiveSheet.Range("$A$1:$J${nov_last}")
    Range(Selection, Selection.End(xlToRight)).Select
    Range(Selection, Selection.End(xlDown)).Select
    Selection.Copy
    Sheets.Add After:=ActiveSheet
    ActiveSheet.Paste
    Application.CutCopyMode = False
    ActiveSheet.Name = "HighAcuity"                     ' EDIT: was Sheets("Sheet1").Select / .Name
    ActiveWorkbook.Worksheets("HighAcuity").Sort.SortFields.Clear
    ' EDIT: keys were Range("E2:E{nov_ha_last}") and Range("C2:C{nov_ha_last}"). One cell is enough
    ' to name the key column. (.Add works in Excel 2007 and later; Microsoft 365 records .Add2.)
    ActiveWorkbook.Worksheets("HighAcuity").Sort.SortFields.Add Key:=Range("E1"), _
        SortOn:=xlSortOnValues, Order:=xlAscending, DataOption:=xlSortNormal
    ActiveWorkbook.Worksheets("HighAcuity").Sort.SortFields.Add Key:=Range("C1"), _
        SortOn:=xlSortOnValues, Order:=xlAscending, DataOption:=xlSortNormal
    With ActiveWorkbook.Worksheets("HighAcuity").Sort
        .SetRange Range("A1").CurrentRegion             ' EDIT: was Range("A1:J{nov_ha_last}")
        .Header = xlYes
        .MatchCase = False
        .Orientation = xlTopToBottom
        .SortMethod = xlPinYin
        .Apply
    End With
    Cells.EntireColumn.AutoFit
    Range("A1").Select
    Sheets("ED_Dec").Select                             ' EDIT: was Sheets("ED_Nov")
    Range("A1").Select
    Selection.AutoFilter
End Sub'''


def _module(name: str, header: list[str], procs: list[str]) -> str:
    lines = [f'Attribute VB_Name = "{name}"', "' " + "=" * 74]
    lines += ["' " + h for h in header]
    lines += ["' " + "=" * 74, "Option Explicit", ""]
    return "\n".join(lines) + "\n" + "\n\n".join(procs) + "\n"


SPOILER = [
    "Lesson 5.1 - Recording Your First Macros - reference solution",
    "SPOILER: record your own macro first, then compare.",
    "",
    "To import: open the VBE with Alt + F11 (Mac: Option + F11, or Developer >",
    "Visual Basic), choose File > Import File..., and pick this .bas file. If your",
    "workbook already has a macro with the same name, rename or delete one of them",
    "first so it's clear which one runs.",
    "An imported macro has no shortcut key (the 'Keyboard Shortcut' line is only a",
    "comment). To set one, open Macros (Alt + F8; Mac: Option + F8), select it, and",
    "click Options...",
    "Save the workbook as .xlsm to keep the code.",
]


# ---------------------------------------------------------------------------
# Data helpers
# ---------------------------------------------------------------------------
def _census(month: int) -> list[dict]:
    rows = [r for r in data.load("daily_census")
            if r["DeptID"] == CENSUS_DEPT and r["CensusDate"].year == 2025 and r["CensusDate"].month == month]
    rows.sort(key=lambda r: r["CensusDate"])
    return rows


def _ed(month: int) -> list[dict]:
    rows = [r for r in data.load("ed_visits")
            if r["FacilityID"] == FACILITY and r["ArrivalDateTime"].year == 2025 and r["ArrivalDateTime"].month == month]
    # The tracking system exports visits sorted by patient, so the macro has to sort them itself.
    rows.sort(key=lambda r: (r["PatientID"], r["ArrivalDateTime"]))
    return rows


def _high_acuity(rows: list[dict]) -> list[dict]:
    """What the macro must produce: ESI <= 2, sorted by ESILevel then ArrivalDateTime (stable, like Excel)."""
    ha = [r for r in rows if r["ESILevel"] <= 2]
    return sorted(ha, key=lambda r: (r["ESILevel"], r["ArrivalDateTime"]))


def _nov_sizes() -> tuple[int, int]:
    """Last row of ED_Nov ($A$1:$J$81) and of the HighAcuity list recorded from it (A1:J21)."""
    ed_nov = _ed(11)
    return len(ed_nov) + 1, len(_high_acuity(ed_nov)) + 1


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="05-automation-vba", slug="01-recording-macros",
        title="Recording Your First Macros", level="Expert", minutes=45,
        objectives=[
            "Show the Developer tab and set macro security safely",
            "Record, run, and save macros in a macro-enabled workbook (.xlsm)",
            "Choose between absolute and relative recording",
            "Run macros from buttons, shortcuts, and the Quick Access Toolbar — and read the recorded code",
        ],
        data_note="Raw daily census exports for the Medical-Surgical unit at Cedar Ridge Medical Center (November and "
                  "December 2025), plus the hospital's raw Emergency Department visit exports for the same two months (bonus).",
    )

    # ------------------------------------------------------------------ data
    nov, dec = _census(11), _census(12)
    ed_nov, ed_dec = _ed(11), _ed(12)
    assert len(nov) == 30 and len(dec) == 31

    raw_fmt = {"CensusDate": "General"}   # an export: dates arrive as serial numbers
    cn = L.add_table_sheet("Census_Nov", nov, columns=CENSUS_COLS, as_table=False, formats=raw_fmt)
    cd = L.add_table_sheet("Census_Dec", dec, columns=CENSUS_COLS, as_table=False, formats=raw_fmt)
    en = L.add_table_sheet("ED_Nov", ed_nov, columns=ED_COLS, as_table=False)
    edd = L.add_table_sheet("ED_Dec", ed_dec, columns=ED_COLS, as_table=False)

    # ------------------------------------------------------------------ answers (computed in Python)
    def totals(rows):
        bed_days = sum(r["StaffedBeds"] for r in rows)
        pat_days = sum(r["MidnightCensus"] for r in rows)
        return bed_days, pat_days, pat_days / bed_days

    nov_beds, nov_pd, nov_occ = totals(nov)
    dec_beds, dec_pd, dec_occ = totals(dec)
    nov_total_row = cn.last_row + 1            # 32
    dec_total_row = cd.last_row + 1            # 33
    assert (nov_total_row, dec_total_row) == (32, 33)

    def total_row_values(sd, total_row):
        """The TOTAL row the recorded macro writes (formulas, exactly as the recorder replays them)."""
        last = total_row - 1
        return (["TOTAL", None, None] + [f"=SUM({c}${sd.first_row}:{c}{last})" for c in "DEFG"]
                + [f"=G{total_row}/D{total_row}"])

    nov_fill = {"range": f"Census_Nov!A{nov_total_row}:H{nov_total_row}", "values": total_row_values(cn, nov_total_row)}
    dec_fill = {"range": f"Census_Dec!A{dec_total_row}:H{dec_total_row}", "values": total_row_values(cd, dec_total_row)}

    def total_cell(sheet, col):
        """Summary: find the row the learner's macro labeled TOTAL and read one of its cells."""
        return f'=IFERROR(INDEX({sheet}!{col}:{col},MATCH("TOTAL",{sheet}!A:A,0)),"")'

    # Predict-the-cell answers (absolute vs relative recording)
    start_col, start_row = 1, 1          # A1 is selected when recording starts (relative case)
    click_col, click_row = 2, 3          # the learner clicks B3
    run_col, run_row = 8, 10             # the macro is later run from H10
    abs_answer = "B3"
    rel_col = run_col + (click_col - start_col)
    rel_row = run_row + (click_row - start_row)
    rel_answer = f"{chr(64 + rel_col)}{rel_row}"          # I12

    # Reading R1C1: =SUM(R[-31]C[-1]:R[-1]C[-1]) evaluated with H33 active
    r1c1_top, r1c1_bottom, r1c1_coff, r1c1_row, r1c1_col = -31, -1, -1, dec_total_row, 8
    r1c1_letter = chr(64 + r1c1_col + r1c1_coff)                        # G
    r1c1_answer = f"{r1c1_letter}{r1c1_row + r1c1_top}:{r1c1_letter}{r1c1_row + r1c1_bottom}"   # G2:G32

    L.practice_intro = ("Tasks 1–6, 9 and 13 are quick checks: type your answer in the yellow cell. Tasks 7, 8, 11 and 12 have gray "
                        "cells that read the TOTAL rows your macro writes on Census_Nov and Census_Dec, so you don't type anything "
                        "there. Task 10 has nothing to check: compare your buttons with the answer key. Save this file as an .xlsm "
                        "before you record anything (task 1).")
    L.start_notes = [
        "This workbook is an .xlsx file, so it can't store macros. Before you record, use File → Save As (F12; Mac: ⌘ + Shift + S) "
        "and choose Excel Macro-Enabled Workbook (*.xlsm).",
        "Ctrl + Z (Mac: ⌘ + Z) can't undo a macro. Save before you run a macro you just recorded, so you can close without saving "
        "if it goes wrong.",
        "Reference macros (.bas files) are in the lesson's solutions/ folder on GitHub. They're spoilers, so record your own first.",
    ]
    # Nothing here is a formula, several Practice cells and every Bonus cell are gray, so the generic
    # "Type a formula or value in each yellow cell" lines don't fit.
    L.practice_how = ("Go to the 'Practice' sheet. Type your answers in the yellow cells. The gray cells read the TOTAL rows your "
                      "macro writes on the census sheets.")
    L.practice_instructions = (
        "Type your answer in each yellow cell. The gray cells fill in by themselves once your macro has run. The Check column "
        "turns green when your answer matches. Stuck? Read the hint, then the lesson guide. Answers: right-click a sheet tab → "
        f"Unhide… → '{L.key_sheet}'.")
    L.bonus_instructions = (
        f"There's nothing to type on this sheet. The gray cells read the {HA_SHEET} sheet your macro builds from ED_Dec, and the "
        "Check column turns green when they match. Stuck? Read the hint, then Guide section 9. Answers: right-click a sheet tab → "
        f"Unhide… → '{L.bonus_key_sheet}'.")
    L.bonus_where = ("Record the macro on the **ED_Nov** sheet and run your edited version on **ED_Dec**. The gray cells on the "
                     f"**Bonus** sheet check the {HA_SHEET} sheet it builds.")
    # ED_Nov and ED_Dec are used only by the bonus, so they sit right after the Bonus tab.
    L.sheet_order = ["Start Here", "Practice", "Census_Nov", "Census_Dec", "Bonus", "ED_Nov", "ED_Dec", "Answer Key", "Bonus Key"]
    # Start Here lists each data sheet as "Data: N rows"; the customize hook adds what each export is for.
    sheet_roles = {
        "Census_Nov": "November's raw census export. Record FormatCensusReport here (tasks 7–8).",
        "Census_Dec": "December's raw census export. Add the button and run your macro here (tasks 10–12).",
        "ED_Nov": "November's raw ED export (bonus only). Record ExtractHighAcuity here.",
        "ED_Dec": f"December's raw ED export (bonus only). Your edited macro builds {HA_SHEET} from it.",
    }

    L.tasks = [
        Task("This workbook is an .xlsx file, which can't store macros. Which file extension does the Excel Macro-Enabled Workbook "
             "format use, the format you save it in before you record? Type the extension, like .xlsx",
             answer=".xlsm", accept=["xlsm", "*.xlsm"],
             hint="**File → Save As** shows each format's extension in the Save as type list",
             solution="**.xlsm**: **File → Save As** (F12; Mac: ⌘ + Shift + S) → *Save as type:* **Excel Macro-Enabled Workbook (*.xlsm)**.",
             explanation="An .xlsx file can't contain VBA code. If you save a workbook that has macros as .xlsx, Excel warns you that "
                         "the *VB project* can't be saved, and if you click **Yes** it throws the code away. .xlsb, .xltm, .xlam and the "
                         "old .xls can hold macros too, but .xlsm is the everyday format for a workbook with macros. Save as .xlsm now, "
                         "before you record, so nothing gets lost."),
        Task("A colleague emails you Census_Report.xlsm. When you open it, a red bar says Microsoft has blocked macros because the "
             "source of this file is untrusted. Which is the safe way to run its macros? Type the letter.\n"
             "A: Change the Trust Center to Enable VBA macros.\n"
             "B: Confirm with the colleague that they sent it, save it to your computer, then right-click the file → **Properties** → "
             "tick **Unblock**.\n"
             "C: Click Enable Content on the red bar.\n"
             "D: Rename the file to .xlsx.",
             answer="B", hint="The red bar has no Enable button. Look at the file's Properties",
             solution="**B.** Confirm the sender, save the file, then File Explorer → right-click → **Properties** → **General** → tick "
                      "**Unblock** → **OK**. Reopen it and click **Enable Content** on the yellow bar.",
             explanation="Email attachments and downloads carry the **Mark of the Web**, so Excel blocks their macros outright and the red "
                         "bar has no Enable Content button (C is impossible). Enabling every macro (A) would let any file run code, "
                         "and renaming to .xlsx (D) doesn't run anything because .xlsx can't hold macros. Unblocking one file you've "
                         "verified is the targeted, safe fix."),
        Task("Which function key do you press with Alt (Mac: Option) to open the Visual Basic Editor? Type just the key, like F5.",
             answer="F11", accept=["alt+f11", "alt + f11", "option+f11", "option + f11"],
             hint="It's also on the ribbon: **Developer → Visual Basic**",
             solution="**F11**: Alt + F11 (Mac: Option + F11, or **Developer → Visual Basic**). On a Mac laptop, add Fn if your "
                      "top-row keys control brightness and volume.",
             explanation="Alt + F11 opens the Visual Basic Editor (VBE), where the recorder puts your code. Pressing it again flips back to "
                         "Excel. Alt + F8 (Mac: Option + F8) opens the Macro dialog instead."),
        Task("On Windows, what is the file name (with extension) of the Personal Macro Workbook, the hidden workbook that makes a macro "
             "available in every workbook you open?",
             answer="PERSONAL.XLSB", hint="It's a binary workbook stored in the XLSTART folder",
             solution="**PERSONAL.XLSB**, stored in `%APPDATA%\\Microsoft\\Excel\\XLSTART`.",
             explanation="Choosing *Store macro in: Personal Macro Workbook* creates PERSONAL.XLSB the first time. Excel opens every file in "
                         "XLSTART automatically and hides this one, so its macros are always available on your computer (but not on "
                         "anyone else's)."),
        Task(f"With Use Relative References OFF, you click Record Macro, click cell B3, type Verified, press Enter, and stop recording. "
             f"Later you select H10 and run the macro. Which cell receives Verified? Type the address, like A1.",
             answer=abs_answer, accept=["$B$3"], hint="Absolute recording stores the address you clicked",
             solution='The recorder wrote `Range("B3").Select`, so the macro always goes to **B3**.',
             explanation='Absolute recording (the default) stores exact addresses: Range("B3").Select, then '
                         'ActiveCell.FormulaR1C1 = "Verified", then Range("B4").Select for the Enter key. The cell you had selected '
                         'before running (H10) doesn\'t matter.'),
        Task(f"This time A1 is selected when you click Record Macro. You turn ON Use Relative References, click B3, type Verified, "
             f"press Enter, and stop. Later you select H10 and run the macro. Which cell receives Verified?",
             answer=rel_answer, accept=[f"${rel_answer[0]}${rel_answer[1:]}"],
             hint="Count how far you moved from A1 to B3, then make the same move from H10",
             solution=f'The recorder wrote `ActiveCell.Offset(2, 1).Range("A1").Select`: 2 rows down and 1 column right of the active '
                      f'cell. From H10 that is **{rel_answer}**.',
             explanation="Relative recording stores moves, not addresses. From A1 to B3 is 2 rows down and 1 column right, so the macro "
                         f"makes that same move from wherever it starts. From H10 that lands on {rel_answer}. The Enter key is stored as "
                         "another move, ActiveCell.Offset(1, 0)."),
        Task("Record the FormatCensusReport macro on the Census_Nov sheet by following the recipe in Guide section 7 (TOTAL row recorded "
             "with relative references, shortcut Ctrl + Shift + R, Mac: Option + ⌘ + Shift + R). The gray cell finds the row your macro labeled TOTAL and reads its "
             "MidnightCensus cell, which should equal November's patient days.",
             answer=nov_pd, title="Record FormatCensusReport on Census_Nov (TOTAL patient days)",
             summary=total_cell("Census_Nov", "G"), fill=nov_fill,
             live=f"=SUM(Census_Nov!G{cn.first_row}:G{cn.last_row})",
             hint="Guide section 7 has the click-by-click recipe",
             solution=CENSUS_RECORDED, solution_lang="vba",
             explanation="This is what the recorder writes for the recipe. Your code may differ in small ways, such as an extra "
                         "`Offset` line for each Tab you pressed. The lines that matter are `Range(\"A1\").Select` and "
                         "`Selection.End(xlDown).Select` (recorded with relative references off, so they always start at A1 and "
                         "jump to the last date), then the relative `ActiveCell.Offset(1, 0)` that steps into the empty row below. "
                         f"November has 30 days in rows {cn.first_row}–{cn.last_row}, so TOTAL lands in row {nov_total_row}. "
                         "The second `Range(\"A1\").Select`, near the end, is the Ctrl + Home step. It scrolls back to the top, so "
                         "`.SplitRow = 1` freezes the header row and not whichever row happens to be at the top of the screen."),
        Task("Step 10 of the recipe in Guide section 7 put the month's occupancy in column H of your TOTAL row on Census_Nov: total patient days "
             "divided by total staffed-bed days, formatted as a percentage with 1 decimal place. The gray cell reads that cell.",
             answer=nov_occ, fmt="0.0%", title="Census_Nov TOTAL row occupancy",
             summary=total_cell("Census_Nov", "H"), fill=nov_fill,
             live=f"=SUM(Census_Nov!G{cn.first_row}:G{cn.last_row})/SUM(Census_Nov!D{cn.first_row}:D{cn.last_row})",
             hint="Patient days are in column G, staffed-bed days in column D",
             solution=f"In H{nov_total_row} type `=G{nov_total_row}/D{nov_total_row}`, then click **Home → Percent Style** and "
                      "**Increase Decimal** once. "
                      'The recorder stores it as `"=RC[-1]/RC[-4]"`: same row, 1 and 4 columns to the left.',
             explanation="Total patient days ÷ total staffed-bed days is the unit's occupancy for the month. The formula only refers "
                         "to cells in its own row, so in R1C1 notation it is the same on every sheet and works on December too."),
        Task("In the Record Macro dialog you typed an uppercase R in the Shortcut key box. Which key combination runs the macro on "
             "Windows? Type it like Ctrl + Alt + X.",
             answer="Ctrl + Shift + R",
             accept=["ctrl+shift+r", "shift+ctrl+r", "shift + ctrl + r", "control+shift+r", "control + shift + r", "ctrl shift r",
                     "ctrl-shift-r"],
             hint="An uppercase letter adds a key",
             solution="**Ctrl + Shift + R** (Mac: Option + ⌘ + Shift + R).",
             explanation="A lowercase letter gives Ctrl + letter, which replaces Excel's own shortcut while this workbook is open "
                         "(Ctrl + r would stop doing Fill Right). An uppercase letter adds Shift, which collides with far fewer "
                         "built-in shortcuts. Change it later with **Macros** (Alt + F8; Mac: Option + F8) → select the macro → "
                         "**Options…**"),
        Task("Add two more ways to run FormatCensusReport: (a) a button on the Census_Dec sheet labeled Build report, and (b) a button on "
             "the Quick Access Toolbar. Don't click them yet.",
             answer=None, check="manual", title="Button on the sheet and on the Quick Access Toolbar",
             hint="Right-click a shape → **Assign Macro…**",
             solution=("**Sheet button**\n\n"
                       "1. On Census_Dec, choose **Insert → Shapes → Rectangle: Rounded Corners** and draw it to the right of the data "
                       "(for example over columns J:K).\n"
                       "2. Type **Build report** while the shape is selected.\n"
                       "3. Right-click the shape's border → **Assign Macro…** → pick **FormatCensusReport** → **OK**.\n"
                       "4. Click a cell to deselect. The pointer becomes a hand over the shape, and one click runs the macro.\n\n"
                       "**Quick Access Toolbar** (Windows)\n\n"
                       "1. **File → Options → Quick Access Toolbar**.\n"
                       "2. *Choose commands from:* **Macros** → select **FormatCensusReport** → **Add >>**.\n"
                       "3. Click **Modify…**, pick an icon, set the display name to *Build census report* → **OK** → **OK**.\n\n"
                       "If you can't see the toolbar, right-click the ribbon → **Show Quick Access Toolbar**."),
             explanation="A shape button belongs to one sheet and travels with the workbook. A Quick Access Toolbar button belongs to "
                         "your copy of Excel and remembers which workbook holds the macro, so it opens that workbook if needed. To edit a "
                         "shape later without running the macro, Ctrl + click it (Mac: ⌘ + click) or right-click it."),
        Task("Go to Census_Dec and run your macro (Ctrl + Shift + R, Mac: Option + ⌘ + Shift + R, or your button). December has 31 days. The gray cell finds your "
             "TOTAL row and counts the daily rows above it. It should show all 31 days.",
             answer=len(dec), title="Run on Census_Dec (days above the TOTAL row)",
             summary=f'=IFERROR(MATCH("TOTAL",Census_Dec!A:A,0)-{cd.first_row},"")', fill=dec_fill,
             live=f"=COUNT(Census_Dec!A{cd.first_row}:A{cd.last_row})",
             hint="If it shows 30, your TOTAL row overwrote a day",
             solution=f"Click the **Census_Dec** tab and press **Ctrl + Shift + R** (Mac: **Option + ⌘ + Shift + R**). "
                      f"`Selection.End(xlDown)` finds the last date "
                      f"(row {cd.last_row}) and `ActiveCell.Offset(1, 0)` steps to row {dec_total_row}.",
             explanation=f"If you recorded the TOTAL steps with relative references off, the code says Range(\"A{nov_total_row}\").Select, "
                         f"so on December it types TOTAL over 12/31/2025 in row {nov_total_row} and the gray cell shows 30. Ctrl + ↓ "
                         "(Mac: ⌘ + ↓) plus a relative one-row move finds the first empty row on any month. Macros can't be undone, so close without "
                         "saving (or reopen your saved copy), fix the recording, and run it again."),
        Task("On Census_Dec, the gray cell reads the MidnightCensus total in your TOTAL row. It should equal December's patient days "
             "(all 31 days).",
             answer=dec_pd, title="Census_Dec TOTAL patient days",
             summary=total_cell("Census_Dec", "G"), fill=dec_fill,
             live=f"=SUM(Census_Dec!G{cd.first_row}:G{cd.last_row})",
             hint="Did you anchor the first row with a $ when you typed the SUM?",
             solution=f"The recorded formula `=SUM(R2C:R[-1]C)` becomes `=SUM(G$2:G{cd.last_row})` in row {dec_total_row}.",
             explanation="The recorder stores formulas in R1C1 notation. With =SUM(D$2:D31) the $ makes the top row absolute (R2C) and "
                         "the bottom row relative (R[-1]C, the row above), so the range grows with the month. With AutoSum or "
                         "=SUM(D2:D31) the recorder stores =SUM(R[-30]C:R[-1]C), which always adds exactly 30 rows. On December "
                         "that skips 12/01 and the total comes out low."),
        Task(f"Read this recorded line: ActiveCell.FormulaR1C1 = \"=SUM(R[{r1c1_top}]C[{r1c1_coff}]:R[{r1c1_bottom}]C[{r1c1_coff}])\". "
             f"If it runs while {chr(64 + r1c1_col)}{r1c1_row} is the active cell, which range does the SUM add up? Type it like A1:A9.",
             answer=r1c1_answer,
             accept=[f"${r1c1_letter}${r1c1_row + r1c1_top}:${r1c1_letter}${r1c1_row + r1c1_bottom}"],
             hint="Square brackets count from the formula's own cell: R for rows, C for columns",
             solution=f"**{r1c1_answer}**: R[{r1c1_top}] is {-r1c1_top} rows above row {r1c1_row} (row {r1c1_row + r1c1_top}), "
                      f"R[{r1c1_bottom}] is the row above (row {r1c1_row + r1c1_bottom}), and C[{r1c1_coff}] is one column left of "
                      f"{chr(64 + r1c1_col)}, which is {r1c1_letter}.",
             explanation=f"It's the R1C1 form of =SUM({r1c1_answer}) typed in {chr(64 + r1c1_col)}{r1c1_row}. Square brackets mean "
                         "*relative to the formula's own cell*. A number without brackets is a fixed row or column, so R2C means row 2 "
                         f"of this column. Every part of this formula is relative, so in {chr(64 + r1c1_col)}{r1c1_row + 1} it would add "
                         f"{r1c1_letter}{r1c1_row + r1c1_top + 1}:{r1c1_letter}{r1c1_row + r1c1_bottom + 1}. Reading R1C1 like this tells "
                         "you whether a recorded formula will still be right when the data has a different number of rows."),
    ]

    # ------------------------------------------------------------------ bonus
    ha_dec = _high_acuity(ed_dec)
    ha_nov = _high_acuity(ed_nov)
    first_id, last_id = ha_dec[0]["EDVisitID"], ha_dec[-1]["EDVisitID"]
    n_esi1 = sum(1 for r in ha_dec if r["ESILevel"] == 1)
    # Live formulas below find these IDs by arrival time; make sure those times are unique in the export.
    arrivals = [r["ArrivalDateTime"] for r in ed_dec]
    assert arrivals.count(ha_dec[0]["ArrivalDateTime"]) == 1 and arrivals.count(ha_dec[-1]["ArrivalDateTime"]) == 1
    assert ha_dec[0]["ESILevel"] == 1 and ha_dec[-1]["ESILevel"] == 2
    ef, el = edd.first_row, edd.last_row
    nov_last, nov_ha_last = en.last_row, len(ha_nov) + 1          # $A$1:$J$81 and row 21 in the recording
    assert (nov_last, nov_ha_last) == _nov_sizes()
    ha_edited = _ha_edited(nov_last, nov_ha_last)
    assert len(ed_dec) > len(ed_nov) and len(ha_dec) > len(ha_nov)  # December outgrows every hard-coded range

    # What the checks see when a learner gets the sort wrong (used in the explanations, and to prove the checks catch it).
    ha_export = [r for r in ed_dec if r["ESILevel"] <= 2]                     # filtered rows, still in export order
    by_level_arrival = lambda r: (r["ESILevel"], r["ArrivalDateTime"])      # noqa: E731
    partial = sorted(ha_export[:len(ha_nov)], key=by_level_arrival) + ha_export[len(ha_nov):]   # sort range stuck at row 21
    arrival_only = sorted(ha_export, key=lambda r: r["ArrivalDateTime"])     # ESILevel sort level left out
    esi2_row = n_esi1 + 2                                                    # first ESI 2 visit (row 5)
    esi2_id = ha_dec[n_esi1]["EDVisitID"]
    assert ha_dec[n_esi1]["ESILevel"] == 2 and ha_dec[n_esi1 - 1]["ESILevel"] == 1
    assert arrivals.count(ha_dec[n_esi1]["ArrivalDateTime"]) == 1
    assert partial[0]["EDVisitID"] != first_id and partial[-1]["EDVisitID"] != last_id
    assert partial[n_esi1]["EDVisitID"] != esi2_id and arrival_only[n_esi1]["EDVisitID"] != esi2_id
    arrival_only_same_ends = (arrival_only[0]["EDVisitID"], arrival_only[-1]["EDVisitID"]) == (first_id, last_id)

    ha_range = f"'{HA_SHEET}'!A:A"
    # ISREF(INDIRECT(...)) is FALSE until the learner's macro creates the sheet. Don't test with COUNTA(INDIRECT(...)):
    # Excel's COUNTA counts the #REF! error value as 1 (LibreOffice propagates it), which would pre-answer checks.
    ha_exists = f'ISREF(INDIRECT("\'{HA_SHEET}\'!A1"))'
    ha_fill = {"range": f"{HA_SHEET}!A1:A1", "values": ["EDVisitID"]}   # the hook builds the sheet; this just exercises it

    assert not {r["EDVisitID"] for r in ed_nov} & {r["EDVisitID"] for r in ed_dec}   # the guard below relies on this

    def ha_from_dec(expr):
        """Show expr only once HighAcuity exists AND was built from December (its A2 is a December EDVisitID), so a
        HighAcuity sheet left over from recording on ED_Nov can't turn B5 green before the macro has touched ED_Dec."""
        return (f'=IFERROR(IF({ha_exists},IF(COUNTIF(ED_Dec!A{ef}:A{el},INDIRECT("\'{HA_SHEET}\'!A2")&"")>0,{expr},""),""),"")')

    L.bonus_title = "Bonus: The high-acuity ED list"
    L.bonus_scenario = (
        "Cedar Ridge's ED medical director reviews every high-acuity visit (ESI 1 or 2) each month. The tracking system exports the "
        "month's visits sorted by PatientID, and someone spends 20 minutes turning that export into a list. Automate it. "
        "The macro must:\n\n"
        "1. Filter the export's ESILevel column to values less than or equal to 2.\n"
        f"2. Copy the visible rows, with the header, to a new sheet named {HA_SHEET}.\n"
        "3. Sort that sheet by ESILevel (smallest first), then ArrivalDateTime (oldest first), and AutoFit the columns.\n"
        "4. Go back to the export and turn its filter off.\n\n"
        "Recording tips: Check that Use Relative References is off. Click A1 and turn the filter on with **Data → Filter**, then "
        "use the ESILevel filter arrow → **Number Filters → Less Than Or Equal To** → 2 → **OK**. Select the data with "
        "Ctrl + Shift + → and then Ctrl + Shift + ↓ (Mac: ⌘ + Shift + arrows), and copy it with Ctrl + C (Mac: ⌘ + C). Add a "
        "sheet with the + (New "
        "sheet) button next to the sheet tabs, paste with Ctrl + V (Mac: ⌘ + V), and rename the sheet by double-clicking its tab. "
        "In **Data → Sort** (Lesson 1.6), sort by ESILevel (Smallest to Largest), then click **Add Level** and pick ArrivalDateTime "
        "(Oldest to Newest). "
        "AutoFit every column: click the Select All button (the triangle where the row and column headings meet), then "
        "double-click any column boundary. Finish by "
        "clicking the ED_Nov tab, then A1, then **Data → Filter** again.\n\n"
        f"Record it on ED_Nov ({len(ed_nov)} visits) as ExtractHighAcuity (shortcut Ctrl + Shift + H, Mac: "
        f"Option + ⌘ + Shift + H). Then delete the {HA_SHEET} sheet it made (right-click its tab → **Delete**), open the VBE with "
        "Alt + F11 (Mac: Option + F11, or **Developer → Visual Basic**), "
        f"and edit the code so it works on ED_Dec ({len(ed_dec)} visits). Guide section 9 lists the edits recorded code usually "
        "needs and what to do if the macro stops with a run-time error. Run it on ED_Dec. The gray cells read the "
        f"{HA_SHEET} sheet your macro builds from December, so they stay blank until it exists. While the sheet from your "
        "November recording is still there, they show ✘ Not yet or stay blank."
    )
    L.bonus = [
        Task(f"How many visits (data rows, not counting the header) are on the {HA_SHEET} sheet built from ED_Dec?",
             answer=len(ha_dec), title=f"Visits on {HA_SHEET} (from ED_Dec)",
             summary=f'=IF({ha_exists},COUNTA(INDIRECT("{ha_range}"))-1,"")', fill=ha_fill,
             live=f'=COUNTIF(ED_Dec!E{ef}:E{el},"<=2")',
             hint="COUNTIF on ED_Dec's ESILevel column gives you a target to compare with",
             solution=ha_edited, solution_lang="vba",
             explanation=(
                 "Recording on ED_Nov bakes November into the code in four places. Edit each one:\n\n"
                 f"1. **The filter range.** `ActiveSheet.Range(\"$A$1:$J${nov_last}\").AutoFilter` names November's exact block, but "
                 f"December runs to row {el}. Change it to `Range(\"A1\").CurrentRegion`, the whole block around A1, which grows and "
                 "shrinks with the export. The unedited line may still work here, but only because the line above it already switched "
                 "the filter on for the whole block. Had the filter been on before you started recording, the recorder would have "
                 f"skipped `Selection.AutoFilter`, and on December the hard-coded line would filter rows 1–{nov_last} only.\n"
                 f"2. **The rename.** `Sheets(\"Sheet1\").Select` and `Sheets(\"Sheet1\").Name = \"{HA_SHEET}\"` use the new sheet's "
                 "temporary name. The next new sheet may be Sheet2, and then the macro stops with run-time error 9 (Subscript out of "
                 f"range). The new sheet is active right after it's added, so use `ActiveSheet.Name = \"{HA_SHEET}\"`.\n"
                 f"3. **The sort ranges.** `Key:=Range(\"E2:E{nov_ha_last}\")`, `Key:=Range(\"C2:C{nov_ha_last}\")` and "
                 f"`.SetRange Range(\"A1:J{nov_ha_last}\")` fit November's {len(ha_nov)} visits. Use `Range(\"E1\")`, `Range(\"C1\")` "
                 "and `Range(\"A1\").CurrentRegion`. One cell is enough to name a sort key's column.\n"
                 "4. **The last sheet.** `Sheets(\"ED_Nov\").Select` sends the macro back to November, where `Selection.AutoFilter` "
                 "switches a filter on, and ED_Dec stays filtered. Change it to `Sheets(\"ED_Dec\")`.\n\n"
                 "Adding `Sheets(\"ED_Dec\").Select` as the first line makes the macro start on the right sheet whichever sheet is "
                 "active. Both versions, as recorded and as edited, are in `solutions/modHighAcuity.bas`. Lesson 5.3 replaces the "
                 "remaining Select lines and hard-coded sheet names with variables.")),
        Task(f"Which EDVisitID is in the first data row (A2) of {HA_SHEET}?",
             answer=first_id, title=f"First EDVisitID on {HA_SHEET}",
             summary=f'=IFERROR(INDIRECT("\'{HA_SHEET}\'!A2")&"","")', fill=ha_fill,
             live=f"=INDEX(ED_Dec!A{ef}:A{el},MATCH(MINIFS(ED_Dec!C{ef}:C{el},ED_Dec!E{ef}:E{el},1),ED_Dec!C{ef}:C{el},0))",
             hint="It should be the earliest ESI 1 arrival in December",
             solution=f"**{first_id}**, the earliest-arriving ESI 1 visit ({ha_dec[0]['ArrivalDateTime']:%m/%d/%Y %H:%M}).",
             explanation=f"Sorting by ESILevel first puts the {n_esi1} ESI 1 visits on top. The second sort level, ArrivalDateTime, orders "
                         f"visits within each level. If you see a different ID, the sort range probably stopped at row {nov_ha_last} "
                         f"(November's size), so only part of the list was sorted. In that case A2 shows {partial[0]['EDVisitID']}. "
                         "The live formula in the key finds the right visit with MINIFS: the earliest arrival among ESI 1 visits."),
        Task(f"Which EDVisitID is in cell A{esi2_row} of {HA_SHEET}?",
             answer=esi2_id, title=f"EDVisitID in A{esi2_row} of {HA_SHEET}",
             summary=f'=IFERROR(INDIRECT("\'{HA_SHEET}\'!A{esi2_row}")&"","")', fill=ha_fill,
             live=f"=INDEX(ED_Dec!A{ef}:A{el},MATCH(MINIFS(ED_Dec!C{ef}:C{el},ED_Dec!E{ef}:E{el},2),ED_Dec!C{ef}:C{el},0))",
             hint="This cell depends on both sort levels",
             solution=f"**{esi2_id}**, the earliest-arriving ESI 2 visit ({ha_dec[n_esi1]['ArrivalDateTime']:%m/%d/%Y %H:%M}).",
             explanation=f"The {n_esi1} ESI 1 visits fill rows 2–{esi2_row - 1}, so row {esi2_row} holds the first ESI 2 visit, which is "
                         "the earliest ESI 2 arrival. This cell shows whether both sort levels worked. "
                         + (f"Sorted by ArrivalDateTime alone, the first and last rows happen to be the same as the correct list, "
                            f"but A{esi2_row} shows {arrival_only[n_esi1]['EDVisitID']}. "
                            if arrival_only_same_ends else
                            f"Sorted by ArrivalDateTime alone, A{esi2_row} shows {arrival_only[n_esi1]['EDVisitID']}. ")
                         + f"With the sort range stuck at row {nov_ha_last}, it shows {partial[n_esi1]['EDVisitID']}."),
        Task(f"Which EDVisitID is in the last data row of {HA_SHEET}?",
             answer=last_id, title=f"Last EDVisitID on {HA_SHEET}",
             summary=f'=IFERROR(IF({ha_exists},INDEX(INDIRECT("{ha_range}"),MAX(1,COUNTA(INDIRECT("{ha_range}"))))&"",""),"")',
             fill=ha_fill,
             live=f"=INDEX(ED_Dec!A{ef}:A{el},MATCH(MAXIFS(ED_Dec!C{ef}:C{el},ED_Dec!E{ef}:E{el},2),ED_Dec!C{ef}:C{el},0))",
             hint=f"Rows below row {nov_ha_last} are only sorted if the sort range adapts",
             solution=f"**{last_id}**, the latest-arriving ESI 2 visit ({ha_dec[-1]['ArrivalDateTime']:%m/%d/%Y %H:%M}).",
             explanation=f"December has {len(ha_dec)} high-acuity visits, so the list ends in row {len(ha_dec) + 1}. A sort range "
                         f"recorded as A1:J{nov_ha_last} leaves the last {len(ha_dec) - len(ha_nov)} rows in export (PatientID) order, "
                         f"so the last row shows {partial[-1]['EDVisitID']} instead. `Range(\"A1\").CurrentRegion` always covers "
                         "every row."),
        Task("After your macro finishes, how many rows of ED_Dec are visible? The gray cell counts them with SUBTOTAL(103, …), which "
             f"skips rows a filter hides. It should be every December visit. It stays blank until {HA_SHEET} holds December visits.",
             answer=len(ed_dec), title="ED_Dec filter turned off at the end",
             summary=ha_from_dec(f"SUBTOTAL(103,ED_Dec!A{ef}:A{el})"), fill=ha_fill,
             live=f"=COUNTA(ED_Dec!A{ef}:A{el})",
             hint="Which sheet does the end of your recorded code go back to?",
             solution=f"End the macro with `Sheets(\"ED_Dec\").Select`, `Range(\"A1\").Select`, `Selection.AutoFilter`. "
                      f"That toggles the filter off, so all {len(ed_dec)} visits show again.",
             explanation=f"The recording ends with Sheets(\"ED_Nov\").Select, so on December it toggles a filter ON on ED_Nov and "
                         f"leaves ED_Dec showing only {len(ha_dec)} rows. A good macro leaves the source data the way it found it. "
                         "Selection.AutoFilter is a toggle (like **Data → Filter**), so it turns the filter off only when one is on."),
    ]

    # ------------------------------------------------------------------ workbook polish + self-test simulation
    @L.customize
    def _raw_exports(wb, lesson, selftest):
        from copy import copy

        # Start Here: add each export's role to its "Data: N rows" line.
        start = wb["Start Here"]
        done = set()
        for r in range(1, start.max_row + 1):
            name, desc = start.cell(row=r, column=2).value, start.cell(row=r, column=3)
            if name in sheet_roles and isinstance(desc.value, str) and desc.value.startswith("Data:"):
                desc.value = f"{desc.value}. {sheet_roles[name]}"
                start.row_dimensions[r].height = 15 * estimate_lines(desc.value, 100) + 3
                done.add(name)
        assert done == set(sheet_roles), done

        plain = Font(name="Calibri", size=11)
        for name in ("Census_Nov", "Census_Dec", "ED_Nov", "ED_Dec"):
            ws = wb[name]
            ws.freeze_panes = None                     # an export arrives with no frozen panes
            for c in ws[1]:
                c.font = plain
                c.fill = PatternFill(fill_type=None)
        for name in ("Census_Nov", "Census_Dec"):
            ws = wb[name]
            for col in "ABCDEFGH":                     # narrow default-ish widths: AutoFit is part of the macro
                ws.column_dimensions[col].width = 10
        if selftest:
            # Simulate the bonus macro: HighAcuity = ED_Dec rows with ESI <= 2, sorted by ESILevel then ArrivalDateTime.
            src = wb["ED_Dec"]
            ha = wb.create_sheet(HA_SHEET, index=wb.sheetnames.index("ED_Dec") + 1)
            for j, h in enumerate(ED_COLS, 1):
                ha.cell(row=1, column=j, value=h)
            for i, r in enumerate(ha_dec, 2):
                for j, h in enumerate(ED_COLS, 1):
                    c = ha.cell(row=i, column=j, value=r[h])
                    c.number_format = copy(src.cell(row=2, column=j).number_format)

    L.extra_files = {
        "solutions/modCensusReport.bas": _crlf(_module("modCensusReport", SPOILER, [CENSUS_RECORDED, CENSUS_CLEAN])),
        "solutions/modHighAcuity.bas": _crlf(_module("modHighAcuity", SPOILER,
                                                      [_ha_recorded(nov_last, nov_ha_last), ha_edited])),
    }
    return L


# ---------------------------------------------------------------------------
# LibreOffice smoke test of the reference macros (best effort; see xlcourse/vba.py)
# ---------------------------------------------------------------------------
def smoke_test():
    """Run the reference macros in LibreOffice and print what they wrote. Returns the vba.run results."""
    import re
    import tempfile
    from pathlib import Path

    from xlcourse import vba

    L = build()
    wb = L.dir / L.workbook_name
    sol = L.dir / "solutions"
    helpers = "\n".join(f"""Sub Go_{n}()
    Worksheets("{n}").Activate
End Sub""" for n in ("Census_Nov", "Census_Dec", "ED_Nov", "ED_Dec"))
    cells = [(s, f"{c}{r}") for s in ("Census_Nov", "Census_Dec") for r in (31, 32, 33, 34) for c in "ADGH"]
    results = {}
    for macro in ("FormatCensusReport", "FormatCensusReport_Clean"):
        res = vba.run(wb, [sol / "modCensusReport.bas"], ["Go_Census_Nov", macro, "Go_Census_Dec", macro],
                      read=cells, extra_code=helpers)
        results[macro] = res
        print(macro, "errors:", res["errors"])
        for k, v in res["cells"].items():
            print("   ", k, v)
    # LibreOffice has no Worksheet.Sort object (Excel 2007+), so the bonus macro is also run with its Sort block
    # swapped for the older Range.Sort call, which sorts the same way. Probe records the ED_Dec filter state.
    probe = helpers + """
Sub Probe()
    Worksheets("HighAcuity").Range("M1").Value = Worksheets("ED_Dec").AutoFilterMode
    Worksheets("HighAcuity").Range("M2").Value = Application.WorksheetFunction.Subtotal(103, Worksheets("ED_Dec").Range("A2:A200"))
End Sub"""
    ha_cells = [(HA_SHEET, a) for a in ("A1", "A2", "E2", "E4", "A5", "E5", "A27", "E27", "A28", "M1", "M2")]
    edited = _ha_edited(*_nov_sizes())
    lo_sort = re.sub(r"    ActiveWorkbook\.Worksheets\(\"HighAcuity\"\)\.Sort\.SortFields\.Clear.*?End With\n",
                     '    Range("A1").CurrentRegion.Sort Key1:=Range("E1"), Order1:=xlAscending, '
                     'Key2:=Range("C1"), Order2:=xlAscending, Header:=xlYes\n', edited, flags=re.S)
    assert lo_sort != edited
    tmp = Path(tempfile.mkdtemp(prefix="smoke51_"))
    for label, code in (("ExtractHighAcuity (as written)", edited), ("ExtractHighAcuity (Range.Sort swap)", lo_sort)):
        bas = tmp / "variant.bas"      # vba.run treats strings containing "/" as paths, so pass a file
        bas.write_text(code, encoding="utf-8")
        res = vba.run(wb, [bas], ["ExtractHighAcuity", "Probe"], read=ha_cells, extra_code=probe)
        results[label] = res
        print(label, "errors:", res["errors"], "sheets:", res["sheets"])
        for k, v in res["cells"].items():
            print("   ", k, v)
    return results
