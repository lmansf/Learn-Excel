"""Lesson 4.1 · Dynamic Arrays: FILTER, SORT, UNIQUE & More.

Spill exercises live on a 'Workspace' sheet (built in a customize hook) so that no spilling formula ever
sits in a Practice answer cell. Gray summary cells on the Practice sheet condense each spill into one
checkable value.

LibreOffice (the verifier) differs from Excel in three ways that matter here, so the SELF-TEST copy
(never the shipped workbook) uses equivalent rewrites:
  * a bare table name (tblEncounters) includes the header row in LibreOffice; Excel means the data body.
    The self-test writes tblEncounters[#Data], which is identical in Excel.
  * the spill-range operator (N6#, stored as _xlfn.ANCHORARRAY) is not supported, so the self-test
    swaps each X# for the explicit range the spill occupies (its size is known from Python).
  * XLOOKUP does not vectorize an array of lookup values, so task 11 is verified through an
    equivalent INDEX/XMATCH live formula instead of the self-test.
"""
from __future__ import annotations

import random
import re
from collections import Counter, defaultdict
from statistics import mean

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from xlcourse import Lesson, Task, data
from xlcourse.lesson import INPUT_BORDER, INPUT_FILL, NAVY
from xlcourse.xlfn import to_file_formula

CODE = "4.1"
MONTH_NAMES = ["", "January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
               "November", "December"]
SEED = 41
SAMPLE_SIZE = 2000
WS = "Workspace"

ENC_COLUMNS = ["EncounterID", "PatientID", "AdmitDate", "DischargeDate", "LOSDays", "EncounterType", "Facility",
               "Department", "ServiceLine", "Attending", "DxCode", "DxCategory", "Payer", "PayerType",
               "TotalCharges", "Readmit30"]


def col_no(name: str) -> int:
    """1-based position of a column in tblEncounters (what CHOOSECOLS needs)."""
    return ENC_COLUMNS.index(name) + 1


# ---------------------------------------------------------------------------- data
def encounter_rows() -> list[dict]:
    """A seeded, representative sample of 2,000 encounters from 2025 with names pre-joined."""
    fac = {f["FacilityID"]: f["FacilityName"] for f in data.load("facilities")}
    dep = {d["DeptID"]: d for d in data.load("departments")}
    prov = {p["ProviderID"]: p for p in data.load("providers")}
    pay = {p["PayerID"]: p for p in data.load("payers")}
    dx = {d["DxCode"]: d for d in data.load("diagnoses")}
    e25 = sorted((e for e in data.load("encounters") if e["AdmitDateTime"].year == 2025), key=lambda e: e["EncounterID"])
    sample = random.Random(SEED).sample(e25, SAMPLE_SIZE)
    sample.sort(key=lambda e: (e["AdmitDateTime"], e["EncounterID"]))
    rows = []
    for e in sample:
        p = prov[e["AttendingProviderID"]]
        admit, disch = e["AdmitDateTime"].date(), e["DischargeDateTime"].date()
        rows.append({
            "EncounterID": e["EncounterID"], "PatientID": e["PatientID"], "AdmitDate": admit, "DischargeDate": disch,
            "LOSDays": (disch - admit).days, "EncounterType": e["EncounterType"], "Facility": fac[e["FacilityID"]],
            "Department": dep[e["DeptID"]]["DeptName"], "ServiceLine": dep[e["DeptID"]]["ServiceLine"],
            "Attending": f"{p['LastName']}, {p['FirstName']}", "DxCode": e["PrimaryDxCode"],
            "DxCategory": dx[e["PrimaryDxCode"]]["DxCategory"], "Payer": pay[e["PayerID"]]["PayerName"],
            "PayerType": pay[e["PayerID"]]["PayerType"], "TotalCharges": e["TotalCharges"], "Readmit30": e["Readmit30"],
        })
    return rows


def provider_rows() -> list[dict]:
    fac = {f["FacilityID"]: f["FacilityName"] for f in data.load("facilities")}
    dep = {d["DeptID"]: d["DeptName"] for d in data.load("departments")}
    out = []
    for p in sorted(data.load("providers"), key=lambda p: p["ProviderID"]):
        out.append({"ProviderID": p["ProviderID"], "Provider": f"{p['LastName']}, {p['FirstName']}",
                    "Credential": p["Credential"], "Specialty": p["Specialty"], "Facility": fac[p["FacilityID"]],
                    "Department": dep[p["PrimaryDeptID"]], "HireDate": p["HireDate"], "FTE": p["FTE"]})
    return out


def department_rows() -> list[dict]:
    """Departments with StaffedBeds = 0 (not blank) where a department has no beds.

    The guide sorts tblDepartments by StaffedBeds largest first. Excel's SORT/SORTBY don't rank an empty cell as 0,
    so blanks could land above the real numbers in a descending sort; a real 0 keeps every example deterministic.
    """
    fac = {f["FacilityID"]: f["FacilityName"] for f in data.load("facilities")}
    return [{"DeptID": d["DeptID"], "Department": d["DeptName"], "Facility": fac[d["FacilityID"]],
             "ServiceLine": d["ServiceLine"], "UnitType": d["UnitType"], "StaffedBeds": d["StaffedBeds"] or 0}
            for d in sorted(data.load("departments"), key=lambda d: d["DeptID"])]


# ---------------------------------------------------------------------------- Workspace layout
# anchor -> (label above the exercise, short note under the label)
ANCHORS = {
    "B6": ("Task 1 · UNIQUE", "Distinct payers ↓"),
    "D6": ("Task 2 · SORT + UNIQUE", "Service lines, A→Z ↓"),
    "F6": ("Task 12 · SEQUENCE", "Month starts ↓"),
    "G6": ("", "Encounters ↓"),
    "I6": ("Task 13 · Long-stay worklist", "Header row, then data ↓ (4 columns)"),
    "N6": ("Bonus · Unit leaderboard", "Header row, then data ↓ (4 columns)"),
}
WIDTHS = {"A": 2, "B": 31, "C": 2, "D": 26, "E": 2, "F": 12, "G": 12, "H": 2, "I": 13, "J": 26, "K": 10, "L": 14,
          "M": 2, "N": 32, "O": 32, "P": 12, "Q": 14}
OBSTRUCTION = ("D11", "draft")
LAST_WS_ROW = 60  # summaries read anchor:row60, far more than any spill here needs

LEADERBOARD_MIN = 50
HEADER_WORKLIST = '{"EncounterID","Department","LOSDays","TotalCharges"}'
HEADER_BOARD = '{"Facility","Department","Encounters","AvgCharge"}'


def lo_compat(formula: str, sizes: dict[str, tuple[int, int]]) -> str:
    """Self-test only: rewrite Excel syntax LibreOffice can't evaluate into an exact equivalent."""
    f = re.sub(r"\btblEncounters\b(?!\[)", "tblEncounters[#Data]", formula)
    for anchor, (nrows, ncols) in sizes.items():
        f = f.replace(f"{WS}!{anchor}#", f"{WS}!{anchor}:{_end_cell(anchor, nrows, ncols)}")
    if "#" in f.replace("[#Data]", ""):
        raise ValueError(f"unresolved spill reference in {formula}")
    return f


def _col_idx(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def _end_cell(anchor: str, nrows: int, ncols: int) -> str:
    col = re.match(r"[A-Z]+", anchor).group(0)
    row = int(anchor[len(col):])
    return f"{get_column_letter(_col_idx(col) + ncols - 1)}{row + nrows - 1}"


# ---------------------------------------------------------------------------- build
def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="04-advanced-analysis", slug="01-dynamic-arrays",
        title="Dynamic Arrays: FILTER, SORT, UNIQUE & More", level="Advanced", minutes=60,
        objectives=[
            "Understand spilling, the # spill reference, and #SPILL! errors",
            "Extract lists with UNIQUE, FILTER, SORT, and SORTBY",
            "Generate sequences with SEQUENCE and reshape with TAKE, DROP, CHOOSECOLS, VSTACK, and HSTACK",
            "Combine functions into one-formula reports",
        ],
        data_note="2,000 encounters sampled from Bluestone Health System's 2025 activity at all four facilities, with "
                  "facility, department, attending, diagnosis category, and payer names already joined in. Also includes "
                  "the provider roster (147) and the department list (31, with StaffedBeds 0 for departments that have "
                  "no inpatient beds).",
    )

    enc = encounter_rows()
    provs = provider_rows()
    depts = department_rows()
    E = L.add_table_sheet("Encounters", enc, table="tblEncounters", columns=ENC_COLUMNS,
                          formats={"TotalCharges": "#,##0.00", "LOSDays": "0"},
                          widths={"Facility": 30, "Department": 30, "Payer": 28, "DxCategory": 24, "Attending": 21})
    L.add_table_sheet("Providers", provs, table="tblProviders",
                      columns=["ProviderID", "Provider", "Credential", "Specialty", "Facility", "Department", "HireDate", "FTE"],
                      formats={"FTE": "0.0"}, widths={"Facility": 30, "Department": 30, "Specialty": 26})
    L.add_table_sheet("Departments", depts, table="tblDepartments",
                      columns=["DeptID", "Department", "Facility", "ServiceLine", "UnitType", "StaffedBeds"],
                      widths={"Department": 32, "Facility": 30})
    assert E.n == SAMPLE_SIZE
    L.sheet_order = ["Start Here", "Practice", WS, "Encounters", "Providers", "Departments", "Bonus",
                     "Answer Key", "Bonus Key"]
    L.start_notes = [
        "Spill formulas go on the Workspace sheet. Type each one in its yellow anchor cell and leave the cells below "
        "and to the right empty so the results have room to spill.",
        "Needs Microsoft 365 or Excel 2024 (VSTACK, HSTACK, TAKE, DROP, CHOOSECOLS). FILTER, SORT, SORTBY, UNIQUE, "
        "SEQUENCE, and XLOOKUP also work in Excel 2021.",
    ]

    # ------------------------------------------------------------------ answers (computed in Python)
    def where(**conds):
        return [r for r in enc if all(r[k] == v for k, v in conds.items())]

    payers = list(dict.fromkeys(r["Payer"] for r in enc))
    service_lines = sorted(set(r["ServiceLine"] for r in enc))
    units = list(dict.fromkeys((r["Facility"], r["Department"]) for r in enc))
    patient_counts = Counter(r["PatientID"] for r in enc)
    once = sum(1 for c in patient_counts.values() if c == 1)
    big_inpatient_units = [d for d in depts if d["UnitType"] == "Inpatient" and d["StaffedBeds"] >= 20]
    big_icus = [d for d in depts if d["UnitType"] == "Critical Care" and d["StaffedBeds"] >= 20]
    assert len(big_icus) == 1, "task 5's explanation names exactly one 20+ bed critical-care unit"
    ed = where(EncounterType="Emergency")
    ed_attendings = set(r["Attending"] for r in ed)
    flu = [r for r in ed if r["AdmitDate"].month <= 2 and r["DxCategory"] in ("Respiratory", "Infectious")]
    ed_top5 = sum(sorted((r["TotalCharges"] for r in ed), reverse=True)[:5])
    ip = where(EncounterType="Inpatient")
    ip_sorted = sorted((r["TotalCharges"] for r in ip), reverse=True)
    trimmed_avg = mean(ip_sorted[10:])
    by_attending = defaultdict(float)
    for r in enc:
        by_attending[r["Attending"]] += r["TotalCharges"]
    ranked_att = sorted(by_attending.items(), key=lambda kv: -kv[1])
    assert ranked_att[0][1] - ranked_att[1][1] > 1, "tie at the top of the attending ranking"
    specialty = {p["Provider"]: p["Specialty"] for p in provs}
    assert len(specialty) == len(provs), "provider names must be unique for the XLOOKUP task"
    ip_specialties = set(specialty[r["Attending"]] for r in ip)
    month_counts = Counter(r["AdmitDate"].month for r in enc)
    assert sum(month_counts.values()) == SAMPLE_SIZE and len(month_counts) == 12
    worklist = sorted([r for r in enc if r["Facility"] == "Cedar Ridge Medical Center" and r["LOSDays"] >= 7],
                      key=lambda r: -r["LOSDays"])
    los_vals = [r["LOSDays"] for r in worklist]
    assert los_vals.count(max(los_vals)) == 1 and los_vals.count(min(los_vals)) == 1, "worklist ends must be unique"
    los_ties = [v for v, c in Counter(los_vals).items() if c > 1]
    assert len(los_ties) == 1 and los_vals.count(los_ties[0]) == 2, "task 13's explanation describes exactly one 2-way tie"
    busiest_month, busiest_n = month_counts.most_common(1)[0]
    assert list(month_counts.values()).count(busiest_n) == 1

    unit_n = Counter((r["Facility"], r["Department"]) for r in enc)
    unit_sum = defaultdict(float)
    for r in enc:
        unit_sum[(r["Facility"], r["Department"])] += r["TotalCharges"]
    board = sorted(((f, d, unit_n[(f, d)], unit_sum[(f, d)] / unit_n[(f, d)]) for (f, d) in units
                    if unit_n[(f, d)] >= LEADERBOARD_MIN), key=lambda x: -x[3])
    assert board[0][3] - board[1][3] > 1
    mem_ed = ("Bluestone Memorial Hospital", "Emergency Department")
    mem_ed_avg = unit_sum[mem_ed] / unit_n[mem_ed]
    board_share = sum(unit_sum[(f, d)] for f, d, _, _ in board) / sum(r["TotalCharges"] for r in enc)
    unit_type = {(d["Facility"], d["Department"]): d["UnitType"] for d in depts}
    off_board = [u for u in units if unit_n[u] < LEADERBOARD_MIN]
    assert sum(unit_type[u] in ("Inpatient", "Critical Care") for u in off_board) > len(off_board) / 2, \
        "B4's explanation says most off-board units are ICUs and inpatient floors"
    assert all(unit_type[(f, d)] == "Inpatient" for f, d, _, _ in board[:5]), "B2's explanation: inpatient units lead"

    # spill sizes (rows, cols) — used to resolve X# references in the self-test copy
    sizes = {"B6": (len(payers), 1), "D6": (len(service_lines), 1), "F6": (12, 1), "G6": (12, 1),
             "I6": (len(worklist) + 1, 4), "N6": (len(board) + 1, 4)}

    # ------------------------------------------------------------------ formulas
    TE = "tblEncounters"
    is_ed = f'{TE}[EncounterType]="Emergency"'
    is_ip = f'{TE}[EncounterType]="Inpatient"'
    f_payers = f"=UNIQUE({TE}[Payer])"
    f_lines = f"=SORT(UNIQUE({TE}[ServiceLine]))"
    f_months = "=DATE(2025,SEQUENCE(12),1)"
    f_counts = (f'=COUNTIFS({TE}[AdmitDate],">="&{WS}!F6#,{TE}[AdmitDate],"<"&DATE(YEAR({WS}!F6#),MONTH({WS}!F6#)+1,1))')
    f_counts_learner = f_counts.replace(f"{WS}!", "")  # what a learner types on the Workspace sheet itself
    cedar_long = f'({TE}[Facility]="Cedar Ridge Medical Center")*({TE}[LOSDays]>=7)'
    f_worklist = (f"=VSTACK({HEADER_WORKLIST},SORT(CHOOSECOLS(FILTER({TE},{cedar_long}),"
                  f"{col_no('EncounterID')},{col_no('Department')},{col_no('LOSDays')},{col_no('TotalCharges')}),3,-1))")
    unit_crit = f"{TE}[Facility],fac,{TE}[Department],dept"
    board_core = (f"LET(units,UNIQUE({TE}[[Facility]:[Department]]),fac,CHOOSECOLS(units,1),dept,CHOOSECOLS(units,2),"
                  f"n,COUNTIFS({unit_crit}),avg,AVERAGEIFS({TE}[TotalCharges],{unit_crit}),"
                  f"VSTACK({HEADER_BOARD},SORT(FILTER(HSTACK(units,n,avg),n>={LEADERBOARD_MIN}),4,-1)))")
    f_board = "=" + board_core
    board_pretty = (
        "=LET(units, UNIQUE(tblEncounters[[Facility]:[Department]]),\n"
        "     fac,   CHOOSECOLS(units, 1),\n"
        "     dept,  CHOOSECOLS(units, 2),\n"
        "     n,     COUNTIFS(tblEncounters[Facility], fac, tblEncounters[Department], dept),\n"
        "     avg,   AVERAGEIFS(tblEncounters[TotalCharges], tblEncounters[Facility], fac, tblEncounters[Department], dept),\n"
        f"     VSTACK({HEADER_BOARD},\n"
        f"            SORT(FILTER(HSTACK(units, n, avg), n >= {LEADERBOARD_MIN}), 4, -1)))"
    )

    def ws_rng(anchor):
        col = re.match(r"[A-Z]+", anchor).group(0)
        return f"{WS}!{anchor}:{col}{LAST_WS_ROW}"

    # ------------------------------------------------------------------ practice tasks
    spill_fills = {}  # task -> (anchor, formula) written as a real multi-cell array in the self-test copy

    def ws_task(anchors_formulas, **kw):
        t = Task(**kw)
        first_anchor = anchors_formulas[0][0]
        t.fill = {"range": f"{WS}!{first_anchor}:{first_anchor}", "values": []}  # no-op marker; see hook
        spill_fills[id(t)] = anchors_formulas
        return t

    L.practice_intro = (
        "Tasks 1, 2, 12, and 13 are spill exercises: build them in the yellow anchor cells on the Workspace sheet, and the "
        "gray cells here read your results. Every other answer cell needs a formula that returns ONE value, so wrap "
        "spilling functions in ROWS, SUM, AVERAGE, INDEX, or TAKE(…,1). Refer to the data with structured references "
        "such as tblEncounters[TotalCharges]."
    )
    t1 = ws_task(
        [("B6", f_payers)],
        prompt="Go to the Workspace sheet. In the yellow cell B6, enter one formula that spills the list of distinct payer "
               "names from the Payer column of tblEncounters. The gray cell counts the names in your list.",
        answer=len(payers), title="Distinct payers (Workspace!B6)",
        solution=f_payers, hint="UNIQUE of one Table column",
        summary=f'=IF(ISBLANK({WS}!B6),"",COUNTA({ws_rng("B6")}))',
        live=f"=ROWS(UNIQUE({TE}[Payer]))",
        explanation="UNIQUE returns each payer once, in the order it first appears in the table. Only B6 holds the "
                    "formula, and the other names spill into B7, B8, and so on. Click one of them and the formula bar shows "
                    "the formula grayed out, and a blue border outlines the whole spill range.",
    )
    lines_answer = f"{len(service_lines)} service lines · first: {service_lines[0]} · last: {service_lines[-1]}"
    t2 = ws_task(
        [("D6", f_lines)],
        prompt="In Workspace!D6, enter one formula that spills the distinct ServiceLine values from tblEncounters sorted "
               "A to Z. It shows #SPILL! at first. Find out why, fix the problem without moving your formula, and the gray "
               "cell summarizes your list.",
        answer=lines_answer, title="Sorted service lines and a #SPILL! fix (Workspace!D6)",
        solution=f_lines, hint="SORT(UNIQUE(…)). Then click the warning icon next to the error",
        # IFERROR: while D6 shows #SPILL!, say where to look instead of echoing the error into the Practice sheet.
        summary=(f'=IF(ISBLANK({WS}!D6),"",IFERROR(COUNTA({ws_rng("D6")})&" service lines · first: "&{WS}!D6'
                 f'&" · last: "&INDEX({ws_rng("D6")},COUNTA({ws_rng("D6")})),'
                 f'"Workspace!D6 shows an error. Fix it there."))'),
        live=(f'=ROWS(UNIQUE({TE}[ServiceLine]))&" service lines · first: "&INDEX(SORT(UNIQUE({TE}[ServiceLine])),1)'
              f'&" · last: "&INDEX(SORT(UNIQUE({TE}[ServiceLine]),,-1),1)'),
        explanation=f"A leftover note ('{OBSTRUCTION[1]}') sits in {OBSTRUCTION[0]}, inside the range the list needs. Excel "
                    "never overwrites data, so it shows #SPILL! in the anchor instead. Click the warning icon, choose "
                    "**Select Obstructing Cells**, press Delete, and the list spills at once. SORT puts the UNIQUE "
                    "results in A to Z order. UNIQUE(SORT(…)) gives the same list.",
    )
    t3 = Task(
        "Department names repeat across hospitals (each hospital has its own 'Emergency Department'). How many distinct "
        "Facility + Department combinations (units) appear in tblEncounters?",
        answer=len(units), title="Distinct units (Facility + Department)",
        solution=f"=ROWS(UNIQUE({TE}[[Facility]:[Department]]))",
        hint="UNIQUE over two adjacent columns, then count the rows",
        explanation=f"Given two columns, UNIQUE compares whole rows, so 'Bluestone Memorial Hospital | Emergency Department' "
                    f"and 'Cedar Ridge Medical Center | Emergency Department' are different units. "
                    f"=ROWS(UNIQUE(tblEncounters[Department])) returns only {len(set(r['Department'] for r in enc))}, "
                    "because it merges the same name across hospitals. ROWS counts the rows of the spilled result and "
                    "returns one number, so it fits in an answer cell. For columns that aren't next to each other, "
                    "use UNIQUE(HSTACK(col1, col2)).",
    )
    t4 = Task(
        "How many patients (PatientID) have exactly one encounter in this extract?",
        answer=once, title="Patients with exactly one encounter",
        solution=f"=ROWS(UNIQUE({TE}[PatientID],,TRUE))", hint="UNIQUE has an optional third argument",
        explanation=f"The third argument, exactly_once, set to TRUE keeps only values that occur once. The two commas skip "
                    f"the second argument (by_col). Compare =ROWS(UNIQUE(tblEncounters[PatientID])), which counts every "
                    f"distinct patient ({len(patient_counts):,}). The difference, {len(patient_counts) - once} patients, "
                    "came back at least twice.",
    )
    t5 = Task(
        "This one uses tblDepartments (on the Departments sheet), but the formula still goes in the yellow cell here. "
        "Using FILTER, how many departments have UnitType \"Inpatient\" and 20 or more StaffedBeds?",
        answer=len(big_inpatient_units), title="Inpatient units with 20+ staffed beds",
        solution='=ROWS(FILTER(tblDepartments[Department],(tblDepartments[UnitType]="Inpatient")*(tblDepartments[StaffedBeds]>=20)))',
        hint="Multiply the two conditions with *, then count the rows FILTER returns",
        explanation="Each condition returns TRUE or FALSE for all 31 rows. Multiplying them turns TRUE/FALSE into 1/0, so "
                    f"only rows where both are 1 pass the filter (AND logic). {big_icus[0]['Facility']}'s "
                    f"{big_icus[0]['Department']} has {big_icus[0]['StaffedBeds']} beds but doesn't count, because its "
                    "UnitType is 'Critical Care'. COUNTIFS gives the same number here. FILTER is worth learning because the same "
                    "include argument can also return the rows themselves.",
    )
    t6 = Task(
        "How many different attending providers (Attending) treated Emergency encounters?",
        answer=len(ed_attendings), title="Distinct ED attendings",
        solution=f"=ROWS(UNIQUE(FILTER({TE}[Attending],{is_ed})))",
        hint="FILTER first, then UNIQUE, then count",
        explanation="Read it from the inside out. FILTER keeps the Attending names of Emergency encounters (with repeats), "
                    "UNIQUE removes the repeats, and ROWS counts what's left. Nesting functions this way is the core "
                    "dynamic-array habit.",
    )
    t7 = Task(
        "Flu-season review: what were the total charges of Emergency encounters admitted in January or February 2025 "
        "whose DxCategory is Respiratory or Infectious? Enter dollars and cents.",
        answer=round(sum(r["TotalCharges"] for r in flu), 2), fmt="#,##0.00", title="Flu-season ED charges",
        solution=(f'=SUM(FILTER({TE}[TotalCharges],({is_ed})*({TE}[AdmitDate]<DATE(2025,3,1))'
                  f'*(({TE}[DxCategory]="Respiratory")+({TE}[DxCategory]="Infectious"))))'),
        hint="AND with *, OR with +. Wrap the OR part in its own parentheses",
        explanation=f"Three conditions are multiplied (AND), and the last one is itself an OR built by adding two tests. "
                    f"The extra parentheses around the OR matter: without them, multiplication happens first and the "
                    f"formula would add every Infectious encounter of the year regardless of type or date. "
                    f"{len(flu)} encounters match. Because the data covers 2025 only, AdmitDate < 3/1/2025 is enough for "
                    "'January or February'.",
    )
    t8 = Task(
        "What is the combined TotalCharges of the five most expensive Emergency encounters?",
        answer=round(ed_top5, 2), fmt="#,##0.00", title="Top five ED charges (TAKE)",
        solution=f"=SUM(TAKE(SORT(FILTER({TE}[TotalCharges],{is_ed}),,-1),5))",
        hint="SORT the ED charges largest first, TAKE the first 5, then add them",
        explanation="SORT(…,,-1) sorts descending (the empty second argument keeps the default sort column). TAKE(array, 5) "
                    "keeps the first five rows, and SUM collapses them to one number. "
                    "=SUM(LARGE(FILTER(…),{1,2,3,4,5})) works too, but TAKE scales better: change 5 to 50 and nothing "
                    "else changes.",
    )
    t9 = Task(
        "Outliers distort averages. What is the average TotalCharges of Inpatient encounters after dropping the 10 most "
        "expensive stays? Enter dollars and cents.",
        answer=round(trimmed_avg, 2), fmt="#,##0.00", title="Inpatient average without the top 10 (DROP)",
        solution=f"=AVERAGE(DROP(SORT(FILTER({TE}[TotalCharges],{is_ip}),,-1),10))",
        hint="Sort largest first, then DROP the first 10 rows",
        explanation=f"DROP is the mirror image of TAKE: it removes rows from the start (or, with a negative number, from the "
                    f"end). The full inpatient average is {mean(ip_sorted):,.2f}, so ten very expensive stays move the "
                    f"mean by about ${mean(ip_sorted) - trimmed_avg:,.0f}. TRIMMEAN (Lesson 2.4) trims a percentage from both "
                    "ends, while DROP lets you decide exactly what to remove.",
    )
    t10 = Task(
        "Which attending provider's encounters add up to the highest TotalCharges? Enter the name exactly as it appears "
        "(Last, First).",
        answer=ranked_att[0][0], title="Attending with the highest total charges (SORTBY)",
        solution=(f"=TAKE(SORTBY(UNIQUE({TE}[Attending]),SUMIFS({TE}[TotalCharges],{TE}[Attending],"
                  f"UNIQUE({TE}[Attending])),-1),1)"),
        hint="SORTBY the UNIQUE names by a SUMIFS that uses those same names as its criteria",
        explanation=f"SUMIFS normally takes one criterion. Hand it the whole UNIQUE list and it returns one total per "
                    f"name ({len(by_attending)} totals). SORTBY then orders the names by those totals, even though the "
                    f"totals never appear in the result, and TAKE(…,1) keeps the top name. The runner-up, "
                    f"{ranked_att[1][0]}, is only ${ranked_att[0][1] - ranked_att[1][1]:,.0f} behind.",
    )
    t11 = Task(
        "Look up the specialty of the attending on every Inpatient encounter (tblProviders has a Specialty for each "
        "Provider name). How many distinct specialties appear?",
        answer=len(ip_specialties), title="Specialties attending inpatient stays (XLOOKUP with an array)",
        solution=(f"=ROWS(UNIQUE(XLOOKUP(FILTER({TE}[Attending],{is_ip}),tblProviders[Provider],"
                  f"tblProviders[Specialty])))"),
        hint="XLOOKUP accepts a whole array of lookup values",
        live=(f"=ROWS(UNIQUE(INDEX(tblProviders[Specialty],XMATCH(FILTER({TE}[Attending],{is_ip}),"
              f"tblProviders[Provider]))))"),
        self_test=False,  # LibreOffice's XLOOKUP doesn't vectorize; the INDEX/XMATCH live formula verifies the answer
        explanation=f"FILTER returns {len(ip)} attending names, one per inpatient stay. XLOOKUP looks up each of them and "
                    "returns an array of the same size, so you get one specialty per stay. UNIQUE and ROWS then count "
                    "the distinct specialties. In Excel 2019 and earlier you'd need a helper column for this.",
    )
    months_answer = f"12 months · {SAMPLE_SIZE} encounters"
    t12 = ws_task(
        [("F6", f_months), ("G6", f_counts)],
        prompt="On the Workspace sheet, make F6 spill the first day of each month of 2025 (01/01/2025 through 12/01/2025) "
               "using SEQUENCE. Then, in G6, write ONE COUNTIFS formula that refers to your month list with the spill "
               "operator (F6#) and spills the number of encounters admitted in each month. The gray cell checks both "
               "columns.",
        answer=months_answer, title="Month calendar with SEQUENCE and F6# (Workspace!F6:G6)",
        solution=f"1. In **F6**: `{f_months}`\n2. In **G6**: `{f_counts_learner}`",
        hint="DATE accepts an array of months. Count AdmitDate ≥ each month start and < the next month's start",
        summary=f'=IF(ISBLANK({WS}!F6),"",COUNT({ws_rng("F6")})&" months · "&SUM({ws_rng("G6")})&" encounters")',
        live=(f'=LET(starts,DATE(2025,SEQUENCE(12),1),counts,COUNTIFS({TE}[AdmitDate],">="&starts,{TE}[AdmitDate],'
              f'"<"&DATE(YEAR(starts),MONTH(starts)+1,1)),ROWS(starts)&" months · "&SUM(counts)&" encounters")'),
        explanation=f"SEQUENCE(12) spills 1 to 12, and DATE turns each number into that month's first day. In G6, F6# means "
                    f"'the whole spill that starts in F6', so COUNTIFS receives 12 start dates and returns 12 counts. The "
                    f"upper bound DATE(YEAR(F6#),MONTH(F6#)+1,1) is the next month's first day (month 13 rolls into January "
                    f"2026). If you change F6 to 24 months, G6 grows with it automatically. {MONTH_NAMES[busiest_month]} is the "
                    f"busiest month ({busiest_n} encounters).",
    )
    worklist_answer = (f"{len(worklist)} stays · longest: {worklist[0]['EncounterID']} ({worklist[0]['LOSDays']} d) · "
                       f"shortest: {worklist[-1]['EncounterID']} ({worklist[-1]['LOSDays']} d)")
    wl = ws_rng("I6")
    k_col = ws_rng("K6")
    t13 = ws_task(
        [("I6", f_worklist)],
        prompt="Build a long-stay worklist in Workspace!I6 with ONE formula: a header row (EncounterID, Department, LOSDays, "
               "TotalCharges) on top of every Cedar Ridge Medical Center encounter with LOSDays of 7 or more, showing only "
               "those four columns, sorted by LOSDays from longest to shortest.",
        answer=worklist_answer, title="One-formula long-stay worklist (Workspace!I6)",
        solution=f_worklist,
        hint="VSTACK(header, SORT(CHOOSECOLS(FILTER(tblEncounters, …), …), …))",
        summary=(f'=IF(ISBLANK({WS}!I6),"",(COUNTA({wl})-1)&" stays · longest: "&{WS}!I7&" ("&{WS}!K7&" d) · shortest: "'
                 f'&INDEX({wl},COUNTA({wl}))&" ("&INDEX({k_col},COUNTA({wl}))&" d)")'),
        live=(f'=LET(keep,{cedar_long},ids,FILTER({TE}[EncounterID],keep),days,FILTER({TE}[LOSDays],keep),'
              f'srt,SORTBY(HSTACK(ids,days),days,-1),nrows,ROWS(srt),nrows&" stays · longest: "&INDEX(srt,1,1)&" ("'
              f'&INDEX(srt,1,2)&" d) · shortest: "&INDEX(srt,nrows,1)&" ("&INDEX(srt,nrows,2)&" d)")'),
        explanation=f"Build it from the inside out. FILTER(tblEncounters, …) returns all 16 columns of the matching rows. "
                    f"CHOOSECOLS keeps columns {col_no('EncounterID')}, {col_no('Department')}, {col_no('LOSDays')}, and "
                    f"{col_no('TotalCharges')} in the order you list them. SORT(…, 3, -1) sorts by the third of those "
                    "columns (LOSDays), largest first. VSTACK puts the header row, an array constant in braces, on top. "
                    f"Two stays tie at {los_ties[0]} days. To break ties by charges, use SORT(…, {{3,4}}, {{-1,-1}}).",
    )
    L.tasks = [t1, t2, t3, t4, t5, t6, t7, t8, t9, t10, t11, t12, t13]

    # ------------------------------------------------------------------ bonus
    L.bonus_title = "Bonus: A one-formula unit leaderboard"
    L.bonus_scenario = (
        "The CFO wants a 'unit leaderboard' she can refresh every month: every unit (Facility + Department) with at least "
        f"{LEADERBOARD_MIN} encounters in the extract, with four columns (Facility, Department, Encounters, AvgCharge), "
        "where Encounters is the unit's number of encounters and AvgCharge is the average TotalCharges of those "
        "encounters. Sort it by AvgCharge from highest to lowest and put a header row on top. Build it in Workspace!N6 as ONE formula. LET "
        "(previewed in the guide) makes it much easier to read. Then answer the questions below with formulas that refer "
        "to your leaderboard through N6#."
    )
    board_ref = f"{WS}!N6#"
    b1 = Task(
        "How many units qualify for the leaderboard? Count data rows only, not the header.",
        answer=len(board), title="Units on the leaderboard",
        solution=(f"1. In **Workspace!N6**:\n\n```\n{board_pretty}\n```\n\n2. In the answer cell: "
                  f"`=ROWS({board_ref})-1`"),
        hint="ROWS of the spill, minus the header row",
        live=f"=ROWS({board_core})-1",
        explanation="UNIQUE over Facility:Department lists every unit once. CHOOSECOLS splits that two-column list so "
                    "COUNTIFS and AVERAGEIFS can use each column as a criteria array, and each returns one value per unit. "
                    "HSTACK glues the four columns together, FILTER keeps units with n ≥ 50, SORT orders by column 4 "
                    "descending, and VSTACK adds the header. LET names each piece once, so UNIQUE runs once. Without LET, "
                    "the same UNIQUE(…) would have to be repeated in every part of the formula that needs it.",
    )
    b2 = Task(
        "Which Department tops the leaderboard?",
        answer=board[0][1], title="Top unit's department",
        solution=f"=INDEX({board_ref},2,2)", hint="Row 2 of the spill is the first data row",
        live=f"=INDEX({board_core},2,2)",
        explanation=f"INDEX works on a spill like on any range: row 2 (row 1 is the header), column 2. The top unit is "
                    f"{board[0][1]} at {board[0][0]}. Inpatient units dominate the top of the board because a stay "
                    "costs far more than a clinic visit.",
    )
    b3 = Task(
        "What is the top unit's average charge? Enter dollars and cents.",
        answer=round(board[0][3], 2), fmt="#,##0.00", title="Top unit's average charge",
        solution=f"=INDEX({board_ref},2,4)", hint="Same row, fourth column",
        live=f"=INDEX({board_core},2,4)",
        explanation=f"Column 4 of the leaderboard is AvgCharge. A {board[0][1]} stay averages "
                    f"${board[0][3]:,.0f}, about {board[0][3] / mem_ed_avg:.0f} times the average Bluestone Memorial "
                    "emergency visit. If your answer has many decimals, that's fine: the check accepts the unrounded value.",
    )
    b4 = Task(
        "What share of all TotalCharges in tblEncounters comes from the units on the leaderboard? Each unit's total is "
        "Encounters × AvgCharge. Enter it as a percentage.",
        answer=board_share, fmt="0.0%", title="Leaderboard units' share of charges",
        solution=(f"=SUM(CHOOSECOLS(DROP({board_ref},1),3)*CHOOSECOLS(DROP({board_ref},1),4))"
                  f"/SUM({TE}[TotalCharges])"),
        hint="DROP the header row, multiply column 3 by column 4, add up, then divide by all charges",
        live=(f"=SUM(CHOOSECOLS(DROP({board_core},1),3)*CHOOSECOLS(DROP({board_core},1),4))"
              f"/SUM({TE}[TotalCharges])"),
        explanation=f"DROP(N6#,1) removes the header so only numbers remain. CHOOSECOLS pulls out the Encounters and "
                    f"AvgCharge columns. Multiplying them row by row rebuilds each unit's total charges, and SUM adds them. "
                    f"The {len(board)} busiest units bring in {board_share:.1%} of all charges. The other "
                    f"{len(units) - len(board)} units see fewer patients, but most of them are ICUs and inpatient floors "
                    f"with high charges per stay, so together they still account for the remaining {1 - board_share:.1%}.",
    )
    L.bonus = [b1, b2, b3, b4]
    hash_tasks = {id(b1): f"=ROWS({board_ref})-1", id(b2): b2.solution, id(b3): b3.solution, id(b4): b4.solution}

    # ------------------------------------------------------------------ Workspace sheet + self-test simulation
    all_tasks = L.tasks + L.bonus

    @L.customize
    def workspace(wb, lesson, selftest):
        ws = wb.create_sheet(WS)
        ws.sheet_properties.tabColor = "BF9000"
        ws["A1"] = "Workspace: spill your dynamic-array formulas here"
        ws["A1"].font = Font(bold=True, size=16, color=NAVY)
        ws["A2"] = ("Type each formula in its yellow anchor cell and press Enter. Leave the cells below and to the right "
                    "empty so the results can spill. The gray cells on the Practice sheet read these results.")
        ws["A2"].font = Font(italic=True, color="404040")
        for col, w in WIDTHS.items():
            ws.column_dimensions[col].width = w
        label_font = Font(bold=True, color="FFFFFF")
        label_fill = PatternFill("solid", fgColor=NAVY)
        for anchor, (label, note) in ANCHORS.items():
            col = re.match(r"[A-Z]+", anchor).group(0)
            if label:
                c = ws[f"{col}4"]
                c.value = label
                c.font = label_font
            n = ws[f"{col}5"]
            n.value = note
            n.font = Font(italic=True, color="595959", size=9)
            ws[anchor].fill = INPUT_FILL
            ws[anchor].border = INPUT_BORDER
        for rng in ("B4:B4", "D4:D4", "F4:G4", "I4:L4", "N4:Q4"):
            for row in ws[rng]:
                for c in row:
                    c.fill = label_fill
                    c.font = label_font
        for r in range(6, LAST_WS_ROW + 1):
            ws[f"F{r}"].number_format = "mmm yyyy"
            ws[f"L{r}"].number_format = "#,##0.00"
            ws[f"Q{r}"].number_format = "#,##0.00"
            ws[f"K{r}"].number_format = "0"
            ws[f"P{r}"].number_format = "#,##0"
        for rng in ("J6:L6", "O6:Q6"):  # header rows of the two report spills
            for row in ws[rng]:
                for c in row:
                    c.font = Font(bold=True)
        oc = ws[OBSTRUCTION[0]]
        oc.value = OBSTRUCTION[1]
        oc.font = Font(italic=True, color="808080")
        oc.alignment = Alignment(horizontal="left")
        ws.freeze_panes = "A6"
        ws.sheet_view.zoomScale = 90
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        # The text summaries are longer than the answer column is wide, so let them wrap.
        practice = wb[lesson.practice_sheet]
        for t in lesson.tasks:
            if t.summary and t.kind() == "text":
                practice[t.answer_cell].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

        if not selftest:
            return
        # Simulate the learner's work: write each spill formula as a real multi-cell array over exactly the range
        # it spills to, so LibreOffice evaluates the learner's formula (not a pasted copy of the answer).
        ws[OBSTRUCTION[0]].value = None  # the learner deletes the obstruction in task 2
        for t in all_tasks:
            for anchor, formula in spill_fills.get(id(t), []):
                nrows, ncols = sizes[anchor]
                ref = f"{anchor}:{_end_cell(anchor, nrows, ncols)}"
                ws[anchor] = ArrayFormula(ref, to_file_formula(lo_compat(formula, sizes)))
        nrows, ncols = sizes["N6"]
        ws["N6"] = ArrayFormula(f"N6:{_end_cell('N6', nrows, ncols)}", to_file_formula(lo_compat(f_board, sizes)))
        # Answer cells whose solutions use N6# get the explicit-range equivalent.
        for t in all_tasks:
            if id(t) in hash_tasks:
                t.fill = {"range": f"'{t._sheet}'!{t.answer_cell}:{t.answer_cell}",
                          "formula": lo_compat(hash_tasks[id(t)], sizes), "array": True}

    return L
