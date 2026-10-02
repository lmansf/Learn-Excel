"""Lesson 2.6 · Lookups: VLOOKUP, INDEX/MATCH & XLOOKUP.

Data story: Dr. Isabella Nguyen (PRV1137, Internal Medicine) runs a primary care panel at the Bluestone Outpatient
Pavilion. Her practice's care manager reviews every 2025 encounter (clinic, ED, observation, and inpatient) for the
panel's patients. The Encounters sheet holds IDs only, exactly as an EHR extract would; the names, descriptions and
amounts live in lookup tables:

  Encounters   all 2025 encounters for the panel's patients (IDs + dates + charges), oldest first
  Patients     the panel roster (MRN as 8-character text, DOB, age on 12/31/2025, BMI)
  Referrals    referrals received by the practice in December 2025 (some patients are not on the panel yet)
  Providers, Diagnoses, Payers, Departments   system-wide lookup tables
  Claims       one claim per encounter (status as of 12/31/2025)
  Budget       2025 annual expense budget: department (rows) x category (columns), for two-way lookups
  BMITiers, AgeBands   tier tables for approximate-match lookups
  Card         (customize, bonus) a two-column "encounter lookup card" the learner builds

Every answer below is computed in Python from data/*.csv.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal

from openpyxl.styles import Alignment, Font, PatternFill

from xlcourse import Lesson, Task, data
from xlcourse.lesson import BOX, HEADER_FILL, INPUT_BORDER, INPUT_FILL, NAVY

CODE = "2.6"
PCP = "PRV1137"                 # Dr. Isabella Nguyen, Internal Medicine, Primary Care Clinic (D400)
YEAR = 2025
AS_OF = date(2025, 12, 31)

CATS = ["Salaries & Wages", "Employee Benefits", "Medical Supplies", "Pharmaceuticals",
        "Purchased Services", "Equipment & Maintenance", "Other Operating"]

# Tier tables: lower bound of each tier, sorted ascending (CDC adult BMI categories; common age bands).
BMI_TIERS = [(0, "Underweight", "Below 18.5"), (18.5, "Healthy weight", "18.5 to under 25"),
             (25, "Overweight", "25 to under 30"), (30, "Obesity class 1", "30 to under 35"),
             (35, "Obesity class 2", "35 to under 40"), (40, "Obesity class 3", "40 and above")]
AGE_BANDS = [(0, "0-17"), (18, "18-44"), (45, "45-64"), (65, "65-74"), (75, "75-84"), (85, "85+")]

# Bonus: the two IDs on the lookup card.
CARD_A = "ENC117295"            # a real encounter on the sheet
CARD_B_REAL = "ENC119088"       # ...and a real one that was copied with a typo:
CARD_B = "ENC119O88"            # the letter O instead of a zero
NOT_FOUND = "Not found"

# Task targets chosen from the data (validated below)
TWO_WAY_DEPT, TWO_WAY_CAT = "D130", "Pharmaceuticals"      # ICU, Bluestone Memorial
ROW_TOTAL_DEPT = "D400"                                    # Dr. Nguyen's own clinic
WILDCARD = "*knee*"
MRN_ZEROS = "000"                # task 8: the first panel patient whose MRN starts with three zeros


# ---------------------------------------------------------------------------
# Python models of the Excel behaviour the tasks rely on
# ---------------------------------------------------------------------------
def round_half_up(x: float, digits: int) -> float:
    """Excel's ROUND (halves away from zero), so a learner who recomputes BMI gets the stored value."""
    q = Decimal(1).scaleb(-digits)
    return float(Decimal(repr(x)).quantize(q, rounding=ROUND_HALF_UP))


def datedif_y(start: date, end: date) -> int:
    """DATEDIF(start, end, "Y"): completed years."""
    return end.year - start.year - ((end.month, end.day) < (start.month, start.day))


def approx(table, value):
    """VLOOKUP(..., TRUE) / XLOOKUP(..., -1): the last tier whose lower bound is <= value."""
    hit = None
    for row in table:
        if row[0] <= value:
            hit = row
    return hit


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="02-formulas-functions", slug="06-lookup-functions",
        title="Lookups: VLOOKUP, INDEX/MATCH & XLOOKUP", level="Beginner → Intermediate", minutes=60,
        objectives=[
            "Look up exact matches with VLOOKUP and XLOOKUP",
            "Use approximate matches for tiers like age bands and BMI categories",
            "Combine INDEX and MATCH for flexible and two-way lookups",
            "Handle missing values with IFNA and XLOOKUP's if_not_found",
        ],
        data_note="Every 2025 encounter (529) for the 185 patients on Dr. Isabella Nguyen's internal medicine panel, stored as "
                  "IDs only, plus the lookup tables that turn those IDs into names: the panel roster, providers, ICD-10 "
                  "diagnoses, payers, departments, and claims. Also a 2025 department budget grid, BMI and age-band tier "
                  "tables, and the practice's December 2025 referral list.",
    )

    # ------------------------------------------------------------------ source data
    all_pts = data.load("patients")
    pts_by_id = {p["PatientID"]: p for p in all_pts}
    prov = {p["ProviderID"]: p for p in data.load("providers")}
    dx = {d["DxCode"]: d for d in data.load("diagnoses")}
    payers = {p["PayerID"]: p for p in data.load("payers")}
    depts = {d["DeptID"]: d for d in data.load("departments")}
    facilities = {f["FacilityID"]: f["FacilityName"] for f in data.load("facilities")}
    claims_by_enc = {c["EncounterID"]: c for c in data.load("claims")}

    # Panel roster (sorted by PatientID) with age on 12/31/2025 and BMI (rounded to 1 decimal, as an EHR shows it).
    panel = sorted((dict(p) for p in all_pts if p["PCPProviderID"] == PCP), key=lambda p: p["PatientID"])
    for p in panel:
        p["Age"] = datedif_y(p["DOB"], AS_OF)
        p["BMI"] = round_half_up(703 * p["WeightLb"] / p["HeightIn"] ** 2, 1)
    panel_ids = {p["PatientID"] for p in panel}
    pidx = {p["PatientID"]: i for i, p in enumerate(panel)}

    # Encounters: every 2025 encounter for the panel, oldest first (so the LAST match is the most recent).
    enc = [dict(e) for e in data.load("encounters") if e["PatientID"] in panel_ids and e["AdmitDateTime"].year == YEAR]
    enc.sort(key=lambda e: (e["AdmitDateTime"], e["EncounterID"]))
    for e in enc:
        e["AdmitDate"] = e["AdmitDateTime"].date()
        e["DischargeDate"] = e["DischargeDateTime"].date()
    eidx = {e["EncounterID"]: i for i, e in enumerate(enc)}

    claims = sorted((dict(claims_by_enc[e["EncounterID"]]) for e in enc), key=lambda c: c["ClaimID"])

    providers = sorted(prov.values(), key=lambda p: p["ProviderID"])
    diagnoses = sorted(dx.values(), key=lambda d: d["DxCode"])
    payer_rows = sorted(payers.values(), key=lambda p: p["PayerID"])
    dept_rows = sorted(depts.values(), key=lambda d: d["DeptID"])

    # Budget grid: 2025 annual expense budget by department x category.
    bsum = defaultdict(int)
    for b in data.load("budget"):
        if b["FiscalYear"] == YEAR and b["LineType"] == "Expense":
            bsum[(b["DeptID"], b["Category"])] += b["BudgetAmount"]
    budget_rows = []
    for d in dept_rows:
        row = {"DeptID": d["DeptID"], "DeptName": d["DeptName"], "FacilityID": d["FacilityID"]}
        for c in CATS:
            row[c] = bsum[(d["DeptID"], c)]
        budget_rows.append(row)

    # Referrals received in December 2025: 17 panel patients + 7 patients with no PCP on file (not on the panel).
    no_pcp = sorted((p for p in all_pts if p["PCPProviderID"] is None), key=lambda p: p["PatientID"])
    on_panel = [panel[3 + 11 * k] for k in range(17)]
    new_pts = [no_pcp[5 + 41 * k] for k in range(7)]
    new_slots = {2, 5, 9, 12, 16, 19, 22}
    sources = ["Bluestone Memorial ED", "Bluestone Memorial Hospital Medicine", "Cedar Ridge Medical Center",
               "Ashby Falls Community Hospital", "Cardiology Clinic", "Endocrinology & Diabetes Clinic"]
    reasons = ["ED follow-up", "Post-discharge follow-up", "Medication review", "Abnormal lab follow-up",
               "Blood pressure check", "Diabetes management", "Establish primary care"]
    biz_days = [date(2025, 12, 1) + timedelta(days=i) for i in range(31)]
    biz_days = [d for d in biz_days if d.weekday() < 5 and d != date(2025, 12, 25)]
    referrals, a, b = [], iter(on_panel), iter(new_pts)
    for i in range(24):
        p = next(b) if i in new_slots else next(a)
        referrals.append({
            "ReferralID": f"RF-2512-{i + 1:02d}", "ReceivedDate": biz_days[i * len(biz_days) // 24], "MRN": p["MRN"],
            "PatientName": f"{p['LastName']}, {p['FirstName']}", "ReferredBy": sources[(i * 5) % len(sources)],
            "Reason": reasons[(i * 3) % len(reasons)], "_new": p["PatientID"] not in panel_ids, "_pid": p["PatientID"],
        })

    bmi_rows = [{"MinBMI": lo, "Category": c, "BMIRange": r} for lo, c, r in BMI_TIERS]
    age_rows = [{"MinAge": lo, "AgeBand": band} for lo, band in AGE_BANDS]

    # ------------------------------------------------------------------ sheets
    es = L.add_table_sheet(
        "Encounters", enc, table="tblEncounters",
        columns=["EncounterID", "PatientID", "EncounterType", "AdmitDate", "DischargeDate", "DeptID",
                 "AttendingProviderID", "PrimaryDxCode", "PayerID", "TotalCharges"],
        extra_cols=["PayerName"], widths={"PayerName": 30, "AttendingProviderID": 20},
    )
    ps = L.add_table_sheet(
        "Patients", panel, table="tblPatients",
        columns=["PatientID", "MRN", "FirstName", "LastName", "Sex", "DOB", "Age", "HeightIn", "WeightLb", "BMI",
                 "PrimaryPayerID"],
        extra_cols=["BMICategory"], formats={"HeightIn": "0.0", "WeightLb": "0.0", "BMI": "0.0"},
        widths={"BMICategory": 18, "MRN": 11},
    )
    rs = L.add_table_sheet(
        "Referrals", referrals, table="tblReferrals",
        columns=["ReferralID", "ReceivedDate", "MRN", "PatientName", "ReferredBy", "Reason"],
        extra_cols=["PatientID"], widths={"PatientID": 14, "MRN": 11},
    )
    prs = L.add_table_sheet(
        "Providers", providers, table="tblProviders",
        columns=["ProviderID", "FirstName", "LastName", "Credential", "Specialty", "PrimaryDeptID", "FacilityID"],
    )
    dxs = L.add_table_sheet(
        "Diagnoses", diagnoses, table="tblDiagnoses",
        columns=["DxCode", "DxDescription", "DxCategory", "ChronicCondition", "ExpectedLOS"],
        formats={"ExpectedLOS": "0.0"}, widths={"DxDescription": 62},
    )
    pys = L.add_table_sheet(
        "Payers", payer_rows, table="tblPayers",
        columns=["PayerID", "PayerName", "PayerType", "AvgAllowedPctOfCharges", "AvgDaysToPay", "TimelyFilingDays"],
        formats={"AvgAllowedPctOfCharges": "0%"},
    )
    ds = L.add_table_sheet(
        "Departments", dept_rows, table="tblDepartments",
        columns=["DeptID", "DeptName", "FacilityID", "ServiceLine", "UnitType", "StaffedBeds", "CostCenter"],
    )
    cls = L.add_table_sheet(
        "Claims", claims, table="tblClaims",
        columns=["ClaimID", "EncounterID", "PayerID", "SubmitDate", "BilledAmount", "PaidAmount", "ClaimStatus",
                 "DenialReason"],
    )
    bs = L.add_table_sheet(
        "Budget", budget_rows, table="tblBudget",
        columns=["DeptID", "DeptName", "FacilityID"] + CATS,
        formats={c: "#,##0" for c in CATS}, widths={c: 15 for c in CATS},
    )
    bts = L.add_table_sheet("BMITiers", bmi_rows, table="tblBMITiers", columns=["MinBMI", "Category", "BMIRange"],
                            formats={"MinBMI": "0.0"}, widths={"Category": 18, "BMIRange": 20})
    ags = L.add_table_sheet("AgeBands", age_rows, table="tblAgeBands", columns=["MinAge", "AgeBand"])

    # ------------------------------------------------------------------ range helpers
    def R(sd, col):
        """Absolute data range, e.g. Patients!$A$2:$A$186 (written the way a learner types it: no quotes, because
        none of this lesson's sheet names contain spaces)."""
        c = sd.col(col)
        return f"{sd.name}!${c}${sd.first_row}:${c}${sd.last_row}"

    def tbl(sd, first_col, last_col):
        """Absolute block from first_col to last_col, e.g. Patients!$A$2:$K$186."""
        return f"{sd.name}!${sd.col(first_col)}${sd.first_row}:${sd.col(last_col)}${sd.last_row}"

    def erow(enc_id):
        return es.first_row + eidx[enc_id]

    def ecell(col, enc_id):
        return f"Encounters!{es.col(col)}{erow(enc_id)}"

    pat_tbl = tbl(ps, "PatientID", "PrimaryPayerID")
    pay_tbl = tbl(pys, "PayerID", "TimelyFilingDays")
    bmi_tbl = tbl(bts, "MinBMI", "Category")
    age_tbl = tbl(ags, "MinAge", "AgeBand")
    grid = tbl(bs, CATS[0], CATS[-1])
    grid_hdr = f"Budget!${bs.col(CATS[0])}$1:${bs.col(CATS[-1])}$1"

    # ------------------------------------------------------------------ choose task targets (deterministic)
    def first(pred, start=0):
        return next(e for e in enc[start:] if pred(e))

    t1 = first(lambda e: e["EncounterType"] == "Inpatient", 6)
    t2 = first(lambda e: e["PayerID"] == "PY02", 16)
    t3 = first(lambda e: e["PrimaryDxCode"] == "E11.65", 30)
    t4 = first(lambda e: e["EncounterType"] == "Inpatient"
               and prov[e["AttendingProviderID"]]["Specialty"] == "Pulmonary & Critical Care", 40)
    t6p = next(p for p in panel if p["Age"] == AGE_BANDS[4][0])                   # exactly on a band boundary (75)
    t8p = next(p for p in panel if p["MRN"].startswith(MRN_ZEROS))
    mrn_number = int(t8p["MRN"])
    enc_count = Counter(e["PatientID"] for e in enc)
    t12_pid = min(enc_count, key=lambda k: (-enc_count[k], k))                      # the panel's most frequent visitor
    t12_last = [e for e in enc if e["PatientID"] == t12_pid][-1]
    t12_first = [e for e in enc if e["PatientID"] == t12_pid][0]
    wild_hits = [d for d in diagnoses if "knee" in d["DxDescription"].lower()]
    assert len(wild_hits) == 1
    assert CARD_A in eidx and CARD_B_REAL in eidx and CARD_B not in eidx

    # ------------------------------------------------------------------ answers
    a1 = pts_by_id[t1["PatientID"]]["LastName"]
    a2 = payers[t2["PayerID"]]["PayerType"]
    a3 = dx[t3["PrimaryDxCode"]]["DxDescription"]
    a4 = prov[t4["AttendingProviderID"]]["Specialty"]
    a5 = sum(1 for e in enc if payers[e["PayerID"]]["PayerName"] == "Medicare")
    a6 = approx(AGE_BANDS, t6p["Age"])[1]
    a7 = sum(1 for p in panel if approx(BMI_TIERS, p["BMI"])[1].startswith("Obesity"))
    a8 = t8p["PatientID"]
    a9 = wild_hits[0]["DxCode"]
    a10 = bsum[(TWO_WAY_DEPT, TWO_WAY_CAT)]
    a11 = sum(bsum[(ROW_TOTAL_DEPT, c)] for c in CATS)
    a12 = t12_last["AdmitDate"]
    a13 = sum(1 for r in referrals if r["_pid"] not in panel_ids)
    assert a13 == 7 and len({r["MRN"] for r in referrals}) == 24
    assert sum(1 for p in panel if p["MRN"] == t8p["MRN"]) == 1

    # helper values for prompts/explanations
    bmi_boundary = sum(1 for p in panel if p["BMI"] in (25.0, 30.0))
    n_bmi = Counter(approx(BMI_TIERS, p["BMI"])[1] for p in panel)
    pay_col, bmi_col, pid_col = es.col("PayerName"), ps.col("BMICategory"), rs.col("PatientID")
    pay_rng = f"Encounters!{pay_col}{es.first_row}:{pay_col}{es.last_row}"
    bmi_rng = f"Patients!{bmi_col}{ps.first_row}:{bmi_col}{ps.last_row}"
    pid_rng = f"Referrals!{pid_col}{rs.first_row}:{pid_col}{rs.last_row}"
    r6 = ps.first_row + pidx[t6p["PatientID"]]

    pay_fill = f"=XLOOKUP({es.col('PayerID')}{es.first_row},{R(pys, 'PayerID')},{R(pys, 'PayerName')})"
    bmi_fill = f"=VLOOKUP({ps.col('BMI')}{ps.first_row},{bmi_tbl},2,TRUE)"
    ref_fill = (f'=XLOOKUP({rs.col("MRN")}{rs.first_row},{R(ps, "MRN")},{R(ps, "PatientID")},"{NOT_FOUND}")')

    L.practice_intro = (
        "The Encounters sheet stores IDs only. Every answer comes from looking an ID up in another sheet. "
        "Write each answer as a formula. Lock lookup ranges with $ (F4, Mac: ⌘ + T) so they don't slide when you copy.")

    L.tasks = [
        Task(f"Encounter {t1['EncounterID']} is on row {erow(t1['EncounterID'])} of the Encounters sheet. Use VLOOKUP with its "
             f"PatientID (cell B{erow(t1['EncounterID'])}) to return the patient's last name from the Patients sheet.",
             answer=a1, hint="VLOOKUP(lookup_value, table_array, col_index_num, FALSE). Count columns from PatientID (column 1) to LastName",
             solution=f"=VLOOKUP({ecell('PatientID', t1['EncounterID'])},{pat_tbl},4,FALSE)",
             explanation="VLOOKUP searches the first column of the table (PatientID) for the ID, then returns the value from the 4th "
                         "column of the same row: PatientID, MRN, FirstName, **LastName**. FALSE asks for an exact match, which is "
                         "what you want for any ID."),
        Task(f"Encounter {t2['EncounterID']} (Encounters row {erow(t2['EncounterID'])}) was billed to payer "
             f"{t2['PayerID']}. Use VLOOKUP on the Payers sheet to return the payer's PayerType.",
             answer=a2, hint="Count the columns of the Payers table to find PayerType's column number",
             solution=f"=VLOOKUP({ecell('PayerID', t2['EncounterID'])},{pay_tbl},3,FALSE)",
             explanation=f"PayerType is the 3rd column of Payers (PayerID, PayerName, **PayerType**). The payer's name, "
                         f"{payers[t2['PayerID']]['PayerName']}, contains the word Medicare, but its type is "
                         f"{a2}, a private plan that contracts with Medicare. That difference matters for billing rules, "
                         "which is why reports look up the type instead of guessing it from the name."),
        Task(f"Encounter {t3['EncounterID']} (Encounters row {erow(t3['EncounterID'])}) has primary diagnosis code "
             f"{t3['PrimaryDxCode']}. Use XLOOKUP to return its description from the Diagnoses sheet.",
             answer=a3, hint="XLOOKUP(lookup_value, lookup_array, return_array)",
             solution=f"=XLOOKUP({ecell('PrimaryDxCode', t3['EncounterID'])},{R(dxs, 'DxCode')},{R(dxs, 'DxDescription')})",
             explanation="XLOOKUP takes two separate columns, one to search and one to return, so there's no column number to count. "
                         "It also defaults to an exact match, so you can't forget the FALSE that VLOOKUP needs."),
        Task(f"Who attended encounter {t4['EncounterID']} (Encounters row {erow(t4['EncounterID'])})? Use INDEX and MATCH "
             f"with its AttendingProviderID to return the provider's Specialty from the Providers sheet.",
             answer=a4, hint="INDEX(Specialty column, MATCH(id, ProviderID column, 0))",
             solution=f"=INDEX({R(prs, 'Specialty')},MATCH({ecell('AttendingProviderID', t4['EncounterID'])},"
                      f"{R(prs, 'ProviderID')},0))",
             explanation="MATCH finds the position of the ProviderID in the ProviderID column (the 0 means exact match). INDEX then "
                         "returns the item at that same position in the Specialty column. Because the two columns are separate "
                         "arguments, INDEX/MATCH works in every Excel version and doesn't care where the columns sit."),
        Task(f"Fill the yellow PayerName column on the Encounters sheet with a lookup that returns each encounter's payer "
             f"name (start in {pay_col}{es.first_row} and copy down to row {es.last_row}). The gray cell counts how many "
             f"encounters your column shows as exactly \"Medicare\".",
             answer=a5, title="PayerName column (count of Medicare encounters)",
             hint="Lock the Payers ranges with $ before you copy the formula down",
             solution=pay_fill,
             summary=f'=IF(COUNTA({pay_rng})=0,"",COUNTIF({pay_rng},"Medicare"))',
             fill={"range": pay_rng, "formula": pay_fill},
             live=f'=COUNTIF({R(es, "PayerID")},"PY01")',
             explanation=f"This is the everyday use of a lookup: adding a readable column to an ID-only extract. The $ signs keep "
                         f"the Payers ranges fixed while the lookup value ({es.col('PayerID')}{es.first_row}, then "
                         f"{es.col('PayerID')}{es.first_row + 1}, …) moves down a row at a time. "
                         f"`=VLOOKUP({es.col('PayerID')}{es.first_row},{pay_tbl},2,FALSE)` works just as well. COUNTIF with "
                         "\"Medicare\" counts exact matches only, so Silverline Medicare Advantage isn't included."),
        Task(f"Patient {t6p['PatientID']} is on Patients row {r6}, and the Age column shows {t6p['Age']}. Use VLOOKUP with an "
             f"approximate match (TRUE) on the AgeBands sheet to return the patient's age band.",
             answer=a6, hint="An approximate match returns the band whose MinAge is the largest one that is ≤ the age",
             solution=f"=VLOOKUP(Patients!{ps.col('Age')}{r6},{age_tbl},2,TRUE)",
             explanation=f"With TRUE, VLOOKUP looks for the largest MinAge that is less than or equal to {t6p['Age']}. MinAge "
                         f"{AGE_BANDS[4][0]} qualifies, and the next one ({AGE_BANDS[5][0]}) is too big, so the answer is the "
                         f"{a6} band. A value exactly on a boundary belongs to the band that **starts** there. That's why a tier "
                         "table lists each band's lower bound, sorted smallest to largest."),
        Task(f"Fill the yellow BMICategory column on the Patients sheet: look up each patient's BMI in the BMITiers table "
             f"with an approximate match (start in {bmi_col}{ps.first_row} and copy down to row {ps.last_row}). The gray cell "
             "counts patients in any obesity class (BMI 30 or higher).",
             answer=a7, title="BMICategory column (patients in an obesity class)",
             hint="VLOOKUP(…, TRUE) or XLOOKUP with match_mode -1 (exact match or next smaller item)",
             solution=bmi_fill,
             summary=f'=IF(COUNTA({bmi_rng})=0,"",COUNTIF({bmi_rng},"Obesity*"))',
             fill={"range": bmi_rng, "formula": bmi_fill},
             live=f'=COUNTIF({R(ps, "BMI")},">=30")',
             explanation=f"Each BMI falls between two lower bounds in BMITiers, and the approximate match returns the tier whose "
                         f"lower bound is the largest one at or below the BMI. {bmi_boundary} patients sit exactly on a boundary "
                         f"(25.0 or 30.0) and correctly land in the higher tier. The XLOOKUP version is "
                         f"`=XLOOKUP({ps.col('BMI')}{ps.first_row},{R(bts, 'MinBMI')},{R(bts, 'Category')},,-1)`. "
                         f"The panel has {n_bmi['Obesity class 1']} patients in class 1, {n_bmi['Obesity class 2']} in class 2, "
                         f"and {n_bmi['Obesity class 3']} in class 3."),
        Task(f"The lab's interface file sent a result for MRN {mrn_number}. It arrived as a number, so its leading zeros were "
             f"dropped, but the Patients sheet stores MRNs as 8-character text. Return the PatientID for this MRN. "
             f"(PatientID is to the left of MRN, so VLOOKUP can't do it.)",
             answer=a8, hint='Rebuild the text MRN with TEXT(number,"00000000"), then use INDEX/MATCH or XLOOKUP',
             solution=f'=XLOOKUP(TEXT({mrn_number},"00000000"),{R(ps, "MRN")},{R(ps, "PatientID")})',
             explanation=f"Looking up the number {mrn_number} returns #N/A, because a number never equals the text "
                         f"\"{t8p['MRN']}\". TEXT(…,\"00000000\") pads it back to 8 characters as text, and then the match works. "
                         f"The INDEX/MATCH version is `=INDEX({R(ps, 'PatientID')},MATCH(TEXT({mrn_number},\"00000000\"),"
                         f"{R(ps, 'MRN')},0))`. Both can return a column to the left of the one they search."),
        Task("A clinic note mentions \"osteoarthritis of the right knee.\" Use a wildcard lookup to return the DxCode whose "
             "DxDescription contains the word knee.",
             answer=a9, hint="An asterisk wildcard stands for any characters, so put one on each side of the word. XLOOKUP needs match_mode 2",
             solution=f'=XLOOKUP("{WILDCARD}",{R(dxs, "DxDescription")},{R(dxs, "DxCode")},,2)',
             explanation=f"The asterisk stands for any number of characters, so `\"{WILDCARD}\"` matches "
                         f"\"{wild_hits[0]['DxDescription']}\". XLOOKUP only uses wildcards when match_mode is 2. MATCH with "
                         f"match_type 0 uses them automatically: `=INDEX({R(dxs, 'DxCode')},MATCH(\"{WILDCARD}\","
                         f"{R(dxs, 'DxDescription')},0))`. If several descriptions matched, you'd get the first one."),
        Task(f"Two-way lookup: on the Budget sheet, what is the 2025 {TWO_WAY_CAT} budget for department {TWO_WAY_DEPT} "
             f"({depts[TWO_WAY_DEPT]['DeptName']} at {facilities[depts[TWO_WAY_DEPT]['FacilityID']]})? Find the row by "
             f"DeptID and the column by category name.",
             answer=a10, fmt="#,##0", hint="INDEX(grid, MATCH(row label…), MATCH(column header…))",
             solution=f'=INDEX({grid},MATCH("{TWO_WAY_DEPT}",{R(bs, "DeptID")},0),MATCH("{TWO_WAY_CAT}",{grid_hdr},0))',
             explanation=f"The first MATCH finds the department's row inside the grid and the second finds the category's column "
                         f"in the header row. INDEX returns the cell where they cross. The nested XLOOKUP version is "
                         f"`=XLOOKUP(\"{TWO_WAY_DEPT}\",{R(bs, 'DeptID')},XLOOKUP(\"{TWO_WAY_CAT}\",{grid_hdr},{grid}))`: the "
                         f"inner XLOOKUP returns the whole {TWO_WAY_CAT} column, and the outer one picks the department's row from it. "
                         "Look up by DeptID, not DeptName. Three facilities each have an \"Intensive Care Unit\", and a lookup "
                         "on that name would return only the first one."),
        Task(f"What is the TOTAL 2025 expense budget (all seven categories) for department {ROW_TOTAL_DEPT}, "
             f"{depts[ROW_TOTAL_DEPT]['DeptName']}, where Dr. Nguyen practices? Use one lookup that returns the department's "
             f"whole row of the grid, wrapped in SUM.",
             answer=a11, fmt="#,##0", hint="XLOOKUP's return_array can be several columns wide",
             solution=f'=SUM(XLOOKUP("{ROW_TOTAL_DEPT}",{R(bs, "DeptID")},{grid}))',
             explanation=f"When return_array is seven columns wide, XLOOKUP returns all seven values from the matching row, and SUM "
                         f"adds them. On its own in an empty cell the same XLOOKUP would spill across seven cells. INDEX can do this "
                         f"too, because a column number of 0 means \"the whole row\": `=SUM(INDEX({grid},MATCH(\"{ROW_TOTAL_DEPT}\","
                         f"{R(bs, 'DeptID')},0),0))`."),
        Task(f"Patient {t12_pid} visited more often than anyone else on the panel ({enc_count[t12_pid]} encounters in 2025). "
             f"The Encounters sheet is sorted oldest to newest. Use XLOOKUP searching from the bottom up to return the "
             f"AdmitDate of this patient's most recent encounter.",
             answer=a12, fmt="mm/dd/yyyy", hint="search_mode is XLOOKUP's 6th argument, and -1 searches last-to-first",
             solution=f'=XLOOKUP("{t12_pid}",{R(es, "PatientID")},{R(es, "AdmitDate")},,0,-1)',
             explanation=f"A normal lookup stops at the first match, which here is the oldest encounter "
                         f"({t12_first['AdmitDate']:%m/%d/%Y}). search_mode -1 starts at the bottom, so the first match it meets is "
                         f"the most recent row. If the result shows a number like {int((a12 - date(1899, 12, 30)).days)}, "
                         "format the cell as a date. `=MAXIFS(…)` (Lesson 2.5) gives the same date, but only the lookup can return "
                         "other columns from that row, like its EncounterID or diagnosis."),
        Task(f"The Referrals sheet lists {len(referrals)} referrals received in December. Fill its yellow PatientID column by "
             f"looking up each MRN in the Patients sheet (start in {pid_col}{rs.first_row}). Patients who aren't on the panel "
             f"yet must show the text {NOT_FOUND} instead of #N/A. The gray cell counts the {NOT_FOUND} rows.",
             answer=a13, title="Referral PatientID column (count of Not found)",
             hint="XLOOKUP's 4th argument (if_not_found), or wrap the lookup in IFNA(…, \"Not found\")",
             solution=ref_fill,
             summary=f'=IF(COUNTA({pid_rng})=0,"",COUNTIF({pid_rng},"{NOT_FOUND}"))',
             fill={"range": pid_rng, "formula": ref_fill},
             live=f'=SUMPRODUCT(--ISNA(MATCH({R(rs, "MRN")},{R(ps, "MRN")},0)))',
             explanation=f"XLOOKUP's if_not_found argument replaces #N/A with your own text, but only when the lookup really finds nothing. The "
                         f"older equivalent is `=IFNA(INDEX({R(ps, 'PatientID')},MATCH({rs.col('MRN')}{rs.first_row},"
                         f"{R(ps, 'MRN')},0)),\"{NOT_FOUND}\")`. Prefer IFNA to IFERROR here. IFERROR would also hide a #REF! "
                         f"or #NAME? caused by a broken formula, and you'd never know. The {a13} {NOT_FOUND} rows are new patients "
                         "for the practice to register."),
    ]

    # ------------------------------------------------------------------ bonus: the encounter lookup card
    ea = enc[eidx[CARD_A]]
    pa = pts_by_id[ea["PatientID"]]
    pr = prov[ea["AttendingProviderID"]]
    ca = claims_by_enc[CARD_A]
    card_rows = [  # (label, formula for column B (Card A); column C is the same formula copied right)
        ("PatientID", f'=XLOOKUP(B$5,{R(es, "EncounterID")},{R(es, "PatientID")},"{NOT_FOUND}")'),
        ("Attending ProviderID", f'=XLOOKUP(B$5,{R(es, "EncounterID")},{R(es, "AttendingProviderID")},"{NOT_FOUND}")'),
        ("Patient name (First Last)",
         f'=IF(B6="{NOT_FOUND}","{NOT_FOUND}",XLOOKUP(B6,{R(ps, "PatientID")},{R(ps, "FirstName")})&" "&'
         f'XLOOKUP(B6,{R(ps, "PatientID")},{R(ps, "LastName")}))'),
        ("Age at admission (years)",
         f'=IF(B6="{NOT_FOUND}","{NOT_FOUND}",DATEDIF(XLOOKUP(B6,{R(ps, "PatientID")},{R(ps, "DOB")}),'
         f'XLOOKUP(B$5,{R(es, "EncounterID")},{R(es, "AdmitDate")}),"Y"))'),
        ("Primary diagnosis",
         f'=XLOOKUP(XLOOKUP(B$5,{R(es, "EncounterID")},{R(es, "PrimaryDxCode")},"{NOT_FOUND}"),{R(dxs, "DxCode")},'
         f'{R(dxs, "DxDescription")},"{NOT_FOUND}")'),
        ("Unit (department)",
         f'=XLOOKUP(XLOOKUP(B$5,{R(es, "EncounterID")},{R(es, "DeptID")},"{NOT_FOUND}"),{R(ds, "DeptID")},'
         f'{R(ds, "DeptName")},"{NOT_FOUND}")'),
        ("Attending (First Last, Credential)",
         f'=IF(B7="{NOT_FOUND}","{NOT_FOUND}",XLOOKUP(B7,{R(prs, "ProviderID")},{R(prs, "FirstName")})&" "&'
         f'XLOOKUP(B7,{R(prs, "ProviderID")},{R(prs, "LastName")})&", "&XLOOKUP(B7,{R(prs, "ProviderID")},{R(prs, "Credential")}))'),
        ("Attending specialty", f'=XLOOKUP(B7,{R(prs, "ProviderID")},{R(prs, "Specialty")},"{NOT_FOUND}")'),
        ("Payer", f'=XLOOKUP(XLOOKUP(B$5,{R(es, "EncounterID")},{R(es, "PayerID")},"{NOT_FOUND}"),{R(pys, "PayerID")},'
                  f'{R(pys, "PayerName")},"{NOT_FOUND}")'),
        ("Claim status", f'=XLOOKUP(B$5,{R(cls, "EncounterID")},{R(cls, "ClaimStatus")},"{NOT_FOUND}")'),
    ]
    CARD_FIRST = 6
    card_row = {label: CARD_FIRST + i for i, (label, _) in enumerate(card_rows)}
    card_formula = dict(card_rows)
    card_a = {
        "PatientID": pa["PatientID"], "Attending ProviderID": pr["ProviderID"],
        "Patient name (First Last)": f"{pa['FirstName']} {pa['LastName']}",
        "Age at admission (years)": datedif_y(pa["DOB"], ea["AdmitDate"]),
        "Primary diagnosis": dx[ea["PrimaryDxCode"]]["DxDescription"],
        "Unit (department)": depts[ea["DeptID"]]["DeptName"],
        "Attending (First Last, Credential)": f"{pr['FirstName']} {pr['LastName']}, {pr['Credential']}",
        "Attending specialty": pr["Specialty"], "Payer": payers[ea["PayerID"]]["PayerName"],
        "Claim status": ca["ClaimStatus"],
    }
    naive_age = ea["AdmitDate"].year - pa["DOB"].year
    days_to_bday = (date(ea["AdmitDate"].year, pa["DOB"].month, pa["DOB"].day) - ea["AdmitDate"]).days
    assert naive_age != card_a["Age at admission (years)"]      # the YEAR-minus-YEAR trap is real for this patient
    card_b_count = len(card_rows)                                 # every output row must say Not found
    card_table = "\n".join(
        ["| Row | Field | Card A (" + CARD_A + ") | Card B (" + CARD_B + ") |", "|:-:|---|---|---|"]
        + [f"| {card_row[k]} | {k} | {v} | {NOT_FOUND} |" for k, v in card_a.items()])

    def card_task(label, prompt, answer, hint, explanation, live, fmt=None):
        r = card_row[label]
        return Task(prompt, answer=answer, hint=hint, fmt=fmt, explanation=explanation,
                    solution=card_formula[label],
                    summary=f'=IF(Card!B{r}="","",Card!B{r})',
                    fill={"range": f"Card!B{r}:C{r}", "formula": card_formula[label]},
                    live=live, title=f"Card A · {label}")

    pid_of_a = f'XLOOKUP("{CARD_A}",{R(es, "EncounterID")},{R(es, "PatientID")})'
    prv_of_a = f'XLOOKUP("{CARD_A}",{R(es, "EncounterID")},{R(es, "AttendingProviderID")})'
    L.bonus_title = "Bonus: An encounter lookup card"
    L.bonus_scenario = (
        "Dr. Nguyen's care manager answers the same questions all day: who is this encounter's patient, how old were they, "
        "what were they treated for, who was the attending, who pays, and where does the claim stand? Build a reusable "
        "lookup card on the Card sheet. Each column takes one EncounterID in row 5 and returns ten facts below it. "
        "Rows 6 and 7 are helper cells (PatientID and Attending ProviderID) that the other rows can reuse. Write every "
        "formula in column B, then copy B6:B15 to C6:C15. "
        f"Card A holds {CARD_A}. Card B holds {CARD_B}, an ID copied from a handwritten note, and every row of Card B must "
        f"show {NOT_FOUND} instead of an error. Keep both IDs in place while you check your answers. The gray cells on the "
        "Bonus sheet read your card.")
    L.bonus = [
        card_task("Patient name (First Last)",
                  f"Card A ({CARD_A}): what does your Patient name row (Card!B8) show? Format: First Last.",
                  card_a["Patient name (First Last)"],
                  "Chain two lookups: EncounterID → PatientID (row 6), then PatientID → names. Join with &\" \"&",
                  "Row 6 does the first hop (EncounterID → PatientID) once, and every patient row reuses it. Chaining through a "
                  "helper cell keeps each formula short and lets you check each hop on its own. The IF in front returns "
                  f"{NOT_FOUND} when row 6 already says so, which keeps Card B clean.",
                  live=f'=XLOOKUP({pid_of_a},{R(ps, "PatientID")},{R(ps, "FirstName")})&" "&'
                       f'XLOOKUP({pid_of_a},{R(ps, "PatientID")},{R(ps, "LastName")})'),
        card_task("Age at admission (years)",
                  f"Card A: how old was the patient, in completed years, on the encounter's AdmitDate (Card!B9)?",
                  card_a["Age at admission (years)"],
                  "Look up DOB and AdmitDate, then DATEDIF(…, …, \"Y\") from Lesson 2.3",
                  f"DATEDIF with \"Y\" counts completed years. The patient was born on {pa['DOB']:%m/%d/%Y} and admitted on "
                  f"{ea['AdmitDate']:%m/%d/%Y}, {days_to_bday} days before turning {naive_age}, so the answer is "
                  f"{card_a['Age at admission (years)']}. "
                  f"YEAR(admit) − YEAR(DOB) would give {naive_age}, which is one year too old. The DOB comes from Patients "
                  "(via row 6) and the AdmitDate from Encounters (via row 5), so this one formula reads two different tables.",
                  live=f'=DATEDIF(XLOOKUP({pid_of_a},{R(ps, "PatientID")},{R(ps, "DOB")}),'
                       f'XLOOKUP("{CARD_A}",{R(es, "EncounterID")},{R(es, "AdmitDate")}),"Y")'),
        card_task("Attending (First Last, Credential)",
                  f"Card A: what does your Attending row (Card!B12) show? Format: First Last, Credential "
                  f"(for example, Isabella Nguyen, MD).",
                  card_a["Attending (First Last, Credential)"],
                  "Use the Attending ProviderID helper in row 7, three lookups, and & to join them",
                  "Three lookups on the same ProviderID return FirstName, LastName, and Credential, and & glues them together "
                  f"with a space and a comma. Showing the credential is safer than putting \"Dr.\" in front of every name, because "
                  f"some attendings are nurse practitioners (NP) or physician assistants (PA). {pr['FirstName']} {pr['LastName']} "
                  f"is a {pr['Credential']}" + (", a doctor of osteopathic medicine." if pr["Credential"] == "DO" else "."),
                  live=f'=XLOOKUP({prv_of_a},{R(prs, "ProviderID")},{R(prs, "FirstName")})&" "&'
                       f'XLOOKUP({prv_of_a},{R(prs, "ProviderID")},{R(prs, "LastName")})&", "&'
                       f'XLOOKUP({prv_of_a},{R(prs, "ProviderID")},{R(prs, "Credential")})'),
        card_task("Claim status",
                  f"Card A: what is the claim status (Card!B15)? Claims are matched by EncounterID, which is column B of the "
                  f"Claims sheet.",
                  card_a["Claim status"],
                  "XLOOKUP can search any column. VLOOKUP would need its table to start at column B",
                  f"The Claims key is ClaimID, but the card knows only the EncounterID. XLOOKUP searches the EncounterID column "
                  f"wherever it is. With VLOOKUP you'd start the table at column B: "
                  f"`=VLOOKUP(B$5,Claims!$B$2:$H${cls.last_row},6,FALSE)`. This claim was {ca['ClaimStatus'].lower()}"
                  + (f" ({ca['DenialReason']})" if ca["DenialReason"] else "") + ", so it goes on the care manager's follow-up list.",
                  live=f'=XLOOKUP("{CARD_A}",{R(cls, "EncounterID")},{R(cls, "ClaimStatus")})'),
        Task(f"Card B ({CARD_B}): how many of the ten output rows (Card!C6:C15) show exactly {NOT_FOUND}? "
             "All ten should. If any shows #N/A or #VALUE!, fix that row's formula in column B and copy it right again.",
             answer=card_b_count, title=f"Card B · rows showing {NOT_FOUND}",
             hint="Give every lookup an if_not_found, and let IF skip calculations (like DATEDIF) when the helper cell says Not found",
             solution=card_formula["PatientID"],
             summary=f'=IF(COUNTA(Card!C{CARD_FIRST}:C{CARD_FIRST + 9})=0,"",COUNTIF(Card!C{CARD_FIRST}:C{CARD_FIRST + 9},"{NOT_FOUND}"))',
             fill={"range": f"Card!B{CARD_FIRST}:C{CARD_FIRST}", "formula": card_formula["PatientID"]},
             live=f'=IF(ISNA(MATCH("{CARD_B}",{R(es, "EncounterID")},0)),{card_b_count},0)',
             explanation=(
                 f"{CARD_B} contains the letter O where the real ID {CARD_B_REAL} has a zero. To Excel they're different text, "
                 "so every lookup comes back empty-handed. That's the most common reason a lookup fails on data that looks "
                 "right. The sample solution shown is row 6. Three techniques make the whole card fail politely:\n\n"
                 f"1. Every XLOOKUP gets an if_not_found of \"{NOT_FOUND}\".\n"
                 f"2. Chained lookups pass \"{NOT_FOUND}\" along. The inner XLOOKUP returns it, and the outer one can't find a "
                 f"DxCode called \"{NOT_FOUND}\", so it returns its own \"{NOT_FOUND}\".\n"
                 f"3. Rows that calculate (the name join and DATEDIF) test the helper cell first with "
                 f"`IF(B6=\"{NOT_FOUND}\",\"{NOT_FOUND}\",…)`, because DATEDIF of a text value would return #VALUE!.\n\n"
                 "Your finished card should match this:\n\n" + card_table + "\n\n"
                 "Every formula on the card (type them in column B, then copy them to column C):\n\n"
                 + "\n".join(f"- Row {card_row[k]} · {k}: `{f}`" for k, f in card_rows))),
    ]

    # ------------------------------------------------------------------ workbook extras
    L.start_notes = [
        "Encounters is sorted by AdmitDate, oldest first. Patients holds only Dr. Nguyen's panel (185 patients), so a patient "
        "who isn't on the panel won't be found there. That's intentional (task 13).",
        "MRN is stored as 8-character TEXT on purpose (leading zeros matter). The green triangles in that column are Excel's "
        "'number stored as text' warning; leave them as they are.",
        "BMI = 703 × WeightLb ÷ HeightIn², rounded to 1 decimal place. Age is the patient's age in completed years on 12/31/2025.",
        "The Card sheet is for the bonus challenge.",
    ]
    L.sheet_order = ["Start Here", "Practice", "Encounters", "Patients", "Referrals", "Providers", "Diagnoses", "Payers",
                     "Departments", "Claims", "Budget", "BMITiers", "AgeBands", "Bonus", "Card", "Answer Key", "Bonus Key"]

    @L.customize
    def _card(wb, lesson, selftest):
        ws = wb.create_sheet("Card")
        ws.sheet_properties.tabColor = "BF9000"
        ws["A1"] = "Encounter lookup card"
        ws["A1"].font = Font(bold=True, size=16, color=NAVY)
        ws["A2"] = ("Type one formula in each yellow cell of column B (B6:B15), then copy B6:B15 to C6:C15. Don't put a $ in front "
                    "of the B when you refer to B5, so the reference becomes C5 when you copy right. Every Card B cell must show "
                    "Not found. The gray cells on the Bonus sheet check your card.")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws["A2"].font = Font(italic=True, color="404040")
        ws.merge_cells("A2:D2")
        ws.row_dimensions[2].height = 34
        for j, h in enumerate(["Field", "Card A", "Card B", "Where it comes from"], 1):
            c = ws.cell(row=4, column=j, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEADER_FILL
            c.border = BOX
        id_fill = PatternFill("solid", fgColor="DDEBF7")
        ws["A5"] = "EncounterID"
        ws["A5"].font = Font(bold=True)
        for col, v in (("B", CARD_A), ("C", CARD_B)):
            c = ws[f"{col}5"]
            c.value = v
            c.font = Font(bold=True, color="1F4E79")
            c.fill = id_fill
            c.border = BOX
        ws["D5"] = "Input: keep these two IDs while you check the bonus"
        sources = {
            "PatientID": "Encounters (helper cell)", "Attending ProviderID": "Encounters (helper cell)",
            "Patient name (First Last)": "Patients, via row 6", "Age at admission (years)": "Patients DOB + Encounters AdmitDate",
            "Primary diagnosis": "Encounters → Diagnoses", "Unit (department)": "Encounters → Departments",
            "Attending (First Last, Credential)": "Providers, via row 7", "Attending specialty": "Providers, via row 7",
            "Payer": "Encounters → Payers", "Claim status": "Claims (search the EncounterID column)",
        }
        for label, _f in card_rows:
            r = card_row[label]
            ws.cell(row=r, column=1, value=label).border = BOX
            for col in (2, 3):
                c = ws.cell(row=r, column=col)
                c.fill = INPUT_FILL
                c.border = INPUT_BORDER
                c.alignment = Alignment(horizontal="left")
            d = ws.cell(row=r, column=4, value=sources[label])
            d.font = Font(italic=True, color="7F7F7F")
            d.border = BOX
        ws.row_dimensions[5].height = 18
        for col, wdt in zip("ABCD", (34, 60, 30, 46)):
            ws.column_dimensions[col].width = wdt
        ws.freeze_panes = "B6"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if selftest:
            # The self-test fills only the rows the bonus tasks read; the other rows feed them (and Card B's count),
            # so write the whole card here first. The task fills then rewrite their own rows with the same formulas.
            for label, f in card_rows:
                r = card_row[label]
                for col in ("B", "C"):
                    lesson.set_formula(ws, f"{col}{r}", _shift(f, col), dynamic=False)

    return L


def _shift(formula: str, col: str) -> str:
    """The card formulas are written for column B; column C is the same formula copied one column right."""
    if col == "B":
        return formula
    from openpyxl.formula.translate import Translator
    return Translator(formula, origin="B6").translate_formula("C6")
