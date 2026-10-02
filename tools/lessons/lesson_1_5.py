"""Lesson 1.5 · Relative, Absolute & Mixed References.

Every data sheet is a plain range, not an Excel Table, on purpose. Inside a Table, Excel writes structured
references such as [@Total] that don't use $ signs, and this lesson is about A1 references and where the
$ goes. (Tables and structured references are Lesson 3.1.)

Sheets
  Expenses        2025 actual operating expense, Bluestone Memorial Hospital (F01), 16 departments x 7
                  expense categories (budget.csv), with formula totals in column J and row 21. Learners fill
                  a "% of Total" column (absolute reference).
  4 West OT       December 2025 overtime for 4 West's hourly staff (shifts.csv), with ONE overtime-multiplier
                  cell (B3). Learners fill an OTPay column (rate-in-one-cell pattern) and run a what-if.
  Staffing Grid   (customize) census values down column A x HPPD targets across row 5. Learners fill the
                  grid with one mixed-reference formula.
  Q4 Summary      (customize) same layout as Oct/Nov/Dec. Learners fill it with one 3-D formula.
  Oct, Nov, Dec   identical layouts: monthly actuals by department x category (budget.csv), for 3-D sums.
  Plan 2026       (customize, bonus) 2025 actual x (1 + category inflation row) x (1 + department growth
                  column), filled with one formula that mixes relative, row-locked, and column-locked refs.

Every answer is computed in Python from data/*.csv. The planning assumptions (overtime multiplier,
category inflation, department volume growth) are inputs shown in the workbook, not answers.
"""
from __future__ import annotations

import re
from collections import defaultdict

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

from xlcourse import Lesson, Task, data
from xlcourse.lesson import BOX, HEADER_FILL, INPUT_BORDER, INPUT_FILL, NAVY
from xlcourse.xlfn import to_file_formula

CODE = "1.5"

FACILITY = "F01"                      # Bluestone Memorial Hospital
CATS = ["Salaries & Wages", "Employee Benefits", "Medical Supplies", "Pharmaceuticals",
        "Purchased Services", "Equipment & Maintenance", "Other Operating"]
MONTHS = {10: "Oct", 11: "Nov", 12: "Dec"}
MONTH_NAMES = {10: "October", 11: "November", 12: "December"}

OT_MULT = 1.5                         # time-and-a-half (fictional Bluestone policy, see sheet note)
WHATIF_MULT = 2.0                     # task 7: double time
HPPD = [7.0, 7.5, 8.0, 8.5, 9.0, 9.5, 10.0]

# Bonus planning assumptions (inputs on the Plan 2026 sheet)
INFLATION = dict(zip(CATS, [0.035, 0.060, 0.040, 0.075, 0.030, 0.025, 0.020]))
GROWTH = {  # 2026 volume plan by department (DeptID -> growth)
    "D100": 0.030, "D110": 0.010, "D111": 0.005, "D120": 0.020, "D130": 0.025, "D140": 0.035,
    "D150": 0.045, "D160": 0.035, "D170": 0.015, "D180": -0.020, "D190": 0.030, "D195": 0.050,
    "D196": 0.040, "D500": 0.020, "D510": 0.025, "D520": 0.030,
}
PHARMA_WHATIF = 0.09

ASSUMPTION_FONT = Font(bold=True, color="0000FF")
ASSUMPTION_FILL = PatternFill("solid", fgColor="DDEBF7")
TOTAL_FONT = Font(bold=True)
TOTAL_BORDER = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13, color=NAVY)
NOTE_FONT = Font(italic=True, size=10, color="595959")
MONEY0 = "#,##0"

_REF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def copied(formula: str, cols: int, rows: int) -> str:
    """What an A1 formula becomes when copied `cols` columns right and `rows` rows down ($ parts stay put)."""
    def move(m):
        cabs, col, rabs, row = m.groups()
        if not cabs:
            col = get_column_letter(column_index_from_string(col) + cols)
        if not rabs:
            row = str(int(row) + rows)
        return f"{cabs}{col}{rabs}{row}"
    return _REF.sub(move, formula)


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="01-foundations", slug="05-cell-references",
        title="Relative, Absolute & Mixed References", level="Beginner", minutes=45,
        objectives=[
            "Predict how relative references change when a formula is copied",
            "Lock references with $ (absolute) and use F4 to toggle",
            "Build two-way grids with mixed references ($A1 and A$1)",
            "Reference other sheets and sum the same cell across sheets (3-D references)",
        ],
        data_note="Bluestone Memorial Hospital's 2025 operating expenses by department and category, the same "
                  "breakdown for October, November, and December 2025, December 2025 overtime for the hourly staff "
                  "of Medical-Surgical 4 West, and a 4 West staffing grid.",
    )

    # ================================================================== source data
    deps = data.index(data.load("departments"), "DeptID")
    f01 = sorted((d for d in deps if deps[d]["FacilityID"] == FACILITY), key=lambda d: deps[d]["CostCenter"])
    assert set(f01) == set(GROWTH), "growth assumptions must cover every Bluestone Memorial department"
    budget = [r for r in data.load("budget")
              if r["FacilityID"] == FACILITY and r["FiscalYear"] == 2025 and r["LineType"] == "Expense"]

    def actuals(months) -> dict:
        out = defaultdict(int)
        for r in budget:
            if r["FiscalMonth"] in months:
                out[(r["DeptID"], r["Category"])] += r["ActualAmount"]
        return out

    annual = actuals(set(range(1, 13)))
    monthly = {m: actuals({m}) for m in MONTHS}

    # Shared layout for Expenses, Oct, Nov, Dec, Q4 Summary, and Plan 2026
    HDR, FIRST = 4, 5
    LAST = FIRST + len(f01) - 1                     # 20
    TOT = LAST + 1                                  # 21: "Hospital total" row
    cat_col = {c: get_column_letter(3 + i) for i, c in enumerate(CATS)}   # C..I
    C0, C1 = cat_col[CATS[0]], cat_col[CATS[-1]]    # "C", "I"
    TCOL = get_column_letter(3 + len(CATS))         # J: department total
    SCOL = get_column_letter(4 + len(CATS))         # K: % of Total (Expenses only)
    row_of = {d: FIRST + i for i, d in enumerate(f01)}
    name_of = {d: deps[d]["DeptName"] for d in f01}
    assert (C0, C1, TCOL, SCOL, LAST, TOT) == ("C", "I", "J", "K", 20, 21)

    def grid_rows(values: dict) -> list[dict]:
        rows = []
        for d in f01:
            r = row_of[d]
            row = {"Department": name_of[d], "CostCenter": deps[d]["CostCenter"]}
            row.update({c: values[(d, c)] for c in CATS})
            row["Total"] = f"=SUM({C0}{r}:{C1}{r})"
            rows.append(row)
        tot = {"Department": "Hospital total", "CostCenter": None}
        for c in CATS:
            tot[c] = f"=SUM({cat_col[c]}{FIRST}:{cat_col[c]}{LAST})"
        tot["Total"] = f"=SUM({TCOL}{FIRST}:{TCOL}{LAST})"
        rows.append(tot)
        return rows

    grid_cols = [("Department", "Department"), ("CostCenter", "Cost Center")] + [(c, c) for c in CATS] + [("Total", "Total")]
    grid_fmt = {c: MONEY0 for c in CATS} | {"Total": MONEY0}
    grid_w = {"Department": 27, "Cost Center": 8, "Total": 13} | {c: 16 for c in CATS}

    exp = L.add_table_sheet(
        "Expenses", grid_rows(annual), columns=grid_cols, as_table=False, start_row=HDR,
        extra_cols=["% of Total"], formats=grid_fmt | {"% of Total": "0.0%"}, widths=grid_w | {"% of Total": 10},
        notes=["Bluestone Memorial Hospital · 2025 operating expense by department (actual, $)",
               "Column J (Total) and row 21 (Hospital total) are formulas: click one and read the formula bar. "
               "Column K is yours to fill."],
    )
    assert exp.col("% of Total") == SCOL and exp.first_row == FIRST
    for m, short in MONTHS.items():
        L.add_table_sheet(
            short, grid_rows(monthly[m]), columns=grid_cols, as_table=False, start_row=HDR,
            formats=grid_fmt, widths=grid_w,
            notes=[f"Bluestone Memorial Hospital · operating expense actuals · {MONTH_NAMES[m]} 2025 ($)",
                   "Oct, Nov, and Dec share one layout: departments in rows 5–20, categories in columns C–I, "
                   "totals in column J and row 21."],
        )

    # --- 4 West overtime, December 2025 (hourly staff; the salaried nurse manager is excluded)
    emp = data.index(data.load("employees"), "EmployeeID")
    dec_shifts = [s for s in data.load("shifts") if s["DeptID"] == "D110" and s["ShiftDate"].year == 2025
                  and s["ShiftDate"].month == 12 and s["ShiftStatus"] == "Worked"]
    agg = defaultdict(lambda: {"n": 0, "worked": 0.0, "ot": 0.0})
    for s in dec_shifts:
        a = agg[s["EmployeeID"]]
        a["n"] += 1
        a["worked"] += s["WorkedHours"]
        a["ot"] += s["HoursOverSchedule"]
    ot_rows = []
    for eid in sorted(agg):
        e = emp[eid]
        if e["JobTitle"] == "Nurse Manager":
            continue
        ot_rows.append({"EmployeeID": eid, "Name": f"{e['FirstName']} {e['LastName']}", "JobTitle": e["JobTitle"],
                        "HourlyRate": e["HourlyRate"], "ShiftsWorked": agg[eid]["n"],
                        "WorkedHours": round(agg[eid]["worked"], 2), "OTHours": round(agg[eid]["ot"], 2)})
    OT_HDR = 5
    ot = L.add_table_sheet(
        "4 West OT", ot_rows, as_table=False, start_row=OT_HDR,
        columns=["EmployeeID", "Name", "JobTitle", "HourlyRate", "ShiftsWorked", "WorkedHours", "OTHours"],
        extra_cols=["OTPay"],
        formats={"HourlyRate": "#,##0.00", "WorkedHours": "0.00", "OTHours": "0.00", "OTPay": "#,##0.00"},
        widths={"EmployeeID": 20, "Name": 22, "JobTitle": 29, "HourlyRate": 11, "OTPay": 12},
        notes=["Medical-Surgical 4 West · overtime for hourly staff · December 2025",
               "OTHours = December's total of paid hours worked beyond each shift's scheduled hours (time clock). "
               "Bluestone pays each OT hour at the base HourlyRate × the overtime multiplier in B3 (a simplified, fictional "
               "pay rule)."],
    )
    MULT_CELL = "B3"
    o_first, o_last = ot.first_row, ot.last_row
    OT_RATE, OT_HRS, OT_PAY = ot.col("HourlyRate"), ot.col("OTHours"), ot.col("OTPay")
    assert (OT_RATE, OT_HRS, OT_PAY) == ("D", "G", "H")
    OTS = "'4 West OT'"

    def otr(col):
        return f"{OTS}!{col}{o_first}:{col}{o_last}"

    ot_total = sum(r["OTHours"] * r["HourlyRate"] * OT_MULT for r in ot_rows)
    ot_total_whatif = sum(r["OTHours"] * r["HourlyRate"] * WHATIF_MULT for r in ot_rows)
    rn_rows = [r for r in ot_rows if r["JobTitle"] == "Registered Nurse"]
    star = max(rn_rows, key=lambda r: (r["OTHours"], r["EmployeeID"]))
    star_row = o_first + ot_rows.index(star)
    star_pay = star["OTHours"] * star["HourlyRate"] * OT_MULT

    # --- Staffing grid: 4 West census values (covering every 2025 census, in steps of 2) x HPPD targets
    c25 = [r["MidnightCensus"] for r in data.load("daily_census") if r["DeptID"] == "D110" and r["CensusDate"].year == 2025]
    lo_c, hi_c = (min(c25) // 2) * 2, -(-max(c25) // 2) * 2
    census_vals = list(range(lo_c, hi_c + 1, 2))
    G_HDR, G_FIRST = 5, 6
    G_LAST = G_FIRST + len(census_vals) - 1
    G_C0 = "B"
    G_C1 = get_column_letter(1 + len(HPPD))         # H
    GS = "'Staffing Grid'"
    grid_rng = f"{GS}!{G_C0}{G_FIRST}:{G_C1}{G_LAST}"
    grid_total = sum(c * h for c in census_vals for h in HPPD)
    grid_total = int(round(grid_total)) if abs(grid_total - round(grid_total)) < 1e-9 else grid_total
    # predict task (the curriculum's own example): a generic grid with labels in row 2 and column A.
    # B3 holds =B$2*$A3; copied two columns right and three rows down (to D6). Deliberately NOT the Staffing
    # Grid formula, so this task doesn't hand out task 9's answer.
    pred_src, pred_dst, pred_formula = "B3", "D6", "B$2*$A3"
    pred_mixed = copied(pred_formula, 2, 3)
    assert pred_mixed == "D$2*$A6"
    # F4 task: a reference the guide doesn't use for its own F4 demo (the guide cycles J21)
    f4_ref = f"{C0}{TOT}"                                   # C21: hospital Salaries & Wages total
    f4_cycle = [f4_ref, f"${C0}${TOT}", f"{C0}${TOT}", f"${C0}{TOT}"]
    f4_answer = f4_cycle[3 % 4]

    # --- 3-D references
    onc = row_of["D150"]
    pharma_col = cat_col["Pharmaceuticals"]
    onc_pharma_q4 = sum(monthly[m][("D150", "Pharmaceuticals")] for m in MONTHS)
    q4_total = sum(v for m in MONTHS for v in monthly[m].values())
    month_total = {m: sum(monthly[m].values()) for m in MONTHS}
    moved_total = month_total[10] + month_total[12]          # Nov dragged to the right of Dec
    QS = "'Q4 Summary'"
    q4_rng = f"{QS}!{C0}{FIRST}:{C1}{LAST}"

    # --- Expenses: % of total
    dept_total = {d: sum(annual[(d, c)] for c in CATS) for d in f01}
    hosp_total = sum(dept_total.values())
    icu = row_of["D130"]
    icu_share = dept_total["D130"] / hosp_total
    exp_share_rng = f"Expenses!{SCOL}{FIRST}:{SCOL}{LAST}"
    relpred_row = FIRST + 9                            # task 1: J5 copied to J14
    assert copied(f"SUM({C0}{FIRST}:{C1}{FIRST})", 0, relpred_row - FIRST) == f"SUM({C0}{relpred_row}:{C1}{relpred_row})"

    # ================================================================== practice
    L.practice_intro = (
        "Tasks 1, 4, and 8 ask you to predict a formula or reference. Type your answer as text without the leading = "
        "(for example AVERAGE(B2:B9)), because Excel would calculate it if you typed the =. The other tasks want a "
        "formula in the yellow cell, or work on another sheet that a gray cell checks."
    )
    L.tasks = [
        Task(f"On the Expenses sheet, cell {TCOL}{FIRST} contains =SUM({C0}{FIRST}:{C1}{FIRST}), the Emergency Department's "
             f"total. If you copy {TCOL}{FIRST} and paste it into {TCOL}{relpred_row}, what formula will "
             f"{TCOL}{relpred_row} contain? Type it without the =, then click {TCOL}{relpred_row} to check.",
             answer=f"SUM({C0}{relpred_row}:{C1}{relpred_row})",
             accept=[f"=SUM({C0}{relpred_row}:{C1}{relpred_row})"],
             answer_display=f"`=SUM({C0}{relpred_row}:{C1}{relpred_row})`",
             solution=f"Count the move: from row {FIRST} to row {relpred_row} is {relpred_row - FIRST} rows down and 0 columns "
                      f"across. Add {relpred_row - FIRST} to every row number in the formula and leave the column letters alone.",
             hint="Relative references keep their distance from the formula cell", live=False,
             title=f"Predict: {TCOL}{FIRST} copied to {TCOL}{relpred_row}",
             explanation=f"Excel stores =SUM({C0}{FIRST}:{C1}{FIRST}) in {TCOL}{FIRST} as \"the seven cells to my left, on my row.\" "
                         f"Pasted into {TCOL}{relpred_row}, the same instruction points at {C0}{relpred_row}:{C1}{relpred_row}. "
                         "That's why one total formula works for every department."),
        Task(f"What share of Bluestone Memorial's 2025 operating expense came from the Intensive Care Unit? Divide the ICU's "
             f"total (Expenses!{TCOL}{icu}) by the hospital total (Expenses!{TCOL}{TOT}). Enter it as a percentage.",
             answer=icu_share, fmt="0.0%",
             solution=f"=Expenses!{TCOL}{icu}/Expenses!{TCOL}{TOT}",
             hint="Type =, click the Expenses tab, click the cell, type /, click the other cell, press Enter",
             title="ICU share of total expense",
             explanation="A reference to another sheet is the sheet name, an exclamation mark, then the cell. If you build it by "
                         "clicking, Excel writes the Expenses! part for you. Press Enter while you're still on the Expenses sheet, "
                         "and Excel takes you back to Practice."),
        Task(f"On the Expenses sheet, fill the yellow % of Total column ({SCOL}{FIRST}:{SCOL}{LAST}) with ONE formula: type it "
             f"in {SCOL}{FIRST}, lock the reference to the hospital total in {TCOL}{TOT}, then drag the fill handle down to row {LAST}. "
             "The gray cell adds up your column. If it shows #DIV/0!, the total reference moved when you copied.",
             answer=1.0, fmt="0.0%", title="% of Total column (absolute reference)",
             solution=f"={TCOL}{FIRST}/${TCOL}${TOT}",
             summary=f'=IF(COUNT({exp_share_rng})=0,"",SUM({exp_share_rng}))',
             fill={"range": exp_share_rng, "formula": f"={TCOL}{FIRST}/${TCOL}${TOT}"},
             live=f"=SUM(Expenses!{TCOL}{FIRST}:{TCOL}{LAST})/Expenses!{TCOL}{TOT}",
             hint=f"Click {TCOL}{TOT} in the formula, then press F4 (Mac: ⌘ + T)",
             explanation=f"Without the $ signs, {SCOL}{FIRST + 1} would hold ={TCOL}{FIRST + 1}/{TCOL}{TOT + 1}. Row {TOT + 1} is empty, "
                         f"so the result is #DIV/0!. ${TCOL}${TOT} stays put wherever you copy it, while {TCOL}{FIRST} moves to each "
                         f"department's row. ={TCOL}{FIRST}/{TCOL}${TOT} also works here, because you only copy down. Every "
                         "% of total column adds up to 100%, so the gray cell doubles as a sanity check."),
        Task(f"You type ={f4_ref} in a cell, leave the cursor in the reference, and press F4 three times "
             "(Mac: ⌘ + T three times). Which reference do you end up with? Type it without the =.",
             answer=f4_answer, accept=[f"={f4_answer}"], answer_display=f"`{f4_answer}`",
             solution=f"F4 cycles {' → '.join(f4_cycle)} → back to {f4_ref}. The third press gives **{f4_answer}**.",
             hint="Absolute, then row locked, then column locked…", live=False, title="F4 pressed three times",
             explanation=f"{f4_answer} locks the column ({C0}) but lets the row change. Try it: type the formula, press F4 three "
                         "times, and watch the formula bar. Press Esc afterwards so you don't leave the test formula in the sheet."),
        Task(f"On the 4 West OT sheet, {star['Name']} (row {star_row}) worked {star['OTHours']:.2f} overtime hours in December. "
             f"On this Practice sheet, write a formula for this employee's overtime pay: OTHours × HourlyRate × the overtime multiplier "
             f"in {MULT_CELL}. Refer to the cells; don't type the numbers.",
             answer=round(star_pay, 6), fmt="#,##0.00",
             solution=f"={OTS}!{OT_HRS}{star_row}*{OTS}!{OT_RATE}{star_row}*{OTS}!{MULT_CELL}",
             hint="The sheet name starts with a digit and contains spaces, so it needs single quotes",
             title=f"Overtime pay for {star['Name']}",
             explanation="Sheet names that contain spaces or start with a digit must be wrapped in single quotes: "
                         f"{OTS}!{MULT_CELL}. When you click the cells instead of typing, Excel adds the quotes for you. "
                         "No $ signs are needed here, because this formula is never copied."),
        Task(f"On the 4 West OT sheet, fill the yellow OTPay column ({OT_PAY}{o_first}:{OT_PAY}{o_last}) with one formula: "
             f"OTHours × HourlyRate × the overtime multiplier in {MULT_CELL}. Type it in {OT_PAY}{o_first} and copy it down. "
             "The gray cell totals your column: what did December's overtime cost?",
             answer=round(ot_total, 6), fmt="#,##0.00", title="OTPay column (rate in one cell)",
             solution=f"={OT_HRS}{o_first}*{OT_RATE}{o_first}*${MULT_CELL[0]}${MULT_CELL[1:]}",
             summary=f'=IF(COUNT({otr(OT_PAY)})=0,"",SUM({otr(OT_PAY)}))',
             fill={"range": otr(OT_PAY), "formula": f"={OT_HRS}{o_first}*{OT_RATE}{o_first}*${MULT_CELL[0]}${MULT_CELL[1:]}"},
             live=f"=SUMPRODUCT({otr(OT_HRS)},{otr(OT_RATE)})*{OTS}!{MULT_CELL}",
             hint=f"Lock {MULT_CELL}; let the row references move",
             explanation=f"{OT_HRS}{o_first} and {OT_RATE}{o_first} are relative, so each row uses its own employee's hours and rate. "
                         f"${MULT_CELL[0]}${MULT_CELL[1:]} is absolute, so all {len(ot_rows)} rows share the one multiplier cell. If you had "
                         f"typed *1.5 into every formula, a policy change would mean editing {len(ot_rows)} formulas."),
        Task(f"Finance asks what December's overtime would have cost at double time. Change {OTS}!{MULT_CELL} to "
             f"{WHATIF_MULT:g}, read task 6's gray cell, and type that amount here as a number. Then set {MULT_CELL} back to "
             f"{OT_MULT:g}.",
             answer=round(ot_total_whatif, 6), fmt="#,##0.00", tol=0.01,
             solution=f"1. Click **{OTS}!{MULT_CELL}**, type `{WHATIF_MULT:g}`, and press **Enter**.\n"
                      "2. Read the gray cell in task 6 and type that amount into this task's yellow cell.\n"
                      f"3. Change {MULT_CELL} back to `{OT_MULT:g}` so tasks 5 and 6 show ✔ again.",
             live=f"=SUMPRODUCT({otr(OT_HRS)},{otr(OT_RATE)})*{WHATIF_MULT:g}",
             hint="One edit updates the whole column. That's the point of the rate cell", title="What-if: double time",
             explanation=f"Because every OTPay formula points at {MULT_CELL}, one edit recalculates all {len(ot_rows)} rows and the "
                         "total. This is the payoff of the rate-in-one-cell pattern: assumptions live in one labeled cell, so "
                         "what-if questions take seconds."),
        Task(f"A grid has labels across row 2 and down column A, and cell {pred_src} holds ={pred_formula}. If you copy "
             f"{pred_src} to {pred_dst}, what formula will {pred_dst} contain? Type it without the =.",
             answer=pred_mixed, accept=["=" + pred_mixed], answer_display=f"`={pred_mixed}`",
             solution=f"The move is 2 columns right (B to D) and 3 rows down (row 3 to row 6). In B$2 the row is locked, so only "
                      f"the column moves and it becomes D$2. In $A3 the column is locked, so only the row moves and it becomes "
                      f"$A6. The answer is **={pred_mixed}**.",
             hint="A $ freezes only the part right after it", live=False,
             title=f"Predict: {pred_src} copied to {pred_dst}",
             explanation="B$2 always reads row 2 (the labels across the top) in the current column. $A3 always reads column A "
                         "(the labels down the side) in the current row. Together they make every cell multiply its own column "
                         "label by its own row label, which is exactly what a two-way grid needs. You'll build one in the next task."),
        Task(f"On the Staffing Grid sheet, fill the yellow grid {G_C0}{G_FIRST}:{G_C1}{G_LAST} with ONE formula: nursing hours "
             f"needed per day = census (column A) × HPPD target (row {G_HDR}). Type it in B{G_FIRST}, copy it across to "
             f"{G_C1}{G_FIRST}, then down to row {G_LAST}. The gray cell adds up the whole grid.",
             answer=grid_total, title="Staffing grid (mixed references)",
             solution=f"=$A{G_FIRST}*B${G_HDR}",
             summary=f'=IF(COUNT({grid_rng})=0,"",SUM({grid_rng}))',
             fill={"range": grid_rng, "formula": f"=$A{G_FIRST}*B${G_HDR}"},
             live=f"=SUM({GS}!A{G_FIRST}:A{G_LAST})*SUM({GS}!B{G_HDR}:{G_C1}{G_HDR})",
             hint="Lock the column of the census and the row of the HPPD",
             explanation=f"$A{G_FIRST} keeps every formula looking at column A, and B${G_HDR} keeps every formula looking at row "
                         f"{G_HDR}. With =$A${G_FIRST}*$B${G_HDR}, every cell would repeat {census_vals[0] * HPPD[0]:g}, the B{G_FIRST} "
                         f"result. With no $ at all, C{G_FIRST} would compute ={G_C0}{G_FIRST}*C{G_HDR}, multiplying the previous "
                         "result by the HPPD instead of using the census. "
                         "Cross-check: the sum of a multiplication grid equals (sum of the census values) × (sum of the HPPD "
                         "targets), which is how the key's live formula works."),
        Task(f"Using one 3-D reference, what did Oncology (row {onc} on the Oct, Nov, and Dec sheets) spend on Pharmaceuticals "
             f"(column {pharma_col}) in Q4 2025?",
             answer=onc_pharma_q4,
             solution=f"=SUM(Oct:Dec!{pharma_col}{onc})",
             hint="Type =SUM(, click the Oct tab, Shift+click the Dec tab, then click the cell",
             title="3-D SUM: Oncology pharmaceuticals in Q4",
             explanation=f"Oct:Dec! means every sheet from the Oct tab through the Dec tab, so the formula adds {pharma_col}{onc} on all "
                         f"three. It's the same as =Oct!{pharma_col}{onc}+Nov!{pharma_col}{onc}+Dec!{pharma_col}{onc}, but it stays "
                         "short no matter how many monthly sheets sit between the two end tabs."),
        Task(f"Fill the yellow grid on the Q4 Summary sheet ({C0}{FIRST}:{C1}{LAST}) with ONE 3-D formula that adds the same "
             f"cell on the Oct, Nov, and Dec sheets. Type it in {C0}{FIRST}, copy it across to {C1}{FIRST}, then down to row "
             f"{LAST}. The gray cell adds up your grid: what was Bluestone Memorial's Q4 operating expense?",
             answer=q4_total, title="Q4 Summary built with a 3-D formula",
             solution=f"=SUM(Oct:Dec!{C0}{FIRST})",
             summary=f'=IF(COUNT({q4_rng})=0,"",SUM({q4_rng}))',
             fill={"range": q4_rng, "formula": f"=SUM(Oct:Dec!{C0}{FIRST})"},
             live=f"=SUM(Oct:Dec!{C0}{FIRST}:{C1}{LAST})",
             hint="3-D references copy like ordinary relative references",
             explanation=f"The cell part of Oct:Dec!{C0}{FIRST} is relative, so copying moves it exactly like a normal reference: "
                         f"{C1}{LAST} ends up as =SUM(Oct:Dec!{C1}{LAST}). This only works because the three monthly sheets share "
                         "one layout. A 3-D reference adds cells by position, not by department name. Q4 Summary sits to the left "
                         "of Oct, outside the Oct:Dec range, so it never adds itself."),
        Task(f"A colleague drags the Nov tab to the right of the Dec tab. What would =SUM(Oct:Dec!{TCOL}{TOT}) return then? "
             f"({TCOL}{TOT} is each month's hospital total.) Predict it with a formula. If you test it by moving the tab, "
             "drag Nov back between Oct and Dec afterwards.",
             answer=moved_total,
             solution=f"=Oct!{TCOL}{TOT}+Dec!{TCOL}{TOT}",
             hint="A 3-D range is defined by tab positions, not by month names",
             title="3-D gotcha: a tab moved out of the range",
             explanation="Oct:Dec! means \"Oct, Dec, and whatever tabs sit between them right now.\" Once Nov is dragged past Dec, "
                         "it's outside the sandwich, so the total silently drops November. A sheet dragged in between the end "
                         "tabs gets added just as silently. Keep 3-D sheets in order, and consider empty \"bookend\" sheets "
                         "(for example First and Last) that mark the range."),
    ]

    # ================================================================== bonus: 2026 operating expense plan
    PS = "'Plan 2026'"
    plan_rng = f"{PS}!{C0}{FIRST}:{C1}{LAST}"
    INF_ROW = 3
    GROW_COL = "B"
    plan = {(d, c): annual[(d, c)] * (1 + INFLATION[c]) * (1 + GROWTH[d]) for d in f01 for c in CATS}
    plan_total = sum(plan.values())
    plan_dept = {d: sum(plan[(d, c)] for c in CATS) for d in f01}
    pct_up = {d: plan_dept[d] / dept_total[d] - 1 for d in f01}
    ranked = sorted(f01, key=lambda d: pct_up[d], reverse=True)
    top, runner = ranked[0], ranked[1]
    assert pct_up[top] - pct_up[runner] > 0.0015, "B4 needs a clear winner"
    fastest_volume = max(f01, key=lambda d: GROWTH[d])
    assert fastest_volume != top, "B4 should not be answerable from the growth column alone"
    lab = row_of["D500"]
    sup_col = cat_col["Medical Supplies"]
    lab_sup = plan[("D500", "Medical Supplies")]
    pharma_rise = sum(annual[(d, "Pharmaceuticals")] * (1 + GROWTH[d]) for d in f01) * (PHARMA_WHATIF - INFLATION["Pharmaceuticals"])
    plan_sumproduct = (f"SUMPRODUCT(Expenses!{C0}{FIRST}:{C1}{LAST}*(1+{PS}!{C0}{INF_ROW}:{C1}{INF_ROW})"
                       f"*(1+{PS}!{GROW_COL}{FIRST}:{GROW_COL}{LAST}))")
    plan_formula = f"=Expenses!{C0}{FIRST}*(1+{C0}${INF_ROW})*(1+${GROW_COL}{FIRST})"

    L.bonus_title = "Bonus: Build the 2026 expense plan"
    L.bonus_scenario = (
        "Finance needs a first-draft 2026 operating expense plan for Bluestone Memorial. The rule: each 2026 cell = the 2025 "
        f"actual for the same department and category (Expenses sheet) × (1 + that category's price inflation in row {INF_ROW} "
        f"of the Plan 2026 sheet) × (1 + that department's volume growth in column {GROW_COL}). The Plan 2026 sheet uses the same "
        f"rows and columns as the Expenses sheet. Column {TCOL} and row {TOT} of Plan 2026 already total whatever you put in the grid."
    )
    L.bonus = [
        Task(f"Fill the yellow grid on the Plan 2026 sheet ({C0}{FIRST}:{C1}{LAST}) with ONE formula typed in {C0}{FIRST} and "
             "copied across and down. The gray cell adds up your grid: what is the 2026 plan total? (Don't round.)",
             answer=round(plan_total, 6), fmt="#,##0.00", title="2026 plan grid (three kinds of reference in one formula)",
             solution=plan_formula,
             summary=f'=IF(COUNT({plan_rng})=0,"",SUM({plan_rng}))',
             fill={"range": plan_rng, "formula": plan_formula},
             live="=" + plan_sumproduct,
             hint=f"The 2025 actual moves both ways, the inflation row is locked, and the growth column is locked",
             explanation=f"Expenses!{C0}{FIRST} is fully relative, because each plan cell needs the matching 2025 cell. {C0}${INF_ROW} "
                         f"locks the row, so every department reads the inflation in row {INF_ROW} of its own category column. "
                         f"${GROW_COL}{FIRST} locks the column, so every category reads the growth in column {GROW_COL} of its own "
                         "department row. One formula, three reference types, 112 correct cells."),
        Task(f"What is the 2026 plan for Laboratory · Medical Supplies (Plan 2026, row {lab}, column {sup_col})? "
             "Reference the cell in your grid.",
             answer=round(lab_sup, 6), fmt="#,##0.00", tol=0.5,
             solution=f"={PS}!{sup_col}{lab}",
             live=f"=Expenses!{sup_col}{lab}*(1+{PS}!{sup_col}{INF_ROW})*(1+{PS}!{GROW_COL}{lab})",
             hint="Check it by hand: 2025 amount × (1 + inflation) × (1 + growth)",
             title="Spot-check: Laboratory · Medical Supplies",
             explanation=f"Spot-checking a cell far from where you typed the formula proves the copy worked. {sup_col}{lab} should hold "
                         f"=Expenses!{sup_col}{lab}*(1+{sup_col}${INF_ROW})*(1+${GROW_COL}{lab}): the inflation is still read from row "
                         f"{INF_ROW} and the growth from column {GROW_COL}."),
        Task("By what percentage would Bluestone Memorial's total operating expense grow from 2025 to 2026 under this plan? "
             "Enter it as a percentage.",
             answer=plan_total / hosp_total - 1, fmt="0.0%",
             solution=f"={PS}!{TCOL}{TOT}/Expenses!{TCOL}{TOT}-1",
             live=f"={plan_sumproduct}/Expenses!{TCOL}{TOT}-1",
             hint="New ÷ old − 1",
             title="Hospital-wide growth, 2025 → 2026",
             explanation=f"{PS}!{TCOL}{TOT} is the 2026 grand total and Expenses!{TCOL}{TOT} is the 2025 grand total. "
                         "New ÷ old − 1 turns the two totals into a growth rate."),
        Task("Which department's total expense is planned to grow by the largest percentage? Type the department name exactly "
             "as it appears in column A.",
             answer=name_of[top],
             solution=f"1. In {PS}!K{FIRST}, type `={TCOL}{FIRST}/Expenses!{TCOL}{FIRST}-1` and copy it down to K{LAST}. Both "
                      "references are relative, so each row compares its own department.\n"
                      f"2. Format K{FIRST}:K{LAST} as a percentage and find the largest value (`=MAX(K{FIRST}:K{LAST})` helps).\n"
                      "3. Read the department name in column A of that row.",
             hint="Add a helper column that divides each department's 2026 total by its 2025 total",
             live=False, title="Fastest-growing department",
             explanation=f"{name_of[top]} grows {pct_up[top]:.1%}, ahead of {name_of[runner]} at {pct_up[runner]:.1%}, even though "
                         f"{name_of[fastest_volume]} has the highest volume growth ({GROWTH[fastest_volume]:.1%}). "
                         f"{name_of[top]}'s spending is heavy in Pharmaceuticals, the category with the steepest inflation "
                         f"({INFLATION['Pharmaceuticals']:.1%}), so its mix pushes the total up faster."),
        Task(f"Pharmaceutical prices are the shakiest assumption. If pharmaceutical inflation were {PHARMA_WHATIF:.1%} instead of "
             f"{INFLATION['Pharmaceuticals']:.1%}, how many dollars higher would the 2026 plan total be? Change the one input cell, "
             "compare the totals, then put it back. Round to the nearest dollar.",
             answer=round(pharma_rise), fmt="#,##0", tol=1,
             solution=f"1. Note the plan total in {PS}!{TCOL}{TOT} (or the gray cell in B1).\n"
                      f"2. Change {PS}!{pharma_col}{INF_ROW} from {INFLATION['Pharmaceuticals']:.1%} to {PHARMA_WHATIF:.1%}.\n"
                      f"3. Subtract the old total from the new one and round to the nearest dollar.\n"
                      f"4. Set {pharma_col}{INF_ROW} back to {INFLATION['Pharmaceuticals']:.1%}.",
             live=f"=ROUND(SUMPRODUCT(Expenses!{pharma_col}{FIRST}:{pharma_col}{LAST}*(1+{PS}!{GROW_COL}{FIRST}:{GROW_COL}{LAST}))"
                  f"*({PHARMA_WHATIF}-{PS}!{pharma_col}{INF_ROW}),0)",
             hint="Because of the $ signs, one edit flows to all 16 Pharmaceuticals cells",
             title="Sensitivity: pharmaceutical inflation",
             explanation="Only the Pharmaceuticals column reads that input, so only those 16 cells change. Each one rises by its 2025 "
                         f"amount × (1 + growth) × {(PHARMA_WHATIF - INFLATION['Pharmaceuticals']) * 100:.1f} percentage points. Testing the shakiest "
                         "assumption like this is called a sensitivity check, and it only takes seconds because the assumption "
                         "lives in one cell."),
    ]

    L.start_notes = [
        "The data sheets are plain ranges, not Excel Tables, so the formulas you write use ordinary A1 references "
        "with $ signs. (Tables use structured references instead; see Lesson 3.1.)",
        "Staffing Grid, Q4 Summary, and Plan 2026 are layout sheets with yellow grids for you to fill with one formula each.",
        "Keep the Oct, Nov, and Dec tabs together and in that order: 3-D formulas depend on tab positions.",
    ]
    L.sheet_order = ["Start Here", "Practice", "Expenses", "4 West OT", "Staffing Grid", "Q4 Summary", "Oct", "Nov", "Dec",
                     "Bonus", "Plan 2026", "Answer Key", "Bonus Key"]

    # ================================================================== layout sheets & styling
    def header_row(ws, row, labels, start_col=1):
        for j, h in enumerate(labels, start_col):
            c = ws.cell(row=row, column=j, value=h)
            c.font = HDR_FONT
            c.fill = HEADER_FILL
            c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def titled(ws, title, *notes):
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        for i, n in enumerate(notes, 2):
            ws.cell(row=i, column=1, value=n).font = NOTE_FONT

    def print_setup(ws):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    def style_grid_sheet(ws):
        """Common look for the department x category sheets (header wrap, total row, widths)."""
        ws.row_dimensions[HDR].height = 32
        for cell in ws[HDR]:
            if cell.value is not None:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in range(1, 3 + len(CATS) + 1):
            c = ws.cell(row=TOT, column=col)
            c.font = TOTAL_FONT
            c.border = TOTAL_BORDER

    def assumption(cell, fmt):
        cell.font = ASSUMPTION_FONT
        cell.fill = ASSUMPTION_FILL
        cell.border = BOX
        cell.number_format = fmt

    def yellow(ws, rng, fmt):
        for row in ws[rng]:
            for c in row:
                c.fill = INPUT_FILL
                c.border = INPUT_BORDER
                c.number_format = fmt

    def dept_block(ws, total_label="Hospital total"):
        """Department + cost center labels, J totals, and the total row, matching the Expenses layout."""
        for d in f01:
            r = row_of[d]
            ws.cell(row=r, column=1, value=name_of[d]).border = BOX
        for r in range(FIRST, LAST + 1):
            c = ws.cell(row=r, column=3 + len(CATS), value=f"=SUM({C0}{r}:{C1}{r})")
            c.number_format = MONEY0
            c.border = BOX
        ws.cell(row=TOT, column=1, value=total_label)
        for k in range(len(CATS) + 1):
            col = get_column_letter(3 + k)
            c = ws.cell(row=TOT, column=3 + k, value=f"=SUM({col}{FIRST}:{col}{LAST})")
            c.number_format = MONEY0
        for col in range(1, 3 + len(CATS) + 1):
            c = ws.cell(row=TOT, column=col)
            c.font = TOTAL_FONT
            c.border = TOTAL_BORDER
        ws.column_dimensions["A"].width = 27
        for k in range(len(CATS)):
            ws.column_dimensions[get_column_letter(3 + k)].width = 16
        ws.column_dimensions[TCOL].width = 13
        ws.freeze_panes = f"B{FIRST}"

    @L.customize
    def _sheets(wb, lesson, selftest):
        # ---------------------------------------------------------- Expenses + monthly sheets
        for name in ["Expenses"] + list(MONTHS.values()):
            ws = wb[name]
            style_grid_sheet(ws)
            ws.freeze_panes = f"B{FIRST}"
        ws = wb["Expenses"]
        k = ws[f"{SCOL}{TOT}"]
        k.fill = PatternFill(fill_type=None)
        k.border = Border()
        for name in MONTHS.values():
            wb[name].sheet_properties.tabColor = "2E75B6"

        # ---------------------------------------------------------- 4 West OT: the multiplier cell
        ws = wb["4 West OT"]
        ws["A3"] = "Overtime multiplier"
        ws["A3"].font = Font(bold=True)
        ws["A3"].alignment = Alignment(horizontal="right")
        ws[MULT_CELL] = OT_MULT
        assumption(ws[MULT_CELL], "0.0")
        ws["C3"] = "← time-and-a-half. Every OTPay formula should point here."
        ws["C3"].font = NOTE_FONT

        # ---------------------------------------------------------- Staffing Grid
        ws = wb.create_sheet("Staffing Grid")
        ws.sheet_properties.tabColor = "548235"
        titled(ws, "Medical-Surgical 4 West · nursing care hours needed per day",
               "Hours needed = midnight census (column A) × HPPD target (row 5). HPPD = nursing hours per patient day.",
               f"Fill the yellow grid B{G_FIRST}:{G_C1}{G_LAST} with ONE formula: type it in B{G_FIRST}, copy it across "
               f"to {G_C1}{G_FIRST}, then down to row {G_LAST}.")
        c = ws.cell(row=G_HDR, column=1, value="Census ↓   HPPD →")
        c.font = HDR_FONT
        c.fill = HEADER_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center")
        for j, h in enumerate(HPPD, 2):
            c = ws.cell(row=G_HDR, column=j, value=h)
            c.font = HDR_FONT
            c.fill = HEADER_FILL
            c.border = BOX
            c.number_format = "0.0"
            c.alignment = Alignment(horizontal="center")
        for i, cv in enumerate(census_vals):
            c = ws.cell(row=G_FIRST + i, column=1, value=cv)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9E1F2")
            c.border = BOX
            c.alignment = Alignment(horizontal="center")
        yellow(ws, f"{G_C0}{G_FIRST}:{G_C1}{G_LAST}", MONEY0)
        ws.column_dimensions["A"].width = 18
        for j in range(len(HPPD)):
            ws.column_dimensions[get_column_letter(2 + j)].width = 9
        ws.freeze_panes = f"B{G_FIRST}"

        # ---------------------------------------------------------- Q4 Summary
        ws = wb.create_sheet("Q4 Summary")
        ws.sheet_properties.tabColor = "548235"
        titled(ws, "Bluestone Memorial Hospital · Q4 2025 operating expense, October–December ($)",
               "Same layout as the Oct, Nov, and Dec sheets. Fill the yellow grid with ONE 3-D formula that adds the same cell "
               "on all three monthly sheets, then copy it across and down.")
        header_row(ws, HDR, ["Department", "Cost Center"] + CATS + ["Total"])
        ws.row_dimensions[HDR].height = 32
        for d in f01:
            ws.cell(row=row_of[d], column=2, value=deps[d]["CostCenter"]).border = BOX
        yellow(ws, f"{C0}{FIRST}:{C1}{LAST}", MONEY0)
        dept_block(ws)
        ws.column_dimensions["B"].width = 8

        # ---------------------------------------------------------- Plan 2026 (bonus)
        ws = wb.create_sheet("Plan 2026")
        ws.sheet_properties.tabColor = "BF9000"
        titled(ws, "Bluestone Memorial Hospital · 2026 operating expense plan ($)",
               f"2026 = 2025 actual from the Expenses sheet (same cell) × (1 + price inflation in row {INF_ROW}) × "
               f"(1 + volume growth in column {GROW_COL}). Blue cells are planning assumptions.")
        c = ws.cell(row=INF_ROW, column=1, value="Price inflation →")
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="right")
        for cat in CATS:
            cc = ws[f"{cat_col[cat]}{INF_ROW}"]
            cc.value = INFLATION[cat]
            assumption(cc, "0.0%")
            cc.alignment = Alignment(horizontal="center")
        header_row(ws, HDR, ["Department", "Volume growth"] + CATS + ["Total 2026"])
        ws.row_dimensions[HDR].height = 32
        for d in f01:
            cc = ws[f"{GROW_COL}{row_of[d]}"]
            cc.value = GROWTH[d]
            assumption(cc, "0.0%")
            cc.alignment = Alignment(horizontal="center")
        yellow(ws, f"{C0}{FIRST}:{C1}{LAST}", MONEY0)
        dept_block(ws, "Hospital total 2026")
        ws.column_dimensions["B"].width = 10

        for name in ["Expenses", "4 West OT", "Staffing Grid", "Q4 Summary", "Plan 2026"] + list(MONTHS.values()):
            print_setup(wb[name])

        # Excel doesn't allow 3-D references in array formulas, and the library writes every key formula as a
        # dynamic-array formula. Rewrite the key's live 3-D formulas as ordinary formulas (they return one value).
        for key_name, tasks in ((lesson.key_sheet, lesson.tasks), (lesson.bonus_key_sheet, lesson.bonus)):
            ks = wb[key_name]
            for t in tasks:
                live = t.live if isinstance(t.live, str) else (t.solution if (t.live and t.is_formula) else None)
                if t.live_cell and live and re.search(r"[A-Za-z0-9_']+:[A-Za-z0-9_' ]+!", live):
                    ks[t.live_cell] = to_file_formula(live)
                    lesson._dynamic.discard((ks.title, t.live_cell))

    return L
