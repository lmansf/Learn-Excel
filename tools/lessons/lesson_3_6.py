"""Lesson 3.6 · What-If Analysis: Goal Seek, Scenarios, Data Tables & Solver.

The lesson revolves around one model: a monthly operating (P&L) model for Bluestone's Primary Care Clinic
(D400, Bluestone Outpatient Pavilion). Its data-derived inputs come from the course datasets:

  * payer mix .......... share of 2025 D400 claims by payer group (claims.csv joined to encounters.csv)
  * $ per visit ........ average AllowedAmount per 2025 D400 claim in each payer group (denied and appealed claims,
                         whose AllowedAmount is 0, count as $0)
  * hourly rates ....... median HourlyRate of ACTIVE employees by job title (employees.csv)
  * benefits load ...... 2025 D400 Employee Benefits / Salaries & Wages actuals (budget.csv)
  * supplies, vaccines . 2025 D400 Medical Supplies / Pharmaceuticals actuals per month / planned visits per month
  * fixed costs ........ 2025 D400 Equipment & Maintenance, Purchased Services, Other Operating, monthly average

Planning assumptions (clinic days, visits per day, FTEs, salaries, billing fee, staffing hours per visit) are stated
inputs, not answers.

Sheets
  Model     (customize) inputs → calculations → outputs, with two yellow output cells the learner completes
            (Operating income, Operating margin), one deliberately hard-coded calculation (task 3), and a
            "Sensitivity area" in F:K where the learner builds a one-variable and a two-variable Data Table.
            Data Table input cells must be on the same sheet as the table, which is why it lives on Model.
  Staffing  (customize) a Solver worksheet: RN/LPN/CNA FTEs (decision cells), linear constraints, cost objective.
  Sources   where every data-derived input came from (an Excel Table).

Self-test simulation (selftest=True only): the hook writes the learner's two Data Tables as real Excel data-table
formulas (<f t="dataTable">, which LibreOffice evaluates), so the self-test checks the workbook's own model against the
Python answers. The Solver answers are checked in the pristine key instead: task 12's live formula rebuilds the LP
optimum in closed form from the Staffing sheet's cells, and task 13's from the Model. LibreOffice can't run a data table
whose input cell sits inside a range argument (SUMPRODUCT(B14:B18,C14:C18) with input C17), so the self-test copy
replaces that one formula with the identical explicit sum B14*C14+…+B18*C18. The learner's workbook keeps SUMPRODUCT.

Python mirrors the workbook model exactly (``model()`` below). The model is linear in visits per day, so the
Goal Seek break-even has a unique answer; the Solver answers come from an exact vertex-enumeration LP solver
(fractions) and an integer brute force.
"""
from __future__ import annotations

import re
import statistics
from collections import defaultdict
from decimal import ROUND_HALF_UP, Decimal
from fractions import Fraction
from itertools import combinations, product

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import DataTableFormula

from xlcourse import Lesson, Task, data
from xlcourse.lesson import BOX, HEADER_FILL, INPUT_BORDER, INPUT_FILL, NAVY

CODE = "3.6"

# ---------------------------------------------------------------------------- planning assumptions (stated inputs)
CLINIC_DAYS = 21
VISITS_PER_DAY = 200
MD_FTE, APP_FTE = 7, 3
MD_VPD, APP_VPD = 22, 18                  # visits per provider per day (scheduling template)
RN_FTE, LPN_FTE, CNA_FTE, FD_FTE = 5, 3, 8, 8
HOURS_PER_FTE = 168                       # 8-hour day x 21 clinic days
MD_SALARY, APP_SALARY = 22000, 10800      # $ per FTE per month (wages)
BILLING_FEE = 0.04
SUPPORT_HRS_PER_VISIT = 0.60              # Staffing sheet: RN + LPN + CNA hours per visit
LICENSED_HRS_PER_VISIT = 0.28             # RN + LPN hours per visit
RN_MIN_SHARE = 0.60
CNA_CAP = 6                               # bonus only
GROWTH_VPD, GROWTH_APP = 220, 4           # bonus growth plan

GROUPS = ["Medicare", "Medicare Advantage", "Medicaid", "Commercial", "Self-Pay"]
GROUP_OF = {"PY01": "Medicare", "PY02": "Medicare Advantage", "PY03": "Medicaid", "PY04": "Commercial",
            "PY05": "Commercial", "PY06": "Commercial", "PY08": "Commercial", "PY07": "Self-Pay"}
GROUP_LABEL = {"Medicare": "Medicare", "Medicare Advantage": "Medicare Advantage", "Medicaid": "Medicaid",
               "Commercial": "Commercial (incl. Workers' Comp)", "Self-Pay": "Self-Pay"}

# Scenario Manager scenarios (task 10/11): changing cells VisitsPerDay, MedicaidShare, CommercialShare, BillingFeePct
SCENARIOS = {
    "Base plan": dict(vpd=200, medicaid=0.149, commercial=0.376, fee=0.040),
    "Downside": dict(vpd=190, medicaid=0.179, commercial=0.346, fee=0.050),
    "Upside": dict(vpd=206, medicaid=0.129, commercial=0.396, fee=0.035),
}

# Sensitivity grids (Model!F:K)
DT1_VISITS = list(range(180, 221, 5))     # F9:F17
DT2_VISITS = list(range(180, 221, 5))     # F23:F31
DT2_RATES = [130, 140, 150, 160, 170]     # G22:K22

MONEY0 = '#,##0;[Red]-#,##0'
MONEY2 = '#,##0.00'
PCT1 = '0.0%'

TITLE_FONT = Font(bold=True, size=15, color=NAVY)
NOTE_FONT = Font(italic=True, color="595959")
SECTION_FONT = Font(bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FONT = Font(color="0000FF")
INPUT_CELL_FILL = PatternFill("solid", fgColor="DDEBF7")
LINK_FONT = Font(color="008000")
DECISION_FILL = PatternFill("solid", fgColor="E2EFDA")
RESULT_FILL = PatternFill("solid", fgColor="F2F2F2")
TOTAL_BORDER = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))


def r2(x: float) -> float:
    """Round half up to cents (avoids binary-float surprises like 19.735 -> 19.73)."""
    return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="03-data-analysis", slug="06-what-if-analysis",
        title="What-If Analysis: Goal Seek, Scenarios, Data Tables & Solver", level="Intermediate", minutes=60,
        objectives=[
            "Structure a model with separate inputs, calculations, and outputs",
            "Find break-even points with Goal Seek",
            "Compare cases with Scenario Manager and sensitivity Data Tables",
            "Optimize a staffing mix with Solver",
        ],
        data_note="A monthly operating model for Bluestone's Primary Care Clinic (D400, Bluestone Outpatient Pavilion). "
                  "Payer mix and reimbursement come from the clinic's 2,079 claims with 2025 service dates, wage rates from "
                  "Bluestone HR records, and benefits, supply, and fixed costs from the clinic's 2025 budget actuals.",
    )

    # ======================================================================== inputs derived from the data
    enc = {e["EncounterID"]: e for e in data.load("encounters")}
    claims = [c for c in data.load("claims")
              if enc[c["EncounterID"]]["DeptID"] == "D400" and c["ServiceDate"].year == 2025]
    n_claims = len(claims)
    # the Sources text says denied and appealed claims carry AllowedAmount 0 (and only they do)
    assert all((c["AllowedAmount"] == 0) == (c["ClaimStatus"] in ("Denied", "Appealed")) for c in claims)
    g_count, g_allowed = defaultdict(int), defaultdict(float)
    for c in claims:
        g = GROUP_OF[c["PayerID"]]
        g_count[g] += 1
        g_allowed[g] += c["AllowedAmount"]
    exact_share = {g: g_count[g] / n_claims for g in GROUPS}
    avg_allowed = {g: g_allowed[g] / g_count[g] for g in GROUPS}
    rate = {g: int(Decimal(str(avg_allowed[g])).quantize(Decimal("1"), rounding=ROUND_HALF_UP)) for g in GROUPS}
    # shares rounded to 0.1 point with the largest-remainder method so the mix totals exactly 100.0%
    tenths = {g: exact_share[g] * 1000 for g in GROUPS}
    base_t = {g: int(tenths[g]) for g in GROUPS}
    for g in sorted(GROUPS, key=lambda g: tenths[g] - base_t[g], reverse=True)[:1000 - sum(base_t.values())]:
        base_t[g] += 1
    share = {g: base_t[g] / 1000 for g in GROUPS}
    assert sum(base_t.values()) == 1000
    assert abs(share["Medicaid"] - SCENARIOS["Base plan"]["medicaid"]) < 1e-12
    assert abs(share["Commercial"] - SCENARIOS["Base plan"]["commercial"]) < 1e-12

    emp = data.load("employees")

    def median_rate(title):
        v = [e["HourlyRate"] for e in emp if e["JobTitle"] == title and e["Status"] == "Active"]
        return r2(statistics.median(v)), len(v)
    rn_rate, n_rn = median_rate("Registered Nurse")
    lpn_rate, n_lpn = median_rate("Licensed Practical Nurse")
    cna_rate, n_cna = median_rate("Certified Nursing Assistant")
    fd_rate, n_fd = median_rate("Unit Secretary")

    bud = defaultdict(float)
    for r in data.load("budget"):
        if r["DeptID"] == "D400" and r["FiscalYear"] == 2025:
            bud[r["Category"]] += r["ActualAmount"]
    benefits_load = round(bud["Employee Benefits"] / bud["Salaries & Wages"], 3)
    planned_vpm = VISITS_PER_DAY * CLINIC_DAYS
    supplies_pv = r2(bud["Medical Supplies"] / 12 / planned_vpm)
    vaccines_pv = r2(bud["Pharmaceuticals"] / 12 / planned_vpm)
    fixed_equip = round(bud["Equipment & Maintenance"] / 12, -2)
    fixed_purch = round(bud["Purchased Services"] / 12, -2)
    fixed_other = round(bud["Other Operating"] / 12, -2)

    BASE = dict(days=CLINIC_DAYS, vpd=VISITS_PER_DAY, md=MD_FTE, app=APP_FTE,
                shares=[share[g] for g in GROUPS], rates=[rate[g] for g in GROUPS],
                rn=RN_FTE, lpn=LPN_FTE, cna=CNA_FTE, fd=FD_FTE, hrs=HOURS_PER_FTE,
                md_sal=MD_SALARY, app_sal=APP_SALARY, rn_rate=rn_rate, lpn_rate=lpn_rate, cna_rate=cna_rate,
                fd_rate=fd_rate, ben=benefits_load, sup=supplies_pv, vac=vaccines_pv, fee=BILLING_FEE,
                fx=fixed_equip + fixed_purch + fixed_other)

    def model(p: dict) -> dict:
        """Mirror of the Model sheet (same order of operations)."""
        vpm = p["vpd"] * p["days"]
        avg = sum(s * r for s, r in zip(p["shares"], p["rates"]))
        rev = vpm * avg
        wages = (p["md"] * p["md_sal"] + p["app"] * p["app_sal"]
                 + (p["rn"] * p["rn_rate"] + p["lpn"] * p["lpn_rate"] + p["cna"] * p["cna_rate"]) * p["hrs"]
                 + p["fd"] * p["fd_rate"] * p["hrs"])
        ben = wages * p["ben"]
        var = vpm * p["sup"] + vpm * p["vac"]
        fee = rev * p["fee"]
        exp = wages + ben + var + fee + p["fx"]
        oi = rev - exp
        return dict(vpm=vpm, avg=avg, rev=rev, wages=wages, ben=ben, var=var, fee=fee, exp=exp, oi=oi,
                    margin=oi / rev, fixed=wages + ben + p["fx"],
                    cm=avg * (1 - p["fee"]) - p["sup"] - p["vac"])

    def with_(**kw) -> dict:
        p = dict(BASE)
        p.update(kw)
        return p

    def with_scenario(s: dict) -> dict:
        shares = list(BASE["shares"])
        shares[GROUPS.index("Medicaid")] = s["medicaid"]
        shares[GROUPS.index("Commercial")] = s["commercial"]
        assert abs(sum(shares) - 1) < 1e-9
        return with_(vpd=s["vpd"], shares=shares, fee=s["fee"])

    def with_commercial_rate(x: float) -> dict:
        rates = list(BASE["rates"])
        rates[GROUPS.index("Commercial")] = x
        return with_(rates=rates)

    base = model(BASE)

    # ---------------------------------------------------------------- Goal Seek answers (exact)
    breakeven_vpd = base["fixed"] / (base["cm"] * CLINIC_DAYS)          # OI is linear in visits/day
    assert abs(model(with_(vpd=breakeven_vpd))["oi"]) < 1e-6
    s_c = share["Commercial"]
    non_fee_exp = base["exp"] - base["fee"]
    target_rev = non_fee_exp / (1 - BILLING_FEE - 0.05)                 # margin 5% <=> rev*(1-fee-0.05) = non-fee costs
    comm_rate_5 = rate["Commercial"] + (target_rev / base["vpm"] - base["avg"]) / s_c
    assert abs(model(with_commercial_rate(comm_rate_5))["margin"] - 0.05) < 1e-12

    # ---------------------------------------------------------------- Data Table answers
    dt1 = {v: model(with_(vpd=v)) for v in DT1_VISITS}
    dt2 = {(v, x): model(dict(with_commercial_rate(x), vpd=v))["oi"] for v in DT2_VISITS for x in DT2_RATES}
    margin_190 = dt1[190]["margin"]
    first_profit = min(v for v in DT1_VISITS if dt1[v]["oi"] > 0)
    oi_210_160 = dt2[(210, 160)]
    n_profitable = sum(1 for v in dt2.values() if v > 0)
    dt2_first = {x: min(v for v in DT2_VISITS if dt2[(v, x)] > 0) for x in DT2_RATES}
    assert min(abs(v) for v in dt2.values()) > 100 and min(abs(d["oi"]) for d in dt1.values()) > 100
    # task 8 explanation: with swapped input cells, the (200 visits, $140) cell would show 140 visits at $200
    assert model(dict(with_commercial_rate(200), vpd=140))["oi"] < dt2[(VISITS_PER_DAY, 140)] - 20000

    # ---------------------------------------------------------------- Scenario answers
    scen = {k: model(with_scenario(s)) for k, s in SCENARIOS.items()}
    assert abs(scen["Base plan"]["oi"] - base["oi"]) < 1e-6
    # what a learner gets in the Downside if the hard-coded 4% fee is NOT fixed
    down_unfixed_oi = scen["Downside"]["oi"] + scen["Downside"]["rev"] * (SCENARIOS["Downside"]["fee"] - BILLING_FEE)

    # ---------------------------------------------------------------- Solver (Staffing sheet)
    def F(x) -> Fraction:
        return Fraction(str(x))
    loaded = {k: F(v) * HOURS_PER_FTE * (1 + F(benefits_load)) for k, v in
              (("RN", rn_rate), ("LPN", lpn_rate), ("CNA", cna_rate))}
    cost_vec = [loaded["RN"], loaded["LPN"], loaded["CNA"]]

    def staffing_constraints(vpd: int, cap: int | None):
        vpm = vpd * CLINIC_DAYS
        h = HOURS_PER_FTE
        cons = [((h, h, h), vpm * F(SUPPORT_HRS_PER_VISIT)),          # total support hours
                ((h, h, 0), vpm * F(LICENSED_HRS_PER_VISIT)),         # licensed hours
                ((1 - F(RN_MIN_SHARE), -F(RN_MIN_SHARE), 0), 0),       # RN - 0.6*(RN+LPN) >= 0
                ((1, 0, 0), 0), ((0, 1, 0), 0), ((0, 0, 1), 0)]       # non-negativity
        if cap is not None:
            cons.append(((0, 0, -1), -cap))                           # CNA <= cap
        return cons

    def solve3(rows):
        (a, b, c), (d, e, f_), (g, h, i) = [r[0] for r in rows]
        det = a * (e * i - f_ * h) - b * (d * i - f_ * g) + c * (d * h - e * g)
        if det == 0:
            return None
        y1, y2, y3 = [F(r[1]) for r in rows]
        x = (y1 * (e * i - f_ * h) - b * (y2 * i - f_ * y3) + c * (y2 * h - e * y3)) / det
        y = (a * (y2 * i - f_ * y3) - y1 * (d * i - f_ * g) + c * (d * y3 - y2 * g)) / det
        z = (a * (e * y3 - y2 * h) - b * (d * y3 - y2 * g) + y1 * (d * h - e * g)) / det
        return (x, y, z)

    def lp_min(cons):
        """Exact LP minimum by enumerating vertices (3 variables, every constraint written as a.x >= b)."""
        best = None
        for trio in combinations(cons, 3):
            x = solve3(trio)
            if x is None or any(sum(F(ai) * xi for ai, xi in zip(a, x)) < F(b) for a, b in cons):
                continue
            cost = sum(ci * xi for ci, xi in zip(cost_vec, x))
            if best is None or cost < best[0]:
                best = (cost, x)
        return best

    def int_min(cons, top=30):
        sols = []
        for x in product(range(top + 1), repeat=3):
            if all(sum(F(ai) * xi for ai, xi in zip(a, x)) >= F(b) for a, b in cons):
                sols.append((sum(ci * xi for ci, xi in zip(cost_vec, x)), x))
        sols.sort()
        assert sols[0][0] < sols[1][0], "integer optimum must be unique"
        return sols

    lp_cost, lp_x = lp_min(staffing_constraints(VISITS_PER_DAY, None))
    # the task 12 live formula assumes this vertex: licensed FTEs = licensed hours / 168, split at the RN minimum,
    # and CNAs cover the rest of the support hours
    _lic = VISITS_PER_DAY * CLINIC_DAYS * F(LICENSED_HRS_PER_VISIT) / HOURS_PER_FTE
    assert lp_x == (F(RN_MIN_SHARE) * _lic, (1 - F(RN_MIN_SHARE)) * _lic,
                    VISITS_PER_DAY * CLINIC_DAYS * (F(SUPPORT_HRS_PER_VISIT) - F(LICENSED_HRS_PER_VISIT)) / HOURS_PER_FTE)
    # shadow price of the total-hours constraint: re-solve with one more required hour
    cons_plus = staffing_constraints(VISITS_PER_DAY, None)
    cons_plus[0] = (cons_plus[0][0], cons_plus[0][1] + 1)
    shadow_total = lp_min(cons_plus)[0] - lp_cost
    assert shadow_total == F(cna_rate) * (1 + F(benefits_load))
    cons_plus = staffing_constraints(VISITS_PER_DAY, None)
    cons_plus[1] = (cons_plus[1][0], cons_plus[1][1] + 1)
    shadow_licensed = lp_min(cons_plus)[0] - lp_cost
    assert shadow_licensed < shadow_total  # the task 13 explanation says "only"
    current_cost = sum(c * x for c, x in zip(cost_vec, (RN_FTE, LPN_FTE, CNA_FTE)))

    growth_cons = staffing_constraints(GROWTH_VPD, CNA_CAP)
    g_lp_cost, g_lp_x = lp_min(growth_cons)
    g_int = int_min(growth_cons)
    g_int_cost, (g_rn, g_lpn, g_cna) = g_int[0]
    naive = tuple(-(-xi.numerator // xi.denominator) for xi in g_lp_x)        # round every LP value up
    naive_cost = sum(c * x for c, x in zip(cost_vec, naive))
    naive_ok = all(sum(F(ai) * xi for ai, xi in zip(a, naive)) >= F(b) for a, b in growth_cons)
    naive_share = naive[0] / (naive[0] + naive[1])
    assert naive_cost > g_int_cost and naive[1] > g_lpn
    runner_up_gap = float((g_int[1][0] - g_int_cost) / g_int_cost)
    assert runner_up_gap > 0.01  # Solver's default 1% Integer Optimality could not stop at a worse plan
    growth = model(with_(vpd=GROWTH_VPD, app=GROWTH_APP, rn=g_rn, lpn=g_lpn, cna=g_cna))

    # ======================================================================== workbook layout (row numbers)
    R = dict(days=7, vpd=8, md=9, app=10, md_vpd=11, app_vpd=12, payer_hdr=13, payer0=14, mix_total=19,
             rn=21, lpn=22, cna=23, fd=24, hrs=25, md_sal=26, app_sal=27, rn_rate=28, lpn_rate=29, cna_rate=30,
             fd_rate=31, ben=32, sup=34, vac=35, fee=36, fx1=37, fx2=38, fx3=39,
             vpm=42, avg=43, rev=44, w_prov=45, w_sup=46, w_fd=47, wages=48, benefits=49, c_sup=50, c_vac=51,
             c_fee=52, c_fx1=53, c_fx2=54, c_fx3=55, exp=56, oi=59, margin=60, cpv=61, cap=62, util=63)
    payer_row = {g: R["payer0"] + i for i, g in enumerate(GROUPS)}
    MED_ROW, COM_ROW = payer_row["Medicaid"], payer_row["Commercial"]
    OI, MARGIN = f"B{R['oi']}", f"B{R['margin']}"
    # Sensitivity area
    DT1_HDR, DT1_FORM, DT1_FIRST = 7, 8, 9
    DT1_LAST = DT1_FIRST + len(DT1_VISITS) - 1
    DT2_CORNER, DT2_FIRST = 22, 23
    DT2_LAST = DT2_FIRST + len(DT2_VISITS) - 1
    DT2_COLS = ["G", "H", "I", "J", "K"]
    dt1_row = {v: DT1_FIRST + i for i, v in enumerate(DT1_VISITS)}
    dt2_row = {v: DT2_FIRST + i for i, v in enumerate(DT2_VISITS)}
    dt2_col = dict(zip(DT2_RATES, DT2_COLS))
    # Staffing sheet
    S = dict(vpd=6, days=7, vpm=8, sup_hpv=9, lic_hpv=10, rn_share=11, cap=12, hrs=13, ben=14,
             rn=18, lpn=19, cna=20, total=21, c_total=25, c_lic=26, c_rn=27, c_cap=28, objective=31)

    def m(key: str, col: str = "B") -> str:
        return f"Model!{col}{R[key]}"

    def money(x: float) -> str:
        return f"${x:,.0f}" if x >= 0 else f"−${-x:,.0f}"

    def money2(x: float) -> str:
        return f"${x:,.2f}"

    # ======================================================================== Sources sheet (data provenance)
    src = []

    def add_src(inp, cell, value, how, dataset):
        src.append({"Input": inp, "ModelCell": cell, "ValueUsed": value, "HowDerived": how, "SourceData": dataset})
    for g in GROUPS:
        add_src(f"{GROUP_LABEL[g]}: share of visits", f"B{payer_row[g]}", f"{share[g]:.1%}",
                f"{g_count[g]:,} of {n_claims:,} D400 claims with 2025 service dates ({exact_share[g]:.2%}); rounded to 0.1 "
                "point so the five shares total exactly 100%", "claims.csv + encounters.csv (DeptID D400)")
    for g in GROUPS:
        add_src(f"{GROUP_LABEL[g]}: $ per visit", f"C{payer_row[g]}", f"${rate[g]}",
                f"Average AllowedAmount per 2025 D400 claim = ${avg_allowed[g]:,.2f} (denied and appealed claims carry an "
                "AllowedAmount of $0 and count as $0), rounded to the dollar", "claims.csv")
    for label, row, val, n_, title in (("RN hourly rate", R["rn_rate"], rn_rate, n_rn, "Registered Nurse"),
                                       ("LPN hourly rate", R["lpn_rate"], lpn_rate, n_lpn, "Licensed Practical Nurse"),
                                       ("CNA hourly rate", R["cna_rate"], cna_rate, n_cna, "Certified Nursing Assistant"),
                                       ("Front desk hourly rate", R["fd_rate"], fd_rate, n_fd, "Unit Secretary")):
        add_src(label, f"B{row}", f"${val:.2f}", f"Median HourlyRate of the {n_:,} active employees with JobTitle = {title}",
                "employees.csv")
    add_src("Benefits load", f"B{R['ben']}", f"{benefits_load:.1%}",
            f"2025 D400 Employee Benefits ÷ Salaries & Wages actuals = ${bud['Employee Benefits']:,.0f} ÷ "
            f"${bud['Salaries & Wages']:,.0f} = {bud['Employee Benefits'] / bud['Salaries & Wages']:.2%}", "budget.csv")
    add_src("Medical supplies per visit", f"B{R['sup']}", f"${supplies_pv:.2f}",
            f"2025 D400 Medical Supplies actual, monthly average ${bud['Medical Supplies'] / 12:,.0f} ÷ {planned_vpm:,} "
            "planned visits per month", "budget.csv")
    add_src("Vaccines & injectables per visit", f"B{R['vac']}", f"${vaccines_pv:.2f}",
            f"2025 D400 Pharmaceuticals actual, monthly average ${bud['Pharmaceuticals'] / 12:,.0f} ÷ {planned_vpm:,} "
            "planned visits per month", "budget.csv")
    for label, row, val, cat in (("Facility & equipment", R["fx1"], fixed_equip, "Equipment & Maintenance"),
                                 ("IT, EHR & purchased services", R["fx2"], fixed_purch, "Purchased Services"),
                                 ("Other operating", R["fx3"], fixed_other, "Other Operating")):
        add_src(label, f"B{row}", f"${val:,.0f}", f"2025 D400 {cat} actual, monthly average ${bud[cat] / 12:,.0f}, "
                "rounded to the nearest $100", "budget.csv")
    add_src("Everything else on the Model", "—", "—",
            "Planning assumptions for FY2026: clinic days, visits per day, provider and staff FTEs, provider salaries, "
            "visits per provider per day, paid hours per FTE, and the billing fee. Staffing hours per visit are on the "
            "Staffing sheet.", "Plan (not from data)")
    L.add_table_sheet("Sources", src, table="tblSources",
                      columns=[("Input", "Input"), ("ModelCell", "Model cell"), ("ValueUsed", "Value used"),
                               ("HowDerived", "How it was derived"), ("SourceData", "Source data")],
                      widths={"Input": 34, "Model cell": 11, "Value used": 12, "How it was derived": 92,
                              "Source data": 36})

    # ======================================================================== tasks
    L.practice_intro = (
        "Every task uses the Model and Staffing sheets. Tasks 1 and 2 read your Model live, so their checks stay green "
        f"only while the Model holds its base-case inputs ({VISITS_PER_DAY} visits per day, Commercial ${rate['Commercial']}, and so on). After each "
        "Goal Seek run, click Cancel in the Goal Seek Status box so the Model keeps those inputs. The Staffing sheet doesn't "
        "feed the Model, so you can keep Solver's solutions there.")

    t_oi = Task(
        f"On the Model sheet, complete the yellow Operating income cell ({OI}): net patient revenue minus total operating "
        "expenses. The gray cell here reads your formula. What is the clinic's base-case operating income per month?",
        answer=round(base["oi"], 2), fmt="#,##0.00;[Red]-#,##0.00", title="Operating income (Model!B59)",
        solution=f"=B{R['rev']}-B{R['exp']}",
        summary=f'=IF(Model!{OI}="","",Model!{OI})',
        fill={"range": f"Model!{OI}:{OI}", "formula": f"=B{R['rev']}-B{R['exp']}"},
        live=f"={m('rev')}-{m('exp')}",
        hint="An output should only point at calculation cells",
        explanation=f"Operating income is revenue minus expenses, and both already exist as calculation rows "
                    f"(B{R['rev']} and B{R['exp']}). Because those cells are named, `=NetRevenue-TotalExpenses` works too and "
                    f"reads better. The clinic loses about {money(-base['oi'])} a month in the base case. That's common for "
                    "hospital-owned primary care, which is why the rest of the lesson asks what it would take to break even.")
    t_margin = Task(
        f"Complete the yellow Operating margin cell ({MARGIN}): operating income as a share of net patient revenue. "
        "What is the base-case operating margin? (Format it as a percentage.)",
        answer=base["margin"], fmt=PCT1, title="Operating margin (Model!B60)",
        solution=f"=B{R['oi']}/B{R['rev']}",
        summary=f'=IF(Model!{MARGIN}="","",Model!{MARGIN})',
        fill={"range": f"Model!{MARGIN}:{MARGIN}", "formula": f"=B{R['oi']}/B{R['rev']}"},
        live=f"=({m('rev')}-{m('exp')})/{m('rev')}",
        hint="Margin = income ÷ revenue",
        explanation="Margin divides operating income by revenue (`=OperatingIncome/NetRevenue`). A margin lets you compare a "
                    "small clinic with a large hospital, which a dollar figure can't. Goal Seek will aim at this cell in task 5.")
    t_hard = Task(
        f"Audit the model. One formula in the CALCULATIONS section (rows {R['vpm']}–{R['exp']}) has a number typed into it "
        "instead of a reference to its input cell. Type that cell's address (for example B99). Then fix the formula so it "
        "points to the input. Task 10 depends on the fix.",
        answer=f"B{R['c_fee']}", accept=[f"$B${R['c_fee']}", f"Model!B{R['c_fee']}", f"Model!$B${R['c_fee']}"],
        title="Find the hard-coded number",
        solution=f"1. Press **Ctrl + `** (Mac: **⌃ + `**) to show formulas, and read down rows {R['vpm']}–{R['exp']}.\n"
                 f"2. B{R['c_fee']} reads `=NetRevenue*0.04`. The 4% is typed in, although the fee has its own input cell, "
                 f"B{R['fee']}.\n"
                 f"3. Change B{R['c_fee']} to `=NetRevenue*BillingFeePct` (or `=B{R['rev']}*B{R['fee']}`), then press Ctrl + ` "
                 "again to show values.",
        live=False, hint="Show Formulas, or Trace Dependents on each input",
        explanation=f"Today both versions return the same number, so nothing looks wrong. The trouble starts when someone changes "
                    f"the fee input in B{R['fee']}: the model ignores it. Another way to catch this is **Formulas → Trace "
                    f"Dependents** on B{R['fee']}. Excel draws no arrows and tells you no formula refers to the active cell, "
                    "because no formula reads that input.")
    t_be = Task(
        "Use Goal Seek to find the break-even volume: set Operating income to 0 by changing Visits per day (Model!B8). "
        f"How many visits per day does the clinic need? Round to 1 decimal place, then click Cancel to restore {VISITS_PER_DAY}.",
        answer=round(breakeven_vpd, 1), fmt="0.0", tol=0.051, title="Goal Seek: break-even visits per day",
        solution=f"1. Choose **Data → What-If Analysis → Goal Seek**.\n"
                 f"2. **Set cell:** `{OI}` · **To value:** `0` · **By changing cell:** `B{R['vpd']}`.\n"
                 f"3. Click **OK**. B{R['vpd']} shows {breakeven_vpd:.1f} (the cell holds {breakeven_vpd:.4f}…). Note it, "
                 "then click **Cancel**.",
        # The fee rate is read as B52/B44 (fee ÷ revenue), never from the fee input B36: task 3 relies on B36 having
        # no dependents at all, and Excel's Trace Dependents would draw an arrow to this hidden key otherwise.
        live=f"=ROUND({m('vpd')}-({m('rev')}-{m('exp')})/(({m('avg')}*(1-{m('c_fee')}/{m('rev')})-{m('sup')}-{m('vac')})"
             f"*{m('days')}),1)",
        hint="Data → What-If Analysis → Goal Seek",
        explanation=f"Each extra visit per day adds one visit on each of the {CLINIC_DAYS} clinic days. Each of those visits brings "
                    f"in {money2(base['cm'])} after the billing fee, supplies, and vaccines (its **contribution margin**). "
                    f"Fixed costs are {money(base['fixed'])} a month, so break-even = {money(base['fixed'])} ÷ "
                    f"({money2(base['cm'])} × {CLINIC_DAYS}) ≈ {breakeven_vpd:.1f}. That is "
                    f"{breakeven_vpd / (MD_FTE * MD_VPD + APP_FTE * APP_VPD):.0%} of the clinic's "
                    f"{MD_FTE * MD_VPD + APP_FTE * APP_VPD}-visit daily capacity, so volume alone is a fragile fix. The live "
                    "result in the key does this algebra, and Goal Seek reaches the same answer by trial and error.")
    t_rate = Task(
        "The CFO is renegotiating commercial contracts. Use Goal Seek to set Operating margin to 5% (type 0.05) by changing "
        f"the Commercial $ per visit (Model!C{COM_ROW}). What commercial reimbursement per visit is needed? Round to the "
        f"nearest dollar, then click Cancel to restore ${rate['Commercial']}.",
        answer=round(comm_rate_5, 2), fmt=MONEY2, tol=0.8, title="Goal Seek: commercial rate for a 5% margin",
        answer_display=f"about ${comm_rate_5:,.0f} (Goal Seek shows {comm_rate_5:,.2f})",
        solution=f"1. **Data → What-If Analysis → Goal Seek**.\n"
                 f"2. **Set cell:** `{MARGIN}` · **To value:** `0.05` · **By changing cell:** `C{COM_ROW}`.\n"
                 f"3. Click **OK**, read C{COM_ROW}, and click **Cancel**.",
        live=(f"=Model!C{COM_ROW}+(({m('exp')}-{m('c_fee')})/(1-{m('c_fee')}/{m('rev')}-0.05)/{m('vpm')}-{m('avg')})"
              f"/Model!B{COM_ROW}"),
        hint="The Set cell must contain a formula, so use the margin cell",
        explanation=f"Commercial plans would have to pay about {money(comm_rate_5)} instead of ${rate['Commercial']}, an increase of "
                    f"{comm_rate_5 / rate['Commercial'] - 1:.0%}. Only {s_c:.1%} of visits are commercial, so each extra "
                    f"commercial dollar moves the average reimbursement by just {s_c * 100:.1f} cents. Margin is a ratio, so Goal Seek "
                    "has to iterate here, and it stops when the margin is within its tolerance of 5%. That's why the check "
                    "accepts a small range around the exact value.")
    t_dt190 = Task(
        f"Build the one-variable Data Table in Model!F{DT1_FORM}:H{DT1_LAST}: put =B{R['oi']} in G{DT1_FORM} and =B{R['margin']} in "
        f"H{DT1_FORM}, select F{DT1_FORM}:H{DT1_LAST}, and use Column input cell B{R['vpd']}. What operating margin does your "
        "table show at 190 visits per day? Enter it as a percentage to 1 decimal place.",
        answer=margin_190, fmt=PCT1, title="One-variable Data Table: margin at 190 visits/day",
        solution=f"=Model!H{dt1_row[190]}", live=False,
        hint="The visits run down a column, so use the Column input cell",
        explanation=f"The visits per day values run **down a column** (F{DT1_FIRST}:F{DT1_LAST}), so B{R['vpd']} is the "
                    f"**column** input cell and the row input box stays empty. For each value, Excel puts it in B{R['vpd']}, "
                    f"recalculates the model, and writes the results of the formulas in row {DT1_FORM} into that row of the "
                    f"table. You can type the number or point at H{dt1_row[190]}.")
    t_dtfirst = Task(
        "Look down the Operating income column of your one-variable table. What is the lowest visits-per-day value in the "
        "table at which the clinic makes a profit (operating income above 0)?",
        answer=first_profit, title="One-variable Data Table: first profitable volume",
        solution=f'=MINIFS(Model!F{DT1_FIRST}:F{DT1_LAST},Model!G{DT1_FIRST}:G{DT1_LAST},">0")', live=False,
        hint="Read the table, or let MINIFS find it",
        explanation=f"The table jumps in steps of 5, so it brackets the break-even point instead of finding it: "
                    f"{first_profit - 5} visits loses money and {first_profit} makes money, which agrees with Goal Seek's "
                    f"{breakeven_vpd:.1f}. Use a Data Table to see the whole curve, and Goal Seek to pin down the exact crossing. "
                    "MINIFS needs Excel 2019 or later. In older versions, read the value off the table and type it.")
    t_dt2 = Task(
        f"Build the two-variable Data Table: put =B{R['oi']} in the corner cell F{DT2_CORNER}, select "
        f"F{DT2_CORNER}:K{DT2_LAST}, and use Row input cell C{COM_ROW} (commercial $ across row {DT2_CORNER}) and Column "
        f"input cell B{R['vpd']} (visits per day down column F). What operating income does the table show at 210 visits "
        "per day and $160 per commercial visit? Round to the nearest dollar.",
        answer=round(oi_210_160), tol=0.51, fmt="#,##0", title="Two-variable Data Table: 210 visits/day × $160",
        solution=f"=Model!{dt2_col[160]}{dt2_row[210]}", live=False,
        hint="Row input = the input whose values run across the top row",
        explanation=f"A two-variable table has exactly one formula, in its top-left corner. Excel substitutes each top-row value "
                    f"into the **row** input cell (C{COM_ROW}) and each left-column value into the **column** input cell "
                    f"(B{R['vpd']}), and fills every intersection. If you swap the two input cells, the table still fills "
                    f"without any warning, but with wrong numbers. So check one cell by hand. The base case is {VISITS_PER_DAY} "
                    f"visits at ${rate['Commercial']} ({money(base['oi'])}), so the cell at {VISITS_PER_DAY} visits and $140 should "
                    f"be about {money(base['oi'] - dt2[(VISITS_PER_DAY, 140)])} lower, because the clinic's "
                    f"{s_c * base['vpm']:,.0f} commercial visits a month each pay $2 less, minus the {BILLING_FEE:.0%} fee on "
                    "those dollars. With swapped input cells that cell would show 140 visits a day at $200, a far bigger loss.")
    t_dtcount = Task(
        f"How many of the 45 combinations in your two-variable table (G{DT2_FIRST}:K{DT2_LAST}) are profitable (operating "
        "income above 0)? Use a formula.",
        answer=n_profitable, title="Two-variable Data Table: profitable combinations",
        solution=f'=COUNTIF(Model!G{DT2_FIRST}:K{DT2_LAST},">0")', live=False, hint="COUNTIF with \">0\"",
        explanation="Ordinary formulas can read a Data Table's results, so you can count them, chart them, or add "
                    "conditional formatting. The profitable combinations sit in the bottom-right of the table, where volume and "
                    f"rate are both high. Reading down each column shows the trade-off the CFO cares about: at "
                    f"${DT2_RATES[0]} per commercial visit the clinic first makes money at {dt2_first[DT2_RATES[0]]} visits a "
                    f"day, but at ${DT2_RATES[-1]} it does at {dt2_first[DT2_RATES[-1]]}. A better contract lowers the "
                    "volume the clinic needs to break even.")
    scen_rows = ". ".join(f"{k}: visits per day {s['vpd']}, Medicaid share {s['medicaid']:.1%}, Commercial share "
                          f"{s['commercial']:.1%}, billing fee {s['fee']:.1%}" for k, s in SCENARIOS.items())
    t_down = Task(
        f"Open Scenario Manager on the Model and add three scenarios that change B{R['vpd']}, B{MED_ROW}, B{COM_ROW} and "
        f"B{R['fee']}. {scen_rows}. Create a Scenario Summary with result cells B{R['oi']} and B{R['margin']}. What is "
        "operating income in the Downside scenario? Round to the nearest dollar.",
        answer=round(scen["Downside"]["oi"]), tol=0.51, fmt="#,##0", title="Scenario Manager: Downside operating income",
        solution=f"1. Select B{R['vpd']}, then Ctrl-click (Mac: ⌘-click) B{MED_ROW}:B{COM_ROW} and B{R['fee']}.\n"
                 "2. **Data → What-If Analysis → Scenario Manager → Add…**. Name it *Base plan* and click **OK**. The values "
                 "box shows the current values, so click **OK** again.\n"
                 f"3. Click **Add…** again for *Downside* and *Upside*. Type the values in the order B{R['vpd']}, B{MED_ROW}, B{COM_ROW}, B{R['fee']} "
                 "(decimals like 0.179 for 17.9%).\n"
                 f"4. Click **Summary…**, choose **Scenario summary**, set **Result cells** to `B{R['oi']},B{R['margin']}`, and "
                 "click **OK**. Read the Downside column on the new Scenario Summary sheet.",
        live=False, hint="Data → What-If Analysis → Scenario Manager",
        explanation=f"Fewer visits, a shift from commercial to Medicaid, and a higher vendor fee all hit income at once. If you "
                    f"skipped the fix in task 3, the model ignores the 5% fee and you'd see {money(down_unfixed_oi)} instead. "
                    "That's exactly the kind of silent error a hard-coded number causes. The summary sheet is a snapshot. It "
                    "doesn't update when the model changes, so create it again after any edit.")
    t_up = Task(
        "From the same Scenario Summary, what operating margin does the Upside scenario produce? Enter it as a percentage "
        "to 1 decimal place.",
        answer=scen["Upside"]["margin"], fmt=PCT1, title="Scenario Manager: Upside margin",
        solution="Read the Upside column of the OperatingMargin row on the Scenario Summary sheet.",
        live=False, hint="Same summary, different column",
        explanation=f"The Upside case brings the clinic to roughly {scen['Upside']['margin']:.1%}. It takes three things at once: "
                    f"{SCENARIOS['Upside']['vpd'] - VISITS_PER_DAY} more visits a day, "
                    f"{(SCENARIOS['Upside']['commercial'] - share['Commercial']) * 100:.0f} points of payer mix shifted from "
                    f"Medicaid to Commercial, and a cheaper vendor contract ({SCENARIOS['Upside']['fee']:.1%}). Because "
                    f"B{R['oi']} and B{R['margin']} are named, the summary labels those rows OperatingIncome and "
                    f"OperatingMargin instead of $B${R['oi']} and $B${R['margin']}.")
    t_solver = Task(
        f"On the Staffing sheet, use Solver to minimize the monthly staff cost (B{S['objective']}) by changing the RN, LPN, "
        f"and CNA FTEs (B{S['rn']}:B{S['cna']}), subject to the three constraints in rows {S['c_total']}–{S['c_rn']} "
        f"(leave out the CNA cap in row {S['c_cap']}). Keep Make Unconstrained Variables Non-Negative ticked and choose "
        "Simplex LP. What is the minimum monthly cost? Round to the nearest dollar.",
        answer=round(float(lp_cost)), tol=0.51, fmt="#,##0", title="Solver: lowest-cost staffing mix (FTEs)",
        solution=f"1. Choose **Data → Solver**.\n"
                 f"2. **Set Objective:** `$B${S['objective']}` · **To:** Min · **By Changing Variable Cells:** "
                 f"`$B${S['rn']}:$B${S['cna']}`.\n"
                 f"3. Click **Add** and enter `$B${S['c_total']}:$B${S['c_rn']}` **>=** `$D${S['c_total']}:$D${S['c_rn']}`, "
                 "then click **OK**.\n"
                 "4. Leave **Make Unconstrained Variables Non-Negative** ticked, choose **Simplex LP**, and click **Solve**.\n"
                 f"5. Choose **Keep Solver Solution**, click **OK**, and type the value of B{S['objective']}, rounded to the "
                 "dollar, in the answer cell.",
        # Live check: the LP optimum in closed form, read from the Staffing sheet's own cells. Licensed FTEs cover the
        # licensed hours exactly and split at the RN minimum; CNAs cover the remaining support hours.
        live=(f"=Staffing!B{S['vpm']}*Staffing!B{S['lic_hpv']}/Staffing!B{S['hrs']}*(Staffing!B{S['rn_share']}"
              f"*Staffing!D{S['rn']}+(1-Staffing!B{S['rn_share']})*Staffing!D{S['lpn']})+Staffing!B{S['vpm']}"
              f"*(Staffing!B{S['sup_hpv']}-Staffing!B{S['lic_hpv']})/Staffing!B{S['hrs']}*Staffing!D{S['cna']}"),
        hint="Data → Solver (enable the Solver add-in first)",
        explanation=f"Solver chooses {float(lp_x[0]):.1f} RN, {float(lp_x[1]):.1f} LPN and "
                    f"{float(lp_x[2]):.1f} CNA FTEs. This answer makes sense: CNAs are the cheapest staff, so they cover every "
                    "hour that doesn't need a licensed nurse. The licensed hours are then split at exactly the 60% RN "
                    f"minimum, because LPNs cost less than RNs. Today's 5/3/8 staffing costs {money(float(current_cost))}, so "
                    f"the plan saves about {money(float(current_cost - lp_cost))} a month. Type the number rather than "
                    f"pointing at B{S['objective']}, because the bonus runs Solver on the same sheet again. The key's live "
                    "result rebuilds this optimum with algebra from the Staffing sheet's own cells.")
    t_shadow = Task(
        f"Run Solver again (same setup) and select Sensitivity under Reports before you click OK. On the Sensitivity Report "
        f"sheet, what is the Shadow Price of the Total support hours constraint (Staffing!B{S['c_total']})? Enter it in "
        "dollars to 2 decimal places.",
        answer=round(float(shadow_total), 2), fmt=MONEY2, tol=0.011, title="Solver Sensitivity Report: shadow price",
        solution=f"1. **Data → Solver**, same parameters as task 12, **Solve**.\n"
                 "2. In **Solver Results**, select **Sensitivity** under *Reports*, then click **OK**.\n"
                 f"3. On the new *Sensitivity Report 1* sheet, find the row for `$B${S['c_total']}` under *Constraints* and "
                 "read **Shadow Price**.",
        live=f"=ROUND({m('cna_rate')}*(1+{m('ben')}),2)",
        hint="Shadow price = cost change per one-unit increase in the constraint's right-hand side",
        explanation=f"Requiring one more support hour raises the minimum cost by {money2(float(shadow_total))}. That's one CNA hour "
                    f"at ${cna_rate:.2f} plus {benefits_load:.1%} benefits, because the cheapest way to cover an hour that "
                    "doesn't need a licensed nurse is a CNA hour. The live result in the key repeats that arithmetic. The "
                    f"licensed-hours constraint (B{S['c_lic']}) has a shadow price of only {money2(float(shadow_licensed))}, "
                    "even though licensed nurses cost more. An extra licensed hour costs 0.6 RN hours plus 0.4 LPN hours, "
                    "but licensed hours also count toward the total, so it replaces a CNA hour that's no longer needed.")
    L.tasks = [t_oi, t_margin, t_hard, t_be, t_rate, t_dt190, t_dtfirst, t_dt2, t_dtcount, t_down, t_up, t_solver,
               t_shadow]

    # ======================================================================== bonus
    L.bonus_title = "Bonus: Staff the 2026 growth plan"
    L.bonus_scenario = (
        f"Bluestone wants the Primary Care Clinic to grow to {GROWTH_VPD} visits per day in 2026 by adding a fourth APP "
        f"(capacity rises from {MD_FTE * MD_VPD + APP_FTE * APP_VPD} to {MD_FTE * MD_VPD + GROWTH_APP * APP_VPD} visits per "
        f"day). HR hires whole people, so every role needs a whole number of 1.0-FTE staff. Only {CNA_CAP} CNA positions "
        f"are approved (Staffing!B{S['cap']}). Find the cheapest staffing plan, then test the full plan in the P&L.")
    L.bonus = [
        Task(f"On the Staffing sheet, change B{S['vpd']} to {GROWTH_VPD}. Re-run Solver with two more constraints: "
             f"B{S['rn']}:B{S['cna']} = int (integer) and B{S['c_cap']} <= D{S['c_cap']} (the CNA cap). In Solver Options, set "
             "Integer Optimality (%) to 0. What is the minimum monthly staff cost? Round to the nearest dollar.",
             answer=round(float(g_int_cost)), tol=0.51, fmt="#,##0", title="Integer staffing plan: minimum cost",
             solution=f"1. Type {GROWTH_VPD} in Staffing!B{S['vpd']}.\n"
                      f"2. **Data → Solver**. Keep the objective, variable cells, and the three constraints from task 12.\n"
                      f"3. **Add** `$B${S['rn']}:$B${S['cna']}` **int** and **Add** `$B${S['c_cap']} <= $D${S['c_cap']}`.\n"
                      "4. **Options** → *All Methods* tab → **Integer Optimality (%)** = 0 → **OK**. Check that **Ignore "
                      "Integer Constraints** is not ticked.\n"
                      f"5. **Solve**, keep the solution, and read B{S['objective']}.",
             live=False, hint="Add an int constraint on the three variable cells",
             explanation=f"With fractions allowed, the cheapest plan costs {money(float(g_lp_cost))} "
                         f"({float(g_lp_x[0]):.1f} RN, {float(g_lp_x[1]):.1f} LPN, {float(g_lp_x[2]):.1f} CNA). Requiring "
                         f"whole people raises that to {money(float(g_int_cost))}, so indivisibility costs "
                         f"{money(float(g_int_cost - g_lp_cost))} a month. Solver handles integer constraints by "
                         "**branch and bound**, solving many LPs with tighter and tighter bounds. Its default Integer "
                         "Optimality of 1% lets it stop at any plan within 1% of the best possible, so set it to 0 when "
                         "the exact answer matters."),
        Task("How many RNs does that plan hire?", answer=g_rn, title="Integer plan: RNs",
             solution=f"Read Staffing!B{S['rn']} after Solver finishes.", live=False,
             hint=f"Look at B{S['rn']}",
             explanation=f"{g_rn} RNs out of {g_rn + g_lpn} licensed staff is {g_rn / (g_rn + g_lpn):.1%}, which meets the "
                         "60% rule."),
        Task("How many LPNs does that plan hire?", answer=g_lpn, title="Integer plan: LPNs",
             solution=f"Read Staffing!B{S['lpn']} after Solver finishes.", live=False,
             hint="Is rounding the fractional answer up the same thing?",
             explanation=f"The fractional plan needs {float(g_lp_x[0] + g_lp_x[1]):.1f} licensed FTEs, and "
                         f"{g_rn + g_lpn} whole people already cover that. A tempting shortcut is to round each fractional "
                         f"value up, which gives {naive[0]} RN, {naive[1]} LPN and {naive[2]} CNA. That plan hires one "
                         f"licensed person too many and costs {money(float(naive_cost - g_int_cost))} a month more. "
                         + ("It also breaks the RN rule, because "
                            f"{naive[0]} of {naive[0] + naive[1]} licensed staff is only {naive_share:.1%}. "
                            if not naive_ok else "")
                         + "Rounding each variable separately can't see how the constraints interact, but Solver's "
                           "integer search can."),
        Task(f"Test the plan in the P&L. On the Model, add two more scenarios that change B{R['vpd']}, B{R['app']} and "
             f"B{R['rn']}:B{R['cna']}: Today ({VISITS_PER_DAY}, {APP_FTE}, {RN_FTE}, {LPN_FTE}, {CNA_FTE}) and Growth 2026 "
             f"({GROWTH_VPD}, {GROWTH_APP}, and your RN, LPN, "
             "and CNA counts). What operating income does Growth 2026 produce? Round to the nearest dollar.",
             answer=round(growth["oi"]), tol=0.51, fmt="#,##0", title="Growth 2026 scenario: operating income",
             solution=f"1. On the Model, select B{R['vpd']}, B{R['app']} and B{R['rn']}:B{R['cna']} (Ctrl-click, or "
                      "⌘-click on a Mac).\n"
                      f"2. **Scenario Manager → Add…** *Today*: {VISITS_PER_DAY}, {APP_FTE}, {RN_FTE}, {LPN_FTE}, {CNA_FTE}.\n"
                      f"3. **Add…** *Growth 2026*: {GROWTH_VPD}, {GROWTH_APP}, {g_rn}, {g_lpn}, {g_cna}.\n"
                      f"4. **Summary…** with result cells `B{R['oi']},B{R['margin']}`, or select Growth 2026, click "
                      f"**Show**, read B{R['oi']}, then **Show** Today to restore the base case.",
             live=False, hint="Scenario Manager can switch five inputs at once, and switch them back",
             explanation=f"The growth plan turns a {money(base['oi'])} monthly loss into {money(growth['oi'])} "
                         f"({growth['margin']:.1%} margin). The new APP costs ${APP_SALARY:,} plus benefits, and the larger "
                         f"staff costs {money(float(g_int_cost - current_cost))} more a month than today's. The extra "
                         f"{(GROWTH_VPD - VISITS_PER_DAY) * CLINIC_DAYS} visits a month each contribute "
                         f"{money2(base['cm'])}, which covers both. Using scenarios instead of overwriting inputs means one "
                         "click (Show Today) puts the base case back for every other task."),
    ]

    L.start_notes = [
        "Model is the clinic's monthly P&L, laid out as INPUTS (blue numbers), CALCULATIONS (black formulas), and OUTPUTS. You "
        "complete two yellow output cells and build two Data Tables in the Sensitivity area (columns F–K).",
        "Staffing is a Solver worksheet. The green cells are the decision variables that Solver changes. Sources lists where "
        "every data-derived input came from.",
        "Goal Seek and Scenario Manager's Show button change the Model's input cells. Put the base case back after each run "
        "so the other tasks stay correct. Solver only changes the Staffing sheet, which doesn't feed the Model.",
        "Solver is a free add-in that ships with Excel. Turn it on once: File → Options → Add-ins → Manage: Excel Add-ins → "
        "Go → tick Solver Add-in (Mac: Tools → Excel Add-ins).",
    ]
    L.sheet_order = ["Start Here", "Practice", "Model", "Staffing", "Sources", "Bonus", "Answer Key", "Bonus Key"]

    # ======================================================================== Model + Staffing sheets
    model_inputs = [
        # (key, label, value, fmt, unit, note)
        ("days", "Clinic days per month", CLINIC_DAYS, "0", "days", "Weekdays the clinic is open, less holidays (plan)"),
        ("vpd", "Visits per day", VISITS_PER_DAY, "0.0", "visits", "FY2026 plan"),
        ("md", "Physician FTEs (MD/DO)", MD_FTE, "0.0", "FTE", "Current staffing"),
        ("app", "APP FTEs (NP/PA)", APP_FTE, "0.0", "FTE", "Current staffing"),
        ("md_vpd", "Visits per physician per day", MD_VPD, "0", "visits", "Scheduling template"),
        ("app_vpd", "Visits per APP per day", APP_VPD, "0", "visits", "Scheduling template"),
    ]
    staff_inputs = [
        ("rn", "RN FTEs", RN_FTE, "0.0", "FTE", "Current staffing"),
        ("lpn", "LPN FTEs", LPN_FTE, "0.0", "FTE", "Current staffing"),
        ("cna", "CNA FTEs", CNA_FTE, "0.0", "FTE", "Current staffing"),
        ("fd", "Front desk & admin FTEs", FD_FTE, "0.0", "FTE", "Current staffing"),
        ("hrs", "Paid hours per FTE per month", HOURS_PER_FTE, "0", "hours", "Full-time paid hours per month (plan)"),
        ("md_sal", "Physician salary per FTE", MD_SALARY, "#,##0", "$/month", "FY2026 compensation plan"),
        ("app_sal", "APP salary per FTE", APP_SALARY, "#,##0", "$/month", "FY2026 compensation plan"),
        ("rn_rate", "RN hourly rate", rn_rate, MONEY2, "$/hour", "Median of active RNs (Sources)"),
        ("lpn_rate", "LPN hourly rate", lpn_rate, MONEY2, "$/hour", "Median of active LPNs (Sources)"),
        ("cna_rate", "CNA hourly rate", cna_rate, MONEY2, "$/hour", "Median of active CNAs (Sources)"),
        ("fd_rate", "Front desk hourly rate", fd_rate, MONEY2, "$/hour", "Median of active unit secretaries (Sources)"),
        ("ben", "Benefits load (% of wages)", benefits_load, PCT1, "", "2025 D400 benefits ÷ wages (Sources)"),
    ]
    other_inputs = [
        ("sup", "Medical supplies per visit", supplies_pv, MONEY2, "$/visit", "2025 D400 actuals (Sources)"),
        ("vac", "Vaccines & injectables per visit", vaccines_pv, MONEY2, "$/visit", "2025 D400 actuals (Sources)"),
        ("fee", "Billing & collections fee", BILLING_FEE, PCT1, "of revenue", "Revenue-cycle vendor contract"),
        ("fx1", "Facility & equipment", fixed_equip, "#,##0", "$/month", "Fixed. 2025 D400 actuals (Sources)"),
        ("fx2", "IT, EHR & purchased services", fixed_purch, "#,##0", "$/month", "Fixed. 2025 D400 actuals (Sources)"),
        ("fx3", "Other operating", fixed_other, "#,##0", "$/month", "Fixed. 2025 D400 actuals (Sources)"),
    ]
    calcs = [
        ("vpm", "Visits per month", "=VisitsPerDay*ClinicDays", "#,##0", "visits"),
        ("avg", "Average reimbursement per visit", f"=SUMPRODUCT(B{R['payer0']}:B{R['payer0'] + 4},"
                                                   f"C{R['payer0']}:C{R['payer0'] + 4})", MONEY2, "$/visit"),
        ("rev", "Net patient revenue", "=VisitsPerMonth*AvgReimbursement", MONEY0, "$/month"),
        ("w_prov", "Physician & APP wages", f"=B{R['md']}*B{R['md_sal']}+B{R['app']}*B{R['app_sal']}", MONEY0, "$/month"),
        ("w_sup", "Clinical support wages (RN, LPN, CNA)",
         f"=(B{R['rn']}*B{R['rn_rate']}+B{R['lpn']}*B{R['lpn_rate']}+B{R['cna']}*B{R['cna_rate']})*B{R['hrs']}", MONEY0,
         "$/month"),
        ("w_fd", "Front desk & admin wages", f"=B{R['fd']}*B{R['fd_rate']}*B{R['hrs']}", MONEY0, "$/month"),
        ("wages", "Total wages", f"=SUM(B{R['w_prov']}:B{R['w_fd']})", MONEY0, "$/month"),
        ("benefits", "Benefits", f"=B{R['wages']}*B{R['ben']}", MONEY0, "$/month"),
        ("c_sup", "Medical supplies", f"=VisitsPerMonth*B{R['sup']}", MONEY0, "$/month"),
        ("c_vac", "Vaccines & injectables", f"=VisitsPerMonth*B{R['vac']}", MONEY0, "$/month"),
        ("c_fee", "Billing & collections fee", "=NetRevenue*0.04", MONEY0, "$/month"),   # deliberate hard-code (task 3)
        ("c_fx1", "Facility & equipment", f"=B{R['fx1']}", MONEY0, "$/month"),
        ("c_fx2", "IT, EHR & purchased services", f"=B{R['fx2']}", MONEY0, "$/month"),
        ("c_fx3", "Other operating", f"=B{R['fx3']}", MONEY0, "$/month"),
        ("exp", "Total operating expenses", f"=SUM(B{R['wages']}:B{R['c_fx3']})", MONEY0, "$/month"),
    ]
    names = {
        "ClinicDays": f"Model!$B${R['days']}", "VisitsPerDay": f"Model!$B${R['vpd']}",
        "MedicaidShare": f"Model!$B${MED_ROW}", "CommercialShare": f"Model!$B${COM_ROW}",
        "CommercialRate": f"Model!$C${COM_ROW}", "BillingFeePct": f"Model!$B${R['fee']}",
        "VisitsPerMonth": f"Model!$B${R['vpm']}", "AvgReimbursement": f"Model!$B${R['avg']}",
        "NetRevenue": f"Model!$B${R['rev']}", "TotalExpenses": f"Model!$B${R['exp']}",
        "OperatingIncome": f"Model!$B${R['oi']}", "OperatingMargin": f"Model!$B${R['margin']}",
    }

    def section(ws, row, text, c0=1, c1=4):
        for c in range(c0, c1 + 1):
            ws.cell(row=row, column=c).fill = HEADER_FILL
        cell = ws.cell(row=row, column=c0, value=text)
        cell.font = SECTION_FONT

    def subheader(ws, row, labels, c0=1):
        for j, h in enumerate(labels, c0):
            cell = ws.cell(row=row, column=j, value=h)
            cell.font = Font(bold=True, color=NAVY)
            cell.fill = SUB_FILL
            cell.border = BOX

    def input_cell(cell, value, fmt):
        cell.value = value
        cell.font = INPUT_FONT
        cell.fill = INPUT_CELL_FILL
        cell.border = BOX
        cell.number_format = fmt

    def labeled(ws, row, label, unit="", note=""):
        ws.cell(row=row, column=1, value=label).border = BOX
        u = ws.cell(row=row, column=3, value=unit or None)
        u.font = Font(color="595959")
        u.border = BOX
        if note:
            ws.cell(row=row, column=4, value=note).font = NOTE_FONT

    def yellow(cell, fmt):
        cell.fill = INPUT_FILL
        cell.border = INPUT_BORDER
        cell.number_format = fmt

    def fit(ws, landscape=True):
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    @L.customize
    def _sheets(wb, lesson, selftest):
        # ------------------------------------------------------------------ Model
        ws = wb.create_sheet("Model")
        ws.sheet_properties.tabColor = "2E75B6"
        ws["A1"] = "Primary Care Clinic (D400) · Monthly Operating Model"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = "Bluestone Outpatient Pavilion · FY2026 plan · $ per month unless stated"
        ws["A2"].font = NOTE_FONT
        ws["A3"] = "Blue = input (change these) · Black = formula (don't type over) · Yellow = you complete it."
        ws["A3"].font = NOTE_FONT
        ws["A4"] = "Named cells (Formulas → Name Manager): " + ", ".join(names) + "."
        ws["A4"].font = Font(italic=True, color="595959", size=9)
        section(ws, 5, "INPUTS")
        subheader(ws, 6, ["Volume & capacity", "Value", "Unit", "Source / note"])
        for key, label, val, fmt, unit, note in model_inputs:
            labeled(ws, R[key], label, unit, note)
            input_cell(ws[f"B{R[key]}"], val, fmt)
        subheader(ws, R["payer_hdr"], ["Payer mix & reimbursement", "Share of visits", "$ per visit", "Source / note"])
        for g in GROUPS:
            r = payer_row[g]
            ws.cell(row=r, column=1, value=GROUP_LABEL[g]).border = BOX
            input_cell(ws[f"B{r}"], share[g], PCT1)
            input_cell(ws[f"C{r}"], rate[g], MONEY2)
            ws.cell(row=r, column=4, value="2025 D400 claims (Sources)").font = NOTE_FONT
        mt = R["mix_total"]
        ws.cell(row=mt, column=1, value="Payer mix total (check)").font = Font(italic=True)
        ws[f"B{mt}"] = f"=SUM(B{R['payer0']}:B{R['payer0'] + 4})"
        ws[f"B{mt}"].number_format = PCT1
        ws[f"C{mt}"] = f'=IF(ROUND(B{mt},6)=1,"✔ totals 100%","✘ must total 100%")'
        ws[f"C{mt}"].font = Font(italic=True, color="595959")
        subheader(ws, R["rn"] - 1, ["Staffing & pay", "Value", "Unit", "Source / note"])
        for key, label, val, fmt, unit, note in staff_inputs:
            labeled(ws, R[key], label, unit, note)
            input_cell(ws[f"B{R[key]}"], val, fmt)
        subheader(ws, R["sup"] - 1, ["Other costs", "Value", "Unit", "Source / note"])
        for key, label, val, fmt, unit, note in other_inputs:
            labeled(ws, R[key], label, unit, note)
            input_cell(ws[f"B{R[key]}"], val, fmt)

        section(ws, R["vpm"] - 1, "CALCULATIONS  (one formula per row · no typed-in numbers)")
        for key, label, formula, fmt, unit in calcs:
            labeled(ws, R[key], label, unit)
            lesson.set_formula(ws, f"B{R[key]}", formula, dynamic=False)
            ws[f"B{R[key]}"].number_format = fmt
            ws[f"B{R[key]}"].border = BOX
        for key in ("wages", "exp"):
            for col in "AB":
                ws[f"{col}{R[key]}"].font = Font(bold=True)
        ws[f"B{R['exp']}"].border = TOTAL_BORDER

        section(ws, R["oi"] - 1, "OUTPUTS")
        labeled(ws, R["oi"], "Operating income", "$/month", "← you complete this (task 1)")
        labeled(ws, R["margin"], "Operating margin", "of revenue", "← you complete this (task 2)")
        ws[f"A{R['oi']}"].font = Font(bold=True)
        ws[f"A{R['margin']}"].font = Font(bold=True)
        yellow(ws[OI], MONEY0)
        yellow(ws[MARGIN], PCT1)
        labeled(ws, R["cpv"], "Cost per visit", "$/visit")
        lesson.set_formula(ws, f"B{R['cpv']}", "=TotalExpenses/VisitsPerMonth", dynamic=False)
        ws[f"B{R['cpv']}"].number_format = MONEY2
        labeled(ws, R["cap"], "Visit capacity per day", "visits", "Providers × visits per provider per day")
        ws[f"B{R['cap']}"] = f"=B{R['md']}*B{R['md_vpd']}+B{R['app']}*B{R['app_vpd']}"
        labeled(ws, R["util"], "Capacity utilization", "of capacity", "Above 100% means the plan needs more providers")
        lesson.set_formula(ws, f"B{R['util']}", f"=VisitsPerDay/B{R['cap']}", dynamic=False)
        ws[f"B{R['util']}"].number_format = PCT1
        for key in ("cpv", "cap", "util"):
            ws[f"B{R[key]}"].border = BOX

        # Sensitivity area (Data Tables must live on the same sheet as their input cells)
        section(ws, 5, "SENSITIVITY AREA  (Data Tables)", c0=6, c1=11)
        ws.cell(row=6, column=6, value="One-variable Data Table · Column input cell: B8").font = Font(bold=True, color=NAVY)
        subheader(ws, DT1_HDR, ["Visits per day", "Operating income", "Operating margin"], c0=6)
        yellow(ws[f"G{DT1_FORM}"], MONEY0)
        yellow(ws[f"H{DT1_FORM}"], PCT1)
        ws[f"F{DT1_FORM}"].border = BOX
        ws[f"I{DT1_FORM}"] = "← type =B59 and =B60 here (task 6)"
        ws[f"I{DT1_FORM}"].font = NOTE_FONT
        for v in DT1_VISITS:
            r = dt1_row[v]
            c = ws.cell(row=r, column=6, value=v)
            c.font = Font(bold=True)
            c.fill = SUB_FILL
            c.border = BOX
            c.alignment = Alignment(horizontal="center")
            for col, fmt in (("G", MONEY0), ("H", PCT1)):
                ws[f"{col}{r}"].fill = RESULT_FILL
                ws[f"{col}{r}"].border = BOX
                ws[f"{col}{r}"].number_format = fmt
        ws.cell(row=DT1_LAST + 1, column=6,
                value=f"Select F{DT1_FORM}:H{DT1_LAST} → Data → What-If Analysis → Data Table → Column input cell B8. "
                      "The gray cells fill in.").font = NOTE_FONT

        ws.cell(row=DT2_CORNER - 2, column=6,
                value="Two-variable Data Table · Row input cell: C17 · Column input cell: B8").font = Font(bold=True, color=NAVY)
        c = ws.cell(row=DT2_CORNER - 1, column=6, value="↓ Visits per day")
        c.font = Font(bold=True, color=NAVY)
        ws.merge_cells(start_row=DT2_CORNER - 1, start_column=7, end_row=DT2_CORNER - 1, end_column=11)
        c = ws.cell(row=DT2_CORNER - 1, column=7, value="Commercial $ per visit →")
        c.font = Font(bold=True, color=NAVY)
        c.alignment = Alignment(horizontal="center")
        yellow(ws[f"F{DT2_CORNER}"], MONEY0)
        for x in DT2_RATES:
            c = ws[f"{dt2_col[x]}{DT2_CORNER}"]
            c.value = x
            c.font = Font(bold=True)
            c.fill = SUB_FILL
            c.border = BOX
            c.number_format = "$#,##0"
            c.alignment = Alignment(horizontal="center")
        for v in DT2_VISITS:
            r = dt2_row[v]
            c = ws.cell(row=r, column=6, value=v)
            c.font = Font(bold=True)
            c.fill = SUB_FILL
            c.border = BOX
            c.alignment = Alignment(horizontal="center")
            for col in DT2_COLS:
                ws[f"{col}{r}"].fill = RESULT_FILL
                ws[f"{col}{r}"].border = BOX
                ws[f"{col}{r}"].number_format = MONEY0
        ws.cell(row=DT2_LAST + 1, column=6,
                value=f"Type =B59 in the yellow corner F{DT2_CORNER}, select F{DT2_CORNER}:K{DT2_LAST} → Data Table → "
                      "Row input cell C17, Column input cell B8.").font = NOTE_FONT

        for col, w in zip("ABCDEFGHIJK", (40, 14, 13, 44, 3, 16, 17, 17, 13, 13, 13)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "B5"
        fit(ws)

        for name, ref in names.items():
            wb.defined_names[name] = DefinedName(name, attr_text=ref)

        # ------------------------------------------------------------------ Staffing
        st = wb.create_sheet("Staffing")
        st.sheet_properties.tabColor = "548235"
        st["A1"] = "Clinical Support Staffing Plan · Solver worksheet"
        st["A1"].font = TITLE_FONT
        st["A2"] = ("Choose RN, LPN, and CNA FTEs that cover the clinic's workload at the lowest monthly cost. "
                    "Blue = input · Green text = linked from the Model · Green cells = decision variables that Solver changes.")
        st["A2"].font = NOTE_FONT
        section(st, 5, "INPUTS", c1=5)
        st_inputs = [
            ("vpd", "Visits per day to staff for", VISITS_PER_DAY, "0", "visits", "Model base case. The bonus changes it."),
            ("sup_hpv", "Support hours needed per visit (RN + LPN + CNA)", SUPPORT_HRS_PER_VISIT, "0.00", "hours",
             "Rooming, vitals, injections, care coordination (time study)"),
            ("lic_hpv", "Licensed hours needed per visit (RN + LPN only)", LICENSED_HRS_PER_VISIT, "0.00", "hours",
             "Triage calls, medications, patient teaching"),
            ("rn_share", "Minimum RN share of licensed FTEs", RN_MIN_SHARE, "0%", "", "Clinic policy"),
            ("cap", "CNA cap", CNA_CAP, "0", "FTEs", "Approved CNA positions (bonus only)"),
        ]
        for key, label, val, fmt, unit, note in st_inputs:
            labeled(st, S[key], label, unit, note)
            input_cell(st[f"B{S[key]}"], val, fmt)
        for key, label, formula, fmt, unit, note, link in (
                ("days", "Clinic days per month", f"=Model!B{R['days']}", "0", "days", "From the Model", True),
                ("vpm", "Visits per month", f"=B{S['vpd']}*B{S['days']}", "#,##0", "visits", "", False),
                ("hrs", "Paid hours per FTE per month", f"=Model!B{R['hrs']}", "0", "hours", "From the Model", True),
                ("ben", "Benefits load", f"=Model!B{R['ben']}", PCT1, "", "From the Model", True)):
            labeled(st, S[key], label, unit, note)
            st[f"B{S[key]}"] = formula
            st[f"B{S[key]}"].number_format = fmt
            st[f"B{S[key]}"].border = BOX
            if link:
                st[f"B{S[key]}"].font = LINK_FONT

        section(st, S["rn"] - 2, "DECISION VARIABLES  (Solver changes the green cells)", c1=5)
        subheader(st, S["rn"] - 1, ["Role", "FTEs", "Hourly rate", "Cost per FTE per month", "Monthly cost"])
        for key, role, fte, rate_row in (("rn", "Registered Nurse (RN)", RN_FTE, R["rn_rate"]),
                                         ("lpn", "Licensed Practical Nurse (LPN)", LPN_FTE, R["lpn_rate"]),
                                         ("cna", "Certified Nursing Assistant (CNA)", CNA_FTE, R["cna_rate"])):
            r = S[key]
            st.cell(row=r, column=1, value=role).border = BOX
            c = st[f"B{r}"]
            c.value = fte
            c.font = Font(bold=True, color="0000FF")
            c.fill = DECISION_FILL
            c.border = BOX
            c.number_format = "0.00"
            st[f"C{r}"] = f"=Model!B{rate_row}"
            st[f"C{r}"].font = LINK_FONT
            st[f"D{r}"] = f"=C{r}*$B${S['hrs']}*(1+$B${S['ben']})"
            st[f"E{r}"] = f"=B{r}*D{r}"
            for col, fmt in (("C", MONEY2), ("D", MONEY2), ("E", MONEY0)):
                st[f"{col}{r}"].number_format = fmt
                st[f"{col}{r}"].border = BOX
        tr = S["total"]
        st.cell(row=tr, column=1, value="Total").font = Font(bold=True)
        st[f"B{tr}"] = f"=SUM(B{S['rn']}:B{S['cna']})"
        st[f"B{tr}"].number_format = "0.00"
        st[f"E{tr}"] = f"=SUM(E{S['rn']}:E{S['cna']})"
        st[f"E{tr}"].number_format = MONEY0
        for col in "ABCDE":
            st[f"{col}{tr}"].border = TOTAL_BORDER
            st[f"{col}{tr}"].font = Font(bold=True)

        section(st, S["c_total"] - 2, "CONSTRAINTS  (each row: Provided  sign  Required)", c1=5)
        subheader(st, S["c_total"] - 1, ["Constraint", "Provided", "Sign", "Required", "Status"])
        rows = [
            ("c_total", "Total support hours", f"=B{S['total']}*B{S['hrs']}", "≥", f"=B{S['vpm']}*B{S['sup_hpv']}", "#,##0.0"),
            ("c_lic", "Licensed hours (RN + LPN)", f"=(B{S['rn']}+B{S['lpn']})*B{S['hrs']}", "≥",
             f"=B{S['vpm']}*B{S['lic_hpv']}", "#,##0.0"),
            ("c_rn", "RN rule: RN FTEs − 60% × (RN + LPN FTEs)", f"=B{S['rn']}-B{S['rn_share']}*(B{S['rn']}+B{S['lpn']})",
             "≥", 0, "0.00"),
            ("c_cap", "CNA FTEs (bonus only)", f"=B{S['cna']}", "≤", f"=B{S['cap']}", "0.00"),
        ]
        for key, label, lhs, sign, rhs, fmt in rows:
            r = S[key]
            st.cell(row=r, column=1, value=label).border = BOX
            st[f"B{r}"] = lhs
            st[f"C{r}"] = sign
            st[f"C{r}"].alignment = Alignment(horizontal="center")
            st[f"D{r}"] = rhs
            ok = (f'=IF(B{r}<=D{r}+0.001,"OK","Over cap (bonus only)")' if sign == "≤"
                  else f'=IF(B{r}>=D{r}-0.001,"OK","Short")')
            st[f"E{r}"] = ok
            for col in "BCDE":
                st[f"{col}{r}"].border = BOX
            st[f"B{r}"].number_format = fmt
            st[f"D{r}"].number_format = fmt
        section(st, S["objective"] - 1, "OBJECTIVE", c1=5)
        st.cell(row=S["objective"], column=1, value="Monthly clinical support staff cost (minimize)").font = Font(bold=True)
        st[f"B{S['objective']}"] = f"=E{S['total']}"
        st[f"B{S['objective']}"].number_format = MONEY0
        st[f"B{S['objective']}"].font = Font(bold=True)
        st[f"B{S['objective']}"].border = TOTAL_BORDER
        for col, w in zip("ABCDE", (46, 14, 14, 22, 22)):
            st.column_dimensions[col].width = w
        st.column_dimensions["F"].width = 50
        for key, *_rest in st_inputs:
            st[f"D{S[key]}"].alignment = Alignment(horizontal="left")
        fit(st)

        # Sources: wrap the long derivation text
        so = wb["Sources"]
        for row in so.iter_rows(min_row=2, max_row=so.max_row):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        fit(so)

        # Task 3 and the guide promise that Trace Dependents on the fee input finds no formulas at all, so nothing in the
        # workbook (including the hidden keys) may read Model!B36 until the learner fixes the hard-coded fee.
        same_sheet = re.compile(rf"(?<![A-Za-z$!])\$?B\$?{R['fee']}(?!\d)")
        other_sheet = re.compile(rf"Model!\$?B\$?{R['fee']}(?!\d)|\bBillingFeePct\b")
        for sh in wb.worksheets:
            for row in sh.iter_rows():
                for c in row:
                    f = c.value.text if hasattr(c.value, "text") else c.value
                    if c.data_type != "f" and not hasattr(c.value, "text") or not isinstance(f, str):
                        continue
                    assert not other_sheet.search(f) and not (sh.title == "Model" and same_sheet.search(f)), \
                        f"{sh.title}!{c.coordinate} reads the fee input: {f}"

        # ------------------------------------------------------------------ self-test: simulate the learner's work
        if selftest:
            # LibreOffice's data-table engine (MULTIPLE.OPERATIONS) returns Err:504 when an input cell sits inside a
            # range argument such as SUMPRODUCT(B14:B18,C14:C18). Excel has no such limit. The self-test copy uses the
            # mathematically identical explicit sum so LibreOffice can evaluate the two-variable table.
            ws[f"B{R['avg']}"] = "=" + "+".join(f"B{payer_row[g]}*C{payer_row[g]}" for g in GROUPS)
            lesson.set_formula(ws, f"G{DT1_FORM}", f"=B{R['oi']}", dynamic=False)
            lesson.set_formula(ws, f"H{DT1_FORM}", f"=B{R['margin']}", dynamic=False)
            ws[f"G{DT1_FIRST}"] = DataTableFormula(ref=f"G{DT1_FIRST}:H{DT1_LAST}", r1=f"B{R['vpd']}")
            lesson.set_formula(ws, f"F{DT2_CORNER}", f"=B{R['oi']}", dynamic=False)
            ws[f"G{DT2_FIRST}"] = DataTableFormula(ref=f"G{DT2_FIRST}:K{DT2_LAST}", dt2D=True,
                                                   r1=f"C{COM_ROW}", r2=f"B{R['vpd']}")

    return L
