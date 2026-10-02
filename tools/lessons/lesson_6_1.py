"""Lesson 6.1 · Capstone: Hospital Performance Review.

The capstone hands the learner a realistic multi-table extract for 2025 (every encounter discharged in 2025, its claim,
every 2025 ED arrival, the 2025 patient-experience surveys, the 2025 unit census, and the four dimension tables) and a
brief from the CMO and CFO. The Practice sheet walks the work plan:

  1–2   data prep: remove a block of duplicate survey rows (the vendor re-sent its October file) and add a LOSDays column
  3–10  KPIs: ED median door-to-provider, ICU occupancy, HCAHPS top-box, readmission rate, O/E LOS index (needs an
        ExpectedLOS lookup column), top service line for readmissions (needs a ServiceLine lookup column), denial rate,
        net collection rate
  11–13 deliverables: a selector-driven Dashboard card, the RefreshReview macro's audit log, and an executive-summary
        sentence built with one formula

The bonus sizes the length-of-stay opportunity at the hospital with the highest O/E index.

Design notes
* Every answer is computed below from the CSVs. Live key formulas never depend on learner columns (they work on the raw
  data, including the duplicate survey rows), so they prove each number independently.
* The self-test simulates the learner's prep work in a customize hook: it removes the duplicate survey rows, fills the
  ExpectedLOS column with XLOOKUP, sets the Dashboard selector to Ashby Falls with a FacilityID lookup, and creates the
  RefreshLog sheet that the macro would create (its values come from the task's fill).
* A hidden 'Reference Dashboard' sheet holds a complete, formula-driven version of the dashboard with two charts.
* The macro (starter + solution) is written to starter/ and solutions/ from the strings in this file.
"""
from __future__ import annotations

import statistics
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from xlcourse import Lesson, Task, data

CODE = "6.1"
READMIT_CUTOFF = datetime(2025, 12, 1)      # index discharges before Dec 1, 2025 have a full 30-day window in the data
TOPBOX_MIN = 9                              # HCAHPS "top box" = overall rating 9 or 10
DUP_MONTH = 10                              # the vendor re-sent the surveys it returned in October 2025

NAVY = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
GIVEN_FILL = PatternFill("solid", fgColor="EDEDED")
KEY_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _pct(x: float, dec: int = 1) -> str:
    """What Excel's TEXT(x, "0.0%") displays (round half away from zero)."""
    q = Decimal(repr(x * 100)).quantize(Decimal(1).scaleb(-dec), rounding=ROUND_HALF_UP)
    return f"{q}%"


def _crlf(text: str) -> bytes:
    text.encode("ascii")  # the VBE reads ANSI, so keep the modules plain ASCII
    return text.replace("\r\n", "\n").replace("\n", "\r\n").encode("ascii")


# ---------------------------------------------------------------------------------------------------------------- VBA
VBA_HEADER = """Attribute VB_Name = "modRefresh"
'==============================================================================
' Lesson 6.1 - Capstone: one-click refresh for the 2025 performance review
' {kind}
'
' Import: in the VBE (Alt+F11; Mac: Option+F11) choose File > Import File...
' Run:    Alt+F8 (Mac: Option+F8) > RefreshReview > Run, or a button on the
'         Dashboard. Save the workbook as .xlsm to keep the code.
'
' RefreshReview does five things, in order:
'   1. Refreshes every Power Query query and PivotTable (Data > Refresh All).
'   2. Removes duplicate rows from the Surveys sheet (SurveyID is the key).
'   3. Recalculates every formula.
'   4. Appends an audit row to the RefreshLog sheet (created on the first run):
'      A RefreshedAt | B Encounters | C ED_Visits | D Claims | E Surveys | F DuplicatesRemoved
'   5. Writes the refresh time into Dashboard!H2.
'==============================================================================
Option Explicit
"""

VBA_HELPERS = """
' Number of data rows under the header row: last used cell in column A, minus the header.
Private Function DataRowCount(ByVal sheetName As String) As Long
    With ThisWorkbook.Worksheets(sheetName)
        DataRowCount = .Cells(.Rows.Count, "A").End(xlUp).Row - 1
    End With
End Function

' Returns the RefreshLog sheet, creating it with headers the first time.
Private Function LogSheet() As Worksheet
    Dim ws As Worksheet
    On Error Resume Next
    Set ws = ThisWorkbook.Worksheets("RefreshLog")
    On Error GoTo 0
    If ws Is Nothing Then
        Set ws = ThisWorkbook.Worksheets.Add(After:=ThisWorkbook.Worksheets(ThisWorkbook.Worksheets.Count))
        ws.Name = "RefreshLog"
        ws.Range("A1:F1").Value = Array("RefreshedAt", "Encounters", "ED_Visits", "Claims", "Surveys", "DuplicatesRemoved")
        ws.Range("A1:F1").Font.Bold = True
        ws.Columns("A").NumberFormat = "mm/dd/yyyy hh:mm"
        ws.Columns("A:F").ColumnWidth = 14
    End If
    Set LogSheet = ws
End Function
"""

VBA_SOLUTION_SUB = """
Public Sub RefreshReview()
    Dim oldCalc As XlCalculation
    Dim surveysBefore As Long, surveysAfter As Long
    Dim wsLog As Worksheet
    Dim r As Long

    oldCalc = Application.Calculation
    On Error GoTo Fail
    Application.ScreenUpdating = False

    ' 1. Refresh queries and PivotTables, then wait for background queries
    ThisWorkbook.RefreshAll
    On Error Resume Next                      ' a failed wait must not stop the review
    Application.CalculateUntilAsyncQueriesDone
    On Error GoTo Fail

    ' 2. Remove duplicate surveys, keeping the first copy of each SurveyID
    surveysBefore = DataRowCount("Surveys")
    ThisWorkbook.Worksheets("Surveys").Range("A1").CurrentRegion.RemoveDuplicates Columns:=1, Header:=xlYes
    surveysAfter = DataRowCount("Surveys")

    ' 3. Recalculate every formula in every open workbook
    Application.Calculation = xlCalculationAutomatic
    Application.CalculateFull

    ' 4. Append the audit row on the next empty row of RefreshLog
    Set wsLog = LogSheet()
    r = wsLog.Cells(wsLog.Rows.Count, "A").End(xlUp).Row + 1
    wsLog.Cells(r, "A").Value = Now
    wsLog.Cells(r, "B").Value = DataRowCount("Encounters")
    wsLog.Cells(r, "C").Value = DataRowCount("ED_Visits")
    wsLog.Cells(r, "D").Value = DataRowCount("Claims")
    wsLog.Cells(r, "E").Value = surveysAfter
    wsLog.Cells(r, "F").Value = surveysBefore - surveysAfter

    ' 5. Stamp the dashboard and bring it to the front
    With ThisWorkbook.Worksheets("Dashboard")
        .Range("H2").Value = Now
        .Activate
    End With

    ' Optional: save a PDF of the dashboard next to the workbook (the workbook must be saved first)
    ' ThisWorkbook.Worksheets("Dashboard").ExportAsFixedFormat Type:=xlTypePDF, _
    '     Filename:=ThisWorkbook.Path & Application.PathSeparator & "Dashboard_" & Format$(Now, "yyyy-mm-dd") & ".pdf"

Done:
    Application.Calculation = oldCalc
    Application.ScreenUpdating = True
    Exit Sub

Fail:
    MsgBox "RefreshReview stopped: " & Err.Description & " (error " & Err.Number & ")", vbExclamation, "Refresh"
    Resume Done
End Sub
"""

VBA_STARTER_SUB = """
Public Sub RefreshReview()
    Dim oldCalc As XlCalculation
    Dim surveysBefore As Long, surveysAfter As Long
    Dim wsLog As Worksheet
    Dim r As Long

    oldCalc = Application.Calculation
    On Error GoTo Fail
    Application.ScreenUpdating = False

    ' STEP 1 (your code): refresh every query and PivotTable in this workbook (one line).
    '         Then wait for background queries with Application.CalculateUntilAsyncQueriesDone
    '         (wrap that line in On Error Resume Next / On Error GoTo Fail so a failed wait
    '         doesn't stop the macro).

    ' STEP 2 (your code): count the Surveys data rows with DataRowCount("Surveys") into surveysBefore,
    '         remove duplicate rows (SurveyID, column 1, is the key; the sheet has a header row),
    '         then count again into surveysAfter.
    '         Hint: Range("A1").CurrentRegion.RemoveDuplicates Columns:=..., Header:=...

    ' STEP 3 (your code): switch calculation to automatic and recalculate everything.

    ' STEP 4 (your code): Set wsLog = LogSheet(), find the next empty row r in column A, and write:
    '         A = Now, B = Encounters rows, C = ED_Visits rows, D = Claims rows,
    '         E = surveysAfter, F = surveysBefore - surveysAfter

    ' STEP 5 (your code): write Now into Dashboard!H2 and activate the Dashboard sheet.

Done:
    Application.Calculation = oldCalc
    Application.ScreenUpdating = True
    Exit Sub

Fail:
    MsgBox "RefreshReview stopped: " & Err.Description & " (error " & Err.Number & ")", vbExclamation, "Refresh"
    Resume Done
End Sub
"""

# The full Sub is too tall for one Answer Key cell (Excel caps a row at 409 points), so the key shows only the lines the
# learner writes for STEP 1-5. The README answer key and solutions/ hold the complete module.
VBA_KEY_STEPS = """' STEP 1
ThisWorkbook.RefreshAll
On Error Resume Next
Application.CalculateUntilAsyncQueriesDone
On Error GoTo Fail
' STEP 2
surveysBefore = DataRowCount("Surveys")
With ThisWorkbook.Worksheets("Surveys")
  .Range("A1").CurrentRegion.RemoveDuplicates _
    Columns:=1, Header:=xlYes
End With
surveysAfter = DataRowCount("Surveys")
' STEP 3
Application.Calculation = xlCalculationAutomatic
Application.CalculateFull
' STEP 4
Set wsLog = LogSheet()
With wsLog
  r = .Cells(.Rows.Count, "A").End(xlUp).Row + 1
  .Cells(r, "A").Value = Now
  .Cells(r, "B").Value = DataRowCount("Encounters")
  .Cells(r, "C").Value = DataRowCount("ED_Visits")
  .Cells(r, "D").Value = DataRowCount("Claims")
  .Cells(r, "E").Value = surveysAfter
  .Cells(r, "F").Value = surveysBefore - surveysAfter
End With
' STEP 5
With ThisWorkbook.Worksheets("Dashboard")
  .Range("H2").Value = Now
  .Activate
End With
' Full module: solutions/RefreshReview_Solution.bas"""

VBA_SOLUTION = VBA_HEADER.format(kind="SOLUTION (spoiler): try starter/RefreshReview_Starter.bas first.") \
    + VBA_SOLUTION_SUB + VBA_HELPERS
VBA_STARTER = VBA_HEADER.format(kind="STARTER: write the code for STEP 1 to STEP 5 in RefreshReview.") \
    + VBA_STARTER_SUB + VBA_HELPERS


# ------------------------------------------------------------------------------------------------------- definitions
# Shown on the Brief sheet (and mirrored in the README's metric table).
DEFINITIONS = [
    ("Throughput", "ALOS (average length of stay)", "Average LOSDays of Inpatient encounters discharged in 2025. "
     "LOSDays = DischargeDateTime − AdmitDateTime, in days with decimals."),
    ("Throughput", "Median door-to-provider (min)", "Median of (ProviderSeenDateTime − ArrivalDateTime) × 1440 over 2025 ED "
     "arrivals that have a ProviderSeenDateTime."),
    ("Throughput", "LWBS %", "ED arrivals with EDDisposition = LWBS (left without being seen) ÷ all ED arrivals."),
    ("Quality", "O/E LOS index", "Total LOSDays ÷ total ExpectedLOS over Inpatient stays that have an ExpectedLOS benchmark "
     "(from Diagnoses, matched on PrimaryDxCode). 1.00 means exactly at benchmark, and above 1.00 means longer than expected."),
    ("Quality", "30-day readmission rate", "Index stays = Inpatient encounters discharged Jan 1 – Nov 30, 2025 whose "
     "DischargeDisposition is not Expired. Rate = index stays with Readmit30 = Y ÷ index stays. (December discharges are "
     "left out because their 30-day window runs past the end of the data.)"),
    ("Experience", "HCAHPS top-box %", "Surveys with OverallRating 9 or 10 ÷ all surveys, after removing duplicate rows."),
    ("Utilization", "Occupancy %", "Total MidnightCensus ÷ total StaffedBeds over the units and days in scope. Inpatient "
     "occupancy uses Inpatient and Critical Care units (not the Observation unit). ICU occupancy uses Critical Care units."),
    ("Finance", "Denial rate", "Claims with ClaimStatus Denied or Appealed ÷ adjudicated claims (every claim except Pending)."),
    ("Finance", "Net collection rate", "Total PaidAmount ÷ total AllowedAmount over adjudicated claims (not Pending)."),
]

# Dashboard rows: (domain, label, kind, target, goal)  kind drives formulas on the reference sheet
DASH_ROWS = [
    ("Throughput", "ED visits (2025 arrivals)", "ed", None, None),
    ("Throughput", "Median door-to-provider (min)", "d2p", 30, "≤"),
    ("Throughput", "LWBS %", "lwbs", 0.02, "≤"),
    ("Throughput", "Inpatient discharges", "ipn", None, None),
    ("Throughput", "ALOS (days)", "alos", None, None),
    ("Quality", "O/E LOS index", "oe", 1.0, "≤"),
    ("Quality", "30-day readmission rate", "readmit", 0.15, "≤"),
    ("Experience", "HCAHPS top-box %", "topbox", 0.50, "≥"),
    ("Utilization", "Inpatient occupancy %", "occ", 0.85, "≤"),
    ("Utilization", "ICU occupancy %", "icu", 0.85, "≤"),
    ("Finance", "Denial rate (system-wide)", "denial", 0.10, "≤"),
    ("Finance", "Net collection rate (system-wide)", "ncr", 0.90, "≥"),
]
DASH_FIRST = 8                                   # first KPI row on the Dashboard sheets
DASH_FMT = {"ed": "#,##0", "d2p": "0.0", "lwbs": "0.0%", "ipn": "#,##0", "alos": "0.00", "oe": "0.00",
            "readmit": "0.0%", "topbox": "0.0%", "occ": "0.0%", "icu": "0.0%", "denial": "0.0%", "ncr": "0.0%"}
LWBS_CELL = f"D{DASH_FIRST + [k for _, _, k, _, _ in DASH_ROWS].index('lwbs')}"


# ------------------------------------------------------------------------------------------------------------- build
def build() -> Lesson:
    L = Lesson(
        code=CODE, module_dir="06-capstone", slug="01-hospital-performance-review",
        title="Capstone: Hospital Performance Review", level="Expert", minutes=180,
        objectives=[
            "Plan an analysis from business questions to deliverables",
            "Prepare multi-table data (cleaning, joins, calculated fields)",
            "Analyze throughput, quality, utilization, and finance KPIs",
            "Deliver a dashboard, an automated refresh macro, and an executive summary",
        ],
        data_note="Bluestone Health System, 2025: 11,196 encounters discharged in 2025 and their 11,196 claims, "
                  "6,217 ED arrivals, 987 patient-experience survey rows (including a duplicated batch), 6,570 unit-days "
                  "of census, and the Facilities, Departments, Diagnoses, and Payers lookup tables.",
    )

    # ------------------------------------------------------------------ source data (2025 slices)
    dx = data.index(data.load("diagnoses"), "DxCode")
    dep = data.index(data.load("departments"), "DeptID")
    fac = data.index(data.load("facilities"), "FacilityID")

    enc = sorted((e for e in data.load("encounters") if e["DischargeDateTime"].year == 2025),
                 key=lambda e: e["EncounterID"])
    claims = sorted((c for c in data.load("claims") if c["ServiceDate"].year == 2025), key=lambda c: c["ClaimID"])
    ed = sorted((v for v in data.load("ed_visits") if v["ArrivalDateTime"].year == 2025), key=lambda v: v["EDVisitID"])
    surveys = sorted((s for s in data.load("patient_satisfaction") if s["DischargeDate"].year == 2025),
                     key=lambda s: s["SurveyID"])
    dup_block = [dict(s) for s in surveys
                 if s["SurveyReceivedDate"].year == 2025 and s["SurveyReceivedDate"].month == DUP_MONTH]
    survey_rows = surveys + dup_block                       # the re-sent October file is appended at the end
    census = sorted((c for c in data.load("daily_census") if c["CensusDate"].year == 2025),
                    key=lambda c: (c["CensusDate"], c["DeptID"]))
    assert len(enc) == len(claims)
    enc_by_id = {e["EncounterID"]: e for e in enc}
    assert all(enc_by_id[c["EncounterID"]]["DischargeDateTime"].date() == c["ServiceDate"] for c in claims)

    # ------------------------------------------------------------------ sheets
    es = L.add_table_sheet(
        "Encounters", enc, table="tblEncounters",
        columns=["EncounterID", "PatientID", "EncounterType", "FacilityID", "DeptID", "AdmitDateTime",
                 "DischargeDateTime", "PrimaryDxCode", "DischargeDisposition", "PayerID", "TotalCharges", "Readmit30"],
        extra_cols=["LOSDays", "ExpectedLOS", "ServiceLine"],
        formats={"LOSDays": "0.00", "ExpectedLOS": "0.0"},
        widths={"AdmitDateTime": 17, "DischargeDateTime": 17, "DischargeDisposition": 26, "LOSDays": 10,
                "ExpectedLOS": 12, "ServiceLine": 18})
    vs = L.add_table_sheet(
        "ED_Visits", ed, table="tblED",
        columns=["EDVisitID", "EncounterID", "FacilityID", "ArrivalDateTime", "ProviderSeenDateTime",
                 "DepartureDateTime", "ESILevel", "ArrivalMode", "EDDisposition"],
        extra_cols=["DoorToProviderMin"], formats={"DoorToProviderMin": "0"},
        widths={"ArrivalDateTime": 17, "ProviderSeenDateTime": 20, "DepartureDateTime": 17, "DoorToProviderMin": 19})
    cs = L.add_table_sheet(
        "Claims", claims, table="tblClaims",
        columns=["ClaimID", "EncounterID", "PayerID", "ServiceDate", "BilledAmount", "AllowedAmount",
                 "PatientResponsibility", "PaidAmount", "ClaimStatus", "DenialReason"],
        widths={"PatientResponsibility": 21, "DenialReason": 24})
    ss = L.add_table_sheet(
        "Surveys", survey_rows, table="tblSurveys",
        columns=["SurveyID", "EncounterID", "FacilityID", "DeptID", "DischargeDate", "SurveyReceivedDate",
                 "NurseCommunication", "DoctorCommunication", "StaffResponsiveness", "Cleanliness", "Quietness",
                 "OverallRating", "WouldRecommend"],
        widths={"SurveyReceivedDate": 19, "WouldRecommend": 16})
    us = L.add_table_sheet(
        "Census", census, table="tblCensus",
        columns=["CensusDate", "FacilityID", "DeptID", "StaffedBeds", "Admissions", "Discharges", "MidnightCensus"],
        extra_cols=["UnitType"], widths={"UnitType": 14})
    fs = L.add_table_sheet("Facilities", sorted(fac.values(), key=lambda f: f["FacilityID"]), table="tblFacilities",
                           columns=["FacilityID", "FacilityName", "City", "FacilityType", "LicensedBeds"])
    ds = L.add_table_sheet("Departments", sorted(dep.values(), key=lambda d: d["DeptID"]), table="tblDepartments",
                           columns=["DeptID", "DeptName", "FacilityID", "ServiceLine", "UnitType", "StaffedBeds"])
    xs = L.add_table_sheet("Diagnoses", sorted(dx.values(), key=lambda d: d["DxCode"]), table="tblDiagnoses",
                           columns=["DxCode", "DxDescription", "DxCategory", "ExpectedLOS"],
                           formats={"ExpectedLOS": "0.0"}, widths={"DxDescription": 60})
    L.add_table_sheet("Payers", sorted(data.load("payers"), key=lambda p: p["PayerID"]), table="tblPayers",
                      columns=["PayerID", "PayerName", "PayerType"])

    def R(sd, col):
        """Plain A1 range for a column, e.g. Encounters!D2:D11197 (what a learner selects). Sheet names here need no quotes."""
        c = sd.col(col)
        return f"{sd.name}!{c}{sd.first_row}:{c}{sd.last_row}"

    def A(sd, col):
        c = sd.col(col)
        return f"{sd.name}!${c}${sd.first_row}:${c}${sd.last_row}"

    # Encounters columns
    eC, eD, eE, eF, eG, eH, eI, eL = (R(es, c) for c in ("EncounterType", "FacilityID", "DeptID", "AdmitDateTime",
                                                         "DischargeDateTime", "PrimaryDxCode", "DischargeDisposition",
                                                         "Readmit30"))
    eM, eN, eO = R(es, "LOSDays"), R(es, "ExpectedLOS"), R(es, "ServiceLine")
    dxA, dxB, dxD = R(xs, "DxCode"), R(xs, "DxDescription"), R(xs, "ExpectedLOS")
    dpA, dpD, dpE = R(ds, "DeptID"), R(ds, "ServiceLine"), R(ds, "UnitType")
    edA, edC, edD, edE, edI = (R(vs, c) for c in ("EDVisitID", "FacilityID", "ArrivalDateTime", "ProviderSeenDateTime",
                                                  "EDDisposition"))
    clA, clF, clH, clI, clJ = (R(cs, c) for c in ("ClaimID", "AllowedAmount", "PaidAmount", "ClaimStatus", "DenialReason"))
    svA, svL = R(ss, "SurveyID"), R(ss, "OverallRating")
    svAll = f"Surveys!A{ss.first_row}:{ss.col('WouldRecommend')}{ss.last_row}"
    cnC, cnD, cnG = R(us, "DeptID"), R(us, "StaffedBeds"), R(us, "MidnightCensus")
    first = es.first_row
    mcol, ncol = es.col("LOSDays"), es.col("ExpectedLOS")

    # ------------------------------------------------------------------ answers (computed in Python)
    def los(e):
        return (e["DischargeDateTime"] - e["AdmitDateTime"]).total_seconds() / 86400

    def expected(e):
        return dx[e["PrimaryDxCode"]]["ExpectedLOS"]

    inpatient = [e for e in enc if e["EncounterType"] == "Inpatient"]
    assert all(expected(e) for e in inpatient), "every 2025 inpatient diagnosis has a benchmark"

    def is_index(e):
        return (e["EncounterType"] == "Inpatient" and e["DischargeDateTime"] < READMIT_CUTOFF
                and e["DischargeDisposition"] != "Expired")

    def readmit_rate(rows):
        rows = [e for e in rows if is_index(e)]
        return sum(1 for e in rows if e["Readmit30"] == "Y") / len(rows)

    def oe(rows):
        rows = [e for e in rows if e["EncounterType"] == "Inpatient" and expected(e)]
        return sum(map(los, rows)) / sum(map(expected, rows))

    def d2p_minutes(rows):
        return [round((v["ProviderSeenDateTime"] - v["ArrivalDateTime"]).total_seconds() / 60)
                for v in rows if v["ProviderSeenDateTime"] is not None]

    def num(x):
        return int(x) if float(x).is_integer() else float(x)

    n_unique_surveys = len({s["SurveyID"] for s in survey_rows})
    assert n_unique_surveys == len(surveys)
    n_dups = len(survey_rows) - n_unique_surveys
    alos = sum(map(los, inpatient)) / len(inpatient)
    d2p_f03 = num(statistics.median(d2p_minutes([v for v in ed if v["FacilityID"] == "F03"])))
    icu_units = sorted(d for d, r in dep.items() if r["UnitType"] == "Critical Care")
    icu_rows = [c for c in census if c["DeptID"] in icu_units]
    icu_occ = sum(c["MidnightCensus"] for c in icu_rows) / sum(c["StaffedBeds"] for c in icu_rows)
    icu_avg_daily = statistics.mean(c["MidnightCensus"] / c["StaffedBeds"] for c in icu_rows)
    topbox = sum(1 for s in surveys if s["OverallRating"] >= TOPBOX_MIN) / len(surveys)
    topbox_with_dups = sum(1 for s in survey_rows if s["OverallRating"] >= TOPBOX_MIN) / len(survey_rows)
    readmit_f01 = readmit_rate([e for e in enc if e["FacilityID"] == "F01"])
    oe_f03 = oe([e for e in enc if e["FacilityID"] == "F03"])
    oe_f01 = oe([e for e in enc if e["FacilityID"] == "F01"])

    sl_n, sl_y = Counter(), Counter()
    for e in enc:
        if is_index(e):
            s = dep[e["DeptID"]]["ServiceLine"]
            sl_n[s] += 1
            sl_y[s] += e["Readmit30"] == "Y"
    sl_rates = sorted(((sl_y[s] / sl_n[s], s) for s in sl_n), reverse=True)
    top_sl, second_sl = sl_rates[0][1], sl_rates[1][1]

    def sl_ranking(keep):
        """Service-line readmission rates, highest first, for an alternative (wrong) population."""
        n, y = Counter(), Counter()
        for e in enc:
            if keep(e):
                s = dep[e["DeptID"]]["ServiceLine"]
                n[s] += 1
                y[s] += e["Readmit30"] == "Y"
        return sorted(((y[s] / n[s], s) for s in n), reverse=True)

    # Leaving Expired stays in flips the ranking; adding December discharges does not (it only narrows the lead).
    sl_with_expired = sl_ranking(lambda e: e["EncounterType"] == "Inpatient" and e["DischargeDateTime"] < READMIT_CUTOFF)
    sl_with_dec = sl_ranking(lambda e: e["EncounterType"] == "Inpatient" and e["DischargeDisposition"] != "Expired")
    assert sl_with_expired[0][1] == second_sl and sl_with_dec[0][1] == top_sl

    adjudicated = [c for c in claims if c["ClaimStatus"] != "Pending"]
    denied = [c for c in adjudicated if c["ClaimStatus"] in ("Denied", "Appealed")]
    denial_rate = len(denied) / len(adjudicated)
    denial_rate_denied_only = sum(1 for c in adjudicated if c["ClaimStatus"] == "Denied") / len(adjudicated)
    ncr = sum(c["PaidAmount"] for c in adjudicated) / sum(c["AllowedAmount"] for c in adjudicated)
    ncr_all = sum(c["PaidAmount"] for c in claims) / sum(c["AllowedAmount"] for c in claims)
    reasons = Counter(c["DenialReason"] for c in denied)
    (top_reason, top_reason_n), (second_reason, _) = reasons.most_common(2)
    any_reason = Counter(c["DenialReason"] for c in claims if c["DenialReason"])
    top_reason_share = top_reason_n / len(denied)
    sentence = f"Denial rate {_pct(denial_rate)}; top reason {top_reason} ({_pct(top_reason_share)} of denials)"

    ed_f02 = [v for v in ed if v["FacilityID"] == "F02"]
    lwbs_f02 = sum(1 for v in ed_f02 if v["EDDisposition"] == "LWBS") / len(ed_f02)
    log_total = len(enc) + len(ed) + len(claims) + n_unique_surveys

    # sanity: the distinctions the checks rely on
    assert abs(oe_f03 - oe_f01) > 0.0011 and abs(topbox - topbox_with_dups) > 0.001
    assert any_reason.most_common(1)[0][0] != top_reason      # counting partial-payment reasons changes the answer

    # bonus: opportunity sizing at the hospital with the highest O/E
    hosp = ["F01", "F02", "F03"]
    oe_by = {f: oe([e for e in enc if e["FacilityID"] == f]) for f in hosp}
    worst = max(hosp, key=oe_by.get)
    worst_ip = [e for e in inpatient if e["FacilityID"] == worst]
    net_gap = sum(map(los, worst_ip)) - sum(map(expected, worst_ip))
    excess = sum(max(0.0, los(e) - expected(e)) for e in worst_ip)
    beds = excess / 365
    by_dx = defaultdict(float)
    for e in worst_ip:
        by_dx[e["PrimaryDxCode"]] += max(0.0, los(e) - expected(e))
    top_dx, top_dx_days = max(by_dx.items(), key=lambda kv: kv[1])
    second_dx_days = sorted(by_dx.values())[-2]
    system_excess = sum(max(0.0, los(e) - expected(e)) for e in inpatient)

    # numbers quoted in explanations
    readmit_sys = readmit_rate(enc)
    live_ip = [e for e in inpatient if e["DischargeDisposition"] != "Expired"]
    readmit_incl_dec = sum(1 for e in live_ip if e["Readmit30"] == "Y") / len(live_ip)
    dec_ip = [e for e in inpatient if e["DischargeDateTime"] >= READMIT_CUTOFF and e["DischargeDisposition"] != "Expired"]
    dec_rate = sum(1 for e in dec_ip if e["Readmit30"] == "Y") / len(dec_ip)

    # ------------------------------------------------------------------ practice
    L.practice_intro = (
        "The tasks follow the work plan in the lesson guide: 1–2 prepare the data, 3–10 compute the KPIs, 11–13 build the "
        "deliverables. Every metric definition is on the Brief sheet. Data rows: Encounters and Claims 2–11197, "
        f"ED_Visits 2–{vs.last_row}, Surveys 2–{ss.last_row}, Census 2–{us.last_row}.")

    mfill = f"={es.col('DischargeDateTime')}{first}-{es.col('AdmitDateTime')}{first}"
    readmit_crit = (f'{eC},"Inpatient",{eD},"F01",{eG},"<"&DATE(2025,12,1),{eI},"<>Expired"')
    oe_live = (f'=(SUMIFS({eG},{eD},"F03",{eC},"Inpatient")-SUMIFS({eF},{eD},"F03",{eC},"Inpatient"))'
               f'/SUMPRODUCT(COUNTIFS({eH},{dxA},{eD},"F03",{eC},"Inpatient"),{dxD})')
    idx_crit = f'{eC},"Inpatient",{eG},"<"&DATE(2025,12,1),{eI},"<>Expired"'
    sl_live = (f"=LET(dept,{dpA},line,{dpD},"
               f"stays,COUNTIFS({eE},dept,{idx_crit}),"
               f'readmits,COUNTIFS({eE},dept,{idx_crit},{eL},"Y"),'
               f"lines,UNIQUE(line),grid,--(lines=TRANSPOSE(line)),"
               f"rt,IFERROR(MMULT(grid,readmits)/MMULT(grid,stays),-1),"
               f"INDEX(lines,MATCH(MAX(rt),rt,0)))")
    sentence_formula = (
        f'=LET(status,{clI},reason,{clJ},'
        f'isDen,(status="Denied")+(status="Appealed"),'
        f'reasons,UNIQUE(FILTER(reason,isDen)),'
        f'counts,COUNTIFS(reason,reasons,status,"Denied")+COUNTIFS(reason,reasons,status,"Appealed"),'
        f'top,INDEX(reasons,MATCH(MAX(counts),counts,0)),'
        f'"Denial rate "&TEXT(SUM(isDen)/SUM(--(status<>"Pending")),"0.0%")'
        f'&"; top reason "&top&" ("&TEXT(MAX(counts)/SUM(isDen),"0.0%")&" of denials)")')
    dash_c5 = "=XLOOKUP(C4,Facilities!B2:B5,Facilities!A2:A5)"
    lwbs_card = (f'=COUNTIFS({A(vs, "FacilityID")},$C$5,{A(vs, "EDDisposition")},"LWBS")'
                 f'/COUNTIF({A(vs, "FacilityID")},$C$5)')

    def ind(rng):
        return f'INDIRECT("\'RefreshLog\'!{rng}")'

    L.tasks = [
        # ---------------------------------------------------------------- 1 · prep: duplicates
        Task("Data prep · cleaning. The survey vendor re-sent one month's file, so some rows on the Surveys sheet are "
             "exact duplicates of earlier rows. Remove the duplicates so each SurveyID appears once, then enter how many "
             "survey rows remain.",
             answer=n_unique_surveys, title="Survey rows left after removing duplicates",
             solution=f"=COUNTA({svA})",
             live=f'=ROWS(UNIQUE(FILTER({svA},{svA}<>"")))',
             hint="Data → Remove Duplicates, then COUNTA the SurveyID column",
             explanation=f"Click any cell in the Surveys table, choose **Data → Remove Duplicates** (or **Table Design → "
                         f"Remove Duplicates**, which is **Table → Remove Duplicates** on a Mac), leave every column ticked, and click OK. Excel reports {n_dups} duplicate "
                         f"values removed and {n_unique_surveys} unique values remaining, so `=COUNTA(Surveys!A2:A{ss.last_row})` "
                         f"now returns {n_unique_surveys}. Removing duplicates matters because the {n_dups} repeated October "
                         f"surveys would otherwise count twice in the HCAHPS score (task 5). To count unique IDs *without* "
                         f"deleting anything, `=ROWS(UNIQUE(Surveys!A2:A{ss.last_row}))` gives the same answer in "
                         "Microsoft 365."),
        # ---------------------------------------------------------------- 2 · prep: LOSDays → ALOS
        Task(f"Data prep · calculated column. In the yellow LOSDays column of Encounters "
             f"({mcol}{first}:{mcol}{es.last_row}), calculate DischargeDateTime − AdmitDateTime for every row, in days "
             "with decimals. The gray cell then shows the system ALOS: the average LOSDays of the Inpatient rows.",
             answer=alos, fmt="0.00", title="LOSDays column → system ALOS",
             solution=mfill,
             summary=f'=IF(COUNT({eM})=0,"",AVERAGEIFS({eM},{eC},"Inpatient"))',
             fill={"range": f"Encounters!{mcol}{first}:{mcol}{es.last_row}", "formula": mfill},
             live=f'=(SUMIFS({eG},{eC},"Inpatient")-SUMIFS({eF},{eC},"Inpatient"))/COUNTIF({eC},"Inpatient")',
             hint="Subtracting two date-times gives days. Keep the column formatted as a number, not a date",
             explanation=f"Excel stores a date-time as days since 1900 with the time as a fraction, so "
                         f"`={es.col('DischargeDateTime')}{first}-{es.col('AdmitDateTime')}{first}` is the stay in days "
                         "(3.36 days = 3 days, 8 hours, and 38 minutes). Because the data is an Excel Table, the formula fills "
                         "the whole column, and Excel may show it as `=[@DischargeDateTime]-[@AdmitDateTime]`. If the "
                         "results look like dates (01/03/1900 08:38), the column picked up a date format, so set it back "
                         "to Number. The gray cell's AVERAGEIFS keeps only the Inpatient rows. The key's live formula "
                         "gets the same ALOS without a helper column, because the sum of (discharge − admit) equals the "
                         "sum of discharges minus the sum of admits."),
        # ---------------------------------------------------------------- 3 · throughput: median door-to-provider
        Task("Throughput. What was the median door-to-provider time, in minutes, for 2025 ED arrivals at Cedar Ridge "
             "Medical Center (FacilityID F03)? Door-to-provider = ProviderSeenDateTime − ArrivalDateTime. Leave out "
             "visits with no ProviderSeenDateTime.",
             answer=d2p_f03, title="Median door-to-provider minutes, Cedar Ridge (F03)",
             solution=f'=MEDIAN(IF(({edC}="F03")*({edE}<>""),({edE}-{edD})*1440))',
             hint="MEDIAN(IF(…)) with × 1440, or fill the DoorToProviderMin helper column first",
             explanation="Times are fractions of a day, so × 1440 converts them to minutes. There is no MEDIANIFS, so the "
                         "IF inside MEDIAN keeps only F03 visits that have a provider time and returns FALSE for the rest, "
                         "which MEDIAN ignores. In Microsoft 365 and Excel 2021 press Enter. In Excel 2019 and earlier, "
                         "confirm it with Ctrl + Shift + Enter (Mac: ⌘ + Shift + Return). Visits with a blank ProviderSeenDateTime are the patients "
                         "who left without being seen. Left in, each one becomes a huge negative number (an empty cell "
                         "minus the arrival time), which pulls the median down and wrecks any average. A PivotTable can't help here, because its value "
                         "summaries offer Average but not Median."),
        # ---------------------------------------------------------------- 4 · utilization: ICU occupancy
        Task("Utilization. What was the combined 2025 occupancy of the three Intensive Care Units (the units whose "
             "UnitType is Critical Care on the Departments sheet)? Occupancy = total MidnightCensus ÷ total StaffedBeds "
             "over all their days. Enter it as a percentage to 1 decimal place.",
             answer=icu_occ, fmt="0.0%", title="Combined ICU occupancy, 2025",
             solution=(f'=SUM(SUMIFS({cnG},{cnC},{{"{icu_units[0]}","{icu_units[1]}","{icu_units[2]}"}}))'
                       f'/SUM(SUMIFS({cnD},{cnC},{{"{icu_units[0]}","{icu_units[1]}","{icu_units[2]}"}}))'),
             live=(f'=SUMPRODUCT(SUMIFS({cnG},{cnC},{dpA})*({dpE}="Critical Care"))'
                   f'/SUMPRODUCT(SUMIFS({cnD},{cnC},{dpA})*({dpE}="Critical Care"))'),
             hint="Find the three DeptIDs first. SUMIFS ÷ SUMIFS, or add UnitType to Census with XLOOKUP",
             explanation=f"The Critical Care units are {', '.join(icu_units)}. With an array constant, "
                         "`SUMIFS(…,{\"D130\",\"D230\",\"D330\"})` returns three sums and SUM adds them, in any Excel "
                         "version. The reusable alternative is a join: fill the yellow UnitType column on Census with "
                         f"`=XLOOKUP(C2,Departments!$A$2:$A$32,Departments!$E$2:$E$32)`, then use "
                         "`=SUMIFS(Census!G:G,Census!H:H,\"Critical Care\")/SUMIFS(Census!D:D,Census!H:H,\"Critical Care\")`. "
                         f"Averaging the daily percentages instead gives {_pct(icu_avg_daily)}, which is wrong because a "
                         "20-bed unit and an 8-bed unit would count equally. Total ÷ total weights every bed-day the same."),
        # ---------------------------------------------------------------- 5 · experience: HCAHPS top-box
        Task("Experience. Using the de-duplicated Surveys sheet from task 1, what was the system's HCAHPS top-box %: "
             "surveys with an OverallRating of 9 or 10 ÷ all surveys? Enter it as a percentage to 1 decimal place.",
             answer=topbox, fmt="0.0%", title="HCAHPS top-box %, system",
             solution=f'=COUNTIF({svL},">=9")/COUNT({svL})',
             live=(f"=LET(u,UNIQUE({svAll}),id,INDEX(u,0,1),ok,ISTEXT(id)*(id<>\"\"),"
                   f"SUM(ok*(INDEX(u,0,{ss.headers.index('OverallRating') + 1})>=9))/SUM(ok))"),
             hint="COUNTIF(…,\">=9\") ÷ COUNT(…). Do task 1 first",
             explanation=f"**Top box** means the best possible answers, 9 or 10 on the 0–10 overall rating. COUNT counts "
                         f"only numeric ratings, so the empty cells that Remove Duplicates leaves below the data "
                         f"(rows {ss.first_row + n_unique_surveys}–{ss.last_row}) don't change the denominator. With the duplicates still in, you get {_pct(topbox_with_dups)} instead of "
                         f"{_pct(topbox)}, because the October surveys are counted twice. The key's live formula de-duplicates "
                         "inside the formula: `UNIQUE` on the whole table returns each distinct row once, and `INDEX(u,0,12)` "
                         "takes its 12th column (OverallRating)."),
        # ---------------------------------------------------------------- 6 · quality: readmission rate F01
        Task("Quality. What was the 30-day readmission rate at Bluestone Memorial Hospital (F01)? Index stays are "
             "Inpatient encounters at F01 discharged January 1 – November 30, 2025 whose DischargeDisposition is not "
             "Expired. The rate is index stays with Readmit30 = Y ÷ index stays. Enter it as a percentage to 1 decimal "
             "place.",
             answer=readmit_f01, fmt="0.0%", title="30-day readmission rate, Bluestone Memorial (F01)",
             solution=f'=COUNTIFS({readmit_crit},{eL},"Y")/COUNTIFS({readmit_crit})',
             hint="COUNTIFS ÷ COUNTIFS with the same conditions. Discharge date-times include a time of day",
             explanation=f"Numerator and denominator share four conditions, and the numerator adds Readmit30 = \"Y\". "
                         f"Every row on Encounters was discharged in 2025, so only the end date needs a condition. Use "
                         f"`\"<\"&DATE(2025,12,1)` rather than `\"<=\"&DATE(2025,11,30)`, because DATE(2025,11,30) is "
                         f"midnight at the *start* of November 30, so stays discharged later that day would drop out. "
                         f"December discharges are excluded because the data ends on 12/31/2025: a patient discharged on "
                         f"December 20 can't yet show a readmission on January 10. December's own rate is only "
                         f"{_pct(dec_rate)}, so including it would pull the system rate down from {_pct(readmit_sys)} "
                         f"to {_pct(readmit_incl_dec)}. Expired stays can't be readmitted, so they leave the denominator."),
        # ---------------------------------------------------------------- 7 · quality: O/E LOS index F03
        Task(f"Quality · join. First fill the yellow ExpectedLOS column of Encounters ({ncol}{first}:{ncol}{es.last_row}) "
             "by looking up each row's PrimaryDxCode on the Diagnoses sheet. Then calculate the O/E LOS index for "
             "Cedar Ridge Medical Center (F03): total LOSDays ÷ total ExpectedLOS over its Inpatient stays. Enter it to "
             "3 decimal places.",
             answer=oe_f03, fmt="0.000", tol=0.00051, title="O/E LOS index, Cedar Ridge (F03)",
             solution=(f'=ROUND(SUMIFS({eM},{eD},"F03",{eC},"Inpatient",{eN},">0")'
                       f'/SUMIFS({eN},{eD},"F03",{eC},"Inpatient"),3)'),
             live=oe_live,
             hint="XLOOKUP into Diagnoses for the column, then SUMIFS ÷ SUMIFS. Ratio of totals, not an average of ratios",
             explanation=f"In {ncol}{first} type `=XLOOKUP(H{first},Diagnoses!$A$2:$A$52,Diagnoses!$D$2:$D$52)` (or "
                         f"`=INDEX(Diagnoses!$D$2:$D$52,MATCH(H{first},Diagnoses!$A$2:$A$52,0))` before Excel 2021) and "
                         "let the Table fill the column. The `\">0\"` condition keeps the numerator to stays that have a benchmark, as the "
                         "definition says. Every 2025 inpatient diagnosis has one, so here it changes nothing, but it protects "
                         "the index when a code without a benchmark appears. **O/E** means observed ÷ expected. Summing both sides first "
                         "weights every stay by its length. An average of each stay's own ratio would let a 0.4-day stay "
                         "with a 2-day benchmark count as much as a 20-day sepsis stay. Cedar Ridge's patients stay about "
                         f"{(oe_f03 - 1) * 100:.0f}% longer than the benchmark predicts for their diagnoses. Bluestone "
                         f"Memorial's index is {oe_f01:.3f}, so at two decimals the two hospitals look identical. The third decimal tells them "
                         "apart, but a difference that small isn't worth a headline. The key's live formula needs no "
                         "helper columns: COUNTIFS counts F03's stays for each of the 51 codes, and SUMPRODUCT multiplies "
                         "those counts by each code's ExpectedLOS."),
        # ---------------------------------------------------------------- 8 · quality: service line
        Task("Quality · join. Fill the yellow ServiceLine column of Encounters with each row's ServiceLine from the "
             "Departments sheet (match on DeptID). Then, using the task 6 index-stay definition for all three hospitals "
             "together, which service line has the highest 30-day readmission rate? Type its name.",
             answer=top_sl, title="Service line with the highest readmission rate",
             solution=(f"1. In {eO.split('!')[1].split(':')[0]}, type `=XLOOKUP(E{first},Departments!$A$2:$A$32,"
                       "Departments!$D$2:$D$32)` and let it fill the column.\n"
                       "2. On a blank sheet, list the service lines in A2:A13 (copy Departments!D2:D32 and use "
                       "**Data → Remove Duplicates**, or type `=UNIQUE(Departments!D2:D32)` in A2).\n"
                       f"3. In B2, count index stays: `=COUNTIFS({eO},A2,{idx_crit})` and copy down.\n"
                       f"4. In C2, count readmissions: `=COUNTIFS({eO},A2,{idx_crit},{eL},\"Y\")` and copy down.\n"
                       "5. In D2, `=IF(B2=0,\"\",C2/B2)` and copy down. Sort by D (largest first), or use "
                       "`=INDEX(A2:A13,MATCH(MAX(D2:D13),D2:D13,0))`."),
             live=sl_live,
             hint="A small summary table (COUNTIFS ÷ COUNTIFS per service line) or a PivotTable on a helper flag column",
             explanation=f"{top_sl} ({_pct(sl_rates[0][0])} of {sl_n[top_sl]} index stays) edges out {second_sl} "
                         f"({_pct(sl_rates[1][0])} of {sl_n[second_sl]}). If you leave Expired stays in the population, "
                         f"{second_sl} comes out on top instead ({_pct(sl_with_expired[0][0])} against "
                         f"{_pct(dict((s, r) for r, s in sl_with_expired)[top_sl])}), which shows why the definition must "
                         "be fixed before anyone ranks anything. Adding December discharges keeps "
                         f"{top_sl} first but shrinks its lead to {_pct(sl_with_dec[0][0])} against "
                         f"{_pct(dict((s, r) for r, s in sl_with_dec)[second_sl])}. A **PivotTable** works too: add two helper columns, "
                         "`IndexStay` (1 when the stay meets the definition) and `ReadmitFlag` (1 when Readmit30 = Y), as "
                         "guide section 6e shows. Then filter IndexStay = 1, put ServiceLine in Rows, and put Average of "
                         "ReadmitFlag in Values. In an executive summary, "
                         f"say that the top two are within half a point and that {top_sl} has fewer stays, so its rate is "
                         "less certain. Don't present it as the clear outlier."),
        # ---------------------------------------------------------------- 9 · finance: denial rate
        Task("Finance. What was the claim denial rate? Count claims whose ClaimStatus is Denied or Appealed, and divide "
             "by adjudicated claims (every claim except Pending). Enter it as a percentage to 1 decimal place.",
             answer=denial_rate, fmt="0.0%", title="Denial rate (adjudicated claims)",
             solution=f'=SUM(COUNTIF({clI},{{"Denied","Appealed"}}))/COUNTIF({clI},"<>Pending")',
             hint="COUNTIF with an array constant for the two statuses, and \"<>Pending\" for the denominator",
             explanation=f"An **Appealed** claim was denied first, and the appeal hasn't paid yet, so it counts as a "
                         f"denial. Counting only Denied gives {_pct(denial_rate_denied_only)}. **Pending** claims have no "
                         "decision yet, so they belong in neither the numerator nor the denominator. Including them "
                         "would make the rate look better simply because the newest claims haven't been decided yet. "
                         "`COUNTIF(range,{\"Denied\",\"Appealed\"})` returns two counts, and SUM adds them."),
        # ---------------------------------------------------------------- 10 · finance: net collection rate
        Task("Finance. What was the net collection rate: total PaidAmount ÷ total AllowedAmount over adjudicated claims "
             "(ClaimStatus is not Pending)? Enter it as a percentage to 1 decimal place.",
             answer=ncr, fmt="0.0%", title="Net collection rate (adjudicated claims)",
             solution=f'=SUMIFS({clH},{clI},"<>Pending")/SUMIFS({clF},{clI},"<>Pending")',
             hint="SUMIFS ÷ SUMIFS with the same \"<>Pending\" condition",
             explanation=f"**AllowedAmount** is what the contract says the hospital is owed, so Paid ÷ Allowed measures "
                         f"how much of that has arrived. Pending claims have an allowed amount but no payment yet. Leave "
                         f"them in and the rate drops to {_pct(ncr_all)}, which reflects claims nobody has worked yet "
                         "rather than lost money. 📋 In this extract PaidAmount is the cash received on the claim, mostly "
                         "from insurers. Patients' copays and deductibles (PatientResponsibility) are billed separately, "
                         "so this rate sits below the 95–99% often quoted when patient payments are included."),
        # ---------------------------------------------------------------- 11 · deliverable: dashboard card
        Task(f"Deliverable · dashboard. On the Dashboard sheet, write a formula in C5 that turns the hospital name "
             f"chosen in C4 into its FacilityID. Then make the LWBS % card ({LWBS_CELL}) a formula driven by C5: ED visits "
             "with EDDisposition = LWBS ÷ all ED visits at that facility. Choose Ashby Falls Community Hospital in C4. "
             f"The gray answer cell here shows Dashboard!{LWBS_CELL}.",
             answer=lwbs_f02, fmt="0.0%", title="Dashboard: LWBS % card for Ashby Falls (F02)",
             solution=(f"1. Dashboard!C5: `{dash_c5}`\n"
                       f"2. Dashboard!{LWBS_CELL}: `{lwbs_card}`\n"
                       "3. Pick **Ashby Falls Community Hospital** from the C4 drop-down."),
             summary=f'=IF(Dashboard!{LWBS_CELL}="","",Dashboard!{LWBS_CELL})',
             fill={"range": f"Dashboard!{LWBS_CELL}:{LWBS_CELL}", "formula": lwbs_card},
             live=f'=COUNTIFS({edC},"F02",{edI},"LWBS")/COUNTIF({edC},"F02")',
             hint="XLOOKUP the name on Facilities. Then COUNTIFS(…,$C$5,…,\"LWBS\") ÷ COUNTIF(…,$C$5)",
             explanation="A **selector-driven** formula refers to the selector cell instead of a typed ID, so one "
                         "formula serves every hospital. Lock it with `$C$5` so you can copy the card formula to the "
                         "other KPI rows. Change C4 back to Bluestone Memorial Hospital and every card should update. If "
                         "a card doesn't change, it still contains a typed ID. The hidden **Reference Dashboard** sheet "
                         "has a finished version of every card."),
        # ---------------------------------------------------------------- 12 · deliverable: refresh macro
        Task("Deliverable · automation. Import starter/RefreshReview_Starter.bas into the VBE, write the code for its five STEP comments, "
             "save the workbook as .xlsm, and run RefreshReview. It creates a RefreshLog sheet whose row 2 records the "
             "first run: B2 = Encounters rows, C2 = ED_Visits rows, D2 = Claims rows, E2 = Surveys rows after "
             "duplicates are removed. The gray cell adds B2:E2. If a test run logged wrong counts, delete the "
             "RefreshLog sheet and run the macro again.",
             answer=log_total, title="RefreshReview audit log (row counts B2:E2)",
             solution=VBA_SOLUTION_SUB.strip("\n"), solution_lang="vba",
             summary=f'=IFERROR(IF(COUNT({ind("B2:E2")})=0,"",SUM({ind("B2:E2")})),"")',
             fill={"range": "RefreshLog!B2:E2", "values": [len(enc), len(ed), len(claims), n_unique_surveys]},
             live=f'=ROWS({R(es, "EncounterID")})+ROWS({edA})+ROWS({clA})+ROWS(UNIQUE(FILTER({svA},{svA}<>"")))',
             hint="Guide section 8b lists the statements: RefreshAll, RemoveDuplicates, CalculateFull, End(xlUp)",
             explanation=f"Expected: {len(enc):,} + {len(ed):,} + {len(claims):,} + {n_unique_surveys:,} = "
                         f"{log_total:,}. The full module, with the DataRowCount and LogSheet helpers, is in "
                         "solutions/RefreshReview_Solution.bas. An **audit row** of record counts is a cheap safety net: "
                         "if next month's Claims count suddenly halves, you know the extract failed before anyone reads "
                         "a wrong dashboard. If you already removed the duplicates by hand in task 1, the macro finds "
                         "none (F2 = 0) but still logs the right counts. If E2 is too high, RemoveDuplicates didn't "
                         "run. If every count is one too low or high, check the `- 1` for the header row."),
        # ---------------------------------------------------------------- 13 · deliverable: executive summary sentence
        Task("Deliverable · executive summary. In the yellow cell, write a formula that builds this sentence, so it "
             "updates whenever the data changes: Denial rate 0.0%; top reason Xxx (0.0% of denials). The first % is the "
             "task 9 denial rate. The top reason is the DenialReason that appears most often among Denied or Appealed "
             "claims, and the second % is its share of those claims. Format both percentages with TEXT(…,\"0.0%\"). "
             "The formula may refer to helper cells, such as a small table of denials by reason.",
             answer=sentence, title="Executive-summary sentence",
             solution=sentence_formula,
             hint="TEXT(x,\"0.0%\") and &. For the top reason, a small COUNTIFS table plus INDEX/MATCH/MAX, or LET + UNIQUE",
             explanation=f"Expected text: **{sentence}**. The simplest route links to cells you already built: "
                         "`=\"Denial rate \"&TEXT(Practice!D14,\"0.0%\")&\"; top reason \"&G2&\" (\"&TEXT(H2,\"0.0%\")&\" "
                         "of denials)\"`, where G2 and H2 hold the top reason and its share from a small COUNTIFS table "
                         "(guide section 6e), and D14 is your task 9 answer. "
                         "The one-formula version above needs no helper cells. It uses **LET** to name each step: it filters the reasons of "
                         "denied claims, lists each reason once with UNIQUE, counts each with COUNTIFS, and picks the "
                         "largest. Watch the population. Counting DenialReason on *every* claim, including Partially Paid "
                         f"ones, makes {any_reason.most_common(1)[0][0]} the top reason, because partial payments carry "
                         f"reasons too. Among real denials, {top_reason} ({top_reason_n:,} claims) leads "
                         f"{second_reason}. Text built with TEXT and & refreshes with the data, so the summary never "
                         "quotes a stale number."),
    ]

    # ------------------------------------------------------------------ bonus
    wname = fac[worst]["FacilityName"]
    b_crit = f'{eD},"{worst}",{eC},"Inpatient"'
    excess_sum = (f'SUM(IF(({eD}="{worst}")*({eC}="Inpatient")*({eM}>{eN}),{eM}-{eN}))')
    raw_k = f'({eD}="{worst}")*({eC}="Inpatient")'
    raw_excess = (f"LET(k,{raw_k},stay,FILTER({eG}-{eF},k),dxc,FILTER({eH},k),"
                  f"bench,INDEX({dxD},MATCH(dxc,{dxA},0)),SUM(IF(stay>bench,stay-bench,0)))")
    L.bonus_title = "Bonus: size the length-of-stay opportunity"
    L.bonus_scenario = (
        "The CFO reads your draft and asks: \"If our least efficient hospital matched the length-of-stay benchmark, how "
        "many beds would that free up, and where should it start?\" Use the LOSDays and ExpectedLOS columns you built "
        "on Encounters (tasks 2 and 7). All parts use 2025 Inpatient stays.")
    L.bonus = [
        Task("Which hospital has the highest O/E LOS index (total LOSDays ÷ total ExpectedLOS over its Inpatient stays)? "
             "Type its FacilityName.",
             answer=wname, accept=[worst, wname.replace(" Community Hospital", "")], title="Hospital with the highest O/E",
             solution=(f'=LET(ids,Facilities!A2:A4,oe,SUMIFS({eM},{eD},ids,{eC},"Inpatient")'
                       f'/SUMIFS({eN},{eD},ids,{eC},"Inpatient"),XLOOKUP(MAX(oe),oe,Facilities!B2:B4))'),
             live=(f'=LET(ids,Facilities!A2:A4,obs,SUMIFS({eG},{eD},ids,{eC},"Inpatient")'
                   f'-SUMIFS({eF},{eD},ids,{eC},"Inpatient"),'
                   f'expd,MMULT(COUNTIFS({eD},ids,{eC},"Inpatient",{eH},TRANSPOSE({dxA})),IF({dxD}="",0,{dxD})),'
                   f'oe,obs/expd,INDEX(Facilities!B2:B4,MATCH(MAX(oe),oe,0)))'),
             hint="Repeat task 7's SUMIFS ÷ SUMIFS for F01, F02, and F03 and compare",
             explanation=f"O/E by hospital: " + ", ".join(f"{f} {oe_by[f]:.3f}" for f in hosp) + ". "
                         f"{wname} is the highest. SUMIFS with a three-cell criteria range (`ids`) returns all three "
                         "indexes at once in Microsoft 365. Three separate SUMIFS ÷ SUMIFS formulas work in any version."),
        Task(f"At that hospital, how many bed-days separate actual from benchmark? Total LOSDays − total ExpectedLOS over "
             "its Inpatient stays, to 1 decimal place. This is the gap that closes if its O/E index falls to exactly "
             "1.00.",
             answer=net_gap, fmt="0.0", tol=0.051, title="Net gap to O/E = 1.00 (bed-days)",
             solution=f'=SUMIFS({eM},{b_crit})-SUMIFS({eN},{b_crit})',
             live=(f'=SUMIFS({eG},{b_crit})-SUMIFS({eF},{b_crit})'
                   f'-SUMPRODUCT(COUNTIFS({eH},{dxA},{b_crit}),{dxD})'),
             hint="SUMIFS − SUMIFS on the two helper columns",
             explanation="O/E = 1.00 means total observed days equal total expected days, so the days saved are simply "
                         "observed minus expected. That treats every stay that finished *early* as a credit that offsets "
                         "a long one."),
        Task("Short stays can't give days back to long ones, so the improvement target is the excess days. For each of "
             "that hospital's Inpatient stays, take MAX(0, LOSDays − ExpectedLOS), and add them up. Enter the total to "
             "1 decimal place.",
             answer=excess, fmt="0.0", tol=0.051, title="Excess days: sum of MAX(0, LOS − expected)",
             solution=f"={excess_sum}",
             live=f"={raw_excess}",
             hint="An ExcessDays helper column with MAX(0, …), then SUMIFS. Or SUM(IF(…)) in one formula",
             explanation=f"The easiest route is a helper column. Type ExcessDays in P1 (the Table grows to include it), "
                         f"put `=IF(C2=\"Inpatient\",MAX(0,{mcol}2-{ncol}2),\"\")` in P2, then use "
                         f"`=SUMIFS(Encounters!P2:P11197,Encounters!D2:D11197,\"{worst}\")`. The one-cell version uses SUM(IF()): "
                         f"the IF keeps the hospital's inpatient stays that ran past their benchmark and returns their "
                         f"excess. The result ({excess:,.1f}) is larger than the net gap ({net_gap:,.1f}) because "
                         "stays that finished early no longer cancel out stays that ran long. When you present the "
                         "opportunity, say which definition you used, because the two answers differ by about a third."),
        Task("Convert the excess days into staffed beds: excess days ÷ 365, because one bed open all year provides 365 "
             "bed-days. Enter it to 2 decimal places.",
             answer=beds, fmt="0.00", title="Equivalent staffed beds",
             solution=f"={excess_sum}/365",
             live=f"={raw_excess}/365",
             hint="Divide the previous answer by 365. You can reference its cell",
             explanation=f"About {beds:.1f} beds at a {fac[worst]['LicensedBeds']}-bed hospital. That's modest for one "
                         f"site, but the same method across all three hospitals gives {system_excess:,.1f} excess days, "
                         f"or {system_excess / 365:.1f} beds, which is a whole small unit. Planners often divide by "
                         f"365 × 0.85 instead (beds at an 85% occupancy target), giving "
                         f"{system_excess / (365 * 0.85):.1f} system-wide."),
        Task("Where should the hospital start? Which primary diagnosis accounts for the most excess days (the MAX(0, …) "
             "method) at that hospital? Type its DxDescription exactly as it appears on the Diagnoses sheet.",
             answer=dx[top_dx]["DxDescription"], accept=[top_dx], title="Diagnosis with the most excess days",
             solution=("1. Add the ExcessDays helper column from the previous part (Inpatient rows only).\n"
                       "2. Insert a PivotTable from tblEncounters: **Filters** FacilityID = "
                       f"{worst} and EncounterType = Inpatient, **Rows** PrimaryDxCode, **Values** Sum of ExcessDays.\n"
                       "3. Sort the values largest to smallest, then look up the top code's DxDescription on Diagnoses."),
             live=(f"=LET(k,{raw_k},stay,FILTER({eG}-{eF},k),dxc,FILTER({eH},k),"
                   f"bench,INDEX({dxD},MATCH(dxc,{dxA},0)),extra,IF(stay>bench,stay-bench,0),"
                   f"dxu,UNIQUE(dxc),tot,MMULT(--(dxu=TRANSPOSE(dxc)),extra),"
                   f"INDEX({dxB},MATCH(INDEX(dxu,MATCH(MAX(tot),tot,0)),{dxA},0)))"),
             hint="PivotTable of the helper column by PrimaryDxCode, sorted descending",
             explanation=f"{dx[top_dx]['DxDescription']} ({top_dx}) accounts for {top_dx_days:,.1f} of the "
                         f"{excess:,.1f} excess days ({_pct(top_dx_days / excess)}), about twice the next diagnosis "
                         f"({second_dx_days:,.1f} days). Sepsis pathways (early antibiotics, daily discharge-readiness "
                         "review) are a common place to start a length-of-stay project. That's a realistic "
                         "recommendation for your executive summary."),
    ]

    # ------------------------------------------------------------------ files
    L.extra_files = {
        "starter/RefreshReview_Starter.bas": _crlf(VBA_STARTER),
        "solutions/RefreshReview_Solution.bas": _crlf(VBA_SOLUTION),
    }
    L.start_notes = [
        "This is the capstone: read the brief in the lesson guide (or the Brief sheet) before you start. Any tool from "
        "Lessons 1.1–5.5 is fair game: formulas, PivotTables, Power Query, the Data Model, or VBA.",
        "Task 12 needs macros: save your copy as an Excel Macro-Enabled Workbook (.xlsm) before you add the module.",
        "The Dashboard sheet is your deliverable template. A finished example is on the hidden 'Reference Dashboard' sheet.",
    ]
    L.sheet_order = ["Start Here", "Brief", "Practice", "Dashboard", "Bonus", "Encounters", "ED_Visits", "Claims",
                     "Surveys", "Census", "Facilities", "Departments", "Diagnoses", "Payers", "Answer Key", "Bonus Key",
                     "Reference Dashboard"]

    # ------------------------------------------------------------------ custom sheets
    ranges = dict(eC=eC, eD=eD, eE=eE, eF=eF, eG=eG, eH=eH, eI=eI, eL=eL, dxA=dxA, dxD=dxD, dpA=dpA, dpE=dpE,
                  edA=edA, edC=edC, edD=edD, edE=edE, edI=edI, clI=clI, clF=clF, clH=clH, svAll=svAll,
                  cnB=R(us, "FacilityID"), cnC=cnC, cnD=cnD, cnG=cnG)

    @L.customize
    def _custom(wb, lesson, selftest):
        _brief_sheet(wb)
        _dashboard_sheet(wb, lesson, reference=False, rg=ranges)
        _dashboard_sheet(wb, lesson, reference=True, rg=ranges)
        # The full RefreshReview Sub doesn't fit in one key cell, so show just the STEP code there.
        key_row = int(lesson.tasks[11].key_cell.rsplit("$", 1)[1])
        kc = wb["Answer Key"].cell(row=key_row, column=4)
        kc.value = VBA_KEY_STEPS
        kc.font = Font(name="Consolas", size=9)
        for name in ("Brief", "Dashboard", "Reference Dashboard"):          # print / PDF: one page wide, landscape
            ws = wb[name]
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        if selftest:
            # Simulate the learner's prep work so the sample solutions evaluate.
            sv = wb["Surveys"]
            for r in range(ss.first_row + n_unique_surveys, ss.last_row + 1):        # Remove Duplicates
                for c in range(1, len(ss.headers) + 1):
                    sv.cell(row=r, column=c).value = None
            en = wb["Encounters"]
            for r in range(es.first_row, es.last_row + 1):                             # ExpectedLOS lookup
                lesson.set_formula(en, f"{ncol}{r}",
                                   f"=XLOOKUP({es.col('PrimaryDxCode')}{r},{A(xs, 'DxCode')},{A(xs, 'ExpectedLOS')})",
                                   dynamic=False)
            db = wb["Dashboard"]
            db["C4"] = fac["F02"]["FacilityName"]
            lesson.set_formula(db, "C5", dash_c5, dynamic=False)
            log = wb.create_sheet("RefreshLog")                                         # what the macro creates
            for j, h in enumerate(["RefreshedAt", "Encounters", "ED_Visits", "Claims", "Surveys", "DuplicatesRemoved"], 1):
                log.cell(row=1, column=j, value=h)
            log["A2"] = datetime(2026, 1, 5, 7, 30)
            log["F2"] = n_dups

    # keep for the README / smoke test
    L._answers = dict(alos=alos, readmit_sys=readmit_sys, dec_rate=dec_rate, oe_by=oe_by, sl_rates=sl_rates,
                      reasons=reasons, worst=worst, net_gap=net_gap, excess=excess, system_excess=system_excess)
    return L


# ------------------------------------------------------------------------------------------------------ Brief sheet
def _brief_sheet(wb):
    ws = wb.create_sheet("Brief")
    ws.sheet_properties.tabColor = NAVY
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 100
    ws["B2"] = "Project brief · 2025 Hospital Performance Review"
    ws["B2"].font = Font(bold=True, size=16, color=NAVY)
    memo = [
        ("From", "Dr. Priya Raman, Chief Medical Officer, and Marcus Hale, Chief Financial Officer (fictional)"),
        ("To", "You, Senior Analyst, Performance Improvement"),
        ("Ask", "Before the January board retreat, give us an honest, numbers-first review of 2025 across our three "
                "hospitals: Bluestone Memorial (F01), Ashby Falls Community (F02), and Cedar Ridge (F03)."),
        ("Questions", "1. Are patients moving through our EDs and inpatient units efficiently?  2. Is our care safe and "
                      "effective: are stays the right length, and do patients come back?  3. Do patients rate us well?  "
                      "4. Are we using our beds well?  5. Are we getting paid for the care we deliver?"),
        ("Deliverables", "(a) A one-page Dashboard with a hospital selector, system totals, targets, and status.  "
                         "(b) A one-click RefreshReview macro so the review can be rerun monthly.  (c) A one-page "
                         "executive summary: a headline, 3–5 findings with numbers, 2–3 recommendations, and caveats."),
        ("Scope", "Encounters discharged in 2025 and their claims, ED arrivals in 2025, surveys for 2025 discharges, and the "
                  "2025 unit census. Data as of 12/31/2025. All data are synthetic."),
    ]
    r = 4
    for label, text in memo:
        ws.cell(row=r, column=2, value=label).font = Font(bold=True)
        c = ws.cell(row=r, column=3, value=text)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        c.alignment = WRAP
        ws.cell(row=r, column=2).alignment = WRAP
        ws.row_dimensions[r].height = 15 * max(1, -(-len(text) // 125)) + 4
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Metric definitions (use these exactly)").font = Font(bold=True, size=13, color=NAVY)
    r += 1
    for j, h in enumerate(["Domain", "Metric", "Definition"], 2):
        c = ws.cell(row=r, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.border = BOX
    r += 1
    for domain, metric, definition in DEFINITIONS:
        for j, v in enumerate([domain, metric, definition], 2):
            c = ws.cell(row=r, column=j, value=v)
            c.alignment = WRAP
            c.border = BOX
        ws.row_dimensions[r].height = 15 * max(1, -(-len(definition) // 95)) + 4
        r += 1
    r += 1
    notes = [
        "Data quality note: the survey vendor re-sent one month's file, so the Surveys sheet contains a block of exact "
        "duplicate rows. Remove them before you report anything from Surveys.",
        "Join keys: Encounters.PrimaryDxCode → Diagnoses.DxCode. Encounters.DeptID and Census.DeptID → "
        "Departments.DeptID. FacilityID → Facilities. Claims.EncounterID and ED_Visits.EncounterID → Encounters.",
        "Keep the raw data sheets as delivered (apart from the duplicate fix and the yellow helper columns). Do your "
        "analysis on new sheets so you can always trace a number back to its source.",
    ]
    for n in notes:
        c = ws.cell(row=r, column=2, value=n)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c.alignment = WRAP
        c.font = Font(italic=True, color="404040")
        ws.row_dimensions[r].height = 32
        r += 1


# -------------------------------------------------------------------------------------------------- Dashboard sheets
def _dashboard_sheet(wb, lesson, reference: bool, rg: dict):
    name = "Reference Dashboard" if reference else "Dashboard"
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = "7B2C2C" if reference else "2E75B6"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", (2, 13, 34, 18, 14, 11, 7, 16)):
        ws.column_dimensions[col].width = w
    ws["B1"] = "Bluestone Health System · 2025 Hospital Performance Review"
    ws["B1"].font = Font(bold=True, size=16, color=NAVY)
    ws["B2"] = ("REFERENCE SOLUTION (spoiler): a finished, formula-driven version of the Dashboard" if reference else
                "Discharges and ED arrivals Jan 1 – Dec 31, 2025 · data as of 12/31/2025 · synthetic data")
    ws["B2"].font = Font(italic=True, color="7B2C2C" if reference else "595959")
    ws["G2"] = "Refreshed"
    ws["G2"].alignment = Alignment(horizontal="right")
    ws["G2"].font = Font(color="595959")
    ws["H2"].number_format = "mm/dd/yyyy hh:mm"
    ws["B4"] = "Hospital"
    ws["B5"] = "FacilityID"
    for c in ("B4", "B5"):
        ws[c].font = Font(bold=True)
    ws["C4"] = "Bluestone Memorial Hospital"
    ws["C4"].fill = PatternFill("solid", fgColor="DDEBF7")
    ws["C4"].font = Font(bold=True, color=NAVY)
    ws["C4"].border = BOX
    ws["D4"] = "◀ pick a hospital from the drop-down"
    ws["D4"].font = Font(italic=True, color="7F7F7F")
    dv = DataValidation(type="list", formula1="Facilities!$B$2:$B$4", allow_blank=False,   # no "=" inside the XML
                        showErrorMessage=True, showDropDown=False)            # showDropDown=False means SHOW the arrow
    dv.error = "Pick one of the three hospitals from the list."
    dv.errorTitle = "Hospital"
    ws.add_data_validation(dv)
    dv.add("C4")
    ws["C5"].border = BOX
    if reference:
        lesson.set_formula(ws, "C5", "=XLOOKUP(C4,Facilities!B2:B5,Facilities!A2:A5)", dynamic=False)
    else:
        ws["C5"].fill = INPUT_FILL
        ws["D5"] = "◀ your formula: the FacilityID for C4"
        ws["D5"].font = Font(italic=True, color="7F7F7F")
    hdr = DASH_FIRST - 1
    for j, h in enumerate(["Domain", "KPI", "Selected hospital", "System", "Target", "Goal", "Status"], 2):
        c = ws.cell(row=hdr, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center" if j > 3 else "left")
    for i, (domain, label, kind, target, goal) in enumerate(DASH_ROWS):
        r = DASH_FIRST + i
        ws.cell(row=r, column=2, value=domain).font = Font(color="595959")
        ws.cell(row=r, column=3, value=label)
        t = ws.cell(row=r, column=6, value=target if target is not None else "—")
        g = ws.cell(row=r, column=7, value=goal if goal else "—")
        t.font = Font(color="0000FF") if target is not None else Font(color="7F7F7F")
        t.alignment = Alignment(horizontal="right" if target is not None else "center")
        g.alignment = Alignment(horizontal="center")
        if target is not None:
            t.number_format = DASH_FMT[kind]
        for col in (4, 5):
            c = ws.cell(row=r, column=col)
            c.number_format = DASH_FMT[kind]
            c.alignment = Alignment(horizontal="right")
            if not reference:
                c.fill = INPUT_FILL
        if not reference and goal:
            ws.cell(row=r, column=8).fill = INPUT_FILL
        for col in range(2, 9):
            ws.cell(row=r, column=col).border = BOX
        if reference:
            sel, sysf = _kpi_formulas(kind, "$C$5", rg)
            lesson.set_formula(ws, f"D{r}", sel)
            lesson.set_formula(ws, f"E{r}", sysf)
            if goal:
                ws.cell(row=r, column=8).value = (f'=IF(D{r}="","",IF(G{r}="≤",IF(D{r}<=F{r},"Met","Missed"),'
                                                  f'IF(D{r}>=F{r},"Met","Missed")))')
                ws.cell(row=r, column=8).alignment = Alignment(horizontal="center")
    last = DASH_FIRST + len(DASH_ROWS) - 1
    status = f"H{DASH_FIRST}:H{last}"
    ws.conditional_formatting.add(status, FormulaRule(formula=[f'H{DASH_FIRST}="Met"'],
                                                      fill=PatternFill("solid", fgColor="C6EFCE"),
                                                      font=Font(bold=True, color="006100")))
    ws.conditional_formatting.add(status, FormulaRule(formula=[f'H{DASH_FIRST}="Missed"'],
                                                      fill=PatternFill("solid", fgColor="FFC7CE"),
                                                      font=Font(bold=True, color="9C0006")))
    r = last + 2
    if not reference:
        tips = [
            "How to build it (lesson guide, section 7):",
            "1. C5: look up the FacilityID of the hospital in C4. Every 'Selected hospital' formula refers to $C$5.",
            "2. Fill D8:E19. Use the Brief sheet's definitions. Finance rows are system-wide, so D and E match there.",
            "3. Status: =IF(D9=\"\",\"\",IF(G9=\"≤\",IF(D9<=F9,\"Met\",\"Missed\"),IF(D9>=F9,\"Met\",\"Missed\"))) and copy it "
            "to the rows that have a goal. The green/red formatting is already set up.",
            "4. Add at least one chart below (for example O/E or readmission rate by hospital) and a 'Refreshed' time "
            "in H2 from your macro. Then set Page Layout → Scale to Fit → Width to 1 page and print or export to PDF.",
            "5. Delete these build notes before you print or share the dashboard.",
        ]
        for k, line in enumerate(tips):
            # Wrap each note inside B:H so it can't run off to the right and shrink the printed page.
            rr = r + k
            ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=8)
            c = ws.cell(row=rr, column=2, value=line)
            c.font = Font(bold=k == 0, italic=k > 0, color="404040")
            c.alignment = WRAP
            ws.row_dimensions[rr].height = 15 * max(1, -(-len(line) // 120)) + 2
        return

    # ---- reference: comparison table + charts
    ws.cell(row=r, column=2, value="Hospital comparison (feeds the charts)").font = Font(bold=True, color=NAVY)
    r += 1
    comp_hdr = r
    heads = ["FacilityID", "Hospital", "O/E LOS index", "Readmission rate", "LWBS %", "HCAHPS top-box %"]
    kinds = ["oe", "readmit", "lwbs", "topbox"]
    for j, h in enumerate(heads, 2):
        c = ws.cell(row=r, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[r].height = 44
    ws.column_dimensions["G"].width = 9
    for k, fid in enumerate(["F01", "F02", "F03"]):
        rr = r + 1 + k
        ws.cell(row=rr, column=2, value=fid).alignment = Alignment(horizontal="center")
        lesson.set_formula(ws, f"C{rr}", f'=XLOOKUP(B{rr},Facilities!A2:A5,Facilities!B2:B5)', dynamic=False)
        for j, kind in enumerate(kinds, 4):
            sel, _ = _kpi_formulas(kind, f"$B{rr}", rg)
            cell = ws.cell(row=rr, column=j)
            lesson.set_formula(ws, cell.coordinate, sel)
            cell.number_format = DASH_FMT[kind]
        for j in range(2, 8):
            ws.cell(row=rr, column=j).border = BOX
    comp_last = r + 3
    chart = BarChart()
    chart.type = "col"
    chart.title = "O/E LOS index by hospital (1.00 = benchmark)"
    chart.y_axis.title = "Observed ÷ expected days"
    chart.add_data(Reference(ws, min_col=4, min_row=comp_hdr, max_row=comp_last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=comp_hdr + 1, max_row=comp_last))
    chart.legend = None
    chart.y_axis.scaling.min = 0                 # bars must start at zero, or small gaps look huge
    chart.y_axis.scaling.max = 1.5
    chart.y_axis.majorUnit = 0.25
    chart.y_axis.number_format = "0.00"
    chart.x_axis.delete = False                  # openpyxl 3.1 hides axes in Excel unless told otherwise
    chart.y_axis.delete = False
    chart.height, chart.width = 7.5, 13
    ws.add_chart(chart, f"B{comp_last + 2}")

    # monthly readmission rate (index discharges Jan–Nov)
    mr = comp_last + 2
    ws.cell(row=mr - 1, column=10, value="Monthly readmission rate, system (index stays)").font = Font(bold=True, color=NAVY)
    for j, h in enumerate(["Month", "Index stays", "Readmits", "Rate", "Target"], 10):
        c = ws.cell(row=mr, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    for m in range(1, 12):
        rr = mr + m
        ws.cell(row=rr, column=10, value=date(2025, m, 1)).number_format = "mmm"
        crit = (f'{rg["eC"]},"Inpatient",{rg["eI"]},"<>Expired",{rg["eG"]},">="&J{rr},'
                f'{rg["eG"]},"<"&EDATE(J{rr},1)')
        ws.cell(row=rr, column=11, value=f"=COUNTIFS({crit})")
        ws.cell(row=rr, column=12, value=f'=COUNTIFS({crit},{rg["eL"]},"Y")')
        ws.cell(row=rr, column=13, value=f"=IF(K{rr}=0,\"\",L{rr}/K{rr})").number_format = "0.0%"
        ws.cell(row=rr, column=14, value=0.15).number_format = "0.0%"
    for col, w in zip("JKLMN", (9, 11, 10, 8, 8)):
        ws.column_dimensions[col].width = w
    line = LineChart()
    line.title = "30-day readmission rate by discharge month, 2025"
    line.y_axis.number_format = "0%"
    line.add_data(Reference(ws, min_col=13, max_col=14, min_row=mr, max_row=mr + 11), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=10, min_row=mr + 1, max_row=mr + 11))
    line.y_axis.scaling.min = 0
    line.x_axis.number_format = "mmm"
    line.x_axis.delete = False
    line.y_axis.delete = False
    for series in line.series:
        series.smooth = False
    line.height, line.width = 7.5, 13
    ws.add_chart(line, f"J{mr + 13}")
    ws.sheet_state = "hidden"


def _kpi_formulas(kind: str, idc: str, rg: dict) -> tuple[str, str]:
    """(selected-hospital formula using the FacilityID in idc, system-wide formula) for one dashboard KPI."""
    g = rg
    ip_sel = f'{g["eD"]},{idc},{g["eC"]},"Inpatient"'
    ip_sys = f'{g["eC"]},"Inpatient"'
    idx = f'{g["eC"]},"Inpatient",{g["eG"]},"<"&DATE(2025,12,1),{g["eI"]},"<>Expired"'
    adj = f'{g["clI"]},"<>Pending"'
    if kind == "ed":
        return f'=COUNTIF({g["edC"]},{idc})', f'=COUNTA({g["edA"]})'
    if kind == "d2p":
        mins = f'({g["edE"]}-{g["edD"]})*1440'
        return (f'=MEDIAN(IF(({g["edC"]}={idc})*({g["edE"]}<>""),{mins}))',
                f'=MEDIAN(IF({g["edE"]}<>"",{mins}))')
    if kind == "lwbs":
        return (f'=COUNTIFS({g["edC"]},{idc},{g["edI"]},"LWBS")/COUNTIF({g["edC"]},{idc})',
                f'=COUNTIF({g["edI"]},"LWBS")/COUNTA({g["edA"]})')
    if kind == "ipn":
        return f"=COUNTIFS({ip_sel})", f"=COUNTIFS({ip_sys})"
    if kind == "alos":
        return (f'=(SUMIFS({g["eG"]},{ip_sel})-SUMIFS({g["eF"]},{ip_sel}))/COUNTIFS({ip_sel})',
                f'=(SUMIFS({g["eG"]},{ip_sys})-SUMIFS({g["eF"]},{ip_sys}))/COUNTIFS({ip_sys})')
    if kind == "oe":
        def oe(crit):
            return (f'=SUMPRODUCT((SUMIFS({g["eG"]},{g["eH"]},{g["dxA"]},{crit})-SUMIFS({g["eF"]},{g["eH"]},{g["dxA"]},{crit}))'
                    f'*({g["dxD"]}>0))/SUMPRODUCT(COUNTIFS({g["eH"]},{g["dxA"]},{crit}),{g["dxD"]})')
        return oe(ip_sel), oe(ip_sys)
    if kind == "readmit":
        return (f'=COUNTIFS({idx},{g["eD"]},{idc},{g["eL"]},"Y")/COUNTIFS({idx},{g["eD"]},{idc})',
                f'=COUNTIFS({idx},{g["eL"]},"Y")/COUNTIFS({idx})')
    if kind == "topbox":
        u = f'UNIQUE({g["svAll"]})'
        return (f'=LET(u,{u},ok,INDEX(u,0,3)={idc},SUM(ok*(INDEX(u,0,12)>=9))/SUM(--ok))',
                f'=LET(u,{u},id,INDEX(u,0,1),ok,ISTEXT(id)*(id<>""),SUM(ok*(INDEX(u,0,12)>=9))/SUM(ok))')
    if kind in ("occ", "icu"):
        cond = f'({g["dpE"]}="Critical Care")' if kind == "icu" else \
            f'(({g["dpE"]}="Inpatient")+({g["dpE"]}="Critical Care"))'

        def occ(extra):
            return (f'=SUMPRODUCT(SUMIFS({g["cnG"]},{g["cnC"]},{g["dpA"]}{extra})*{cond})'
                    f'/SUMPRODUCT(SUMIFS({g["cnD"]},{g["cnC"]},{g["dpA"]}{extra})*{cond})')
        return occ(f',{g["cnB"]},{idc}'), occ("")
    if kind == "denial":
        f = f'=SUM(COUNTIF({g["clI"]},{{"Denied","Appealed"}}))/COUNTIF({g["clI"]},"<>Pending")'
        return f, f
    if kind == "ncr":
        f = f'=SUMIFS({g["clH"]},{adj})/SUMIFS({g["clF"]},{adj})'
        return f, f
    raise ValueError(kind)
